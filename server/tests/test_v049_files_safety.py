"""v0.49 real file types (Magika), MarkItDown reading, Indian PII gate, keep-alive steps."""
import io
import json

import httpx
import pytest

ELF = open("/bin/ls", "rb").read()[:200000]


def test_program_disguised_as_pdf_is_refused_for_owner_and_declined_from_peers(env):
    from app import chatfiles, filecheck
    if not filecheck.available():
        pytest.skip("magika not installed")
    rt = filecheck.real_type("bill.pdf", ELF)
    assert rt["danger"] and "program" in rt["note"]
    with pytest.raises(chatfiles.FileError):
        chatfiles.store("default", "bill.pdf", ELF)
    ok = chatfiles.store("default", "notes.txt", b"#!/bin/bash\necho hello\n" * 5)     # a script in a .txt is just text
    assert ok["file_id"]
    from tests.test_net import _connect, ev
    from tests.test_net_share import _file
    net = _connect()
    r = net.send("default", "Kunal", "file", {"file_id": _file("default", "invoice.pdf", ELF)})
    assert r["status"] == "declined"
    assert "khatarnak" in ev("kunal", "hoola_msg")[-1]["text"]


def test_png_named_jpg_is_only_noted(env):
    from PIL import Image
    from app import filecheck
    if not filecheck.available():
        pytest.skip("magika not installed")
    b = io.BytesIO()
    Image.new("RGB", (40, 40), (1, 2, 3)).save(b, "PNG")
    rt = filecheck.real_type("photo.jpg", b.getvalue())
    assert not rt["danger"] and rt["label"] == "png"


def test_markitdown_reads_word_and_excel_tables(env):
    from app.docread import read_text
    from docx import Document
    d = Document()
    d.add_heading("FNW Rates", 1)
    t = d.add_table(rows=2, cols=2)
    t.cell(0, 0).text, t.cell(0, 1).text = "Item", "Price"
    t.cell(1, 0).text, t.cell(1, 1).text = "ONT", "1200"
    b = io.BytesIO()
    d.save(b)
    out = read_text("rates.docx", "", b.getvalue())
    assert "ONT" in out["text"] and "1200" in out["text"]
    import openpyxl
    wb = openpyxl.Workbook()
    wb.active.append(["Name", "Mobile"])
    wb.active.append(["Kunal", "9876543210"])
    b = io.BytesIO()
    wb.save(b)
    out = read_text("list.xlsx", "", b.getvalue())
    assert "Kunal" in out["text"] and "|" in out["text"]


def test_pii_gate_in_memory(env):
    import asyncio
    from app import pii
    from app.agent.life_tools import remember
    from app.agent.tools import Ctx, ToolError
    assert pii.find("Mera Aadhaar 2341 2341 2346") == ["aadhaar"]
    assert "pan" in pii.find("PAN ABCPE1234F hai")
    assert "bank_account" in pii.find("SBI account no 123456789012")
    assert pii.find("Mobile 9876543210, pincode 444101") == []
    ctx = Ctx(tenant_id="default", conversation_id="c")
    with pytest.raises(ToolError):
        asyncio.run(remember(ctx, "2341 2341 2346 yaad rakho"))
    assert asyncio.run(remember(ctx, "Beti ka birthday 12 March"))["saved"]


def test_keepalive_steps_cached_and_credited(env, tmp_path, monkeypatch):
    from app import keepalive
    monkeypatch.setattr(keepalive, "CACHE", tmp_path)
    calls = []

    def handler(req):
        calls.append(str(req.url))
        return httpx.Response(200, json={"name": "Xiaomi", "award": 5, "user_solution": "<p>Settings &gt; Apps</p><ul><li>Autostart ON</li></ul>",
                                         "explanation": "<p>MIUI kills apps.</p>"})
    monkeypatch.setattr(keepalive, "_transport", httpx.MockTransport(handler))
    s = keepalive.steps("Redmi")
    assert s["vendor"] == "xiaomi" and "Settings > Apps" in s["steps"] and any("Autostart" in x for x in s["steps"])
    assert "CC BY" in s["credit"] and calls[0].endswith("/api/v2/xiaomi.json")
    keepalive.steps("xiaomi")
    assert len(calls) == 1                     # cached
    assert keepalive.vendor_for("iQOO") == "vivo"
