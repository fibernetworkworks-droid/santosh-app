"""Contracts Sahayak (contract manager): clause reading from a real Word file, risk review against HER playbook, redline as a NEW
Word copy with track changes (original untouched), renewal / obligation register + reminders, cards that ONLY her tap decides
(send / e-sign / e-stamp with fingerprint / notice), stamp duty only from her rates, DPDP forget, tools and API."""
import asyncio
import hashlib
import io
from datetime import date, datetime, timezone
from pathlib import Path

import pytest
from fastapi import FastAPI
from fastapi.testclient import TestClient

import app.contracts.models  # noqa: F401  (tables for the env fixture's init_db)
from tests.test_life import conv_ctx

T = "default"
DAY = date(2026, 10, 23)

SAMPLE = [
    "SERVICE AGREEMENT",
    "This Agreement is made on 1st April 2026 by and between Ingle Tech Pvt Ltd, a company (hereinafter referred to as the \"Vendor\") "
    "and Acme Retail LLP, a limited liability partnership (hereinafter referred to as the \"Client\").",
    "1. Term. This Agreement shall remain in force for a period of one (1) year from the Effective Date and shall automatically renew "
    "for successive periods of one (1) year unless either Party gives ninety (90) days prior written notice.",
    "2. Payment. The Client shall pay each invoice within ninety (90) days of receipt of invoice. The fees are Rs 12,00,000 per year.",
    "3. Indemnity. The Vendor shall indemnify and hold harmless the Client against all losses.",
    "4. Liability. The Vendor shall have unlimited liability under this Agreement.",
    "5. Non-compete. The Vendor shall not engage in any similar business with competitors of the Client during the term and for two "
    "years after termination.",
    "6. Termination. The Client may terminate this Agreement for convenience by giving thirty (30) days notice.",
    "7. Governing law. This Agreement shall be governed by the laws of Singapore. The seat of arbitration shall be Singapore.",
]


@pytest.fixture
def ct(env, monkeypatch):
    from app.contracts import service as S
    events = []
    monkeypatch.setattr(S, "_publish", lambda tenant, kind, payload, conv="": events.append((kind, payload)))
    monkeypatch.setattr(S, "today", lambda t: DAY)
    monkeypatch.setattr(S, "_ask_json", lambda system, prompt: {})
    S.EVENTS = events
    return S


def kinds(S):
    return [k for k, _ in S.EVENTS]


def _docx(paras) -> bytes:
    from docx import Document
    d = Document()
    for p in paras:
        r = d.add_paragraph().add_run(p)
        r.bold = p.isupper()
    bio = io.BytesIO()
    d.save(bio)
    return bio.getvalue()


def _file(name: str, data: bytes) -> str:
    from app import chatfiles
    return chatfiles.store(T, name, data, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")["file_id"]


def base(S):
    S.save_settings(T, party_name="Ingle Tech Pvt Ltd", default_role="vendor", seat_city="Pune", state="Maharashtra", msme=True)
    fid = _file("Acme_service_agreement.docx", _docx(SAMPLE))
    return S.add(T, "Acme service agreement", file_id=fid, kind="service", their_contact="legal@acme.example")


def test_rules_extract_and_review():
    from app.contracts import rules as R
    ex = R.extract("\n\n".join(SAMPLE))
    v = ex["values"]
    assert v["parties"][0]["name"] == "Ingle Tech Pvt Ltd" and v["parties"][1]["label"] == "Client"
    assert v["start"] == "2026-04-01" and v["term_months"] == 12 and v["auto_renew"] and v["renew_months"] == 12
    assert v["notice_days"] == 90 and v["payment_days"] == 90 and v["payment_phrase"] == "ninety (90) days"
    assert v["governing_law"] == "Singapore" and v["seat"] == "Singapore" and v["unlimited_liability"] and v["value"] == 1200000
    meta = {"kind": "service", "our_role": "vendor", "our_label": "Vendor", "their_label": "Client", "seat_city": "Pune", "msme": True}
    risks = R.review(ex, R.playbook({}), meta)
    by = {r["key"]: r for r in risks}
    assert by["governing_law"]["level"] == "high" and by["liability_cap"]["level"] == "high" and by["indemnity"]["level"] == "high"
    assert by["non_compete"]["level"] == "high" and "s.27" in by["non_compete"]["why"] and by["non_compete"]["source"].startswith("https://")
    assert "MSMED" in by["payment"]["why"] and by["payment"]["fix"]["replace"] == "forty five (45) days"
    assert by["arbitration"]["fix"] == {"op": "replace", "find": "Singapore", "replace": "Pune"}
    assert by["termination"]["title"].startswith("Bina wajah nikalne")
    assert by["stamp"]["level"] == "medium" and "s.35" in by["stamp"]["why"]
    assert risks[0]["level"] == "high" and all(r["draft"] == R.DRAFT_LINE for r in risks)
    # her playbook wins: notice up to 120 days is fine for her, non-compete check off
    mine = R.playbook({"notice": {"value": 120}, "non_compete": {"off": True}, "payment": {"value": 90}})
    keys = {r["key"] for r in R.review(ex, mine, meta)}
    assert "non_compete" not in keys and "payment" not in keys and "notice" not in keys
    # small readers
    assert R.duration_months("a period of thirty six (36) months") == 36 and R.duration_months("2 saal") == 24
    assert R.days_in("sixty (60) days written notice") == 60 and R.days_in("3 months' prior written notice") == 90
    assert R.parse_date("April 5, 2026") == date(2026, 4, 5) and R.parse_date("05/04/2026") == date(2026, 4, 5)
    assert R.add_months(date(2026, 1, 31), 1) == date(2026, 2, 28)
    assert R.notice_deadline(date(2027, 3, 31), 90) == date(2026, 12, 31)
    assert R.mask("Aadhaar 2345 6789 0123 PAN ABCPE1234F otp 445566") == "Aadhaar [Aadhaar hataya] PAN [PAN hataya] otp [hataya]"
    # stamp duty: only her rate
    assert R.stamp_calc("service", "Maharashtra", 1200000, [])["amount"] is None
    rates = [{"kind": "service", "state": "Maharashtra", "mode": "pct", "rate": 0.1, "min": 100, "max": 500}]
    assert R.stamp_calc("service", "maharashtra", 1200000, rates)["amount"] == 500
    assert R.stamp_calc("service", "Maharashtra", None, rates)["amount"] is None
    assert R.stamp_calc("nda", "Goa", None, [{"kind": "any", "mode": "fixed", "rate": 100}])["amount"] == 100
    # IT Act First Schedule: no e-sign for immovable sale / POA
    assert any(r["key"] == "signing" for r in R.review(R.extract("Agreement for sale of flat"), R.playbook({}), {"kind": "sale_immovable"}))


def test_add_from_word_file_redline_original_untouched(ct):
    S = ct
    v = base(S)
    assert v["number"] == "C-0001" and v["our_label"] == "Vendor" and v["their_label"] == "Client"
    assert v["counterparty"] == "Acme Retail LLP" and v["start"] == "2026-04-01" and v["end"] == "2027-03-31"
    assert v["auto_renew"] and v["notice_days"] == 90 and v["notice_by"] == "2026-12-31" and v["high"] >= 4
    assert "contracts_review" in kinds(S) and "vakil" in v["say"]
    # renewal register made from the contract's own dates
    obs = {o["kind"]: o for o in v["obligations"]}
    assert obs["renewal_notice"]["due"] == "2026-12-31" and obs["expiry"]["due"] == "2027-03-31"
    # duplicate upload is recognised
    assert S.add(T, "again", text="\n\n".join(SAMPLE))["duplicate"] is True
    from app.contracts.models import Contract
    from app.db import session
    with session() as s:
        stored = s.get(Contract, v["id"]).stored
    before = hashlib.sha256(Path(stored).read_bytes()).hexdigest()
    r = S.redline(T, "C-0001")
    assert r["original_untouched"] and r["applied"] >= 5 and r["version"] == 1
    assert hashlib.sha256(Path(stored).read_bytes()).hexdigest() == before
    from app.contracts import redline as RL
    with session() as s:
        path = [x for x in __import__("json").loads(s.get(Contract, v["id"]).versions) if x["kind"] == "redline"][0]["path"]
    tr = RL.tracked(Path(path).read_bytes())
    assert "forty five (45) days" in tr["inserted"] and "ninety (90) days" in tr["deleted"] and "India" in tr["inserted"]
    assert "Pune" in tr["inserted"] and any("aggregate liability" in t for t in tr["inserted"]) and tr["authors"] == ["Hoola (draft)"]
    assert any("Non-compete" in d or "similar business" in d for d in tr["deleted"])
    # her own edit must exist in the contract; a model edit with a made-up number is refused
    with pytest.raises(S.ContractError):
        S.redline(T, "C-0001", keys=["nothing"], edits=[{"op": "replace", "find": "words not there", "replace": "x"}])
    r2 = S.redline(T, "C-0001", keys=["none"], edits=[{"op": "replace", "find": "all losses", "replace": "third-party claims"},
                                                       {"op": "replace", "find": "Rs 12,00,000", "replace": "Rs 99,99,999"}])
    assert r2["applied"] == 1 and r2["refused"] and r2["version"] == 2
    # a text-only contract still gets a Word copy
    t2 = S.add(T, "Paste NDA", text="NON-DISCLOSURE AGREEMENT\n\nThe Recipient shall keep all information confidential.\n\n"
                                    "This Agreement shall be governed by the laws of England.", kind="nda", our_role="recipient")
    assert S.redline(T, t2["number"])["version"] == 1


def test_cards_only_her_tap(ct, monkeypatch):
    S = ct
    base(S)
    S.redline(T, "C-0001")
    card = S.card_request(T, "C-0001", "send", via="email")
    assert card["status"] == "waiting" and card["show"]["to"] == "legal@acme.example" and card["show"]["version"] == 1
    assert "contracts_card" in kinds(S)
    with pytest.raises(S.ContractError):
        S.decide(T, card["id"], True, card["hash"])                       # not a tap
    with pytest.raises(S.ContractError):
        S.decide(T, card["id"], True, "bad", by_tap=True)
    r = S.decide(T, card["id"], True, card["hash"], by_tap=True)
    assert r["links"]["email"].startswith("mailto:legal%40acme.example") and r["card"]["result"]["file"]["file_id"]
    assert S.decide(T, card["id"], True, card["hash"], by_tap=True)["card"]["status"] == "approved"   # decided once
    # a change after the card = stale
    c2 = S.card_request(T, "C-0001", "send")
    S.update(T, "C-0001", notice_days=60)
    S.add_version(T, "C-0001", text="\n\n".join(SAMPLE).replace("ninety (90) days of receipt", "forty-five (45) days of receipt"))
    with pytest.raises(S.ContractError):
        S.decide(T, c2["id"], True, c2["hash"], by_tap=True)
    assert S.cards(T, "stale")
    # e-sign blocked until stamp is in; never for First Schedule kinds
    es = S.card_request(T, "C-0001", "esign", signers=[{"name": "Meghana Ingle", "role": "Director", "contact": "9876543210"}])
    assert es["status"] == "blocked" and any("stamp" in b for b in es["show"]["blockers"])
    with pytest.raises(S.ContractError):
        S.decide(T, es["id"], True, es["hash"], by_tap=True)
    # e-stamp: amount only from her rate; fingerprint once a phone key is active
    st = S.card_request(T, "C-0001", "estamp")
    assert st["status"] == "blocked" and "rate" in st["show"]["blockers"][0]
    S.set_rates(T, [{"kind": "service", "state": "Maharashtra", "mode": "fixed", "rate": 500}])
    st = S.card_request(T, "C-0001", "estamp")
    assert st["status"] == "waiting" and st["show"]["amount"] == 500
    from app import approvekey as ak
    monkeypatch.setattr(ak, "active_on", lambda tenant: True)
    with pytest.raises(S.ContractError) as e:
        S.decide(T, st["id"], True, st["hash"], by_tap=True)
    assert str(e.value) == "fingerprint_needed"
    assert S.card_challenge(T, st["id"])["message"].startswith(f"hoola-approve|v1|{T}|contracts|{st['id']}|")
    monkeypatch.setattr(S, "_verify_fingerprint", lambda *a, **k: None)
    r = S.decide(T, st["id"], True, st["hash"], {"challenge_id": "x"}, by_tap=True)
    assert "shcilestamp" in r["links"]["portal"] and S.contract_view(T, "C-0001")["stamp"]["paid"] == "pending"
    S.stamped(T, "C-0001", cert_no="IN-MH12345678901234A")
    assert "stamp" not in {x["key"] for x in S.contract_view(T, "C-0001")["risks"]}
    es = S.card_request(T, "C-0001", "esign", signers=[{"name": "Meghana Ingle", "role": "Director"}])
    assert es["status"] == "waiting"
    r = S.decide(T, es["id"], True, es["hash"], by_tap=True)
    assert "OTP" in " ".join(r["card"]["result"]["steps"]) and S.contract_view(T, "C-0001")["status"] == "waiting_sign"
    # immovable sale: e-sign refused
    sale = S.add(T, "Flat sale", text="AGREEMENT FOR SALE between A and B for flat no 5.", kind="sale_immovable")
    S.stamped(T, sale["number"], not_needed=True)
    x = S.card_request(T, sale["number"], "esign", signers=[{"name": "A"}])
    assert x["status"] == "blocked" and "First Schedule" in x["show"]["blockers"][0]


def test_signed_notice_register_and_tick(ct):
    S = ct
    base(S)
    S.stamped(T, "C-0001", cert_no="IN-MH12345678901234A")
    v = S.signed(T, "C-0001", on="2026-04-01")
    assert v["status"] == "active" and v["notice_by"] == "2026-12-31"
    n = S.card_request(T, "C-0001", "notice", notice_kind="stop_renewal", to="legal@acme.example")
    assert "do not wish to renew" in n["show"]["letter"] and "DRAFT" in n["show"]["letter"]
    with pytest.raises(S.ContractError):
        S.card_done(T, n["id"])                                       # not approved yet
    r = S.decide(T, n["id"], True, n["hash"], by_tap=True)
    assert r["card"]["result"]["file"]["name"].endswith(".docx")
    S.card_done(T, n["id"], "email bheja")
    obs = {o["kind"]: o for o in S.contract_view(T, "C-0001")["obligations"]}
    assert obs["renewal_notice"]["status"] == "done" and not S.contract_view(T, "C-0001")["auto_renew"]
    # obligations: recurring rolls forward
    o = S.oblig_add(T, "C-0001", "Quarterly SLA report", "2026-10-30", "us", "report", 3)
    assert S.oblig_done(T, o["number"])["due"] == "2027-01-30"
    # reminders: once per window
    t2 = S.add(T, "Office rent", text="LEASE AGREEMENT\n\nThis lease is valid till 30/11/2026 and shall automatically renew for twelve (12) "
                                      "months unless either party gives thirty (30) days prior written notice.", kind="lease", our_role="tenant")
    assert t2["notice_by"] == "2026-10-31"
    S.EVENTS.clear()
    when = datetime(2026, 10, 23, 5, 0, tzinfo=timezone.utc)      # 10:30 IST
    out = S.tick(when)
    assert out["alerts"] >= 1 and "contracts_alert" in kinds(S) and "contracts_day" in kinds(S)
    assert S.tick(when)["alerts"] == 0                             # never twice
    # an active auto-renew contract past its end rolls; register Excel
    S.update(T, "C-0001", end="2026-10-01", auto_renew=True, renew_months=12)
    assert S.tick(datetime(2026, 10, 24, 5, 0, tzinfo=timezone.utc))["rolled"] == 1
    assert S.contract_view(T, "C-0001")["end"] == "2027-10-01"
    assert S.register_file(T)["file"]["name"] == "Contract_register.xlsx"


def test_forget_and_playbook(ct):
    S = ct
    base(S)
    with pytest.raises(S.ContractError):
        S.forget(T, "C-0001")
    p = S.save_playbook(T, "payment", position="90 din chalega", value=90)
    assert "payment" in p["mine"]
    assert "payment" not in {x["key"] for x in S.review_again(T, "C-0001")["risks"]}
    S.save_playbook(T, "payment", reset=True)
    assert "payment" in {x["key"] for x in S.review_again(T, "C-0001")["risks"]}
    S.set_status(T, "C-0001", "terminated")
    from app.contracts.models import Contract
    from app.db import session
    with session() as s:
        stored = s.get(Contract, S.contracts(T)[0]["id"]).stored
    r = S.forget(T, "C-0001")
    v = S.contract_view(T, "C-0001")
    assert r["forgotten"] and v["forgotten"] and not v["clauses"] and not v["their_contact"] and not Path(stored).exists()
    assert v["number"] == "C-0001" and v["end"] == "2027-03-31"


def test_tools_and_api(ct, monkeypatch):
    from app.agent import contracts_tools as AT
    from app.agent.tools import ToolError
    S = ct
    ctx = conv_ctx()
    run = lambda c, tool, **a: asyncio.run(AT.CONTRACTS_HANDLERS[tool](c, **a))   # noqa: E731
    run(ctx, "contracts", action="settings", party_name="Ingle Tech Pvt Ltd", seat_city="Pune")
    v = run(ctx, "contracts", action="add", text="\n\n".join(SAMPLE), title="Acme", kind="service", our_role="vendor")
    assert v["number"] == "C-0001"
    assert run(ctx, "contracts_dekho", what="risks", contract="C-0001")["risks"]
    assert run(ctx, "contracts_dekho", what="renewals")["renewals"]
    card = run(ctx, "contracts", action="send", contract="C-0001", to="9876543210")
    assert card["status"] == "waiting" and card["show"]["via"] == "whatsapp"
    staff = conv_ctx()
    staff.actor = "staff:abc"
    with pytest.raises(ToolError):
        run(staff, "contracts", action="stamped", contract="C-0001", cert_no="IN-MH1")
    cust = conv_ctx()
    cust.actor = "customer:wa:1"
    with pytest.raises(ToolError):
        run(cust, "contracts_dekho", what="list")
    assert len(AT.CONTRACTS_TOOLS) == 2 and set(AT.CONTRACTS_LABELS) == set(AT.CONTRACTS_HANDLERS)
    # no tool can decide a card
    assert "decide" not in str(AT.CONTRACTS_TOOLS[0]["input_schema"]["properties"]["action"]["enum"])
    from app.contracts.api import make_router
    app = FastAPI()
    app.include_router(make_router(lambda: None, lambda: "default"))
    c = TestClient(app)
    d = c.get("/api/contracts").json()
    assert d["contracts"][0]["number"] == "C-0001" and d["cards"][0]["id"] == card["id"] and d["never"]
    assert c.post(f"/api/contracts/cards/{card['id']}/decide", json={}).status_code == 400
    assert c.post(f"/api/contracts/cards/{card['id']}/decide", json={"yes": True, "hash": "bad"}).status_code == 400
    ok = c.post(f"/api/contracts/cards/{card['id']}/decide", json={"yes": True, "hash": card["hash"]})
    assert ok.status_code == 200 and ok.json()["links"]["whatsapp"].startswith("https://wa.me/919876543210")
    assert c.put("/api/contracts/playbook/notice", json={"value": 120}).json()["mine"] == ["notice"]
    assert c.post("/api/contracts/contracts/C-0001/redline", json={}).json()["version"] == 1
    assert c.post("/api/contracts/contracts/C-0001/obligations", json={"what": "Insurance copy", "due": "2026-11-15"}).json()["number"]
    S.set_rates(T, [{"kind": "any", "mode": "fixed", "rate": 100}])
    st = c.post("/api/contracts/contracts/C-0001/cards", json={"kind": "estamp"}).json()
    from app import approvekey as ak
    monkeypatch.setattr(ak, "active_on", lambda tenant: True)
    assert c.post(f"/api/contracts/cards/{st['id']}/decide", json={"yes": True, "hash": st["hash"]}).status_code == 428
    assert c.post(f"/api/contracts/cards/{st['id']}/challenge").json()["challenge_id"]
    assert c.get("/api/contracts/rules").json()["first_schedule"]["poa"]
    assert c.get("/api/contracts/contracts/C-0001").json()["clauses"]
