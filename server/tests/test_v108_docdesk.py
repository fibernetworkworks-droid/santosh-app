"""v0.108 Document desk + Sheet Milaan on her PC: the REAL Santosh Desktop actions behind a fake bridge. Files are found on
the PC, read in full, numbered per job; two sheets are compared by key in code; her Haan saves an updated COPY (formatting
kept) into the Santosh Inbox and opens it; the original is untouched."""
import asyncio
import io
import json
import sys
from pathlib import Path

import pytest

sys.path.insert(0, str(Path(__file__).resolve().parents[2] / "desktop"))

from app.db import Conversation, session  # noqa: E402


def conv():
    with session() as s:
        c = Conversation(tenant_id="default")
        s.add(c)
        s.flush()
        return c.id


@pytest.fixture
def pc(env, tmp_path, monkeypatch):
    monkeypatch.setenv("SANTOSH_DESKTOP_HOME", str(tmp_path / "desk"))
    from santosh_desktop import actions as A
    from santosh_desktop.config import Config
    from santosh_desktop.index import Index
    from santosh_desktop.safety import unique
    root = tmp_path / "pcfiles"
    (root / "Downloads").mkdir(parents=True)
    inbox = tmp_path / "Santosh Inbox"
    inbox.mkdir()
    cfg = Config(server="http://x", folders=[str(root)], inbox=str(inbox))
    acts = A.Actions(cfg, Index(tmp_path / "i.db"), log=lambda m: None)
    opened = []
    import subprocess as _sp
    from types import SimpleNamespace as _NS
    # only the desktop's own `subprocess` name is replaced (Chromium for PDFs still starts normally)
    monkeypatch.setattr(A, "subprocess", _NS(Popen=lambda args, **k: opened.append(args[-1]), run=_sp.run))
    monkeypatch.setattr(A.sys, "platform", "linux")
    from app.mailbox import service as ms
    from app.pc import service as pcs
    from app.pc.hub import hub, PcOffline
    state = {"online": True, "calls": []}

    async def call(tenant, action, args=None, device_id=None, timeout=None, log=True):
        args = args or {}
        state["calls"].append(action)
        if not state["online"]:
            raise PcOffline("PC_OFFLINE")
        if action == "send_file":
            p = acts._file(args["path"])
            row = ms.store_bytes(tenant, p.name, p.read_bytes(), "to_phone")
            with session() as s:
                from app.db import PcFile
                s.get(PcFile, row.id).pc_path = str(p)
            return {"file_id": row.id, "name": p.name}
        if action == "save_file":
            f = pcs.get_file(args["file_id"])
            dest = unique(inbox, args["name"])
            dest.write_bytes(Path(f.stored_path).read_bytes())
            return {"saved_path": str(dest), "name": dest.name}
        if action == "search":
            words = [w for w in args["query"].lower().split() if len(w) > 1]
            res = [{"path": str(p), "name": p.name, "modified": "2026-10-22 10:00"} for p in root.rglob("*")
                   if p.is_file() and all(w in p.name.lower() for w in words)]
            return {"results": res}
        return acts.handle(action, args)
    monkeypatch.setattr(hub, "call", call)
    return {"root": root, "inbox": inbox, "acts": acts, "opened": opened, "state": state}


def make_old(path):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    wb = Workbook()
    ws = wb.active
    ws.title = "Customers"
    ws["A1"] = "Patil Fibernet - ONU list"
    ws.append([])
    ws.append(["Sr No", "Customer ID", "Name", "MAC", "Plan"])
    for c in ws[3]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1F4E78")
    rows = [(1, "C101", "Ramesh", "AA:BB:CC:00:00:01", 499), (2, "C102", "Suresh", "AA:BB:CC:00:00:02", 699),
            (3, "C103", "Mahesh", "AA:BB:CC:00:00:03", 499), (4, "C104", "Ganesh", "AA:BB:CC:00:00:04", 999)]
    for r in rows:
        ws.append(list(r))
    for row in ws.iter_rows(min_row=4, max_row=7):
        row[4].number_format = "₹#,##0"
        row[0].fill = PatternFill("solid", fgColor="E2EFDA")
    ws.column_dimensions["D"].width = 24
    wb.save(path)


def make_new(path):
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.append(["customer id", "Name", "MAC Address", "Plan"])
    ws.append(["C101", "Ramesh", "aa-bb-cc-00-00-01", 499])        # same MAC, other format: not a change
    ws.append(["C102", "Suresh", "aa-bb-cc-00-00-99", 699])        # MAC changed
    ws.append(["C103", "Mahesh", "AABB.CC00.0003", 799])            # plan changed (not asked), MAC same
    ws.append(["C105", "Dinesh", "AA:BB:CC:00:00:05", 499])         # new customer
    wb.save(path)                                                    # C104 missing


def test_sheet_milaan_end_to_end_on_her_pc(pc):
    from app import docdesk as dd
    from openpyxl import load_workbook
    old_p = pc["root"] / "MAC data.xlsx"
    make_old(old_p)
    before = old_p.read_bytes()
    make_new(pc["root"] / "Downloads" / "new list from whatsapp.xlsx")
    c = conv()
    first = asyncio.run(dd.read("default", c, "MAC data"))
    assert first["n"] == 1 and "Ramesh" in first["text"] and first["where"].endswith("MAC data.xlsx")
    rec = asyncio.run(dd.recent_on_pc("default", "excel"))
    assert rec[0]["name"] == "new list from whatsapp.xlsx"
    second = asyncio.run(dd.read("default", c, rec[0]["path"]))
    assert second["n"] == 2
    out = asyncio.run(dd.sheet_compare("default", c, "pehli wali", "doosri wali", update=["MAC"]))
    assert out["status"] == "waiting_owner"
    assert out["key"]["old"] == "Customer ID"
    assert out["counts"]["changed"] == 1 and out["counts"]["added"] == 1 and out["counts"]["missing"] == 1
    assert out["changed"][0]["cells"] == [{"col": "MAC", "old": "AA:BB:CC:00:00:02", "new": "aa-bb-cc-00-00-99"}]
    assert not list(pc["inbox"].iterdir())                       # nothing saved before her tap
    card = asyncio.run(dd.decide("default", out["edit_id"], True, {"missing": "mark"}))
    assert card["status"] == "saved"
    saved = Path(card["out"]["saved_path"])
    assert saved.parent == pc["inbox"] and saved.name.startswith("MAC data (Hoola ")
    assert pc["opened"] == [str(saved)]
    assert old_p.read_bytes() == before                           # original untouched
    ws = load_workbook(saved)["Customers"]
    assert ws["D5"].value == "aa-bb-cc-00-00-99"
    assert ws["E6"].value == 499                                   # plan not asked -> not changed
    assert ws["B8"].value == "C105" and ws["A8"].value == 5 and ws["E8"].number_format == "₹#,##0"
    assert ws["A8"].fill.fgColor.rgb.endswith("E2EFDA")            # style copied from the row above
    assert ws["C3"].font.bold and ws["C3"].fill.fgColor.rgb.endswith("1F4E78")
    assert ws.column_dimensions["D"].width == 24
    assert ws["F3"].value == "Hoola milaan" and ws["F7"].value == "naye mein nahi"
    # the new copy is on the job list and in the diary
    files = dd.job_files("default", c)
    assert len(files) == 3 and files[-1].pc_path == str(saved)
    from app import yaad
    blk = yaad.diary_block("default", c)
    assert "1. MAC data.xlsx" in blk and "Sheet Milaan" in blk


def test_mac_as_key_and_every_common_column(pc):
    from app import sheetmilaan as sm
    buf_old, buf_new = io.BytesIO(), io.BytesIO()
    from openpyxl import Workbook
    for buf, rows in ((buf_old, [["MAC", "Location", "Status"], ["AA:BB:CC:DD:EE:01", "Akola", "active"],
                                 ["AA:BB:CC:DD:EE:02", "Amravati", "active"]]),
                      (buf_new, [["mac address", "location", "status"], ["aabb.ccdd.ee01", "Akola", "inactive"],
                                 ["AA-BB-CC-DD-EE-02", "Amravati", "active"]])):
        wb = Workbook()
        for r in rows:
            wb.active.append(r)
        wb.save(buf)
    old = sm.load("a.xlsx", buf_old.getvalue())
    new = sm.load("b.xlsx", buf_new.getvalue())
    res = sm.compare(old, new)
    assert res["key"]["old"] == "MAC" and res["counts"]["changed"] == 1 and res["counts"]["same"] == 1
    assert res["changed"][0]["cells"] == [{"col": "Status", "old": "active", "new": "inactive"}]
    with pytest.raises(sm.SheetError):
        sm.compare(old, new, key="Nope")


def test_csv_milaan_and_word_edit_keep_formatting(pc):
    from app import docdesk as dd
    import docx
    c = conv()
    (pc["root"] / "old.csv").write_text("Serial,Owner\nS1,Ram\nS2,Shyam\n")
    (pc["root"] / "new.csv").write_text("serial,owner\nS1,Ram\nS2,Ghanshyam\nS3,Mohan\n")
    out = asyncio.run(dd.sheet_compare("default", c, str(pc["root"] / "old.csv"), str(pc["root"] / "new.csv")))
    card = asyncio.run(dd.decide("default", out["edit_id"], True))
    text = Path(card["out"]["saved_path"]).read_text(encoding="utf-8-sig")
    assert "S2,Ghanshyam" in text and "S3,Mohan" in text
    # Word: the price is split across two runs, one bold
    d = docx.Document()
    p = d.add_paragraph("Quotation rate: ")
    r = p.add_run("Rs 5")
    r.bold = True
    r2 = p.add_run("00 per month")
    r2.bold = True
    d.add_paragraph("Valid till 31 Oct.")
    d.save(pc["root"] / "quote.docx")
    before = (pc["root"] / "quote.docx").read_bytes()
    res = asyncio.run(dd.propose_edit("default", c, "quote", [{"find": "Rs 500", "replace": "Rs 650"},
                                                                {"find": "not there", "replace": "x"},
                                                                {"after": "Valid till", "insert": "GST extra."}]))
    assert res["status"] == "waiting_owner" and res["changes"] == 2 and res["not_found"] == ["not there"]
    card = asyncio.run(dd.decide("default", res["edit_id"], True))
    new = docx.Document(card["out"]["saved_path"])
    para = new.paragraphs[0]
    assert para.text == "Quotation rate: Rs 650 per month"
    assert para.runs[1].bold and "Rs 650" in para.runs[1].text
    assert [x.text for x in new.paragraphs][-1] == "GST extra."
    assert (pc["root"] / "quote.docx").read_bytes() == before
    # rejected: nothing saved
    res2 = asyncio.run(dd.propose_edit("default", c, "quote", [{"find": "Valid", "replace": "Good"}]))
    n = len(list(pc["inbox"].iterdir()))
    assert asyncio.run(dd.decide("default", res2["edit_id"], False))["status"] == "rejected"
    assert len(list(pc["inbox"].iterdir())) == n


def test_search_across_files_and_numbers(pc):
    from app import docdesk as dd
    c = conv()
    (pc["root"] / "notes a.txt").write_text("Meeting with Kulkarni on Monday.\n\nRouter price 2400.")
    (pc["root"] / "notes b.txt").write_text("Payment from Joshi pending.\n\nKulkarni wants 3 routers.")
    asyncio.run(dd.read("default", c, str(pc["root"] / "notes a.txt")))
    asyncio.run(dd.read("default", c, str(pc["root"] / "notes b.txt")))
    hits = dd.search("default", c, "Kulkarni")
    assert {h["n"] for h in hits} == {1, 2}
    assert dd.find("default", c, "aakhri").name == "notes b.txt"
    assert dd.find("default", c, "2").name == "notes b.txt"
    big = "line of text\n" * 3000
    (pc["root"] / "big.txt").write_text(big)
    r1 = asyncio.run(dd.read("default", c, str(pc["root"] / "big.txt")))
    assert r1["parts"] >= 3 and "next" in r1
    r3 = asyncio.run(dd.read("default", c, "3", part=3))
    assert r3["part"] == 3


def test_pc_offline_copy_comes_to_phone_and_link_download(pc, monkeypatch):
    import httpx
    from app import docdesk as dd, web
    from openpyxl import Workbook
    c = conv()
    buf = io.BytesIO()
    wb = Workbook()
    wb.active.append(["ID", "Name"])
    wb.active.append(["1", "A"])
    wb.save(buf)

    def handler(req):
        if "export" in str(req.url):
            return httpx.Response(200, content=buf.getvalue(), headers={
                "content-type": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                "content-disposition": 'attachment; filename="Stock list.xlsx"'})
        return httpx.Response(200, text="<html>Sign in</html>", headers={"content-type": "text/html"})
    monkeypatch.setattr(web, "_transport", httpx.MockTransport(handler))
    monkeypatch.setattr(web, "_public_host", lambda h: None)
    got = asyncio.run(dd.read("default", c, "https://docs.google.com/spreadsheets/d/abc123/edit#gid=0"))
    assert got["name"] == "Stock list.xlsx" and "Name" in got["text"]
    with pytest.raises(dd.DeskError, match="LOGIN_NEEDED"):
        asyncio.run(dd.read("default", c, "https://example.com/private/file"))
    pc["state"]["online"] = False
    out = asyncio.run(dd.save_to_pc("default", "x.txt", b"hello"))
    assert out["where"] == "phone" and out["url"]


def test_desktop_open_file_and_recent_are_safe(pc, tmp_path):
    from santosh_desktop.actions import ActionError
    acts = pc["acts"]
    (pc["root"] / "a.xlsx").write_bytes(b"x")
    (pc["root"] / "run.bat").write_text("echo")
    outside = tmp_path / "secret.xlsx"
    outside.write_bytes(b"x")
    assert acts.handle("open_file", {"path": str(pc["root"] / "a.xlsx")})["name"] == "a.xlsx"
    with pytest.raises(ActionError):
        acts.handle("open_file", {"path": str(pc["root"] / "run.bat")})
    with pytest.raises(ActionError):
        acts.handle("open_file", {"path": str(outside)})
    names = [f["name"] for f in acts.handle("recent", {})["files"]]
    assert "a.xlsx" in names and "secret.xlsx" not in names


def test_api_decide_and_tools(pc):
    from fastapi.testclient import TestClient
    from app.main import app
    from app.agent.tools import HANDLERS, Ctx
    c = conv()
    (pc["root"] / "o.csv").write_text("ID,Qty\n1,5\n2,7\n")
    (pc["root"] / "n.csv").write_text("ID,Qty\n1,6\n2,7\n")
    ctx = Ctx(tenant_id="default", conversation_id=c)
    out = asyncio.run(HANDLERS["sheet_milaan"](ctx, old=str(pc["root"] / "o.csv"), new=str(pc["root"] / "n.csv")))
    assert out["counts"]["changed"] == 1
    lst = asyncio.run(HANDLERS["doc_desk"](ctx, action="files"))
    assert len(lst["files"]) == 2
    cl = TestClient(app)
    full = cl.get(f"/api/docdesk/edits/{out['edit_id']}").json()
    assert full["result"]["changed"][0]["cells"][0]["new"] == "6"
    r = cl.post(f"/api/docdesk/edits/{out['edit_id']}/decide", json={"yes": True}).json()
    assert r["status"] == "saved"
    assert cl.get("/api/docdesk", params={"conv": c}).json()["edits"][0]["status"] == "saved"
    found = asyncio.run(HANDLERS["yaad_khojo"](ctx, query="o.csv"))
    assert found["count"] >= 1


def test_pdf_becomes_a_corrected_new_file(pc):
    from app import docdesk as dd
    from reportlab.pdfgen import canvas
    c = conv()
    p = pc["root"] / "bill.pdf"
    cv = canvas.Canvas(str(p))
    cv.drawString(72, 720, "Customer: Patil Steel")
    cv.drawString(72, 700, "Total amount Rs 500 only")
    cv.save()
    res = asyncio.run(dd.propose_edit("default", c, str(p), [{"find": "Rs 500", "replace": "Rs 560"}]))
    assert res["changes"] == 1 and any("saada" in w for w in res["warnings"])
    card = asyncio.run(dd.decide("default", res["edit_id"], True))
    out = Path(card["out"]["saved_path"])
    assert out.suffix in (".pdf", ".docx") and out.stat().st_size > 500
    if out.suffix == ".pdf":
        from pypdf import PdfReader
        assert "Rs 560" in PdfReader(str(out)).pages[0].extract_text()
