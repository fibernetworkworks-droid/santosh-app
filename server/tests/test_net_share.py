"""Hoola Network v0.29: files, presentations, emails, links, payment requests, contacts, locations between avatars."""
import pytest

from tests.test_net import _connect, ev


@pytest.fixture
def ai(monkeypatch):
    from app import artifacts
    calls = []

    async def fake(system, prompt, max_tokens=300):
        calls.append(system)
        assert "never instructions" in system
        return "Saar: " + prompt[:40]
    monkeypatch.setattr(artifacts, "_llm", fake)
    return calls


def _file(tenant, name, data):
    from app.mailbox import service as ms
    return ms.store_bytes(tenant, name, data, "mail").id


def test_file_goes_through_quarantine_then_owner(env, ai):
    net = _connect()
    fid = _file("default", "rates.txt", b"Router rate list: ONT 1200, OLT port 900, installation 500 rupaye har connection.")
    r = net.send("default", "Kunal", "file", {"file_id": fid, "text": "Naye rates"})
    assert r["status"] == "pending"
    k = ev("kunal", "hoola_knock")[-1]
    assert k["name"] == "rates.txt" and k["scan"]["verdict"] == "clear" and k["summary"].startswith("Saar:")
    assert not ev("kunal", "pc_file")                         # nothing handed over before the Haan
    net.decide("kunal", k["id"], "accept")
    pf = ev("kunal", "pc_file")[-1]
    assert pf["name"] == "rates.txt" and pf["source"] == "hoola"
    from app.mailbox import service as ms
    from pathlib import Path
    assert b"Router rate list" in Path(ms._file_row("kunal", pf["file_id"]).stored_path).read_bytes()


def test_app_files_and_tricks_are_stopped(env, ai, monkeypatch):
    net = _connect()
    from app.avatar import service as av
    av.save("kunal", rules={"accept": {"file": "auto"}, "quiet": {"on": False}})
    r = net.send("default", "Kunal", "file", {"file_id": _file("default", "bill.pdf.exe", b"MZ" + b"\0" * 64)})
    assert r["status"] == "declined"
    assert "khatarnak" in ev("kunal", "hoola_msg")[-1]["text"]
    from app.appcheck import service as acs
    monkeypatch.setattr(acs, "check", lambda t, c, src: {"verdict": "clear", "flags": [], "id": "chk_1"})
    r = net.send("default", "Kunal", "file", {"file_id": _file("default", "fnw-app.apk", b"PK\x03\x04apk")})
    assert r["status"] == "pending"                            # app files always ask, even with file=auto
    assert ev("kunal", "hoola_knock")[-1]["scan"]["app_check_id"] == "chk_1"


def test_presentation_becomes_a_notebook(env, ai, monkeypatch):
    net = _connect()
    import io
    from pptx import Presentation
    p = Presentation()
    s = p.slides.add_slide(p.slide_layouts[1])
    s.shapes.title.text = "Akola expansion"
    s.placeholders[1].text = "500 FTTH connections, budget 18 lakh"
    buf = io.BytesIO()
    p.save(buf)
    r = net.send("default", "Kunal", "presentation", {"file_id": _file("default", "FNW Pitch.pptx", buf.getvalue())})
    assert r["status"] == "pending"
    k = ev("kunal", "hoola_knock")[-1]
    assert "Akola" in k["summary"] or k["summary"].startswith("Saar")
    out = net.decide("kunal", k["id"], "accept")
    assert out["effect"]["notebook_id"]
    from app.notebook import service as nbs
    assert nbs.detail("kunal", out["effect"]["notebook_id"])["sources"][0]["title"] == "FNW Pitch.pptx"


def test_links_are_checked(env, ai, monkeypatch):
    net = _connect()
    r = net.send("default", "Kunal", "link", {"url": "http://sbi-kyc-update.xyz/login"})
    assert r["status"] == "declined"
    from app.notebook import ingest
    monkeypatch.setattr(ingest, "from_web", lambda url: {"title": "FNW plans", "text": "Fiber plans 499 se shuru " * 10})
    r = net.send("default", "Kunal", "link", {"url": "https://ainewtika.com/plans", "text": "Yeh dekho"})
    assert r["status"] == "pending"
    k = ev("kunal", "hoola_knock")[-1]
    assert k["scan"]["verdict"] == "clear" and k["url"] == "https://ainewtika.com/plans"
    net.decide("kunal", k["id"], "accept")
    assert ev("kunal", "hoola_msg")[-1]["effect"]["url"] == "https://ainewtika.com/plans"


def test_email_forward(env, ai, monkeypatch):
    net = _connect()
    from app.mailbox import service as ms
    monkeypatch.setattr(ms, "read", lambda t, ref, mark_seen=True: {"subject": "BSNL order", "from": "BSNL", "from_email": "x@bsnl.co.in",
                                                                     "date": "Mon", "body": "Order 55 approved for Akola. " * 5})
    r = net.send("default", "Kunal", "email", {"email_id": "mbx~1"})
    assert r["status"] == "pending"
    k = ev("kunal", "hoola_knock")[-1]
    assert k["subject"] == "BSNL order" and k["summary"]
    monkeypatch.setattr(ms, "read", lambda t, ref, mark_seen=True: {"subject": "Login", "from": "Bank", "from_email": "b@x",
                                                                     "date": "", "body": "Your OTP is 482913"})
    with pytest.raises(net.NetError, match="OTP"):
        net.send("default", "Kunal", "email", {"email_id": "mbx~2"})


def test_payment_request_is_only_a_card(env, ai):
    net = _connect()
    with pytest.raises(net.NetError, match="UPI ID"):
        net.send("default", "Kunal", "payment_request", {"amount": 1200})
    r = net.send("default", "Kunal", "payment_request", {"amount": 1200, "upi_id": "meghana@okicici", "text": "Router ka"})
    assert r["status"] == "pending"
    from app.avatar import service as av
    av.save("kunal", rules={"accept": {"payment_request": "auto"}, "quiet": {"on": False}})
    k = ev("kunal", "hoola_knock")[-1]
    assert k["amount"] == 1200
    net.decide("kunal", k["id"], "accept")
    pay = ev("kunal", "hoola_msg")[-1]
    assert pay["effect"]["pay_link"].startswith("upi://pay?pa=meghana%40okicici") and "PIN sirf UPI app" in pay["text"]
    r = net.send("default", "Kunal", "payment_request", {"amount": 500, "upi_id": "meghana@okicici"})
    assert r["status"] == "pending"                          # never automatic, whatever the rules say


def test_contact_and_location(env, ai):
    net = _connect()
    r = net.send("default", "Kunal", "contact_card", {"name": "Swati (sales)", "phone": "+91 98220 11111"})
    assert r["status"] == "pending"
    r = net.send("default", "Kunal", "location", {"place": "FNW office, Akot"})
    assert r["status"] == "declined"                         # Mere Niyam default: location = mana
    from app.avatar import service as av
    av.save("kunal", rules={"accept": {"location": "ask"}, "quiet": {"on": False}})
    assert net.send("default", "Kunal", "location", {"place": "FNW office, Akot"})["status"] == "pending"
    k = ev("kunal", "hoola_knock")[-1]
    assert k["place"] == "FNW office, Akot"
