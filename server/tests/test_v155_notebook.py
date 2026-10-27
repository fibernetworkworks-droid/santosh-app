"""v0.155 Hoola Notebook = NotebookLM / Gemini Notebook parity: sources on / off + source guide + refresh, chat with saved
history / Configure / follow-ups, notes (and note -> source), Discover (fast + deep research), and every Studio item
(report, flashcards, quiz, data table, mind map, infographic, slides + revise, video, audio + Join) with real Chromium
renders, real ffmpeg video and signed file links."""
import asyncio
import io
import json
import re
import subprocess
import time

import pytest
from fastapi.testclient import TestClient

import app.notebook.models  # noqa: F401
from app.agent.tools import dispatch
from tests.test_bazar import _events
from tests.test_life import conv_ctx

SRC1 = "FNW Q3 report.\nRevenue was Rs 42 lakh, up 18 percent.\nGoa SSA launched in September. MTTR fell to 4 hours."
SRC2 = "Hiring plan: 5 field engineers in Pune. Training at ITPC in November."
SRC3 = "Secret plan: open Nashik office in December 2026."


class FakeAI:
    """Answers each kind of notebook prompt with a plausible, valid answer (no network)."""

    def __init__(self):
        self.prompts = []

    async def __call__(self, system, prompt, max_tokens=4000):
        self.prompts.append((system, prompt))
        s, p = system, prompt
        if "search keywords" in p:
            return "revenue income growth"
        if "source guides" in s:
            return json.dumps({"summary": "Q3 revenue 42 lakh.", "topics": ["Revenue", "Goa"], "questions": ["Revenue kitna?"]})
        if "overview card" in s:
            return json.dumps({"title": "FNW Q3", "overview": "**Revenue** and hiring.", "questions": ["Q1?", "Q2?"]})
        if "study flashcards" in s:
            return json.dumps({"title": "FNW cards", "cards": [{"front": "Q3 revenue?", "back": "Rs 42 lakh", "source": 1},
                                                              {"front": "MTTR?", "back": "4 hours", "source": 1}]})
        if "multiple-choice" in s:
            return "```json\n" + json.dumps({"title": "FNW quiz", "questions": [
                {"q": "Q3 revenue?", "options": ["10 lakh", "42 lakh", "1 crore", "5 lakh"], "answer": 1, "hint": "lakh",
                 "why": "Report says 42 lakh", "source": 1}, {"q": "bad", "options": ["x"], "answer": 3}]}) + "\n```"
        if "data table" in s:
            return json.dumps({"title": "SSA table", "columns": ["SSA", "Fact", "Source"],
                               "rows": [["Goa", "Launched Sept", "1"], ["Pune", "5 engineers", "2"]], "notes": ""})
        if "mind maps" in s:
            return json.dumps({"root": {"label": "FNW Q3", "children": [{"label": "Revenue", "children": [{"label": "42 lakh"}]},
                                                                       {"label": "Team", "children": [{"label": "Pune"}]}]}})
        if "design infographics" in s:
            return json.dumps({"title": "FNW Q3", "subtitle": "ek nazar", "hero": {"value": "₹42 lakh", "label": "revenue"},
                               "stats": [{"value": "18%", "label": "growth"}], "sections": [{"icon": "📡", "heading": "Goa",
                               "text": "Launched", "points": ["Sept"]}], "steps": ["a", "b"], "takeaway": "Grow.", "sources": [1]})
        if "revise one slide" in s:
            return json.dumps({"layout": "bullets", "title": "Naya title", "bullets": ["changed"]})
        if "presentation designer" in s:
            return json.dumps({"title": "FNW deck", "slides": [
                {"layout": "title", "title": "FNW Q3", "subtitle": "Revenue 42 lakh", "visual": "fiber cables"},
                {"layout": "stats", "title": "Numbers", "stats": [{"value": "₹42L", "label": "Revenue"}], "notes": "say it"},
                {"layout": "bullets", "title": "Next", "bullets": ["Nashik", "Pune hiring"]}]})
        if "narrated explainer videos" in s:
            return json.dumps({"title": "FNW video", "scenes": [
                {"headline": "FNW Q3", "sub": "42 lakh", "visual": "fiber", "narration": "Namaste. FNW ne Q3 mein 42 lakh kamaye."},
                {"headline": "Goa", "points": ["Sept launch"], "visual": "goa beach", "narration": "Goa SSA September mein shuru hua."}]})
        if "podcast scripts" in s:
            if "joined the show live" in p:
                return "Priya: Achha sawaal!\nRavi: Report ke hisaab se revenue 42 lakh tha.\nPriya: Chaliye, wahi baat aage badhaate hain."
            return ("Ravi: Namaste! Aaj FNW ki report.\nPriya: Revenue 42 lakh hua.\nRavi: 18 percent zyada.\n"
                    "Priya: Goa launch hua.\nRavi: Recap: 42 lakh, Goa.\nPriya: Dhanyavaad!")
        if "plan web research" in s:
            return json.dumps({"queries": ["ftth india growth", "bsnl bharat fiber 2026"]})
        if "deep research reports" in s:
            return "# FTTH in India\n\nFiber is growing [1].\n\n## Sources\n[1] Page - https://example.in/a"
        if "suggest useful report formats" in s:
            return json.dumps({"formats": [{"title": "Board memo", "description": "for directors", "prompt": "Write a memo"}]})
        if "Question:" in p:
            nums = re.findall(r"^\[(\d+)\]", p, re.M)
            hit = next((n for n in nums if "42 lakh" in p.split(f"[{n}]")[1].split("\n[")[0]), nums[0] if nums else "1")
            return f"Revenue **Rs 42 lakh** tha [{hit}].\n>>> Growth kitni thi? | Goa kab shuru hua? | MTTR kya hai?"
        if "Write a dense summary" in p:
            return "Summary"
        return "## Briefing doc\n- Revenue 42 lakh (S1)\n\n| SSA | Fact |\n|---|---|\n| Goa | Launch |"


@pytest.fixture
def ai(env, monkeypatch):
    from app.notebook import service as ns
    fake = FakeAI()
    monkeypatch.setattr(ns, "_llm", fake)
    monkeypatch.setenv("NOTEBOOK_DIR", "nbdata")
    return fake


def run(name, args, ctx):
    return asyncio.run(dispatch(name, args, ctx))


def _nb(ctx, title="FNW Q3"):
    out = run("notebook_create", {"title": title, "sources": [{"text": SRC1, "title": "Q3 report"},
                                                              {"text": SRC2, "title": "HR"},
                                                              {"text": SRC3, "title": "Secret"}]}, ctx)
    assert len(out["added"]) == 3
    return out["id"]


def _made(oid):
    from app.notebook import studio
    asyncio.run(studio.run(oid))
    o = studio._row(oid)
    assert o.status == "ready", o.error
    return o


def test_sources_on_off_guide_view_rename(ai):
    from app.notebook import service as ns
    ctx = conv_ctx()
    nid = _nb(ctx)
    d = ns.set_enabled("default", nid, [ns._sources(nid)[2].id], False)
    assert [s["enabled"] for s in d["sources"]] == [True, True, False]
    out = asyncio.run(ns.ask("default", nid, "Revenue kitna tha?"))
    assert "42 lakh" in out["answer"] and out["citations"][0]["title"] == "Q3 report"
    assert "Nashik" not in ai.prompts[-1][1]                         # an unticked source never reaches the AI
    assert out["followups"] == ["Growth kitni thi?", "Goa kab shuru hua?", "MTTR kya hai?"] and ">>>" not in out["answer"]
    ns.set_enabled("default", nid, [ns._sources(nid)[1].id], True, only=True)
    assert [s.enabled for s in ns._sources(nid)] == [False, True, False]
    ns.set_enabled("default", nid, [], True, only=True)                                   # nothing ticked
    with pytest.raises(ns.NotebookError, match="tick"):
        asyncio.run(ns.ask("default", nid, "x"))
    ns.set_enabled("default", nid, None, True)
    g = asyncio.run(ns.source_guide("default", nid, "1"))
    n = len(ai.prompts)
    assert g["topics"] == ["Revenue", "Goa"] and asyncio.run(ns.source_guide("default", nid, "1")) == g and len(ai.prompts) == n
    v = ns.source_view("default", nid, "Q3 report")
    assert v["passages"] and v["guide"]["summary"]
    assert ns.rename_source("default", nid, "2", "Hiring")["sources"][1]["title"] == "Hiring"
    gg = asyncio.run(ns.notebook_guide("default", nid))
    assert gg["questions"] and asyncio.run(ns.notebook_guide("default", nid)) == gg        # kept until sources change
    with pytest.raises(ns.NotebookError, match="dobara padha"):
        ns.refresh_source("default", nid, "1")                                            # text sources can't be re-read


def test_refresh_web_source(ai):
    import httpx
    from app.notebook import ingest
    from app.notebook import service as ns
    body = {"t": "GPON splitters 1:32. " * 30}

    def handler(req):
        return httpx.Response(200, text=f"<html><head><title>Guide</title></head><body><p>{body['t']}</p></body></html>",
                              headers={"content-type": "text/html"})
    ingest.set_http(httpx.Client(transport=httpx.MockTransport(handler)))
    try:
        nb = ns.create("default", "", "Web")
        ns.add_source("default", nb["id"], {"url": "https://example.in/guide"})
        assert ns.refresh_source("default", nb["id"], "1")["changed"] is False
        body["t"] = "GPON splitters 1:64 now. " * 30
        r = ns.refresh_source("default", nb["id"], "1")
        assert r["changed"] and "1:64" in ns._sources(nb["id"])[0].text
    finally:
        ingest.set_http(None)


def test_chat_history_config_and_notes(ai):
    from app.notebook import service as ns
    ctx = conv_ctx()
    nid = _nb(ctx)
    asyncio.run(ns.ask("default", nid, "Revenue kitna tha?"))
    asyncio.run(ns.ask("default", nid, "Aur growth?"))
    assert "Earlier in this chat" in ai.prompts[-1][1] and "Revenue kitna tha?" in ai.prompts[-1][1]
    h = ns.chat_history("default", nid)
    assert [m["role"] for m in h] == ["user", "assistant", "user", "assistant"] and h[1]["followups"]
    cfg = ns.set_config("default", nid, goal="learning", length="shorter", language="mr")
    assert cfg == {"goal": "learning", "custom": "", "length": "shorter", "language": "mr"}
    asyncio.run(ns.ask("default", nid, "MTTR samjhao"))
    sysmsg = ai.prompts[-1][0]
    assert "LEARNING GUIDE" in sysmsg and "Marathi" in sysmsg and "short" in sysmsg
    with pytest.raises(ns.NotebookError, match="Custom goal"):
        ns.set_config("default", nid, goal="custom")
    ns.set_config("default", nid, goal="custom", custom="GST exam ka teacher", language="auto")
    asyncio.run(ns.ask("default", nid, "x?"))
    assert "GST exam ka teacher" in ai.prompts[-1][0] and "Always answer in" not in ai.prompts[-1][0]
    # notes: save an answer, her own note, edit, turn into a source, delete
    saved = ns.save_message("default", nid, h[1]["id"])
    assert saved["kind"] == "answer" and saved["title"] == "Revenue kitna tha?"
    mine = ns.note_create("default", nid, "", "Meri baat: Nashik office ka budget 20 lakh.")
    assert mine["kind"] == "user" and mine["title"].startswith("Meri baat")
    ns.note_update("default", nid, mine["id"], title="Budget")
    d = ns.note_to_source("default", nid, mine["id"])
    assert d["sources"][-1]["kind"] == "note" and d["sources"][-1]["title"] == "Note: Budget"
    assert ns.note_delete("default", nid, saved["id"]) == "Revenue kitna tha?"
    assert ns.clear_chat("default", nid) == 8 and ns.chat_history("default", nid) == []


def test_text_studio_items(ai):
    from app.notebook import service as ns
    from app.notebook import studio
    ctx = conv_ctx()
    nid = _nb(ctx)
    fc = _made(studio.start("default", ctx.conversation_id, nid, "flashcards", {"count": "fewer", "difficulty": "hard", "topic": "money"})["id"])
    assert len(fc.data["cards"]) == 2 and "flashcards.csv" in fc.files and "hard" in ai.prompts[-1][1]
    qz = _made(studio.start("default", "", nid, "quiz", {})["id"])
    assert len(qz.data["questions"]) == 1 and qz.data["questions"][0]["answer"] == 1      # the broken question was dropped
    tb = _made(studio.start("default", "", nid, "table", {"columns": "SSA, Fact"})["id"])
    assert tb.data["columns"] == ["SSA", "Fact", "Source"] and {"table.csv", "table.xlsx"} <= set(tb.files)
    from openpyxl import load_workbook
    p = ns.out_dir("default", nid, tb.id) / "table.xlsx"
    assert load_workbook(p).active.max_row == 3
    rp = _made(studio.start("default", "", nid, "report", {"type": "study_guide", "language": "hi"})["id"])
    assert {"report.md", "report.docx", "report.pdf"} <= set(rp.files) and "Hindi" in ai.prompts[-1][1]
    assert (ns.out_dir("default", nid, rp.id) / "report.pdf").read_bytes()[:4] == b"%PDF"
    with pytest.raises(ns.NotebookError, match="Custom report"):
        studio.start("default", "", nid, "report", {"type": "custom"})
    cu = _made(studio.start("default", "", nid, "report", {"type": "custom", "prompt": "Board memo for directors"})["id"])
    assert "Board memo for directors" in ai.prompts[-1][1] and cu.data["type"] == "custom"
    mm = _made(studio.start("default", "", nid, "mind_map", {})["id"])
    assert mm.data["nodes"] == 5 and "mindmap.png" in mm.files
    ig = _made(studio.start("default", "", nid, "infographic", {"orientation": "landscape", "detail": "concise", "style": "heritage"})["id"])
    png = (ns.out_dir("default", nid, ig.id) / "infographic.png").read_bytes()
    from PIL import Image
    assert Image.open(io.BytesIO(png)).size == (2000, 1200) and ig.data["overflow"] == 0
    fm = asyncio.run(studio.suggest_reports("default", nid))
    n = len(ai.prompts)
    assert fm[0]["title"] == "Board memo" and asyncio.run(studio.suggest_reports("default", nid)) == fm and len(ai.prompts) == n
    d = ns.detail("default", nid)
    assert len(d["outputs"]) == 7 and all(o["status"] == "ready" for o in d["outputs"])
    assert any(o["cover"] for o in d["outputs"] if o["kind"] == "infographic")
    ev = [e for e in _events(ctx.conversation_id, "notebook_output") if e["status"] == "ready"]
    assert ev and ev[-1]["push"] and ev[-1]["kind"] == "flashcards"
    from app import kaamwatch
    assert kaamwatch.shape("notebook_output", ev[-1])["state"] == "done"


def test_slides_pdf_pptx_and_revise(ai):
    from app.notebook import service as ns
    from app.notebook import studio
    nid = _nb(conv_ctx())
    o = _made(studio.start("default", "", nid, "slides", {"format": "presenter", "length": "short", "style": "kawaii", "audience": "investors"})["id"])
    assert "Presenter slides" in ai.prompts[-1][1] and "investors" in ai.prompts[-1][1]
    assert {"slide-01.png", "slide-02.png", "slide-03.png", "deck.pdf", "deck.pptx"} <= set(o.files)
    d = ns.out_dir("default", nid, o.id)
    from pptx import Presentation
    deck = Presentation(str(d / "deck.pptx"))
    assert len(deck.slides) == 3 and deck.slides[1].notes_slide.notes_text_frame.text == "say it"
    before = (d / "slide-02.png").read_bytes()
    out = asyncio.run(studio.revise_slides("default", nid, o.id, [{"n": 2, "instruction": "title badlo"}]))
    assert out["data"]["slides"][1]["title"] == "Naya title" and out["data"]["revised"] == [2]
    assert (d / "slide-02.png").read_bytes() != before and out["status"] == "ready"
    with pytest.raises(ns.NotebookError, match="Slide number"):
        asyncio.run(studio.revise_slides("default", nid, o.id, [{"n": 9, "instruction": "x"}]))


def _picture(subject, st, custom, aspect="16:9"):
    from PIL import Image
    b = io.BytesIO()
    Image.new("RGB", (640, 360), (120, 80, 200)).save(b, "JPEG")
    return b.getvalue()


def test_video_overview_real_ffmpeg(ai, monkeypatch):
    monkeypatch.setenv("GEMINI_API_KEY", "g-test")
    from app.config import get_settings
    get_settings.cache_clear()
    from app.notebook import media, studio
    from app.notebook import service as ns
    said = []
    monkeypatch.setattr(media, "tts_chunk", lambda text, voices=None, voice="": said.append((text, voice)) or b"\x00\x08" * 24000)
    monkeypatch.setattr(studio, "_picture", _picture)
    nid = _nb(conv_ctx())
    with pytest.raises(ns.NotebookError, match="Custom style"):
        studio.start("default", "", nid, "video", {"style": "custom"})
    o = _made(studio.start("default", "", nid, "video", {"format": "brief", "style": "whiteboard", "voice": "male", "language": "hinglish"})["id"])
    d = ns.out_dir("default", nid, o.id)
    assert {"overview.mp4", "captions.srt", "cover.jpg", "frame-01.jpg", "frame-02.jpg"} <= set(o.files)
    dur = float(subprocess.run(["ffprobe", "-v", "error", "-show_entries", "format=duration", "-of", "csv=p=0",
                                str(d / "overview.mp4")], capture_output=True, text=True).stdout)
    assert 2.5 < dur < 4.5 and o.data["seconds"] >= 2
    assert said[0][1] == "Charon" and "Hinglish" in said[0][0]
    assert "00:00:00,000 -->" in (d / "captions.srt").read_text()


def test_audio_and_join_and_signed_links(ai, monkeypatch):
    monkeypatch.setenv("GEMINI_API_KEY", "g-test")
    from app.config import get_settings
    get_settings.cache_clear()
    from app.notebook import media, studio
    monkeypatch.setattr(media, "tts_chunk", lambda text, voices=None, voice="": b"\x01\x00" * 24000 * 2)
    nid = _nb(conv_ctx())
    o = _made(studio.start("default", "", nid, "audio", {"format": "critique", "length": "longer", "language": "ta"})["id"])
    assert "Critique" in ai.prompts[-1][1] and "Tamil" in ai.prompts[-1][1] and "about 1700 words" in ai.prompts[-1][1]
    assert o.data["seconds"] >= 2 and o.data["audio"] in o.files
    j = asyncio.run(studio.join_audio("default", nid, o.id, "Revenue kitna tha?", 5))
    assert "42 lakh" in j["script"] and j["url"].startswith(f"/api/notebook-files/{o.id}/join-01")
    assert "joined the show live" in ai.prompts[-1][1]
    from app import main
    with TestClient(main.app) as c:
        r = c.get(j["url"])
        assert r.status_code == 200 and len(r.content) > 100
        bad = j["url"].replace("sig=", "sig=0")
        assert c.get(bad).status_code == 404
        exp = int(time.time()) - 5
        from app.pc.service import sign
        name = o.data["audio"]
        assert c.get(f"/api/notebook-files/{o.id}/{name}?exp={exp}&sig={sign(f'nb:{o.id}/{name}', exp)}").status_code == 404
        assert c.get(f"/api/notebook-files/{o.id}/..%2Fx?exp={exp + 999}&sig=x").status_code == 404


def test_discover_and_deep_research(ai, monkeypatch):
    from app import web
    from app.notebook import service as ns
    from app.notebook import studio

    async def search(q):
        return {"answer": "FTTH is growing", "sources": [{"title": "A", "url": "https://example.in/a"},
                                                         {"title": "B", "url": "https://example.in/b"}, {"url": "ftp://x"}]}

    async def page(url):
        return {"url": url, "title": "Page", "text": f"Page text of {url}. Fiber subscribers grew."}
    monkeypatch.setattr(web, "search_any", search)
    monkeypatch.setattr(web, "open_website", page)
    nid = _nb(conv_ctx())
    out = asyncio.run(ns.discover("default", nid, "FTTH India"))
    assert [s["url"] for s in out["sources"]] == ["https://example.in/a", "https://example.in/b"] and out["summary"]
    o = _made(studio.start("default", "", nid, "research", {"query": "FTTH growth India"})["id"])
    assert o.data["sources"][0]["url"] == "https://example.in/a" and o.data["queries"] == ["ftth india growth", "bsnl bharat fiber 2026"]
    assert "Page text of https://example.in/a" in ai.prompts[-1][1]
    src = ns._sources(nid)[-1]
    assert src.kind == "research" and src.title == "Deep research: FTTH growth India"


def test_tools_and_api(ai):
    ctx = conv_ctx()
    nid = _nb(ctx)
    from app.notebook import studio
    out = run("notebook_studio", {"notebook": "FNW Q3", "kind": "flashcards", "count": "fewer", "prompt": "only money"}, ctx)
    assert out["status"] == "making"
    asyncio.run(studio.run(out["output_id"]))
    show = run("notebook_output", {"notebook": "FNW Q3", "output_id": out["output_id"], "action": "show"}, ctx)
    assert show["status"] == "ready" and show["data"]["cards"]
    f = run("notebook_output", {"notebook": "FNW Q3", "output_id": out["output_id"], "action": "file", "name": "flashcards.csv"}, ctx)
    assert f["file_id"] and _events(ctx.conversation_id, "pc_file")
    s = run("notebook_sources", {"notebook": "FNW Q3", "action": "off", "sources": ["3"]}, ctx)
    assert s["sources"][2]["enabled"] is False
    cfg = run("notebook_chat_config", {"notebook": "FNW Q3", "goal": "learning"}, ctx)
    assert cfg["chat_config"]["goal"] == "learning"
    n = run("notebook_note", {"notebook": "FNW Q3", "action": "create", "content": "apni baat"}, ctx)
    assert run("notebook_note", {"notebook": "FNW Q3", "action": "read", "note_id": n["id"]}, ctx)["content"] == "apni baat"
    opened = run("notebook_open", {"notebook": "FNW Q3"}, ctx)
    assert opened["studio"][0]["kind"] == "flashcards" and "content" not in opened["notes"][0]
    from app import main
    with TestClient(main.app) as c:
        cat = c.get("/api/notebooks-catalogue").json()
        assert "anime" in cat["video"]["styles"] and cat["languages"]["mr"] == "मराठी" and len(cat["kinds"]) == 9
        r = c.post(f"/api/notebooks/{nid}/studio", json={"kind": "quiz", "options": {"count": "more"}})
        assert r.status_code == 200 and r.json()["status"] == "queued"
        oid = r.json()["id"]
        for _ in range(100):
            v = c.get(f"/api/notebooks/{nid}/studio/{oid}").json()
            if v["status"] in ("ready", "failed"):
                break
            time.sleep(0.1)
        assert v["status"] == "ready" and v["data"]["questions"]
        assert c.post(f"/api/notebooks/{nid}/ask", json={"question": "revenue?"}).json()["followups"]
        assert len(c.get(f"/api/notebooks/{nid}/chat").json()["messages"]) == 2
        assert c.put(f"/api/notebooks/{nid}/chat/config", json={"length": "longer"}).json()["chat_config"]["length"] == "longer"
        assert c.post(f"/api/notebooks/{nid}/sources/select", json={"ids": None, "enabled": True}).json()["sources"][2]["enabled"]
        assert c.get(f"/api/notebooks/{nid}/sources/1").json()["passages"]
        up = c.post(f"/api/notebooks/{nid}/upload", files={"file": ("ch1.md", io.BytesIO(b"# Chapter 1\nOLT setup steps."), "text/markdown")},
                    data={"title": "Chapter one"}).json()
        assert up["sources"][-1]["title"] == "Chapter one"
        cp = c.post(f"/api/notebooks/{nid}/copy").json()
        assert cp["title"] == "FNW Q3 (copy)" and len(cp["sources"]) == 4
        assert c.patch(f"/api/notebooks/{nid}", json={"title": "FNW Q3 2026", "emoji": "📡"}).json()["emoji"] == "📡"
        assert c.delete(f"/api/notebooks/{nid}/studio/{oid}").json()["deleted"]
        from app.notebook import service as ns
        folder = ns.root_dir() / "default" / nid
        assert folder.exists()
        assert c.delete(f"/api/notebooks/{nid}").json()["deleted"] == "FNW Q3 2026"
        assert not folder.exists()
