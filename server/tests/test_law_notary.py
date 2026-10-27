"""Notary Sahayak: templates in English / Hindi / Marathi filled only from facts (blanks otherwise), stamp only from HER rate
table, IDs only as Pehchaan Locker references, final print ONLY on her tap (hash), Form XV register + fee receipt with Rule 10
caps, desk mode, apostille steps, Rentdesk link, translation via Lekhak, DPDP forget, reminders, tools, API."""
import asyncio
import io
from datetime import date, datetime, timezone

import pytest
from fastapi import FastAPI
from fastapi.testclient import TestClient

import app.notary.models  # noqa: F401  (tables for the env fixture's init_db)
from tests.test_life import conv_ctx

T = "default"
DAY = date(2026, 10, 23)


def _aadhaar() -> str:
    from app.pii import verhoeff_ok
    for x in range(10):
        n = f"23456789012{x}"
        if verhoeff_ok(n):
            return n
    raise AssertionError


@pytest.fixture
def nt(env, monkeypatch):
    from app import indicpdf
    from app.notary import service as S
    events = []
    monkeypatch.setattr(S, "_publish", lambda tenant, kind, payload, conv="": events.append((kind, payload)))
    monkeypatch.setattr(S, "today", lambda t: DAY)
    monkeypatch.setattr(S, "local_now", lambda t: datetime(2026, 10, 23, 11, 15))
    monkeypatch.setattr(S, "_ai_json", lambda system, prompt: {})
    monkeypatch.setattr(indicpdf, "_chromium", lambda html: None)      # fast ReportLab fallback in tests
    S.EVENTS = events
    return S


def kinds(S):
    return [k for k, _ in S.EVENTS]


PARTY = {"name": "Asha Patil", "relation": "wo", "relative": "Ramesh Patil", "age": "34", "occupation": "Tailor",
         "address": "Plot 12, Khanapur Ves, Akot 444101", "gender": "f"}


def test_rules_templates_fees_stamp_apostille():
    from app.notary import rules as R
    for k in R.KINDS:
        for lg in R.LANGS:
            b = R.build(k, lg, PARTY, {})
            assert b["title"] == R.TITLES[k][lg] and "{" not in b["text"]
            if k != "rent_summary":
                assert "Asha Patil" in b["text"]
    full = R.build("aff_address", "hi", PARTY, {"since": "2015", "purpose": "gas connection"})
    assert full["blanks"] == [] and "करती हूँ" in full["text"] and "रही हूँ" in full["text"] and "पत्नी Ramesh Patil" in full["text"]
    male = R.build("aff_address", "mr", {**PARTY, "gender": "m"}, {"since": "2015", "purpose": "gas"})
    assert "देतो" in male["text"] and "पूर्वीचा" not in male["text"]          # optional previous address left out, not guessed
    missing = R.build("aff_lost_document", "en", PARTY, {"document": "10th marksheet"})
    assert {"Kisne diya tha", "Kab khoya", "Kahan khoya", "Police / lost report no."} <= set(missing["blanks"])
    assert "[____ Kab khoya]" in missing["text"]
    inc = R.build("aff_income", "en", PARTY, {"annual_income": "1.2 lakh", "source": "tailoring", "year": "2025-26", "purpose": "scholarship"})
    assert "₹1,20,000 (Rupees one lakh twenty thousand only)" in inc["text"]
    poa = R.build("poa_general", "en", PARTY, {"attorney_name": "Ramesh Patil", "attorney_address": "Akot",
                                               "powers": "Operate my SBI savings account\nSign and file my papers"})
    assert "1. Operate my SBI savings account." in poa["text"] and "2. Sign and file my papers." in poa["text"] and "WITNESSES" in poa["text"]
    # Rule 10 caps (2014 table) + her own update
    assert R.fee_check("affidavit", 35)["ok"] and not R.fee_check("affidavit", 50)["ok"]
    assert R.fee_check("affidavit", 50, caps={"affidavit": 60})["ok"]
    assert R.fee_check("true_copy", 30, pages=3)["ok"] and not R.fee_check("true_copy", 40, pages=3)["ok"]
    assert R.fee_check("true_copy", 20, pages=1)["ok"] and R.fee_check("other", 999)["ok"] and not R.fee_check("attest", None)["ok"]
    # stamp: her table only
    assert R.stamp_for({}, "MH", "affidavit")["amount"] is None
    s = R.stamp_for({"MH": {"affidavit": 500}}, "Maharashtra", "affidavit")
    assert s["amount"] == 500 and s["mode"] == "esbtr" and R.stamp_for({"DL": {"affidavit": 10}}, "DL", "affidavit")["mode"] == "e_stamp"
    assert R.stamp_for({"MH": {"self_declaration": 0}}, "MH", "self_declaration")["mode"] == "none"
    assert any("[stamp ₹ ?]" in x for x in R.stamp_steps({"mode": "e_stamp", "amount": None}, ["Asha"], "Affidavit"))
    # apostille: notarised affidavit -> notarise, state authentication, (hague unknown -> check) ...
    st = [x["key"] for x in R.apostille_steps("affidavit", None, "Germany")]
    assert st[:3] == ["notarise", "first_auth", "check_hague"] and "mea" in st and "embassy" in st
    assert [x["key"] for x in R.apostille_steps("commercial", False, "UAE")] == ["first_auth", "mea_attest", "embassy", "collect"]
    assert "embassy" not in [x["key"] for x in R.apostille_steps("educational", True, "USA")]


def test_draft_ids_in_locker_stamp_blank_and_print_only_on_tap(nt):
    S = nt
    aad = _aadhaar()
    r = S.draft(T, "aff_address", "hi", {**PARTY, "aadhaar": aad}, {"since": "2015"}, state="MH")
    v = r["doc"]
    assert v["number"] == "ND-0001" and v["status"] == "draft"
    p = v["parties"][0]
    assert p["id_kind"] == "aadhaar" and p["id_masked"].endswith(aad[-4:]) and p["id_ref"].startswith("{id:")
    with S.session() as s:
        row = s.get(S.NotaryDoc, v["id"])
        assert aad not in row.parties and aad not in row.text
    from app import idvault
    assert idvault.resolve(T, p["id_ref"])[0] == aad                     # the real number lives only in the locker
    assert "Kis kaam ke liye" in v["blanks"] and v["stamp"]["amount"] is None and v["files"].get("draft_pdf")
    blocked = S.print_request(T, "ND-0001")
    assert blocked.get("blocked") and "notary_print" not in kinds(S)
    S.update(T, "ND-0001", facts={"purpose": "gas connection"})
    assert S.print_request(T, "ND-0001").get("blocked")                  # stamp still not in her table
    S.set_stamp(T, "MH", "affidavit", 500)
    card = S.print_request(T, "ND-0001")
    assert card["doc"]["status"] == "waiting" and card["doc"]["hash"] and "notary_print" in kinds(S)
    h = card["doc"]["hash"]
    with pytest.raises(S.NotaryError):
        S.decide_print(T, "ND-0001", True, h)                            # never without her tap
    with pytest.raises(S.NotaryError):
        S.decide_print(T, "ND-0001", True, "x" * 64, by_tap=True)
    # a change after the card voids it
    S.stamp_ref(T, "ND-0001", ref="MH0123456789")
    assert S.get_doc(T, "ND-0001")["status"] == "draft"
    assert S.decide_print(T, "ND-0001", True, h, by_tap=True)["doc"]["status"] == "draft"     # old card does nothing
    card = S.print_request(T, "ND-0001")
    out = S.decide_print(T, "ND-0001", True, card["doc"]["hash"], by_tap=True)
    assert out["doc"]["status"] == "final" and out["doc"]["files"]["final_pdf"] and out["doc"]["files"]["final_docx"]
    with pytest.raises(S.NotaryError):
        S.update(T, "ND-0001", facts={"since": "2016"})                  # final documents are locked: copy instead
    c = S.copy(T, "ND-0001", "en")
    assert c["doc"]["number"] == "ND-0002" and c["doc"]["lang"] == "en" and c["doc"]["parties"][0]["id_ref"] == p["id_ref"]


def test_register_receipt_fee_caps_correction_excel_desk(nt):
    S = nt
    S.save_settings(T, notary_name="Adv. Meghana Ingle", reg_no="N-1234", area="Akola district", state="MH", cop_valid_to="2027-03-01")
    S.set_stamp(T, "MH", "affidavit", 500)
    with pytest.raises(S.NotaryError):
        S.set_fee(T, "affidavit", 100)                                    # above Rule 10(1)(g) cap
    S.set_fee(T, "affidavit", 35)
    d = S.draft(T, "aff_gap_year", "en", PARTY, {"from_year": "2022", "to_year": "2023", "activity": "preparing for exams",
                                                 "purpose": "XYZ College"})["doc"]
    with pytest.raises(S.NotaryError):
        S.notarise(T, d["number"])                                       # not final yet
    card = S.print_request(T, d["number"])
    S.decide_print(T, d["number"], True, card["doc"]["hash"], by_tap=True)
    with pytest.raises(S.NotaryError):
        S.notarise(T, d["number"], fee=200)
    r = S.notarise(T, d["number"], payer="Asha Patil", mode="upi", ref="UTR 4421", charges={"typing": 50, "xerox": 10})
    e, rc = r["entry"], r["receipt"]
    assert e["serial"] == "2026/0001" and e["fee"] == 35 and e["receipt_no"] == rc["number"] == "NR/2627/0001"
    assert e["executants"][0]["name"] == "Asha Patil" and "id_ref" not in e["executants"][0]
    assert rc["total"] == 95 and {x["kind"] for x in rc["lines"]} == {"notarial", "desk"} and r["file"]["file_id"]
    assert S.get_doc(T, d["number"])["status"] == "executed" and "notary_entry" in kinds(S)
    with pytest.raises(S.NotaryError):
        S.notarise(T, d["number"])                                       # never twice
    # a document not made here: executants + contents directly, address required (Form XV col. 4)
    with pytest.raises(S.NotaryError):
        S.notarise(T, act="true_copy", fee=20, executants=[{"name": "Ravi"}], contents="True copy of sale deed")
    r2 = S.notarise(T, act="true_copy", fee=30, pages=3, executants=[{"name": "Ravi Deshmukh", "address": "Akot"}],
                    contents="True copy of sale deed (3 pages)")
    assert r2["entry"]["serial"] == "2026/0002"
    cor = S.correct_entry(T, "2026/0001", "Executant address: Plot 21, not Plot 12")
    assert cor["entry"]["corrects"] == "2026/0001" and cor["entry"]["serial"] == "2026/0003"
    reg = S.register(T, 2026)
    assert reg["total_acts"] == 2 and reg["fees_total"] == 125 and len(reg["entries"]) == 3
    saved = {}
    real_store = S._store

    def grab(tenant, name, data, ctype, conv="", title=""):
        saved[name] = data
        return real_store(tenant, name, data, ctype, conv, title)
    S._store = grab
    try:
        f = S.register_file(T, 2026)
    finally:
        S._store = real_store
    assert f["file"]["file_id"]
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(saved["Notarial_register_2026.xlsx"]))
    assert wb.sheetnames == ["Form XV", "Fee register", "Annual count", "About"]
    assert [c.value for c in wb["Form XV"][1]][:3] == ["Serial No.", "Date", "Name of notarial act"]
    assert wb["Form XV"].max_row == 4 and wb["Form XV"].cell(2, 1).value == "2026/0001"
    rec = S.receipt(T, "Ravi Deshmukh", {"typing": 40})
    assert rec["receipt"]["number"] == "NR/2627/0003" and rec["receipt"]["total"] == 40
    assert S.void_receipt(T, rec["receipt"]["number"], "galti")["receipt"]["void"]
    # desk mode: no notarial register
    S.save_settings(T, mode="desk")
    d2 = S.draft(T, "noc", "en", PARTY, {"to_whom": "Society", "subject": "flat 4", "purpose": "loan"})["doc"]
    S.set_stamp(T, "MH", "noc", 100)
    c2 = S.print_request(T, d2["number"])
    S.decide_print(T, d2["number"], True, c2["doc"]["hash"], by_tap=True)
    with pytest.raises(S.NotaryError):
        S.notarise(T, d2["number"], fee=35)
    assert S.executed(T, d2["number"], "Adv. X")["doc"]["status"] == "executed"


def test_words_understanding_keeps_only_her_facts(nt, monkeypatch):
    S = nt
    monkeypatch.setattr(S, "_ai_json", lambda system, prompt: {
        "party": {"name": "Sunil Gawai", "relation": "so", "relative": "Baban Gawai", "age": "41", "address": "Mumbai"},   # Mumbai invented
        "facts": {"annual_income": "180000", "source": "kheti", "year": "2025-26", "purpose": "college fees concession"}})
    words = ("Sunil Gawai s/o Baban Gawai, 41 saal, Akot ka, income affidavit chahiye, saal ki aay 180000 kheti se, 2025-26, "
             "college fees concession ke liye. Aadhaar " + _aadhaar())
    v = S.draft(T, words=words, lang="en")["doc"]
    assert v["kind"] == "aff_income"
    p = v["parties"][0]
    assert p["name"] == "Sunil Gawai" and "address" not in [k for k, x in p.items() if x] and p["id_kind"] == "aadhaar"
    assert "Pura pata" in v["blanks"]
    full = S.get_doc(T, v["number"])
    assert "₹1,80,000" in full["text"] and _aadhaar() not in full["text"]


def test_rentdesk_link_translate_apostille_forget(nt, monkeypatch):
    S = nt
    from app.rentdesk import service as RD
    monkeypatch.setattr(RD, "tenancy_view", lambda t, ref: {
        "number": "T-0001", "tenant_name": "Kiran Joshi", "start": "2026-11-01", "months": 11, "rent": 12000, "deposit": 50000,
        "kind_label": "Leave and licence", "unit_label": "Flat 4", "unit_view": {"address": "Flat 4, Sai Heights, Akola", "owner_name": "Meera Rao"},
        "person": {"phone": "9876543210"}})
    monkeypatch.setattr(RD, "settings", lambda t: {"landlord_name": ""})
    r = S.from_rentdesk(T, "T-0001", "en")["doc"]
    full = S.get_doc(T, r["number"])
    assert r["kind"] == "rent_summary" and r["source"] == "rentdesk:T-0001"
    assert "Meera Rao" in full["text"] and "Kiran Joshi" in full["text"] and "₹12,000" in full["text"] and "not the agreement" in full["text"]
    from app.lekhak import service as L
    monkeypatch.setattr(L, "translate", lambda tenant, text="", to="", register="formal", conv="", **k: {"doc": {"number": "L-0001"},
                                                                                                    "say": "Gujarati mein taiyaar."})
    tr = S.translate(T, r["number"], "gujarati")
    assert tr["translation"]["number"] == "L-0001" and "certified nahi" in tr["say"]
    assert S.translate(T, r["number"], "marathi")["doc"]["lang"] == "mr"
    ap = S.apostille(T, "", "affidavit", "Germany", None, "Kiran Joshi", doc=r["number"])["apostille"]
    assert ap["number"] == "A-001" and ap["doc_type"] == "affidavit"
    assert ap["steps"][0]["key"] == "notarise"
    u = S.apostille_update(T, "A-001", hague="yes")["apostille"]
    assert "embassy" not in [x["key"] for x in u["steps"]] and u["hague"] == "yes"
    for k in [x["key"] for x in u["steps"]]:
        u = S.apostille_update(T, "A-001", step=k)["apostille"]
    assert u["status"] == "done"
    aad = _aadhaar()
    d = S.draft(T, "aff_general", "en", {**PARTY, "aadhaar": aad}, {"statements": "I am a resident of Akot"})["doc"]
    from app import idvault
    assert idvault.resolve(T, d["parties"][0]["id_ref"])
    f = S.forget(T, d["number"])
    assert "mita di" in f["say"] and S.get_doc(T, d["number"])["forgotten"]
    assert idvault.resolve(T, d["parties"][0]["id_ref"]) is None          # the locker ID went too (no other use)


def test_tick_reminders_and_morning(nt, monkeypatch):
    S = nt
    S.save_settings(T, notary_name="Adv. M", cop_valid_to="2027-03-01", morning_time="09:00")
    S.draft(T, "aff_general", "en", PARTY, {})
    out = S.tick(datetime(2026, 10, 23, 5, 0, tzinfo=timezone.utc))      # 10:30 IST, 129 days left
    assert out["alerts"] == 1 and out["morning"] == 1 and "notary_alert" in kinds(S) and "notary_day" in kinds(S)
    assert S.tick(datetime(2026, 10, 23, 6, 0, tzinfo=timezone.utc)) == {"alerts": 0, "morning": 0}
    S.EVENTS.clear()
    jan = S.tick(datetime(2027, 1, 2, 5, 0, tzinfo=timezone.utc))
    says = [p["say"] for k, p in S.EVENTS if k == "notary_alert"]
    assert any("Form XIV" in x for x in says) and jan["alerts"] >= 1
    d = S.day(T)
    assert d["blanks"] and "Certificate of practice" in " ".join(d["dates"])


def test_tools_owner_only_and_print_is_a_card(nt):
    from app.agent import notary_tools as NT
    from app.agent.tools import Ctx, ToolError
    ctx = conv_ctx()
    S = nt
    S.set_stamp(T, "MH", "noc", 100)
    r = asyncio.run(NT.notary(ctx, "draft", kind="noc", lang="mr", party=PARTY, facts={"to_whom": "सोसायटी", "subject": "flat", "purpose": "loan"}))
    num = r["doc"]["number"]
    c = asyncio.run(NT.notary(ctx, "print", doc=num))
    assert c["doc"]["status"] == "waiting" and "tap" in c["say"]
    staff = Ctx(T, ctx.conversation_id, actor="staff:1")
    with pytest.raises(ToolError):
        asyncio.run(NT.notary(staff, "notarise", doc=num, fee=35))
    with pytest.raises(ToolError):
        asyncio.run(NT.notary(staff, "stamp_rate", state="MH", key="noc", amount=1))
    with pytest.raises(ToolError):
        asyncio.run(NT.notary_dekho(staff, "register"))
    with pytest.raises(ToolError):
        asyncio.run(NT.notary(Ctx(T, ctx.conversation_id, actor="customer:wa:9"), "draft", kind="noc"))
    assert asyncio.run(NT.notary_dekho(staff, "doc", doc=num))["status"] == "waiting"
    k = asyncio.run(NT.notary_dekho(ctx, "kinds"))
    assert "aff_name_change" in k["kinds"] and NT.NOTARY_LABELS["notary"]
    assert {t["name"] for t in NT.NOTARY_TOOLS} == set(NT.NOTARY_HANDLERS)


def test_api_routes(nt):
    from app.notary.api import make_router
    app = FastAPI()
    app.include_router(make_router(lambda: None, lambda: "default"))
    c = TestClient(app)
    assert "aff_income" in c.get("/api/notary/kinds").json()["kinds"]
    assert c.post("/api/notary/stamp_rates", json={"state": "MH", "key": "affidavit", "amount": 500}).status_code == 200
    r = c.post("/api/notary/docs", json={"kind": "aff_address", "lang": "en", "party": PARTY,
                                         "facts": {"since": "2015", "purpose": "passport"}}).json()
    num = r["doc"]["number"]
    card = c.post(f"/api/notary/docs/{num}/print").json()
    assert card["doc"]["status"] == "waiting"
    assert c.post(f"/api/notary/docs/{num}/decide", json={"yes": True, "hash": "bad"}).status_code == 400
    ok = c.post(f"/api/notary/docs/{num}/decide", json={"yes": True, "hash": card["doc"]["hash"]}).json()
    assert ok["doc"]["status"] == "final"
    assert c.post("/api/notary/register", json={"doc": num, "fee": 50}).status_code == 400          # above cap
    reg = c.post("/api/notary/register", json={"doc": num, "fee": 35}).json()
    assert reg["entry"]["serial"].endswith("/0001")
    ov = c.get("/api/notary").json()
    assert ov["register"]["total_acts"] == 1 and ov["never"]
    assert c.post("/api/notary/apostilles", json={"doc_title": "Birth certificate", "doc_type": "personal", "country": "USA",
                                                  "hague": "yes"}).json()["apostille"]["number"] == "A-001"
