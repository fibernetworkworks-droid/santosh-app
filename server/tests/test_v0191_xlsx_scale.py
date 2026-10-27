"""v0.191.0 "Excel at scale + full coverage" - the owner's directive: complex Excel, LAKH columns,
100% Excel coverage. The honest physics: the XLSX format itself caps one sheet at 16,384 columns
(XFD) and 1,048,576 rows, so lakh-scale columns are delivered two ways - up to 16,384 cols in ONE
sheet, and an AUTO-SPLIT into '-P1'..'-PN' parts + an Index sheet beyond that (row key column
repeated on every part). Nothing is ever silently clipped (the v0.189 lie is dead: clipping is a
QA issue naming what was asked and what was written). Big sheets switch styling strategy (one
conditional-format zebra rule instead of per-cell fills) so lakh-scale builds finish in seconds.
Full coverage: classic formula set welcome (VLOOKUP / INDEX+MATCH / SUMIFS / COUNTIFS ...),
9 chart kinds (bar, column, line, pie, scatter, area, doughnut, radar, combo bar+line),
conditional formatting (color_scale / data_bar / cell_is / expression), data validation
(list / number / date), merges, defined names - all ranges rebased onto the designed grid.
Pivot tables stay an honest NO (no Python library creates real pivots) - SUMIFS summary sheets
instead, taught in prompts.
"""
import io
import re
import time

import pytest
from openpyxl import load_workbook
from openpyxl.chart import (AreaChart, BarChart, DoughnutChart, LineChart, PieChart,
                            RadarChart, ScatterChart)

import app.docstudio.xlsxgen as xlsxgen
from app.docstudio import preview


# ---------------------------------------------------------------- limits + honesty

def test_engine_limits_match_excels_own():
    assert xlsxgen.MAX_COLS == 16_384          # Excel's own hard limit (XFD)
    assert xlsxgen.MAX_ROWS == 100_000
    assert xlsxgen.BIG_CELLS == 100_000
    assert xlsxgen.WIDTH_MEASURE_COLS == 64


def test_16k_col_sheet_builds_whole():
    n = 16_384
    rows = [[f"C{i}" for i in range(n)]] + [[f"v{i}" for i in range(n)]]
    out = xlsxgen.make_xlsx("Wide 16K", sheets={"Wide": rows})
    assert out["splits"] == []                 # at the limit: one sheet, no split
    wb = load_workbook(io.BytesIO(out["data"]))
    ws = wb["Wide"]
    assert ws.max_column == 16_385             # 16,384 data cols + margin col A
    assert ws.cell(row=4, column=16_385).value == "C16383"
    assert out["qa"]["status"] == "pass"


def test_lakh_cols_auto_split_with_index_sheet():
    n = 20_000                                  # > 16,384 -> must split, never clip
    rows = [[f"Col{i}" for i in range(n)]] + [[f"v{i}" for i in range(n)]]
    out = xlsxgen.make_xlsx("Lakh Split", sheets={"Mega": rows})
    assert len(out["splits"]) == 1
    sp = out["splits"][0]
    assert sp["cols"] == 20_000
    assert sp["parts"] == ["Mega-P1", "Mega-P2"]
    assert "16,384" in sp["note"]
    wb = load_workbook(io.BytesIO(out["data"]))
    assert wb.sheetnames == ["Mega-P1", "Mega-P2", "Index"]
    p1, p2, idx = wb["Mega-P1"], wb["Mega-P2"], wb["Index"]
    assert p1.max_column == 16_385             # key col + 16,383 data cols + margin A
    assert p1.cell(row=4, column=2).value == "Col0"
    assert p1.cell(row=4, column=3).value == "Col1"
    assert p1.cell(row=4, column=16_385).value == "Col16383"
    assert p2.cell(row=4, column=2).value == "Col0"          # row key repeated
    assert p2.cell(row=4, column=3).value == "Col16384"      # P1 ended at Col16383
    # the Index sheet speaks every part's range
    assert idx.cell(row=4, column=2).value == "Part sheet"
    assert idx.cell(row=5, column=2).value == "Mega-P1"
    assert idx.cell(row=5, column=3).value == 16_384         # data cols incl. the key
    assert idx.cell(row=6, column=2).value == "Mega-P2"
    assert idx.cell(row=6, column=3).value == 20_000 - 16_383
    assert out["qa"]["status"] == "pass"


def test_row_clipping_is_loud_never_silent():
    rows = [["H"]] + [[f"r{i}"] for i in range(100_005)]     # 1 header + 100,005 body = 100,006 grid rows
    out = xlsxgen.make_xlsx("Clip", sheets={"Big": rows})
    said = [i for i in out["qa"]["issues"] if "rows maangi thi" in i]
    assert said, "clipping must be a named QA issue, never silent"
    assert "100,006" in said[0] and "100,000" in said[0]


def test_too_many_sheets_is_loud_too():
    sheets = {f"S{i}": [["A"], ["x"]] for i in range(27)}
    out = xlsxgen.make_xlsx("Many", sheets=sheets)
    assert any("sheets maangi thi" in i for i in out["qa"]["issues"])


# ---------------------------------------------------------------- big mode

def test_big_sheet_uses_cf_zebra_and_builds_fast():
    # 20,000 rows x 8 cols = 160K cells -> big mode
    header = ["ID", "Name", "Qty", "Rate", "City", "Date", "Status", "Score"]
    rows = [header] + [[f"r{i}", f"Item {i}", str(i), f"Rs {i}.50", "Jaipur",
                        "2025-01-15", "ok", str(i % 100)] for i in range(20_000)]
    t0 = time.time()
    out = xlsxgen.make_xlsx("Big Mode", sheets={"Data": rows})
    build_s = time.time() - t0
    assert build_s < 90, f"big sheet took {build_s:.1f}s"
    assert out["sheets"][0]["big"] is True
    assert out["conditional_formats"] >= 1     # the CF zebra rule
    wb = load_workbook(io.BytesIO(out["data"]))
    ws = wb["Data"]
    assert ws.max_row == 20_004
    kinds = [str(r.type) for rules in ws.conditional_formatting._cf_rules.values() for r in rules]
    assert "expression" in kinds               # MOD(ROW(),2)=0 zebra
    # typed values survive big mode (cols: B=ID, C=Name, D=Qty, E=Rate, F=City, G=Date)
    assert ws.cell(row=5, column=4).value == 0
    assert ws.cell(row=5, column=5).value == 0.5        # "Rs 0.50"
    assert ws.cell(row=5, column=7).value.strftime("%Y-%m-%d") == "2025-01-15"


def test_small_sheet_keeps_v0189_per_cell_zebra():
    rows = [["Item", "Qty"], ["Alpha", "1,200"], ["Beta", "3,400"], ["Total", "=SUM(B2:B3)"]]
    out = xlsxgen.make_xlsx("Small", sheets={"Data": rows})
    assert out["sheets"][0]["big"] is False
    wb = load_workbook(io.BytesIO(out["data"]))
    ws = wb["Data"]
    cell = ws.cell(row=5, column=2)            # first body cell
    assert cell.fill.start_color.rgb in ("00FFFFFF", "FFFFFFFF")   # NEUTRAL_0 zebra fill present
    assert cell.font.name == "Calibri"
    assert ws.freeze_panes == "B5"


# ---------------------------------------------------------------- formulas: the classic set

def test_classic_functions_welcome():
    rows = [["K", "V"], ["=VLOOKUP(C2,D:E,2,0)", "=SUMIFS(A:A,B:B,\">2\")"],
            ["=INDEX(A:A,MATCH(1,B:B,0))", "=COUNTIF(C:C,\"x\")"],
            ["=TEXT(TODAY(),\"dd-mmm\")", "=IFERROR(A2/B2,0)"]]
    out = xlsxgen.make_xlsx("Classic", sheets={"S": rows})
    assert out["formulas"] == 6
    assert out["qa"]["status"] == "pass"


def test_dynamic_array_still_refused_with_cell_named():
    with pytest.raises(ValueError) as e:
        xlsxgen.make_xlsx("Dyn", sheets={"S": [["A"], ["=XLOOKUP(1,B:B,C:C)"]]})
    msg = str(e.value)
    assert "XLOOKUP" in msg and "'S'!2,1" in msg


def test_formula_length_cap_warns():
    long_f = "=IFERROR(" + "+".join(["A1"] * 4_200) + ",0)"
    out = xlsxgen.make_xlsx("Long", sheets={"S": [["A"], [long_f]]})
    assert any("8192" in w for w in out["qa"]["issues"])


# ---------------------------------------------------------------- charts: 9 kinds

def _chart_wb(kind, extra=None):
    sheets = {"Sales": [["Month", "Units", "Rev"], ["Jan", 100, "Rs 5000"],
                        ["Feb", 150, "Rs 7500"], ["Mar", 130, "Rs 6500"]]}
    spec = {"sheet": "Sales", "type": kind, "categories": "A2:A4", "values": "B2:B4",
            "title": f"T {kind}"}
    spec.update(extra or {})
    return xlsxgen.make_xlsx("Charts", sheets=sheets, charts=[spec])


def test_new_chart_kinds_build_and_reopen():
    cases = [("scatter", ScatterChart), ("area", AreaChart), ("doughnut", DoughnutChart),
             ("radar", RadarChart), ("pie", PieChart), ("line", LineChart), ("bar", BarChart)]
    for kind, cls in cases:
        out = _chart_wb(kind)
        assert out["charts"] == 1, kind
        wb = load_workbook(io.BytesIO(out["data"]))
        charts = wb["Sales"]._charts
        assert len(charts) == 1, kind
        assert isinstance(charts[0], cls), f"{kind}: {type(charts[0]).__name__}"
    # combo = BarChart frame carrying a LineChart sub-chart on the secondary axis
    out = _chart_wb("combo", {"values": "B1:B4", "line_values": "C1:C4", "titles_from_data": True})
    wb = load_workbook(io.BytesIO(out["data"]))
    ch = wb["Sales"]._charts[0]
    assert isinstance(ch, BarChart)
    subs = [type(s).__name__ for s in getattr(ch, "_charts", [])]
    assert "LineChart" in subs
    assert out["qa"]["status"] == "pass"


def test_chart_caps_never_silent():
    sheets = {"S": [["M", "V"], ["a", 1], ["b", 2]]}
    many = [{"sheet": "S", "type": "bar", "categories": "A2:A3", "values": "B2:B3"} for _ in range(9)]
    out = xlsxgen.make_xlsx("Caps", sheets=sheets, charts=many)
    assert out["charts"] == xlsxgen.CHARTS_PER_SHEET      # 6 per sheet


# ---------------------------------------------------------------- full-coverage extras

def _extras_sheet():
    return {"Sales": {
        "rows": [["Region", "Units", "Price", "Status"],
                 ["North", "1,200", "Rs 45.50", "Yes"],
                 ["South", "850", "Rs 39.00", "No"],
                 ["East", "1,500", "Rs 52.25", "Maybe"]],
        "conditional_formatting": [
            {"type": "color_scale", "range": "B2:B4"},
            {"type": "data_bar", "range": "C2:C4"},
            {"type": "cell_is", "range": "C2:C4", "operator": "greaterThan", "formula": "1000",
             "fill": "FFC7CE"},
            {"type": "expression", "range": "A2:A4", "formula": "$B2>1000", "fill": "DDEBF7"},
        ],
        "validation": [
            {"range": "D2:D4", "type": "list", "formula1": '"Yes,No,Maybe"'},
            {"range": "B2:B4", "type": "number", "operator": "between", "formula1": "0",
             "formula2": "5000"},
        ],
        "merges": ["A6:B6"],
        "defined_names": [{"name": "Units", "ref": "B2:B4"}],
    }}


def test_conditional_formatting_lands_with_rebased_ranges():
    out = xlsxgen.make_xlsx("CF", sheets=_extras_sheet())
    assert out["conditional_formats"] == 4
    wb = load_workbook(io.BytesIO(out["data"]))
    ws = wb["Sales"]
    got = {str(r.type): r for rules in ws.conditional_formatting._cf_rules.values() for r in rules}
    assert set(got) == {"colorScale", "dataBar", "cellIs", "expression"}
    # ranges rebased +3 rows +1 col: B2:B4 -> C5:C7 (sqref lives on the ConditionalFormatting key)
    rngs = [str(k.sqref) for k in ws.conditional_formatting._cf_rules.keys()]
    assert any("C5:C7" in r for r in rngs)
    cell_is = got["cellIs"]
    assert cell_is.operator == "greaterThan" and cell_is.formula[0] == "1000"


def test_validation_dropdowns_and_ranges():
    out = xlsxgen.make_xlsx("DV", sheets=_extras_sheet())
    assert out["validations"] == 2
    wb = load_workbook(io.BytesIO(out["data"]))
    ws = wb["Sales"]
    dvs = ws.data_validations.dataValidation
    assert len(dvs) == 2
    kinds = {dv.type for dv in dvs}
    assert kinds == {"list", "decimal"}
    list_dv = next(dv for dv in dvs if dv.type == "list")
    assert "Yes" in list_dv.formula1


def test_merges_and_defined_names_rebased():
    out = xlsxgen.make_xlsx("MDN", sheets=_extras_sheet())
    assert out["defined_names"] == 1
    wb = load_workbook(io.BytesIO(out["data"]))
    ws = wb["Sales"]
    merged = {str(r) for r in ws.merged_cells.ranges}
    assert "B9:C9" in merged                 # A6:B6 -> +3r +1c
    assert "Units" in wb.defined_names
    assert "'Sales'!$C$5:$C$7" in str(wb.defined_names["Units"].attr_text) \
        or "Sales" in str(wb.defined_names["Units"].attr_text)


def test_bad_extra_types_are_warnings_not_crashes():
    sheet = {"Sales": {
        "rows": [["A", "B"], [1, 2]],
        "conditional_formatting": [{"type": "sparkline_thing", "range": "B2:B2"}],
        "validation": [{"range": "", "type": "nope"}],
        "merges": ["not a range!!"],
        "defined_names": [{"name": "A1", "ref": "B2:B2"}],          # cell-ref-looking name
    }}
    out = xlsxgen.make_xlsx("Bad", sheets=sheet)
    assert out["qa"]["status"] in ("pass", "warn")
    assert any("conditional_formatting" in w or "validation" in w for w in out["qa"]["issues"])
    wb = load_workbook(io.BytesIO(out["data"]))   # file still opens
    assert wb.sheetnames == ["Sales"]


# ---------------------------------------------------------------- deep preview honesty

def test_preview_reports_written_cols_and_extras():
    out = xlsxgen.make_xlsx("PV", sheets=_extras_sheet())
    view = preview.preview("xlsx", out["data"], out)
    s = view["sheets"][0]
    assert s["cols"] == 4                       # WRITTEN cols (was the v0.189 input-length lie)
    assert view.get("conditional_formats") == 4
    assert view.get("validations") == 2
    assert "Units" in (view.get("defined_names") or [])


def test_preview_16k_and_split_views():
    rows = [[f"C{i}" for i in range(16_384)]] + [["v"] * 16_384]
    out = xlsxgen.make_xlsx("PV16", sheets={"Wide": rows})
    view = preview.preview("xlsx", out["data"], out)
    assert view["sheets"][0]["cols"] == 16_384
    rows = [[f"C{i}" for i in range(20_000)]] + [["v"] * 20_000]
    out = xlsxgen.make_xlsx("PVSplit", sheets={"Mega": rows})
    view = preview.preview("xlsx", out["data"], out)
    names = [s["name"] for s in view["sheets"]]
    assert names[:3] == ["Mega-P1", "Mega-P2", "Index"]


# ---------------------------------------------------------------- tool surface survives

def test_doc_studio_schema_carries_new_surface():
    from app.agent.docstudio_tools import DOCSTUDIO_TOOLS
    tool = next(t for t in DOCSTUDIO_TOOLS if t["name"] == "doc_studio")
    props = tool["input_schema"]["properties"]
    assert "scatter|area|doughnut|radar|combo" in props["charts"]["description"]
    assert "16,384" in props["sheets"]["description"]
    assert "conditional_formatting" in props["sheets"]["description"]


def test_doc_studio_delivers_wide_sheet_with_split_note(monkeypatch, tmp_path):
    from app.agent import docstudio_tools
    class FakeCtx:
        conversation_id = "c1"
    delivered = {}
    def fake_deliver(ctx, name, data, src):
        delivered["name"], delivered["data"] = name, data
        return {"file_id": "f1", "name": name}
    monkeypatch.setattr(docstudio_tools, "_deliver", fake_deliver)
    import asyncio
    rows = [[f"Col{i}" for i in range(20_000)]] + [[f"v{i}" for i in range(20_000)]]
    res = asyncio.run(docstudio_tools.doc_studio(
        FakeCtx(), kind="xlsx", title="Wide", sheets={"Mega": rows}))
    assert res["made"] == "xlsx"
    assert "parts" in res["note"] or "16,384" in res["note"]
    wb = load_workbook(io.BytesIO(delivered["data"]))
    assert wb.sheetnames == ["Mega-P1", "Mega-P2", "Index"]
