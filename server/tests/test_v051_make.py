"""v0.51 real Office files, product photo cleanup (stubbed), ID photo service client, meaning-based recall."""
import asyncio
import io

import pytest


def _ctx():
    from app.agent.tools import Ctx
    return Ctx(tenant_id="default", conversation_id="c1")


def test_make_real_office_files(env):
    from app.agent.make_tools import make_document
    from app.mailbox import service as ms
    from pathlib import Path
    d = asyncio.run(make_document(_ctx(), "docx", "FNW Quotation", text="# Rates\n- ONT fitting\n| Item | Price |\n|---|---|\n| ONT | 1200 |\n**Total** 1200"))
    from docx import Document
    doc = Document(io.BytesIO(Path(ms._file_row("default", d["file_id"]).stored_path).read_bytes()))
    assert any("ONT" in c.text for t in doc.tables for row in t.rows for c in row.cells)
    x = asyncio.run(make_document(_ctx(), "xlsx", "Bills", sheets={"Oct": [["Item", "Amt"], ["Bijli", 900], ["Net", 499], ["Total", "=SUM(B2:B3)"]]}))
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(Path(ms._file_row("default", x["file_id"]).stored_path).read_bytes()))
    assert wb["Oct"]["B4"].value == "=SUM(B2:B3)"
    p = asyncio.run(make_document(_ctx(), "pptx", "FNW Pitch", slides=[{"title": "Akola", "bullets": ["500 homes"], "notes": "say hi"}]))
    from pptx import Presentation
    pr = Presentation(io.BytesIO(Path(ms._file_row("default", p["file_id"]).stored_path).read_bytes()))
    assert len(pr.slides) == 2 and p["name"].endswith(".pptx")


def test_photo_clean_and_id_photo(env, monkeypatch):
    from PIL import Image
    from app import imagekit
    from app.agent.make_tools import id_photo, photo_clean
    from app.agent.tools import ToolError
    from app.mailbox import service as ms
    b = io.BytesIO()
    Image.new("RGB", (64, 64), (200, 10, 10)).save(b, "JPEG")
    fid = ms.store_bytes("default", "saree.jpg", b.getvalue(), "mail").id
    monkeypatch.setattr(imagekit, "clean_background", lambda data, color="white": b"\xff\xd8clean")
    out = asyncio.run(photo_clean(_ctx(), fid))
    assert out["name"].endswith(".jpg")
    with pytest.raises(ToolError) as e:
        asyncio.run(id_photo(_ctx(), fid))
    assert "HIVISION_URL" in str(e.value)
    monkeypatch.setattr(imagekit, "id_photo", lambda data, size="", color="": {"photo": b"\xff\xd8one", "sheet": b"\xff\xd8sheet"})
    r = asyncio.run(id_photo(_ctx(), fid, "us_visa"))
    assert r["photo"]["file_id"] and r["print_sheet"]["file_id"]
    assert imagekit.SIZES["india_passport"] == (531, 413)


def test_memory_search_by_meaning_or_words(env):
    from app.agent.life_tools import remember
    from app.agent.make_tools import memory_search
    for f in ["Riya ki class teacher Mrs Rao hai", "Gaadi ka insurance 12 March ko khatam", "Papa ko BP ki dawa subah"]:
        asyncio.run(remember(_ctx(), f))
    out = asyncio.run(memory_search(_ctx(), "gaadi insurance kab"))
    assert out["found"] and "insurance" in out["found"][0]["text"]
    assert out["searched"] == 3


def test_tools_in_chat_and_voice(env):
    from app.agent.tools import HANDLERS
    from app.voice import VOICE_TOOL_NAMES
    for t in ("make_document", "photo_clean", "id_photo", "memory_search"):
        assert t in HANDLERS and t in VOICE_TOOL_NAMES
