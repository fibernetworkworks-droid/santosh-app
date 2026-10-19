"""Build app/bazar/agents_directory.json.gz (Santosh Agent Army) from the open-source AI agents master sheet.

    python tools/build_agent_directory.py path/to/Open_Source_AI_Agents_All_Fields.xlsx

Keeps every row of the 'All agents' sheet (all 18 columns) plus the sheet's own field table and notes.
"""
import gzip
import json
import sys
from pathlib import Path

import openpyxl

COLS = ["n", "name", "link", "field", "also", "does", "kind", "found_in", "lang", "stars", "npm_dl", "trending_days",
        "last_release", "licence", "lists", "packages", "score", "tier"]


def clean(v):
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        return int(v)
    return v if isinstance(v, (int, float)) else str(v).strip()


def main(path: str) -> None:
    wb = openpyxl.load_workbook(path, read_only=True)
    rows = list(wb["All agents"].iter_rows(values_only=True))
    header, body = rows[0], rows[1:]
    assert len(header) == len(COLS), header
    agents = []
    for r in body:
        rec = [clean(v) for v in r]
        rec[17] = (rec[17] or "").split(" ")[0]          # "A - proven" -> "A"
        agents.append(rec)
    fields = [dict(zip(["field", "agents", "tier_a", "tier_b", "best"], [clean(v) for v in r]))
              for r in list(wb["Fields"].iter_rows(values_only=True))[1:] if r and r[0]]
    notes = {str(r[0]): clean(r[1]) for r in wb["Summary"].iter_rows(values_only=True) if r and r[0] and r[1] is not None}
    curated = sum(1 for r in list(wb["Curated-list agents"].iter_rows(values_only=True))[1:] if r and r[1])
    out = {"columns": COLS, "agents": agents, "fields": fields, "notes": notes, "curated": curated,
           "source": Path(path).name}
    dest = Path(__file__).resolve().parents[1] / "app" / "bazar" / "agents_directory.json.gz"
    with gzip.open(dest, "wt", encoding="utf-8", compresslevel=9) as fh:
        json.dump(out, fh, ensure_ascii=False, separators=(",", ":"))
    print(f"{len(agents)} agents, {len(fields)} fields -> {dest} ({dest.stat().st_size // 1024} KB)")


if __name__ == "__main__":
    main(sys.argv[1])
