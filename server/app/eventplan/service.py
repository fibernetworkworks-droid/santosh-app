"""Event Sahayak: Hoola as an event planner's coworker (corporate events, conferences, launches, birthdays, college fests,
exhibitions), end to end.

  1. Client brief -> event file (date, time, venue, headcount, client's budget, facts: sound / food / music / alcohol /
     pandal / ticketed ...) -> a task plan counted back from the event date + a permissions checklist from those facts.
  2. Budget by heads with HER vendors' quotes (never invented: a missing cost is a blank), her margin per head or a fee on
     cost; vendor list, quote comparison per head (cheapest / GST included? / advance / cancellation / what is missing).
  3. Vendor booking cards (+ advance), client quotation with GST (CGST + SGST inside her state, IGST outside; pure-agent
     reimbursements outside GST), run-of-show / cue sheet minute by minute with clash / owner / Noise Rules checks, guest list +
     RSVP from her / guests' words, invite drafts, day-of timeline, post-event report + client invoice draft.

Waits for HER tap (only the app API sets by_tap=True; tools only make the card):
  * any vendor booking / advance (`eventplan_booking`; fingerprint once a phone key is active when an advance goes),
  * the client quotation / invoice going out (`eventplan_doc`),
  * invites and any message to guests / vendors / the client (`eventplan_message`: the WhatsApp / SMS links appear after it),
  * any other spend (`eventplan_spend`; fingerprint once a phone key is active).
Hoola never pays, never collects, never sends a message by itself."""
from __future__ import annotations

import csv
import io
import json
import logging
import re
import secrets
import threading
from datetime import date, datetime, timedelta, timezone
from typing import Any
from urllib.parse import quote as urlq

from sqlalchemy import select

from ..db import now, session
from . import rules as R
from .models import (
    EpAlert,
    EpBooking,
    EpDoc,
    EpEvent,
    EpGuest,
    EpMessage,
    EpSettings,
    EpSpend,
    EpTask,
    EpVendor,
    EpVendorQuote,
)

log = logging.getLogger("santosh.eventplan")
_LOCK = threading.RLock()
STATUS = {"enquiry": "Enquiry", "planning": "Planning", "confirmed": "Pakka", "live": "Aaj event", "done": "Ho gaya",
          "cancelled": "Radd"}
MAX_GUESTS = 3000


class EventError(Exception):
    pass


# ---------------------------------------------------------------- helpers
def _j(s: str | None, default):
    try:
        v = json.loads(s or "")
        return v if isinstance(v, type(default)) else default
    except (TypeError, ValueError):
        return default


def _dump(v) -> str:
    return json.dumps(v, ensure_ascii=False, default=str)


def _aware(d: datetime | None) -> datetime | None:
    if d is None:
        return None
    return d if d.tzinfo else d.replace(tzinfo=timezone.utc)


def clock() -> datetime:
    return now()


def _tz(tenant: str):
    from ..life import tz_of
    return tz_of(tenant)


def today(tenant: str) -> date:
    return clock().astimezone(_tz(tenant)).date()


def _date(v) -> date | None:
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    if not v:
        return None
    s = str(v).strip()[:10]
    for f in ("%Y-%m-%d", "%d-%m-%Y", "%d/%m/%Y", "%d.%m.%Y", "%d-%m-%y", "%d/%m/%y"):
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            pass
    return None


def _iso(d: date | None) -> str:
    return d.isoformat() if d else ""


def _rs(v: float | None) -> str:
    if v is None:
        return "[daam?]"
    try:
        from ..india.money import inr
        return inr(v)
    except Exception:
        return f"₹{v:,.0f}"


def _pdf_rs(v: float | None) -> str:
    return "[ ? ]" if v is None else "Rs. " + f"{v:,.2f}"


def _words(v: float) -> str:
    try:
        from ..india.money import rupees_words
        return rupees_words(v, "en")
    except Exception:
        return ""


def _publish(tenant: str, kind: str, payload: dict, conv: str = "") -> None:
    from ..events import bus
    try:
        if not conv:
            from ..net.service import home_conv
            conv = home_conv(tenant)
        bus.publish(conv, kind, payload)
    except Exception:
        log.exception("eventplan publish failed")


def is_owner(actor: str) -> bool:
    a = actor or "owner"
    return a == "owner" or a.startswith("approved:")


def actor_name(tenant: str, actor: str) -> str:
    if is_owner(actor):
        return "Owner"
    try:
        from ..employee.service import actor_label
        return actor_label(tenant, actor)
    except Exception:
        return str(actor)[:40]


def _st(s, tenant: str) -> EpSettings:
    r = s.get(EpSettings, tenant)
    if r is None:
        r = EpSettings(tenant_id=tenant)
        s.add(r)
        s.flush()
    return r


def _load_file(tenant: str, file_id: str) -> tuple[str, bytes]:
    from ..mailbox import service as ms
    try:
        name, _m, data = ms.load_files(tenant, [{"type": "file", "file_id": file_id}])[0]
        return name, data
    except Exception as e:
        raise EventError(f"File {file_id} nahi mili: {e}")


def _store(tenant: str, name: str, data: bytes, mime: str, conv: str = "", title: str = "") -> dict:
    from .. import chatfiles
    f = chatfiles.store(tenant, name, data, mime)
    _publish(tenant, "pc_file", {**f, "source": "eventplan", "title": title or name}, conv)
    return f


def _wa(phone: str, text: str) -> str:
    d = R.phone(phone)
    return f"https://wa.me/91{d}?text={urlq(text)}" if d else ""


def _sms(phone: str, text: str) -> str:
    d = R.phone(phone)
    return f"sms:+91{d}?body={urlq(text)}" if d else ""


def _upi_link(upi: str, name: str, amount: float, note: str) -> str:
    if not upi or amount <= 0:
        return ""
    return f"upi://pay?pa={urlq(upi)}&pn={urlq(name)}&am={amount:.2f}&cu=INR&tn={urlq(note[:40])}"


def _her_upi(tenant: str) -> tuple[str, str]:
    with session() as s:
        st = _st(s, tenant)
        upi, nm = st.upi_id, st.firm
    if not upi:
        try:
            from ..employee import service as emp
            p = emp.profile(tenant) or {}
            upi = p.get("upi_id") or ""
            nm = nm or p.get("upi_name") or p.get("business_name") or ""
        except Exception:
            upi = ""
    return upi, nm


def _norm(h) -> str:
    return re.sub(r"[^a-z0-9]", "", str(h or "").lower())


def _rows(name: str, data: bytes) -> list[list]:
    if name.lower().endswith((".xlsx", ".xlsm")):
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        return [list(r) for r in wb.worksheets[0].iter_rows(values_only=True)]
    text = data.decode("utf-8-sig", errors="replace")
    return [r for r in csv.reader(io.StringIO(text))]


def _fy(d: date) -> str:
    y = d.year if d.month >= 4 else d.year - 1
    return f"{y % 100:02d}{(y + 1) % 100:02d}"


# ---------------------------------------------------------------- settings
_SET_KEYS = ("firm", "gstin", "state_code", "address", "phone", "email", "upi_id", "margin_pct", "head_margins", "fee_mode",
             "fee_pct", "gst_pct", "sac", "validity_days", "terms", "sound_till", "morning_time", "team", "rules")


def _settings_view(r: EpSettings) -> dict:
    return {"firm": r.firm, "gstin": r.gstin, "state_code": r.state_code, "address": r.address, "phone": r.phone, "email": r.email,
            "upi_id": r.upi_id, "margin_pct": r.margin_pct, "head_margins": _j(r.head_margins, {}), "fee_mode": r.fee_mode,
            "fee_pct": r.fee_pct, "gst_pct": r.gst_pct, "sac": r.sac, "validity_days": r.validity_days, "terms": r.terms,
            "sound_till": r.sound_till, "morning_time": r.morning_time, "team": _j(r.team, []), "rules": r.rules,
            "heads": {k: v["label"] for k, v in R.HEADS.items()}, "kinds": {k: v["label"] for k, v in R.KINDS.items()},
            "permits": {k: v["label"] for k, v in R.PERMITS.items()}}


def settings(tenant: str) -> dict:
    with session() as s:
        return _settings_view(_st(s, tenant))


def save_settings(tenant: str, **kw: Any) -> dict:
    with _LOCK, session() as s:
        r = _st(s, tenant)
        for k in _SET_KEYS:
            if k not in kw or kw[k] is None:
                continue
            v = kw[k]
            if k == "gstin":
                g = str(v).strip().upper()
                if g:
                    from ..india.money import gstin_ok
                    if not gstin_ok(g):
                        raise EventError("GSTIN sahi nahi lagta (check digit / format).")
                    r.state_code = g[:2]
                r.gstin = g
            elif k in ("margin_pct", "fee_pct", "gst_pct"):
                n = R.num(v)
                if n is None or n < 0 or n > (40 if k == "gst_pct" else 300):
                    raise EventError(f"{k} sahi number nahi.")
                setattr(r, k, n)
            elif k == "head_margins":
                if not isinstance(v, dict):
                    raise EventError("head_margins = {head: %}.")
                cur = _j(r.head_margins, {})
                for hk, hv in v.items():
                    hk = hk if hk in R.HEADS else R.head_of(hk)
                    n = R.num(hv)
                    if hv in (None, ""):
                        cur.pop(hk, None)
                    elif n is None or n < 0 or n > 300:
                        raise EventError(f"{hk} ka margin sahi nahi.")
                    else:
                        cur[hk] = n
                r.head_margins = _dump(cur)
            elif k == "fee_mode":
                if v not in ("markup", "fee", "none"):
                    raise EventError("fee_mode = markup | fee | none")
                r.fee_mode = v
            elif k == "validity_days":
                r.validity_days = max(1, min(int(R.num(v) or 10), 90))
            elif k in ("sound_till", "morning_time"):
                t = R.hhmm(str(v))
                if not t:
                    raise EventError(f"{k} HH:MM mein likhiye.")
                if k == "sound_till" and not ("22:00" <= t <= "23:59" or t == "00:00"):
                    raise EventError("Sound ki seema 22:00 se raat 12 tak hi (Noise Rules 2000; 22:00 ke baad sirf state ki chhoot se).")
                setattr(r, k, t)
            elif k == "team":
                if not isinstance(v, list):
                    raise EventError("team = [{name, phone, role}]")
                r.team = _dump([{"name": R.clean(x.get("name"), 60), "phone": R.phone(x.get("phone")), "role": R.clean(x.get("role"), 40)}
                                for x in v if isinstance(x, dict) and str(x.get("name") or "").strip()])
            elif k == "phone":
                r.phone = R.phone(v) or str(v)[:20]
            elif k == "state_code":
                r.state_code = str(v).strip()[:2]
            else:
                setattr(r, k, R.clean(v, 2000) if k in ("terms", "rules", "address") else R.clean(v, 160))
        view = _settings_view(r)
    return {"settings": view, "say": "Niyam save ho gaye."}


# ---------------------------------------------------------------- vendors
def _vendor_view(v: EpVendor) -> dict:
    return {"id": v.id, "name": v.name, "head": v.head, "head_label": R.HEADS.get(v.head, R.HEADS["misc"])["label"], "phone": v.phone,
            "upi_id": v.upi_id, "gstin": v.gstin, "fssai": v.fssai, "notes": v.notes}


def vendor_save(tenant: str, name: str, head: str = "", phone: str = "", upi_id: str = "", gstin: str = "", fssai: str = "",
                notes: str = "", vendor_id: str = "") -> dict:
    name = R.clean(name, 120)
    with _LOCK, session() as s:
        v = s.get(EpVendor, vendor_id) if vendor_id else None
        if v is None and name:
            v = s.scalars(select(EpVendor).where(EpVendor.tenant_id == tenant, EpVendor.name == name)).first()
        if v is None:
            if not name:
                raise EventError("Vendor ka naam chahiye.")
            v = EpVendor(tenant_id=tenant, name=name)
            s.add(v)
        elif v.tenant_id != tenant:
            raise EventError("Vendor nahi mila.")
        if name:
            v.name = name
        if head:
            v.head = head if head in R.HEADS else R.head_of(head)
        elif not v.head or v.head == "misc":
            v.head = R.head_of(name + " " + notes)
        if phone:
            p = R.phone(phone)
            if not p:
                raise EventError("Phone 10 ankon ka mobile number nahi lagta.")
            v.phone = p
        if upi_id:
            u = str(upi_id).strip()
            if not re.fullmatch(r"[\w.\-]{2,64}@[A-Za-z]{2,32}", u):
                raise EventError("UPI ID sahi nahi lagta (naam@bank).")
            v.upi_id = u
        if gstin:
            from ..india.money import gstin_ok
            g = gstin.strip().upper()
            if not gstin_ok(g):
                raise EventError("Vendor ka GSTIN sahi nahi lagta.")
            v.gstin = g
        if fssai:
            f = re.sub(r"\D", "", fssai)
            if len(f) != 14:
                raise EventError("FSSAI number 14 ankon ka hota hai.")
            v.fssai = f
        if notes:
            v.notes = R.clean(notes, 500)
        s.flush()
        view = _vendor_view(v)
    return {"vendor": view, "say": f"Vendor {view['name']} ({view['head_label']}) save."}


def vendors(tenant: str, head: str = "") -> list[dict]:
    with session() as s:
        q = select(EpVendor).where(EpVendor.tenant_id == tenant, EpVendor.active.is_(True))
        if head:
            q = q.where(EpVendor.head == head)
        return [_vendor_view(v) for v in s.scalars(q.order_by(EpVendor.name)).all()]


def _vendor(s, tenant: str, ref: str) -> EpVendor:
    v = s.get(EpVendor, ref) if ref else None
    if v is None and ref:
        rows = s.scalars(select(EpVendor).where(EpVendor.tenant_id == tenant)).all()
        low = ref.strip().lower()
        v = next((x for x in rows if x.name.lower() == low), None) or next((x for x in rows if low in x.name.lower()), None)
    if v is None or v.tenant_id != tenant:
        raise EventError(f"Vendor '{ref}' nahi mila; pehle vendor jodiye.")
    return v


# ---------------------------------------------------------------- events
_EV_KEYS = ("kind", "client", "client_phone", "client_email", "client_gstin", "client_state", "day", "start", "end", "venue", "city",
            "headcount", "budget", "brief", "facts", "status")
_FACT_KEYS = ("venue_kind", "sound", "food", "music", "alcohol", "pandal", "stage", "fireworks", "drone", "road", "ticketed")


def _ev(s, tenant: str, ref: str) -> EpEvent:
    e = s.get(EpEvent, ref) if ref else None
    if e is None and ref:
        rows = s.scalars(select(EpEvent).where(EpEvent.tenant_id == tenant)).all()
        low = ref.strip().lower()
        e = next((x for x in rows if x.number.lower() == low or x.name.lower() == low), None) or \
            next((x for x in rows if low in x.name.lower()), None)
    if e is None and not ref:
        rows = s.scalars(select(EpEvent).where(EpEvent.tenant_id == tenant, EpEvent.status.notin_(("done", "cancelled")))
                         .order_by(EpEvent.created_at.desc())).all()
        if len(rows) == 1:
            e = rows[0]
    if e is None or e.tenant_id != tenant:
        raise EventError("Event nahi mila (number / naam bataiye).")
    return e


def _log(e: EpEvent, text: str, by: str = "") -> None:
    lg = _j(e.log, [])
    lg.append({"at": now().isoformat(timespec="minutes"), "by": by, "text": text[:200]})
    e.log = _dump(lg[-80:])


def _ev_view(s, e: EpEvent, full: bool = False) -> dict:
    v = {"id": e.id, "number": e.number, "name": e.name, "kind": e.kind, "kind_label": R.KINDS.get(e.kind, R.KINDS["other"])["label"],
         "client": e.client, "client_phone": e.client_phone, "client_email": e.client_email, "client_gstin": e.client_gstin,
         "client_state": e.client_state, "day": e.day, "start": e.start, "end": e.end, "venue": e.venue, "city": e.city,
         "headcount": e.headcount or None, "budget": e.budget, "budget_text": _rs(e.budget) if e.budget is not None else "",
         "facts": _j(e.facts, {}), "status": e.status, "status_label": STATUS.get(e.status, e.status), "forgotten": e.forgotten}
    if full:
        v["brief"] = e.brief
        v["log"] = _j(e.log, [])[-20:]
        v["client_paid"] = _j(e.client_paid, [])
        v["report"] = _j(e.report, {})
    return v


def _refresh_permits(s, tenant: str, e: EpEvent) -> None:
    facts = {**_j(e.facts, {}), "headcount": e.headcount}
    cur = {p["key"]: p for p in _j(e.permits, [])}
    for k in R.permits_for(facts):
        if k not in cur:
            cur[k] = {"key": k, "label": R.PERMITS[k]["label"], "status": "pending", "ref": "", "note": "", "custom": False}
    e.permits = _dump(list(cur.values()))


def _plan_tasks(s, tenant: str, e: EpEvent) -> None:
    d = _date(e.day)
    if not d:
        return
    old = {t.what: t for t in s.scalars(select(EpTask).where(EpTask.event_id == e.id, EpTask.auto.is_(True))).all()}
    for t in R.task_plan(e.kind, d, today(tenant)):
        o = old.get(t["what"])
        if o is None:
            s.add(EpTask(tenant_id=tenant, event_id=e.id, what=t["what"], area=t["area"], due=t["due"], auto=True))
        elif o.status == "open":
            o.due = t["due"]


def event_save(tenant: str, name: str = "", event: str = "", actor: str = "owner", conv: str = "", **kw: Any) -> dict:
    with _LOCK, session() as s:
        if event:
            e = _ev(s, tenant, event)
        else:
            if not R.clean(name, 160):
                raise EventError("Event ka naam chahiye.")
            st = _st(s, tenant)
            st.ev_seq += 1
            e = EpEvent(tenant_id=tenant, number=f"E-{st.ev_seq:04d}", name=R.clean(name, 160), conv=conv,
                        kind=R.kind_of(name + " " + str(kw.get("brief") or "")))
            s.add(e)
            s.flush()
        old_day, notes = e.day, []
        if name and event:
            e.name = R.clean(name, 160)
        for k in _EV_KEYS:
            if k not in kw or kw[k] is None:
                continue
            v = kw[k]
            if k == "kind":
                e.kind = v if v in R.KINDS else R.kind_of(str(v))
            elif k == "day":
                d = _date(v)
                if not d:
                    raise EventError("Tareekh samajh nahi aayi (DD-MM-YYYY).")
                if d < today(tenant) and e.status not in ("done", "cancelled"):
                    raise EventError("Event ki tareekh beet chuki hai.")
                e.day = d.isoformat()
            elif k in ("start", "end"):
                t = R.hhmm(str(v))
                if not t:
                    raise EventError(f"{k} samay HH:MM mein.")
                setattr(e, k, t)
            elif k == "headcount":
                n = R.num(v)
                if n is None or n < 0 or n > 200000:
                    raise EventError("Headcount sahi number nahi.")
                e.headcount = int(n)
            elif k == "budget":
                e.budget = R.num(v)
            elif k == "client_phone":
                e.client_phone = R.phone(v) or ""
                if v and not e.client_phone:
                    notes.append("client ka phone 10 ankon ka mobile nahi lagta, save nahi kiya")
            elif k == "client_gstin":
                g = str(v).strip().upper()
                if g:
                    from ..india.money import gstin_ok
                    if not gstin_ok(g):
                        raise EventError("Client ka GSTIN sahi nahi lagta.")
                    e.client_state = g[:2]
                e.client_gstin = g
            elif k == "client_state":
                e.client_state = str(v).strip()[:2]
            elif k == "facts":
                if not isinstance(v, dict):
                    raise EventError("facts = {venue_kind, sound, food, music, alcohol, pandal, stage, fireworks, drone, road, ticketed}")
                f = _j(e.facts, {})
                for fk, fv in v.items():
                    if fk in _FACT_KEYS:
                        f[fk] = (str(fv) if fk == "venue_kind" else bool(fv))
                e.facts = _dump(f)
            elif k == "status":
                if v not in STATUS:
                    raise EventError("status = " + " | ".join(STATUS))
                if v == "confirmed" and not is_owner(actor):
                    raise EventError("Event pakka sirf owner karti hain.")
                e.status = v
            elif k == "brief":
                e.brief = R.clean(v, 4000)
            else:
                setattr(e, k, R.clean(v, 400 if k == "venue" else 120))
        if e.end and e.start and e.end < e.start:
            notes.append("event aadhi raat ke baad tak (end < start) maana")
        _refresh_permits(s, tenant, e)
        if e.day and e.day != old_day:
            _plan_tasks(s, tenant, e)
        _log(e, "event badla" if event else "event bana", actor_name(tenant, actor))
        view = _ev_view(s, e, True)
    say = f"{view['number']} {view['name']} save." + (" " + "; ".join(notes) + "." if notes else "")
    if not view["day"]:
        say += " Tareekh batayiye to task plan ban jayega."
    return {"event": view, "say": say}


def events(tenant: str, status: str = "") -> list[dict]:
    with session() as s:
        q = select(EpEvent).where(EpEvent.tenant_id == tenant)
        if status:
            q = q.where(EpEvent.status == status)
        rows = s.scalars(q.order_by(EpEvent.day.desc(), EpEvent.created_at.desc())).all()[:100]
        return [_ev_view(s, e) for e in rows]


# ---------------------------------------------------------------- vendor quotes + comparison
def _vq_total(q: EpVendorQuote) -> tuple[float | None, float | None, str]:
    """(base before GST, GST amount or None when unknown, note)."""
    amt = q.amount
    if q.per in ("plate", "person", "hour", "day"):
        if q.qty is None:
            return None, None, f"per {q.per} rate hai; kitne {q.per} chahiye pata nahi"
        amt = q.amount * q.qty
    if q.gst_included is True:
        if q.gst_pct is None:
            return R.r2(amt), None, "GST shamil bataya, par kitne % pata nahi"
        base = R.r2(amt / (1 + q.gst_pct / 100))
        return base, R.r2(amt - base), ""
    if q.gst_included is False:
        if q.gst_pct is None:
            return R.r2(amt), None, "GST alag lagega, % pata nahi"
        return R.r2(amt), R.r2(amt * q.gst_pct / 100), ""
    return R.r2(amt), None, "GST shamil hai ya alag, vendor se poochhiye"


def _vq_view(s, q: EpVendorQuote) -> dict:
    v = s.get(EpVendor, q.vendor_id)
    base, gst, note = _vq_total(q)
    return {"id": q.id, "event_id": q.event_id, "vendor_id": q.vendor_id, "vendor": v.name if v else "", "head": q.head,
            "head_label": R.HEADS.get(q.head, R.HEADS["misc"])["label"], "what": q.what, "amount": q.amount, "per": q.per, "qty": q.qty,
            "gst_included": q.gst_included, "gst_pct": q.gst_pct, "base": base, "gst": gst,
            "total": R.r2(base + gst) if base is not None and gst is not None else None, "note": note, "advance_pct": q.advance_pct,
            "cancel_terms": q.cancel_terms, "includes": q.includes, "valid_till": q.valid_till, "file_id": q.file_id, "chosen": q.chosen}


def vquote_add(tenant: str, event: str, vendor: str, amount=None, head: str = "", what: str = "", gst_included=None, gst_pct=None,
               per: str = "", qty=None, advance_pct=None, cancel_terms: str = "", includes: str = "", valid_till: str = "",
               file_id: str = "") -> dict:
    """A vendor's quote exactly as SHE got it. No amount = refused (Hoola never guesses a vendor's price)."""
    a = R.num(amount)
    if a is None or a <= 0:
        raise EventError("Vendor ne kitna bataya? Rakam chahiye (Hoola daam nahi gadhta).")
    with _LOCK, session() as s:
        e = _ev(s, tenant, event)
        v = _vendor(s, tenant, vendor)
        per = per if per in ("plate", "person", "hour", "day") else ""
        q = EpVendorQuote(tenant_id=tenant, event_id=e.id, vendor_id=v.id, head=(head if head in R.HEADS else (R.head_of(head) if head else v.head)),
                          what=R.clean(what, 400), amount=a, per=per, qty=R.num(qty),
                          gst_included=None if gst_included in (None, "") else bool(gst_included), gst_pct=R.num(gst_pct),
                          advance_pct=R.num(advance_pct), cancel_terms=R.clean(cancel_terms, 400), includes=R.clean(includes, 600),
                          valid_till=_iso(_date(valid_till)), file_id=str(file_id or "")[:60])
        if per and q.qty is None and e.headcount and per in ("plate", "person"):
            q.qty = float(e.headcount)
        s.add(q)
        if e.status == "enquiry":
            e.status = "planning"
        s.flush()
        view = _vq_view(s, q)
    return {"quote": view, "say": f"{view['vendor']} ka {view['head_label']} quote: {_rs(view['total'] if view['total'] is not None else view['base'])}"
                                  + (f" ({view['note']})" if view["note"] else "") + "."}


def compare(tenant: str, event: str, head: str = "") -> dict:
    with session() as s:
        e = _ev(s, tenant, event)
        q = select(EpVendorQuote).where(EpVendorQuote.event_id == e.id)
        if head:
            q = q.where(EpVendorQuote.head == (head if head in R.HEADS else R.head_of(head)))
        rows = [_vq_view(s, x) for x in s.scalars(q).all()]
        lines = _j(e.budget_lines, [])
    out = []
    for h in sorted({r["head"] for r in rows}):
        qs = sorted([r for r in rows if r["head"] == h], key=lambda r: (r["total"] is None, r["total"] or r["base"] or 0))
        priced = [r for r in qs if r["total"] is not None]
        if priced:
            priced[0]["tag"] = "sabse sasta (GST samet)"
        for r in qs:
            miss = []
            if r["gst"] is None:
                miss.append("GST saaf nahi")
            if r["advance_pct"] is None:
                miss.append("advance kitna")
            if not r["cancel_terms"]:
                miss.append("radd karne ki shart")
            if not r["includes"]:
                miss.append("kya shamil hai")
            r["missing"] = miss
        planned = sum((ln.get("cost") or 0) for ln in lines if ln.get("head") == h)
        out.append({"head": h, "label": R.HEADS.get(h, R.HEADS["misc"])["label"], "quotes": qs, "budget_cost": planned or None})
    say = "; ".join(f"{g['label']}: " + ", ".join(f"{q['vendor']} {_rs(q['total'] if q['total'] is not None else q['base'])}"
                                                  for q in g["quotes"]) for g in out) or "Abhi koi vendor quote nahi."
    return {"heads": out, "say": say}


def choose_quote(tenant: str, quote_id: str, actor: str = "owner") -> dict:
    """Her pick: the quote becomes the budget line's cost for that head (base before GST; GST noted)."""
    with _LOCK, session() as s:
        q = s.get(EpVendorQuote, quote_id)
        if q is None or q.tenant_id != tenant:
            raise EventError("Quote nahi mila.")
        e = s.get(EpEvent, q.event_id)
        for o in s.scalars(select(EpVendorQuote).where(EpVendorQuote.event_id == e.id, EpVendorQuote.head == q.head)).all():
            o.chosen = o.id == q.id
        view = _vq_view(s, q)
        lines = _j(e.budget_lines, [])
        ln = next((x for x in lines if x.get("quote") == q.id), None) or \
            next((x for x in lines if x.get("head") == q.head and not x.get("quote")), None)
        if ln is None:
            ln = {"id": f"B{len(lines) + 1}", "head": q.head, "item": view["what"] or view["head_label"]}
            lines.append(ln)
        ln.update({"quote": q.id, "vendor": view["vendor"], "cost": view["base"], "vendor_gst": view["gst"]})
        e.budget_lines = _dump(lines)
        _log(e, f"quote chuna: {view['vendor']} ({view['head_label']})", actor_name(tenant, actor))
    b = budget(tenant, e.id)
    return {"quote": view, "budget": b, "say": f"{view['vendor']} chuna; budget mein {view['head_label']} {_rs(view['base'])} (GST se pehle)."}


# ---------------------------------------------------------------- budget
def _margin(st: EpSettings, head: str) -> float:
    hm = _j(st.head_margins, {})
    return float(hm.get(head, st.margin_pct))


def _line(st: EpSettings, ln: dict) -> dict:
    cost = ln.get("cost")
    pure = bool(ln.get("pure_agent"))
    m = ln.get("margin_pct")
    m = float(m) if m is not None else (0.0 if pure or st.fee_mode != "markup" else _margin(st, ln.get("head", "misc")))
    price = None if cost is None else R.r2(cost * (1 + m / 100))
    return {**ln, "pure_agent": pure, "margin_pct": m, "client_price": price, "head_label": R.HEADS.get(ln.get("head"), R.HEADS["misc"])["label"],
            "blanks": [] if cost is not None else ["kharcha (vendor ka quote / aapka rate)"]}


def budget_set(tenant: str, event: str, lines: list | None, replace: bool = False, actor: str = "owner") -> dict:
    with _LOCK, session() as s:
        e = _ev(s, tenant, event)
        if e.accepted_doc and not is_owner(actor):
            raise EventError("Quote client ne maan liya hai; budget sirf owner badalti hain.")
        cur = [] if replace else _j(e.budget_lines, [])
        by_id = {x["id"]: x for x in cur}
        n = max([int(re.sub(r"\D", "", x["id"]) or 0) for x in cur] + [0])
        for raw in lines or []:
            if not isinstance(raw, dict):
                continue
            item = R.clean(raw.get("item") or raw.get("name"), 200)
            head = raw.get("head")
            head = head if head in R.HEADS else R.head_of(str(head or "") + " " + item)
            cost = R.num(raw.get("cost"))
            ln = by_id.get(str(raw.get("id") or ""))
            if ln is None:
                n += 1
                ln = {"id": f"B{n}"}
                cur.append(ln)
            ln.update({"head": head, "item": item or R.HEADS[head]["label"], "qty": R.num(raw.get("qty")), "unit": R.clean(raw.get("unit"), 20),
                       "cost": cost if cost is not None else ln.get("cost"), "pure_agent": bool(raw.get("pure_agent", ln.get("pure_agent")))})
            if raw.get("margin_pct") not in (None, ""):
                ln["margin_pct"] = R.num(raw.get("margin_pct"))
        e.budget_lines = _dump(cur)
        if e.status == "enquiry":
            e.status = "planning"
    return budget(tenant, e.id)


def budget_remove(tenant: str, event: str, line_id: str) -> dict:
    with _LOCK, session() as s:
        e = _ev(s, tenant, event)
        lines = [x for x in _j(e.budget_lines, []) if x.get("id") != line_id]
        e.budget_lines = _dump(lines)
    return budget(tenant, e.id)


def _actuals(s, e: EpEvent) -> dict[str, float]:
    act: dict[str, float] = {}
    for b in s.scalars(select(EpBooking).where(EpBooking.event_id == e.id, EpBooking.status.in_(("approved", "advance_paid")))).all():
        act[b.head] = act.get(b.head, 0) + b.amount
    for x in s.scalars(select(EpSpend).where(EpSpend.event_id == e.id, EpSpend.status.in_(("approved", "paid")))).all():
        act[x.head] = act.get(x.head, 0) + x.amount
    return act


def budget(tenant: str, event: str) -> dict:
    with session() as s:
        e = _ev(s, tenant, event)
        st = _st(s, tenant)
        lines = [_line(st, x) for x in _j(e.budget_lines, [])]
        act = _actuals(s, e)
        heads: dict[str, dict] = {}
        for ln in lines:
            h = heads.setdefault(ln["head"], {"head": ln["head"], "label": ln["head_label"], "cost": 0.0, "client": 0.0, "blank": 0})
            if ln["cost"] is None:
                h["blank"] += 1
            else:
                h["cost"] += ln["cost"]
                h["client"] += ln["client_price"] or 0
        for h, v in heads.items():
            v["actual"] = R.r2(act.get(h, 0))
            v["over"] = v["actual"] > v["cost"] + 0.5 if v["cost"] else False
        cost = R.r2(sum(ln["cost"] for ln in lines if ln["cost"] is not None))
        client = R.r2(sum(ln["client_price"] for ln in lines if ln["client_price"] is not None))
        fee = R.r2(cost * st.fee_pct / 100) if st.fee_mode == "fee" and st.fee_pct else 0.0
        blanks = [f"{ln['id']} {ln['item']}" for ln in lines if ln["blanks"]]
        warn = []
        if e.budget is not None and client + fee > e.budget:
            warn.append(f"Client ka budget {_rs(e.budget)}, plan {_rs(client + fee)} (GST se pehle) - {_rs(client + fee - e.budget)} zyada.")
        for v in heads.values():
            if v["over"]:
                warn.append(f"{v['label']}: asli kharcha {_rs(v['actual'])} plan {_rs(v['cost'])} se zyada.")
        view = {"event": e.number, "event_id": e.id, "lines": lines, "heads": list(heads.values()), "cost": cost, "client": client,
                "fee": fee, "fee_mode": st.fee_mode, "margin": R.r2(client + fee - cost), "client_budget": e.budget, "blanks": blanks,
                "warnings": warn, "actual": R.r2(sum(act.values()))}
    view["say"] = (f"{view['event']}: kharcha {_rs(cost)}, client ko {_rs(client + fee)} (GST se pehle), margin {_rs(view['margin'])}."
                   + (f" {len(blanks)} line ka kharcha khaali." if blanks else "") + (" " + " ".join(warn) if warn else ""))
    return view


# ---------------------------------------------------------------- client quotation / invoice (her tap)
def _doc_lines(st: EpSettings, lines: list[dict]) -> tuple[list[dict], list[str]]:
    out, blanks = [], []
    for ln in lines:
        L = _line(st, ln)
        if L["cost"] is None:
            blanks.append(f"{L['id']} {L['item']}: kharcha")
        out.append({"id": L["id"], "head": L["head"], "item": L["item"], "qty": L.get("qty"), "unit": L.get("unit", ""),
                    "amount": L["client_price"] if st.fee_mode == "markup" else L["cost"], "pure_agent": L["pure_agent"]})
    return out, blanks


def _doc_totals(st: EpSettings, e: EpEvent, lines: list[dict]) -> dict:
    intra = not (st.state_code and e.client_state and st.state_code != e.client_state)
    taxable = sum(c["amount"] for c in lines if c["amount"] is not None and not c["pure_agent"])
    pure = sum(c["amount"] for c in lines if c["amount"] is not None and c["pure_agent"])
    fee = 0.0
    if st.fee_mode == "fee" and st.fee_pct:
        fee = R.r2(sum(c["amount"] for c in lines if c["amount"] is not None) * st.fee_pct / 100)
    taxable = R.r2(taxable + fee)
    g = R.gst_split(taxable, st.gst_pct, intra)
    total = R.r2(taxable + g["tax"] + pure)
    return {"taxable": taxable, "fee": fee, "fee_pct": st.fee_pct if fee else 0, "pure_agent": R.r2(pure), "gst_pct": st.gst_pct,
            "sac": st.sac, "intra": intra, **g, "total": total, "words": _words(total)}


def _doc_show(d: EpDoc, e: EpEvent) -> dict:
    t = _j(d.totals, {})
    return {"kind": d.kind, "number": d.number, "event": e.number, "client": e.client, "phone": e.client_phone, "total": t.get("total"),
            "lines": [[c["id"], c["item"], c["amount"], c["pure_agent"]] for c in _j(d.lines, [])], "valid_until": d.valid_until,
            "due": t.get("balance")}


def _doc_view(s, d: EpDoc) -> dict:
    e = s.get(EpEvent, d.event_id)
    t = _j(d.totals, {})
    return {"id": d.id, "kind": d.kind, "number": d.number or "(haan par number milega)", "event": e.number if e else "", "event_id": d.event_id,
            "event_name": e.name if e else "", "client": e.client if e else "", "lines": _j(d.lines, []), "totals": t,
            "total_text": _rs(t.get("total")), "blanks": _j(d.blanks, []), "status": d.status, "hash": d.hash, "valid_until": d.valid_until,
            "file_id": d.file_id, "reason": d.reason}


def quote_make(tenant: str, event: str, conv: str = "", actor: str = "owner") -> dict:
    """A new quotation revision from the budget (draft; nothing goes to the client)."""
    with _LOCK, session() as s:
        e = _ev(s, tenant, event)
        if e.accepted_doc:
            raise EventError("Client ne quote maan liya hai; ab badlav budget + invoice mein.")
        st = _st(s, tenant)
        src = _j(e.budget_lines, [])
        if not src:
            raise EventError("Pehle budget banaiye (vendor quotes / aapke rate).")
        lines, blanks = _doc_lines(st, src)
        if not e.client:
            blanks.append("client ka naam")
        if not e.day:
            blanks.append("event ki tareekh")
        totals = _doc_totals(st, e, lines)
        olds = s.scalars(select(EpDoc).where(EpDoc.event_id == e.id, EpDoc.kind == "quote")).all()
        for o in olds:
            if o.status in ("draft", "waiting", "approved"):
                o.status = "superseded"
        rev = len(olds) + 1
        if rev == 1:
            st.doc_seq += 1
            base = f"EQ-{st.doc_seq:04d}"
        else:
            base = olds[0].number.split("/")[0]
        d = EpDoc(tenant_id=tenant, event_id=e.id, kind="quote", number=f"{base}/R{rev}", lines=_dump(lines), totals=_dump(totals),
                  blanks=_dump(blanks), valid_until=_iso(today(tenant) + timedelta(days=st.validity_days)))
        s.add(d)
        s.flush()
        d.hash = R.digest(_doc_show(d, e))
        view = _doc_view(s, d)
    say = f"{view['number']} draft: {view['total_text']} (GST samet)." + (
        f" {len(blanks)} khaali jagah bhariye, tabhi client ko jayega." if blanks else " Client ko bhejne ke liye haan ka card maangiye.")
    return {"doc": view, "say": say}


def _doc(s, tenant: str, ref: str) -> EpDoc:
    d = s.get(EpDoc, ref)
    if d is None:
        d = s.scalars(select(EpDoc).where(EpDoc.tenant_id == tenant, EpDoc.number == ref)).first()
    if d is None or d.tenant_id != tenant:
        raise EventError("Quotation / invoice nahi mila.")
    return d


def doc_view(tenant: str, ref: str) -> dict:
    with session() as s:
        return _doc_view(s, _doc(s, tenant, ref))


def doc_request(tenant: str, ref: str, conv: str = "", actor: str = "owner") -> dict:
    """Card for HER tap; the quotation / invoice goes to the client only after it."""
    with _LOCK, session() as s:
        d = _doc(s, tenant, ref)
        e = s.get(EpEvent, d.event_id)
        if d.status not in ("draft", "waiting"):
            raise EventError(f"{d.number} '{d.status}' hai.")
        blanks = _j(d.blanks, [])
        if blanks:
            raise EventError("Pehle khaali jagah bhariye: " + "; ".join(blanks[:6]))
        d.status = "waiting"
        d.hash = R.digest(_doc_show(d, e))
        view = _doc_view(s, d)
    what = "quotation" if view["kind"] == "quote" else "invoice"
    say = f"{view['event']} ki {what} ({view['total_text']}) client {view['client']} ko bhejne se pehle aapki haan ka card bheja hai."
    _publish(tenant, "eventplan_doc", {"doc": view, "say": say}, conv)
    return {"doc": view, "status": "needs_tap", "say": say}


def decide_doc(tenant: str, ref: str, yes: bool, card_hash: str = "", reason: str = "", by_tap: bool = False, conv: str = "") -> dict:
    if not by_tap:
        raise EventError("Quotation / invoice client ko sirf aapke tap se (app ka card).")
    with _LOCK, session() as s:
        d = _doc(s, tenant, ref)
        e = s.get(EpEvent, d.event_id)
        if d.status != "waiting":
            return {"doc": _doc_view(s, d), "say": f"{d.number} pehle hi '{d.status}' hai."}
        if not card_hash or card_hash != d.hash or R.digest(_doc_show(d, e)) != d.hash:
            raise EventError("Card badal gaya hai; naya card dekhiye.")
        if not yes:
            d.status, d.reason = "rejected", R.clean(reason, 300)
            return {"doc": _doc_view(s, d), "say": "Theek hai, nahi bheja."}
        if d.kind == "invoice" and not d.number:
            st = _st(s, tenant)
            st.inv_seq += 1
            d.number = f"EI/{_fy(today(tenant))}/{st.inv_seq:04d}"
        d.status = "approved"
        did = d.id
    pdf = doc_pdf(tenant, did, final=True)
    with session() as s:
        d = _doc(s, tenant, did)
        e = s.get(EpEvent, d.event_id)
        st = _st(s, tenant)
        f = _store(tenant, f"{d.number.replace('/', '-')}_{d.kind}.pdf", pdf, "application/pdf", conv, f"{d.number} {d.kind}")
        d.file_id = f.get("file_id", "")
        t = _j(d.totals, {})
        if d.kind == "quote":
            text = (f"Namaste {e.client} ji, {st.firm or 'hamari team'} ki taraf se {e.name} ka quotation {d.number}: kul "
                    f"{_rs(t.get('total'))} (GST samet), {d.valid_until} tak valid. Koi sawaal ho to bataiye.")
            upi = ""
        else:
            bal = t.get("balance", t.get("total")) or 0
            upi_id, nm = _her_upi(tenant)
            upi = _upi_link(upi_id, nm, bal, f"{d.number}")
            text = (f"Namaste {e.client} ji, {e.name} ka invoice {d.number}: kul {_rs(t.get('total'))}, baaki {_rs(bal)}."
                    + (f" UPI: {upi}" if upi else "") + f" Dhanyavaad, {st.firm or ''}".rstrip(", "))
        draft = {"text": text, "wa": _wa(e.client_phone, text), "sms": _sms(e.client_phone, text), "file_id": d.file_id, "upi": upi}
        view = _doc_view(s, d)
    return {"doc": view, "draft": draft, "say": f"{view['number']} final PDF bana. Message draft taiyaar: aap khud bhejiye (PDF saath lagaiye)."}


def doc_mark(tenant: str, ref: str, what: str, actor: str = "owner") -> dict:
    """sent (she sent it) | accepted (client agreed: owner only) | rejected (client said no)."""
    with _LOCK, session() as s:
        d = _doc(s, tenant, ref)
        e = s.get(EpEvent, d.event_id)
        if what == "sent":
            if d.status != "approved":
                raise EventError("Pehle aapki haan (card), phir bhejna.")
            d.status = "sent"
        elif what == "accepted":
            if not is_owner(actor):
                raise EventError("Client ki manzoori sirf owner darj karti hain.")
            if d.kind != "quote" or d.status not in ("approved", "sent"):
                raise EventError("Yeh quotation aapki haan se client ko gaya hi nahi.")
            d.status = "accepted"
            e.accepted_doc = d.id
            if e.status in ("enquiry", "planning"):
                e.status = "confirmed"
        elif what == "rejected":
            d.status = "rejected"
        else:
            raise EventError("what = sent | accepted | rejected")
        _log(e, f"{d.kind} {d.number} {what}", actor_name(tenant, actor))
        view = _doc_view(s, d)
    return {"doc": view, "say": f"{view['number']}: {what}." + (" Event pakka." if what == "accepted" else "")}


def client_paid(tenant: str, event: str, amount=None, ref: str = "", day: str = "", actor: str = "owner") -> dict:
    if not is_owner(actor):
        raise EventError("Client se paise mile, yeh sirf owner darj karti hain.")
    a = R.num(amount)
    if a is None or a <= 0:
        raise EventError("Kitne mile?")
    with _LOCK, session() as s:
        e = _ev(s, tenant, event)
        paid = _j(e.client_paid, [])
        paid.append({"amount": a, "ref": R.clean(ref, 80), "on": (_date(day) or today(tenant)).isoformat()})
        e.client_paid = _dump(paid)
        tot = sum(x["amount"] for x in paid)
        num = e.number
    return {"paid": paid, "say": f"{num}: {_rs(a)} mila darj (kul {_rs(tot)}). Hoola paise nahi leta, sirf hisaab."}


def invoice_make(tenant: str, event: str, conv: str = "", actor: str = "owner") -> dict:
    """Client invoice DRAFT after the event: current budget lines at client prices (changes after the quote shown), GST,
    minus what the client already paid (her word). Number only at her tap."""
    with _LOCK, session() as s:
        e = _ev(s, tenant, event)
        st = _st(s, tenant)
        d0 = _date(e.day)
        if not d0 or d0 > today(tenant):
            raise EventError("Invoice event ke baad (ya event ke din) banta hai.")
        src = _j(e.budget_lines, [])
        if not src:
            raise EventError("Budget khaali hai.")
        lines, blanks = _doc_lines(st, src)
        if not e.client:
            blanks.append("client ka naam")
        totals = _doc_totals(st, e, lines)
        paid = R.r2(sum(x["amount"] for x in _j(e.client_paid, [])))
        totals["paid"] = paid
        totals["balance"] = R.r2(totals["total"] - paid)
        if e.accepted_doc:
            q = s.get(EpDoc, e.accepted_doc)
            qt = (_j(q.totals, {}).get("total") or 0) if q else 0
            totals["quote_total"] = qt
            totals["vs_quote"] = R.r2(totals["total"] - qt)
        for o in s.scalars(select(EpDoc).where(EpDoc.event_id == e.id, EpDoc.kind == "invoice", EpDoc.status.in_(("draft", "waiting")))).all():
            o.status = "superseded"
        d = EpDoc(tenant_id=tenant, event_id=e.id, kind="invoice", lines=_dump(lines), totals=_dump(totals), blanks=_dump(blanks))
        s.add(d)
        s.flush()
        d.hash = R.digest(_doc_show(d, e))
        view = _doc_view(s, d)
    say = f"Invoice draft: {view['total_text']}, baaki {_rs(totals['balance'])}."
    if totals.get("vs_quote"):
        say += f" Quote se {_rs(abs(totals['vs_quote']))} {'zyada' if totals['vs_quote'] > 0 else 'kam'} (client ko batayiye)."
    return {"doc": view, "say": say + " Bhejne ke liye haan ka card maangiye."}


# ---------------------------------------------------------------- PDFs
def _pdf(title: str, blocks: list[tuple], watermark: str = "") -> bytes:
    from xml.sax.saxutils import escape

    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    from .. import indicpdf
    pdf = indicpdf.blocks_pdf(title, blocks, watermark)       # Hindi / Marathi names: printed by Chromium (shaped)
    if pdf:
        return pdf
    ss = getSampleStyleSheet()
    small = ss["BodyText"].clone("small", fontSize=8.5, leading=10.5)

    def txt(t: str) -> str:
        return indicpdf.rl_markup(escape(str(t or "").replace("₹", "Rs. ")))
    story: list = []
    for b in blocks:
        if b[0] == "h":
            story.append(Paragraph(txt(b[1]), ss["Heading2"]))
        elif b[0] == "h1":
            story.append(Paragraph(txt(b[1]), ss["Title"]))
        elif b[0] == "p":
            story.append(Paragraph(txt(b[1]), ss["BodyText"]))
        elif b[0] == "table":
            rows = [[Paragraph(txt(c), small) for c in r] for r in b[1]]
            t = Table(rows, colWidths=b[2] if len(b) > 2 else None, repeatRows=1)
            t.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.3, colors.grey), ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#EEEEEE")),
                                   ("VALIGN", (0, 0), (-1, -1), "TOP")]))
            story.append(t)
        story.append(Spacer(1, 6))
    bio = io.BytesIO()

    def mark(c, _d):
        if watermark:
            c.saveState()
            c.setFont("Helvetica-Bold", 46)
            c.setFillColor(colors.Color(0.85, 0.2, 0.2, alpha=0.18))
            c.translate(300, 420)
            c.rotate(35)
            c.drawCentredString(0, 0, watermark)
            c.restoreState()
    doc = SimpleDocTemplate(bio, pagesize=A4, title=title, leftMargin=36, rightMargin=36, topMargin=36, bottomMargin=36)
    doc.build(story, onFirstPage=mark, onLaterPages=mark)
    return bio.getvalue()


def _head(st: EpSettings, title: str) -> list[tuple]:
    who = " · ".join(x for x in (st.firm or "Event planner", f"GSTIN {st.gstin}" if st.gstin else "") if x)
    contact = " · ".join(x for x in (st.address, st.phone, st.email) if x)
    return [("h1", title), ("p", who)] + ([("p", contact)] if contact else [])


def doc_pdf(tenant: str, ref: str, final: bool = False) -> bytes:
    with session() as s:
        d = _doc(s, tenant, ref)
        e = s.get(EpEvent, d.event_id)
        st = _st(s, tenant)
        t = _j(d.totals, {})
        name = "Quotation" if d.kind == "quote" else "Tax invoice"
        blocks = _head(st, f"{name} {d.number or '(draft)'}")
        blocks += [("p", f"Client: {e.client}" + (f" (GSTIN {e.client_gstin})" if e.client_gstin else "")),
                   ("p", f"Event {e.number}: {e.name} · {e.day} {e.start}-{e.end} · {e.venue} {e.city} · {e.headcount or '-'} guests"),
                   ("p", f"Date {_aware(d.created_at).date().isoformat() if d.created_at else ''}"
                    + (f" · Valid till {d.valid_until}" if d.kind == "quote" else "") + f" · SAC {t.get('sac', '')}")]
        rows = [["#", "Head", "Item", "Qty", "Amount", "GST?"]]
        for c in _j(d.lines, []):
            rows.append([c["id"], R.HEADS.get(c["head"], R.HEADS["misc"])["label"], c["item"],
                         "" if c.get("qty") is None else f"{c['qty']:g} {c.get('unit') or ''}", _pdf_rs(c["amount"]),
                         "No (reimbursement at actual)" if c["pure_agent"] else "Yes"])
        blocks.append(("table", rows, [28, 90, 160, 55, 80, 80]))
        tr = [["", "Amount"]]
        if t.get("fee"):
            tr.append([f"Event management fee {t.get('fee_pct'):g}%", _pdf_rs(t["fee"])])
        tr.append(["Taxable value", _pdf_rs(t.get("taxable"))])
        if t.get("intra"):
            tr += [[f"CGST {t['gst_pct'] / 2:g}%", _pdf_rs(t.get("cgst"))], [f"SGST {t['gst_pct'] / 2:g}%", _pdf_rs(t.get("sgst"))]]
        else:
            tr.append([f"IGST {t['gst_pct']:g}%", _pdf_rs(t.get("igst"))])
        if t.get("pure_agent"):
            tr.append(["Reimbursements at actual (pure agent, outside GST)", _pdf_rs(t["pure_agent"])])
        tr.append(["TOTAL", _pdf_rs(t.get("total"))])
        if d.kind == "invoice":
            tr += [["Received", _pdf_rs(t.get("paid"))], ["Balance due", _pdf_rs(t.get("balance"))]]
        blocks.append(("table", tr, [330, 120]))
        if t.get("words"):
            blocks.append(("p", f"In words: {t['words']}"))
        if _j(d.blanks, []):
            blocks.append(("p", "Blanks to fill: " + "; ".join(_j(d.blanks, []))))
        blocks.append(("h", "Terms"))
        blocks.append(("p", st.terms or "Advance as agreed; balance after the event. Changes in scope / headcount are charged extra. "
                       "Cancellation charges as per the vendors' own terms. GST as applicable."))
        blocks.append(("p", "Prepared with Hoola. Costs are from the planner's vendor quotes."))
    return _pdf(f"{name} {d.number}", blocks, "" if final else "DRAFT - NOT SENT")


# ---------------------------------------------------------------- vendor bookings (her tap + fingerprint for an advance)
def _bk_show(s, b: EpBooking) -> dict:
    e = s.get(EpEvent, b.event_id)
    v = s.get(EpVendor, b.vendor_id)
    return {"number": b.number, "event": e.number if e else "", "day": e.day if e else "", "vendor": v.name if v else "",
            "vendor_phone": v.phone if v else "", "vendor_upi": v.upi_id if v else "", "head": b.head, "amount": b.amount,
            "advance": b.advance, "advance_due": b.advance_due, "arrive": b.arrive, "terms": b.terms}


def _bk_view(s, b: EpBooking) -> dict:
    return {"id": b.id, **_bk_show(s, b), "event_id": b.event_id, "head_label": R.HEADS.get(b.head, R.HEADS["misc"])["label"],
            "amount_text": _rs(b.amount), "advance_text": _rs(b.advance), "checks": _j(b.checks, []), "status": b.status, "hash": b.hash,
            "paid": b.paid, "paid_ref": b.paid_ref, "asked_by": b.asked_by, "reason": b.reason}


def booking_request(tenant: str, event: str, vendor: str = "", quote: str = "", amount=None, advance=None, advance_due: str = "",
                    arrive: str = "", terms: str = "", head: str = "", actor: str = "owner", conv: str = "") -> dict:
    """Only a card. Amount = the vendor's quote (chosen quote, or her word). Venue double-booked on the same day = refused."""
    with _LOCK, session() as s:
        e = _ev(s, tenant, event)
        q = s.get(EpVendorQuote, quote) if quote else None
        if q is not None and (q.tenant_id != tenant or q.event_id != e.id):
            q = None
        if q is None and not quote and vendor:
            v0 = _vendor(s, tenant, vendor)
            q = s.scalars(select(EpVendorQuote).where(EpVendorQuote.event_id == e.id, EpVendorQuote.vendor_id == v0.id)
                          .order_by(EpVendorQuote.chosen.desc(), EpVendorQuote.created_at.desc())).first()
        v = s.get(EpVendor, q.vendor_id) if q is not None else _vendor(s, tenant, vendor)
        amt = R.num(amount)
        checks = []
        if amt is None and q is not None:
            qv = _vq_view(s, q)
            amt = qv["total"] if qv["total"] is not None else qv["base"]
            if qv["note"]:
                checks.append(qv["note"])
        if amt is None or amt <= 0:
            raise EventError("Booking ki rakam chahiye (vendor ka quote). Hoola daam nahi gadhta.")
        adv = R.num(advance)
        if adv is None and q is not None and q.advance_pct:
            adv = R.r2(amt * q.advance_pct / 100)
        adv = adv or 0.0
        if adv > amt:
            raise EventError("Advance kul rakam se zyada nahi ho sakta.")
        h = head if head in R.HEADS else (q.head if q is not None else v.head)
        # double booking: the same venue on the same day for another of her live events
        if h == "venue" and e.day:
            for o in s.scalars(select(EpBooking).where(EpBooking.tenant_id == tenant, EpBooking.vendor_id == v.id,
                                                       EpBooking.status.in_(("waiting", "approved", "advance_paid")),
                                                       EpBooking.event_id != e.id)).all():
                oe = s.get(EpEvent, o.event_id)
                if oe and oe.day == e.day and oe.status not in ("cancelled", "done"):
                    raise EventError(f"{v.name} {e.day} ko pehle se {oe.number} ke liye book hai: double booking nahi.")
        for o in s.scalars(select(EpBooking).where(EpBooking.event_id == e.id, EpBooking.vendor_id == v.id,
                                                   EpBooking.status == "waiting")).all():
            return {"booking": _bk_view(s, o), "existing": True, "say": f"{o.number} ka card pehle se aapki haan ka intezaar kar raha hai."}
        if h == "catering" and not v.fssai:
            checks.append("Caterer ka FSSAI number nahi likha (vendor mein jodiye).")
        if not v.phone:
            checks.append("Vendor ka phone nahi; confirmation draft nahi banega.")
        if not e.accepted_doc:
            checks.append("Client ne abhi quotation nahi maana.")
        planned = sum((ln.get("cost") or 0) for ln in _j(e.budget_lines, []) if ln.get("head") == h)
        if planned and amt > planned * 1.18 + 1:
            checks.append(f"Budget mein {R.HEADS[h]['label']} {_rs(planned)} (GST se pehle) tha; booking {_rs(amt)}.")
        ad = _date(advance_due)
        if advance_due and not ad:
            raise EventError("Advance ki tareekh samajh nahi aayi.")
        ar = R.hhmm(arrive) if arrive else ""
        st = _st(s, tenant)
        st.bk_seq += 1
        b = EpBooking(tenant_id=tenant, event_id=e.id, vendor_id=v.id, quote_id=q.id if q is not None else "", number=f"VB-{st.bk_seq:04d}",
                      head=h, amount=R.r2(amt), advance=R.r2(adv), advance_due=_iso(ad), arrive=ar, terms=R.clean(terms or (q.cancel_terms if q is not None else ""), 600),
                      checks=_dump(checks), asked_by=actor_name(tenant, actor))
        s.add(b)
        s.flush()
        b.hash = R.digest(_bk_show(s, b))
        view = _bk_view(s, b)
    say = (f"{view['number']}: {view['vendor']} ({view['head_label']}) {view['amount_text']}"
           + (f", advance {view['advance_text']}" if view["advance"] else "") + " - aapki haan ka card bheja hai."
           + (" Dhyaan: " + " ".join(checks) if checks else ""))
    _publish(tenant, "eventplan_booking", {"booking": view, "say": say}, conv)
    return {"booking": view, "status": "needs_tap", "say": say}


def _challenge(tenant: str, target: str, digest: str, title: str, amount: str, payee: str) -> dict:
    from .. import approvekey as ak
    nonce = secrets.token_urlsafe(18)
    exp = datetime.now(timezone.utc) + ak.CHALLENGE_TTL
    msg = f"hoola-approve|v1|{tenant}|eventplan|{target}|{digest}|{nonce}|{int(exp.timestamp())}"
    with session() as s:
        ch = ak.ApproveChallenge(tenant_id=tenant, kind="eventplan", target=target, digest=digest, message=msg, expires_at=exp)
        s.add(ch)
        s.flush()
        chid = ch.id
    return {"challenge_id": chid, "message": msg, "show": {"title": title, "amount": amount, "payee": payee}, "expires_at": exp.isoformat()}


def _verify(tenant: str, target: str, digest: str, proof: dict | None) -> None:
    from .. import approvekey as ak
    if not proof:
        raise EventError("fingerprint_needed")
    with session() as s:
        ak._settle(s, tenant)
        ok = ak._check(s, tenant, "eventplan", target, proof, digest)
    if not ok:
        raise EventError("fingerprint_invalid")


def booking_challenge(tenant: str, booking_id: str) -> dict:
    with session() as s:
        b = s.get(EpBooking, booking_id)
        if b is None or b.tenant_id != tenant or b.status != "waiting":
            raise EventError("Yeh card ab baaki nahi hai.")
        v = _bk_view(s, b)
    return _challenge(tenant, booking_id, v["hash"], f"Vendor booking {v['number']}", v["advance_text"] if v["advance"] else v["amount_text"], v["vendor"])


def decide_booking(tenant: str, booking_id: str, yes: bool, card_hash: str = "", reason: str = "", proof: dict | None = None,
                   by_tap: bool = False) -> dict:
    if not by_tap:
        raise EventError("Vendor booking sirf aapke tap se (app ka card).")
    with session() as s:
        b = s.get(EpBooking, booking_id)
        if b is None or b.tenant_id != tenant:
            raise EventError("Card nahi mila.")
        if b.status != "waiting":
            return {"booking": _bk_view(s, b), "say": f"{b.number} pehle hi '{b.status}' hai."}
        if not card_hash or card_hash != b.hash or R.digest(_bk_show(s, b)) != b.hash:
            raise EventError("Card badal gaya hai; naya card dekhiye.")
        adv = b.advance
    if yes and adv > 0:
        from .. import approvekey as ak
        if ak.active_on(tenant):
            _verify(tenant, booking_id, card_hash, proof)
    with _LOCK, session() as s:
        b = s.get(EpBooking, booking_id)
        b.status = "approved" if yes else "rejected"
        b.reason = R.clean(reason, 300)
        if yes and b.quote_id:
            q = s.get(EpVendorQuote, b.quote_id)
            if q is not None and not q.chosen:
                for o in s.scalars(select(EpVendorQuote).where(EpVendorQuote.event_id == b.event_id, EpVendorQuote.head == q.head)).all():
                    o.chosen = o.id == q.id
        e = s.get(EpEvent, b.event_id)
        _log(e, f"booking {b.number} {'haan' if yes else 'nahi'}", "Owner")
        view = _bk_view(s, b)
        firm = _st(s, tenant).firm
    if not yes:
        return {"booking": view, "say": f"{view['number']}: booking nahi."}
    text = (f"Namaste {view['vendor']} ji, {view['event']} ({view['day']}) ke liye {view['head_label']} booking pakki: {view['amount_text']}"
            + (f", advance {view['advance_text']}" + (f" {view['advance_due']} tak" if view["advance_due"] else "") if view["advance"] else "")
            + (f". Aapko {view['arrive']} baje pahunchna hai" if view["arrive"] else "") + f". - {firm or ''}".rstrip(" -"))
    upi = _upi_link(view["vendor_upi"], view["vendor"], view["advance"], f"{view['event']} advance") if view["advance"] else ""
    return {"booking": view, "draft": {"text": text, "wa": _wa(view["vendor_phone"], text), "sms": _sms(view["vendor_phone"], text), "upi": upi},
            "say": f"{view['number']} maana. Vendor ko message draft" + (" + advance ka UPI link (aap apne UPI app se dengi)" if upi else "")
                   + ". Hoola khud paise nahi deta."}


def booking_paid(tenant: str, booking_id: str, amount=None, ref: str = "", actor: str = "owner") -> dict:
    if not is_owner(actor):
        raise EventError("Paise diye, yeh sirf owner darj karti hain.")
    with _LOCK, session() as s:
        b = s.get(EpBooking, booking_id)
        if b is None or b.tenant_id != tenant:
            raise EventError("Booking nahi mili.")
        if b.status not in ("approved", "advance_paid"):
            raise EventError("Pehle aapki haan (card), phir paisa.")
        a = R.num(amount)
        b.paid = R.r2(b.paid + (a if a is not None else b.advance))
        b.paid_ref = R.clean(ref, 80)
        b.status = "advance_paid"
        v = _bk_view(s, b)
    return {"booking": v, "say": f"{v['number']}: {_rs(v['paid'])} diya darj (kul {v['amount_text']})."}


def bookings(tenant: str, event: str = "", status: str = "") -> list[dict]:
    with session() as s:
        q = select(EpBooking).where(EpBooking.tenant_id == tenant)
        if event:
            q = q.where(EpBooking.event_id == _ev(s, tenant, event).id)
        if status:
            q = q.where(EpBooking.status == status)
        return [_bk_view(s, b) for b in s.scalars(q.order_by(EpBooking.created_at)).all()]


# ---------------------------------------------------------------- other spends (her tap + fingerprint)
def _sp_show(s, x: EpSpend) -> dict:
    e = s.get(EpEvent, x.event_id)
    return {"number": x.number, "event": e.number if e else "", "head": x.head, "what": x.what, "payee": x.payee, "amount": x.amount}


def _sp_view(s, x: EpSpend) -> dict:
    return {"id": x.id, **_sp_show(s, x), "event_id": x.event_id, "amount_text": _rs(x.amount), "checks": _j(x.checks, []),
            "status": x.status, "hash": x.hash, "paid_ref": x.paid_ref, "asked_by": x.asked_by,
            "head_label": R.HEADS.get(x.head, R.HEADS["misc"])["label"]}


def spend_request(tenant: str, event: str, what: str, amount=None, payee: str = "", head: str = "", actor: str = "owner",
                  conv: str = "") -> dict:
    a = R.num(amount)
    if a is None or a <= 0:
        raise EventError("Kitna kharcha? Rakam chahiye (bill / quote se).")
    what = R.clean(what, 300)
    if not what:
        raise EventError("Kis cheez ka kharcha?")
    with _LOCK, session() as s:
        e = _ev(s, tenant, event)
        h = head if head in R.HEADS else R.head_of(head + " " + what)
        for o in s.scalars(select(EpSpend).where(EpSpend.event_id == e.id, EpSpend.status == "waiting")).all():
            if o.what == what and abs(o.amount - a) < 0.5:
                return {"spend": _sp_view(s, o), "existing": True, "say": f"{o.number} ka card pehle se intezaar mein."}
        checks = []
        planned = sum((ln.get("cost") or 0) for ln in _j(e.budget_lines, []) if ln.get("head") == h)
        act = _actuals(s, e).get(h, 0)
        if planned and act + a > planned + 0.5:
            checks.append(f"{R.HEADS[h]['label']}: plan {_rs(planned)}, is kharche ke baad {_rs(act + a)}.")
        if not planned:
            checks.append(f"{R.HEADS[h]['label']} budget mein nahi tha.")
        st = _st(s, tenant)
        st.sp_seq += 1
        x = EpSpend(tenant_id=tenant, event_id=e.id, number=f"ES-{st.sp_seq:04d}", head=h, what=what, payee=R.clean(payee, 120), amount=R.r2(a),
                    checks=_dump(checks), asked_by=actor_name(tenant, actor))
        s.add(x)
        s.flush()
        x.hash = R.digest(_sp_show(s, x))
        view = _sp_view(s, x)
    say = f"{view['number']}: {view['what']} {view['amount_text']} - aapki haan ka card bheja hai." + (" " + " ".join(checks) if checks else "")
    _publish(tenant, "eventplan_spend", {"spend": view, "say": say}, conv)
    return {"spend": view, "status": "needs_tap", "say": say}


def spend_challenge(tenant: str, spend_id: str) -> dict:
    with session() as s:
        x = s.get(EpSpend, spend_id)
        if x is None or x.tenant_id != tenant or x.status != "waiting":
            raise EventError("Yeh card ab baaki nahi hai.")
        v = _sp_view(s, x)
    return _challenge(tenant, spend_id, v["hash"], f"Kharcha {v['number']}", v["amount_text"], v["payee"] or v["what"])


def decide_spend(tenant: str, spend_id: str, yes: bool, card_hash: str = "", proof: dict | None = None, by_tap: bool = False) -> dict:
    if not by_tap:
        raise EventError("Kharcha sirf aapke tap se (app ka card).")
    with session() as s:
        x = s.get(EpSpend, spend_id)
        if x is None or x.tenant_id != tenant:
            raise EventError("Card nahi mila.")
        if x.status != "waiting":
            return {"spend": _sp_view(s, x), "say": f"{x.number} pehle hi '{x.status}' hai."}
        if not card_hash or card_hash != x.hash or R.digest(_sp_show(s, x)) != x.hash:
            raise EventError("Card badal gaya hai; naya card dekhiye.")
    if yes:
        from .. import approvekey as ak
        if ak.active_on(tenant):
            _verify(tenant, spend_id, card_hash, proof)
    with _LOCK, session() as s:
        x = s.get(EpSpend, spend_id)
        x.status = "approved" if yes else "rejected"
        v = _sp_view(s, x)
    return {"spend": v, "say": f"{v['number']}: " + ("haan; aap khud dijiye, phir 'diya' darj kijiye." if yes else "nahi.")}


def spend_paid(tenant: str, spend_id: str, ref: str = "", actor: str = "owner") -> dict:
    if not is_owner(actor):
        raise EventError("Sirf owner.")
    with _LOCK, session() as s:
        x = s.get(EpSpend, spend_id)
        if x is None or x.tenant_id != tenant:
            raise EventError("Kharcha nahi mila.")
        if x.status != "approved":
            raise EventError("Pehle aapki haan (card).")
        x.status, x.paid_ref = "paid", R.clean(ref, 80)
        v = _sp_view(s, x)
    return {"spend": v, "say": f"{v['number']} diya darj."}


def spends(tenant: str, event: str = "", status: str = "") -> list[dict]:
    with session() as s:
        q = select(EpSpend).where(EpSpend.tenant_id == tenant)
        if event:
            q = q.where(EpSpend.event_id == _ev(s, tenant, event).id)
        if status:
            q = q.where(EpSpend.status == status)
        return [_sp_view(s, x) for x in s.scalars(q.order_by(EpSpend.created_at)).all()]


# ---------------------------------------------------------------- guests + RSVP
def _g_view(g: EpGuest) -> dict:
    return {"id": g.id, "name": g.name, "phone": g.phone, "group": g.group, "invited": g.invited, "status": g.status, "plus": g.plus,
            "note": g.note, "no_messages": g.no_messages, "came": g.came}


def guests_add(tenant: str, event: str, rows: list | None = None, text: str = "", file_id: str = "", data: bytes | None = None,
               name: str = "guests.csv") -> dict:
    """From her list: rows [{name, phone, group}], lines 'Rahul Sharma 9876543210 office', or a CSV / Excel sheet."""
    items: list[dict] = [r for r in (rows or []) if isinstance(r, dict)]
    for raw in str(text or "").splitlines():
        line = raw.strip(" •-\t")
        if not line:
            continue
        ph = re.search(r"(?:\+?91[ -]?)?[6-9]\d{4}[ -]?\d{5}", line)
        rest = line.replace(ph.group(0), " ") if ph else line
        parts = [p.strip() for p in re.split(r"[,|;]", rest) if p.strip()]
        nm = parts[0] if parts else ""
        items.append({"name": nm, "phone": ph.group(0) if ph else "", "group": parts[1] if len(parts) > 1 else ""})
    if file_id and data is None:
        name, data = _load_file(tenant, file_id)
    bad = []
    if data:
        rows2 = _rows(name, data)
        hdr = [_norm(h) for h in (rows2[0] if rows2 else [])]

        def col(*names):
            for i, h in enumerate(hdr):
                if any(n in h for n in names):
                    return i
            return None
        ci, cp, cg = col("name", "naam", "guest"), col("phone", "mobile", "number", "whatsapp"), col("group", "side", "company", "team")
        if ci is None:
            raise EventError("Sheet mein 'Name' column nahi mila.")
        for r in rows2[1:]:
            if ci < len(r) and r[ci]:
                items.append({"name": str(r[ci]), "phone": str(r[cp]) if cp is not None and cp < len(r) and r[cp] else "",
                              "group": str(r[cg]) if cg is not None and cg < len(r) and r[cg] else ""})
    added = 0
    with _LOCK, session() as s:
        e = _ev(s, tenant, event)
        have = {(g.name.lower(), g.phone) for g in s.scalars(select(EpGuest).where(EpGuest.event_id == e.id)).all()}
        n = len(have)
        for it in items:
            nm = R.clean(it.get("name"), 120)
            ph = R.phone(it.get("phone"))
            if not nm:
                continue
            if it.get("phone") and not ph:
                bad.append(f"{nm}: phone sahi nahi")
            if (nm.lower(), ph) in have:
                continue
            if n >= MAX_GUESTS:
                bad.append("guest list ki seema 3000")
                break
            s.add(EpGuest(tenant_id=tenant, event_id=e.id, name=nm, phone=ph, group=R.clean(it.get("group"), 60)))
            have.add((nm.lower(), ph))
            added += 1
            n += 1
    return {"added": added, "bad": bad, "say": f"{added} mehmaan jude." + (f" {len(bad)} line mein dikkat." if bad else "")}


def _guest(s, tenant: str, e: EpEvent, ref: str) -> EpGuest:
    g = s.get(EpGuest, ref) if ref else None
    if g is None and ref:
        rows = s.scalars(select(EpGuest).where(EpGuest.event_id == e.id)).all()
        low = ref.strip().lower()
        ph = R.phone(ref)
        cands = [x for x in rows if (ph and x.phone == ph) or x.name.lower() == low] or [x for x in rows if low in x.name.lower()]
        if len(cands) > 1:
            raise EventError(f"'{ref}' naam ke {len(cands)} mehmaan; phone ya poora naam bataiye.")
        g = cands[0] if cands else None
    if g is None or g.tenant_id != tenant or g.event_id != e.id:
        raise EventError(f"Mehmaan '{ref}' nahi mila.")
    return g


def rsvp(tenant: str, event: str, guest: str, words: str = "", status: str = "", plus=None, came=None) -> dict:
    """Her / the guest's reply ('haan +2', 'nahi aa payenge') -> status. Guest words are data."""
    st, pl = R.rsvp_of(words) if words else ("", None)
    if status:
        if status not in ("yes", "no", "maybe", ""):
            raise EventError("status = yes | no | maybe")
        st = status
    if plus is not None and R.num(plus) is not None:
        pl = int(R.num(plus) or 0)
    if not st and pl is None and came is None:
        raise EventError("Jawab samajh nahi aaya (haan / nahi / shayad, +N).")
    with _LOCK, session() as s:
        e = _ev(s, tenant, event)
        g = _guest(s, tenant, e, guest)
        if st:
            g.status = st
        if pl is not None:
            g.plus = max(0, min(pl, 20))
        if came is not None:
            g.came = bool(came)
        if words:
            g.note = R.clean(words, 300)
        v = _g_view(g)
    return {"guest": v, "say": f"{v['name']}: {v['status'] or '-'}" + (f" +{v['plus']}" if v["plus"] else "") + "."}


def guests(tenant: str, event: str, status: str = "") -> dict:
    with session() as s:
        e = _ev(s, tenant, event)
        rows = s.scalars(select(EpGuest).where(EpGuest.event_id == e.id).order_by(EpGuest.name)).all()
        gl = [_g_view(g) for g in rows]
        num = e.number
    yes = [g for g in gl if g["status"] == "yes"]
    sm = {"total": len(gl), "invited": sum(1 for g in gl if g["invited"]), "yes": len(yes), "heads_yes": len(yes) + sum(g["plus"] for g in yes),
          "no": sum(1 for g in gl if g["status"] == "no"), "maybe": sum(1 for g in gl if g["status"] == "maybe"),
          "pending": sum(1 for g in gl if g["invited"] and not g["status"]), "not_invited": sum(1 for g in gl if not g["invited"]),
          "came": sum(1 for g in gl if g["came"])}
    if status:
        gl = [g for g in gl if (g["status"] or "pending") == status]
    return {"event": num, "guests": gl, "summary": sm,
            "say": f"{num}: {sm['total']} mehmaan, {sm['yes']} haan (log {sm['heads_yes']}), {sm['no']} nahi, {sm['maybe']} shayad, "
                   f"{sm['pending']} jawab baaki. Caterer ko {sm['heads_yes']} (+ shayad {sm['maybe']}) bataiye."}


# ---------------------------------------------------------------- messages: invites / vendors / client (her tap)
def _invite_text(st: EpSettings, e: EpEvent) -> str:
    when = e.day or "[tareekh?]"
    t = (f"{e.start} baje" if e.start else "[samay?]")
    where = e.venue or "[jagah?]"
    return (f"Namaste {{name}} ji, aapko {e.name} mein amantrit karte hain: {when}, {t}, {where}. Kripya bataiye aap aa rahe hain "
            f"(haan / nahi, saath mein kitne log). - {st.firm or e.client or ''}").rstrip(" -")


def _msg_show(m: EpMessage) -> dict:
    return {"kind": m.kind, "text": m.text, "to": [[x.get("name"), x.get("phone")] for x in _j(m.to, [])]}


def _msg_view(s, m: EpMessage) -> dict:
    e = s.get(EpEvent, m.event_id)
    to = _j(m.to, [])
    return {"id": m.id, "event": e.number if e else "", "event_id": m.event_id, "kind": m.kind, "text": m.text, "count": len(to),
            "sample": [x.get("name") for x in to[:8]], "status": m.status, "hash": m.hash, "asked_by": m.asked_by}


def message_request(tenant: str, event: str, kind: str = "invite", text: str = "", group: str = "", to: str = "",
                    actor: str = "owner", conv: str = "") -> dict:
    """invite / reminder: to her guests with a phone (not opted out); vendor / client: one person. Only a card."""
    if kind not in ("invite", "reminder", "vendor", "client", "custom"):
        raise EventError("kind = invite | reminder | vendor | client | custom")
    with _LOCK, session() as s:
        e = _ev(s, tenant, event)
        st = _st(s, tenant)
        body = R.clean(text, 1200) or (_invite_text(st, e) if kind in ("invite", "custom") else "")
        if kind == "reminder" and not body:
            body = f"Namaste {{name}} ji, yaad dila rahe hain: {e.name} {e.day or '[tareekh?]'} {e.start or ''} {e.venue or ''}. Aap aa rahe hain na?"
        if not body:
            raise EventError("Message ke shabd chahiye.")
        blanks = R.blanks_in(body)
        if blanks:
            raise EventError("Message mein khaali jagah: " + ", ".join(blanks) + " (event ki details bhariye).")
        rec: list[dict] = []
        skipped = 0
        if kind in ("invite", "reminder", "custom") and not to:
            q = select(EpGuest).where(EpGuest.event_id == e.id)
            for g in s.scalars(q).all():
                if group and g.group.lower() != group.lower():
                    continue
                if kind == "invite" and g.invited:
                    continue
                if kind == "reminder" and (g.status in ("yes", "no") or not g.invited):
                    continue
                if g.no_messages or not g.phone:
                    skipped += 1
                    continue
                rec.append({"name": g.name, "phone": g.phone, "guest_id": g.id})
        elif kind == "client" or to == "client":
            if not e.client_phone:
                raise EventError("Client ka phone nahi.")
            rec = [{"name": e.client, "phone": e.client_phone}]
        else:
            v = _vendor(s, tenant, to)
            if not v.phone:
                raise EventError(f"{v.name} ka phone nahi.")
            rec = [{"name": v.name, "phone": v.phone, "vendor_id": v.id}]
        if not rec:
            raise EventError("Kisi ko bhejne layak nahi mila (phone nahi / pehle bhej chuke / mana kiya)." )
        m = EpMessage(tenant_id=tenant, event_id=e.id, kind=kind, text=body, to=_dump(rec), asked_by=actor_name(tenant, actor))
        s.add(m)
        s.flush()
        m.hash = R.digest(_msg_show(m))
        view = _msg_view(s, m)
    say = f"{view['count']} logon ko {kind} ka draft - aapki haan ka card bheja hai." + (f" {skipped} ko nahi (phone nahi / mana)." if skipped else "")
    _publish(tenant, "eventplan_message", {"message": view, "say": say}, conv)
    return {"message": view, "status": "needs_tap", "say": say}


def decide_message(tenant: str, message_id: str, yes: bool, card_hash: str = "", by_tap: bool = False) -> dict:
    """Her Haan -> one WhatsApp / SMS link per person (she sends each from her own phone). Nothing is sent by Hoola."""
    if not by_tap:
        raise EventError("Message sirf aapke tap se (app ka card).")
    with _LOCK, session() as s:
        m = s.get(EpMessage, message_id)
        if m is None or m.tenant_id != tenant:
            raise EventError("Card nahi mila.")
        if m.status != "waiting":
            return {"message": _msg_view(s, m), "say": f"Yeh card '{m.status}' hai."}
        if not card_hash or card_hash != m.hash or R.digest(_msg_show(m)) != m.hash:
            raise EventError("Card badal gaya hai; naya card dekhiye.")
        if not yes:
            m.status = "rejected"
            return {"message": _msg_view(s, m), "say": "Theek hai, nahi bheja."}
        m.status = "approved"
        links = []
        for x in _j(m.to, []):
            t = m.text.replace("{name}", x.get("name") or "")
            links.append({"name": x.get("name"), "phone": x.get("phone"), "text": t, "wa": _wa(x.get("phone", ""), t), "sms": _sms(x.get("phone", ""), t)})
            if m.kind == "invite" and x.get("guest_id"):
                g = s.get(EpGuest, x["guest_id"])
                if g is not None:
                    g.invited = True
        view = _msg_view(s, m)
    return {"message": view, "links": links, "say": f"{len(links)} draft link taiyaar: har ek aap apne phone se bhejiye."}


def message_sent(tenant: str, message_id: str) -> dict:
    with _LOCK, session() as s:
        m = s.get(EpMessage, message_id)
        if m is None or m.tenant_id != tenant:
            raise EventError("Card nahi mila.")
        if m.status != "approved":
            raise EventError("Pehle aapki haan (card).")
        m.status = "sent"
        v = _msg_view(s, m)
    return {"message": v, "say": "Bhej diya darj."}


def guest_optout(tenant: str, event: str, guest: str) -> dict:
    with _LOCK, session() as s:
        e = _ev(s, tenant, event)
        g = _guest(s, tenant, e, guest)
        g.no_messages = True
        nm = g.name
    return {"say": f"{nm} ko aage koi message nahi."}


# ---------------------------------------------------------------- tasks
def _t_view(t: EpTask) -> dict:
    return {"id": t.id, "event_id": t.event_id, "what": t.what, "area": t.area, "owner": t.owner, "due": t.due, "status": t.status, "auto": t.auto}


def tasks(tenant: str, event: str) -> dict:
    with session() as s:
        e = _ev(s, tenant, event)
        rows = [_t_view(t) for t in s.scalars(select(EpTask).where(EpTask.event_id == e.id).order_by(EpTask.due)).all()]
        on = today(tenant).isoformat()
    late = [t for t in rows if t["status"] == "open" and t["due"] and t["due"] < on]
    no_owner = [t for t in rows if t["status"] == "open" and not t["owner"]]
    return {"tasks": rows, "late": len(late), "no_owner": len(no_owner),
            "say": f"{len(rows)} kaam, {len(late)} der se, {len(no_owner)} ka koi zimmedar nahi."}


def task_add(tenant: str, event: str, what: str, owner: str = "", due: str = "") -> dict:
    what = R.clean(what, 200)
    if not what:
        raise EventError("Kaam kya hai?")
    d = _date(due)
    if due and not d:
        raise EventError("Tareekh samajh nahi aayi.")
    with _LOCK, session() as s:
        e = _ev(s, tenant, event)
        if d is None and _date(e.day):
            d = _date(e.day)
        t = EpTask(tenant_id=tenant, event_id=e.id, what=what, owner=R.clean(owner, 80), due=_iso(d))
        s.add(t)
        s.flush()
        v = _t_view(t)
    return {"task": v, "say": f"Kaam juda: {what}" + (f" ({v['owner']}, {v['due']})" if v["owner"] or v["due"] else "") + "."}


def task_update(tenant: str, task_id: str, status: str = "", owner: str = "", due: str = "") -> dict:
    with _LOCK, session() as s:
        t = s.get(EpTask, task_id)
        if t is None or t.tenant_id != tenant:
            raise EventError("Kaam nahi mila.")
        if status:
            if status not in ("open", "done", "skipped"):
                raise EventError("status = open | done | skipped")
            t.status = status
        if owner:
            t.owner = R.clean(owner, 80)
        if due:
            d = _date(due)
            if not d:
                raise EventError("Tareekh samajh nahi aayi.")
            t.due = d.isoformat()
        v = _t_view(t)
    return {"task": v, "say": f"{v['what']}: {v['status']}."}


# ---------------------------------------------------------------- permissions
def permits(tenant: str, event: str) -> dict:
    with session() as s:
        e = _ev(s, tenant, event)
        rows = _j(e.permits, [])
        d = _date(e.day)
        num = e.number
    on = today(tenant)
    for p in rows:
        info = R.PERMITS.get(p["key"], {})
        p["why"] = info.get("why", p.get("note", ""))
        p["who"] = info.get("who", "")
        lead = info.get("lead", 10)
        p["apply_by"] = (d - timedelta(days=lead)).isoformat() if d else ""
        p["late"] = bool(d and p["status"] in ("pending",) and d - timedelta(days=lead) < on)
    open_ = [p for p in rows if p["status"] in ("pending", "applied")]
    return {"event": num, "permits": rows, "note": "Yeh starter list hai: aapke shahar / thane / venue ke niyam faisla karte hain.",
            "say": f"{num}: {len(rows)} permission, {len(open_)} baaki" + (f", {sum(1 for p in rows if p['late'])} der se." if any(p['late'] for p in rows) else ".")}


def permit_update(tenant: str, event: str, key: str, status: str = "", ref: str = "", note: str = "", label: str = "") -> dict:
    if status and status not in ("pending", "applied", "granted", "refused", "not_needed"):
        raise EventError("status = pending | applied | granted | refused | not_needed")
    with _LOCK, session() as s:
        e = _ev(s, tenant, event)
        rows = _j(e.permits, [])
        p = next((x for x in rows if x["key"] == key or x["label"].lower() == str(key).lower()), None)
        if p is None:
            if not label and key not in R.PERMITS:
                raise EventError("Yeh permission list mein nahi; naya jodne ke liye label dijiye.")
            p = {"key": key if key in R.PERMITS else f"c{len(rows) + 1}", "label": R.PERMITS.get(key, {}).get("label") or R.clean(label, 120),
                 "status": "pending", "ref": "", "note": "", "custom": key not in R.PERMITS}
            rows.append(p)
        if status:
            p["status"] = status
        if ref:
            p["ref"] = R.clean(ref, 80)
        if note:
            p["note"] = R.clean(note, 300)
        e.permits = _dump(rows)
        lbl, stt = p["label"], p["status"]
    return {"say": f"{lbl}: {stt}."}


# ---------------------------------------------------------------- run of show
def show_set(tenant: str, event: str, text: str = "", cues: list | None = None, replace: bool = True) -> dict:
    parsed, bad = R.parse_cues(text) if text else ([], [])
    for c in cues or []:
        if isinstance(c, dict) and R.hhmm(str(c.get("time") or "")):
            parsed.append({"time": R.hhmm(str(c["time"])), "mins": int(R.num(c.get("mins")) or 0) or None, "what": R.clean(c.get("what"), 160),
                           "who": R.clean(c.get("who"), 60), "where": R.clean(c.get("where"), 60),
                           "av": [a for a in (c.get("av") or []) if a in R._AV_WORDS]})
    if not parsed:
        raise EventError("Koi cue nahi samjha. Har line samay se shuru: '18:30 Welcome (10 min) - Riya | stage'.")
    with _LOCK, session() as s:
        e = _ev(s, tenant, event)
        cur = [] if replace else _j(e.cues, [])
        cur = sorted(cur + parsed, key=lambda c: c["time"])
        e.cues = _dump(cur)
    out = show(tenant, event)
    if bad:
        out["bad"] = bad
        out["say"] += f" {len(bad)} line mein samay nahi tha (chhodi)."
    return out


def show(tenant: str, event: str) -> dict:
    with session() as s:
        e = _ev(s, tenant, event)
        st = _st(s, tenant)
        cues = _j(e.cues, [])
        arrivals = [{"time": b.arrive, "what": f"{(s.get(EpVendor, b.vendor_id) or EpVendor(name='?')).name} pahunche "
                                              f"({R.HEADS.get(b.head, R.HEADS['misc'])['label']})", "who": "vendor", "kind": "arrival"}
                    for b in s.scalars(select(EpBooking).where(EpBooking.event_id == e.id, EpBooking.status.in_(("approved", "advance_paid")))).all()
                    if b.arrive]
        sound_till = st.sound_till or "22:00"
        num = e.number
    problems = R.cue_problems(cues, sound_till)
    timeline = sorted([{**c, "kind": "cue"} for c in cues] + arrivals, key=lambda c: c["time"])
    return {"event": num, "cues": cues, "timeline": timeline, "problems": problems, "sound_till": sound_till,
            "say": f"{num}: {len(cues)} cue" + (f", {len(problems)} dikkat: " + "; ".join(problems[:3]) if problems else ", sab theek.")}


def show_pdf(tenant: str, event: str) -> bytes:
    v = show(tenant, event)
    with session() as s:
        e = _ev(s, tenant, event)
        st = _st(s, tenant)
        blocks = _head(st, f"Run of show: {e.name}")
        blocks.append(("p", f"{e.number} · {e.day} · {e.venue} · {e.headcount or '-'} guests"))
    rows = [["Time", "Min", "What", "Who", "Where", "AV"]]
    for c in v["timeline"]:
        rows.append([c["time"], str(c.get("mins") or ""), c["what"], c.get("who", ""), c.get("where", ""), ", ".join(c.get("av") or [])])
    blocks.append(("table", rows, [40, 30, 200, 80, 80, 70]))
    if v["problems"]:
        blocks.append(("h", "Check"))
        for p in v["problems"]:
            blocks.append(("p", p))
    return _pdf(f"Run of show {e.number}", blocks)


def show_file(tenant: str, event: str, conv: str = "") -> dict:
    pdf = show_pdf(tenant, event)
    with session() as s:
        num = _ev(s, tenant, event).number
    f = _store(tenant, f"{num}_run_of_show.pdf", pdf, "application/pdf", conv, f"{num} run of show")
    return {"file": f, "say": "Cue sheet PDF bani."}


# ---------------------------------------------------------------- post-event report
def report(tenant: str, event: str, attended=None, notes: str = "", issues: list | None = None, conv: str = "",
           actor: str = "owner") -> dict:
    b = budget(tenant, event)
    g = guests(tenant, event)["summary"]
    tk = tasks(tenant, event)
    with _LOCK, session() as s:
        e = _ev(s, tenant, event)
        d = _date(e.day)
        if not d or d > today(tenant):
            raise EventError("Report event ke baad.")
        att = R.num(attended)
        rep = {"attended": int(att) if att is not None else (g["came"] or None), "rsvp_yes": g["heads_yes"], "invited": g["invited"],
               "plan_cost": b["cost"], "actual": b["actual"], "client": b["client"] + b["fee"], "margin": b["margin"],
               "heads": [{"label": h["label"], "plan": h["cost"], "actual": h["actual"]} for h in b["heads"]],
               "tasks_late": tk["late"], "notes": R.clean(notes, 2000), "issues": [R.clean(x, 200) for x in (issues or [])][:20],
               "client_paid": R.r2(sum(x["amount"] for x in _j(e.client_paid, [])))}
        e.report = _dump(rep)
        if e.status not in ("cancelled",):
            e.status = "done"
        _log(e, "post-event report", actor_name(tenant, actor))
        st = _st(s, tenant)
        blocks = _head(st, f"Post-event report: {e.name}")
        blocks.append(("p", f"{e.number} · {e.day} · {e.venue} · client {e.client}"))
        blocks.append(("p", f"Guests: invited {rep['invited']}, RSVP yes (with plus-ones) {rep['rsvp_yes']}, attended {rep['attended'] or '-'}"))
        rows = [["Head", "Plan (cost)", "Actual"]] + [[h["label"], _pdf_rs(h["plan"]), _pdf_rs(h["actual"])] for h in rep["heads"]]
        blocks.append(("table", rows, [200, 120, 120]))
        blocks.append(("p", f"Planned cost {_pdf_rs(rep['plan_cost'])} · actual booked {_pdf_rs(rep['actual'])} · client value "
                       f"{_pdf_rs(rep['client'])} (before GST) · received {_pdf_rs(rep['client_paid'])}"))
        if rep["issues"]:
            blocks.append(("h", "Issues"))
            for x in rep["issues"]:
                blocks.append(("p", "- " + x))
        if rep["notes"]:
            blocks.append(("h", "Notes"))
            blocks.append(("p", rep["notes"]))
        num = e.number
    pdf = _pdf(f"Report {num}", blocks)
    f = _store(tenant, f"{num}_report.pdf", pdf, "application/pdf", conv, f"{num} post-event report")
    say = (f"{num} report: aaye {rep['attended'] or '?'} / RSVP {rep['rsvp_yes']}, kharcha {_rs(rep['actual'])} (plan {_rs(rep['plan_cost'])}), "
           f"margin {_rs(rep['margin'])}. Invoice ke liye 'invoice' boliye.")
    return {"report": rep, "file": f, "say": say}


def forget(tenant: str, event: str, actor: str = "owner") -> dict:
    """DPDP: after the event, guests' and the client's personal details go; numbers / amounts stay for accounts."""
    if not is_owner(actor):
        raise EventError("Sirf owner.")
    with _LOCK, session() as s:
        e = _ev(s, tenant, event)
        if e.status not in ("done", "cancelled"):
            raise EventError("Event poora / radd hone ke baad hi.")
        n = 0
        for g in s.scalars(select(EpGuest).where(EpGuest.event_id == e.id)).all():
            s.delete(g)
            n += 1
        for m in s.scalars(select(EpMessage).where(EpMessage.event_id == e.id)).all():
            m.to, m.text = "[]", ""
        e.client, e.client_phone, e.client_email, e.forgotten = "Client (hataya)", "", "", True
        num = e.number
    return {"say": f"{num}: {n} mehmaanon aur client ki personal details hata di (hisaab rakha)."}


# ---------------------------------------------------------------- views / day / worker
def event_view(tenant: str, ref: str) -> dict:
    with session() as s:
        e = _ev(s, tenant, ref)
        v = _ev_view(s, e, full=True)
        v["docs"] = [_doc_view(s, d) for d in s.scalars(select(EpDoc).where(EpDoc.event_id == e.id).order_by(EpDoc.created_at)).all()]
        v["vquotes"] = [_vq_view(s, q) for q in s.scalars(select(EpVendorQuote).where(EpVendorQuote.event_id == e.id)).all()]
        eid = e.id
    v["budget"] = budget(tenant, eid)
    v["bookings"] = bookings(tenant, eid)
    v["spends"] = spends(tenant, eid)
    v["guests"] = guests(tenant, eid)
    v["tasks"] = tasks(tenant, eid)
    v["permits"] = permits(tenant, eid)
    v["show"] = show(tenant, eid)
    return v


def waiting(tenant: str) -> dict:
    with session() as s:
        ds = [_doc_view(s, d) for d in s.scalars(select(EpDoc).where(EpDoc.tenant_id == tenant, EpDoc.status == "waiting")).all()]
        bs = [_bk_view(s, b) for b in s.scalars(select(EpBooking).where(EpBooking.tenant_id == tenant, EpBooking.status == "waiting")).all()]
        ss = [_sp_view(s, x) for x in s.scalars(select(EpSpend).where(EpSpend.tenant_id == tenant, EpSpend.status == "waiting")).all()]
        ms = [_msg_view(s, m) for m in s.scalars(select(EpMessage).where(EpMessage.tenant_id == tenant, EpMessage.status == "waiting")).all()]
    return {"docs": ds, "bookings": bs, "spends": ss, "messages": ms}


def day(tenant: str) -> dict:
    on = today(tenant)
    soon = (on + timedelta(days=7)).isoformat()
    with session() as s:
        evs = {e.id: e for e in s.scalars(select(EpEvent).where(EpEvent.tenant_id == tenant, EpEvent.status.notin_(("done", "cancelled")))).all()}
        upcoming = sorted([(e.day, e.number, e.name) for e in evs.values() if e.day and on.isoformat() <= e.day <= soon])
        late = [(evs[t.event_id].number, t.what, t.due) for t in s.scalars(select(EpTask).where(EpTask.tenant_id == tenant, EpTask.status == "open")).all()
                if t.event_id in evs and t.due and t.due < on.isoformat()]
        due_today = [(evs[t.event_id].number, t.what, t.owner) for t in s.scalars(select(EpTask).where(EpTask.tenant_id == tenant, EpTask.status == "open")).all()
                     if t.event_id in evs and t.due == on.isoformat()]
        adv = [(_bk_view(s, b)) for b in s.scalars(select(EpBooking).where(EpBooking.tenant_id == tenant, EpBooking.status == "approved")).all()
               if b.advance > 0 and b.advance_due and b.advance_due <= on.isoformat()]
        ids = list(evs)
    permit_late = []
    for eid in ids:
        for p in permits(tenant, eid)["permits"]:
            if p["late"]:
                permit_late.append(f"{p['label']}")
    w = waiting(tenant)
    parts = []
    if upcoming:
        parts.append("Is hafte: " + "; ".join(f"{n} {nm} ({d})" for d, n, nm in upcoming[:5]))
    if due_today:
        parts.append("Aaj ka kaam: " + "; ".join(f"{n} {w_}" + (f" ({o})" if o else "") for n, w_, o in due_today[:6]))
    if late:
        parts.append("Der se: " + "; ".join(f"{n} {w_}" for n, w_, _d in late[:6]))
    if permit_late:
        parts.append("Permission ki arzi der se: " + ", ".join(permit_late[:5]))
    if adv:
        parts.append("Advance dene ki tareekh: " + "; ".join(f"{b['vendor']} {b['advance_text']}" for b in adv[:4]))
    nw = sum(len(v) for v in w.values())
    if nw:
        parts.append(f"{nw} card aapki haan ka intezaar kar rahe")
    return {"events": len(evs), "upcoming": upcoming, "due_today": due_today, "late": late, "permits_late": permit_late, "advances": adv,
            "waiting": w, "say": (". ".join(parts) + ".") if parts else "Events ka aaj sab theek."}


def overview(tenant: str) -> dict:
    return {"settings": settings(tenant), "events": events(tenant), "vendors": vendors(tenant), "day": day(tenant),
            "heads": {k: v["label"] for k, v in R.HEADS.items()}, "kinds": {k: v["label"] for k, v in R.KINDS.items()}}


def _once(s, tenant: str, key: str) -> bool:
    k = f"{tenant}:{key}"[:200]
    if s.get(EpAlert, k) is not None:
        return False
    s.add(EpAlert(key=k, tenant_id=tenant))
    return True


def _tenants() -> list[str]:
    with session() as s:
        return sorted(set(s.scalars(select(EpEvent.tenant_id)).all()) | set(s.scalars(select(EpSettings.tenant_id)).all()))


def tick(when: datetime | None = None) -> dict:
    """Worker, every minute: overdue tasks / late permission arzi / event tomorrow (once each), morning note at her time."""
    n = _aware(when) or now()
    out = {"alerts": 0, "morning": 0}
    for tenant in _tenants():
        try:
            loc = n.astimezone(_tz(tenant))
            on = loc.date()
            alerts = []
            with _LOCK, session() as s:
                evs = {e.id: e for e in s.scalars(select(EpEvent).where(EpEvent.tenant_id == tenant,
                                                                         EpEvent.status.notin_(("done", "cancelled")))).all()}
                for t in s.scalars(select(EpTask).where(EpTask.tenant_id == tenant, EpTask.status == "open")).all():
                    if t.event_id in evs and t.due and t.due < on.isoformat() and _once(s, tenant, f"late:{t.id}:{t.due}"):
                        alerts.append({"kind": "late", "say": f"{evs[t.event_id].number}: '{t.what}' {t.due} tak hona tha"
                                                              + (f" ({t.owner})" if t.owner else " (koi zimmedar nahi)") + "."})
                for e in evs.values():
                    d = _date(e.day)
                    if d and d == on + timedelta(days=1) and _once(s, tenant, f"eve:{e.id}:{e.day}"):
                        alerts.append({"kind": "tomorrow", "say": f"Kal {e.number} {e.name}: cue sheet, vendor timings aur emergency number dekh lijiye."})
                    if d:
                        for p in _j(e.permits, []):
                            lead = R.PERMITS.get(p["key"], {}).get("lead", 10)
                            if p["status"] == "pending" and d - timedelta(days=lead) <= on and _once(s, tenant, f"permit:{e.id}:{p['key']}"):
                                alerts.append({"kind": "permit", "say": f"{e.number}: '{p['label']}' ki arzi ka samay aa gaya ({R.PERMITS.get(p['key'], {}).get('who', '')})."})
            for a in alerts:
                _publish(tenant, "eventplan_alert", a)
                out["alerts"] += 1
            with session() as s:
                r = s.get(EpSettings, tenant)
                due_m = False
                if r is not None and r.last_morning != on.isoformat():
                    try:
                        hh, mm = (int(z) for z in (r.morning_time or "08:30").split(":"))
                    except ValueError:
                        hh, mm = 8, 30
                    due_m = (loc.hour, loc.minute) >= (hh, mm)
                    if due_m:
                        r.last_morning = on.isoformat()
            if due_m:
                d = day(tenant)
                if d["say"] != "Events ka aaj sab theek.":
                    _publish(tenant, "eventplan_day", {"say": d["say"], "events": d["events"], "late": len(d["late"])})
                    out["morning"] += 1
        except Exception:
            log.exception("eventplan tick failed for %s", tenant)
    return out
