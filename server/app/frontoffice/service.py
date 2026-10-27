"""Front Office Sahayak: Hoola as a small / mid hotel's front desk (hotel, lodge, guest house).

  1. Room types (HER rates: weekday / weekend nights / seasons, extra adult) + rooms; housekeeping status board
     (dirty / clean / inspected / out_of_order). Out of order never makes a future night overbooked (refused).
  2. Availability calendar per room type per night = rooms - out of order - held nights. A booking (walk-in, phone, WhatsApp,
     agent, or HER pasted OTA mail / exported sheet: MakeMyTrip, Goibibo, Booking.com, Agoda, Airbnb...) that would sell a night
     twice is REFUSED in code; a pasted OTA booking that doesn't fit becomes an "overbook" alert (stop-sell / move the guest).
  3. Check-in: guest register (IDs straight into the Pehchaan Locker; only {id:...} + masked kept), room assignment (free,
     not out of order, clean first; dirty only when she says), foreign guest -> passport + visa refs + Form III (C-Form) arrival
     entry due in 24 h (she submits on indianfrro.gov.in), departure entry 24 h after check-out.
  4. Folio: room nights at HER rates, extras (restaurant / room service 5% or 18% by 'specified premises', laundry, extra bed ...),
     payments she marks received. GST bill: room value per unit per night <= Rs 7,500 -> 5%, above -> 18% (from 22 Sep 2025),
     only when her GSTIN is set. Bill becomes final ONLY on her tap (INV/2627/0001 + PDF).
  5. Advance / balance = UPI links to HER UPI ID (the guest pays in their own app). Refunds, discounts, complimentary stays,
     cancellation charges and bill finalisation = cards that only HER tap decides (refund: fingerprint once a phone key is set).
  6. Confirmation / pre-arrival / payment / review-request / bill / cancellation messages = DRAFTS; only her tap sends them.
  7. Daily report: occupancy, ADR, RevPAR, by source; morning note (arrivals, departures, Form C due, holds, dirty rooms).
  Never: overbooking, plain ID numbers, OTA logins / scraping, moving money, inventing rates / availability / dates."""
from __future__ import annotations

import csv
import hashlib
import io
import json
import logging
import re
import secrets
import threading
from datetime import date, datetime, timedelta, timezone
from pathlib import Path
from typing import Any
from urllib.parse import quote

from sqlalchemy import select

from ..db import now, session
from . import rules as R
from .models import (
    FoAlert,
    FoBill,
    FoBooking,
    FoCard,
    FoCharge,
    FoFormC,
    FoGuest,
    FoMessage,
    FoPayment,
    FoRoom,
    FoRoomType,
    FoSettings,
)

log = logging.getLogger("santosh.frontoffice")
_LOCK = threading.RLock()
HOLDING = ("hold", "confirmed", "in_house")
STATUS_LABEL = {"hold": "Hold (advance baaki)", "confirmed": "Pakki", "in_house": "Ruke hue", "checked_out": "Chale gaye",
                "cancelled": "Radd", "no_show": "Nahi aaye"}
CARD_LABEL = {"discount": "Discount", "comp": "Complimentary (muft)", "cancel_fee": "Cancellation charge", "refund": "Refund",
              "bill": "Bill final"}
MONEY_KINDS = ("refund",)
FRRO_URL = "https://indianfrro.gov.in/frro/FormC"


class FoError(Exception):
    pass


# ---------------------------------------------------------------- helpers
def _j(s: str, default):
    try:
        v = json.loads(s or "")
        return v if isinstance(v, type(default)) else default
    except (TypeError, ValueError):
        return default


def _dump(v) -> str:
    return json.dumps(v, ensure_ascii=False, default=str, sort_keys=True)


def _aware(d: datetime | None) -> datetime | None:
    if d is None:
        return None
    return d if d.tzinfo else d.replace(tzinfo=timezone.utc)


def clock() -> datetime:
    return now()


def _tz(tenant: str):
    from ..life import tz_of
    return tz_of(tenant)


def today(tenant: str) -> date:
    return clock().astimezone(_tz(tenant)).date()


def _date(v, tenant: str = "") -> date | None:
    return R.parse_date(v, today(tenant) if tenant else None)


def _rs(v: float | None) -> str:
    if v is None:
        return "[daam?]"
    try:
        from ..india.money import inr
        return inr(v)
    except Exception:
        return f"₹{v:,.0f}"


def _r2(v: float) -> float:
    return round(float(v or 0) + 1e-9, 2)


def _hash(v) -> str:
    return hashlib.sha256(_dump(v).encode()).hexdigest()


def _publish(tenant: str, kind: str, payload: dict, conv: str = "") -> None:
    from ..events import bus
    try:
        if not conv:
            from ..net.service import home_conv
            conv = home_conv(tenant)
        bus.publish(conv, kind, payload)
    except Exception:
        log.exception("frontoffice publish failed")


def is_owner(actor: str) -> bool:
    a = actor or "owner"
    return a == "owner" or a.startswith("approved:")


def actor_name(tenant: str, actor: str) -> str:
    if is_owner(actor):
        return "Owner"
    try:
        from ..employee.service import actor_label
        return actor_label(tenant, actor)
    except Exception:
        return str(actor)[:40]


def _owner_only(actor: str, what: str) -> None:
    if not is_owner(actor):
        raise FoError(f"{what} sirf owner kar sakti hain.")


def _st(s, tenant: str) -> FoSettings:
    r = s.get(FoSettings, tenant)
    if r is None:
        r = FoSettings(tenant_id=tenant)
        s.add(r)
        s.flush()
    return r


def _next(s, tenant: str, key: str) -> int:
    st = _st(s, tenant)
    seq = _j(st.seq, {})
    seq[key] = int(seq.get(key, 0)) + 1
    st.seq = _dump(seq)
    return seq[key]


def _clean(text: Any, n: int = 500) -> str:
    return R.scrub(text, n)


def _once(s, tenant: str, key: str) -> bool:
    k = f"{tenant}:{key}"[:200]
    if s.get(FoAlert, k) is not None:
        return False
    s.add(FoAlert(key=k, tenant_id=tenant))
    s.flush()
    return True


# ---------------------------------------------------------------- settings
_SET_KEYS = ("hotel_name", "address", "city", "state", "phone", "email", "gstin", "specified_premises", "extra_gst", "upi_id",
             "upi_name", "check_in_time", "check_out_time", "weekend_days", "cancel_policy", "hold_hours", "review_link",
             "morning_time", "rules")


def _settings_view(r: FoSettings) -> dict:
    pol = _j(r.cancel_policy, {})
    return {"hotel_name": r.hotel_name, "address": r.address, "city": r.city, "state": r.state, "phone": r.phone, "email": r.email,
            "gstin": r.gstin, "specified_premises": r.specified_premises, "extra_gst": _j(r.extra_gst, {}), "upi_id": r.upi_id,
            "upi_name": r.upi_name, "check_in_time": r.check_in_time, "check_out_time": r.check_out_time,
            "weekend_days": _j(r.weekend_days, [4, 5]), "cancel_policy": pol, "cancel_policy_text": R.policy_text(pol),
            "hold_hours": r.hold_hours, "review_link": r.review_link, "morning_time": r.morning_time, "rules": r.rules,
            "gst_note": ("GST: kamra Rs 7,500 / raat tak 5% (ITC nahi), upar 18%; restaurant "
                         + ("18% (specified premises)." if r.specified_premises else "5% (ITC nahi).")) if r.gstin
            else "GSTIN set nahi: bill par GST nahi lagega (registered hain to Niyam mein GSTIN daaliye)."}


def _upi(r: FoSettings) -> tuple[str, str]:
    if r.upi_id:
        return r.upi_id, r.upi_name or r.hotel_name
    try:
        from ..employee import service as emp
        p = emp.profile(r.tenant_id) or {}
        return p.get("upi_id") or "", p.get("upi_name") or p.get("business_name") or r.hotel_name
    except Exception:
        return "", ""


def settings(tenant: str) -> dict:
    with session() as s:
        return _settings_view(_st(s, tenant))


def _hhmm(v: Any) -> str:
    m = re.fullmatch(r"(\d{1,2}):(\d{2})", str(v).strip())
    if not m or int(m.group(1)) > 23 or int(m.group(2)) > 59:
        raise FoError("Samay HH:MM mein (jaise 11:00).")
    return f"{int(m.group(1)):02d}:{m.group(2)}"


def save_settings(tenant: str, actor: str = "owner", **kw: Any) -> dict:
    _owner_only(actor, "Hotel ki settings")
    with _LOCK, session() as s:
        r = _st(s, tenant)
        for k in _SET_KEYS:
            if k not in kw or kw[k] is None:
                continue
            v = kw[k]
            if k in ("check_in_time", "check_out_time", "morning_time"):
                setattr(r, k, _hhmm(v))
            elif k == "specified_premises":
                r.specified_premises = bool(v)
            elif k == "hold_hours":
                r.hold_hours = max(1, min(168, int(R.money(v)) or 24))
            elif k == "weekend_days":
                days = v if isinstance(v, list) else re.split(r"[,\s]+", str(v))
                names = {"mon": 0, "tue": 1, "wed": 2, "thu": 3, "fri": 4, "sat": 5, "sun": 6}
                out = sorted({names.get(str(d).lower()[:3], int(R.money(d)) if str(d).strip().isdigit() else -1) for d in days} - {-1})
                r.weekend_days = _dump([d for d in out if 0 <= d <= 6])
            elif k == "cancel_policy":
                p = v if isinstance(v, dict) else {}
                fee = str(p.get("fee") or "first_night")
                if fee not in ("first_night", "pct", "full"):
                    raise FoError("Cancellation fee: first_night | pct | full")
                r.cancel_policy = _dump({"free_days": max(0, int(R.money(p.get("free_days")))), "fee": fee,
                                         "pct": max(0.0, min(100.0, R.money(p.get("pct")))) if fee == "pct" else 0})
            elif k == "extra_gst":
                d = v if isinstance(v, dict) else {}
                r.extra_gst = _dump({str(a): int(R.money(b)) for a, b in d.items() if a in R.EXTRA_KINDS and R.money(b) in (0, 5, 12, 18, 28, 40)})
            elif k == "gstin":
                g = str(v).strip().upper()
                if g:
                    from ..india.money import gstin_ok
                    if not gstin_ok(g):
                        raise FoError("GSTIN sahi nahi (check digit).")
                r.gstin = g
            elif k == "phone":
                r.phone = R.phone10(v) or ""
            elif k == "review_link":
                u = str(v).strip()
                if u and not u.startswith("https://"):
                    raise FoError("Review link https:// se shuru hona chahiye.")
                r.review_link = u[:300]
            else:
                setattr(r, k, _clean(v, 2000))
        view = _settings_view(r)
    return {"settings": view, "say": "Hotel ki settings save."}


# ---------------------------------------------------------------- room types / rooms
_TYPE_KEYS = ("code", "name", "aliases", "max_adults", "base", "weekend", "seasons", "extra_adult", "extra_child", "meal_plan", "archived")


def _type(s, tenant: str, ref: str) -> FoRoomType:
    ref = str(ref or "").strip()
    rows = s.scalars(select(FoRoomType).where(FoRoomType.tenant_id == tenant, FoRoomType.archived.is_(False))).all()
    for t in rows:
        if ref in (t.id, t.code) or ref.lower() in (t.name.lower(), t.code.lower()):
            return t
    hit = R.match_room_type(ref, [{"id": t.id, "name": t.name, "code": t.code, "aliases": t.aliases} for t in rows])
    if hit:
        return s.get(FoRoomType, hit["id"])
    names = ", ".join(f"{t.code} {t.name}" for t in rows) or "koi room type nahi (pehle jodiye)"
    raise FoError(f"Room type '{ref}' nahi mila. Hain: {names}.")


def _type_view(s, t: FoRoomType) -> dict:
    rooms = s.scalars(select(FoRoom).where(FoRoom.type_id == t.id, FoRoom.archived.is_(False)).order_by(FoRoom.number)).all()
    return {"id": t.id, "code": t.code, "name": t.name, "aliases": t.aliases, "max_adults": t.max_adults, "base": t.base or None,
            "weekend": t.weekend or None, "seasons": _j(t.seasons, []), "extra_adult": t.extra_adult or None,
            "extra_child": t.extra_child or None, "meal_plan": t.meal_plan, "rooms": [r.number for r in rooms],
            "rate_missing": not t.base}


def save_type(tenant: str, type_ref: str = "", actor: str = "owner", **kw: Any) -> dict:
    """A room type with HER rates. Rates are never guessed; a type without a base rate shows a blank on quotes."""
    _owner_only(actor, "Room type / rate")
    with _LOCK, session() as s:
        if type_ref:
            t = _type(s, tenant, type_ref)
        else:
            name = _clean(kw.get("name"), 80)
            if not name:
                raise FoError("Room type ka naam chahiye (Deluxe / Standard ...).")
            t = FoRoomType(tenant_id=tenant, name=name)
            s.add(t)
        for k in _TYPE_KEYS:
            if k not in kw or kw[k] is None:
                continue
            v = kw[k]
            if k in ("base", "weekend", "extra_adult", "extra_child"):
                setattr(t, k, max(0.0, R.money(v)))
            elif k == "max_adults":
                t.max_adults = max(1, min(12, int(R.money(v)) or 2))
            elif k == "seasons":
                ss = []
                for x in (v if isinstance(v, list) else []):
                    if not isinstance(x, dict) or R.money(x.get("rate")) <= 0:
                        continue
                    f, to = str(x.get("from") or ""), str(x.get("to") or "")
                    if not ((re.fullmatch(r"\d{4}-\d{2}-\d{2}", f) and re.fullmatch(r"\d{4}-\d{2}-\d{2}", to)) or
                            (R._mmdd(f) and R._mmdd(to))):
                        raise FoError("Season ki tareekh YYYY-MM-DD ya DD-MM (har saal).")
                    ss.append({"name": _clean(x.get("name"), 40), "from": f, "to": to, "rate": R.money(x.get("rate"))})
                t.seasons = _dump(ss[:20])
            elif k == "meal_plan":
                t.meal_plan = str(v).upper() if str(v).upper() in R.MEAL_PLANS else "EP"
            elif k == "archived":
                t.archived = bool(v)
            elif k == "code":
                t.code = re.sub(r"[^A-Za-z0-9]", "", str(v)).upper()[:12]
            else:
                setattr(t, k, _clean(v, 400))
        if not t.code:
            t.code = re.sub(r"[^A-Za-z]", "", t.name).upper()[:3] or "RM"
        s.flush()
        view = _type_view(s, t)
    say = f"{view['code']} {view['name']} save." + ("" if view["base"] else " Rate abhi nahi: quote par khaali rahega.")
    return {"room_type": view, "say": say}


def types(tenant: str) -> list[dict]:
    with session() as s:
        rows = s.scalars(select(FoRoomType).where(FoRoomType.tenant_id == tenant, FoRoomType.archived.is_(False))
                         .order_by(FoRoomType.code)).all()
        return [_type_view(s, t) for t in rows]


def add_rooms(tenant: str, type_ref: str, numbers: Any, floor: str = "", actor: str = "owner") -> dict:
    """'101-105, 201' -> rooms of that type. A number already used is refused."""
    _owner_only(actor, "Kamre jodna")
    raw = numbers if isinstance(numbers, list) else re.split(r"[,\s]+", str(numbers or ""))
    nums: list[str] = []
    for x in raw:
        x = str(x).strip()
        m = re.fullmatch(r"(\d+)-(\d+)", x)
        if m and int(m.group(2)) >= int(m.group(1)) and int(m.group(2)) - int(m.group(1)) < 200:
            nums += [str(i) for i in range(int(m.group(1)), int(m.group(2)) + 1)]
        elif re.fullmatch(r"[A-Za-z0-9-]{1,12}", x):
            nums.append(x.upper())
    if not nums:
        raise FoError("Room number chahiye (jaise '101-105, 201').")
    with _LOCK, session() as s:
        t = _type(s, tenant, type_ref)
        have = set(s.scalars(select(FoRoom.number).where(FoRoom.tenant_id == tenant, FoRoom.archived.is_(False))).all())
        dup = [n for n in nums if n in have]
        if dup:
            raise FoError("Yeh room number pehle se hain: " + ", ".join(dup[:10]))
        for n in nums:
            s.add(FoRoom(tenant_id=tenant, number=n, type_id=t.id, floor=_clean(floor, 12)))
        s.flush()
        view = _type_view(s, t)
    return {"room_type": view, "added": nums, "say": f"{len(nums)} kamre {view['name']} mein jode: {', '.join(nums[:12])}."}


def _room(s, tenant: str, ref: str) -> FoRoom:
    ref = str(ref or "").strip().upper()
    r = s.scalars(select(FoRoom).where(FoRoom.tenant_id == tenant, FoRoom.archived.is_(False),
                                       (FoRoom.number == ref) | (FoRoom.id == ref.lower()) | (FoRoom.id == ref))).first()
    if r is None:
        raise FoError(f"Room {ref} nahi mila.")
    return r


def _ooo_on(r: FoRoom, d: date) -> bool:
    if r.status != "out_of_order":
        return False
    f = R.parse_date(r.ooo_from) or date.min
    u = R.parse_date(r.ooo_until)
    return f <= d and (u is None or d <= u)


# ---------------------------------------------------------------- availability (never overbook)
def _nights_of(b: FoBooking) -> list[date]:
    a, d = R.parse_date(b.arrive), R.parse_date(b.depart)
    return R.nights(a, d) if a and d and d > a else []


def _holding(s, tenant: str, exclude: str = "") -> list[FoBooking]:
    rows = s.scalars(select(FoBooking).where(FoBooking.tenant_id == tenant, FoBooking.status.in_(HOLDING))).all()
    out = []
    n = clock()
    for b in rows:
        if b.id == exclude:
            continue
        if b.status == "hold" and b.hold_until and _aware(b.hold_until) < n:
            continue          # an expired hold no longer holds a room (tick releases it)
        out.append(b)
    return out


def _sellable(s, tenant: str, type_id: str, d: date) -> int:
    rooms = s.scalars(select(FoRoom).where(FoRoom.tenant_id == tenant, FoRoom.type_id == type_id, FoRoom.archived.is_(False))).all()
    return sum(1 for r in rooms if not _ooo_on(r, d))


def _free_counts(s, tenant: str, type_id: str, arrive: date, depart: date, exclude: str = "") -> list[dict]:
    held = _holding(s, tenant, exclude)
    out = []
    for d in R.nights(arrive, depart):
        used = sum(1 for b in held if b.type_id == type_id and d in _nights_of(b))
        total = _sellable(s, tenant, type_id, d)
        out.append({"date": d.isoformat(), "rooms": total, "sold": used, "free": total - used})
    return out


def _check_avail(s, tenant: str, type_id: str, arrive: date, depart: date, rooms: int = 1, exclude: str = "") -> None:
    short = [x for x in _free_counts(s, tenant, type_id, arrive, depart, exclude) if x["free"] < rooms]
    if short:
        t = s.get(FoRoomType, type_id)
        days = ", ".join(f"{x['date']} ({max(0, x['free'])} khaali)" for x in short[:6])
        raise FoError(f"OVERBOOK: {t.name if t else 'is type'} mein in raaton kamra nahi bachta: {days}. Booking nahi bani "
                      "(ek raat do logon ko nahi di jaati).")


def _room_busy(s, tenant: str, room_id: str, arrive: date, depart: date, exclude: str = "") -> FoBooking | None:
    for b in _holding(s, tenant, exclude):
        if b.room_id == room_id and R.overlaps(arrive, depart, R.parse_date(b.arrive), R.parse_date(b.depart)):
            return b
    return None


def _free_room(s, tenant: str, type_id: str, arrive: date, depart: date, exclude: str = "", allow_dirty: bool = True) -> FoRoom | None:
    rank = {"inspected": 0, "clean": 1, "dirty": 2}
    rooms = s.scalars(select(FoRoom).where(FoRoom.tenant_id == tenant, FoRoom.type_id == type_id, FoRoom.archived.is_(False))).all()
    ok = []
    for r in rooms:
        if any(_ooo_on(r, d) for d in R.nights(arrive, depart)) or r.status == "out_of_order":
            continue
        if not allow_dirty and r.status == "dirty":
            continue
        if _room_busy(s, tenant, r.id, arrive, depart, exclude):
            continue
        ok.append(r)
    ok.sort(key=lambda r: (rank.get(r.status, 3), r.number))
    return ok[0] if ok else None


def availability(tenant: str, arrive: str = "", depart: str = "", days: int = 7) -> dict:
    a = _date(arrive, tenant) or today(tenant)
    d = _date(depart, tenant) or a + timedelta(days=max(1, min(60, int(days or 7))))
    if d <= a:
        raise FoError("Jaane ki tareekh aane ke baad honi chahiye.")
    with session() as s:
        out = []
        for t in s.scalars(select(FoRoomType).where(FoRoomType.tenant_id == tenant, FoRoomType.archived.is_(False))
                           .order_by(FoRoomType.code)).all():
            cal = _free_counts(s, tenant, t.id, a, d)
            out.append({"type": t.code, "name": t.name, "nights": cal, "free_all_nights": min((x["free"] for x in cal), default=0)})
    say = "; ".join(f"{x['name']}: {max(0, x['free_all_nights'])} kamre poore samay khaali" for x in out) or "Pehle room types jodiye."
    return {"from": a.isoformat(), "to": d.isoformat(), "types": out, "say": say}


# ---------------------------------------------------------------- housekeeping
def set_room_status(tenant: str, room: str, status: str, reason: str = "", until: str = "", from_: str = "",
                    actor: str = "owner") -> dict:
    """Housekeeping board. Out of order is refused when it would take a room that is sold / occupied."""
    status = str(status or "").lower().replace(" ", "_")
    words = {"ganda": "dirty", "saaf": "clean", "ready": "inspected", "jaancha": "inspected", "kharab": "out_of_order", "band": "out_of_order",
             "ooo": "out_of_order"}
    status = words.get(status, status)
    if status not in R.ROOM_STATUS:
        raise FoError("Status: dirty | clean | inspected | out_of_order")
    with _LOCK, session() as s:
        r = _room(s, tenant, room)
        if status == "out_of_order":
            f = _date(from_, tenant) or today(tenant)
            u = _date(until, tenant)
            if u and u < f:
                raise FoError("Kab tak band: shuru ke baad ki tareekh.")
            end = (u or f + timedelta(days=60)) + timedelta(days=1)
            busy = _room_busy(s, tenant, r.id, f, end)
            if busy:
                raise FoError(f"Room {r.number} {busy.number} ({busy.arrive} se {busy.depart}) ke liye diya hua hai; pehle kamra badliye.")
            old = (r.status, r.ooo_from, r.ooo_until)
            r.status, r.ooo_from, r.ooo_until, r.ooo_reason = status, f.isoformat(), u.isoformat() if u else "", _clean(reason, 200)
            s.flush()
            short = [x for x in _free_counts(s, tenant, r.type_id, f, end) if x["free"] < 0]
            if short:
                r.status, r.ooo_from, r.ooo_until = old
                raise FoError(f"Room {r.number} band karne se in raaton overbooking hogi: "
                              + ", ".join(x["date"] for x in short[:6]) + ". Pehle bookings dekhiye.")
        else:
            r.status, r.ooo_reason = status, ("" if status != "out_of_order" else r.ooo_reason)
            r.ooo_from = r.ooo_until = ""
        r.updated_at = now()
        r.notes = (r.notes + f"\n{today(tenant).isoformat()} {status} ({actor_name(tenant, actor)})").strip()[-2000:]
        num, st = r.number, r.status
    return {"room": num, "status": st, "say": f"Room {num}: {R.ROOM_STATUS_LABEL[st]}."}


def board(tenant: str) -> dict:
    td = today(tenant)
    with session() as s:
        rooms = s.scalars(select(FoRoom).where(FoRoom.tenant_id == tenant, FoRoom.archived.is_(False)).order_by(FoRoom.number)).all()
        held = _holding(s, tenant)
        out = []
        for r in rooms:
            t = s.get(FoRoomType, r.type_id)
            occ = next((b for b in held if b.room_id == r.id and b.status == "in_house"), None)
            arr = next((b for b in held if b.room_id == r.id and b.status != "in_house" and b.arrive == td.isoformat()), None)
            g = s.get(FoGuest, occ.guest_id) if occ else None
            out.append({"id": r.id, "number": r.number, "type": t.code if t else "", "type_name": t.name if t else "", "floor": r.floor,
                        "status": r.status, "status_label": R.ROOM_STATUS_LABEL.get(r.status, r.status),
                        "ooo": {"from": r.ooo_from, "until": r.ooo_until, "reason": r.ooo_reason} if r.status == "out_of_order" else None,
                        "occupied": bool(occ), "guest": g.name if g and not g.forgotten else "", "booking": occ.number if occ else "",
                        "departs": occ.depart if occ else "", "arriving": arr.number if arr else ""})
    counts = {k: sum(1 for x in out if x["status"] == k) for k in R.ROOM_STATUS}
    return {"rooms": out, "counts": counts, "say": ", ".join(f"{R.ROOM_STATUS_LABEL[k]} {v}" for k, v in counts.items() if v) or "Koi kamra nahi."}


# ---------------------------------------------------------------- guests
def _guest_view(g: FoGuest | None, full: bool = False) -> dict:
    if g is None:
        return {}
    v = {"id": g.id, "number": g.number, "name": g.name, "phone": g.phone, "nationality": g.nationality,
         "foreign": R.is_foreign(g.nationality), "forgotten": g.forgotten, "no_messages": g.no_messages,
         "ids": {k: {"masked": x.get("masked"), "ref": x.get("ref"), "record": x.get("record")} for k, x in _j(g.ids, {}).items()}}
    if full:
        visa = _j(g.visa, {})
        v.update({"email": g.email, "address": g.address, "company": g.company, "gstin": g.gstin,
                  "visa": {"type": visa.get("type", ""), "expiry": visa.get("expiry", ""), "issued_at": visa.get("issued_at", ""),
                           "number": (visa.get("number") or {}).get("masked", ""), "record": (visa.get("number") or {}).get("record", "")},
                  "arrived_india_on": g.arrived_india_on, "notes": g.notes})
    return v


def _lock_id(tenant: str, kind: str, value: str, gnum: str, note: str) -> dict:
    from .. import idvault
    k = kind if kind in idvault.KINDS else "other"
    try:
        rec = idvault.save(tenant, k, value, f"mehmaan {gnum}", note[:120])
    except Exception as e:
        raise FoError(f"{idvault.KINDS.get(k, 'ID')} sahi nahi: {e}")
    return {"ref": rec["ref"], "masked": rec["masked"], "record": rec["id"], "kind": k}


_ID_KIND = {"aadhaar": "aadhaar", "aadhar": "aadhaar", "passport": "passport", "dl": "dl", "driving": "dl", "licence": "dl",
            "voter": "voter", "epic": "voter", "pan": "pan", "other": "other"}


def save_guest(tenant: str, guest: str = "", name: str = "", phone: str = "", id_kind: str = "", id_number: str = "",
               visa_number: str = "", **kw: Any) -> dict:
    """A guest. ID numbers go straight into the Pehchaan Locker; only {id:...} + masked are kept here."""
    with _LOCK, session() as s:
        g = None
        if guest:
            g = s.scalars(select(FoGuest).where(FoGuest.tenant_id == tenant, (FoGuest.id == guest) | (FoGuest.number == guest.upper()))).first()
            if g is None:
                raise FoError("Mehmaan nahi mila.")
        ph = R.phone10(phone) if phone else ""
        if phone and not ph and not R.is_foreign(str(kw.get("nationality") or "")):
            raise FoError("Mobile 10 ank ka (6-9 se shuru).")
        if g is None and ph:
            g = s.scalars(select(FoGuest).where(FoGuest.tenant_id == tenant, FoGuest.phone == ph, FoGuest.forgotten.is_(False))).first()
        new = g is None
        if new:
            if not str(name).strip():
                raise FoError("Mehmaan ka naam chahiye.")
            g = FoGuest(tenant_id=tenant, number=f"G-{_next(s, tenant, 'guest'):04d}")
            s.add(g)
        if g.forgotten:
            raise FoError("Is mehmaan ki jaankari mita di gayi thi.")
        if name:
            g.name = _clean(name, 160)
        if ph:
            g.phone = ph
        for k in ("email", "address", "company", "notes"):
            if kw.get(k) is not None:
                setattr(g, k, _clean(kw[k], 600))
        if kw.get("nationality"):
            g.nationality = _clean(kw["nationality"], 40)
        if kw.get("gstin"):
            from ..india.money import gstin_ok
            gs = str(kw["gstin"]).strip().upper()
            if not gstin_ok(gs):
                raise FoError("Company ka GSTIN sahi nahi.")
            g.gstin = gs
        if kw.get("no_messages") is not None:
            g.no_messages = bool(kw["no_messages"])
        if kw.get("arrived_india_on"):
            d = R.parse_date(kw["arrived_india_on"])
            g.arrived_india_on = d.isoformat() if d else ""
        visa = _j(g.visa, {})
        for k in ("visa_type", "visa_expiry", "visa_issued_at"):
            if kw.get(k):
                key = k.split("_", 1)[1]
                v = R.parse_date(kw[k]) if key == "expiry" else None
                visa[key] = v.isoformat() if v else _clean(kw[k], 60)
        g.visa = _dump(visa)
        s.flush()
        gid, gnum, gname = g.id, g.number, g.name
    got = {}
    if id_number:
        kind = _ID_KIND.get(str(id_kind or "").lower().split()[0] if id_kind else "", "")
        if not kind:
            kind = "aadhaar" if re.fullmatch(r"[\d\s-]{12,14}", id_number) else "other"
        got[kind] = _lock_id(tenant, kind, id_number, gnum, f"Hotel guest {gname}")
    visa_ref = _lock_id(tenant, "other", visa_number, gnum, f"Visa {gname}") if visa_number else None
    with _LOCK, session() as s:
        g = s.get(FoGuest, gid)
        if got:
            g.ids = _dump({**_j(g.ids, {}), **got})
        if visa_ref:
            g.visa = _dump({**_j(g.visa, {}), "number": visa_ref})
        view = _guest_view(g, True)
    return {"guest": view, "new": new, "say": f"{gnum} {gname} save." + (" ID locker mein (masked)." if got or visa_ref else "")}


def guests(tenant: str, q: str = "") -> list[dict]:
    with session() as s:
        rows = s.scalars(select(FoGuest).where(FoGuest.tenant_id == tenant, FoGuest.forgotten.is_(False)).order_by(FoGuest.number.desc())).all()
        return [_guest_view(g) for g in rows if not q or q.lower() in f"{g.name} {g.phone} {g.number}".lower()][:200]


# ---------------------------------------------------------------- rates / folio
def _night_lines(s, tenant: str, t: FoRoomType, arrive: date, depart: date, adults: int, child_beds: int = 0,
                 rate: Any = None) -> list[dict]:
    st = _st(s, tenant)
    plan = {"base": t.base, "weekend": t.weekend, "seasons": _j(t.seasons, [])}
    out = []
    for d in R.nights(arrive, depart):
        if rate not in (None, "") and R.money(rate) > 0:
            r, why = R.money(rate), "aapka rate"
        else:
            r, why = R.night_rate(plan, d, _j(st.weekend_days, [4, 5]))
        extra = max(0, adults - 2) * (t.extra_adult or 0) + max(0, child_beds) * (t.extra_child or 0)
        if max(0, adults - 2) and not t.extra_adult:
            why += " (extra adult ka rate nahi)"
        out.append({"date": d.isoformat(), "rate": r, "extra": _r2(extra), "why": why})
    return out


def _payments(s, b: FoBooking) -> list[FoPayment]:
    return s.scalars(select(FoPayment).where(FoPayment.booking_id == b.id, FoPayment.void.is_(False)).order_by(FoPayment.created_at)).all()


def _charges(s, b: FoBooking) -> list[FoCharge]:
    return s.scalars(select(FoCharge).where(FoCharge.booking_id == b.id, FoCharge.void.is_(False)).order_by(FoCharge.created_at)).all()


def _folio(s, b: FoBooking) -> dict:
    """Room nights + extras + GST by slab + payments. Rates that are blank stay blank (bill can't be final)."""
    st = _st(s, b.tenant_id)
    gst_on = bool(st.gstin)
    nights = _j(b.nights, [])
    blanks = [n["date"] for n in nights if n.get("rate") is None]
    room_gross = sum((n["rate"] or 0) + (n.get("extra") or 0) for n in nights)
    disc = room_gross if b.complimentary else min(b.discount or 0, room_gross)
    share = (disc / room_gross) if room_gross else 0.0
    t = s.get(FoRoomType, b.type_id)
    room = s.get(FoRoom, b.room_id) if b.room_id else None
    lines, slabs = [], []
    for n in nights:
        if n.get("rate") is None:
            lines.append({"kind": "room", "text": f"Kamra {n['date']}", "value": None, "gst_pct": None, "blank": True})
            continue
        gross = n["rate"] + (n.get("extra") or 0)
        val = _r2(gross * (1 - share))
        pct = R.room_gst_pct(val) if gst_on else 0
        slabs.append(pct)
        lines.append({"kind": "room", "text": f"Kamra {room.number + ' ' if room else ''}{t.name if t else ''} {n['date']}".strip(),
                      "value": val, "gross": _r2(gross), "gst_pct": pct, "sac": "996311"})
    room_slab = max(slabs) if slabs else 5
    for c in _charges(s, b):
        val = _r2(c.qty * c.rate)
        if c.kind in ("cancel_fee",):
            pct = room_slab
        else:
            pct = R.extra_gst_pct(c.kind, st.specified_premises, _j(st.extra_gst, {}))
            pct = room_slab if pct == -1 else pct
        lines.append({"kind": c.kind, "text": c.text or c.kind, "value": val, "gst_pct": pct if gst_on else 0, "id": c.id,
                      "qty": c.qty, "rate": c.rate, "on": c.on})
    taxable = cgst = sgst = 0.0
    for ln in lines:
        if ln.get("value") is None:
            continue
        tx = R.split_tax(ln["value"], ln["gst_pct"] or 0, False)
        ln.update({"cgst": tx["cgst"], "sgst": tx["sgst"], "total": _r2(ln["value"] + tx["tax"])})
        taxable += ln["value"]
        cgst += tx["cgst"]
        sgst += tx["sgst"]
    total = _r2(taxable + cgst + sgst)
    paid = _r2(sum(p.amount for p in _payments(s, b)))
    return {"lines": lines, "blanks": blanks, "room_gross": _r2(room_gross), "discount": _r2(disc), "complimentary": b.complimentary,
            "taxable": _r2(taxable), "cgst": _r2(cgst), "sgst": _r2(sgst), "total": total, "paid": paid,
            "balance": _r2(total - paid) if not blanks else None, "gst_on": gst_on,
            "note": "Kamre ki GST jagah ke hisaab se CGST + SGST (place of supply = hotel ka state)." if gst_on else "GSTIN nahi: GST nahi lagaya."}


def _bview(s, b: FoBooking, full: bool = False) -> dict:
    g = s.get(FoGuest, b.guest_id) if b.guest_id else None
    t = s.get(FoRoomType, b.type_id) if b.type_id else None
    r = s.get(FoRoom, b.room_id) if b.room_id else None
    f = _folio(s, b)
    v = {"id": b.id, "number": b.number, "status": b.status, "status_label": STATUS_LABEL.get(b.status, b.status),
         "source": b.source, "source_label": R.SOURCE_LABEL.get(b.source, b.source), "ota_ref": b.ota_ref,
         "guest": g.name if g and not g.forgotten else ("(mitaya)" if g else ""), "guest_number": g.number if g else "",
         "phone": g.phone if g and not g.forgotten else "", "foreign": bool(g and R.is_foreign(g.nationality)),
         "type": t.code if t else "", "type_name": t.name if t else "", "room": r.number if r else "",
         "arrive": b.arrive, "depart": b.depart, "nights_count": len(_j(b.nights, [])), "adults": b.adults, "children": b.children,
         "meal_plan": b.meal_plan, "total": f["total"] if not f["blanks"] else None, "paid": f["paid"], "balance": f["balance"],
         "blanks": f["blanks"], "discount": f["discount"], "complimentary": b.complimentary,
         "hold_until": _aware(b.hold_until).isoformat(timespec="minutes") if b.hold_until and b.status == "hold" else "",
         "ota_amount": b.ota_amount, "paid_online": b.paid_online, "pay_at_hotel": b.pay_at_hotel, "commission": b.commission}
    if full:
        bill = s.scalars(select(FoBill).where(FoBill.booking_id == b.id).order_by(FoBill.created_at.desc())).first()
        fc = s.scalars(select(FoFormC).where(FoFormC.booking_id == b.id)).all()
        v.update({"folio": f, "nights": _j(b.nights, []), "guest_full": _guest_view(g, True) if g and not g.forgotten else {},
                  "payments": [{"id": p.id, "amount": p.amount, "mode": p.mode, "on": p.on, "receipt_no": p.receipt_no,
                                "ref_tail": p.ref_tail, "by": p.by} for p in _payments(s, b)],
                  "policy": _j(b.policy, {}), "policy_text": R.policy_text(_j(b.policy, {})), "companions": _j(b.companions, []),
                  "purpose": b.purpose, "coming_from": b.coming_from, "going_to": b.going_to, "notes": b.notes,
                  "bill": {"id": bill.id, "number": bill.number, "status": bill.status, "file_id": bill.file_id} if bill else None,
                  "formc": [_fc_view(s, x) for x in fc], "cancel_fee": b.cancel_fee,
                  "ota_note": "OTA ka amount: tax / commission shamil hai ya nahi, extranet se milaiye." if b.ota_amount is not None else ""})
    return v


def _booking(s, tenant: str, ref: str) -> FoBooking:
    ref = str(ref or "").strip()
    b = s.scalars(select(FoBooking).where(FoBooking.tenant_id == tenant,
                                          (FoBooking.id == ref) | (FoBooking.number == ref.upper()) | (FoBooking.ota_ref == ref.upper()))).first()
    if b is None:
        raise FoError(f"Booking {ref} nahi mili.")
    return b


# ---------------------------------------------------------------- bookings
def book(tenant: str, room_type: str, arrive: str, depart: str = "", nights: Any = None, name: str = "", phone: str = "",
         guest: str = "", adults: Any = 1, children: Any = 0, child_beds: Any = 0, rate: Any = None, discount: Any = None,
         complimentary: bool = False, source: str = "walk_in", hold: bool = False, room: str = "", meal_plan: str = "",
         notes: str = "", nationality: str = "", ota_ref: str = "", ota_amount: Any = None, paid_online: bool = False,
         pay_at_hotel: Any = None, commission: Any = None, actor: str = "owner", conv: str = "", publish: bool = True,
         draft_confirmation: bool = True) -> dict:
    """A reservation. Refused in code when any night would be sold twice. Discount / complimentary only via her card."""
    td = today(tenant)
    a = _date(arrive, tenant)
    if not a:
        raise FoError("Aane ki tareekh chahiye (24-10-2026 / kal).")
    d = _date(depart, tenant) if depart else (a + timedelta(days=int(R.money(nights))) if R.money(nights) > 0 else None)
    if not d or d <= a:
        raise FoError("Jaane ki tareekh (ya kitni raatein) chahiye; aane ke baad.")
    if (d - a).days > 90:
        raise FoError("90 raat se lambi booking yahan nahi (long stay alag se).")
    if a < td - timedelta(days=1) and source in ("walk_in", "phone", "whatsapp", "direct", "agent"):
        raise FoError("Beeti hui tareekh ki booking nahi banti.")
    src = source if source in R.SOURCE_LABEL else "other"
    ad = max(1, int(R.money(adults)) or 1)
    ch = max(0, int(R.money(children)))
    gid = guest
    if not gid:
        if not (name or "").strip():
            raise FoError("Mehmaan ka naam chahiye.")
        gid = save_guest(tenant, "", name, phone if R.phone10(phone) else "", nationality=nationality or "Indian")["guest"]["id"]
    with _LOCK, session() as s:
        t = _type(s, tenant, room_type)
        if ad > max(t.max_adults, 1):
            raise FoError(f"{t.name} mein zyada se zyada {t.max_adults} log; ek aur kamra lijiye.")
        g = s.scalars(select(FoGuest).where(FoGuest.tenant_id == tenant, (FoGuest.id == gid) | (FoGuest.number == str(gid).upper()))).first()
        if g is None or g.forgotten:
            raise FoError("Mehmaan nahi mila.")
        _check_avail(s, tenant, t.id, a, d)
        rm = None
        if room:
            rm = _room(s, tenant, room)
            if rm.type_id != t.id:
                raise FoError(f"Room {rm.number} {t.name} ka nahi.")
            if any(_ooo_on(rm, x) for x in R.nights(a, d)):
                raise FoError(f"Room {rm.number} un dino band (out of order) hai.")
            busy = _room_busy(s, tenant, rm.id, a, d)
            if busy:
                raise FoError(f"Room {rm.number} {busy.number} ko diya hua hai.")
        st = _st(s, tenant)
        per_night = None
        if ota_amount not in (None, "") and R.money(ota_amount) > 0 and rate in (None, ""):
            per_night = _r2(R.money(ota_amount) / (d - a).days)
        lines = _night_lines(s, tenant, t, a, d, ad, int(R.money(child_beds)), rate if rate not in (None, "") else per_night)
        if per_night is not None:
            for ln in lines:
                ln["why"] = "OTA mail ka amount"
        b = FoBooking(tenant_id=tenant, number=f"R-{_next(s, tenant, 'booking'):04d}", guest_id=g.id, source=src,
                      ota_ref=_clean(ota_ref, 40).upper(), type_id=t.id, room_id=rm.id if rm else "", arrive=a.isoformat(),
                      depart=d.isoformat(), adults=ad, children=ch, meal_plan=(meal_plan or t.meal_plan or "EP").upper()[:4],
                      nights=_dump(lines), status="hold" if hold else "confirmed",
                      hold_until=clock() + timedelta(hours=st.hold_hours) if hold else None,
                      ota_amount=R.money(ota_amount) if ota_amount not in (None, "") else None, paid_online=bool(paid_online),
                      pay_at_hotel=R.money(pay_at_hotel) if pay_at_hotel not in (None, "") else None,
                      commission=R.money(commission) if commission not in (None, "") else None,
                      policy=st.cancel_policy, notes=_clean(notes, 1000), created_by=actor_name(tenant, actor))
        s.add(b)
        s.flush()
        view = _bview(s, b, True)
        bid = b.id
    say = (f"{view['number']}: {view['guest']} · {view['type_name']} · {view['arrive']} se {view['depart']} "
           f"({view['nights_count']} raat) · " + (f"kul {_rs(view['total'])}" if view["total"] is not None else
                                                  "rate khaali: " + ", ".join(view["blanks"][:4])) + ".")
    if hold:
        say += f" Hold {view['hold_until']} tak (advance aane par pakki)."
    out = {"booking": view, "say": say}
    if publish:
        _publish(tenant, "frontoffice_booking", {"booking": view, "say": say}, conv)
    disc = R.money(discount) if discount not in (None, "") else 0
    if complimentary or disc > 0:
        c = request_discount(tenant, bid, disc, bool(complimentary), "booking ke waqt", actor, conv)
        out["card"] = c["card"]
        out["say"] += " " + c["say"]
    if draft_confirmation and src not in R.SOURCES_OTA and view["phone"]:
        try:
            out["message"] = draft_message(tenant, bid, "confirmation", conv=conv, publish=publish)["message"]
            out["say"] += " Confirmation ka draft taiyaar (aapke tap se jaayega)."
        except FoError:
            pass
    return out


def confirm_hold(tenant: str, booking: str, actor: str = "owner", conv: str = "") -> dict:
    with _LOCK, session() as s:
        b = _booking(s, tenant, booking)
        if b.status != "hold":
            raise FoError(f"{b.number} hold par nahi ({STATUS_LABEL.get(b.status)}).")
        _check_avail(s, tenant, b.type_id, R.parse_date(b.arrive), R.parse_date(b.depart), exclude=b.id)
        b.status, b.hold_until = "confirmed", None
        view = _bview(s, b)
    _publish(tenant, "frontoffice_booking", {"booking": view, "say": f"{view['number']} pakki."}, conv)
    return {"booking": view, "say": f"{view['number']} pakki ho gayi."}


def amend(tenant: str, booking: str, arrive: str = "", depart: str = "", room_type: str = "", adults: Any = None, rate: Any = None,
          notes: str = "", actor: str = "owner", conv: str = "") -> dict:
    """Change dates / type / people. Availability is checked again; rates are re-read from HER plan (or her new rate)."""
    with _LOCK, session() as s:
        b = _booking(s, tenant, booking)
        if b.status not in ("hold", "confirmed", "in_house"):
            raise FoError(f"{b.number} {STATUS_LABEL.get(b.status)} hai; badal nahi sakte.")
        if s.scalars(select(FoBill).where(FoBill.booking_id == b.id, FoBill.status == "final")).first():
            raise FoError("Bill final ho chuka hai.")
        a = _date(arrive, tenant) or R.parse_date(b.arrive)
        d = _date(depart, tenant) or R.parse_date(b.depart)
        if b.status == "in_house" and a != R.parse_date(b.arrive):
            raise FoError("Mehmaan ruke hue hain: aane ki tareekh nahi badalti.")
        if d <= a:
            raise FoError("Jaane ki tareekh aane ke baad honi chahiye.")
        t = _type(s, tenant, room_type) if room_type else s.get(FoRoomType, b.type_id)
        _check_avail(s, tenant, t.id, a, d, exclude=b.id)
        if b.room_id:
            busy = _room_busy(s, tenant, b.room_id, a, d, exclude=b.id)
            rm = s.get(FoRoom, b.room_id)
            if busy or (rm and rm.type_id != t.id) or (rm and any(_ooo_on(rm, x) for x in R.nights(a, d))):
                if b.status == "in_house":
                    raise FoError(f"Room {rm.number if rm else ''} in raaton khaali nahi; pehle kamra badliye (action=move).")
                b.room_id = ""
        ad = max(1, int(R.money(adults))) if adults not in (None, "") else b.adults
        old = {n["date"]: n for n in _j(b.nights, [])}
        lines = _night_lines(s, tenant, t, a, d, ad, 0, rate)
        if rate in (None, ""):
            for ln in lines:      # keep rates already agreed for nights that stay (OTA / her rate)
                if ln["date"] in old and old[ln["date"]].get("rate") is not None and t.id == b.type_id:
                    ln.update({"rate": old[ln["date"]]["rate"], "why": old[ln["date"]].get("why", "")})
        b.type_id, b.arrive, b.depart, b.adults, b.nights = t.id, a.isoformat(), d.isoformat(), ad, _dump(lines)
        if notes:
            b.notes = (b.notes + "\n" + _clean(notes, 500)).strip()
        view = _bview(s, b)
    say = f"{view['number']} badli: {view['arrive']} se {view['depart']}, {view['type_name']}."
    _publish(tenant, "frontoffice_booking", {"booking": view, "say": say}, conv)
    return {"booking": view, "say": say}


def bookings(tenant: str, status: str = "", on: str = "", q: str = "") -> list[dict]:
    d = _date(on, tenant) if on else None
    with session() as s:
        qy = select(FoBooking).where(FoBooking.tenant_id == tenant)
        if status:
            qy = qy.where(FoBooking.status.in_(status.split(",")))
        rows = s.scalars(qy.order_by(FoBooking.arrive.desc(), FoBooking.number.desc())).all()
        out = []
        for b in rows:
            if d and not (R.parse_date(b.arrive) <= d <= R.parse_date(b.depart)):
                continue
            v = _bview(s, b)
            if q and q.lower() not in f"{v['number']} {v['guest']} {v['phone']} {v['ota_ref']} {v['room']}".lower():
                continue
            out.append(v)
            if len(out) >= 200:
                break
        return out


def booking_view(tenant: str, booking: str) -> dict:
    with session() as s:
        return _bview(s, _booking(s, tenant, booking), True)


# ---------------------------------------------------------------- check-in / room move / check-out
def check_in(tenant: str, booking: str, room: str = "", id_kind: str = "", id_number: str = "", nationality: str = "",
             passport: str = "", visa_number: str = "", visa_type: str = "", visa_expiry: str = "", arrived_india_on: str = "",
             address: str = "", coming_from: str = "", going_to: str = "", purpose: str = "", companions: Any = None,
             force: bool = False, reason: str = "", actor: str = "owner", conv: str = "") -> dict:
    """Guest register + room. IDs into the Pehchaan Locker. Foreign guest -> Form C (Form III) arrival entry due in 24 h."""
    with session() as s:
        b = _booking(s, tenant, booking)
        if b.status not in ("hold", "confirmed"):
            raise FoError(f"{b.number} {STATUS_LABEL.get(b.status)}; check-in nahi.")
        a = R.parse_date(b.arrive)
        if a > today(tenant):
            raise FoError(f"{b.number} {b.arrive} ki hai; pehle aane ki tareekh badliye (action=amend).")
        gid = b.guest_id
        g = s.get(FoGuest, gid)
        foreign = R.is_foreign(nationality or (g.nationality if g else ""))
        have_ids = _j(g.ids, {}) if g else {}
    if force and not is_owner(actor):
        raise FoError("Bina ID / ganda kamra check-in sirf owner ki marzi se.")
    # guest register (IDs into the locker first)
    kw: dict[str, Any] = {}
    if nationality:
        kw["nationality"] = nationality
    if address:
        kw["address"] = address
    if arrived_india_on:
        kw["arrived_india_on"] = arrived_india_on
    if visa_type:
        kw["visa_type"] = visa_type
    if visa_expiry:
        kw["visa_expiry"] = visa_expiry
    if foreign:
        if not passport and "passport" not in have_ids:
            raise FoError("Videshi mehmaan: passport number chahiye (Form C ke liye; Pehchaan Locker mein jaayega).")
        if not visa_number and not _j(g.visa, {}).get("number") and not force:
            raise FoError("Videshi mehmaan: visa number (ya OCI / e-visa no.) chahiye. Na ho to owner force=true + wajah.")
        save_guest(tenant, gid, id_kind="passport" if passport else "", id_number=passport, visa_number=visa_number, **kw)
    else:
        if not id_number and not have_ids and not force:
            raise FoError("Mehmaan ka photo ID chahiye (Aadhaar / DL / voter / passport): police register ke liye. "
                          "Owner force=true + wajah se bina ID bhi kar sakti hain.")
        save_guest(tenant, gid, id_kind=id_kind, id_number=id_number, **kw)
    with _LOCK, session() as s:
        b = _booking(s, tenant, booking)
        a, d = R.parse_date(b.arrive), R.parse_date(b.depart)
        rm = None
        if room:
            rm = _room(s, tenant, room)
            if rm.type_id != b.type_id:
                t2 = s.get(FoRoomType, rm.type_id)
                if not force:
                    raise FoError(f"Room {rm.number} {t2.name if t2 else ''} ka hai, booking {s.get(FoRoomType, b.type_id).name} ki. "
                                  "Upgrade ho to owner force=true kahein.")
                _check_avail(s, tenant, rm.type_id, a, d, exclude=b.id)
        elif b.room_id:
            rm = s.get(FoRoom, b.room_id)
        else:
            rm = _free_room(s, tenant, b.type_id, a, d, exclude=b.id)
        if rm is None:
            raise FoError("Is type ka koi kamra khaali / taiyaar nahi mila.")
        if rm.status == "out_of_order" or any(_ooo_on(rm, x) for x in R.nights(a, d)):
            raise FoError(f"Room {rm.number} band (out of order) hai.")
        busy = _room_busy(s, tenant, rm.id, a, d, exclude=b.id)
        if busy:
            raise FoError(f"Room {rm.number} {busy.number} ko diya hua hai.")
        if rm.status == "dirty" and not force:
            raise FoError(f"Room {rm.number} abhi ganda hai (housekeeping). Saaf karwaiye ya owner force=true.")
        b.room_id, b.type_id = rm.id, rm.type_id
        b.status, b.hold_until, b.checked_in_at = "in_house", None, clock()
        b.coming_from, b.going_to, b.purpose = _clean(coming_from, 120), _clean(going_to, 120), _clean(purpose, 80)
        if companions is not None:
            comp = companions if isinstance(companions, list) else str(companions).split(",")
            b.companions = _dump([_clean(x, 80) for x in comp if str(x).strip()][:12])
        if force and reason:
            b.notes = (b.notes + f"\nCheck-in force: {_clean(reason, 200)}").strip()
        g = s.get(FoGuest, b.guest_id)
        fc = None
        if R.is_foreign(g.nationality):
            fc = FoFormC(tenant_id=tenant, booking_id=b.id, guest_id=g.id, kind="arrival", event_at=b.checked_in_at,
                         due_at=R.formc_due(b.checked_in_at))
            s.add(fc)
        s.flush()
        view = _bview(s, b, True)
        fcv = _fc_view(s, fc) if fc else None
    say = f"{view['number']} {view['guest']}: check-in, room {view['room']}."
    if fcv:
        say += f" Videshi mehmaan: Form C {fcv['due_local']} tak indianfrro.gov.in par bhariye (data taiyaar)."
    if not id_number and not have_ids and not foreign:
        say += " ID baaki hai (owner ne bina ID kiya)."
    _publish(tenant, "frontoffice_booking", {"booking": view, "say": say}, conv)
    if fcv:
        _publish(tenant, "frontoffice_alert", {"kind": "formc", "formc": fcv, "say": say}, conv)
    return {"booking": view, "formc": fcv, "say": say}


def move_room(tenant: str, booking: str, room: str, actor: str = "owner", conv: str = "") -> dict:
    with _LOCK, session() as s:
        b = _booking(s, tenant, booking)
        if b.status not in HOLDING:
            raise FoError(f"{b.number} {STATUS_LABEL.get(b.status)}.")
        rm = _room(s, tenant, room)
        a, d = R.parse_date(b.arrive), R.parse_date(b.depart)
        if b.status == "in_house":
            a = max(a, today(tenant))
        if rm.status == "out_of_order" or any(_ooo_on(rm, x) for x in R.nights(a, d)):
            raise FoError(f"Room {rm.number} band hai.")
        busy = _room_busy(s, tenant, rm.id, a, d, exclude=b.id)
        if busy:
            raise FoError(f"Room {rm.number} {busy.number} ko diya hua hai.")
        if rm.type_id != b.type_id:
            _owner_only(actor, "Doosre type ka kamra (upgrade)")
            _check_avail(s, tenant, rm.type_id, a, d, exclude=b.id)
        old = s.get(FoRoom, b.room_id) if b.room_id else None
        if old and b.status == "in_house":
            old.status = "dirty"
        b.room_id, b.type_id = rm.id, rm.type_id
        view = _bview(s, b)
    say = f"{view['number']}: ab room {view['room']}" + (f" (room {old.number} ganda)." if old and view["status"] == "in_house" else ".")
    _publish(tenant, "frontoffice_booking", {"booking": view, "say": say}, conv)
    return {"booking": view, "say": say}


def check_out(tenant: str, booking: str, actor: str = "owner", conv: str = "") -> dict:
    """Room -> dirty, bill draft -> her 'bill final' card, Form C departure entry for foreigners (24 h)."""
    with _LOCK, session() as s:
        b = _booking(s, tenant, booking)
        if b.status != "in_house":
            raise FoError(f"{b.number} ruke hue nahi ({STATUS_LABEL.get(b.status)}).")
        td = today(tenant)
        if td < R.parse_date(b.depart):          # early departure: nights not stayed come off (she can add a charge herself)
            lines = [n for n in _j(b.nights, []) if R.parse_date(n["date"]) < max(td, R.parse_date(b.arrive) + timedelta(days=1))]
            b.nights, b.depart = _dump(lines), max(td, R.parse_date(b.arrive) + timedelta(days=1)).isoformat()
        b.status, b.checked_out_at = "checked_out", clock()
        rm = s.get(FoRoom, b.room_id)
        if rm:
            rm.status, rm.updated_at = "dirty", now()
        g = s.get(FoGuest, b.guest_id)
        fc = None
        if g and R.is_foreign(g.nationality):
            fc = FoFormC(tenant_id=tenant, booking_id=b.id, guest_id=g.id, kind="departure", event_at=b.checked_out_at,
                         due_at=R.formc_due(b.checked_out_at))
            s.add(fc)
        s.flush()
        view = _bview(s, b, True)
        fcv = _fc_view(s, fc) if fc else None
    out = {"booking": view, "say": f"{view['number']} check-out; room {view['room']} ganda (housekeeping)."}
    try:
        c = bill_request(tenant, booking, actor, conv)
        out["card"] = c["card"]
        out["say"] += " " + c["say"]
    except FoError as e:
        out["say"] += f" Bill: {e}"
    if view["balance"]:
        out["say"] += f" Baaki {_rs(view['balance'])} (UPI link: action=pay_link)."
    if fcv:
        out["formc"] = fcv
        out["say"] += f" Form C departure entry {fcv['due_local']} tak."
        _publish(tenant, "frontoffice_alert", {"kind": "formc", "formc": fcv, "say": "Form C departure entry baaki."}, conv)
    _publish(tenant, "frontoffice_booking", {"booking": view, "say": out["say"]}, conv)
    return out


# ---------------------------------------------------------------- folio: charges / payments
def add_charge(tenant: str, booking: str, kind: str, amount: Any, qty: Any = 1, text: str = "", on: str = "",
               actor: str = "owner", conv: str = "") -> dict:
    kind = kind if kind in R.EXTRA_KINDS else "other"
    rate = R.money(amount)
    if rate <= 0:
        raise FoError("Amount chahiye (aapka / bill ka number).")
    q = R.money(qty) or 1
    with _LOCK, session() as s:
        b = _booking(s, tenant, booking)
        if b.status not in ("in_house", "checked_out", "confirmed"):
            raise FoError(f"{b.number} {STATUS_LABEL.get(b.status)}; charge nahi.")
        if s.scalars(select(FoBill).where(FoBill.booking_id == b.id, FoBill.status == "final")).first():
            raise FoError("Bill final ho chuka hai: nayi line nahi (credit / debit note apne CA se).")
        d = _date(on, tenant) or today(tenant)
        s.add(FoCharge(tenant_id=tenant, booking_id=b.id, kind=kind, text=_clean(text or kind.replace("_", " ").title(), 200), qty=q,
                       rate=rate, on=d.isoformat(), by=actor_name(tenant, actor)))
        s.flush()
        _stale_cards(s, b.id, ("bill",))
        view = _bview(s, b)
    return {"booking": view, "say": f"{view['number']}: {kind} {_rs(q * rate)} joda. Kul {_rs(view['total'])}."}


def void_charge(tenant: str, charge: str, actor: str = "owner") -> dict:
    _owner_only(actor, "Line hatana")
    with _LOCK, session() as s:
        c = s.get(FoCharge, charge)
        if c is None or c.tenant_id != tenant:
            raise FoError("Line nahi mili.")
        if s.scalars(select(FoBill).where(FoBill.booking_id == c.booking_id, FoBill.status == "final")).first():
            raise FoError("Bill final ho chuka hai.")
        c.void = True
        _stale_cards(s, c.booking_id, ("bill",))
        view = _bview(s, s.get(FoBooking, c.booking_id))
    return {"booking": view, "say": "Line hatayi."}


def received(tenant: str, booking: str, amount: Any, mode: str = "upi", ref: str = "", on: str = "", actor: str = "owner",
             conv: str = "") -> dict:
    """Only SHE says money came (Hoola never takes money). A hold with an advance becomes confirmed."""
    _owner_only(actor, "Paise milna darj karna")
    amt = R.money(amount)
    if amt <= 0:
        raise FoError("Kitne mile?")
    mode = mode if mode in ("upi", "cash", "card", "bank", "ota") else "upi"
    with _LOCK, session() as s:
        b = _booking(s, tenant, booking)
        if b.status in ("cancelled", "no_show") and amt > 0 and not b.cancel_fee:
            raise FoError(f"{b.number} radd hai.")
        st = _st(s, tenant)
        d = _date(on, tenant) or today(tenant)
        p = FoPayment(tenant_id=tenant, booking_id=b.id, amount=_r2(amt), mode=mode, ref_tail=re.sub(r"\W", "", ref)[-4:],
                      on=d.isoformat(), by=actor_name(tenant, actor),
                      receipt_no=f"RC/{R.fy_of(d)}/{_next(s, tenant, 'receipt:' + R.fy_of(d)):04d}")
        s.add(p)
        if b.status == "hold":
            b.status, b.hold_until = "confirmed", None
        s.flush()
        _stale_cards(s, b.id, ("bill", "refund"))
        view = _bview(s, b)
        _ = st
    say = f"{view['number']}: {_rs(amt)} mile ({mode}), raseed {p.receipt_no}."
    if view["balance"] is not None:
        say += f" Baaki {_rs(view['balance'])}."
    _publish(tenant, "frontoffice_booking", {"booking": view, "say": say}, conv)
    return {"booking": view, "receipt_no": p.receipt_no, "say": say}


def pay_link(tenant: str, booking: str, amount: Any = None, what: str = "balance") -> dict:
    """upi://pay to HER UPI ID. The guest pays in their own app; Hoola never takes money."""
    with session() as s:
        b = _booking(s, tenant, booking)
        st = _st(s, tenant)
        upi, nm = _upi(st)
        v = _bview(s, b)
    if not upi:
        raise FoError("Aapka UPI ID nahi hai (Niyam ya Employee profile mein daaliye).")
    amt = R.money(amount) if amount not in (None, "") else (v["balance"] or 0)
    if amt <= 0:
        return {"url": "", "say": "Abhi kuch baaki nahi." if v["balance"] is not None else "Rate khaali hai: pehle rate / amount bataiye."}
    note = f"{'Advance' if what == 'advance' else 'Room'} {v['number']}"[:40]
    url = f"upi://pay?pa={quote(upi)}&pn={quote(nm or '')}&am={amt:.2f}&cu=INR&tn={quote(note)}"
    return {"url": url, "amount": _r2(amt), "upi_id": upi, "say": f"{_rs(amt)} ka UPI link (paise seedhe aapke UPI ID par)."}


# ---------------------------------------------------------------- cards (HER tap only)
def _basis(s, kind: str, target: str) -> dict:
    """What a card was made from, read NOW: any change after the card makes it stale."""
    b = s.get(FoBooking, target)
    if b is None:
        return {}
    pays = [(p.id, p.amount) for p in _payments(s, b)]
    chg = [(c.id, c.qty, c.rate) for c in _charges(s, b)]
    base = {"status": b.status, "nights": _j(b.nights, []), "arrive": b.arrive, "depart": b.depart, "type": b.type_id}
    if kind == "bill":
        return {**base, "pays": pays, "chg": chg, "disc": b.discount, "comp": b.complimentary}
    if kind == "refund":
        return {**base, "pays": pays}
    return base


def _card_view(c: FoCard) -> dict:
    show = _j(c.show, {})
    return {"id": c.id, "number": c.number, "kind": c.kind, "kind_label": CARD_LABEL.get(c.kind, c.kind), "target": c.target,
            "show": show, "hash": c.hash, "status": c.status, "reason": c.reason, "asked_by": c.asked_by,
            "money": c.kind in MONEY_KINDS, "created_at": _aware(c.created_at).isoformat(timespec="minutes") if c.created_at else ""}


def _make_card(s, tenant: str, kind: str, target: str, show: dict, actor: str, conv: str) -> tuple[dict, bool]:
    h = _hash({"show": show, "basis": _basis(s, kind, target)})
    for o in s.scalars(select(FoCard).where(FoCard.tenant_id == tenant, FoCard.kind == kind, FoCard.target == target,
                                            FoCard.status == "waiting")).all():
        if o.hash == h:
            return _card_view(o), True
        o.status = "stale"
    c = FoCard(tenant_id=tenant, number=f"C-{_next(s, tenant, 'card'):04d}", kind=kind, target=target, show=_dump(show), hash=h,
               asked_by=actor_name(tenant, actor), conv=conv)
    s.add(c)
    s.flush()
    return _card_view(c), False


def _stale_cards(s, target: str, kinds: tuple) -> None:
    for o in s.scalars(select(FoCard).where(FoCard.target == target, FoCard.kind.in_(kinds), FoCard.status == "waiting")).all():
        o.status = "stale"


def _ask(tenant: str, view: dict, existing: bool, conv: str) -> dict:
    if existing:
        return {"card": view, "status": "needs_tap", "existing": True, "say": f"{view['number']} pehle se aapki haan ka intezaar kar raha hai."}
    say = f"{view['number']}: {view['kind_label']} - {view['show'].get('title', '')}. Aapki haan ka card bheja hai."
    if view["show"].get("checks"):
        say += " Dhyaan: " + " ".join(view["show"]["checks"])
    _publish(tenant, "frontoffice_card", {"card": view, "say": say}, conv)
    return {"card": view, "status": "needs_tap", "say": say}


def request_discount(tenant: str, booking: str, amount: Any = None, complimentary: bool = False, reason: str = "",
                     actor: str = "owner", conv: str = "", pct: Any = None) -> dict:
    with _LOCK, session() as s:
        b = _booking(s, tenant, booking)
        if b.status not in HOLDING:
            raise FoError(f"{b.number} {STATUS_LABEL.get(b.status)}.")
        f = _folio(s, b)
        if f["blanks"]:
            raise FoError("Pehle har raat ka rate bharen (khaali: " + ", ".join(f["blanks"][:4]) + ").")
        gross = f["room_gross"]
        if complimentary:
            amt, kind = gross, "comp"
        else:
            amt = _r2(gross * R.money(pct) / 100) if pct not in (None, "") else R.money(amount)
            kind = "discount"
            if amt <= 0:
                raise FoError("Discount kitna (rupaye ya %)?")
            if amt > gross:
                raise FoError(f"Discount kamre ke kiraye ({_rs(gross)}) se zyada nahi.")
        g = s.get(FoGuest, b.guest_id)
        checks = []
        if amt >= 0.3 * gross and not complimentary:
            checks.append(f"Kiraye ka {round(100 * amt / gross)}% discount.")
        show = {"title": f"{b.number} {g.name if g else ''}: " + ("poora kamra muft" if complimentary else f"{_rs(amt)} discount"),
                "booking": b.number, "amount": _r2(amt), "amount_text": _rs(amt), "room_gross": gross, "reason": _clean(reason, 200),
                "checks": checks}
        view, existing = _make_card(s, tenant, kind, b.id, show, actor, conv)
    return _ask(tenant, view, existing, conv)


def bill_request(tenant: str, booking: str, actor: str = "owner", conv: str = "") -> dict:
    with _LOCK, session() as s:
        b = _booking(s, tenant, booking)
        if b.status not in ("in_house", "checked_out", "cancelled", "no_show"):
            raise FoError("Bill check-out (ya cancellation charge) ke baad.")
        if s.scalars(select(FoBill).where(FoBill.booking_id == b.id, FoBill.status == "final")).first():
            raise FoError("Bill pehle hi final hai.")
        f = _folio(s, b)
        if f["blanks"]:
            raise FoError("Rate khaali hai in raaton ka: " + ", ".join(f["blanks"][:4]) + " (action=amend rate=...).")
        g = s.get(FoGuest, b.guest_id)
        checks = []
        if f["balance"] and f["balance"] > 0:
            checks.append(f"Abhi {_rs(f['balance'])} baaki.")
        if f["balance"] and f["balance"] < 0:
            checks.append(f"{_rs(-f['balance'])} zyada mila hai (refund card banaiye?).")
        show = {"title": f"{b.number} {g.name if g and not g.forgotten else ''}: bill {_rs(f['total'])}", "booking": b.number,
                "lines": [{"text": x["text"], "value": x["value"], "gst_pct": x["gst_pct"], "total": x.get("total")} for x in f["lines"]],
                "taxable": f["taxable"], "cgst": f["cgst"], "sgst": f["sgst"], "total": f["total"], "paid": f["paid"],
                "balance": f["balance"], "gst_on": f["gst_on"], "checks": checks, "amount_text": _rs(f["total"])}
        view, existing = _make_card(s, tenant, "bill", b.id, show, actor, conv)
    return _ask(tenant, view, existing, conv)


def refund_request(tenant: str, booking: str, amount: Any, reason: str = "", actor: str = "owner", conv: str = "") -> dict:
    amt = R.money(amount)
    if amt <= 0:
        raise FoError("Kitna refund?")
    with _LOCK, session() as s:
        b = _booking(s, tenant, booking)
        f = _folio(s, b)
        keep = 0.0 if f["blanks"] and b.status in ("cancelled", "no_show") else (f["total"] if b.status in ("in_house", "checked_out")
                                                                                   else b.cancel_fee)
        can = _r2(f["paid"] - (keep or 0))
        if amt > can + 0.01:
            raise FoError(f"Refund {_rs(can)} se zyada nahi (mila {_rs(f['paid'])}, rakhna {_rs(keep or 0)}).")
        g = s.get(FoGuest, b.guest_id)
        pays = _payments(s, b)
        last = pays[-1] if pays else None
        show = {"title": f"{b.number} {g.name if g and not g.forgotten else ''}: {_rs(amt)} wapas", "booking": b.number,
                "amount": _r2(amt), "amount_text": _rs(amt), "payee": g.name if g else "", "reason": _clean(reason, 200),
                "how": f"Jis tareeke se aaya ({last.mode if last else 'upi'}): aap apne app / bank se bhejiye.", "checks": []}
        view, existing = _make_card(s, tenant, "refund", b.id, show, actor, conv)
    return _ask(tenant, view, existing, conv)


def cancel(tenant: str, booking: str, reason: str = "", by_ota: bool = False, no_show: bool = False, actor: str = "owner",
           conv: str = "") -> dict:
    """Cancel. No fee by HER policy -> cancelled at once. A fee -> her card (Haan = with fee, Nahi = free)."""
    with _LOCK, session() as s:
        b = _booking(s, tenant, booking)
        if b.status not in ("hold", "confirmed"):
            raise FoError(f"{b.number} {STATUS_LABEL.get(b.status)}; radd nahi ho sakti.")
        a = R.parse_date(b.arrive)
        nights = _j(b.nights, [])
        first = (nights[0]["rate"] or 0) + (nights[0].get("extra") or 0) if nights else 0
        total = sum((n["rate"] or 0) + (n.get("extra") or 0) for n in nights)
        pol = _j(b.policy, {})
        if no_show:
            kind = str(pol.get("no_show") or pol.get("fee") or "first_night")
            fee, why = (R.cancel_fee({**pol, "free_days": 0, "fee": kind}, a, a, first, total) if pol else
                        (0.0, "Policy set nahi; no-show charge nahi."))
        else:
            fee, why = R.cancel_fee(pol, a, today(tenant), first, total)
        if by_ota or b.source in R.SOURCES_OTA:
            fee, why = 0.0, "OTA booking: cancellation charge OTA ki policy se OTA hi leta hai."
        if fee <= 0:
            b.status = "no_show" if no_show else "cancelled"
            b.notes = (b.notes + f"\nRadd: {_clean(reason, 200)} ({why})").strip()
            view = _bview(s, b)
            paid = view["paid"]
            card = None
        else:
            g = s.get(FoGuest, b.guest_id)
            show = {"title": f"{b.number} {g.name if g else ''}: {'no-show' if no_show else 'cancellation'} charge {_rs(fee)}",
                    "booking": b.number, "amount": _r2(fee), "amount_text": _rs(fee), "why": why, "no_show": no_show,
                    "reason": _clean(reason, 200), "policy": R.policy_text(pol),
                    "checks": ["Haan = charge ke saath radd; Nahi = bina charge radd."]}
            card, existing = _make_card(s, tenant, "cancel_fee", b.id, show, actor, conv)
            view = _bview(s, b)
            paid = view["paid"]
    if card is None:
        say = f"{view['number']} {'no-show' if no_show else 'radd'}: {why}"
        _publish(tenant, "frontoffice_booking", {"booking": view, "say": say}, conv)
        out = {"booking": view, "say": say}
        if paid > 0:
            r = refund_request(tenant, view["id"], paid, "radd booking ka advance", actor, conv)
            out["card"] = r["card"]
            out["say"] += " " + r["say"]
        if view["phone"] and not by_ota:
            try:
                out["message"] = draft_message(tenant, view["id"], "cancel", conv=conv)["message"]
            except FoError:
                pass
        return out
    return _ask(tenant, card, existing, conv) | {"booking": view}


def challenge(tenant: str, cid: str) -> dict:
    """approvekey's message format (kind 'frontoffice'): the phone signs this card as it is NOW."""
    from .. import approvekey as ak
    c = card(tenant, cid)
    if not c or c["status"] != "waiting":
        raise FoError("Yeh card ab baaki nahi hai.")
    nonce = secrets.token_urlsafe(18)
    exp = datetime.now(timezone.utc) + ak.CHALLENGE_TTL
    msg = f"hoola-approve|v1|{tenant}|frontoffice|{cid}|{c['hash']}|{nonce}|{int(exp.timestamp())}"
    with session() as s:
        ch = ak.ApproveChallenge(tenant_id=tenant, kind="frontoffice", target=cid, digest=c["hash"], message=msg, expires_at=exp)
        s.add(ch)
        s.flush()
        chid = ch.id
    show = {"title": f"{c['kind_label']}: {c['show'].get('title', '')}"[:80], "amount": c["show"].get("amount_text", ""),
            "payee": c["show"].get("payee", "")}
    return {"challenge_id": chid, "message": msg, "show": show, "expires_at": exp.isoformat()}


def _verify_fp(tenant: str, cid: str, digest: str, proof: dict | None) -> None:
    from .. import approvekey as ak
    if not proof:
        raise FoError("fingerprint_needed")
    with session() as s:
        ak._settle(s, tenant)
        ok = ak._check(s, tenant, "frontoffice", cid, proof, digest)
    if not ok:
        raise FoError("fingerprint_invalid")


def card(tenant: str, cid: str) -> dict | None:
    with session() as s:
        c = s.get(FoCard, cid)
        return _card_view(c) if c and c.tenant_id == tenant else None


def cards(tenant: str, status: str = "waiting") -> list[dict]:
    with session() as s:
        q = select(FoCard).where(FoCard.tenant_id == tenant)
        if status:
            q = q.where(FoCard.status == status)
        return [_card_view(c) for c in s.scalars(q.order_by(FoCard.created_at.desc())).all()[:100]]


def decide(tenant: str, cid: str, yes: bool, card_hash: str = "", reason: str = "", proof: dict | None = None,
           by_tap: bool = False, conv: str = "") -> dict:
    """ONLY her tap in the app (API). Refunds need her fingerprint once a phone key is active."""
    if not by_tap:
        raise FoError("Yeh faisla sirf aapke tap se (app ka card).")
    with session() as s:
        c = s.get(FoCard, cid)
        if c is None or c.tenant_id != tenant:
            raise FoError("Card nahi mila.")
        if c.status == "stale":
            raise FoError("Card ke baad booking badal gayi (purana card); naya card banaiye.")
        if c.status != "waiting":
            return {"card": _card_view(c), "say": f"{c.number} pehle hi '{c.status}' hai."}
        now_hash = _hash({"show": _j(c.show, {}), "basis": _basis(s, c.kind, c.target)})
        if not card_hash or card_hash != c.hash or now_hash != c.hash:
            if now_hash != c.hash:
                c.status = "stale"
            raise FoError("Card ke baad booking badal gayi; naya card banaiye.")
        kind = c.kind
    if yes and kind in MONEY_KINDS:
        from .. import approvekey as ak
        if ak.active_on(tenant):
            _verify_fp(tenant, cid, card_hash, proof)
    follow: list[tuple] = []
    with _LOCK, session() as s:
        c = s.get(FoCard, cid)
        c.status = "approved" if yes else "rejected"
        c.reason = _clean(reason, 300)
        c.decided_at = now()
        b = s.get(FoBooking, c.target)
        show = _j(c.show, {})
        say = f"{c.number}: {'haan' if yes else 'nahi'}."
        extra: dict[str, Any] = {}
        if kind in ("discount", "comp") and yes:
            if kind == "comp":
                b.complimentary, b.discount = True, 0.0
            else:
                b.discount, b.complimentary = float(show.get("amount") or 0), False
            say = f"{b.number}: {'complimentary' if kind == 'comp' else _rs(b.discount) + ' discount'} laga."
        elif kind == "cancel_fee":
            b.status = "no_show" if show.get("no_show") else "cancelled"
            if yes:
                b.cancel_fee = float(show.get("amount") or 0)
                s.add(FoCharge(tenant_id=tenant, booking_id=b.id, kind="cancel_fee", text="Cancellation / no-show charge",
                               rate=b.cancel_fee, on=today(tenant).isoformat(), by="Owner"))
                say = f"{b.number} radd, charge {_rs(b.cancel_fee)}."
            else:
                say = f"{b.number} bina charge radd."
            s.flush()
            paid = _r2(sum(p.amount for p in _payments(s, b)))
            back = _r2(paid - (b.cancel_fee if yes else 0))
            if back > 0:
                follow.append(("refund", b.id, back))
            follow.append(("msg", b.id, "cancel"))
        elif kind == "refund" and yes:
            say = f"{b.number}: {_rs(show.get('amount'))} refund maana. Aap apne app / bank se bhejiye, phir 'Refund bhej diya' dabaiye."
        elif kind == "bill" and yes:
            st = _st(s, tenant)
            d = today(tenant)
            fy = R.fy_of(d)
            f = _folio(s, b)
            bill = FoBill(tenant_id=tenant, booking_id=b.id, status="final", final_at=now(),
                          number=f"INV/{fy}/{_next(s, tenant, 'invoice:' + fy):04d}", data=_dump(f))
            s.add(bill)
            s.flush()
            extra["bill"] = {"id": bill.id, "number": bill.number}
            say = f"Bill {bill.number} final: {_rs(f['total'])}."
            follow.append(("bill_pdf", bill.id, ""))
            follow.append(("msg", b.id, "bill"))
            _ = st
        view = _card_view(c)
        bview = _bview(s, b) if b else {}
    out: dict[str, Any] = {"card": view, "booking": bview, "say": say, **extra}
    for kind2, a1, a2 in follow:
        try:
            if kind2 == "refund":
                r = refund_request(tenant, a1, a2, "radd booking ka paisa", "owner", conv)
                out["refund_card"] = r["card"]
                out["say"] += f" {_rs(a2)} refund ka card bheja."
            elif kind2 == "bill_pdf":
                out["file"] = bill_file(tenant, a1, conv)["file"]
            elif kind2 == "msg":
                bv = booking_view(tenant, a1)
                if bv["phone"] and (a2 == "bill" or bv["source"] not in R.SOURCES_OTA):
                    out["message"] = draft_message(tenant, a1, a2, conv=conv)["message"]
                    out["say"] += " Mehmaan ke liye message draft taiyaar (aapke tap se jaayega)."
        except FoError as e:
            out["say"] += f" ({e})"
    _publish(tenant, "frontoffice_card", {"card": view, "say": out["say"]}, conv)
    return out


def refund_done(tenant: str, card_id: str, ref: str = "", mode: str = "upi", actor: str = "owner") -> dict:
    """She sent the approved refund herself: recorded as money out."""
    _owner_only(actor, "Refund darj karna")
    with _LOCK, session() as s:
        c = s.get(FoCard, card_id)
        if c is None or c.tenant_id != tenant or c.kind != "refund":
            raise FoError("Refund card nahi mila.")
        if c.status != "approved":
            raise FoError("Pehle refund card par aapki haan chahiye.")
        show = _j(c.show, {})
        if show.get("done"):
            raise FoError("Yeh refund pehle hi darj hai.")
        d = today(tenant)
        s.add(FoPayment(tenant_id=tenant, booking_id=c.target, amount=-_r2(show.get("amount") or 0), mode=mode if mode in
                        ("upi", "cash", "card", "bank") else "upi", ref_tail=re.sub(r"\W", "", ref)[-4:], on=d.isoformat(), by="Owner",
                        receipt_no=f"RF/{R.fy_of(d)}/{_next(s, tenant, 'refund:' + R.fy_of(d)):04d}"))
        show["done"] = True
        c.show = _dump(show)
        b = s.get(FoBooking, c.target)
        view = _bview(s, b)
    return {"booking": view, "say": f"{view['number']}: refund {_rs(show.get('amount'))} darj."}


# ---------------------------------------------------------------- bill PDF / files
def _pdf(title: str, blocks: list[tuple]) -> bytes:
    from .. import indicpdf
    pdf = indicpdf.blocks_pdf(title, blocks)
    if pdf:
        return pdf
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
    from xml.sax.saxutils import escape
    buf = io.BytesIO()
    ss = getSampleStyleSheet()
    story: list[Any] = []
    for b in blocks:
        k = b[0]
        if k == "h1":
            story.append(Paragraph(escape(str(b[1])), ss["Title"]))
        elif k == "h":
            story.append(Paragraph(escape(str(b[1])), ss["Heading3"]))
        elif k == "p":
            story.append(Paragraph(escape(str(b[1])).replace("\n", "<br/>"), ss["Normal"]))
        elif k == "table":
            t = Table([[str(c if c is not None else "") for c in r] for r in (b[1] or [])], repeatRows=1)
            t.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.3, colors.grey), ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#eeeeee")),
                                   ("FONTSIZE", (0, 0), (-1, -1), 8)]))
            story.append(t)
        story.append(Spacer(1, 4))
    SimpleDocTemplate(buf, pagesize=A4, title=title).build(story)
    return buf.getvalue()


def bill_pdf(tenant: str, bill: str) -> bytes:
    with session() as s:
        bl = s.get(FoBill, bill)
        if bl is None or bl.tenant_id != tenant:
            raise FoError("Bill nahi mila.")
        b = s.get(FoBooking, bl.booking_id)
        g = s.get(FoGuest, b.guest_id)
        st = _st(s, tenant)
        rm = s.get(FoRoom, b.room_id) if b.room_id else None
        f = _j(bl.data, {})
        pays = _payments(s, b)
    head = "TAX INVOICE" if st.gstin else "BILL OF SUPPLY / BILL"
    blocks: list[tuple] = [("h1", st.hotel_name or "Hotel"),
                           ("p", "\n".join(x for x in (st.address, f"GSTIN: {st.gstin}" if st.gstin else "", st.phone) if x)),
                           ("h", f"{head} {bl.number}  ·  {(_aware(bl.final_at) or now()).astimezone(_tz(tenant)).date().isoformat()}"),
                           ("p", f"Guest: {g.name if g and not g.forgotten else ''}" + (f"\nCompany: {g.company} GSTIN {g.gstin}" if g and g.gstin else "")
                            + f"\nBooking {b.number} · Room {rm.number if rm else ''} · {b.arrive} to {b.depart} · {b.adults} adult(s)")]
    rows = [["Item", "SAC", "Value", "GST %", "CGST", "SGST", "Total"]]
    for ln in f.get("lines", []):
        rows.append([ln.get("text"), ln.get("sac", "996311" if ln.get("kind") == "room" else ""), f"{ln.get('value') or 0:.2f}",
                     ln.get("gst_pct"), f"{ln.get('cgst') or 0:.2f}", f"{ln.get('sgst') or 0:.2f}", f"{ln.get('total') or 0:.2f}"])
    rows.append(["Total", "", f"{f.get('taxable', 0):.2f}", "", f"{f.get('cgst', 0):.2f}", f"{f.get('sgst', 0):.2f}", f"{f.get('total', 0):.2f}"])
    blocks.append(("table", rows))
    try:
        from ..india.money import rupees_words
        blocks.append(("p", f"Amount in words: {rupees_words(f.get('total', 0), 'en')}"))
    except Exception:
        pass
    if f.get("discount"):
        blocks.append(("p", f"Discount applied: {f['discount']:.2f}" + (" (complimentary)" if f.get("complimentary") else "")))
    prow = [["Receipt", "Date", "Mode", "Amount"]] + [[p.receipt_no, p.on, p.mode, f"{p.amount:.2f}"] for p in pays]
    blocks += [("h", "Payments"), ("table", prow), ("p", f"Paid: {f.get('paid', 0):.2f}   Balance: {(f.get('total', 0) - sum(p.amount for p in pays)):.2f}")]
    blocks.append(("p", f.get("note", "")))
    return _pdf(f"Bill {bl.number}", blocks)


def bill_file(tenant: str, bill: str, conv: str = "") -> dict:
    from .. import chatfiles
    data = bill_pdf(tenant, bill)
    with session() as s:
        bl = s.get(FoBill, bill)
        name = f"Bill_{bl.number.replace('/', '-')}.pdf"
    fl = chatfiles.store(tenant, name, data, "application/pdf")
    with session() as s:
        s.get(FoBill, bill).file_id = fl.get("file_id", "")
    _publish(tenant, "pc_file", {**fl, "source": "frontoffice", "title": name}, conv)
    return {"file": fl, "say": f"{name} taiyaar."}


# ---------------------------------------------------------------- messages (DRAFTS; her tap sends)
def _msg_view(s, m: FoMessage) -> dict:
    g = s.get(FoGuest, m.guest_id) if m.guest_id else None
    ph = g.phone if g and not g.forgotten else ""
    b = s.get(FoBooking, m.booking_id)
    return {"id": m.id, "kind": m.kind, "text": m.text, "hash": m.hash, "status": m.status, "via": m.via,
            "to": g.name if g else "", "phone": ph, "no_messages": bool(g and g.no_messages), "booking": b.number if b else "",
            "wa": f"https://wa.me/91{ph}?text={quote(m.text)}" if ph else "", "sms": f"sms:{ph}?body={quote(m.text)}" if ph else "",
            "sent_at": _aware(m.sent_at).isoformat(timespec="minutes") if m.sent_at else ""}


def _msg_text(s, b: FoBooking, kind: str, text: str, tenant: str) -> str:
    st = _st(s, tenant)
    g = s.get(FoGuest, b.guest_id)
    t = s.get(FoRoomType, b.type_id)
    f = _folio(s, b)
    hotel = st.hotel_name or "[hotel ka naam]"
    nm = (g.name.split()[0] if g and g.name else "ji")
    total = _rs(f["total"]) if not f["blanks"] else "[daam?]"
    pol = R.policy_text(_j(b.policy, {}))
    if kind == "confirmation":
        upi, _n = _upi(st)
        lines = [f"Namaste {nm} ji, {hotel} mein aapki booking {b.number} pakki hai.",
                 f"{t.name if t else 'Kamra'} · {b.arrive} se {b.depart} ({len(_j(b.nights, []))} raat) · {b.adults} log.",
                 f"Kul: {total}" + (" (GST sahit)" if st.gstin and not f["blanks"] else "") + ".",
                 f"Check-in {st.check_in_time}, check-out {st.check_out_time}."]
        if f["balance"] and f["balance"] > 0 and upi:
            lines.append(f"Advance / baaki UPI se: {upi} (note mein {b.number} likhiye).")
        lines.append(f"Cancellation: {pol}" if pol else "Cancellation policy: [policy]")
        lines.append("Aate waqt photo ID saath laaiye (videshi mehmaan: passport + visa).")
        if st.address:
            lines.append(f"Pata: {st.address}")
        return "\n".join(lines)
    if kind == "pre_arrival":
        return (f"Namaste {nm} ji, kal {b.arrive} ko {hotel} mein aapka intezaar hai (booking {b.number}). Check-in {st.check_in_time} "
                f"ke baad. Photo ID saath laaiye." + (f"\nPata: {st.address}" if st.address else "")
                + (f"\nPhone: {st.phone}" if st.phone else ""))
    if kind == "payment":
        link = pay_link(tenant, b.id)
        upi, _n = _upi(st)
        return (f"Namaste {nm} ji, booking {b.number}: baaki {_rs(link.get('amount'))}. UPI: {upi} (note mein {b.number}).\n"
                "Hum kabhi OTP / PIN nahi maangte.")
    if kind == "review":
        link = st.review_link or "[review link]"
        return (f"Namaste {nm} ji, {hotel} mein rukne ke liye dhanyavaad! Aapka anubhav kaisa raha? Samay ho to apni raay yahan likhiye: "
                f"{link}\n(Achhi ho ya buri, har raay se hum seekhte hain.)")
    if kind == "bill":
        bl = s.scalars(select(FoBill).where(FoBill.booking_id == b.id, FoBill.status == "final")).first()
        return (f"Namaste {nm} ji, {hotel} ka bill {bl.number if bl else ''}: {total}. Mila {_rs(f['paid'])}"
                + (f", baaki {_rs(f['balance'])}" if f["balance"] and f["balance"] > 0 else "") + ". Dhanyavaad!")
    if kind == "cancel":
        return (f"Namaste {nm} ji, aapki booking {b.number} ({b.arrive} se {b.depart}) radd kar di gayi hai."
                + (f" Cancellation charge {_rs(b.cancel_fee)}." if b.cancel_fee else "")
                + " Koi paisa wapas banta ho to alag se bataenge.")
    return _clean(text, 1500)


def draft_message(tenant: str, booking: str, kind: str = "custom", text: str = "", conv: str = "", publish: bool = True) -> dict:
    """A DRAFT for the guest. Nothing is sent here: only `send_message(by_tap=True)` from her app."""
    from ..net import policy
    if text and (policy.secret_found(text) or re.search(r"\b(otp|pin|cvv|password|mpin)\b", text, re.I)):
        raise FoError("Message mein OTP / PIN / card / Aadhaar jaisi cheez hai; yeh kabhi nahi bheja jaata.")
    kind = kind if kind in ("confirmation", "pre_arrival", "payment", "review", "bill", "cancel", "custom") else "custom"
    with _LOCK, session() as s:
        b = _booking(s, tenant, booking)
        g = s.get(FoGuest, b.guest_id)
        if g is None or g.forgotten:
            raise FoError("Mehmaan ki jaankari nahi.")
        body = _msg_text(s, b, kind, text, tenant)
        if not body.strip():
            raise FoError("Message khaali hai.")
        if policy.secret_found(body):
            raise FoError("Message mein ID / card jaisi cheez hai; yeh kabhi nahi bheja jaata.")
        for o in s.scalars(select(FoMessage).where(FoMessage.booking_id == b.id, FoMessage.kind == kind, FoMessage.status == "draft")).all():
            if kind != "custom":
                o.status = "dropped"
        m = FoMessage(tenant_id=tenant, booking_id=b.id, guest_id=g.id, kind=kind, text=body[:3000])
        s.add(m)
        s.flush()
        m.hash = _hash({"id": m.id, "text": m.text, "phone": g.phone})
        view = _msg_view(s, m)
    say = f"{view['to']} ke liye {kind} draft: aapke tap se hi jaayega."
    if "[" in view["text"] and "]" in view["text"]:
        say += " (Khaali jagah [ ] bhariye.)"
    if publish:
        _publish(tenant, "frontoffice_message", {"message": view, "say": say}, conv)
    return {"message": view, "status": "needs_tap", "say": say}


def edit_message(tenant: str, msg: str, text: str) -> dict:
    from ..net import policy
    with _LOCK, session() as s:
        m = s.get(FoMessage, msg)
        if m is None or m.tenant_id != tenant or m.status != "draft":
            raise FoError("Draft nahi mila.")
        if policy.secret_found(text) or re.search(r"\b(otp|pin|cvv|password|mpin)\b", text, re.I):
            raise FoError("Message mein OTP / PIN / card jaisi cheez hai; yeh kabhi nahi bheja jaata.")
        m.text = _clean(text, 3000)
        g = s.get(FoGuest, m.guest_id)
        m.hash = _hash({"id": m.id, "text": m.text, "phone": g.phone if g else ""})
        return {"message": _msg_view(s, m), "say": "Draft badla."}


async def send_message(tenant: str, msg: str, via: str = "self", card_hash: str = "", by_tap: bool = False) -> dict:
    """ONLY her tap. via=self: she sends from her own WhatsApp / SMS (the app opens the link; this marks it sent).
    via=business: her business WhatsApp number, only if the guest wrote to it (Meta's 24 h window, enforced by the inbox)."""
    if not by_tap:
        raise FoError("Mehmaan ko message sirf aapke tap se jaata hai.")
    with session() as s:
        m = s.get(FoMessage, msg)
        if m is None or m.tenant_id != tenant:
            raise FoError("Message nahi mila.")
        if m.status != "draft":
            raise FoError(f"Yeh message {m.status} hai.")
        g = s.get(FoGuest, m.guest_id)
        now_hash = _hash({"id": m.id, "text": m.text, "phone": g.phone if g else ""})
        if not card_hash or card_hash != m.hash or now_hash != m.hash:
            raise FoError("Message badal gaya hai; naya draft dekhiye.")
        if g is None or not g.phone:
            raise FoError("Mehmaan ka mobile nahi.")
        if re.search(r"\[(daam\?|samay|policy|review link|hotel ka naam)\]", m.text):
            raise FoError("Message mein khaali jagah [ ] hai; pehle bhariye.")
        text, phone = m.text, g.phone
    if via == "business":
        from ..employee import inbox
        from ..employee.models import BizContact
        with session() as s:
            bc = s.scalars(select(BizContact).where(BizContact.tenant_id == tenant, BizContact.wa_id == "91" + phone)).first()
            bid = bc.id if bc else ""
        if not bid:
            raise FoError("Mehmaan ne business number par kabhi message nahi kiya: apne phone se bhejiye.")
        try:
            await inbox.send(tenant, bid, text, by="owner")
        except inbox.InboxError as e:
            raise FoError(str(e))
    elif via != "self":
        raise FoError("via: self | business")
    with session() as s:
        m = s.get(FoMessage, msg)
        m.status, m.via, m.sent_at = "sent", via, now()
        return {"message": _msg_view(s, m), "say": "Bhej diya (aapke tap se)." if via == "business" else "Aapke phone se bheja, darj kiya."}


def drop_message(tenant: str, msg: str) -> dict:
    with session() as s:
        m = s.get(FoMessage, msg)
        if m is None or m.tenant_id != tenant:
            raise FoError("Message nahi mila.")
        if m.status == "draft":
            m.status = "dropped"
        return {"message": _msg_view(s, m), "say": "Draft hataya."}


def messages(tenant: str, status: str = "draft") -> list[dict]:
    with session() as s:
        q = select(FoMessage).where(FoMessage.tenant_id == tenant)
        if status:
            q = q.where(FoMessage.status == status)
        return [_msg_view(s, m) for m in s.scalars(q.order_by(FoMessage.created_at.desc())).all()[:100]]


# ---------------------------------------------------------------- OTA mails / sheets (her paste = data)
def import_ota(tenant: str, text: str, room_type: str = "", actor: str = "owner", conv: str = "") -> dict:
    """One pasted OTA confirmation / modification / cancellation mail. Text is DATA, never instructions."""
    p = R.parse_ota(text, today(tenant))
    return _import_one(tenant, p, room_type, actor, conv)


def _import_one(tenant: str, p: dict, room_type: str, actor: str, conv: str) -> dict:
    if p.get("cancelled"):
        if not p.get("ref"):
            raise FoError("Cancellation mail mein booking ID nahi mila.")
        with session() as s:
            rows = s.scalars(select(FoBooking).where(FoBooking.tenant_id == tenant, FoBooking.ota_ref.like(p["ref"] + "%"),
                                                     FoBooking.status.in_(("hold", "confirmed")))).all()
            ids = [b.id for b in rows]
        if not ids:
            return {"parsed": p, "status": "not_found", "say": f"{p['ref']} ki koi chalu booking nahi mili (pehle hi radd ya kabhi aayi hi nahi)."}
        outs = [cancel(tenant, i, "OTA cancellation mail", by_ota=True, actor=actor, conv=conv) for i in ids]
        return {"parsed": p, "status": "cancelled", "bookings": [o["booking"] for o in outs],
                "say": f"{p['source']} {p['ref']}: {len(ids)} booking radd (OTA ne)."}
    miss = [m for m in ("arrive", "depart", "guest") if m in p.get("missing", [])]
    if miss:
        return {"parsed": p, "status": "needs_info", "say": "Mail se yeh nahi mila: " + ", ".join(miss) + ". Bataiye ya poora mail paste kijiye."}
    with session() as s:
        rows = s.scalars(select(FoRoomType).where(FoRoomType.tenant_id == tenant, FoRoomType.archived.is_(False))).all()
        tlist = [{"id": t.id, "name": t.name, "code": t.code, "aliases": t.aliases} for t in rows]
        existing = s.scalars(select(FoBooking).where(FoBooking.tenant_id == tenant, FoBooking.ota_ref == p["ref"])).first() if p.get("ref") else None
        ex_id = existing.id if existing else ""
        ex_status = existing.status if existing else ""
    tt = None
    if room_type:
        tt = {"id": room_type}
    else:
        tt = R.match_room_type(p.get("room_type", ""), tlist)
    if not tt:
        return {"parsed": p, "status": "needs_type", "types": [t["code"] + " " + t["name"] for t in tlist],
                "say": f"OTA ka room '{p.get('room_type') or '?'}' aapke kis room type ka hai? (room_type=...; aliases mein jod dein to agli baar khud)."}
    if ex_id:
        if ex_status not in ("hold", "confirmed"):
            return {"parsed": p, "status": "exists", "say": f"{p['ref']} pehle se hai ({STATUS_LABEL.get(ex_status)})."}
        try:
            r = amend(tenant, ex_id, p["arrive"], p["depart"], tt["id"], p.get("adults") or None, actor=actor, conv=conv)
        except FoError as e:
            _overbook_alert(tenant, p, str(e), conv)
            raise
        return {"parsed": p, "status": "updated", "booking": r["booking"], "say": f"{p['ref']} badli (OTA modification): " + r["say"]}
    n = max(1, int(p.get("rooms") or 1))
    amt = p.get("amount")
    per_room = _r2(amt / n) if amt else None
    made, err = [], ""
    for i in range(n):
        try:
            r = book(tenant, tt["id"], p["arrive"], p["depart"], name=p["guest"], phone=p.get("phone") or "",
                     adults=max(1, (p.get("adults") or n) // n), children=p.get("children") or 0, source=p["source"],
                     ota_ref=p["ref"] + (f"-{i + 1}" if n > 1 else ""), ota_amount=per_room, paid_online=p.get("paid_online"),
                     pay_at_hotel=p.get("pay_at_hotel"), commission=p.get("commission"), meal_plan=p.get("meal_plan") or "",
                     notes=f"{R.SOURCE_LABEL.get(p['source'], p['source'])} mail se", actor=actor, conv=conv, draft_confirmation=False)
            made.append(r["booking"])
        except FoError as e:
            err = str(e)
            break
    if err:
        _overbook_alert(tenant, p, err, conv)
        return {"parsed": p, "status": "refused", "bookings": made, "say": err + " OTA extranet par turant stop-sell / mehmaan ko shift."}
    return {"parsed": p, "status": "booked", "bookings": made,
            "say": f"{R.SOURCE_LABEL.get(p['source'], p['source'])} {p['ref']}: {len(made)} booking ({p['arrive']} se {p['depart']})."}


def _overbook_alert(tenant: str, p: dict, err: str, conv: str) -> None:
    if err.startswith("OVERBOOK"):
        _publish(tenant, "frontoffice_alert", {"kind": "overbook", "ref": p.get("ref"), "source": p.get("source"),
                                              "say": f"{R.SOURCE_LABEL.get(p.get('source'), '')} {p.get('ref')} nahi samaayi: {err} "
                                                     "OTA par stop-sell kijiye / mehmaan ko doosre hotel mein shift (Hoola ne booking nahi banayi)."}, conv)


def import_sheet(tenant: str, data: bytes, name: str = "", source: str = "", room_type: str = "", actor: str = "owner",
                 conv: str = "") -> dict:
    """Her exported OTA / channel-manager sheet (CSV / Excel). Rows are DATA."""
    rows: list[list[Any]] = []
    if name.lower().endswith((".xlsx", ".xlsm")):
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        ws = wb.active
        rows = [list(r) for r in ws.iter_rows(values_only=True)][:2000]
    else:
        txt = data.decode("utf-8-sig", errors="replace")
        rows = [r for r in csv.reader(io.StringIO(txt))][:2000]
    hdr_i, cols = -1, {}
    for i, r in enumerate(rows[:15]):
        c = R.map_columns([str(x or "") for x in r])
        if "arrive" in c and "depart" in c:
            hdr_i, cols = i, c
            break
    if hdr_i < 0:
        raise FoError("Sheet mein check-in / check-out column nahi mila.")
    src = source if source in R.SOURCE_LABEL else (R.ota_source(name) or "other")
    done = {"booked": 0, "cancelled": 0, "updated": 0, "refused": 0, "skipped": 0}
    notes: list[str] = []
    for r in rows[hdr_i + 1:]:
        if not any(str(x or "").strip() for x in r):
            continue
        p = R.row_booking(list(r), cols, src, today(tenant))
        p["missing"] = [k for k in ("arrive", "depart", "guest") if not p.get(k)]
        try:
            out = _import_one(tenant, p, room_type, actor, conv)
            st = out.get("status", "")
            done[st if st in done else "skipped"] += 1
            if st not in ("booked", "cancelled", "updated"):
                notes.append(f"{p.get('ref') or p.get('guest')}: {out.get('say')}")
        except FoError as e:
            done["refused"] += 1
            notes.append(f"{p.get('ref') or p.get('guest')}: {e}")
    say = ", ".join(f"{k} {v}" for k, v in done.items() if v) or "Kuch nahi mila."
    return {"result": done, "notes": notes[:30], "say": "Sheet: " + say}


def import_file(tenant: str, file_id: str, source: str = "", room_type: str = "", actor: str = "owner", conv: str = "") -> dict:
    from ..mailbox import service as ms
    try:
        f = ms._file_row(tenant, str(file_id))
        data = Path(f.stored_path).read_bytes()
    except Exception as e:
        raise FoError(f"File nahi mili: {e}")
    if f.name.lower().endswith((".csv", ".xlsx", ".xlsm")):
        return import_sheet(tenant, data, f.name, source, room_type, actor, conv)
    try:
        from .. import docread
        text = docread.read_text(f.name, "", data).get("text") or ""
    except Exception:
        text = data.decode("utf-8", "replace")
    return import_ota(tenant, text, room_type, actor, conv)


# ---------------------------------------------------------------- Form C (Form III) for foreign guests
def _fc_view(s, x: FoFormC | None) -> dict:
    if x is None:
        return {}
    g = s.get(FoGuest, x.guest_id)
    b = s.get(FoBooking, x.booking_id)
    due = _aware(x.due_at)
    tz = _tz(x.tenant_id)
    return {"id": x.id, "kind": x.kind, "status": x.status, "ref": x.ref, "booking": b.number if b else "",
            "guest": g.name if g and not g.forgotten else "", "due_at": due.isoformat() if due else "",
            "due_local": due.astimezone(tz).strftime("%d-%m %H:%M") if due else "",
            "overdue": bool(due and x.status == "due" and clock() > due)}


def formc_list(tenant: str, status: str = "due") -> list[dict]:
    with session() as s:
        q = select(FoFormC).where(FoFormC.tenant_id == tenant)
        if status:
            q = q.where(FoFormC.status == status)
        return [_fc_view(s, x) for x in s.scalars(q.order_by(FoFormC.due_at)).all()[:200]]


def formc_data(tenant: str, fid: str) -> dict:
    """The fields she types into the e-FRRO Form III (C-Form). ID numbers masked; the app reveals them from her locker."""
    with session() as s:
        x = s.get(FoFormC, fid)
        if x is None or x.tenant_id != tenant:
            raise FoError("Form C entry nahi mili.")
        g = s.get(FoGuest, x.guest_id)
        b = s.get(FoBooking, x.booking_id)
        st = _st(s, tenant)
        if g is None or g.forgotten:
            raise FoError("Mehmaan ki jaankari nahi.")
        ids = _j(g.ids, {})
        visa = _j(g.visa, {})
        tz = _tz(tenant)
        ev = _aware(x.event_at)
        fields = [
            ("Entry", "Arrival" if x.kind == "arrival" else "Check Out / Departure"),
            ("Guest name", g.name), ("Nationality", g.nationality),
            ("Passport no.", (ids.get("passport") or {}).get("masked", "[passport?]")),
            ("Visa / OCI no.", (visa.get("number") or {}).get("masked", "[visa?]")), ("Visa type", visa.get("type") or "[visa type?]"),
            ("Visa valid till", visa.get("expiry") or "[?]"), ("Arrived in India on", g.arrived_india_on or "[?]"),
            ("Arrived at hotel", ev.astimezone(tz).strftime("%d-%m-%Y %H:%M") if ev and x.kind == "arrival" else (b.checked_in_at and _aware(b.checked_in_at).astimezone(tz).strftime("%d-%m-%Y %H:%M")) or ""),
            ("Departed hotel", ev.astimezone(tz).strftime("%d-%m-%Y %H:%M") if ev and x.kind == "departure" else ""),
            ("Coming from", b.coming_from or "[?]"), ("Going to", b.going_to or "[?]"), ("Purpose", b.purpose or "[?]"),
            ("Address in India (hotel)", st.address or "[hotel ka pata]"), ("Indian phone", g.phone or ""),
        ]
        reveal = {k: (ids.get(k) or {}).get("record") for k in ("passport",)}
        if (visa.get("number") or {}).get("record"):
            reveal["visa"] = visa["number"]["record"]
        view = _fc_view(s, x)
    return {"formc": view, "fields": [{"label": a, "value": v} for a, v in fields], "reveal_records": reveal, "portal": FRRO_URL,
            "say": f"Form C ({view['kind']}) {view['due_local']} tak: indianfrro.gov.in par login karke bhariye; "
                   "passport / visa number Pehchaan Locker se dekhiye. Hoola khud submit nahi karta."}


def formc_done(tenant: str, fid: str, ref: str = "", actor: str = "owner") -> dict:
    with _LOCK, session() as s:
        x = s.get(FoFormC, fid)
        if x is None or x.tenant_id != tenant:
            raise FoError("Form C entry nahi mili.")
        x.status, x.ref, x.submitted_at = "submitted", _clean(ref, 60), now()
        view = _fc_view(s, x)
    return {"formc": view, "say": f"Form C {view['kind']} ({view['guest']}) bhara gaya darj" + (f", ref {view['ref']}." if view["ref"] else ".")}


# ---------------------------------------------------------------- reports
def report(tenant: str, start: str = "", end: str = "") -> dict:
    """Occupancy, ADR, RevPAR (room revenue before GST, after discount) per night and for the range; by source."""
    s0 = _date(start, tenant) or today(tenant)
    e0 = _date(end, tenant) or s0
    if e0 < s0 or (e0 - s0).days > 92:
        raise FoError("Report ki range sahi nahi (92 din tak).")
    days = [s0 + timedelta(days=i) for i in range((e0 - s0).days + 1)]
    with session() as s:
        types_ = s.scalars(select(FoRoomType).where(FoRoomType.tenant_id == tenant)).all()
        rows = s.scalars(select(FoBooking).where(FoBooking.tenant_id == tenant,
                                                 FoBooking.status.in_(("confirmed", "in_house", "checked_out")))).all()
        per_day, src = [], {}
        tot_av = tot_sold = 0
        tot_rev = 0.0
        for d in days:
            av = sum(_sellable(s, tenant, t.id, d) for t in types_ if not t.archived)
            sold, rev = 0, 0.0
            for b in rows:
                for n in _j(b.nights, []):
                    if n["date"] != d.isoformat():
                        continue
                    sold += 1
                    gross = sum((x["rate"] or 0) + (x.get("extra") or 0) for x in _j(b.nights, []))
                    v = (n["rate"] or 0) + (n.get("extra") or 0)
                    if b.complimentary:
                        v = 0
                    elif gross and b.discount:
                        v = v * (1 - min(b.discount, gross) / gross)
                    rev += v
                    sv = src.setdefault(b.source, {"room_nights": 0, "revenue": 0.0})
                    sv["room_nights"] += 1
                    sv["revenue"] = _r2(sv["revenue"] + v)
            per_day.append({"date": d.isoformat(), **R.kpis(av, sold, rev)})
            tot_av, tot_sold, tot_rev = tot_av + av, tot_sold + sold, tot_rev + rev
    k = R.kpis(tot_av, tot_sold, tot_rev)
    by_src = [{"source": a, "label": R.SOURCE_LABEL.get(a, a), **v} for a, v in sorted(src.items(), key=lambda x: -x[1]["room_nights"])]
    say = (f"{s0.isoformat()}" + (f" se {e0.isoformat()}" if e0 != s0 else "") + f": occupancy {k['occupancy_pct']}%, "
           f"ADR {_rs(k['adr'])}, RevPAR {_rs(k['revpar'])} ({k['rooms_sold']}/{k['rooms_available']} room-nights).")
    return {"from": s0.isoformat(), "to": e0.isoformat(), "total": k, "days": per_day, "by_source": by_src, "say": say,
            "note": "Revenue = kamre ka kiraya, GST se pehle, discount ke baad; complimentary 0. OTA amount jaisa mail mein tha."}


def register_file(tenant: str, start: str = "", end: str = "", conv: str = "") -> dict:
    """Guest register (police / audit) + bookings Excel. ID numbers only masked."""
    from .. import chatfiles, officegen
    s0 = _date(start, tenant) or today(tenant) - timedelta(days=30)
    e0 = _date(end, tenant) or today(tenant)
    with session() as s:
        reg = [["Booking", "Guest", "Nationality", "Phone", "ID (masked)", "Room", "Check-in", "Check-out", "From", "To", "Purpose", "Companions"]]
        bk = [["Booking", "Source", "OTA ref", "Guest", "Type", "Room", "Arrive", "Depart", "Status", "Total", "Paid"]]
        for b in s.scalars(select(FoBooking).where(FoBooking.tenant_id == tenant).order_by(FoBooking.arrive)).all():
            a = R.parse_date(b.arrive)
            if not a or not (s0 <= a <= e0):
                continue
            g = s.get(FoGuest, b.guest_id)
            r = s.get(FoRoom, b.room_id) if b.room_id else None
            t = s.get(FoRoomType, b.type_id)
            f = _folio(s, b)
            gn = g.name if g and not g.forgotten else "(mitaya)"
            bk.append([b.number, R.SOURCE_LABEL.get(b.source, b.source), b.ota_ref, gn, t.name if t else "", r.number if r else "",
                       b.arrive, b.depart, STATUS_LABEL.get(b.status, b.status), f["total"], f["paid"]])
            if b.checked_in_at:
                ids = _j(g.ids, {}) if g else {}
                reg.append([b.number, gn, g.nationality if g else "", g.phone if g and not g.forgotten else "",
                            ", ".join(f"{k}: {v.get('masked')}" for k, v in ids.items()), r.number if r else "",
                            _aware(b.checked_in_at).astimezone(_tz(tenant)).strftime("%d-%m-%Y %H:%M"),
                            _aware(b.checked_out_at).astimezone(_tz(tenant)).strftime("%d-%m-%Y %H:%M") if b.checked_out_at else "",
                            b.coming_from, b.going_to, b.purpose, ", ".join(_j(b.companions, []))])
    data = officegen.make_xlsx("Guest register", {"Guest register": reg, "Bookings": bk})
    fl = chatfiles.store(tenant, f"Guest_register_{s0.isoformat()}_{e0.isoformat()}.xlsx", data,
                         "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    _publish(tenant, "pc_file", {**fl, "source": "frontoffice", "title": "Guest register"}, conv)
    return {"file": fl, "rows": len(reg) - 1, "say": f"Guest register ({len(reg) - 1} check-in) aur bookings ka Excel."}


# ---------------------------------------------------------------- day / overview
def day(tenant: str) -> dict:
    td = today(tenant).isoformat()
    with session() as s:
        rows = s.scalars(select(FoBooking).where(FoBooking.tenant_id == tenant)).all()
        arr = [_bview(s, b) for b in rows if b.arrive == td and b.status in ("hold", "confirmed")]
        dep = [_bview(s, b) for b in rows if b.depart == td and b.status == "in_house"]
        inh = [_bview(s, b) for b in rows if b.status == "in_house"]
        late = [_bview(s, b) for b in rows if b.status in ("hold", "confirmed") and b.arrive < td]
        holds = [_bview(s, b) for b in rows if b.status == "hold"]
        unpaid = [_bview(s, b) for b in rows if b.status == "checked_out"]
        unpaid = [v for v in unpaid if v["balance"] and v["balance"] > 0]
        fc = [_fc_view(s, x) for x in s.scalars(select(FoFormC).where(FoFormC.tenant_id == tenant, FoFormC.status == "due")).all()]
        dirty = s.scalars(select(FoRoom.number).where(FoRoom.tenant_id == tenant, FoRoom.status == "dirty", FoRoom.archived.is_(False))).all()
        waiting = s.scalars(select(FoCard).where(FoCard.tenant_id == tenant, FoCard.status == "waiting")).all()
        drafts = s.scalars(select(FoMessage).where(FoMessage.tenant_id == tenant, FoMessage.status == "draft")).all()
    parts = []
    if arr:
        parts.append(f"Aaj aane wale {len(arr)}: " + ", ".join(f"{v['number']} {v['guest']}" for v in arr[:5]))
    if dep:
        parts.append(f"Aaj jaane wale {len(dep)}: " + ", ".join(f"{v['number']} room {v['room']}" for v in dep[:5]))
    if fc:
        parts.append(f"Form C baaki {len(fc)}" + (" (kuch der se!)" if any(x["overdue"] for x in fc) else "") + ": "
                     + ", ".join(f"{x['guest']} {x['kind']} {x['due_local']}" for x in fc[:5]))
    if late:
        parts.append(f"No-show? {len(late)}: " + ", ".join(v["number"] for v in late[:5]))
    if holds:
        parts.append(f"Hold {len(holds)} (advance baaki)")
    if dirty:
        parts.append(f"Gande kamre {len(dirty)}: " + ", ".join(dirty[:10]))
    if unpaid:
        parts.append(f"Baaki paisa {len(unpaid)}: " + ", ".join(f"{v['number']} {_rs(v['balance'])}" for v in unpaid[:5]))
    if waiting:
        parts.append(f"Aapki haan ke {len(waiting)} card")
    if drafts:
        parts.append(f"{len(drafts)} message draft")
    return {"arrivals": arr, "departures": dep, "in_house": inh, "formc_due": fc, "no_show_maybe": late, "holds": holds,
            "dirty": list(dirty), "unpaid": unpaid, "waiting": len(waiting), "drafts": len(drafts),
            "say": ". ".join(parts) + "." if parts else "Front desk par sab theek."}


def overview(tenant: str) -> dict:
    return {"settings": settings(tenant), "types": types(tenant), "board": board(tenant), "day": day(tenant),
            "upcoming": [b for b in bookings(tenant, "hold,confirmed") if b["arrive"] >= today(tenant).isoformat()][:50],
            "cards": cards(tenant, "waiting"), "messages": messages(tenant, "draft"), "formc": formc_list(tenant, "due"),
            "report": report(tenant), "never": list(R.NEVER)}


# ---------------------------------------------------------------- DPDP forget
def forget(tenant: str, guest: str, actor: str = "owner") -> dict:
    """Guest's personal details out (DPDP), after their stays ended. Foreign guests' Form C records are kept 1 year
    (Immigration and Foreigners Rules 2025 r.17(1)); a final GST bill keeps the name (GST records 72 months)."""
    _owner_only(actor, "Mehmaan ki jaankari mitana")
    with _LOCK, session() as s:
        g = s.scalars(select(FoGuest).where(FoGuest.tenant_id == tenant, (FoGuest.id == guest) | (FoGuest.number == str(guest).upper()))).first()
        if g is None:
            raise FoError("Mehmaan nahi mila.")
        bs = s.scalars(select(FoBooking).where(FoBooking.guest_id == g.id)).all()
        if any(b.status in HOLDING for b in bs):
            raise FoError("Is mehmaan ki booking abhi chalu hai.")
        if R.is_foreign(g.nationality):
            last = max((_aware(b.checked_out_at) for b in bs if b.checked_out_at), default=None)
            if last and clock() - last < timedelta(days=365):
                raise FoError("Videshi mehmaan ka record 1 saal rakhna zaroori hai (Immigration and Foreigners Rules 2025 r.17(1)).")
        has_bill = bool(s.scalars(select(FoBill).where(FoBill.booking_id.in_([b.id for b in bs] or [""]), FoBill.status == "final")).first())
        recs = [x.get("record") for x in _j(g.ids, {}).values()] + [(_j(g.visa, {}).get("number") or {}).get("record")]
        if not has_bill:
            g.name = f"Purana mehmaan {g.number}"
        g.phone = g.email = g.address = g.notes = g.company = ""
        g.ids, g.visa, g.forgotten = "{}", "{}", True
        for b in bs:
            b.companions, b.coming_from, b.going_to = "[]", "", ""
        num = g.number
    from .. import idvault
    for r in recs:
        if r:
            try:
                idvault.delete(tenant, r)
            except Exception:
                pass
    return {"guest": num, "say": f"{num} ki niji jaankari mita di" + (" (GST bill ke liye naam rakha)." if has_bill else ".")}


# ---------------------------------------------------------------- worker
def _tenants() -> list[str]:
    with session() as s:
        return sorted(set(s.scalars(select(FoSettings.tenant_id)).all()) | set(s.scalars(select(FoBooking.tenant_id)).all()))


def _morning_due(tenant: str, loc: datetime) -> bool:
    with session() as s:
        r = s.get(FoSettings, tenant)
        if r is None or r.last_morning == loc.date().isoformat():
            return False
        hh, mm = (int(x) for x in (r.morning_time or "08:00").split(":"))
        return (loc.hour, loc.minute) >= (hh, mm)


def tick(when: datetime | None = None) -> dict:
    """Worker, every minute: release expired holds, Form C reminders (6 h left / overdue), pre-arrival drafts the day before,
    review-request drafts the day after check-out, no-show alerts, morning note at her time. Each once."""
    n = _aware(when) or now()
    out = {"released": 0, "formc": 0, "drafts": 0, "alerts": 0, "morning": 0}
    for tenant in _tenants():
        try:
            loc = n.astimezone(_tz(tenant))
            td = loc.date()
            pubs: list[tuple[str, dict]] = []
            jobs: list[tuple[str, str]] = []
            with _LOCK, session() as s:
                for b in s.scalars(select(FoBooking).where(FoBooking.tenant_id == tenant, FoBooking.status == "hold")).all():
                    if b.hold_until and _aware(b.hold_until) < n:
                        b.status = "cancelled"
                        b.notes = (b.notes + "\nHold ka samay khatam (advance nahi aaya): kamra chhoda.").strip()
                        out["released"] += 1
                        pubs.append(("frontoffice_alert", {"kind": "hold_released", "booking": b.number,
                                                           "say": f"{b.number}: hold khatam, advance nahi aaya; kamra wapas bikri mein."}))
                for x in s.scalars(select(FoFormC).where(FoFormC.tenant_id == tenant, FoFormC.status == "due")).all():
                    due = _aware(x.due_at)
                    if not due:
                        continue
                    left = due - n
                    stage = "overdue" if left.total_seconds() < 0 else ("soon" if left <= timedelta(hours=6) else "")
                    if stage and _once(s, tenant, f"fc:{x.id}:{stage}"):
                        v = _fc_view(s, x)
                        out["formc"] += 1
                        pubs.append(("frontoffice_alert", {"kind": "formc", "formc": v,
                                     "say": (f"Form C {v['kind']} {v['guest']}: samay nikal gaya ({v['due_local']}). Turant indianfrro.gov.in par bhariye."
                                             if stage == "overdue" else f"Form C {v['kind']} {v['guest']}: {v['due_local']} tak bharna hai.")}))
                for b in s.scalars(select(FoBooking).where(FoBooking.tenant_id == tenant, FoBooking.status.in_(("confirmed", "hold", "checked_out")))).all():
                    g = s.get(FoGuest, b.guest_id)
                    can_msg = g and not g.forgotten and g.phone and not g.no_messages
                    if b.status in ("confirmed", "hold") and b.arrive == (td + timedelta(days=1)).isoformat() and can_msg \
                            and _once(s, tenant, f"pre:{b.id}"):
                        jobs.append((b.id, "pre_arrival"))
                    if b.status in ("confirmed", "hold") and b.arrive < (td - timedelta(days=0)).isoformat() and loc.hour >= 12 \
                            and _once(s, tenant, f"noshow:{b.id}"):
                        pubs.append(("frontoffice_alert", {"kind": "no_show", "booking": b.number,
                                     "say": f"{b.number} {g.name if g else ''} ({b.arrive}) nahi aaye. No-show karein? (action=no_show)"}))
                    if b.status == "checked_out" and b.checked_out_at and can_msg and \
                            _aware(b.checked_out_at).astimezone(_tz(tenant)).date() < td and loc.hour >= 10 and _once(s, tenant, f"rev:{b.id}"):
                        jobs.append((b.id, "review"))
            for kind, payload in pubs:
                _publish(tenant, kind, payload)
                out["alerts"] += 1
            for bid, k in jobs:
                try:
                    draft_message(tenant, bid, k)
                    out["drafts"] += 1
                except FoError:
                    pass
            if _morning_due(tenant, loc):
                d = day(tenant)
                with session() as s:
                    _st(s, tenant).last_morning = td.isoformat()
                if d["say"] != "Front desk par sab theek.":
                    _publish(tenant, "frontoffice_day", {"say": d["say"], "arrivals": len(d["arrivals"]), "departures": len(d["departures"]),
                                                         "formc": len(d["formc_due"]), "waiting": d["waiting"], "drafts": d["drafts"],
                                                         "report": report(tenant, (td - timedelta(days=1)).isoformat())["total"]})
                    out["morning"] += 1
        except Exception:
            log.exception("frontoffice tick failed for %s", tenant)
    return out
