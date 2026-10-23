"""Shikayat Sahayak: Hoola as the business's complaint handler, end to end.

  1. Gets the job: a complaint she tells in chat / on a call ("Ramesh ji ka geyser 3 din mein band, photo bheji hai"), a file
     or screenshot she forwards, a Hoola Network / business-WhatsApp ticket, or a WhatsApp customer writing to her business
     number (customer-safe tools `biz_shikayat` / `biz_shikayat_status`).
  2. Logs it with proof: number SH-0001, customer, product / order, channel, root cause tags and severity read from the words,
     photos / videos / documents / screenshots kept in her own data folder (not 24 h chat files) with EXIF / GPS / author
     metadata removed; card / OTP / PIN / PAN values in the text are blanked.
  3. Tracks it to closure: stages, who owns it, acknowledgement and final-reply deadlines by the sector's rule (E-Commerce Rules,
     RBI, IRDAI, TRAI, electricity, SEBI, DPDP; her targets may be shorter, never longer), reminders before / after each deadline,
     repeat-customer and repeat-product patterns, escalation records, reopen within N days.
  4. Ready as a card / draft / file: acknowledgement, update and polite closing message drafts (wa.me / SMS, or business WhatsApp
     only inside Meta's 24 h window and only on her tap) with the honest next step for regulated sectors; monthly register Excel
     with a summary she can give a regulator.
  5. Waits for her Haan on COMPENSATION: refund / replacement / discount / goodwill / free service / credit = a card; only her tap in
     the app approves it (hash of the card as it is now; her phone's fingerprint for money once a phone key is set up). Hoola never
     pays and never promises what she has not approved.
  6. Delivers and learns: closure checklist (resolution, proof of fix, customer confirmed or silent N days after the message, no
     compensation pending), satisfaction ask, days-to-close and compensation per root cause.
"""
from __future__ import annotations

import hashlib
import io
import json
import logging
import re
import secrets
import shutil
import subprocess
import threading
from datetime import date, datetime, timedelta, timezone
from pathlib import Path
from typing import Any
from urllib.parse import quote

from sqlalchemy import select

from ..db import now, session
from . import rules as R
from .models import ShikayatCase, ShikayatComp, ShikayatProof, ShikayatSettings

log = logging.getLogger("santosh.shikayat")
_LOCK = threading.Lock()
MAX_PROOF_MB = 25
MAX_VIDEO_MB = 60
MAX_PROOFS = 40
DATA_ROOT = Path("data") / "shikayat"


class ShikayatError(Exception):
    pass


def _now() -> datetime:
    return now()


# ---------------------------------------------------------------- helpers
def _j(s: str | None, default):
    try:
        v = json.loads(s or "")
        return v if isinstance(v, type(default)) else default
    except (TypeError, ValueError):
        return default


def _dump(v) -> str:
    return json.dumps(v, ensure_ascii=False)


def _utc(d: datetime | None) -> datetime | None:
    if d is None:
        return None
    return d if d.tzinfo else d.replace(tzinfo=timezone.utc)


def _digits(v: Any) -> str:
    return re.sub(r"\D", "", str(v or ""))


def _phone(v: Any) -> str:
    d = _digits(v)
    return d[-10:] if len(d) >= 10 else d


def _num(v) -> float:
    if v in (None, ""):
        return 0.0
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).replace(",", "").replace("₹", "")
    m = re.search(r"-?\d+(\.\d+)?", s)
    n = float(m.group(0)) if m else 0.0
    low = s.lower()
    if "lakh" in low or "lac" in low:
        n *= 1e5
    elif "crore" in low:
        n *= 1e7
    return n


def _rs(v) -> str:
    v = float(v or 0)
    whole = round(v)
    s = str(abs(whole))
    if len(s) > 3:
        s = re.sub(r"(\d)(?=(\d\d)+$)", r"\1,", s[:-3]) + "," + s[-3:]
    return ("-" if whole < 0 else "") + "₹" + s


def _tz(tenant: str):
    from ..life import tz_of
    return tz_of(tenant)


def _local(tenant: str, d: datetime | None) -> str:
    d = _utc(d)
    return d.astimezone(_tz(tenant)).strftime("%d %b %Y, %I:%M %p") if d else ""


def _local_day(tenant: str, d: datetime | None) -> str:
    d = _utc(d)
    return d.astimezone(_tz(tenant)).strftime("%d %b %Y") if d else ""


def today(tenant: str) -> date:
    return _now().astimezone(_tz(tenant)).date()


def _key(text: str) -> str:
    t = re.sub(r"[^a-z0-9ऀ-ॿ ]", " ", str(text or "").lower())
    words = [w for w in t.split() if len(w) > 1 and w not in ("ka", "ki", "ke", "the", "and", "aur", "wala", "wali")]
    return " ".join(words[:4])[:80]


def _publish(tenant: str, kind: str, payload: dict, conv: str = "") -> None:
    from ..events import bus
    try:
        if not conv:
            from ..net.service import home_conv
            conv = home_conv(tenant)
        bus.publish(conv, kind, payload)
    except Exception:
        log.exception("shikayat publish failed")


def _note(tenant: str, text: str, conv: str = "") -> None:
    from ..events import notify
    try:
        if not conv:
            from ..net.service import home_conv
            conv = home_conv(tenant)
        notify(conv, text)
    except Exception:
        log.exception("shikayat note failed")


def _ai(system: str, prompt: str, max_tokens: int = 700) -> str:
    from ..net.ai import llm
    return llm(system, prompt, max_tokens, timeout=60)


def _st(s, tenant: str) -> ShikayatSettings:
    r = s.get(ShikayatSettings, tenant)
    if r is None:
        r = ShikayatSettings(tenant_id=tenant)
        s.add(r)
        s.flush()
    return r


def _log(c: ShikayatCase, text: str, by: str = "owner") -> None:
    lg = _j(c.log, [])
    lg.append({"at": _now().isoformat(timespec="seconds"), "by": by, "text": str(text)[:400]})
    c.log = _dump(lg[-120:])
    c.updated_at = _now()


def _business(tenant: str, st: ShikayatSettings) -> str:
    if st.business:
        return st.business
    try:
        from ..employee.models import EmployeeProfile
        with session() as s:
            p = s.get(EmployeeProfile, tenant)
            if p and p.business_name:
                return p.business_name
    except Exception:
        pass
    return ""


# ---------------------------------------------------------------- settings
def settings(tenant: str) -> dict[str, Any]:
    with session() as s:
        r = _st(s, tenant)
        t = R.targets(r.sector, r.ack_hours or None, r.resolve_days or None)
        return {"business": r.business, "sector": r.sector, "sector_label": R.SECTORS[R.norm_sector(r.sector)]["label"],
                "ack_hours": r.ack_hours, "resolve_days": r.resolve_days, "effective": {"ack_hours": t["ack_hours"], "resolve_days": t["resolve_days"]},
                "reopen_days": r.reopen_days, "silence_days": r.silence_days, "officer": _j(r.officer, {}), "comp_limit": r.comp_limit,
                "month_limit": r.month_limit, "close_template": r.close_template, "rules": _j(r.rules, []), "team": _j(r.team, []),
                "morning_time": r.morning_time, "learned": _j(r.learned, {}), "law": R.SECTORS[R.norm_sector(r.sector)]["law"],
                "notes": t["notes"]}


def save_settings(tenant: str, **kw: Any) -> dict[str, Any]:
    with session() as s:
        r = _st(s, tenant)
        if kw.get("business") is not None:
            r.business = str(kw["business"]).strip()[:160]
        if kw.get("sector"):
            r.sector = R.norm_sector(kw["sector"])
        for k in ("ack_hours", "resolve_days", "comp_limit", "month_limit"):
            if kw.get(k) not in (None, ""):
                v = _num(kw[k])
                if v < 0:
                    raise ShikayatError(f"{k} 0 se kam nahi.")
                setattr(r, k, v)
        for k, lo, hi in (("reopen_days", 1, 365), ("silence_days", 3, 90)):
            if kw.get(k) not in (None, ""):
                setattr(r, k, max(lo, min(hi, int(_num(kw[k])))))
        if isinstance(kw.get("officer"), dict):
            o = kw["officer"]
            r.officer = _dump({"name": str(o.get("name") or "")[:80], "phone": _phone(o.get("phone")), "email": str(o.get("email") or "")[:120]})
        if kw.get("close_template") is not None:
            r.close_template = str(kw["close_template"])[:1500]
        if kw.get("rules") is not None:
            rl = kw["rules"] if isinstance(kw["rules"], list) else [x for x in str(kw["rules"]).split("\n")]
            r.rules = _dump([str(x).strip()[:300] for x in rl if str(x).strip()][:30])
        if kw.get("add_rule"):
            rl = _j(r.rules, [])
            rl.append(str(kw["add_rule"]).strip()[:300])
            r.rules = _dump(rl[-30:])
        if kw.get("team") is not None and isinstance(kw["team"], list):
            r.team = _dump([{"name": str(m.get("name") or "")[:60], "phone": _phone(m.get("phone"))}
                            for m in kw["team"] if isinstance(m, dict) and m.get("name")][:30])
        if kw.get("morning_time"):
            mt = str(kw["morning_time"])
            if not re.fullmatch(r"\d{1,2}:\d{2}", mt):
                raise ShikayatError("Samay HH:MM mein.")
            r.morning_time = mt.zfill(5)
    out = settings(tenant)
    out["say"] = f"Niyam: {out['sector_label']} - pavti {out['effective']['ack_hours']:g} ghante, hal {out['effective']['resolve_days']:g} din." + \
        (" " + " ".join(out["notes"]) if out["notes"] else "")
    return out


# ---------------------------------------------------------------- views
def _get(s, tenant: str, ref: str) -> ShikayatCase:
    ref = str(ref or "").strip()
    if not ref:
        raise ShikayatError("Kaun si shikayat? Number (SH-0001) bataiye.")
    c = s.get(ShikayatCase, ref)
    if c is not None and c.tenant_id == tenant:
        return c
    up = ref.upper()
    c = s.scalars(select(ShikayatCase).where(ShikayatCase.tenant_id == tenant, ShikayatCase.number == up)).first()
    if c is None and re.fullmatch(r"\d{1,5}", ref):
        c = s.scalars(select(ShikayatCase).where(ShikayatCase.tenant_id == tenant, ShikayatCase.number == f"SH-{int(ref):04d}")).first()
    if c is None and len(_phone(ref)) == 10:
        c = s.scalars(select(ShikayatCase).where(ShikayatCase.tenant_id == tenant, ShikayatCase.phone == _phone(ref))
                      .order_by(ShikayatCase.created_at.desc())).first()
    if c is None:
        rows = s.scalars(select(ShikayatCase).where(ShikayatCase.tenant_id == tenant, ShikayatCase.stage.in_(R.OPEN))).all()
        hits = [x for x in rows if ref.lower() in (x.customer or "").lower()]
        if len(hits) == 1:
            c = hits[0]
        elif len(hits) > 1:
            raise ShikayatError("Kai shikayatein: " + ", ".join(f"{h.number} ({h.customer})" for h in hits[:6]))
    if c is None:
        raise ShikayatError(f"Shikayat '{ref}' nahi mili.")
    return c


def _proof_view(p: ShikayatProof) -> dict[str, Any]:
    return {"id": p.id, "kind": p.kind, "stage": p.stage, "name": p.name, "mime": p.mime, "size": p.size, "cleaned": p.cleaned,
            "note": p.note, "url": f"/api/shikayat/proof/{p.id}", "at": _utc(p.created_at).isoformat() if p.created_at else ""}


def _comp_view(tenant: str, x: ShikayatComp) -> dict[str, Any]:
    return {"id": x.id, "number": x.number, "case_id": x.case_id, "kind": x.kind, "kind_label": R.COMP_KINDS.get(x.kind, x.kind),
            "amount": x.amount, "amount_text": _rs(x.amount) if x.amount else "", "item": x.item, "reason": x.reason,
            "checks": _j(x.checks, []), "status": x.status, "hash": x.hash, "given_how": x.given_how,
            "decided_at": _local(tenant, x.decided_at), "given_at": _local(tenant, x.given_at)}


def _state(tenant: str, c: ShikayatCase) -> dict[str, Any]:
    """Where the clock stands: ack late?, due in / overdue by."""
    n = _now()
    out: dict[str, Any] = {}
    if c.stage in R.OPEN:
        ab = _utc(c.ack_by)
        if not c.ack_at and ab:
            out["ack_late"] = n > ab
            out["ack_left_h"] = round((ab - n).total_seconds() / 3600, 1)
        du = _utc(c.due_at)
        if du:
            left = (du - n).total_seconds() / 86400
            out["due_left_days"] = round(left, 1)
            out["overdue"] = left < 0
    return out


def _view(s, tenant: str, c: ShikayatCase, full: bool = True) -> dict[str, Any]:
    v: dict[str, Any] = {
        "id": c.id, "number": c.number, "customer": c.customer, "phone": c.phone, "channel": c.channel, "product": c.product,
        "order_ref": c.order_ref, "amount": c.amount, "text": c.text, "causes": _j(c.causes, []),
        "cause_labels": [R.CAUSE_LABEL.get(x, x) for x in _j(c.causes, [])], "flags": _j(c.flags, []), "severity": c.severity,
        "sector": c.sector, "stage": c.stage, "stage_label": R.STAGE_LABEL.get(c.stage, c.stage), "owner": c.owner,
        "received": _local(tenant, c.received_at), "ack_by": _local(tenant, c.ack_by), "ack_at": _local(tenant, c.ack_at),
        "due": _local(tenant, c.due_at), "resolved_at": _local(tenant, c.resolved_at), "closed_at": _local(tenant, c.closed_at),
        "told_at": _local(tenant, c.told_at), "resolution": c.resolution, "fix_proof": c.fix_proof, "confirmed": c.confirmed,
        "satisfaction": c.satisfaction, "feedback": c.feedback, "reopen_count": c.reopen_count, "parent": c.parent,
        "escalated": _j(c.escalated, {}), "external_ref": c.external_ref, "ticket_ref": c.ticket_ref, "contact_id": c.contact_id,
        **_state(tenant, c)}
    if full:
        proofs = s.scalars(select(ShikayatProof).where(ShikayatProof.case_id == c.id).order_by(ShikayatProof.created_at)).all()
        comps = s.scalars(select(ShikayatComp).where(ShikayatComp.case_id == c.id).order_by(ShikayatComp.created_at)).all()
        v["proofs"] = [_proof_view(p) for p in proofs]
        v["comps"] = [_comp_view(tenant, x) for x in comps]
        v["log"] = _j(c.log, [])[-40:]
        v["checklist"] = _checklist(s, tenant, c)
        v["routes"] = R.routes_for(c.sector)
        if c.amount:
            v["commission"] = R.commission_for(c.amount)
    return v


def case(tenant: str, ref: str) -> dict[str, Any]:
    with session() as s:
        return _view(s, tenant, _get(s, tenant, ref))


def cases(tenant: str, stage: str = "", q: str = "", limit: int = 100) -> list[dict[str, Any]]:
    stage = R.norm_stage(stage)
    with session() as s:
        qq = select(ShikayatCase).where(ShikayatCase.tenant_id == tenant)
        if stage == "open":
            qq = qq.where(ShikayatCase.stage.in_(R.OPEN))
        elif stage:
            qq = qq.where(ShikayatCase.stage == stage)
        rows = s.scalars(qq.order_by(ShikayatCase.created_at.desc()).limit(500)).all()
        if q:
            ql = q.lower()
            rows = [r for r in rows if ql in (r.customer or "").lower() or ql in (r.product or "").lower() or ql in r.number.lower()
                    or (len(_phone(q)) >= 6 and _phone(q) in r.phone) or ql in (r.text or "").lower()]
        out = [_view(s, tenant, r, False) for r in rows[:limit]]
    rank = {"urgent": 0, "high": 1, "normal": 2, "low": 3}
    return sorted(out, key=lambda v: (v["stage"] not in R.OPEN, rank.get(v["severity"], 2), v.get("due_left_days", 9999)))


# ---------------------------------------------------------------- logging a complaint
def log_complaint(tenant: str, text: str, customer: str = "", phone: str = "", channel: str = "chat", product: str = "",
                  order_ref: str = "", amount: Any = None, sector: str = "", severity: str = "", owner: str = "",
                  file_ids: list[str] | None = None, contact_id: str = "", ticket_ref: str = "", received: Any = None,
                  by: str = "owner", conv: str = "") -> dict[str, Any]:
    text, scrubbed = R.scrub(str(text or "").strip()[:4000])
    if len(text) < 3:
        raise ShikayatError("Shikayat kya hai, woh likhiye / boliye.")
    ph = _phone(phone)
    channel = channel if channel in R.CHANNELS else "other"
    rec = _utc(received) if isinstance(received, datetime) else _now()
    if isinstance(received, str) and received.strip():
        from ..theka.service import _date
        d = _date(received)
        if d and d <= today(tenant):
            rec = datetime(d.year, d.month, d.day, 10, 0, tzinfo=_tz(tenant)).astimezone(timezone.utc)
    with _LOCK, session() as s:
        st = _st(s, tenant)
        sec = R.norm_sector(sector or st.sector)
        # the same complaint twice (a retry, the same forward) = the first one
        if ph:
            recent = s.scalars(select(ShikayatCase).where(ShikayatCase.tenant_id == tenant, ShikayatCase.phone == ph)
                               .order_by(ShikayatCase.created_at.desc()).limit(20)).all()
        else:
            recent = []
        for r0 in recent:
            if r0.text == text and _utc(r0.created_at) and _now() - _utc(r0.created_at) < timedelta(minutes=30):
                return {"case": _view(s, tenant, r0), "duplicate": True, "say": f"Yeh shikayat pehle hi {r0.number} mein darj hai."}
        sev0, flags = R.guess_severity(text)
        sev = severity if severity in R.SEVERITIES else sev0
        causes = R.guess_causes(text + " " + product)
        pkey = _key(product)
        n180 = _now() - timedelta(days=180)
        earlier = [r0 for r0 in recent if _utc(r0.created_at) and _utc(r0.created_at) >= n180]
        if earlier:
            flags.append("repeat_customer")
        reopen_hint = ""
        for r0 in earlier:
            ca = _utc(r0.closed_at)
            if r0.stage == "closed" and ca and _now() - ca <= timedelta(days=st.reopen_days) and \
                    (set(_j(r0.causes, [])) & set(causes) or (pkey and r0.product_key == pkey)):
                reopen_hint = r0.number
                break
        same_product = 0
        if pkey:
            same_product = len(s.scalars(select(ShikayatCase.id).where(ShikayatCase.tenant_id == tenant, ShikayatCase.product_key == pkey,
                                                                        ShikayatCase.created_at >= _now() - timedelta(days=30))).all())
            if same_product >= 2:
                flags.append("product_pattern")
        dl = R.deadlines(rec, sec, st.ack_hours or None, st.resolve_days or None, sev)
        st.seq += 1
        c = ShikayatCase(tenant_id=tenant, number=f"SH-{st.seq:04d}", customer=str(customer or "").strip()[:120], phone=ph,
                         contact_id=str(contact_id or "")[:40], channel=channel, ticket_ref=str(ticket_ref or "")[:40],
                         product=str(product or "").strip()[:160], product_key=pkey, order_ref=str(order_ref or "").strip()[:80],
                         amount=_num(amount), text=text, causes=_dump(causes), flags=_dump(sorted(set(flags))), severity=sev,
                         sector=sec, stage="new", owner=str(owner or "").strip()[:80], received_at=rec, ack_by=dl["ack_by"],
                         due_at=dl["due_at"])
        _log(c, f"Darj ({channel})" + (" - card / OTP jaisi jaankari hatayi" if scrubbed else ""), by)
        s.add(c)
        s.flush()
        cid, num = c.id, c.number
    attached, failed = [], []
    for fid in file_ids or []:
        try:
            attached.append(attach_file(tenant, cid, fid, "complaint"))
        except ShikayatError as e:
            failed.append(f"{fid}: {e}")
    with session() as s:
        v = _view(s, tenant, s.get(ShikayatCase, cid))
    say = (f"Shikayat {num} darj" + (f" ({v['customer']})" if v["customer"] else "") + f": {', '.join(v['cause_labels'])}. "
           f"Pavti {v['ack_by']} tak, hal {v['due']} tak ({R.SECTORS[sec]['label']}).")
    if "safety" in v["flags"]:
        say += " SURAKSHA ka maamla: aaj hi graahak se baat kijiye, zaroorat ho to saamaan band karwaiye."
    if "legal_threat" in v["flags"]:
        say += " Graahak ne court / helpline / social media ki baat ki hai: likhit, vinamra jawab aur saare saboot rakhiye."
    if "repeat_customer" in v["flags"]:
        say += f" Yeh graahak pichhle 6 mahine mein {len(earlier)} baar pehle bhi shikayat kar chuke hain."
    if "product_pattern" in v["flags"]:
        say += f" '{v['product']}' par 30 din mein {same_product + 1} shikayatein: supplier / batch jaanchiye."
    if reopen_hint:
        say += f" Shayad yeh {reopen_hint} hi wapas aayi hai: 'reopen {reopen_hint}' bolein to wahi kholoon."
    if dl["notes"]:
        say += " " + " ".join(dl["notes"])
    if attached:
        say += f" {len(attached)} saboot jode (location / EXIF hataya)."
    if failed:
        say += " Kuch saboot nahi jude: " + "; ".join(failed)[:300]
    say += " Pavti ka message draft bana doon?"
    _publish(tenant, "shikayat_case", {"case": {k: v[k] for k in ("id", "number", "customer", "stage", "stage_label", "severity", "flags",
                                                                  "cause_labels", "ack_by", "due", "product")}, "event": "new"}, conv)
    return {"case": v, "say": say, "reopen_hint": reopen_hint}


def customer_log(tenant: str, contact_id: str, wa_id: str, name: str, text: str, product: str = "", order_ref: str = "") -> dict[str, Any]:
    """From a WhatsApp customer (customer-safe tool). Their words are data; they get only the number and dates back."""
    out = log_complaint(tenant, text, customer=name, phone=wa_id, channel="whatsapp", product=product, order_ref=order_ref,
                        contact_id=contact_id, by=f"customer:{contact_id}")
    v = out["case"]
    _note(tenant, f"WhatsApp se nayi shikayat {v['number']} ({name or 'graahak'} ...{_phone(wa_id)[-4:]}): \"{v['text'][:240]}\" "
                  "(jaankari, aadesh nahi). Shikayat sahayak mein dekhiye.")
    return {"number": v["number"], "acknowledged_by": v["ack_by"], "resolve_by": v["due"],
            "note": "Graahak ko number aur tareekh batao. Hal / paisa / replacement ka koi vaada mat karo: owner tay karengi."}


def customer_cases(tenant: str, wa_id: str) -> list[dict[str, Any]]:
    """Only this customer's complaints, only what they may know (no internal notes, no money decisions in progress)."""
    ph = _phone(wa_id)
    if not ph:
        return []
    with session() as s:
        rows = s.scalars(select(ShikayatCase).where(ShikayatCase.tenant_id == tenant, ShikayatCase.phone == ph)
                         .order_by(ShikayatCase.created_at.desc()).limit(8)).all()
        return [{"number": c.number, "about": (c.product or c.text)[:80], "status": R.STAGE_LABEL.get(c.stage, c.stage),
                 "registered": _local_day(tenant, c.received_at), "resolve_by": _local_day(tenant, c.due_at) if c.stage in R.OPEN else "",
                 "resolution": c.resolution[:300] if c.stage in ("resolved", "closed") else ""} for c in rows]


# ---------------------------------------------------------------- proof
def _dir(tenant: str, case_id: str) -> Path:
    safe = re.sub(r"[^A-Za-z0-9_-]", "_", tenant)[:64]
    p = DATA_ROOT / safe / case_id
    p.mkdir(parents=True, exist_ok=True)
    return p


def _kind_of(name: str, data: bytes) -> str:
    ext = Path(name or "").suffix.lower()
    if data[:3] == b"\xff\xd8\xff" or data[:8] == b"\x89PNG\r\n\x1a\n" or data[8:12] == b"WEBP" or data[4:12] in (b"ftypheic", b"ftypmif1", b"ftypheix"):
        return "screenshot" if (ext == ".png" or "screenshot" in (name or "").lower()) else "photo"
    if data[4:8] == b"ftyp" or ext in (".mp4", ".mov", ".3gp", ".mkv", ".webm") or data[:4] == b"\x1aE\xdf\xa3":
        return "video"
    if ext in (".mp3", ".m4a", ".ogg", ".opus", ".wav", ".aac", ".amr") or data[:4] in (b"OggS", b"RIFF", b"ID3\x03", b"ID3\x04"):
        return "audio"
    return "doc"


def _clean_image(data: bytes, keep_png: bool) -> tuple[bytes, str, str]:
    """Re-encode: EXIF (camera, time, GPS location) and any hidden chunk are dropped; orientation applied first."""
    from PIL import Image, ImageOps
    try:
        img = Image.open(io.BytesIO(data))
        img.load()
    except Exception:
        raise ShikayatError("Yeh photo khul nahi rahi.")
    try:
        img = ImageOps.exif_transpose(img)
    except Exception:
        pass
    img.thumbnail((2400, 2400))
    out = io.BytesIO()
    if keep_png:
        mode = "RGBA" if img.mode in ("RGBA", "LA", "P") else "RGB"
        src = img.convert(mode)
        clean = Image.frombytes(mode, src.size, src.tobytes())       # pixels only: no text / exif / time chunks
        clean.save(out, "PNG", optimize=True)
        return out.getvalue(), ".png", "image/png"
    src = img.convert("RGB")
    Image.frombytes("RGB", src.size, src.tobytes()).save(out, "JPEG", quality=86)   # no exif= -> none written
    return out.getvalue(), ".jpg", "image/jpeg"


def _clean_video(src: Path, dst: Path) -> bool:
    """Container metadata (GPS / location / creation tags) removed without re-encoding; False when ffmpeg is missing / fails."""
    ff = shutil.which("ffmpeg")
    if not ff:
        return False
    try:
        r = subprocess.run([ff, "-y", "-v", "error", "-i", str(src), "-map", "0", "-map_metadata", "-1", "-map_chapters", "-1",
                            "-c", "copy", "-movflags", "+faststart", str(dst)], capture_output=True, timeout=180)
        return r.returncode == 0 and dst.exists() and dst.stat().st_size > 0
    except Exception:
        return False


def _clean_doc(name: str, data: bytes) -> tuple[bytes, bool]:
    ext = Path(name or "").suffix.lower()
    try:
        if ext == ".pdf" or data[:5] == b"%PDF-":
            import pikepdf
            with pikepdf.open(io.BytesIO(data)) as pdf:
                try:
                    del pdf.Root.Metadata
                except Exception:
                    pass
                try:
                    del pdf.trailer.Info
                except Exception:
                    pass
                out = io.BytesIO()
                pdf.save(out)
                return out.getvalue(), True
        if ext == ".docx":
            import docx
            d = docx.Document(io.BytesIO(data))
            cp = d.core_properties
            for k in ("author", "last_modified_by", "comments", "keywords", "subject", "title", "category"):
                try:
                    setattr(cp, k, "")
                except Exception:
                    pass
            out = io.BytesIO()
            d.save(out)
            return out.getvalue(), True
    except Exception:
        log.info("doc metadata clean failed for %s", name)
    return data, False


def add_proof(tenant: str, case_ref: str, data: bytes, name: str = "proof", stage: str = "complaint", note: str = "",
              by: str = "owner") -> dict[str, Any]:
    from .. import filecheck
    if not data:
        raise ShikayatError("File khaali hai.")
    stage = stage if stage in ("complaint", "fix", "other") else "complaint"
    name = (str(name or "proof").replace("\\", "/").rsplit("/", 1)[-1] or "proof")[:160]
    kind = _kind_of(name, data)
    cap = MAX_VIDEO_MB if kind == "video" else MAX_PROOF_MB
    if len(data) > cap * 1024 * 1024:
        raise ShikayatError(f"File {cap} MB se badi hai.")
    if data[:2] == b"MZ" or data[:4] == b"\x7fELF" or data[:2] == b"#!" or \
            Path(name).suffix.lower() in (".exe", ".apk", ".bat", ".cmd", ".sh", ".js", ".vbs", ".ps1", ".msi", ".dll", ".jar", ".scr"):
        raise ShikayatError("Yeh program / script hai, saboot nahi; nahi rakhi (kisi ne bheji ho to mat kholiye).")
    rt = filecheck.real_type(name, data)
    if rt.get("danger"):
        raise ShikayatError(f"Yeh file khatarnak lag rahi hai ({rt.get('note')}); nahi rakhi.")
    with session() as s:
        c = _get(s, tenant, case_ref)
        cid, num = c.id, c.number
        n = len(s.scalars(select(ShikayatProof.id).where(ShikayatProof.case_id == cid)).all())
    if n >= MAX_PROOFS:
        raise ShikayatError(f"Ek shikayat mein {MAX_PROOFS} saboot tak.")
    folder = _dir(tenant, cid)
    stem = f"{stage}_{n + 1:02d}_{secrets.token_hex(3)}"
    cleaned, mime = False, ""
    if kind in ("photo", "screenshot"):
        body, ext, mime = _clean_image(data, keep_png=kind == "screenshot")
        path = folder / (stem + ext)
        path.write_bytes(body)
        cleaned = True
    elif kind == "video":
        ext = Path(name).suffix.lower() if Path(name).suffix.lower() in (".mp4", ".mov", ".3gp", ".mkv", ".webm") else ".mp4"
        raw = folder / (stem + "_raw" + ext)
        raw.write_bytes(data)
        path = folder / (stem + ext)
        cleaned = _clean_video(raw, path)
        if cleaned:
            raw.unlink(missing_ok=True)
        else:
            raw.rename(path)
        mime = {".mov": "video/quicktime", ".webm": "video/webm", ".3gp": "video/3gpp", ".mkv": "video/x-matroska"}.get(ext, "video/mp4")
    else:
        body, cleaned = _clean_doc(name, data) if kind == "doc" else (data, False)
        ext = Path(name).suffix.lower()[:8] or ".bin"
        path = folder / (stem + ext)
        path.write_bytes(body)
        import mimetypes
        mime = mimetypes.guess_type(name)[0] or "application/octet-stream"
    body = path.read_bytes()
    with session() as s:
        p = ShikayatProof(tenant_id=tenant, case_id=cid, kind=kind, stage=stage, name=name, path=str(path), mime=mime, size=len(body),
                          sha256=hashlib.sha256(body).hexdigest(), cleaned=cleaned, note=str(note or "")[:300])
        s.add(p)
        c = s.get(ShikayatCase, cid)
        _log(c, f"Saboot ({'theek hone ka' if stage == 'fix' else 'shikayat ka'} {kind}) joda" + ("" if cleaned or kind == "audio" else " - metadata nahi hata"), by)
        s.flush()
        v = _proof_view(p)
    v["say"] = f"{num}: {kind} saboot joda" + (" (location / camera jaankari hatayi)." if cleaned else
                                              (". Video ki metadata nahi hat payi (ffmpeg nahi)." if kind == "video" else "."))
    return v


def attach_file(tenant: str, case_ref: str, file_id: str, stage: str = "complaint", note: str = "", by: str = "owner") -> dict[str, Any]:
    """A file she sent in chat / on a call (file_id) is copied into the complaint's own store (chat files vanish after 24 h)."""
    from ..mailbox import service as ms
    try:
        name, _mime, data = ms.load_files(tenant, [{"type": "file", "file_id": str(file_id)}])[0]
    except Exception as e:
        raise ShikayatError(f"File {file_id} nahi mili: {e}")
    return add_proof(tenant, case_ref, data, name, stage, note, by)


def proof_file(tenant: str, proof_id: str) -> tuple[Path, str]:
    with session() as s:
        p = s.get(ShikayatProof, proof_id)
        if p is None or p.tenant_id != tenant or not Path(p.path).exists():
            raise ShikayatError("Saboot nahi mila.")
        return Path(p.path), p.mime or "application/octet-stream"


# ---------------------------------------------------------------- moving the case
def ack(tenant: str, case_ref: str, how: str = "", by: str = "owner") -> dict[str, Any]:
    with session() as s:
        c = _get(s, tenant, case_ref)
        if c.ack_at:
            return {"case": _view(s, tenant, c, False), "say": f"{c.number} ki pavti pehle hi {_local(tenant, c.ack_at)} ko di gayi."}
        c.ack_at = _now()
        if c.stage == "new":
            c.stage = "acknowledged"
        late = _utc(c.ack_by) and c.ack_at > _utc(c.ack_by)
        _log(c, "Pavti di" + (f" ({how})" if how else "") + (" - samay ke baad" if late else ""), by)
        s.flush()
        v = _view(s, tenant, c, False)
    return {"case": v, "say": f"{v['number']}: pavti darj." + (" Samay-seema nikal chuki thi; register mein dikhega." if late else "")}


def update(tenant: str, case_ref: str, stage: str = "", owner: str | None = None, note: str = "", severity: str = "",
           causes: list[str] | None = None, due: str = "", external_ref: str = "", product: str = "", by: str = "owner") -> dict[str, Any]:
    stage = R.norm_stage(stage)
    if stage and stage not in R.STAGES:
        raise ShikayatError("Stage: " + " | ".join(R.STAGES))
    if stage in ("resolved", "closed"):
        raise ShikayatError("Hal / band ke liye 'resolve' / 'close' kijiye (checklist ke saath).")
    if stage == "reopened":
        return reopen(tenant, case_ref, note, by)
    note, _ = R.scrub(note)
    with session() as s:
        st = _st(s, tenant)
        c = _get(s, tenant, case_ref)
        if c.stage in ("closed", "rejected") and stage:
            raise ShikayatError(f"{c.number} band hai; 'reopen' kijiye.")
        bits = []
        if stage == "rejected":
            if len(note.strip()) < 8:
                raise ShikayatError("Shikayat na maanne ki saaf wajah likhiye (graahak ko bhi batani hogi).")
            c.stage, c.resolution, c.closed_at = "rejected", note.strip()[:2000], _now()
            bits.append("Nahi maani: " + note[:120])
        elif stage:
            if stage == "acknowledged" and not c.ack_at:
                c.ack_at = _now()
            c.stage = stage
            bits.append("Stage: " + R.STAGE_LABEL[stage])
        if owner is not None:
            team = [m["name"].lower() for m in _j(st.team, [])]
            o = str(owner).strip()
            if o and team and o.lower() not in team and o.lower() not in ("main", "owner", "aap", "me"):
                raise ShikayatError("Team mein yeh naam nahi: " + ", ".join(m["name"] for m in _j(st.team, [])))
            c.owner = o[:80]
            bits.append(f"Zimmedaar: {o or 'aap'}")
        if severity in R.SEVERITIES:
            c.severity = severity
            bits.append(f"Gambhirta: {severity}")
        if causes:
            cs = [x for x in causes if x in R.CAUSES or x == "other"]
            if cs:
                c.causes = _dump(cs[:3])
                bits.append("Wajah: " + ", ".join(R.CAUSE_LABEL.get(x, x) for x in cs))
        if due:
            from ..theka.service import _date
            d = _date(due)
            if not d:
                raise ShikayatError("Tareekh samajh nahi aayi.")
            new_due = datetime(d.year, d.month, d.day, 18, 0, tzinfo=_tz(tenant)).astimezone(timezone.utc)
            lim = R.deadlines(_utc(c.received_at), c.sector, None, None, "normal")["due_at"]
            if R.SECTORS[c.sector]["legal"] and new_due > lim:
                bits.append(f"Nayi tareekh niyam ({_local_day(tenant, lim)}) ke baad hai - graahak ko likhit mein wajah dijiye")
            c.due_at = new_due
            bits.append(f"Nayi tareekh {d.isoformat()}")
        if external_ref:
            c.external_ref = str(external_ref)[:80]
            bits.append(f"Docket: {external_ref}")
        if product:
            c.product, c.product_key = str(product)[:160], _key(product)
        if note and stage != "rejected":
            bits.append(note[:300])
        if not bits:
            raise ShikayatError("Kya badalna hai?")
        _log(c, "; ".join(bits), by)
        s.flush()
        v = _view(s, tenant, c, False)
    say = f"{v['number']}: " + "; ".join(bits) + "."
    if v["stage"] == "rejected":
        say += " Graahak ko wajah aur aage ka rasta batane wala message draft karoon?"
    return {"case": v, "say": say}


def note(tenant: str, case_ref: str, text: str, by: str = "owner") -> dict[str, Any]:
    t, _ = R.scrub(text)
    if not t.strip():
        raise ShikayatError("Note khaali hai.")
    with session() as s:
        c = _get(s, tenant, case_ref)
        _log(c, t, by)
        return {"say": f"{c.number}: note likha."}


def resolve(tenant: str, case_ref: str, resolution: str, fix_proof: str = "", by: str = "owner", conv: str = "") -> dict[str, Any]:
    resolution, _ = R.scrub(resolution)
    if len(resolution.strip()) < 5:
        raise ShikayatError("Kya hal kiya, woh likhiye (graahak ko yahi bataya jayega).")
    with session() as s:
        c = _get(s, tenant, case_ref)
        if c.stage in ("closed", "rejected"):
            raise ShikayatError(f"{c.number} pehle hi band hai.")
        c.stage, c.resolution, c.resolved_at = "resolved", resolution.strip()[:2000], _now()
        if fix_proof:
            c.fix_proof = R.scrub(fix_proof)[0][:500]
        if not c.ack_at:
            c.ack_at = c.resolved_at
        _log(c, "Hal: " + resolution[:200], by)
        s.flush()
        v = _view(s, tenant, c)
    missing = [x["say"] for x in v["checklist"] if not x["ok"]]
    say = f"{v['number']} hal kiya." + (" Band karne se pehle: " + "; ".join(missing) + "." if missing else "") + \
        " Graahak ke liye vinamra closing message draft karoon?"
    _publish(tenant, "shikayat_case", {"case": {k: v[k] for k in ("id", "number", "customer", "stage", "stage_label", "severity")},
                                       "event": "resolved"}, conv)
    return {"case": v, "say": say}


def _checklist(s, tenant: str, c: ShikayatCase) -> list[dict[str, Any]]:
    st = _st(s, tenant)
    proofs = s.scalars(select(ShikayatProof).where(ShikayatProof.case_id == c.id, ShikayatProof.stage == "fix")).all()
    pend = s.scalars(select(ShikayatComp).where(ShikayatComp.case_id == c.id, ShikayatComp.status == "pending")).all()
    appr = s.scalars(select(ShikayatComp).where(ShikayatComp.case_id == c.id, ShikayatComp.status == "approved")).all()
    told = _utc(c.told_at)
    silent = bool(told and _now() - told >= timedelta(days=st.silence_days) and c.confirmed is None)
    out = [
        {"item": "resolution", "ok": bool(c.resolution), "say": "kya hal kiya woh likhiye"},
        {"item": "proof_of_fix", "ok": bool(proofs or c.fix_proof), "say": "theek hone ka saboot (photo / bill / graahak ka sandesh)"},
        {"item": "customer_told", "ok": bool(told), "say": "graahak ko closing message bhejiye (aapke tap se)"},
        {"item": "customer_confirmed", "ok": c.confirmed is True or silent,
         "say": f"graahak se 'theek ho gaya' ki pushti (ya message ke {st.silence_days} din baad tak jawab na aaye)"},
        {"item": "compensation_decided", "ok": not pend, "say": "muavze ka card abhi aapki haan / na ka intezaar kar raha hai"},
        {"item": "compensation_given", "ok": not appr, "say": "jo muavza aapne maana, woh diya? ('de diya' dabaiye)"},
    ]
    if c.confirmed is False:
        out[3]["ok"] = False
        out[3]["say"] = "graahak ne kaha theek nahi hua"
    return out


# ---------------------------------------------------------------- messages to the customer (drafts; her tap sends)
def _approved_comps(s, case_id: str) -> list[ShikayatComp]:
    return s.scalars(select(ShikayatComp).where(ShikayatComp.case_id == case_id, ShikayatComp.status.in_(("approved", "given")))).all()


def _template(kind: str, biz: str, c: ShikayatCase, tenant: str, comps: list[ShikayatComp], officer: dict) -> str:
    name = (c.customer.split()[0] + " ji") if c.customer else "Namaste"
    head = f"{name}, namaste." if c.customer else "Namaste."
    sign = f"\n- {biz}" if biz else ""
    contact = ""
    if officer.get("name") or officer.get("phone"):
        contact = f"\nShikayat adhikari: {officer.get('name') or ''} {officer.get('phone') or ''} {officer.get('email') or ''}".rstrip()
    about = f" ({c.product})" if c.product else ""
    if kind == "ack":
        return (f"{head} Aapki shikayat{about} humein mil gayi hai. Shikayat number: {c.number}. Hum ise dekh rahe hain aur "
                f"{_local_day(tenant, c.due_at)} tak hal karke aapko batayenge. Asuvidha ke liye khed hai.{contact}{sign}")
    if kind == "update":
        return (f"{head} Shikayat {c.number}{about} par kaam chal raha hai: {R.STAGE_LABEL.get(c.stage, c.stage).lower()}. "
                f"Hum {_local_day(tenant, c.due_at)} tak aapko poori jaankari denge. Dhairya ke liye dhanyavaad.{sign}")
    if kind == "satisfaction":
        return (f"{head} Shikayat {c.number} ke hal se aap kitne santusht hain? 1 se 5 mein jawab de dijiye (5 = bahut achha). "
                f"Aapki raay se hum aur behtar karenge.{sign}")
    if kind == "reject":
        return (f"{head} Shikayat {c.number}{about} hamne dhyan se dekhi. {c.resolution} Isliye is baar hum ise maan nahi pa rahe, "
                f"iske liye khed hai. {R.route_line(c.sector) or R.route_line('general')}{contact}{sign}")
    # close
    gave = ""
    if comps:
        gave = " " + " ".join(
            f"{R.COMP_KINDS.get(x.kind, x.kind)}" + (f" {_rs(x.amount)}" if x.amount and x.kind != "replacement" else "")
            + (f" ({x.item})" if x.item else "") + (" diya gaya hai." if x.status == "given" else " diya jayega.") for x in comps)
    route = (" " + R.route_line(c.sector)) if R.SECTORS.get(c.sector, {}).get("legal") else ""
    return (f"{head} Shikayat {c.number}{about} ke baare mein: {c.resolution.rstrip('.')}.{gave} Aapko jo pareshani hui, uske liye "
            f"hum maafi chahte hain. Agar ab bhi koi dikkat ho to isi number par bataiye, hum shikayat dobara kholenge.{route}{contact}{sign}")


def _polish(base: str, kind: str, number: str, allowed_money: list[str], rules: list[str]) -> str:
    """Optional: the model makes the template warmer. Refused (template kept) if it adds money / promises / drops facts."""
    try:
        out = _ai("You rewrite a short customer-service message for an Indian small business. Keep the SAME language mix "
                  "(Hinglish in Latin letters), polite (aap), warm, 3-6 short sentences. Keep EVERY fact exactly: complaint "
                  "number, dates, amounts, names, phone numbers, the route line. Do NOT add any offer, refund, discount, "
                  "free item, compensation, promise or date that is not already in it. No emojis. Reply with only the message. "
                  f"Her rules: {'; '.join(rules)[:600]}", base, 600).strip()
    except Exception:
        return base
    if not out or len(out) > 1200:
        return base
    if number not in out:
        return base
    for amt in re.findall(r"₹[\d,]+", out):
        if amt not in allowed_money:
            return base
    if R.MONEY_RE.search(out) and not R.MONEY_RE.search(base):
        return base
    if R.PROMISE_RE.search(out) and not R.PROMISE_RE.search(base):
        return base
    for d in re.findall(r"\d{1,2} [A-Z][a-z]{2} \d{4}", base):
        if d not in out:
            return base
    return out


def message(tenant: str, case_ref: str, kind: str = "close", polish: bool = True, conv: str = "") -> dict[str, Any]:
    kind = kind if kind in ("ack", "update", "close", "satisfaction", "reject") else "close"
    from ..employee import inbox
    from ..employee.models import BizContact
    with session() as s:
        st = _st(s, tenant)
        c = _get(s, tenant, case_ref)
        if kind == "close" and not c.resolution:
            raise ShikayatError(f"{c.number}: pehle 'resolve' kijiye (kya hal kiya), phir closing message.")
        if kind == "reject" and c.stage != "rejected":
            raise ShikayatError("Pehle shikayat 'rejected' kijiye (wajah ke saath).")
        pend = s.scalars(select(ShikayatComp).where(ShikayatComp.case_id == c.id, ShikayatComp.status == "pending")).all()
        comps = _approved_comps(s, c.id) if kind == "close" else []
        biz = _business(tenant, st)
        base = (st.close_template.replace("{number}", c.number).replace("{name}", c.customer or "").replace("{resolution}", c.resolution)
                if kind == "close" and st.close_template else _template(kind, biz, c, tenant, comps, _j(st.officer, {})))
        allowed = [_rs(x.amount) for x in comps if x.amount]
        rules = _j(st.rules, [])
        num, ph, cid_contact, cid = c.number, c.phone, c.contact_id, c.id
        if not cid_contact and ph:
            bc = [x for x in s.scalars(select(BizContact).where(BizContact.tenant_id == tenant)).all() if _phone(x.wa_id) == ph]
            cid_contact = bc[0].id if bc else ""
        window = False
        if cid_contact:
            bc0 = s.get(BizContact, cid_contact)
            window = bool(bc0 and bc0.tenant_id == tenant and not bc0.opted_out and inbox._window_open(bc0))
    text = _polish(base, kind, num, allowed, rules) if polish else base
    text = R.scrub(text)[0]
    out = {"case_id": cid, "number": num, "kind": kind, "text": text,
           "wa_link": f"https://wa.me/91{ph}?text={quote(text)}" if len(ph) == 10 else "",
           "sms_link": f"sms:+91{ph}?body={quote(text)}" if len(ph) == 10 else "",
           "business_ok": window, "contact_id": cid_contact,
           "say": f"{num}: {'closing' if kind == 'close' else kind} message ka draft taiyaar. Aap bhejiye"
                  + (" (ya business WhatsApp se, aapke tap par)." if window else " (apne WhatsApp / SMS se).")}
    if pend:
        out["say"] += " Dhyaan: muavze ka card abhi baaki hai; message mein uska vaada nahi likha."
    _publish(tenant, "shikayat_message", {k: out[k] for k in ("case_id", "number", "kind", "text", "wa_link", "sms_link", "business_ok")}, conv)
    return out


async def send_business(tenant: str, case_ref: str, text: str, kind: str = "close") -> dict[str, Any]:
    """Her tap on 'Business WhatsApp se bhejo' (API only): the official Cloud API, inside the 24 h window only."""
    from ..employee import inbox
    text = R.scrub(str(text or "").strip())[0]
    if not text:
        raise ShikayatError("Message khaali hai.")
    with session() as s:
        c = _get(s, tenant, case_ref)
        if R.MONEY_RE.search(text):
            allowed = [_rs(x.amount) for x in _approved_comps(s, c.id) if x.amount]
            amounts = re.findall(r"₹[\d,]+", text)
            if any(a not in allowed for a in amounts):
                raise ShikayatError("Message mein aisa paisa likha hai jo aapne muavze mein maana nahi. Pehle card par haan kijiye.")
        contact = c.contact_id
        if not contact and c.phone:
            from ..employee.models import BizContact
            bc = [x for x in s.scalars(select(BizContact).where(BizContact.tenant_id == tenant)).all() if _phone(x.wa_id) == c.phone]
            contact = bc[0].id if bc else ""
        cid = c.id
    if not contact:
        raise ShikayatError("Is graahak ne business WhatsApp par kabhi message nahi kiya: draft apne WhatsApp / SMS se bhejiye.")
    try:
        await inbox.send(tenant, contact, text, by="owner")
    except inbox.InboxError as e:
        raise ShikayatError(str(e))
    return mark_told(tenant, cid, "business WhatsApp", kind)


def mark_told(tenant: str, case_ref: str, how: str = "whatsapp", kind: str = "close") -> dict[str, Any]:
    with session() as s:
        c = _get(s, tenant, case_ref)
        if kind == "ack" and not c.ack_at:
            c.ack_at = _now()
            if c.stage == "new":
                c.stage = "acknowledged"
        if kind in ("close", "reject"):
            c.told_at = _now()
        _log(c, f"Graahak ko {kind} message bheja ({how})")
        return {"say": f"{c.number}: graahak ko bataya ({how}).", "case_id": c.id}


def confirm(tenant: str, case_ref: str, fixed: bool = True, satisfaction: Any = None, feedback: str = "", by: str = "owner") -> dict[str, Any]:
    """The customer's answer: fixed or not, and a 1-5 rating (asked after closing)."""
    with session() as s:
        c = _get(s, tenant, case_ref)
        if not fixed:
            c.confirmed = False
            if c.stage in ("resolved", "closed"):
                c.stage = "reopened"
                c.reopen_count += 1
                c.closed_at = None
                st = _st(s, tenant)
                c.due_at = _now() + timedelta(days=R.targets(c.sector, st.ack_hours or None, st.resolve_days or None)["resolve_days"])
            _log(c, "Graahak: theek nahi hua" + (f" - {feedback[:200]}" if feedback else ""), by)
            s.flush()
            return {"case": _view(s, tenant, c, False), "say": f"{c.number} phir khuli: graahak ne kaha theek nahi hua. Nayi tareekh {_local_day(tenant, c.due_at)}."}
        c.confirmed = True
        if satisfaction not in (None, ""):
            sat = int(_num(satisfaction))
            if not 1 <= sat <= 5:
                raise ShikayatError("Rating 1 se 5.")
            c.satisfaction = sat
        if feedback:
            c.feedback = R.scrub(feedback)[0][:1000]
        _log(c, "Graahak: theek ho gaya" + (f", rating {c.satisfaction}/5" if c.satisfaction else ""), by)
        s.flush()
        v = _view(s, tenant, c, False)
    say = f"{v['number']}: graahak ne pushti ki." + (f" Rating {v['satisfaction']}/5." if v["satisfaction"] else "")
    if v["satisfaction"] and v["satisfaction"] <= 2:
        say += " Rating kam hai: ek baar khud baat kar lijiye."
    return {"case": v, "say": say}


def close(tenant: str, case_ref: str, force: bool = False, reason: str = "", by: str = "owner", conv: str = "") -> dict[str, Any]:
    with session() as s:
        c = _get(s, tenant, case_ref)
        if c.stage == "closed":
            return {"case": _view(s, tenant, c, False), "say": f"{c.number} pehle hi band hai."}
        if not c.resolution:
            raise ShikayatError("Pehle 'resolve' kijiye: kya hal kiya.")
        ck = _checklist(s, tenant, c)
        left = [x for x in ck if not x["ok"]]
        if any(x["item"] == "compensation_decided" for x in left):
            raise ShikayatError("Muavze ka card abhi baaki hai: pehle haan / na kijiye.")
        if left and not force:
            raise ShikayatError("Band karne se pehle: " + "; ".join(x["say"] for x in left) + ". (Aap chahein to wajah likh kar 'phir bhi band' kar sakti hain.)")
        if left and len(str(reason).strip()) < 5:
            raise ShikayatError("Checklist poori nahi; 'phir bhi band' ke liye wajah likhiye.")
        c.stage, c.closed_at = "closed", _now()
        _log(c, "Band" + (f" (checklist adhoori: {', '.join(x['item'] for x in left)}; wajah: {reason[:150]})" if left else ""), by)
        st = _st(s, tenant)
        learned = _j(st.learned, {})
        days = ((_utc(c.closed_at) - _utc(c.received_at)).total_seconds() / 86400) if c.received_at else 0
        comp_total = sum(x.amount for x in _approved_comps(s, c.id))
        for cause in _j(c.causes, []) or ["other"]:
            L = learned.setdefault(cause, {"cases": 0, "days": 0.0, "comp": 0.0, "late": 0})
            L["cases"] += 1
            L["days"] = round(L["days"] + days, 2)
            L["comp"] = round(L["comp"] + comp_total, 2)
            if _utc(c.due_at) and c.closed_at > _utc(c.due_at):
                L["late"] += 1
        st.learned = _dump(learned)
        s.flush()
        v = _view(s, tenant, c, False)
    say = f"{v['number']} band ({days:.1f} din mein)."
    if v["satisfaction"] is None:
        say += " Graahak se 1-5 rating poochhne ka draft bana doon?"
    if v["ticket_ref"]:
        say += f" Hoola ticket {v['ticket_ref']} bhi update karna ho to app mein 'Ticket update' dabaiye."
    _publish(tenant, "shikayat_case", {"case": {k: v[k] for k in ("id", "number", "customer", "stage", "stage_label", "severity")},
                                       "event": "closed"}, conv)
    return {"case": v, "say": say}


def reopen(tenant: str, case_ref: str, reason: str = "", by: str = "owner") -> dict[str, Any]:
    reason = R.scrub(reason)[0]
    with session() as s:
        st = _st(s, tenant)
        c = _get(s, tenant, case_ref)
        if c.stage in R.OPEN:
            return {"case": _view(s, tenant, c, False), "say": f"{c.number} abhi khuli hi hai."}
        ca = _utc(c.closed_at) or _utc(c.resolved_at) or _now()
        if _now() - ca <= timedelta(days=st.reopen_days):
            c.stage, c.reopen_count, c.closed_at, c.confirmed = "reopened", c.reopen_count + 1, None, None
            c.due_at = _now() + timedelta(days=R.targets(c.sector, st.ack_hours or None, st.resolve_days or None)["resolve_days"])
            _log(c, "Phir khuli" + (f": {reason[:200]}" if reason else ""), by)
            s.flush()
            v = _view(s, tenant, c, False)
            return {"case": v, "say": f"{v['number']} phir khuli ({v['reopen_count']} baar). Nayi tareekh {v['due']}."}
        parent = c.number
        data = {"customer": c.customer, "phone": c.phone, "product": c.product, "order_ref": c.order_ref, "channel": c.channel,
                "contact_id": c.contact_id, "amount": c.amount}
    out = log_complaint(tenant, (reason or "Pichhli shikayat jaisi dikkat phir aayi") + f" (pichhli: {parent})", by=by, **data)
    with session() as s:
        n = s.get(ShikayatCase, out["case"]["id"])
        n.parent = parent
    out["say"] = f"{parent} ko band hue {st.reopen_days} din se zyada ho gaye, isliye nayi shikayat {out['case']['number']} banayi (pichhli se judi)."
    return out


def escalate(tenant: str, case_ref: str, to: str, ref: str = "", by: str = "owner") -> dict[str, Any]:
    """The customer went outside (NCH 1915, ombudsman, consumer commission, social media). Recorded; the case gets priority."""
    to = str(to or "").strip()[:80]
    if not to:
        raise ShikayatError("Graahak kahan gaye? (1915 / ombudsman / consumer court / social media ...)")
    with session() as s:
        c = _get(s, tenant, case_ref)
        c.escalated = _dump({"to": to, "ref": str(ref or "")[:80], "on": today(tenant).isoformat()})
        c.severity = "urgent" if c.severity == "urgent" else "high"
        fl = set(_j(c.flags, []))
        fl.add("escalated")
        c.flags = _dump(sorted(fl))
        if c.stage in ("closed", "resolved", "rejected"):
            c.stage = "reopened"
        _log(c, f"Graahak bahar gaye: {to}" + (f" ({ref})" if ref else ""), by)
        s.flush()
        v = _view(s, tenant, c)
    low = to.lower()
    tip = "Har notice / message ka jawab samay par, likhit mein; saare saboot aur log taiyaar rakhiye."
    if "1915" in low or "nch" in low or "helpline" in low:
        tip = "NCH company ko shikayat bhejta hai: wahan jawab dijiye aur hal ka saboot lagaiye. " + tip
    elif "court" in low or "commission" in low or "jagriti" in low or "forum" in low:
        tip = "Commission ka notice aaye to tay samay mein jawab (vakil se salah lijiye; Hoola kanooni salah nahi deta). " + tip
    elif "ombudsman" in low or "lokpal" in low or "rbi" in low or "irdai" in low:
        tip = "Ombudsman ko apna paksh aur saboot samay par dijiye; award aaye to us ki samay-seema dekhiye. " + tip
    return {"case": v, "say": f"{v['number']}: graahak {to} gaye - zaroori maamla. {tip}"}


# ---------------------------------------------------------------- compensation (HER tap only)
def _comp_show(x: ShikayatComp, c: ShikayatCase) -> dict[str, Any]:
    return {"title": f"{c.number} muavza: {R.COMP_KINDS.get(x.kind, x.kind)}", "case": c.number, "customer": c.customer,
            "kind": x.kind, "amount": _rs(x.amount) if x.amount else "", "item": x.item, "reason": x.reason[:300], "comp": x.number}


def _comp_hash(show: dict) -> str:
    return hashlib.sha256(json.dumps(show, sort_keys=True, ensure_ascii=False).encode()).hexdigest()


def comp_request(tenant: str, case_ref: str, kind: str, amount: Any = None, item: str = "", reason: str = "", conv: str = "",
                 by: str = "owner") -> dict[str, Any]:
    """Prepare a compensation card. Nothing is approved, promised or paid here."""
    kind = str(kind or "").lower().strip()
    kind = {"paisa wapas": "refund", "wapas": "refund", "badal": "replacement", "naya": "replacement", "chhoot": "discount",
            "maafi": "goodwill", "free": "free_service", "muft": "free_service", "voucher": "credit"}.get(kind, kind)
    if kind not in R.COMP_KINDS:
        raise ShikayatError("Muavza: " + " | ".join(R.COMP_KINDS))
    amt = _num(amount)
    if amt < 0 or amt > 1_00_00_000:
        raise ShikayatError("Rakam sahi nahi.")
    if kind in R.MONEY_KINDS and amt <= 0:
        raise ShikayatError("Kitne rupaye? (rakam ke bina card nahi banta)")
    if kind == "replacement" and not str(item).strip():
        raise ShikayatError("Kya badal kar dena hai? (item)")
    reason = R.scrub(reason)[0]
    with _LOCK, session() as s:
        st = _st(s, tenant)
        c = _get(s, tenant, case_ref)
        if c.stage == "closed":
            raise ShikayatError(f"{c.number} band hai; pehle reopen kijiye.")
        checks = []
        if st.comp_limit and amt > st.comp_limit:
            checks.append(f"Aapki seema {_rs(st.comp_limit)} se zyada")
        if c.amount and amt > c.amount:
            checks.append(f"Kharid ki keemat {_rs(c.amount)} se zyada")
        month0 = datetime(today(tenant).year, today(tenant).month, 1, tzinfo=_tz(tenant)).astimezone(timezone.utc)
        month_total = sum(x.amount for x in s.scalars(select(ShikayatComp).where(
            ShikayatComp.tenant_id == tenant, ShikayatComp.status.in_(("approved", "given")), ShikayatComp.decided_at >= month0)).all())
        if st.month_limit and month_total + amt > st.month_limit:
            checks.append(f"Is mahine ka muavza {_rs(month_total + amt)} ho jayega (seema {_rs(st.month_limit)})")
        if c.phone:
            ids = [x.id for x in s.scalars(select(ShikayatCase).where(ShikayatCase.tenant_id == tenant, ShikayatCase.phone == c.phone,
                                                                         ShikayatCase.id != c.id)).all()]
            if ids:
                past = s.scalars(select(ShikayatComp).where(ShikayatComp.case_id.in_(ids), ShikayatComp.status.in_(("approved", "given")),
                                                            ShikayatComp.decided_at >= _now() - timedelta(days=365))).all()
                if past:
                    checks.append(f"Is graahak ko saal bhar mein {len(past)} baar muavza mila ({_rs(sum(p.amount for p in past))})")
        learned = _j(st.learned, {})
        for cause in _j(c.causes, []):
            L = learned.get(cause)
            if L and L.get("cases", 0) >= 3 and L.get("comp"):
                checks.append(f"'{R.CAUSE_LABEL.get(cause, cause)}' mein ab tak ausat {_rs(L['comp'] / L['cases'])} muavza")
                break
        for old in s.scalars(select(ShikayatComp).where(ShikayatComp.case_id == c.id, ShikayatComp.status == "pending")).all():
            old.status = "cancelled"
            old.note = "naya card bana"
        st.comp_seq += 1
        x = ShikayatComp(tenant_id=tenant, case_id=c.id, number=f"C-{st.comp_seq:04d}", kind=kind, amount=amt,
                         item=str(item or "").strip()[:200], reason=reason.strip()[:1000], checks=_dump(checks), conv=str(conv or "")[:40])
        s.add(x)
        s.flush()
        x.hash = _comp_hash(_comp_show(x, c))
        _log(c, f"Muavze ka card {x.number}: {R.COMP_KINDS[kind]}" + (f" {_rs(amt)}" if amt else "") + " (aapki haan ka intezaar)", by)
        s.flush()
        v = _comp_view(tenant, x)
        show = _comp_show(x, c)
    v["money"] = amt > 0
    _publish(tenant, "shikayat_comp", {"comp": v, "show": show}, conv)
    say = (f"{show['case']}: {R.COMP_KINDS[kind]}" + (f" {_rs(amt)}" if amt else "") + (f" ({v['item']})" if v["item"] else "") +
           " ka card bheja. Aapki haan ke baad hi yeh tay hoga; graahak se abhi vaada nahi." + (" Dhyaan: " + "; ".join(checks) + "." if checks else ""))
    return {"comp": v, "show": show, "say": say, "needs_tap": True}


def comp_card(tenant: str, comp_id: str) -> dict[str, Any] | None:
    with session() as s:
        x = s.get(ShikayatComp, comp_id)
        if x is None or x.tenant_id != tenant:
            return None
        c = s.get(ShikayatCase, x.case_id)
        return _comp_show(x, c) | {"status": x.status, "hash": x.hash}


def comp_needs_fingerprint(tenant: str, comp_id: str) -> bool:
    from .. import approvekey as ak
    card = comp_card(tenant, comp_id)
    if not card:
        return False
    with session() as s:
        x = s.get(ShikayatComp, comp_id)
        money = bool(x and x.amount > 0)
    return money and ak.active_on(tenant)


def comp_challenge(tenant: str, comp_id: str) -> dict[str, Any]:
    """Same signed-challenge format as approvekey (kind 'shikayat'): the phone signs this card as it is NOW."""
    from .. import approvekey as ak
    card = comp_card(tenant, comp_id)
    if not card or card["status"] != "pending":
        raise ShikayatError("Yeh card ab baaki nahi hai.")
    digest = _comp_hash({k: v for k, v in card.items() if k not in ("status", "hash")})
    nonce = secrets.token_urlsafe(18)
    exp = datetime.now(timezone.utc) + ak.CHALLENGE_TTL
    msg = f"hoola-approve|v1|{tenant}|shikayat|{comp_id}|{digest}|{nonce}|{int(exp.timestamp())}"
    with session() as s:
        ch = ak.ApproveChallenge(tenant_id=tenant, kind="shikayat", target=comp_id, digest=digest, message=msg, expires_at=exp)
        s.add(ch)
        s.flush()
        chid = ch.id
    show = {"title": card["title"], "amount": card["amount"], "payee": card["customer"]}
    return {"challenge_id": chid, "message": msg, "show": show, "expires_at": exp.isoformat()}


def _verify_fingerprint(tenant: str, comp_id: str, proof: dict | None) -> None:
    from .. import approvekey as ak
    if not proof:
        raise ShikayatError("fingerprint_needed")
    card = comp_card(tenant, comp_id)
    digest = _comp_hash({k: v for k, v in (card or {}).items() if k not in ("status", "hash")})
    with session() as s:
        ak._settle(s, tenant)
        ok = ak._check(s, tenant, "shikayat", comp_id, proof, digest)
    if not ok:
        raise ShikayatError("fingerprint_invalid")


def decide_comp(tenant: str, comp_id: str, yes: bool, card_hash: str = "", proof: dict | None = None, note: str = "",
                by_tap: bool = False) -> dict[str, Any]:
    """ONLY her tap in the app (API). The tool can never reach this."""
    if not by_tap:
        raise ShikayatError("Muavze ki haan sirf aapke tap se (app ka card).")
    with session() as s:
        x = s.get(ShikayatComp, comp_id)
        if x is None or x.tenant_id != tenant:
            raise ShikayatError("Card nahi mila.")
        if x.status != "pending":
            raise ShikayatError(f"Yeh card pehle hi {x.status} hai.")
        c = s.get(ShikayatCase, x.case_id)
        if not card_hash or card_hash != x.hash or _comp_hash(_comp_show(x, c)) != x.hash:
            raise ShikayatError("Card badal gaya hai; naya card dekhiye.")
        money = x.amount > 0
    if yes and money:
        from .. import approvekey as ak
        if ak.active_on(tenant):
            _verify_fingerprint(tenant, comp_id, proof)
    with session() as s:
        x = s.get(ShikayatComp, comp_id)
        c = s.get(ShikayatCase, x.case_id)
        x.status = "approved" if yes else "rejected"
        x.decided_at = _now()
        x.note = R.scrub(note)[0][:500]
        _log(c, f"Muavza {x.number} " + ("MANA (aapki haan)" if yes else "NAHI maana") + (f": {x.note}" if x.note else ""), "owner")
        s.flush()
        v = _comp_view(tenant, x)
        num = c.number
        conv = x.conv
    if yes:
        say = (f"{num}: {v['kind_label']}" + (f" {v['amount_text']}" if v["amount_text"] else "") + " maana. Aap khud dijiye "
               "(UPI / cash / naya saamaan) aur 'de diya' dabaiye; closing message mein ab yeh likha jayega.")
    else:
        say = f"{num}: muavza nahi. Graahak ko vinamra jawab mein wajah batayiye."
    _publish(tenant, "shikayat_comp", {"comp": v, "event": "decided"}, conv)
    return {"comp": v, "say": say}


def comp_given(tenant: str, comp_id: str, how: str = "", ref: str = "") -> dict[str, Any]:
    with session() as s:
        x = s.get(ShikayatComp, comp_id)
        if x is None or x.tenant_id != tenant:
            x = s.scalars(select(ShikayatComp).where(ShikayatComp.tenant_id == tenant, ShikayatComp.number == str(comp_id).upper())).first()
        if x is None:
            raise ShikayatError("Card nahi mila.")
        if x.status != "approved":
            raise ShikayatError("Sirf aapki haan wala muavza 'de diya' ho sakta hai.")
        x.status, x.given_at = "given", _now()
        x.given_how = (str(how or "")[:120] + (f" ref ...{_digits(ref)[-4:]}" if _digits(ref) else "")).strip()
        c = s.get(ShikayatCase, x.case_id)
        _log(c, f"Muavza {x.number} diya ({x.given_how or 'aapne bataya'})")
        s.flush()
        return {"comp": _comp_view(tenant, x), "say": f"{c.number}: muavza diya darj."}


# ---------------------------------------------------------------- Hoola Network / WhatsApp tickets
def import_tickets(tenant: str) -> dict[str, Any]:
    """Open hoola_tickets (Hoola Network customers and WhatsApp biz_complaint) become complaints here, once each."""
    from ..net.business import Ticket
    with session() as s:
        have = set(s.scalars(select(ShikayatCase.ticket_ref).where(ShikayatCase.tenant_id == tenant, ShikayatCase.ticket_ref != "")).all())
        rows = s.scalars(select(Ticket).where(Ticket.tenant_id == tenant, Ticket.status.in_(("open", "assigned", "in_progress")))).all()
        todo = [(t.number, t.customer, t.customer_label, t.text, t.category, _utc(t.created_at)) for t in rows if t.number not in have]
    made = []
    for number, cust, label, text, cat, at in todo:
        wa = cust.startswith("wa:")
        try:
            out = log_complaint(tenant, text, customer=re.sub(r"\s*\(.*\)$", "", label or "")[:80], phone=cust[3:] if wa else "",
                                channel="whatsapp" if wa else "network", product=cat or "", ticket_ref=number, received=at, by="hoola")
            made.append(out["case"]["number"])
        except ShikayatError:
            continue
    return {"imported": made, "say": (f"{len(made)} ticket shikayat mein aaye: " + ", ".join(made)) if made else "Koi naya ticket nahi."}


def ticket_sync(tenant: str, case_ref: str) -> dict[str, Any]:
    """Her tap: the Hoola ticket gets the same status (the customer's avatar / WhatsApp is told by the ticket system)."""
    from ..net import business
    with session() as s:
        c = _get(s, tenant, case_ref)
        if not c.ticket_ref:
            raise ShikayatError("Is shikayat ka koi Hoola ticket nahi.")
        st = {"closed": "closed", "resolved": "resolved", "rejected": "closed", "new": "open", "acknowledged": "assigned"}.get(c.stage, "in_progress")
        ref, text = c.ticket_ref, (c.resolution or R.STAGE_LABEL.get(c.stage, ""))[:400]
    try:
        business.update_ticket(tenant, ref, st, text)
    except Exception as e:
        raise ShikayatError(f"Ticket update nahi hua: {e}")
    return {"say": f"Ticket {ref}: {st}."}


# ---------------------------------------------------------------- privacy (DPDP): forget a customer
def forget_customer(tenant: str, phone: str) -> dict[str, Any]:
    ph = _phone(phone)
    if len(ph) != 10:
        raise ShikayatError("10 ank ka number.")
    n = 0
    with session() as s:
        for c in s.scalars(select(ShikayatCase).where(ShikayatCase.tenant_id == tenant, ShikayatCase.phone == ph)).all():
            for p in s.scalars(select(ShikayatProof).where(ShikayatProof.case_id == c.id)).all():
                Path(p.path).unlink(missing_ok=True)
                s.delete(p)
            c.customer, c.phone, c.contact_id, c.text, c.feedback = "[hataya]", "", "", "[graahak ke kehne par hataya]", ""
            c.log = _dump([{"at": _now().isoformat(timespec="seconds"), "by": "owner", "text": "Personal jaankari hatayi (DPDP)"}])
            n += 1
    return {"forgotten": n, "say": f"{n} shikayaton se graahak ki personal jaankari aur saboot hata diye (ginti register mein rahegi)."}


# ---------------------------------------------------------------- register + summary
def _month_bounds(tenant: str, month: str) -> tuple[datetime, datetime, str]:
    if month and re.fullmatch(r"\d{4}-\d{2}", month):
        y, m = int(month[:4]), int(month[5:])
    else:
        d = today(tenant)
        y, m = d.year, d.month
    tz = _tz(tenant)
    start = datetime(y, m, 1, tzinfo=tz).astimezone(timezone.utc)
    end = datetime(y + (m == 12), m % 12 + 1, 1, tzinfo=tz).astimezone(timezone.utc)
    return start, end, f"{y}-{m:02d}"


def summary(tenant: str, month: str = "") -> dict[str, Any]:
    start, end, label = _month_bounds(tenant, month)
    with session() as s:
        st = _st(s, tenant)
        rows = s.scalars(select(ShikayatCase).where(ShikayatCase.tenant_id == tenant)).all()
        comps = s.scalars(select(ShikayatComp).where(ShikayatComp.tenant_id == tenant, ShikayatComp.decided_at >= start,
                                                     ShikayatComp.decided_at < end)).all()

        def closed_by(c, t):
            ca = _utc(c.closed_at)
            return c.stage in ("closed", "rejected") and ca is not None and ca < t

        opening = [c for c in rows if _utc(c.received_at) < start and not closed_by(c, start)]
        received = [c for c in rows if start <= _utc(c.received_at) < end]
        disposed = [c for c in rows if c.stage in ("closed", "rejected") and _utc(c.closed_at) and start <= _utc(c.closed_at) < end]
        pending = [c for c in rows if _utc(c.received_at) < end and not closed_by(c, end)]
        ack_ok = [c for c in received if c.ack_at and _utc(c.ack_at) <= _utc(c.ack_by)]
        in_time = [c for c in disposed if _utc(c.due_at) and _utc(c.closed_at) <= _utc(c.due_at)]
        ages = {"0-15": 0, "16-30": 0, "31-90": 0, "90+": 0}
        for c in pending:
            d = (min(end, _now()) - _utc(c.received_at)).days
            ages["0-15" if d <= 15 else "16-30" if d <= 30 else "31-90" if d <= 90 else "90+"] += 1
        by_cause: dict[str, int] = {}
        by_channel: dict[str, int] = {}
        for c in received:
            for x in _j(c.causes, []) or ["other"]:
                by_cause[x] = by_cause.get(x, 0) + 1
            by_channel[c.channel] = by_channel.get(c.channel, 0) + 1
        days = [(_utc(c.closed_at) - _utc(c.received_at)).total_seconds() / 86400 for c in disposed]
        sats = [c.satisfaction for c in disposed if c.satisfaction]
        products: dict[str, int] = {}
        for c in received:
            if c.product_key:
                products[c.product] = products.get(c.product, 0) + 1
        out = {
            "month": label, "sector": st.sector, "sector_label": R.SECTORS[R.norm_sector(st.sector)]["label"],
            "opening_pending": len(opening), "received": len(received), "disposed": len(disposed), "closing_pending": len(pending),
            "rejected": len([c for c in disposed if c.stage == "rejected"]),
            "ack_in_time_pct": round(100 * len(ack_ok) / len(received)) if received else None,
            "closed_in_time_pct": round(100 * len(in_time) / len(disposed)) if disposed else None,
            "avg_days_to_close": round(sum(days) / len(days), 1) if days else None,
            "pending_age": ages, "by_cause": {R.CAUSE_LABEL.get(k, k): v for k, v in sorted(by_cause.items(), key=lambda kv: -kv[1])},
            "by_channel": by_channel, "satisfaction_avg": round(sum(sats) / len(sats), 1) if sats else None, "rated": len(sats),
            "comp_approved": round(sum(x.amount for x in comps if x.status in ("approved", "given")), 2),
            "comp_given": round(sum(x.amount for x in comps if x.status == "given"), 2),
            "comp_count": len([x for x in comps if x.status in ("approved", "given")]),
            "reopened": len([c for c in received + disposed if c.reopen_count]),
            "escalated": len([c for c in received if "escalated" in _j(c.flags, [])]),
            "repeat_customers": len({c.phone for c in received if "repeat_customer" in _j(c.flags, []) and c.phone}),
            "top_products": sorted(products.items(), key=lambda kv: -kv[1])[:5],
            "law": R.SECTORS[R.norm_sector(st.sector)]["law"],
        }
    tips = []
    if out["closed_in_time_pct"] is not None and out["closed_in_time_pct"] < 80:
        tips.append(f"Sirf {out['closed_in_time_pct']}% samay par band hui")
    if out["by_cause"]:
        k, v = next(iter(out["by_cause"].items()))
        if v >= 3:
            tips.append(f"Sabse zyada: {k} ({v}) - jad se theek kijiye")
    if out["pending_age"]["31-90"] + out["pending_age"]["90+"]:
        tips.append(f"{out['pending_age']['31-90'] + out['pending_age']['90+']} shikayatein 30 din se purani")
    out["tips"] = tips
    out["say"] = (f"{label}: {out['received']} aayi, {out['disposed']} band, {out['closing_pending']} baaki"
                  + (f", {out['closed_in_time_pct']}% samay par" if out["closed_in_time_pct"] is not None else "")
                  + (f", rating {out['satisfaction_avg']}/5" if out["satisfaction_avg"] else "")
                  + (f", muavza {_rs(out['comp_approved'])}" if out["comp_approved"] else "") + "." + (" " + "; ".join(tips) + "." if tips else ""))
    return out


def register_file(tenant: str, month: str = "", conv: str = "") -> dict[str, Any]:
    import openpyxl
    from openpyxl.styles import Font
    from .. import chatfiles
    start, end, label = _month_bounds(tenant, month)
    sm = summary(tenant, label)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Register"
    head = ["No", "Received", "Customer", "Phone", "Channel", "Product", "Order", "Complaint", "Cause", "Severity", "Stage", "Owner",
            "Ack by", "Ack at", "Due", "Closed", "Within time", "Resolution", "Proof of fix", "Confirmed", "Rating", "Compensation",
            "Reopened", "Escalated", "Docket / ticket"]
    ws.append(head)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    with session() as s:
        rows = s.scalars(select(ShikayatCase).where(ShikayatCase.tenant_id == tenant).order_by(ShikayatCase.received_at)).all()
        for c in rows:
            rec = _utc(c.received_at)
            ca = _utc(c.closed_at)
            in_month = start <= rec < end or (ca and start <= ca < end) or (rec < end and c.stage in R.OPEN)
            if not in_month:
                continue
            comps = _approved_comps(s, c.id)
            fixp = bool(c.fix_proof) or bool(s.scalars(select(ShikayatProof.id).where(ShikayatProof.case_id == c.id, ShikayatProof.stage == "fix")).first())
            ws.append([c.number, _local(tenant, c.received_at), c.customer, ("..." + c.phone[-4:]) if c.phone else "", c.channel, c.product, c.order_ref,
                       c.text[:500], ", ".join(R.CAUSE_LABEL.get(x, x) for x in _j(c.causes, [])), c.severity, R.STAGE_LABEL.get(c.stage, c.stage),
                       c.owner, _local(tenant, c.ack_by), _local(tenant, c.ack_at), _local(tenant, c.due_at), _local(tenant, c.closed_at),
                       ("Haan" if ca and _utc(c.due_at) and ca <= _utc(c.due_at) else "Nahi") if ca else "",
                       c.resolution[:500], "Haan" if fixp else "", {True: "Haan", False: "Nahi"}.get(c.confirmed, ""), c.satisfaction or "",
                       "; ".join(f"{x.kind} {x.amount:g}" for x in comps), c.reopen_count or "", _j(c.escalated, {}).get("to", ""),
                       c.external_ref or c.ticket_ref])
    ws2 = wb.create_sheet("Summary")
    ws2.append(["Month", sm["month"]])
    ws2.append(["Sector", sm["sector_label"]])
    for k, lab in (("opening_pending", "Pending at start of month"), ("received", "Received during month"), ("disposed", "Disposed during month"),
                   ("rejected", "  of which not accepted (with reasons)"), ("closing_pending", "Pending at end of month"),
                   ("ack_in_time_pct", "Acknowledged in time %"), ("closed_in_time_pct", "Resolved in time %"),
                   ("avg_days_to_close", "Average days to close"), ("satisfaction_avg", "Average rating (1-5)"),
                   ("comp_approved", "Compensation approved (Rs)"), ("comp_given", "Compensation given (Rs)"), ("reopened", "Reopened"),
                   ("escalated", "Escalated by customer"), ("repeat_customers", "Repeat complainants")):
        ws2.append([lab, sm[k] if sm[k] is not None else ""])
    ws2.append([])
    ws2.append(["Pending by age (days)"])
    for k, v in sm["pending_age"].items():
        ws2.append([k, v])
    ws2.append([])
    ws2.append(["By root cause"])
    for k, v in sm["by_cause"].items():
        ws2.append([k, v])
    ws2.append([])
    ws2.append(["By channel"])
    for k, v in sm["by_channel"].items():
        ws2.append([k, v])
    ws2.append([])
    ws2.append(["Rule followed", sm["law"]])
    ws2.append(["Note", "Prepared by Hoola from her own records. For a regulator's return, copy these numbers into that regulator's own form."])
    for cell in ws2["A"]:
        cell.font = Font(bold=cell.row <= 2 or cell.value in ("Pending by age (days)", "By root cause", "By channel"))
    bio = io.BytesIO()
    wb.save(bio)
    f = chatfiles.store(tenant, f"Shikayat_register_{label}.xlsx", bio.getvalue(),
                        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    _publish(tenant, "pc_file", {**f, "source": "shikayat", "title": f"Shikayat register {label}"}, conv)
    return {"file": f, "summary": sm, "say": f"Shikayat register {label} Excel mein (Register + Summary). " + sm["say"]}


# ---------------------------------------------------------------- the day, reminders
def day(tenant: str) -> dict[str, Any]:
    n = _now()
    parts: list[str] = []
    with session() as s:
        rows = s.scalars(select(ShikayatCase).where(ShikayatCase.tenant_id == tenant, ShikayatCase.stage.in_(R.OPEN + ("resolved",)))).all()
        pend_comp = s.scalars(select(ShikayatComp).where(ShikayatComp.tenant_id == tenant, ShikayatComp.status == "pending")).all()
        appr = s.scalars(select(ShikayatComp).where(ShikayatComp.tenant_id == tenant, ShikayatComp.status == "approved")).all()
        open_ = [c for c in rows if c.stage in R.OPEN]
        no_ack = [c for c in open_ if not c.ack_at]
        ack_late = [c for c in no_ack if _utc(c.ack_by) and n > _utc(c.ack_by)]
        overdue = [c for c in open_ if _utc(c.due_at) and n > _utc(c.due_at)]
        due_soon = [c for c in open_ if _utc(c.due_at) and timedelta(0) <= _utc(c.due_at) - n <= timedelta(days=2)]
        urgent = [c for c in open_ if c.severity == "urgent"]
        resolved = [c for c in rows if c.stage == "resolved"]
        not_told = [c for c in resolved if not c.told_at]
        if urgent:
            parts.append("Zaroori: " + ", ".join(f"{c.number} ({c.customer or c.product or 'graahak'})" for c in urgent[:4]))
        if ack_late:
            parts.append(f"{len(ack_late)} ki pavti der se: " + ", ".join(c.number for c in ack_late[:5]))
        elif no_ack:
            parts.append(f"{len(no_ack)} ki pavti baaki")
        if overdue:
            parts.append(f"{len(overdue)} samay-seema paar: " + ", ".join(c.number for c in overdue[:5]))
        if due_soon:
            parts.append(f"{len(due_soon)} ki tareekh 2 din mein: " + ", ".join(c.number for c in due_soon[:5]))
        if not_told:
            parts.append(f"{len(not_told)} hal ho gayi par graahak ko bataya nahi: " + ", ".join(c.number for c in not_told[:5]))
        if pend_comp:
            parts.append(f"{len(pend_comp)} muavze ke card aapki haan ke intezaar mein")
        if appr:
            parts.append(f"{len(appr)} maana hua muavza abhi diya nahi")
        v = {"open": len(open_), "no_ack": len(no_ack), "ack_late": len(ack_late), "overdue": len(overdue), "due_soon": len(due_soon),
             "urgent": len(urgent), "resolved_not_told": len(not_told), "comp_pending": len(pend_comp), "comp_to_give": len(appr)}
    return {"parts": parts, **v, "say": ("Shikayat: " + "; ".join(parts) + ".") if parts else f"Shikayat: {v['open']} khuli, sab samay par."}


def _remind(tenant: str, c: ShikayatCase, tag: str, text: str, out: list[str]) -> None:
    done = _j(c.reminded, [])
    if tag in done:
        return
    done.append(tag)
    c.reminded = _dump(done[-30:])
    out.append(text)


def tick(now_at: datetime | None = None) -> int:
    """Worker, every minute or so: import new tickets, reminders before / after each deadline (once each), the morning note."""
    global _now
    real = _now
    if now_at is not None:
        _now = lambda: _utc(now_at)          # noqa: E731
    sent = 0
    try:
        with session() as s:
            tenants = set(s.scalars(select(ShikayatSettings.tenant_id)).all()) | set(s.scalars(select(ShikayatCase.tenant_id).distinct()).all())
        for tenant in tenants:
            try:
                import_tickets(tenant)
            except Exception:
                log.exception("shikayat ticket import failed")
            lines: list[str] = []
            n = _now()
            with session() as s:
                for c in s.scalars(select(ShikayatCase).where(ShikayatCase.tenant_id == tenant, ShikayatCase.stage.in_(R.OPEN))).all():
                    ab, du = _utc(c.ack_by), _utc(c.due_at)
                    who = f"{c.number} ({c.customer or c.product or 'graahak'})"
                    if not c.ack_at and ab:
                        if n > ab:
                            _remind(tenant, c, "ack_late", f"{who}: pavti ki samay-seema nikal gayi - abhi pavti bhejiye", lines)
                        elif ab - n <= timedelta(hours=6):
                            _remind(tenant, c, "ack_soon", f"{who}: pavti {_local(tenant, ab)} tak deni hai", lines)
                    if du:
                        if n > du:
                            _remind(tenant, c, "due_late", f"{who}: hal ki tareekh nikal gayi - graahak ko likhit update aur nayi tareekh dijiye", lines)
                        elif du - n <= timedelta(days=1):
                            _remind(tenant, c, "due_1", f"{who}: kal tak hal karna hai", lines)
                        elif du - n <= timedelta(days=3):
                            _remind(tenant, c, "due_3", f"{who}: hal ke 3 din bache", lines)
                st = s.get(ShikayatSettings, tenant)
                morning = False
                if st:
                    local = n.astimezone(_tz(tenant))
                    try:
                        hh, mm = [int(x) for x in (st.morning_time or "09:30").split(":")]
                    except ValueError:
                        hh, mm = 9, 30
                    late = (local.hour * 60 + local.minute) - (hh * 60 + mm)
                    if st.last_morning != local.date().isoformat() and 0 <= late <= 180:
                        st.last_morning = local.date().isoformat()
                        morning = True
            if lines:
                text = "Shikayat yaad dilana: " + "; ".join(lines[:8]) + "."
                _note(tenant, text)
                _publish(tenant, "shikayat_alert", {"text": text, "count": len(lines)})
                sent += 1
            if morning:
                d = day(tenant)
                if d["parts"]:
                    _note(tenant, d["say"])
                    _publish(tenant, "shikayat_day", {"text": d["say"], **{k: d[k] for k in ("open", "overdue", "ack_late", "comp_pending")}})
                    sent += 1
    finally:
        _now = real
    return sent


def home(tenant: str) -> dict[str, Any]:
    return {"day": day(tenant), "cases": cases(tenant, "open"), "recent": cases(tenant, "", limit=30), "settings": settings(tenant),
            "sectors": {k: {"label": v["label"], "ack_hours": v["ack_hours"], "resolve_days": v["resolve_days"], "law": v["law"]}
                        for k, v in R.SECTORS.items()},
            "causes": R.CAUSE_LABEL, "comp_kinds": R.COMP_KINDS, "never": list(R.NEVER)}
