"""Courier Sahayak (courier desk), end to end.

  1. Gets the job: "Riya ko kurti bhejni hai, COD 899, Akola 444001", an order she forwards, a customer asking "mera parcel
     kahan hai?" on the business WhatsApp, a carrier's COD remittance sheet, a Wapsi return that needs a return label.
  2. Does the work: parcel P-0001 with the customer's address, dead + volumetric weight (L x B x H / 5000), COD / prepaid,
     value; the PIN code checked with HER carrier's API (serviceable? COD allowed?), the carrier's OWN rate quotes ranked by
     her rules; e-way bill needed above the Rule 138 limit; prohibited goods refused.
  3. Waits for her tap: booking a pickup can charge her courier wallet, so `book_request` only makes a card (hash of the
     parcel + the quote); ONLY `decide(by_tap=True)` from her app's POST /api/courier/cards/{id}/decide books it (fingerprint
     too once a phone key is active). Same for return labels (reverse pickup), NDR re-attempt / RTO and cancelling a booking.
  4. Delivers: AWB + label PDF as her file, pickup request; tracking polled by the worker -> the carrier's own words kept
     exactly + our 8 states; NDR -> what to ask the customer + a draft she sends; "where is my parcel?" for her (search by
     P-number / order / AWB / phone / name) and for a customer (only their own parcels, by their number); COD remittance
     sheet matched to delivered COD parcels; daily manifest.
  5. Never: books / charges without her tap, invents a status / rate / date, shows a customer anyone else's parcel, touches
     COD money, generates an e-way bill.
"""
from __future__ import annotations

import csv
import io
import json
import logging
import os
import re
import threading
from datetime import date, datetime, timedelta, timezone
from pathlib import Path
from typing import Any

from sqlalchemy import or_, select

from ..db import now, session
from . import carriers as K
from . import rules as R
from .models import CourierCard, CourierParcel, CourierRemit, CourierSettings

log = logging.getLogger("santosh.courier")
_LOCK = threading.Lock()
POLL_MIN = int(os.environ.get("COURIER_POLL_MIN") or 30)
KEY_FIELDS = {p: c["keys"] for p, c in R.CARRIERS.items() if c["keys"]}
API_PROVIDERS = ("shiprocket", "delhivery")


class CourierError(Exception):
    pass


# ---------------------------------------------------------------- helpers
def _j(s: str, default):
    try:
        v = json.loads(s or "")
        return v if isinstance(v, type(default)) else default
    except (TypeError, ValueError):
        return default


def _dump(v) -> str:
    return json.dumps(v, ensure_ascii=False, default=str)


def _rs(v) -> str:
    v = round(float(v or 0), 2)
    whole = int(abs(v))
    s = str(whole)
    if len(s) > 3:
        s = re.sub(r"(\d)(?=(\d\d)+$)", r"\1,", s[:-3]) + "," + s[-3:]
    paise = round(abs(v) - whole, 2)
    return ("-" if v < 0 else "") + "₹" + s + (f"{paise:.2f}"[1:] if paise >= 0.005 else "")


def _utc(d: datetime | None) -> datetime | None:
    if d is None:
        return None
    return d if d.tzinfo else d.replace(tzinfo=timezone.utc)


def local_now(tenant: str) -> datetime:
    from ..life import tz_of
    return datetime.now(tz_of(tenant))


def today(tenant: str) -> date:
    return local_now(tenant).date()


def _publish(tenant: str, kind: str, payload: dict, conv: str = "") -> None:
    from ..events import bus
    try:
        if not conv:
            from ..net.service import home_conv
            conv = home_conv(tenant)
        bus.publish(conv, kind, payload)
    except Exception:
        log.exception("courier publish failed")


def _note(tenant: str, text: str) -> None:
    try:
        from ..events import notify
        from ..net.service import home_conv
        notify(home_conv(tenant), text)
    except Exception:
        log.exception("courier note failed")


def _dir(tenant: str) -> Path:
    safe = re.sub(r"[^A-Za-z0-9_-]", "_", tenant)[:64]
    p = Path(os.environ.get("COURIER_DIR") or "data/courier") / safe
    p.mkdir(parents=True, exist_ok=True)
    return p


def _log(p: CourierParcel, text: str) -> None:
    lg = _j(p.log, [])
    lg.append({"at": now().isoformat(), "text": text[:300]})
    p.log = _dump(lg[-80:])
    p.updated_at = now()


def _st(s, tenant: str) -> CourierSettings:
    r = s.get(CourierSettings, tenant)
    if r is None:
        r = CourierSettings(tenant_id=tenant, rules=_dump(R.DEFAULT_RULES))
        s.add(r)
        s.flush()
    return r


def _clean_str(v, n: int = 200) -> str:
    return re.sub(r"\s+", " ", str(v or "")).strip()[:n]


# ---------------------------------------------------------------- settings, rules, keys
def _joined(st: CourierSettings) -> dict[str, bool]:
    keys = _j(st.keys, {})
    out = {p: bool(keys.get(p)) for p in KEY_FIELDS}
    try:
        from ..wapsi.models import WapsiSettings
        with session() as s:
            w = s.get(WapsiSettings, st.tenant_id)
            if w and w.courier_key and w.courier in API_PROVIDERS and not out.get(w.courier):
                out[w.courier] = True               # her Wapsi keys count (pickup location still needed for forward)
    except Exception:
        pass
    return out


def settings(tenant: str) -> dict[str, Any]:
    with session() as s:
        r = _st(s, tenant)
        rules = R.clean_rules(_j(r.rules, {}))
        return {"business": r.business, "gstin": r.gstin, "warehouse": _j(r.warehouse, {}), "default_carrier": r.default_carrier,
                "joined": _joined(r), "rules": rules, "rules_text": r.rules_text, "rule_lines": R.rule_lines(rules),
                "divisor": r.divisor, "remit_days": r.remit_days, "manifest_time": r.manifest_time, "morning_time": r.morning_time,
                "carriers": [{"id": k, "label": v["label"], "api": v["api"], "keys": list(v["keys"]), "note": v["note"]} for k, v in R.CARRIERS.items()]}


def save_settings(tenant: str, **kw: Any) -> dict[str, Any]:
    with session() as s:
        r = _st(s, tenant)
        if "business" in kw:
            r.business = _clean_str(kw["business"], 160)
        if "gstin" in kw:
            g = str(kw["gstin"] or "").strip().upper()
            if g:
                from ..india.money import gstin_ok
                if not gstin_ok(g):
                    raise CourierError("GSTIN sahi nahi lag raha.")
            r.gstin = g
        if "warehouse" in kw and isinstance(kw["warehouse"], dict):
            w = {k: _clean_str(kw["warehouse"].get(k), 300 if k == "address" else 80)
                 for k in ("name", "phone", "address", "city", "state", "pincode", "email")}
            if w["pincode"] and not R.pincode_ok(w["pincode"]):
                raise CourierError("Pickup ka PIN code 6 ankon ka chahiye.")
            if w["phone"]:
                w["phone"] = R.phone10(w["phone"])
            r.warehouse = _dump(w)
        if "default_carrier" in kw:
            c = str(kw["default_carrier"] or "")
            if c and c not in R.CARRIERS:
                raise CourierError("Courier: " + " | ".join(R.CARRIERS))
            r.default_carrier = c
        if "divisor" in kw and kw["divisor"]:
            d = R.num(kw["divisor"])
            if not 1000 <= d <= 10000:
                raise CourierError("Volumetric divisor 1000-10000 ke beech (aam taur par 5000).")
            r.divisor = d
        if "remit_days" in kw:
            r.remit_days = int(max(0, min(R.num(kw["remit_days"]), 60)))
        for k in ("manifest_time", "morning_time"):
            if kw.get(k):
                if not re.fullmatch(r"([01]\d|2[0-3]):[0-5]\d", str(kw[k])):
                    raise CourierError("Samay HH:MM mein likhiye.")
                setattr(r, k, str(kw[k]))
    return settings(tenant) | {"say": "Courier ki setting save."}


def set_rules(tenant: str, text: str = "", rules: dict | None = None) -> dict[str, Any]:
    if text:
        parsed, got = R.parse_rules(text)
    else:
        parsed, got = R.clean_rules(rules), []
    with session() as s:
        r = _st(s, tenant)
        r.rules = _dump(R.clean_rules(parsed))
        if text:
            r.rules_text = text[:1000]
    lines = R.rule_lines(parsed)
    return {"rules": R.clean_rules(parsed), "lines": lines, "understood": got,
            "say": "Courier ke niyam: " + "; ".join(lines) + "."}


def set_keys(tenant: str, provider: str, creds: dict[str, Any]) -> dict[str, Any]:
    """Her carrier keys, typed only in the app; checked with the carrier, sealed with the vault, never returned."""
    from ..bazar.security import vault
    if provider not in KEY_FIELDS:
        raise CourierError("provider: " + " | ".join(KEY_FIELDS))
    c = {k: str(creds.get(k) or "").strip() for k in KEY_FIELDS[provider]}
    if not all(c.values()):
        raise CourierError("Yeh sab chahiye: " + ", ".join(KEY_FIELDS[provider]))
    try:
        if provider == "shiprocket":
            K.ShiprocketDesk(c).token()
        elif provider == "bluedart":
            K.BlueDart(c).token()
    except K.CarrierError as e:
        raise CourierError(str(e))
    with session() as s:
        r = _st(s, tenant)
        keys = _j(r.keys, {})
        keys[provider] = vault().seal(_dump(c))
        r.keys = _dump(keys)
        if not r.default_carrier and provider in API_PROVIDERS:
            r.default_carrier = provider
    return {"ok": True, "provider": provider, "say": f"{R.CARRIERS[provider]['label']} jud gaya (chaabi band tijori mein, Hoola ko nahi dikhti)."}


def forget_keys(tenant: str, provider: str) -> dict[str, Any]:
    with session() as s:
        r = _st(s, tenant)
        keys = _j(r.keys, {})
        keys.pop(provider, None)
        r.keys = _dump(keys)
        if r.default_carrier == provider:
            r.default_carrier = ""
    return {"ok": True, "say": "Chaabi hata di."}


def _creds(tenant: str, provider: str) -> dict[str, Any]:
    from ..bazar.security import vault
    with session() as s:
        r = _st(s, tenant)
        sealed = _j(r.keys, {}).get(provider, "")
    if sealed:
        try:
            return _j(vault().open(sealed), {})
        except Exception:
            return {}
    if provider in API_PROVIDERS:                          # her Wapsi courier keys (same Shiprocket / Delhivery account)
        try:
            from ..wapsi import service as ws
            from ..wapsi.models import WapsiSettings
            with session() as s:
                w = s.get(WapsiSettings, tenant)
                if w and w.courier == provider and w.courier_key:
                    return ws._creds(w.courier_key)
        except Exception:
            return {}
    return {}


def _adapter(tenant: str, provider: str) -> K.Adapter:
    if provider in ("", "manual") or provider not in K.ADAPTERS:
        return K.Manual()
    creds = _creds(tenant, provider)
    if not creds:
        return K.Manual()
    return K.get(provider, creds)


def _api_providers(tenant: str) -> list[str]:
    return [p for p in API_PROVIDERS if _creds(tenant, p)]


# ---------------------------------------------------------------- parcels
def _clean_customer(c: dict | None) -> dict[str, str]:
    c = c or {}
    out = {k: _clean_str(c.get(k), 300 if k == "address" else 80) for k in ("name", "phone", "email", "address", "city", "state", "pincode")}
    out["phone"] = R.phone10(out["phone"]) or out["phone"][:15]
    out["pincode"] = re.sub(r"\D", "", out["pincode"])[:6]
    return out


def _clean_items(items) -> list[dict[str, Any]]:
    if isinstance(items, str):
        out = []
        for part in re.split(r"[,;\n]+", items):
            part = part.strip()
            if not part:
                continue
            m = re.match(r"(.+?)\s*(?:x\s*(\d+))?\s*(?:@\s*₹?\s*([\d,.]+))?$", part, re.I)
            name, qty, price = (m.group(1), m.group(2), m.group(3)) if m else (part, None, None)
            out.append({"name": name.strip()[:100], "qty": int(qty or 1), "price": R.num(price) if price else 0.0})
        return out[:30]
    out = []
    for i in (items or [])[:30]:
        if isinstance(i, dict) and str(i.get("name") or "").strip():
            out.append({"name": _clean_str(i.get("name"), 100), "qty": int(R.num(i.get("qty")) or 1), "price": R.num(i.get("price")),
                        "sku": _clean_str(i.get("sku"), 40)})
    return out


def _next_number(s, tenant: str, card: bool = False) -> str:
    st = _st(s, tenant)
    if card:
        st.bk_seq += 1
        return f"B-{st.bk_seq:04d}"
    st.seq += 1
    return f"P-{st.seq:04d}"


def _get(s, tenant: str, ref: str) -> CourierParcel:
    ref = str(ref or "").strip()
    if not ref:
        raise CourierError("Parcel batayiye (P-0001 / AWB / order number).")
    p = s.get(CourierParcel, ref)
    if p is None or p.tenant_id != tenant:
        p = s.scalars(select(CourierParcel).where(CourierParcel.tenant_id == tenant, or_(
            CourierParcel.number == ref.upper(), CourierParcel.awb == ref, CourierParcel.order_ref == ref))
                      .order_by(CourierParcel.created_at.desc())).first()
    if p is None:
        raise CourierError(f"Parcel '{ref}' nahi mila.")
    return p


def _problems(p: CourierParcel, st: CourierSettings) -> tuple[list[str], list[str]]:
    """(blocking, warnings) before a booking card can be made."""
    c = _j(p.customer, {})
    wh = _j(st.warehouse, {})
    w = _j(p.weight, {})
    block, warn = [], []
    if p.kind == "forward":
        if not c.get("address"):
            block.append("Graahak ka poora pata chahiye.")
        if not R.pincode_ok(c.get("pincode", "")):
            block.append("Graahak ka 6 ankon ka PIN code chahiye.")
        if len(R.phone10(c.get("phone", ""))) != 10:
            block.append("Graahak ka 10 ankon ka phone chahiye (courier wale call karte hain).")
    if not (wh.get("address") and R.pincode_ok(wh.get("pincode", "")) and R.phone10(wh.get("phone", ""))):
        block.append("Pickup ka pata / PIN / phone Courier > Niyam mein likhiye.")
    if not R.num(w.get("dead_kg")):
        block.append("Parcel ka wazan (kg) chahiye.")
    if p.payment == "cod" and p.cod_amount <= 0:
        block.append("COD parcel par graahak se kitna lena hai, woh rakam chahiye.")
    goods = R.goods_check(" ".join(i.get("name", "") for i in _j(p.items, [])) + " " + p.note)
    if goods["prohibited"]:
        block.append("Yeh courier se nahi jaata: " + ", ".join(goods["prohibited"]) + ".")
    if goods["restricted"]:
        warn.append("Courier ke niyam dekhiye: " + ", ".join(goods["restricted"]) + " (packing / declaration).")
    if p.kind == "forward":
        from ..india.money import STATE_CODES
        from ..thok.rules import eway_needed
        seller = (st.gstin or "")[:2]
        buyer = next((k for k, v in STATE_CODES.items() if v.lower() == str(c.get("state", "")).strip().lower()), "")
        need, limit = eway_needed(p.value, seller, buyer if buyer else "xx")
        if need and not p.eway_no:
            block.append(f"Keemat {_rs(p.value)} hai ({_rs(limit)} se zyada): pehle e-way bill banaiye (Rule 138) aur uska number likhiye.")
    sv = _j(p.serviceable, {})
    if sv and sv.get("ok") is False:
        block.append(f"{sv.get('provider', 'Courier')} is PIN par deliver nahi karta.")
    if sv and p.payment == "cod" and sv.get("cod") is False:
        block.append("Is PIN par COD nahi hota: prepaid kijiye ya doosra courier.")
    warn += [n for n in w.get("notes", []) if "wazan (kg)" not in n]
    return block, warn


def _view(s, p: CourierParcel, full: bool = True) -> dict[str, Any]:
    c = _j(p.customer, {})
    st = _st(s, p.tenant_id)
    block, warn = _problems(p, st)
    nd = _j(p.ndr, {})
    v = {"id": p.id, "number": p.number, "kind": p.kind, "order_ref": p.order_ref, "status": p.status,
         "status_label": R.STATE_LABEL.get(p.status, p.status), "raw": p.raw, "edd": p.edd, "provider": p.provider, "carrier": p.carrier,
         "awb": p.awb, "payment": p.payment, "cod_amount": p.cod_amount, "value": p.value, "weight": _j(p.weight, {}),
         "customer_name": c.get("name", ""), "city": c.get("city", "") or R.mask_address(c.get("address", "")), "pincode": c.get("pincode", ""),
         "phone_masked": R.mask_phone(c.get("phone", "")), "ndr": nd, "cod": _j(p.cod, {}), "wapsi_ref": p.wapsi_ref,
         "track_link": track_link(p.carrier if p.provider in ("", "manual") else p.provider, p.awb), "blocking": block, "warnings": warn,
         "checked_at": _utc(p.checked_at).isoformat() if p.checked_at else None,
         "created_at": _utc(p.created_at).isoformat(), "updated_at": _utc(p.updated_at).isoformat()}
    if full:
        v.update({"customer": c, "items": _j(p.items, []), "eway_no": p.eway_no, "serviceable": _j(p.serviceable, {}),
                  "quotes": _j(p.quotes, []), "pick": _j(p.pick, {}), "booking": _j(p.booking, {}), "label": _j(p.label, {}),
                  "events": _j(p.events, []), "note": p.note, "log": _j(p.log, [])[-20:],
                  "cards": [card_view(x) for x in s.scalars(select(CourierCard).where(CourierCard.parcel_id == p.id)
                                                              .order_by(CourierCard.created_at.desc()).limit(6)).all()]})
    return v


def track_link(provider: str, awb: str) -> str:
    k = R.carrier_key(provider) if provider not in R.CARRIERS else provider
    info = R.CARRIERS.get(k) or {}
    t = info.get("track", "")
    if not t or not awb:
        return t
    return t.replace("{awb}", awb) if "{awb}" in t else t


def parcel_view(tenant: str, ref: str) -> dict[str, Any]:
    with session() as s:
        return _view(s, _get(s, tenant, ref))


def parcels(tenant: str, status: str = "", q: str = "", limit: int = 100) -> list[dict[str, Any]]:
    with session() as s:
        qq = select(CourierParcel).where(CourierParcel.tenant_id == tenant)
        if status == "active":
            qq = qq.where(CourierParcel.status.in_(R.ACTIVE))
        elif status:
            qq = qq.where(CourierParcel.status == status)
        if q:
            like = f"%{q.strip()[:40]}%"
            qq = qq.where(or_(CourierParcel.number.ilike(like), CourierParcel.awb.ilike(like), CourierParcel.order_ref.ilike(like),
                              CourierParcel.customer.ilike(like), CourierParcel.phone.like(like)))
        rows = s.scalars(qq.order_by(CourierParcel.created_at.desc()).limit(min(limit, 300))).all()
        return [_view(s, p, full=False) for p in rows]


def new_parcel(tenant: str, *, order_ref: str = "", customer: dict | None = None, items=None, value=None, payment: str = "prepaid",
               cod_amount=None, weight_kg=None, dims: dict | None = None, eway_no: str = "", note: str = "", kind: str = "forward",
               wapsi_ref: str = "", conv: str = "") -> dict[str, Any]:
    pay = "cod" if str(payment or "").lower() in ("cod", "cash", "cash on delivery") else "prepaid"
    its = _clean_items(items)
    val = R.num(value) or round(sum(i["price"] * i["qty"] for i in its), 2)
    cod = R.num(cod_amount) if pay == "cod" else 0.0
    if pay == "cod" and not cod:
        cod = val
    d = {k: R.num((dims or {}).get(k) or (dims or {}).get({"l": "length", "b": "breadth", "h": "height"}[k])) for k in ("l", "b", "h")}
    with _LOCK, session() as s:
        st = _st(s, tenant)
        w = R.weights(R.num(weight_kg), d, st.divisor) | {"dims": d}
        p = CourierParcel(tenant_id=tenant, number=_next_number(s, tenant), kind=kind if kind in ("forward", "reverse") else "forward",
                          order_ref=_clean_str(order_ref, 80), customer=_dump(_clean_customer(customer)), items=_dump(its), value=val,
                          payment=pay, cod_amount=cod, weight=_dump(w), eway_no=_clean_str(eway_no, 20), note=_clean_str(note, 500),
                          wapsi_ref=_clean_str(wapsi_ref, 40))
        p.phone = R.phone10(_j(p.customer, {}).get("phone", ""))
        s.add(p)
        s.flush()
        _log(p, "Parcel bana")
        v = _view(s, p)
    _publish(tenant, "courier_parcel", {"parcel": v, "event": "new"}, conv)
    say = f"{v['number']} bana: {v['customer_name'] or 'graahak'} {v['pincode']}, {('COD ' + _rs(cod)) if pay == 'cod' else 'prepaid'}, " \
          f"chargeable {w['chargeable_kg']:g} kg."
    if v["blocking"]:
        say += " Pehle: " + " ".join(v["blocking"][:2])
    else:
        say += " Ab PIN aur rate check karta hoon (`check`)."
    return {"parcel": v, "say": say}


def update_parcel(tenant: str, ref: str, **kw: Any) -> dict[str, Any]:
    with _LOCK, session() as s:
        p = _get(s, tenant, ref)
        if p.status not in ("draft", "waiting", "failed"):
            raise CourierError(f"{p.number} book ho chuka: badlav courier ke panel mein (ya cancel ka card).")
        st = _st(s, tenant)
        if kw.get("customer") is not None:
            c = _j(p.customer, {})
            c.update({k: v for k, v in _clean_customer({**c, **kw["customer"]}).items()})
            p.customer = _dump(c)
            p.phone = R.phone10(c.get("phone", ""))
            p.serviceable, p.quotes, p.pick = "{}", "[]", "{}"
        if kw.get("items") is not None:
            p.items = _dump(_clean_items(kw["items"]))
        if kw.get("value") is not None:
            p.value = R.num(kw["value"])
        if kw.get("payment"):
            p.payment = "cod" if str(kw["payment"]).lower() == "cod" else "prepaid"
            if p.payment == "prepaid":
                p.cod_amount = 0.0
        if kw.get("cod_amount") is not None and p.payment == "cod":
            p.cod_amount = R.num(kw["cod_amount"])
        if kw.get("weight_kg") is not None or kw.get("dims") is not None:
            w = _j(p.weight, {})
            d = w.get("dims") or {}
            if kw.get("dims"):
                d = {k: R.num(kw["dims"].get(k) or kw["dims"].get({"l": "length", "b": "breadth", "h": "height"}[k]) or d.get(k)) for k in ("l", "b", "h")}
            dead = R.num(kw["weight_kg"]) if kw.get("weight_kg") is not None else R.num(w.get("dead_kg"))
            p.weight = _dump(R.weights(dead, d, st.divisor) | {"dims": d})
            p.quotes, p.pick = "[]", "{}"
        if kw.get("eway_no") is not None:
            p.eway_no = _clean_str(kw["eway_no"], 20)
        if kw.get("order_ref") is not None:
            p.order_ref = _clean_str(kw["order_ref"], 80)
        if kw.get("note") is not None:
            p.note = _clean_str(kw["note"], 500)
        stale = _stale_cards(s, p.id)
        if p.status == "waiting":
            p.status = "draft"
        _log(p, "Badlav" + (" (purana card radd)" if stale else ""))
        v = _view(s, p)
    return {"parcel": v, "say": f"{v['number']} badal diya" + (": purana booking card ab nahi chalega, naya banaiye." if stale else ".")}


def _stale_cards(s, parcel_id: str) -> int:
    n = 0
    for c in s.scalars(select(CourierCard).where(CourierCard.parcel_id == parcel_id, CourierCard.status == "waiting")).all():
        c.status = "stale"
        n += 1
    return n


# ---------------------------------------------------------------- serviceability + rates (read-only, her keys)
def check(tenant: str, ref: str, provider: str = "") -> dict[str, Any]:
    with session() as s:
        p = _get(s, tenant, ref)
        st = _st(s, tenant)
        wh = _j(st.warehouse, {})
        c = _j(p.customer, {})
        w = _j(p.weight, {})
        rules = R.clean_rules(_j(st.rules, {}))
        pid, cod, kg, value = p.id, p.payment == "cod", R.num(w.get("chargeable_kg")) or R.num(w.get("dead_kg")) or 0.5, p.value
        dst = c.get("pincode", "") if p.kind == "forward" else wh.get("pincode", "")
        src = wh.get("pincode", "") if p.kind == "forward" else c.get("pincode", "")
    if not (R.pincode_ok(src) and R.pincode_ok(dst)):
        raise CourierError("Dono PIN code chahiye (pickup aur graahak).")
    provs = [provider] if provider else _api_providers(tenant)
    if not provs:
        return {"parcel": parcel_view(tenant, pid), "manual": True,
                "say": "Koi courier API nahi juda: PIN aur rate apne courier app mein dekhiye (Courier > Niyam mein Shiprocket / Delhivery jod sakti hain)."}
    options, svc, errors = [], {"ok": False, "cod": None, "provider": ", ".join(provs), "at": now().isoformat()}, []
    for pv in provs:
        try:
            out = _adapter(tenant, pv).check(src, dst, kg, cod, value)
        except K.CarrierError as e:
            errors.append(f"{pv}: {e}")
            continue
        options += out.get("options") or []
        svc["ok"] = svc["ok"] or bool(out.get("serviceable"))
        if cod:
            svc["cod"] = bool(svc["cod"]) or bool(out.get("cod_ok"))
    if errors and not options and not svc["ok"]:
        raise CourierError("Courier se jawab nahi mila: " + "; ".join(errors)[:300])
    pick = R.choose(options, rules, cod, kg)
    with _LOCK, session() as s:
        p = _get(s, tenant, pid)
        p.serviceable = _dump(svc)
        p.quotes = _dump(options[:30])
        p.pick = _dump({"best": pick["best"], "why": pick["why"], "at": now().isoformat()})
        _log(p, f"PIN check: {'theek' if svc['ok'] else 'nahi'}; {len(options)} rate")
        v = _view(s, p)
    if not svc["ok"]:
        say = f"{v['number']}: {svc['provider']} is PIN ({dst}) par deliver nahi karta."
    elif cod and svc["cod"] is False:
        say = f"{v['number']}: PIN par delivery hoti hai par COD nahi."
    elif pick["best"]:
        b = pick["best"]
        say = (f"{v['number']}: sabse theek {b['courier']}" + (f" {_rs(b['rate'])}" if b.get("rate") else " (rate courier ne nahi bataya)") +
               (f", ~{b['days']} din" if b.get("days") else "") + f" ({pick['why']}). Book karne ke liye card bhejoon?")
    else:
        say = f"{v['number']}: {pick['why']}"
    if errors:
        say += " (" + "; ".join(errors)[:160] + ")"
    return {"parcel": v, "ranked": pick["ranked"], "best": pick["best"], "why": pick["why"], "say": say}


# ---------------------------------------------------------------- cards (only her tap decides)
def card_view(c: CourierCard) -> dict[str, Any]:
    return {"id": c.id, "number": c.number, "kind": c.kind, "parcel_id": c.parcel_id, "status": c.status, "hash": c.hash, "money": c.money,
            "body": _j(c.body, {}), "result": _j(c.result, {}), "created_at": _utc(c.created_at).isoformat(),
            "decided_at": _utc(c.decided_at).isoformat() if c.decided_at else None}


def _parcel_fingerprint(p: CourierParcel, st: CourierSettings) -> dict[str, Any]:
    c = _j(p.customer, {})
    w = _j(p.weight, {})
    return {"number": p.number, "kind": p.kind, "to": [c.get("name"), c.get("address"), c.get("pincode"), c.get("phone")],
            "from": [_j(st.warehouse, {}).get(k) for k in ("address", "pincode", "phone")], "payment": p.payment, "cod": p.cod_amount,
            "value": p.value, "kg": w.get("chargeable_kg"), "dims": w.get("dims"), "eway": p.eway_no, "items": _j(p.items, []),
            "wapsi": p.wapsi_ref, "awb": p.awb}


def _make_card(s, tenant: str, p: CourierParcel, kind: str, body: dict, money: bool, conv: str = "") -> CourierCard:
    st = _st(s, tenant)
    _stale_cards(s, p.id)
    body = {**body, "parcel_number": p.number, "kind": kind}
    h = R.card_hash({"body": body, "parcel": _parcel_fingerprint(p, st)})
    c = CourierCard(tenant_id=tenant, parcel_id=p.id, number=_next_number(s, tenant, card=True), kind=kind, body=_dump(body),
                    hash=h, money=money)
    s.add(c)
    s.flush()
    return c


def book_request(tenant: str, ref: str, courier: str = "", provider: str = "", conv: str = "") -> dict[str, Any]:
    """The booking card: carrier, rate, weight, address, COD. Booking happens only on her tap (it may charge her wallet)."""
    with _LOCK, session() as s:
        p = _get(s, tenant, ref)
        if p.kind != "forward":
            raise CourierError("Return label ke liye `return` action.")
        if p.status not in ("draft", "waiting", "failed"):
            raise CourierError(f"{p.number} pehle se {R.STATE_LABEL.get(p.status, p.status)}.")
        st = _st(s, tenant)
        block, warn = _problems(p, st)
        if block:
            raise CourierError("Book se pehle: " + " ".join(block))
        quotes = _j(p.quotes, [])
        best = _j(p.pick, {}).get("best")
        q = None
        if courier:
            ck = R.carrier_key(courier)
            q = next((o for o in quotes if R.carrier_key(o.get("courier", "")) == ck or str(o.get("courier_id")) == str(courier)), None)
            if q is None:
                raise CourierError(f"'{courier}' ka rate nahi mila: pehle `check` (rate courier se aate hain, Hoola nahi banata).")
        elif provider and provider in API_PROVIDERS:
            q = next((o for o in quotes if o.get("provider") == provider), None)
        else:
            q = best
        if q is None:
            if _api_providers(tenant):
                raise CourierError("Pehle PIN aur rate check kijiye (`check`): bina courier ke rate ke card nahi banta.")
            p.provider = "manual"
            p.status = "manual"
            _log(p, "Courier API nahi: aap book karein")
            v = _view(s, p)
            return {"parcel": v, "manual": _manual_steps(p, st), "say":f"{v['number']}: courier API nahi juda. Apne courier app mein book kijiye aur AWB batayiye."}
        if q.get("rate") in (None, 0):
            warn.append("Courier ne rate nahi bataya: asli kharcha unke panel / bill par dikhega.")
        rules = R.clean_rules(_j(st.rules, {}))
        if rules["max_rate"] and q.get("rate") and q["rate"] > rules["max_rate"]:
            warn.append(f"Rate {_rs(q['rate'])} aapki seema {_rs(rules['max_rate'])} se zyada.")
        c = _j(p.customer, {})
        w = _j(p.weight, {})
        body = {"title": f"{p.number} book karein?", "provider": q.get("provider"), "courier": q.get("courier"), "courier_id": q.get("courier_id"),
                "rate": q.get("rate"), "days": q.get("days"), "etd": q.get("etd", ""), "to": f"{c.get('name', '')}, {c.get('city', '') or R.mask_address(c.get('address', ''))} {c.get('pincode', '')}".strip(),
                "payment": p.payment, "cod_amount": p.cod_amount, "value": p.value, "chargeable_kg": w.get("chargeable_kg"),
                "dims": w.get("dims"), "eway_no": p.eway_no, "warnings": warn,
                "say": f"{q.get('courier')}" + (f" {_rs(q['rate'])}" if q.get("rate") else "") + " - aapke courier wallet se kat sakta hai."}
        card = _make_card(s, tenant, p, "book", body, money=True, conv=conv)
        p.status = "waiting"
        _log(p, f"Booking card {card.number}")
        cv = card_view(card)
        v = _view(s, p)
    _publish(tenant, "courier_card", {"card": cv, "parcel": v}, conv)
    return {"card": cv, "parcel": v, "say": f"{v['number']} ka booking card bheja ({body['say']}). Book tabhi hoga jab aap card par haan karengi.",
            "next": "Card uske phone par hai; woh khud tap karegi. Tum book nahi kar sakte, dobara mat poochho."}


def _manual_steps(p: CourierParcel, st: CourierSettings) -> list[str]:
    c = _j(p.customer, {})
    w = _j(p.weight, {})
    return [f"Courier app / counter mein naya shipment: {c.get('name', '')}, {c.get('address', '')} {c.get('pincode', '')}".strip(),
            f"Wazan {w.get('dead_kg', 0):g} kg, dabba {'x'.join(str(int(x)) for x in (w.get('dims') or {}).values() if x) or '-'} cm"
            f" (chargeable {w.get('chargeable_kg', 0):g} kg)",
            ("COD: " + _rs(p.cod_amount)) if p.payment == "cod" else "Prepaid",
            (f"E-way bill: {p.eway_no}" if p.eway_no else "E-way bill: keemat ₹50,000 se kam ho to zaroorat nahi"),
            "AWB / tracking number Hoola ko bataiye (`awb`)"]


def approval_card(tenant: str, card_id: str) -> dict[str, Any] | None:
    """What approvekey (kind 'courier') signs: the card as it is NOW."""
    with session() as s:
        c = s.get(CourierCard, card_id)
        if c is None or c.tenant_id != tenant or c.status != "waiting":
            return None
        b = _j(c.body, {})
        return {"title": b.get("title", ""), "payee": b.get("courier", "") or b.get("provider", ""),
                "amount": _rs(b["rate"]) if b.get("rate") else "", "hash": c.hash}


def cards(tenant: str, status: str = "waiting") -> list[dict[str, Any]]:
    with session() as s:
        q = select(CourierCard).where(CourierCard.tenant_id == tenant)
        if status:
            q = q.where(CourierCard.status == status)
        return [card_view(c) for c in s.scalars(q.order_by(CourierCard.created_at.desc()).limit(100)).all()]


def decide(tenant: str, card_id: str, yes: bool, card_hash: str, by_tap: bool = False, reason: str = "", conv: str = "") -> dict[str, Any]:
    """ONLY her tap in the app (the API passes by_tap=True). The hash must be the card she saw AND still match the parcel."""
    if not by_tap:
        raise CourierError("Yeh sirf owner app mein card par tap karke karti hain.")
    with _LOCK, session() as s:
        c = s.get(CourierCard, card_id)
        if c is None or c.tenant_id != tenant:
            raise CourierError("Card nahi mila.")
        if c.status != "waiting":
            raise CourierError(f"Card {c.number} ab {c.status} hai.")
        if not card_hash or card_hash != c.hash:
            raise CourierError("Card badal gaya: naya card dekhiye.")
        p = s.get(CourierParcel, c.parcel_id) if c.parcel_id else None
        if p is None:
            raise CourierError("Parcel nahi mila.")
        st = _st(s, tenant)
        body = {k: v for k, v in _j(c.body, {}).items()}
        stale = R.card_hash({"body": body, "parcel": _parcel_fingerprint(p, st)}) != c.hash
        if stale:
            c.status = "stale"
            kind = "stale"
        elif not yes:
            c.decided_at = now()
            c.status = "rejected"
            c.result = _dump({"reason": _clean_str(reason, 200)})
            if p.status == "waiting":
                p.status = "draft"
            _log(p, f"Card {c.number}: aapne mana kiya")
            cv, v = card_view(c), _view(s, p)
            kind = ""
        else:
            c.decided_at = now()
            c.status = "doing"
            kind, pid = c.kind, p.id
    if kind == "stale":
        raise CourierError("Parcel card ke baad badla: naya card banaiye.")
    if not kind:
        _publish(tenant, "courier_card", {"card": cv, "parcel": v}, conv)
        return {"card": cv, "parcel": v, "say": f"{cv['number']} roka. Kuch book nahi hua."}
    try:
        if kind == "book":
            res = _do_book(tenant, pid, body)
        elif kind == "reverse":
            res = _do_reverse(tenant, pid, body)
        elif kind == "ndr":
            res = _do_ndr(tenant, pid, body)
        elif kind == "cancel":
            res = _do_cancel(tenant, pid)
        else:
            raise CourierError("card galat")
        status, say = "done", res.get("say", "Ho gaya.")
    except (CourierError, K.CarrierError) as e:
        res, status, say = {"error": str(e)}, "failed", f"Courier ne mana kiya: {e}"
        with _LOCK, session() as s:
            p = s.get(CourierParcel, pid)
            if kind in ("book", "reverse") and p.status in ("waiting", "draft"):
                p.status = "failed"
            _log(p, f"Card {kind}: {str(e)[:120]}")
    with _LOCK, session() as s:
        c = s.get(CourierCard, card_id)
        c.status = status
        c.result = _dump({k: v for k, v in res.items() if k != "parcel"})
        cv = card_view(c)
        v = _view(s, s.get(CourierParcel, pid))
    _publish(tenant, "courier_card", {"card": cv, "parcel": v}, conv)
    _publish(tenant, "courier_parcel", {"parcel": v, "event": kind}, conv)
    return {"card": cv, "parcel": v, "say": say, **({"file": res["file"]} if res.get("file") else {})}


def _store_label(tenant: str, pid: str, number: str, pdf: bytes, conv: str = "") -> dict[str, Any]:
    from .. import chatfiles
    path = _dir(tenant) / f"{pid}.pdf"
    path.write_bytes(pdf)
    try:
        f = chatfiles.store(tenant, f"Label-{number}.pdf", pdf, "application/pdf")
    except Exception as e:                       # the copy in her courier folder is kept anyway
        log.warning("courier label store: %s", e)
        f = {}
    if f:
        _publish(tenant, "pc_file", {**f, "source": "courier", "title": f"Label {number}"}, conv)
    return {"path": str(path), **({"file_id": f.get("file_id") or f.get("id"), "name": f.get("name")} if f else {})}


def _do_book(tenant: str, pid: str, body: dict) -> dict[str, Any]:
    with session() as s:
        p = s.get(CourierParcel, pid)
        st = _st(s, tenant)
        parcel = {"number": p.number, "customer": _j(p.customer, {}), "items": _j(p.items, []), "value": p.value, "payment": p.payment,
                  "cod_amount": p.cod_amount, "weight": _j(p.weight, {}), "eway_no": p.eway_no}
        wh = _j(st.warehouse, {})
    prov = body.get("provider") or ""
    ad = _adapter(tenant, prov)
    if isinstance(ad, K.Manual):
        raise CourierError(f"{prov or 'Courier'} ki chaabi nahi mili: Courier > Niyam mein jodiye.")
    res = ad.book(parcel, wh, {"courier_id": body.get("courier_id"), "courier": body.get("courier")})
    label = {}
    lab_note = ""
    if res.get("awb"):
        try:
            lu = ad.label(res).get("url", "")
            res["label_url"] = lu
            if lu:
                label = _store_label(tenant, pid, parcel["number"], K.fetch_pdf(lu))
        except K.CarrierError as e:
            lab_note = f" Label abhi nahi mila ({e}); courier panel se download kijiye."
    with _LOCK, session() as s:
        p = s.get(CourierParcel, pid)
        p.provider = prov
        p.carrier = _clean_str(res.get("carrier") or body.get("courier"), 40)
        p.awb = _clean_str(res.get("awb"), 40)
        p.booking = _dump({k: v for k, v in res.items() if k not in ("label",)} | {"booked_at": now().isoformat(), "quoted": body.get("rate")})
        p.label = _dump(label)
        p.status = "booked" if p.awb else "failed"
        p.raw = "Booked" if p.awb else ""
        _log(p, f"Book: {p.carrier} AWB {p.awb or '-'}" + (f" (kharcha {_rs(res['charged'])})" if res.get("charged") else ""))
        number, awb, carrier = p.number, p.awb, p.carrier
    if not awb:
        raise CourierError(res.get("note") or "AWB nahi mila.")
    pk = res.get("pickup") or {}
    say = f"{number} book: {carrier}, AWB {awb}." + (f" Pickup {pk['date']}." if pk.get("date") else "") + \
          (" " + res["pickup_note"] if res.get("pickup_note") else "") + (" Label aapki files mein." if label.get("file_id") else "") + lab_note
    return {"awb": awb, "carrier": carrier, "pickup": pk, "label": label, "say": say}


# ---------------------------------------------------------------- manual AWB, tracking
def set_awb(tenant: str, ref: str, awb: str, carrier: str = "") -> dict[str, Any]:
    if not R.awb_ok(awb):
        raise CourierError("AWB / tracking number sahi likhiye (6-30 akshar / ank).")
    with _LOCK, session() as s:
        p = _get(s, tenant, ref)
        if p.status in ("delivered", "cancelled"):
            raise CourierError(f"{p.number} {R.STATE_LABEL[p.status]}.")
        dup = s.scalars(select(CourierParcel).where(CourierParcel.tenant_id == tenant, CourierParcel.awb == awb.strip(),
                                                     CourierParcel.id != p.id)).first()
        if dup:
            raise CourierError(f"AWB {awb} pehle se {dup.number} par hai.")
        _stale_cards(s, p.id)
        ck = carrier if carrier in R.CARRIERS else R.carrier_key(carrier)
        p.awb, p.carrier = awb.strip(), _clean_str(carrier, 40) or p.carrier
        if not p.provider or p.provider == "manual":
            p.provider = ck if ck in ("delhivery", "bluedart") else (p.provider or "manual")
        if p.status in ("draft", "waiting", "manual", "failed"):
            p.status = "booked"
        _log(p, f"AWB {p.awb} ({p.carrier or 'courier'}) aapne likha")
        v = _view(s, p)
    return {"parcel": v, "say": f"{v['number']}: AWB {awb} likh liya." + (" Tracking Hoola dekhta rahega." if v["provider"] in K.ADAPTERS and v["provider"] != "manual" else
                                                                       " Tracking courier ke page par: " + (v["track_link"] or "unka app") + ".")}


def track(tenant: str, ref: str, conv: str = "") -> dict[str, Any]:
    with session() as s:
        p = _get(s, tenant, ref)
        pid, prov, booking, awb = p.id, p.provider, _j(p.booking, {}), p.awb
    if not awb:
        raise CourierError("Abhi AWB nahi: pehle book kijiye ya AWB likhiye.")
    ad = _adapter(tenant, prov)
    if isinstance(ad, K.Manual):
        v = parcel_view(tenant, pid)
        return {"parcel": v, "say": f"{v['number']}: tracking courier ke page par dekhiye ({v['track_link'] or 'unka app'}), AWB {awb}. "
                                    "Wahan jo likha ho mujhe bataiye to main yahan likh doon."}
    t = ad.track({**booking, "awb": awb})
    return _apply_track(tenant, pid, t, conv)


def _apply_track(tenant: str, pid: str, t: dict, conv: str = "") -> dict[str, Any]:
    raw = _clean_str(t.get("raw"), 200)
    new = R.state_of(raw) if raw else None
    with _LOCK, session() as s:
        p = s.get(CourierParcel, pid)
        old, old_raw = p.status, p.raw
        if t.get("events"):
            p.events = _dump(t["events"][:15])
        if t.get("edd"):
            p.edd = _clean_str(t["edd"], 30)
        p.checked_at = now()
        if raw:
            p.raw = raw
        changed = False
        if new and new != old and not (old == "delivered" and new != "rto"):
            p.status = new
            changed = True
            _log(p, f"{R.STATE_LABEL.get(new, new)}: {raw}")
            if new == "delivered":
                p.delivered_at = now()
            if new == "exception":
                nd = _j(p.ndr, {})
                reason = R.ndr_reason(raw + " " + " ".join(e.get("what", "") for e in (t.get("events") or [])[:2]))
                nd.update({"reason": reason, "label": R.NDR_REASONS.get(reason, {}).get("label", "Delivery nahi hui"), "raw": raw,
                           "attempts": int(nd.get("attempts") or 0) + 1, "at": now().isoformat(), "answered": False})
                p.ndr = _dump(nd)
        elif raw and raw != old_raw:
            _log(p, f"Courier: {raw}")
        v = _view(s, p)
    if changed and v["status"] in ("out_for_delivery", "delivered", "exception", "rto", "picked"):
        _publish(tenant, "courier_parcel", {"parcel": v, "event": v["status"]}, conv)
    say = f"{v['number']}: {raw or 'courier ne abhi kuch naya nahi bataya'}" + (f" (ummeed: {v['edd']})" if v["edd"] else "") + "."
    if raw and not new:
        say += " (Yeh shabd main kisi haalat se nahi jod paya; jo courier ne likha wahi bataya.)"
    if v["status"] == "exception":
        say += f" Delivery nahi hui: {v['ndr'].get('label', '')}. Graahak ke liye message taiyaar (`ndr`)."
    return {"parcel": v, "changed": changed, "say": say}


# ---------------------------------------------------------------- NDR
def ndr_message(tenant: str, ref: str) -> dict[str, Any]:
    with session() as s:
        p = _get(s, tenant, ref)
        st = _st(s, tenant)
        nd = _j(p.ndr, {})
        c = _j(p.customer, {})
        reason = nd.get("reason") or "other"
        ask = R.NDR_REASONS.get(reason, {}).get("ask", "Kab aur kahan delivery theek rahegi? Bataiye.")
        biz = st.business or "hamari dukaan"
        text = (f"Namaste {c.get('name', '')}, {biz} se. Aapka parcel ({p.order_ref or p.number}, AWB {p.awb}) courier ne deliver nahi kar paya"
                f"{(' - ' + nd.get('label', '')) if nd.get('label') else ''}. {ask} Hum courier ko dobara bhejne ko bolenge. "
                f"Dhyaan rahe: hum kabhi OTP / PIN nahi maangte.")
        phone = R.phone10(c.get("phone", ""))
        v = _view(s, p)
    from urllib.parse import quote
    return {"parcel": v, "text": text, "whatsapp": f"https://wa.me/91{phone}?text={quote(text)}" if phone else "",
            "sms": f"sms:+91{phone}?body={quote(text)}" if phone else "", "note": R.MAX_ATTEMPTS_NOTE,
            "say": "Graahak ke liye message taiyaar (aap bhejiye). Jawab aaye to dobara bhejne ka card bana doonga."}


def ndr_request(tenant: str, ref: str, action: str = "reattempt", phone: str = "", address: str = "", when: str = "", note: str = "",
                conv: str = "", by: str = "owner") -> dict[str, Any]:
    action = "rto" if action in ("rto", "return", "wapas") else "reattempt"
    if when and not re.fullmatch(r"\d{4}-\d{2}-\d{2}", when.strip()):
        raise CourierError("Tareekh YYYY-MM-DD mein.")
    with _LOCK, session() as s:
        p = _get(s, tenant, ref)
        if p.status != "exception":
            raise CourierError(f"{p.number} par abhi koi NDR nahi ({R.STATE_LABEL.get(p.status, p.status)}).")
        fix = {"phone": R.phone10(phone) if phone else "", "address": _clean_str(address, 180), "date": when.strip(), "note": _clean_str(note, 200)}
        if when:
            try:
                if date.fromisoformat(when.strip()) < today(tenant):
                    raise CourierError("Beeti tareekh nahi.")
            except ValueError:
                raise CourierError("Tareekh galat.")
        nd = _j(p.ndr, {})
        body = {"title": f"{p.number}: " + ("dobara delivery" if action == "reattempt" else "wapas mangwao (RTO)"), "action": action, "fix": fix,
                "provider": p.provider, "awb": p.awb, "reason": nd.get("label", ""), "by": by,
                "say": ("Courier ko dobara bhejne ko bolein" if action == "reattempt" else "Parcel wapas mangwayein (RTO kharcha lag sakta hai)")}
        card = _make_card(s, tenant, p, "ndr", body, money=action == "rto", conv=conv)
        nd.update({"answered": True, "answer": {"action": action, **{k: v for k, v in fix.items() if k != "phone"}, "by": by}})
        p.ndr = _dump(nd)
        _log(p, f"NDR card {card.number} ({action})")
        cv, v = card_view(card), _view(s, p)
    _publish(tenant, "courier_card", {"card": cv, "parcel": v}, conv)
    return {"card": cv, "parcel": v, "say": f"{body['title']}: card bheja, courier ko tabhi bolenge jab aap tap karengi."}


def _do_ndr(tenant: str, pid: str, body: dict) -> dict[str, Any]:
    with session() as s:
        p = s.get(CourierParcel, pid)
        prov, booking, awb, number = p.provider, _j(p.booking, {}), p.awb, p.number
    ad = _adapter(tenant, prov)
    if isinstance(ad, K.Manual):
        raise CourierError("Courier API nahi juda: NDR ka jawab courier ke panel mein dijiye.")
    ad.ndr({**booking, "awb": awb}, body.get("action", "reattempt"), body.get("fix") or {})
    with _LOCK, session() as s:
        p = s.get(CourierParcel, pid)
        nd = _j(p.ndr, {})
        nd["sent"] = {"action": body.get("action"), "at": now().isoformat()}
        p.ndr = _dump(nd)
        _log(p, f"NDR: courier ko bola ({body.get('action')})")
    return {"say": f"{number}: courier ko " + ("dobara delivery ke liye bol diya." if body.get("action") == "reattempt" else "wapas bhejne ko bol diya.")}


# ---------------------------------------------------------------- cancel
def cancel_request(tenant: str, ref: str, conv: str = "") -> dict[str, Any]:
    with _LOCK, session() as s:
        p = _get(s, tenant, ref)
        if p.status in ("draft", "waiting", "failed", "manual"):
            _stale_cards(s, p.id)
            p.status = "cancelled"
            _log(p, "Cancel (book nahi hua tha)")
            v = _view(s, p)
            return {"parcel": v, "say": f"{v['number']} cancel (kuch book nahi hua tha)."}
        if p.status not in ("booked",):
            raise CourierError(f"{p.number} {R.STATE_LABEL.get(p.status, p.status)}: ab cancel courier se baat karke hi.")
        body = {"title": f"{p.number} ki booking cancel karein?", "provider": p.provider, "awb": p.awb, "carrier": p.carrier,
                "say": "Booking radd; wallet ka paisa courier ke niyam se wapas."}
        card = _make_card(s, tenant, p, "cancel", body, money=False, conv=conv)
        cv, v = card_view(card), _view(s, p)
    _publish(tenant, "courier_card", {"card": cv, "parcel": v}, conv)
    return {"card": cv, "parcel": v, "say": f"{v['number']}: cancel ka card bheja (aapke tap par hi)."}


def _do_cancel(tenant: str, pid: str) -> dict[str, Any]:
    with session() as s:
        p = s.get(CourierParcel, pid)
        prov, booking, awb, number = p.provider, _j(p.booking, {}), p.awb, p.number
    ad = _adapter(tenant, prov)
    if isinstance(ad, K.Manual):
        raise CourierError("Courier API nahi juda: cancel courier ke panel mein kijiye.")
    ad.cancel({**booking, "awb": awb})
    with _LOCK, session() as s:
        p = s.get(CourierParcel, pid)
        p.status = "cancelled"
        _log(p, "Booking cancel (aapke tap se)")
    return {"say": f"{number} ki booking cancel."}


# ---------------------------------------------------------------- return labels (reverse pickup)
def return_request(tenant: str, wapsi_ref: str = "", ref: str = "", customer: dict | None = None, items=None, weight_kg=None,
                   dims: dict | None = None, reason: str = "", conv: str = "") -> dict[str, Any]:
    """A return label on her tap. For a Wapsi return the Wapsi desk's own approved return is booked (its rules: only an
    approved return gets a pickup); otherwise a reverse parcel from a delivered parcel / the customer's details."""
    if wapsi_ref:
        from ..wapsi import service as ws
        try:
            r = ws.return_view(tenant, wapsi_ref)
        except ws.WapsiError as e:
            raise CourierError(str(e))
        if r.get("status") not in ("approved", "pickup_manual"):
            raise CourierError(f"Wapsi {r.get('number')} abhi {r.get('status_label', r.get('status'))}: return pehle Wapsi mein manzoor kijiye.")
        cust = r.get("customer") or {}
        made = new_parcel(tenant, order_ref=r.get("order_ref") or r.get("number", ""), customer=cust, items=r.get("items") or [],
                          value=sum(R.num(i.get("price")) * (R.num(i.get("qty")) or 1) for i in r.get("items") or []),
                          weight_kg=weight_kg or 0.5, dims=dims, kind="reverse", wapsi_ref=r.get("id", ""), note=f"Wapsi {r.get('number')}",
                          conv=conv)
        pid = made["parcel"]["id"]
    else:
        if ref:
            src = parcel_view(tenant, ref)
            if src["kind"] != "forward" or src["status"] not in ("delivered",):
                raise CourierError(f"{src['number']} abhi delivered nahi: return label delivered parcel ka hi.")
            cust, its = src["customer"], src["items"]
            w = src["weight"]
            made = new_parcel(tenant, order_ref=src["order_ref"] or src["number"], customer=cust, items=its, value=src["value"],
                              weight_kg=weight_kg or w.get("dead_kg"), dims=dims or w.get("dims"), kind="reverse",
                              note=f"Return of {src['number']}: {reason}"[:300], conv=conv)
        else:
            made = new_parcel(tenant, customer=customer, items=items, weight_kg=weight_kg, dims=dims, kind="reverse",
                              note=f"Return: {reason}"[:300], conv=conv)
        pid = made["parcel"]["id"]
    with _LOCK, session() as s:
        p = s.get(CourierParcel, pid)
        st = _st(s, tenant)
        c = _j(p.customer, {})
        block, warn = _problems(p, st)
        if not (c.get("address") and R.pincode_ok(c.get("pincode", "")) and R.phone10(c.get("phone", ""))):
            block.append("Graahak ka pata, PIN aur phone chahiye (wahan se pickup hoga).")
        if block:
            v = _view(s, p)
            return {"parcel": v, "say": f"{p.number} (return) bana. Pehle: " + " ".join(block)}
        prov = "wapsi" if p.wapsi_ref else (st.default_carrier if st.default_carrier in API_PROVIDERS else (_api_providers(tenant) or ["manual"])[0])
        body = {"title": f"{p.number}: return label / reverse pickup", "provider": prov, "from": f"{c.get('name', '')}, {c.get('pincode', '')}",
                "to": _j(st.warehouse, {}).get("pincode", ""), "wapsi": p.wapsi_ref, "warnings": warn,
                "say": "Graahak se saamaan wapas uthwana (courier wallet se kat sakta hai)" + (" - Wapsi desk ke zariye" if p.wapsi_ref else "")}
        card = _make_card(s, tenant, p, "reverse", body, money=prov != "manual", conv=conv)
        p.status = "waiting"
        cv, v = card_view(card), _view(s, p)
    _publish(tenant, "courier_card", {"card": cv, "parcel": v}, conv)
    return {"card": cv, "parcel": v, "say": f"{v['number']} ka return label card bheja; pickup tabhi book hoga jab aap tap karengi.",
            "next": "Card uske phone par hai; tum book nahi kar sakte."}


def _do_reverse(tenant: str, pid: str, body: dict) -> dict[str, Any]:
    with session() as s:
        p = s.get(CourierParcel, pid)
        st = _st(s, tenant)
        number, wref = p.number, p.wapsi_ref
        cust, items, w = _j(p.customer, {}), _j(p.items, []), _j(p.weight, {})
        wh = _j(st.warehouse, {})
    if wref:
        from ..wapsi import service as ws
        d = w.get("dims") or {}
        try:
            out = ws.book_pickup(tenant, wref, {"weight": w.get("dead_kg"), "length": d.get("l"), "breadth": d.get("b"), "height": d.get("h")})
        except ws.WapsiError as e:
            raise CourierError(str(e))
        res = out.get("pickup") or {}
    else:
        prov = body.get("provider") or "manual"
        ad = _adapter(tenant, prov)
        d = w.get("dims") or {}
        res = ad.reverse(number, cust, wh, items, {"weight": w.get("dead_kg"), "length": d.get("l"), "breadth": d.get("b"), "height": d.get("h")},
                         qc=False)
    with _LOCK, session() as s:
        p = s.get(CourierParcel, pid)
        p.provider = res.get("provider") or body.get("provider") or "manual"
        p.awb = _clean_str(res.get("awb"), 40)
        p.carrier = _clean_str(res.get("courier") or p.provider, 40)
        p.booking = _dump(res | {"booked_at": now().isoformat()})
        p.status = "booked" if p.awb else ("manual" if p.provider == "manual" else "booked")
        _log(p, "Return pickup: " + (f"AWB {p.awb}" if p.awb else ("aap book karein" if p.provider == "manual" else "AWB abhi nahi")))
        awb, prov = p.awb, p.provider
    if prov == "manual":
        return {"steps": res.get("steps", []), "say": f"{number}: courier API nahi, reverse pickup aap courier app mein book kijiye (steps card mein)."}
    return {"awb": awb, "say": f"{number}: return pickup book" + (f", AWB {awb}" if awb else " (AWB courier dega)") + "."}


# ---------------------------------------------------------------- where is my parcel?
def _answer(v: dict[str, Any]) -> str:
    bits = [f"{v['number']}" + (f" ({v['order_ref']})" if v["order_ref"] else "") + f" {v['customer_name']}".rstrip()]
    if not v["awb"]:
        bits.append(R.STATE_LABEL.get(v["status"], v["status"]))
    else:
        bits.append(f"{v['carrier'] or v['provider']} AWB {v['awb']}: {v['raw'] or R.STATE_LABEL.get(v['status'], v['status'])}")
        if v["edd"] and v["status"] not in ("delivered", "rto", "cancelled"):
            bits.append(f"courier ki ummeed {v['edd']}")
        if v["checked_at"]:
            bits.append("dekha " + v["checked_at"][:16].replace("T", " ") + " UTC")
    return " - ".join(bits)


def where(tenant: str, q: str = "", refresh: bool = True) -> dict[str, Any]:
    q = str(q or "").strip()
    rows = parcels(tenant, status="" if q else "active", q=q, limit=5) if q else parcels(tenant, status="active", limit=10)
    if not rows and q and R.phone10(q):
        rows = parcels(tenant, q=R.phone10(q), limit=5)
    if refresh:
        out = []
        for v in rows[:5]:
            if v["awb"] and v["provider"] in K.ADAPTERS and v["provider"] != "manual" and v["status"] not in ("delivered", "cancelled"):
                try:
                    v = track(tenant, v["id"])["parcel"]
                except (CourierError, K.CarrierError) as e:
                    v = {**v, "track_error": str(e)[:120]}
            out.append(v)
        rows = out + rows[5:]
    if not rows:
        return {"parcels": [], "say": "Aisa koi parcel nahi mila." if q else "Abhi koi parcel raste mein nahi."}
    return {"parcels": rows, "say": "\n".join(_answer(v) for v in rows[:5])}


_TO_LOGI = {"booked": "booked", "manual": "booked", "picked": "picked", "in_transit": "in_transit",
            "out_for_delivery": "out_for_delivery", "delivered": "delivered", "rto": "rto", "exception": "failed",
            "cancelled": "cancelled", "failed": "failed"}


def track_awb(tenant: str, courier: str = "", awb: str = "") -> dict[str, Any]:
    """For the Logistics desk: {status, events: [{at, what, where}], raw, track_link} in the courier's own words.
    Her parcel with that AWB first (tracked through her own carrier key when she has one), else her key for that courier
    directly. Nothing known and no key -> CourierError naming the courier's own tracking page."""
    awb = str(awb or "").strip()
    if not awb:
        raise CourierError("AWB nahi diya.")
    with session() as s:
        p = s.scalars(select(CourierParcel).where(CourierParcel.tenant_id == tenant, CourierParcel.awb == awb)
                      .order_by(CourierParcel.created_at.desc())).first()
        pid = p.id if p is not None else ""
    if pid:
        v = parcel_view(tenant, pid)
        if v["provider"] in K.ADAPTERS and v["provider"] != "manual" and v["status"] not in ("delivered", "cancelled"):
            v = track(tenant, pid)["parcel"]
        return {"status": _TO_LOGI.get(v["status"], ""), "raw": v.get("raw", ""),
                "events": [{"at": e.get("at", ""), "what": e.get("what", ""), "where": e.get("where", "")} for e in (v.get("events") or [])],
                "track_link": v.get("track_link", "")}
    prov = R.carrier_key(courier) if courier else ""
    ad = _adapter(tenant, prov) if prov else K.Manual()
    link = track_link(prov or courier, awb)
    if isinstance(ad, K.Manual):
        raise CourierError(f"Is courier ki chaabi Courier desk mein nahi: AWB {awb} " + (f"yahan dekhiye: {link}" if link else "courier ke app / site par dekhiye") + ".")
    t = ad.track({"awb": awb})
    raw = _clean_str(t.get("raw"), 200)
    st = R.state_of(raw) if raw else None
    return {"status": _TO_LOGI.get(st or "", ""), "raw": raw, "track_link": link,
            "events": [{"at": e.get("at", ""), "what": e.get("what", ""), "where": e.get("where", "")} for e in (t.get("events") or [])[:15]]}


def customer_parcels(tenant: str, wa_id: str, ref: str = "") -> list[dict[str, Any]]:
    """THIS customer's own parcels (by their WhatsApp number): no address, no other customer, no internal notes."""
    ph = R.phone10(wa_id)
    if len(ph) != 10:
        return []
    with session() as s:
        q = select(CourierParcel).where(CourierParcel.tenant_id == tenant, CourierParcel.phone == ph,
                                        CourierParcel.status.notin_(("draft", "waiting", "failed")))
        rows = s.scalars(q.order_by(CourierParcel.created_at.desc()).limit(10)).all()
        out = []
        for p in rows:
            if ref and ref.strip().upper() not in (p.number, p.awb.upper(), p.order_ref.upper()):
                continue
            out.append({"number": p.number, "order_ref": p.order_ref, "kind": "return" if p.kind == "reverse" else "order",
                        "status": R.STATE_LABEL.get(p.status, p.status), "courier_said": p.raw, "courier": p.carrier,
                        "awb": p.awb, "expected": p.edd if p.status not in ("delivered", "rto", "cancelled") else "",
                        "last_checked": _utc(p.checked_at).isoformat()[:16] if p.checked_at else "",
                        "cod_to_pay": _rs(p.cod_amount) if p.payment == "cod" and p.kind == "forward" and p.status not in ("delivered", "rto", "cancelled") else "",
                        "track_link": track_link(p.carrier if p.provider in ("", "manual") else p.provider, p.awb)})
        return out


def customer_fix(tenant: str, wa_id: str, ref: str, when: str = "", address: str = "", phone: str = "", note: str = "") -> dict[str, Any]:
    """THIS customer answers a failed delivery: recorded on their own parcel; her card is prepared, nothing is sent."""
    ph = R.phone10(wa_id)
    rows = customer_parcels(tenant, wa_id, ref)
    if not rows:
        raise CourierError("Is number par yeh parcel nahi mila.")
    with session() as s:
        p = s.scalars(select(CourierParcel).where(CourierParcel.tenant_id == tenant, CourierParcel.phone == ph,
                                                  CourierParcel.number == rows[0]["number"])).first()
        if p is None or p.status != "exception":
            return {"say": "Is parcel ki delivery mein abhi koi dikkat darj nahi; courier jaldi pahunchayega."}
        pid = p.id
    out = ndr_request(tenant, pid, "reattempt", phone=phone, address=address, when=when, note=f"Graahak: {note}"[:200], by="customer")
    return {"say": "Shukriya, aapki baat dukaan tak pahuncha di. Courier ko dobara bhejne ka faisla dukaan karegi.",
            "number": out["parcel"]["number"]}


def message(tenant: str, ref: str, kind: str = "shipped") -> dict[str, Any]:
    if kind == "ndr":
        return ndr_message(tenant, ref)
    with session() as s:
        p = _get(s, tenant, ref)
        st = _st(s, tenant)
        c = _j(p.customer, {})
        biz = st.business or "hamari dukaan"
        link = track_link(p.carrier if p.provider in ("", "manual") else p.provider, p.awb)
        if kind == "delivered":
            text = f"Namaste {c.get('name', '')}, {biz} se. Courier ke hisaab se aapka parcel ({p.order_ref or p.number}) pahunch gaya. Kuch dikkat ho to batayiye."
        else:
            if not p.awb:
                raise CourierError("Abhi AWB nahi: book hone ke baad message banega.")
            text = (f"Namaste {c.get('name', '')}, {biz} se. Aapka order {p.order_ref or p.number} {p.carrier or 'courier'} se bhej diya. "
                    f"AWB {p.awb}." + (f" Tracking: {link}" if link else "") +
                    (f" Delivery par {_rs(p.cod_amount)} dena hai." if p.payment == "cod" else "") + " Hum kabhi OTP / PIN nahi maangte.")
        phone = R.phone10(c.get("phone", ""))
    from urllib.parse import quote
    return {"text": text, "whatsapp": f"https://wa.me/91{phone}?text={quote(text)}" if phone else "",
            "say": "Message taiyaar (aap bhejiye)."}


# ---------------------------------------------------------------- COD remittance
def _rows_of(data: bytes, name: str) -> list[list[str]]:
    n = (name or "").lower()
    if n.endswith((".xlsx", ".xlsm")):
        import openpyxl
        wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        ws = wb.active
        return [[("" if c is None else str(c)) for c in row] for row in ws.iter_rows(values_only=True)][:5000]
    text = data.decode("utf-8-sig", errors="replace")
    return [r for r in csv.reader(io.StringIO(text))][:5000]


def remit_import(tenant: str, data: bytes, name: str = "remittance.csv") -> dict[str, Any]:
    rows = _rows_of(data, name)
    head_i, cols = -1, {}
    for i, r in enumerate(rows[:15]):
        c = R.remit_header(r)
        if "awb" in c and "amount" in c:
            head_i, cols = i, c
            break
    if head_i < 0:
        raise CourierError("Sheet mein AWB aur amount ke column nahi mile.")
    added, matched, unknown, dup = 0, 0, [], 0
    with _LOCK, session() as s:
        for r in rows[head_i + 1:]:
            def g(k):
                i = cols.get(k)
                return str(r[i]).strip() if i is not None and i < len(r) and r[i] is not None else ""
            awb = re.sub(r"\s", "", g("awb"))
            amt = R.num(g("amount"))
            if not awb or not R.awb_ok(awb):
                continue
            utr = _clean_str(g("utr"), 60)
            if s.scalars(select(CourierRemit).where(CourierRemit.tenant_id == tenant, CourierRemit.awb == awb, CourierRemit.utr == utr,
                                                     CourierRemit.amount == amt)).first():
                dup += 1
                continue
            p = s.scalars(select(CourierParcel).where(CourierParcel.tenant_id == tenant, CourierParcel.awb == awb)).first()
            row = CourierRemit(tenant_id=tenant, awb=awb, order_ref=_clean_str(g("order"), 80), amount=amt, utr=utr,
                               date=_clean_str(g("date"), 20), source=_clean_str(name, 120), parcel_id=p.id if p else "")
            s.add(row)
            added += 1
            if p is None:
                unknown.append(awb)
                continue
            matched += 1
            cod = _j(p.cod, {})
            got = round(R.num(cod.get("amount")) + amt, 2)
            cod.update({"amount": got, "utr": utr or cod.get("utr", ""), "date": row.date or cod.get("date", ""),
                        "remitted": got >= p.cod_amount - 0.5 if p.payment == "cod" else True,
                        "diff": round(got - p.cod_amount, 2) if p.payment == "cod" else 0.0})
            p.cod = _dump(cod)
            _log(p, f"COD remittance {_rs(amt)}" + (f" UTR {utr}" if utr else ""))
    rep = cod_report(tenant)
    say = f"Remittance sheet: {added} line, {matched} parcel se mili" + (f", {len(unknown)} AWB hamare record mein nahi" if unknown else "") + \
          (f", {dup} pehle se thi" if dup else "") + f". Ab baaki COD: {rep['pending_count']} parcel {_rs(rep['pending_amount'])}."
    if rep["mismatch"]:
        say += f" {len(rep['mismatch'])} mein rakam alag."
    return {"added": added, "matched": matched, "unknown": unknown[:30], "duplicates": dup, "report": rep, "say": say}


def cod_report(tenant: str) -> dict[str, Any]:
    with session() as s:
        st = _st(s, tenant)
        days = st.remit_days
        rows = s.scalars(select(CourierParcel).where(CourierParcel.tenant_id == tenant, CourierParcel.payment == "cod",
                                                     CourierParcel.kind == "forward", CourierParcel.status == "delivered")).all()
        pending, mismatch, done = [], [], []
        for p in rows:
            cod = _j(p.cod, {})
            age = (now() - _utc(p.delivered_at)).days if p.delivered_at else None
            line = {"number": p.number, "awb": p.awb, "carrier": p.carrier, "cod_amount": p.cod_amount, "remitted": R.num(cod.get("amount")),
                    "utr": cod.get("utr", ""), "days_since_delivery": age, "late": bool(days and age is not None and age > days)}
            if cod.get("remitted"):
                if abs(R.num(cod.get("diff"))) >= 0.5:
                    mismatch.append(line | {"diff": cod.get("diff")})
                else:
                    done.append(line)
            elif R.num(cod.get("amount")) > 0:
                mismatch.append(line | {"diff": round(R.num(cod.get("amount")) - p.cod_amount, 2)})
            else:
                pending.append(line)
        unknown = [{"awb": r.awb, "amount": r.amount, "utr": r.utr} for r in
                   s.scalars(select(CourierRemit).where(CourierRemit.tenant_id == tenant, CourierRemit.parcel_id == "").limit(50)).all()]
    pending.sort(key=lambda x: -(x["days_since_delivery"] or 0))
    return {"pending": pending[:100], "pending_count": len(pending), "pending_amount": round(sum(x["cod_amount"] for x in pending), 2),
            "late": [x for x in pending if x["late"]], "mismatch": mismatch[:100], "remitted_count": len(done),
            "remitted_amount": round(sum(x["remitted"] for x in done), 2), "unknown": unknown, "remit_days": days,
            "note": "COD ka paisa courier aapke khaate mein bhejta hai; Hoola sirf sheet milata hai." +
                    ("" if days else " 'Kitne din mein aana chahiye' Niyam mein likhiye to der wale alag dikhenge.")}


# ---------------------------------------------------------------- manifest + day
def manifest(tenant: str, on: str = "") -> dict[str, Any]:
    d = date.fromisoformat(on) if on else today(tenant)
    from ..life import tz_of
    tz = tz_of(tenant)
    with session() as s:
        rows = s.scalars(select(CourierParcel).where(CourierParcel.tenant_id == tenant, CourierParcel.status.in_(("booked", "manual")))
                         .order_by(CourierParcel.created_at)).all()
        out = []
        for p in rows:
            bk = _j(p.booking, {})
            booked = bk.get("booked_at", "")
            try:
                bd = datetime.fromisoformat(booked).astimezone(tz).date() if booked else _utc(p.created_at).astimezone(tz).date()
            except ValueError:
                bd = d
            if bd > d:
                continue
            c = _j(p.customer, {})
            w = _j(p.weight, {})
            out.append({"number": p.number, "kind": p.kind, "order_ref": p.order_ref, "awb": p.awb, "carrier": p.carrier or p.provider,
                        "to": f"{c.get('name', '')} {c.get('pincode', '')}".strip(), "kg": w.get("dead_kg"), "chargeable_kg": w.get("chargeable_kg"),
                        "payment": p.payment, "cod_amount": p.cod_amount if p.payment == "cod" else 0.0, "pickup": (bk.get("pickup") or {}).get("date", "")})
    by = {}
    for x in out:
        by.setdefault(x["carrier"] or "courier", []).append(x)
    return {"date": d.isoformat(), "parcels": out, "by_carrier": {k: len(v) for k, v in by.items()},
            "cod_total": round(sum(x["cod_amount"] for x in out), 2), "say": f"Aaj ke pickup: {len(out)} parcel" +
            (" (" + ", ".join(f"{k} {len(v)}" for k, v in by.items()) + ")" if by else "") + "."}


def manifest_file(tenant: str, on: str = "", conv: str = "") -> dict[str, Any]:
    import openpyxl
    from .. import chatfiles
    m = manifest(tenant, on)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Manifest"
    with session() as s:
        st = _st(s, tenant)
        ws.append([f"{st.business or 'Pickup'} - manifest {m['date']}"])
    ws.append(["Parcel", "Order", "Courier", "AWB", "To", "Wazan kg", "Chargeable kg", "Payment", "COD ₹", "Pickup", "Courier boy sign"])
    for x in m["parcels"]:
        ws.append([x["number"], x["order_ref"], x["carrier"], x["awb"], x["to"], x["kg"], x["chargeable_kg"], x["payment"], x["cod_amount"], x["pickup"], ""])
    ws.append([])
    ws.append(["Kul", len(m["parcels"]), "", "", "", "", "", "COD", m["cod_total"]])
    bio = io.BytesIO()
    wb.save(bio)
    f = chatfiles.store(tenant, f"Manifest_{m['date']}.xlsx", bio.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    _publish(tenant, "pc_file", {**f, "source": "courier", "title": f"Manifest {m['date']}"}, conv)
    return {"file": f, "manifest": m, "say": m["say"] + " Excel aapki files mein (courier boy se sign karwa lijiye)."}


def day(tenant: str) -> dict[str, Any]:
    with session() as s:
        q = select(CourierParcel).where(CourierParcel.tenant_id == tenant)
        rows = s.scalars(q.where(CourierParcel.status.in_(R.ACTIVE + ("draft", "waiting", "failed")))).all()
        by = {}
        for p in rows:
            by.setdefault(p.status, []).append(p.number)
        waiting = s.scalars(select(CourierCard).where(CourierCard.tenant_id == tenant, CourierCard.status == "waiting")).all()
    rep = cod_report(tenant)
    parts = []
    if waiting:
        parts.append(f"{len(waiting)} card aapki haan ke liye")
    if by.get("exception"):
        parts.append(f"Delivery nahi hui (NDR): {', '.join(by['exception'][:5])} - graahak se baat karke jawab dijiye")
    if by.get("rto"):
        parts.append(f"Wapas aa rahe (RTO): {', '.join(by['rto'][:5])}")
    if by.get("out_for_delivery"):
        parts.append(f"Aaj delivery: {len(by['out_for_delivery'])}")
    if by.get("booked") or by.get("manual"):
        parts.append(f"Pickup baaki: {len(by.get('booked', [])) + len(by.get('manual', []))}")
    if by.get("draft") or by.get("failed"):
        parts.append(f"Book karne baaki: {len(by.get('draft', [])) + len(by.get('failed', []))}")
    if rep["late"]:
        parts.append(f"COD der se: {len(rep['late'])} parcel {_rs(sum(x['cod_amount'] for x in rep['late']))}")
    elif rep["pending_count"]:
        parts.append(f"COD aana baaki: {rep['pending_count']} parcel {_rs(rep['pending_amount'])}")
    if rep["mismatch"]:
        parts.append(f"COD rakam alag: {len(rep['mismatch'])}")
    return {"parts": parts, "counts": {k: len(v) for k, v in by.items()}, "waiting": len(waiting),
            "text": "Courier: " + ("; ".join(parts) if parts else "sab theek") + "."}


def home(tenant: str) -> dict[str, Any]:
    return {"day": day(tenant), "settings": settings(tenant), "parcels": parcels(tenant, limit=40), "cards": cards(tenant, "waiting"),
            "cod": cod_report(tenant), "law": R.LAW, "never": list(R.NEVER)}


def forget(tenant: str, ref: str) -> dict[str, Any]:
    """DPDP erasure: a finished parcel's personal details go; the AWB, amounts and dates stay for her accounts."""
    with _LOCK, session() as s:
        p = _get(s, tenant, ref)
        if p.status not in ("delivered", "cancelled", "rto", "failed", "draft"):
            raise CourierError(f"{p.number} abhi chal raha hai; poora hone ke baad mitaiye.")
        p.customer = _dump({"name": "", "phone": "", "email": "", "address": "", "city": "", "state": "", "pincode": _j(p.customer, {}).get("pincode", "")})
        p.phone = ""
        p.ndr = "{}"
        p.note = ""
        _log(p, "Graahak ki jaankari mita di (DPDP)")
        number = p.number
    return {"say": f"{number}: graahak ka naam / pata / phone mita diya; AWB aur rakam hisaab ke liye bachi."}


def prompt_block(tenant: str) -> str:
    try:
        d = day(tenant)
    except Exception:
        return ""
    return "" if not d["parts"] else "COURIER aaj: " + "; ".join(d["parts"][:4])


# ---------------------------------------------------------------- worker
def _hm(s: str) -> tuple[int, int]:
    try:
        h, m = s.split(":")
        return int(h), int(m)
    except (ValueError, AttributeError):
        return 9, 30


def tick(at: datetime | None = None) -> dict[str, Any]:
    """Every minute from the worker: tracking poll every POLL_MIN for parcels on her API carriers, morning note, manifest."""
    at = at or now()
    out = {"polled": 0, "notes": 0}
    with session() as s:
        tenants = [r.tenant_id for r in s.scalars(select(CourierSettings)).all()]
    from ..life import tz_of
    for t in tenants:
        try:
            with session() as s:
                st = _st(s, t)
                last = _utc(st.last_poll)
                due_poll = last is None or at - last >= timedelta(minutes=POLL_MIN)
                if due_poll:
                    st.last_poll = at
                ids = [p.id for p in s.scalars(select(CourierParcel).where(CourierParcel.tenant_id == t, CourierParcel.status.in_(
                    ("booked", "picked", "in_transit", "out_for_delivery", "exception", "rto")), CourierParcel.awb != "")).all()] if due_poll else []
                loc = at.astimezone(tz_of(t))
                ds = loc.date().isoformat()
                morning = st.last_morning != ds and (loc.hour, loc.minute) >= _hm(st.morning_time)
                man = st.last_manifest != ds and (loc.hour, loc.minute) >= _hm(st.manifest_time)
                if morning:
                    st.last_morning = ds
                if man:
                    st.last_manifest = ds
            for pid in ids[:200]:
                try:
                    with session() as s:
                        p = s.get(CourierParcel, pid)
                        prov = p.provider
                    if prov in K.ADAPTERS and prov != "manual" and _creds(t, prov):
                        track(t, pid)
                        out["polled"] += 1
                except (CourierError, K.CarrierError) as e:
                    log.info("courier poll %s: %s", pid, e)
            if morning:
                d = day(t)
                if d["parts"]:
                    _publish(t, "courier_day", {"text": d["text"], "parts": d["parts"]})
                    out["notes"] += 1
            if man:
                m = manifest(t, loc.date().isoformat())
                if m["parcels"]:
                    _publish(t, "courier_day", {"text": m["say"] + " Manifest Courier > Aaj mein.", "manifest": True})
                    out["notes"] += 1
        except Exception:
            log.exception("courier tick %s", t)
    return out
