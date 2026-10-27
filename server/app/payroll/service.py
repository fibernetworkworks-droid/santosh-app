"""v0.94 Payroll Officer: Hoola runs the monthly payroll of a business end to end.

  1. Gets the job: "October ki salary banao", "Ramesh 2 din chhutti, Suresh 6 ghante OT", an attendance sheet (Excel /
     CSV) or a photo of the register, a new joiner's details, "Sunita ne naukri chhodi".
  2. Works it out, shown step by step: earned pay after loss-of-pay days, overtime at twice the rate, Code on Wages
     "wages" (50% rule), EPF on the Rs 25,000 ceiling (Sept 2026 split), EPS only for eligible members, EDLI + admin,
     ESIC for the whole contribution period, professional tax by state, TDS under Sec 392 spread over the year (new or
     old regime), advances recovered, net pay.
  3. Checked against the rules and earlier months: minimum wage, deductions above 50% of wages, ESIC coverage period,
     PF / ESI numbers missing, pay jumping vs last month, bank details missing, salary by the 7th, dues of a leaver in
     2 working days.
  4. Ready as files and drafts: payslips (PDF), salary register (Excel), the bank bulk-payment file, the EPFO ECR text
     file, the ESIC upload sheet, PT summary, payslip messages for each employee (she sends).
  5. Waits for her Haan: paying salaries. Hoola prepares the run, she approves; she uploads the bank file and pays.
     Hoola never pays, never moves money, never logs in to her bank.
  6. Delivers and learns: paid runs are kept as the year-to-date base for TDS and the next month; names she uses for
     people ("Ramu" = Ramesh Patil) are remembered; the due-date calendar (PF / ESI 15th, TDS 7th, Form 138, ESIC
     return, Form 130, bonus) runs every morning.
"""
from __future__ import annotations

import csv
import io
import json
import logging
import re
import threading
from datetime import date, datetime, timedelta
from typing import Any

from sqlalchemy import Boolean, DateTime, Float, Integer, String, Text, select
from sqlalchemy.orm import Mapped, mapped_column

from ..db import Base, new_id, now, session
from . import rules as R

log = logging.getLogger("santosh.payroll")
_LOCK = threading.Lock()
COMPONENTS = ("basic", "da", "hra", "conveyance", "special", "other")
STATUS_LABEL = {"draft": "Ban raha", "ready": "Aapki haan baaki", "approved": "Haan mili: bank file taiyaar", "paid": "Salary di gayi",
                "rejected": "Mana"}


class PayError(Exception):
    pass


class PaySettings(Base):
    __tablename__ = "pay_settings"
    tenant_id: Mapped[str] = mapped_column(String(64), primary_key=True)
    enabled: Mapped[bool] = mapped_column(Boolean, default=True)
    company: Mapped[str] = mapped_column(String(160), default="")
    state: Mapped[str] = mapped_column(String(2), default="27")
    pf_on: Mapped[bool] = mapped_column(Boolean, default=False)              # registered with EPFO
    pf_code: Mapped[str] = mapped_column(String(30), default="")
    esi_on: Mapped[bool] = mapped_column(Boolean, default=False)
    esi_code: Mapped[str] = mapped_column(String(30), default="")
    tan: Mapped[str] = mapped_column(String(12), default="")
    min_wage: Mapped[float] = mapped_column(Float, default=0.0)              # monthly minimum wage she follows (state notification)
    days_basis: Mapped[str] = mapped_column(String(8), default="calendar")   # calendar | 26 | 30
    pay_account: Mapped[str] = mapped_column(Text, default="")               # her debit account (sealed) for the bank file
    morning_time: Mapped[str] = mapped_column(String(5), default="09:45")
    last_morning: Mapped[str] = mapped_column(String(10), default="")


class PayEmployee(Base):
    __tablename__ = "pay_employees"
    id: Mapped[str] = mapped_column(String(40), primary_key=True, default=lambda: new_id("pe"))
    tenant_id: Mapped[str] = mapped_column(String(64), index=True)
    code: Mapped[str] = mapped_column(String(12), default="")
    name: Mapped[str] = mapped_column(String(120), default="")
    aliases: Mapped[str] = mapped_column(Text, default="[]")
    phone: Mapped[str] = mapped_column(String(20), default="")
    gender: Mapped[str] = mapped_column(String(1), default="")
    dob: Mapped[str] = mapped_column(String(10), default="")
    doj: Mapped[str] = mapped_column(String(10), default="")
    dol: Mapped[str] = mapped_column(String(10), default="")
    designation: Mapped[str] = mapped_column(String(80), default="")
    pay: Mapped[str] = mapped_column(Text, default="{}")                     # monthly {basic, da, hra, conveyance, special, other}
    uan: Mapped[str] = mapped_column(String(12), default="")
    esic_ip: Mapped[str] = mapped_column(String(12), default="")
    pan: Mapped[str] = mapped_column(Text, default="")                       # sealed
    bank: Mapped[str] = mapped_column(Text, default="")                      # sealed {"account", "ifsc"}
    regime: Mapped[str] = mapped_column(String(3), default="new")
    declared: Mapped[float] = mapped_column(Float, default=0.0)              # old regime deductions for the year (80C, 80D, HRA...)
    prev_income: Mapped[float] = mapped_column(Float, default=0.0)           # this year from a previous employer
    prev_tds: Mapped[float] = mapped_column(Float, default=0.0)
    pf_member: Mapped[bool] = mapped_column(Boolean, default=True)
    eps_member: Mapped[bool] = mapped_column(Boolean, default=True)
    pf_on_full: Mapped[bool] = mapped_column(Boolean, default=False)
    esi_state: Mapped[str] = mapped_column(Text, default="{}")              # {"period": "2026-10", "covered": true}
    fixed_term: Mapped[bool] = mapped_column(Boolean, default=False)
    active: Mapped[bool] = mapped_column(Boolean, default=True)
    created_at: Mapped[datetime] = mapped_column(DateTime(timezone=True), default=now)


class PayAttendance(Base):
    __tablename__ = "pay_attendance"
    id: Mapped[str] = mapped_column(String(40), primary_key=True, default=lambda: new_id("pa"))
    tenant_id: Mapped[str] = mapped_column(String(64), index=True)
    employee_id: Mapped[str] = mapped_column(String(40), index=True)
    period: Mapped[str] = mapped_column(String(7), index=True)
    lop_days: Mapped[float] = mapped_column(Float, default=0.0)
    ot_hours: Mapped[float] = mapped_column(Float, default=0.0)
    extra: Mapped[float] = mapped_column(Float, default=0.0)                 # incentive / arrears this month (added to pay)
    deduct: Mapped[float] = mapped_column(Float, default=0.0)                # other deduction this month (canteen, damage...)
    note: Mapped[str] = mapped_column(String(200), default="")


class PayAdvance(Base):
    __tablename__ = "pay_advances"
    id: Mapped[str] = mapped_column(String(40), primary_key=True, default=lambda: new_id("pv"))
    tenant_id: Mapped[str] = mapped_column(String(64), index=True)
    employee_id: Mapped[str] = mapped_column(String(40), index=True)
    given_on: Mapped[str] = mapped_column(String(10), default="")
    amount: Mapped[float] = mapped_column(Float, default=0.0)
    emi: Mapped[float] = mapped_column(Float, default=0.0)
    balance: Mapped[float] = mapped_column(Float, default=0.0)
    note: Mapped[str] = mapped_column(String(200), default="")


class PayRun(Base):
    __tablename__ = "pay_runs"
    id: Mapped[str] = mapped_column(String(40), primary_key=True, default=lambda: new_id("pr"))
    tenant_id: Mapped[str] = mapped_column(String(64), index=True)
    period: Mapped[str] = mapped_column(String(7), index=True)
    kind: Mapped[str] = mapped_column(String(8), default="monthly")          # monthly | fnf
    status: Mapped[str] = mapped_column(String(10), default="draft")
    lines: Mapped[str] = mapped_column(Text, default="[]")
    totals: Mapped[str] = mapped_column(Text, default="{}")
    checks: Mapped[str] = mapped_column(Text, default="[]")
    approved_at: Mapped[datetime | None] = mapped_column(DateTime(timezone=True), nullable=True)
    paid_on: Mapped[str] = mapped_column(String(10), default="")
    log: Mapped[str] = mapped_column(Text, default="[]")
    created_at: Mapped[datetime] = mapped_column(DateTime(timezone=True), default=now)
    updated_at: Mapped[datetime] = mapped_column(DateTime(timezone=True), default=now, onupdate=now)


# ---------------------------------------------------------------- helpers
def _j(s: str, default):
    try:
        v = json.loads(s or "")
        return v if isinstance(v, type(default)) else default
    except (TypeError, ValueError):
        return default


def _f(v) -> float:
    if v is None or v == "":
        return 0.0
    if isinstance(v, (int, float)):
        return float(v)
    s = re.sub(r"[^\d.\-]", "", str(v).replace(",", ""))
    try:
        return float(s) if s not in ("", "-", ".") else 0.0
    except ValueError:
        return 0.0


def _rs(v) -> str:
    v = round(float(v or 0))
    s = str(abs(v))
    if len(s) > 3:
        head = re.sub(r"(\d)(?=(\d\d)+$)", r"\1,", s[:-3])
        s = f"{head},{s[-3:]}"
    return ("-" if v < 0 else "") + "₹" + s


def _date(v) -> date | None:
    if not v:
        return None
    if isinstance(v, date):
        return v
    s = str(v).strip()
    for pat, order in ((r"(\d{4})-(\d{1,2})-(\d{1,2})", "ymd"), (r"(\d{1,2})[-/.](\d{1,2})[-/.](\d{4})", "dmy")):
        m = re.fullmatch(pat, s)
        if m:
            a, b, c = (int(x) for x in m.groups())
            try:
                return date(a, b, c) if order == "ymd" else date(c, b, a)
            except ValueError:
                return None
    return None


def _seal(v: str) -> str:
    from .. import idvault
    return idvault._seal(v) if v else ""


def _open(v: str) -> str:
    from .. import idvault
    try:
        return idvault._open(v) if v else ""
    except Exception:
        return ""


def _mask(v: str, keep: int = 4) -> str:
    return ("X" * max(0, len(v) - keep) + v[-keep:]) if v else ""


def _tz(tenant: str):
    from ..life import tz_of
    return tz_of(tenant)


def today(tenant: str) -> date:
    return datetime.now(_tz(tenant)).date()


def _publish(tenant: str, kind: str, payload: dict, conv: str = "") -> None:
    from ..events import bus
    try:
        if not conv:
            from ..net.service import home_conv
            conv = home_conv(tenant)
        bus.publish(conv, kind, payload)
    except Exception:
        log.exception("payroll publish failed")


def _runlog(r: PayRun, what: str, by: str = "hoola") -> None:
    lg = _j(r.log, [])
    lg.append({"at": now().isoformat(timespec="seconds"), "what": what[:300], "by": by})
    r.log = json.dumps(lg[-100:], ensure_ascii=False)


# ---------------------------------------------------------------- settings
def _st(s, tenant: str) -> PaySettings:
    r = s.get(PaySettings, tenant)
    if not r:
        r = PaySettings(tenant_id=tenant)
        s.add(r)
        s.flush()
    return r


def settings(tenant: str) -> dict[str, Any]:
    with session() as s:
        r = _st(s, tenant)
        acct = _j(_open(r.pay_account), {}) if r.pay_account else {}
        return {"enabled": r.enabled, "company": r.company, "state": r.state, "state_name": R.PT_STATES.get(r.state, r.state),
                "pf_on": r.pf_on, "pf_code": r.pf_code, "esi_on": r.esi_on, "esi_code": r.esi_code, "tan": r.tan,
                "min_wage": r.min_wage, "days_basis": r.days_basis, "pay_account": _mask(acct.get("account", "")),
                "morning_time": r.morning_time}


def save_settings(tenant: str, **kw: Any) -> dict[str, Any]:
    with session() as s:
        r = _st(s, tenant)
        for k in ("company", "pf_code", "esi_code", "tan"):
            if kw.get(k) is not None:
                setattr(r, k, str(kw[k]).strip()[:160])
        if kw.get("state"):
            sc = R.state_code(kw["state"])
            if not sc:
                raise PayError("State samajh nahi aaya (jaise Maharashtra / 27).")
            r.state = sc
        for k in ("pf_on", "esi_on", "enabled"):
            if k in kw and kw[k] is not None:
                setattr(r, k, bool(kw[k]))
        if kw.get("min_wage") not in (None, ""):
            r.min_wage = _f(kw["min_wage"])
        if kw.get("days_basis"):
            if str(kw["days_basis"]) not in ("calendar", "26", "30"):
                raise PayError("days_basis: calendar | 26 | 30")
            r.days_basis = str(kw["days_basis"])
        if kw.get("pay_account"):
            a = re.sub(r"\s", "", str(kw["pay_account"]))
            if not re.fullmatch(r"\d{9,18}", a):
                raise PayError("Bank account number 9-18 ank ka hota hai.")
            r.pay_account = _seal(json.dumps({"account": a, "ifsc": str(kw.get("pay_ifsc") or "").upper()}))
        if kw.get("morning_time"):
            r.morning_time = str(kw["morning_time"]).zfill(5)[:5]
    return settings(tenant)


# ---------------------------------------------------------------- employees
def _emp_view(e: PayEmployee) -> dict[str, Any]:
    bank = _j(_open(e.bank), {}) if e.bank else {}
    pan = _open(e.pan) if e.pan else ""
    pay = _j(e.pay, {})
    return {"id": e.id, "code": e.code, "name": e.name, "aliases": _j(e.aliases, []), "phone": e.phone, "gender": e.gender,
            "dob": e.dob, "doj": e.doj, "dol": e.dol, "designation": e.designation, "pay": pay,
            "monthly_gross": round(sum(_f(pay.get(k)) for k in COMPONENTS)), "uan": e.uan, "esic_ip": e.esic_ip,
            "pan": _mask(pan), "bank": {"account": _mask(bank.get("account", "")), "ifsc": bank.get("ifsc", "")},
            "regime": e.regime, "declared": e.declared, "pf_member": e.pf_member, "eps_member": e.eps_member,
            "pf_on_full": e.pf_on_full, "fixed_term": e.fixed_term, "active": e.active}


def _emp(s, tenant: str, ref: str) -> PayEmployee:
    ref = str(ref or "").strip()
    if not ref:
        raise PayError("Kaun sa employee? Naam ya code bataiye.")
    e = s.get(PayEmployee, ref)
    if e and e.tenant_id == tenant:
        return e
    rows = s.scalars(select(PayEmployee).where(PayEmployee.tenant_id == tenant)).all()
    low = ref.lower()
    for e in rows:
        if e.code.lower() == low or e.name.lower() == low or low in [a.lower() for a in _j(e.aliases, [])]:
            return e
    hits = [e for e in rows if low in e.name.lower() and e.active]
    if len(hits) == 1:
        return hits[0]
    if len(hits) > 1:
        raise PayError(f"'{ref}' naam ke {len(hits)} log: " + ", ".join(f"{h.name} ({h.code})" for h in hits[:6]) + ". Code bataiye.")
    raise PayError(f"Employee '{ref}' nahi mila.")


def save_employee(tenant: str, name: str = "", employee: str = "", phone: str = "", gender: str = "", dob: str = "", doj: str = "",
                  designation: str = "", pay: dict | None = None, gross=None, uan: str = "", esic_ip: str = "", pan: str = "",
                  account: str = "", ifsc: str = "", regime: str = "", declared=None, prev_income=None, prev_tds=None,
                  pf_member: bool | None = None, eps_member: bool | None = None, pf_on_full: bool | None = None,
                  fixed_term: bool | None = None, alias: str = "") -> dict[str, Any]:
    """Add or update an employee. PAN and bank account are kept sealed; Hoola shows only the last 4."""
    with _LOCK, session() as s:
        e = _emp(s, tenant, employee) if employee else None
        new = e is None
        if new:
            if not str(name).strip():
                raise PayError("Naam chahiye.")
            n = len(s.scalars(select(PayEmployee.id).where(PayEmployee.tenant_id == tenant)).all()) + 1
            e = PayEmployee(tenant_id=tenant, code=f"E-{n:03d}")
            s.add(e)
        if name:
            e.name = re.sub(r"\s+", " ", str(name)).strip()[:120]
        if phone:
            e.phone = re.sub(r"[^\d+]", "", str(phone))[:20]
        if gender:
            e.gender = str(gender).strip()[:1].upper() if str(gender).strip()[:1].upper() in ("M", "F", "T") else \
                ("F" if str(gender).lower() in ("mahila", "aurat", "female", "stree", "ladies") else "M")
        for k, v in (("dob", dob), ("doj", doj)):
            if v:
                d = _date(v)
                if not d:
                    raise PayError(f"{k} date samajh nahi aayi: {v}")
                setattr(e, k, d.isoformat())
        if designation:
            e.designation = str(designation)[:80]
        p = _j(e.pay, {})
        if pay:
            p |= {k: _f(v) for k, v in pay.items() if k in COMPONENTS}
        if gross not in (None, "") and not pay:
            g = _f(gross)
            p = {"basic": round(g * 0.5), "da": 0, "hra": round(g * 0.2), "conveyance": 0, "special": g - round(g * 0.5) - round(g * 0.2), "other": 0}
        e.pay = json.dumps(p)
        if uan:
            u = re.sub(r"\D", "", str(uan))
            if len(u) != 12:
                raise PayError("UAN 12 ank ka hota hai.")
            e.uan = u
        if esic_ip:
            ip = re.sub(r"\D", "", str(esic_ip))
            if len(ip) != 10:
                raise PayError("ESIC IP number 10 ank ka hota hai.")
            e.esic_ip = ip
        if pan:
            pv = str(pan).strip().upper()
            if not re.fullmatch(r"[A-Z]{5}\d{4}[A-Z]", pv):
                raise PayError("PAN galat (jaise ABCDE1234F).")
            e.pan = _seal(pv)
        if account:
            a = re.sub(r"\s", "", str(account))
            if not re.fullmatch(r"\d{9,18}", a):
                raise PayError("Bank account number 9-18 ank ka hota hai.")
            f = str(ifsc or _j(_open(e.bank), {}).get("ifsc", "")).strip().upper()
            if not re.fullmatch(r"[A-Z]{4}0[A-Z0-9]{6}", f):
                raise PayError("IFSC galat (jaise SBIN0001234).")
            e.bank = _seal(json.dumps({"account": a, "ifsc": f}))
        if regime in ("new", "old"):
            e.regime = regime
        for k, v in (("declared", declared), ("prev_income", prev_income), ("prev_tds", prev_tds)):
            if v not in (None, ""):
                setattr(e, k, _f(v))
        for k, v in (("pf_member", pf_member), ("eps_member", eps_member), ("pf_on_full", pf_on_full), ("fixed_term", fixed_term)):
            if v is not None:
                setattr(e, k, bool(v))
        if new and pf_member is None:
            gross_now = sum(_f(p.get(k)) for k in COMPONENTS)
            wage = R.code_wages(_f(p.get("basic")) + _f(p.get("da")), gross_now)
            # above the ceiling at joining: not compulsorily covered, and never EPS
            e.pf_member = wage <= 25000
            e.eps_member = e.pf_member and (not e.dob or _age(e.dob, date.today()) < 58)
        if alias:
            al = _j(e.aliases, [])
            if alias.strip() and alias.strip().lower() not in [a.lower() for a in al]:
                al.append(alias.strip()[:40])
            e.aliases = json.dumps(al[-8:])
        s.flush()
        v = _emp_view(e)
    say = f"{'Naya employee' if new else 'Update'}: {v['name']} ({v['code']}), mahine ka {_rs(v['monthly_gross'])}."
    miss = [x for x, ok in (("bank account", v["bank"]["account"]), ("UAN", v["uan"] or not v["pf_member"]), ("PAN", v["pan"])) if not ok]
    if miss:
        say += " Abhi chahiye: " + ", ".join(miss) + "."
    return {"employee": v, "say": say}


def _age(dob: str, on: date) -> int:
    d = _date(dob)
    if not d:
        return 0
    return on.year - d.year - ((on.month, on.day) < (d.month, d.day))


def employees(tenant: str, active: bool = True) -> list[dict[str, Any]]:
    with session() as s:
        q = select(PayEmployee).where(PayEmployee.tenant_id == tenant)
        if active:
            q = q.where(PayEmployee.active.is_(True))
        return [_emp_view(e) for e in s.scalars(q.order_by(PayEmployee.code)).all()]


# ---------------------------------------------------------------- attendance, advances
def set_attendance(tenant: str, period: str, entries: list[dict], by: str = "owner") -> dict[str, Any]:
    """entries: [{employee, lop_days, ot_hours, extra, deduct, note}] (only what she said; the rest stay as they are)."""
    p = R.norm_period(period)
    done, errs = [], []
    with _LOCK, session() as s:
        for x in entries or []:
            try:
                e = _emp(s, tenant, str(x.get("employee") or x.get("name") or ""))
            except PayError as ex:
                errs.append(str(ex))
                continue
            a = s.scalars(select(PayAttendance).where(PayAttendance.tenant_id == tenant, PayAttendance.employee_id == e.id,
                                                      PayAttendance.period == p)).first()
            if not a:
                a = PayAttendance(tenant_id=tenant, employee_id=e.id, period=p)
                s.add(a)
            for k in ("lop_days", "ot_hours", "extra", "deduct"):
                if x.get(k) not in (None, ""):
                    setattr(a, k, max(0.0, _f(x[k])))
            if a.lop_days > R.days_in(p):
                errs.append(f"{e.name}: {a.lop_days:g} din chhutti mahine se zyada.")
                a.lop_days = float(R.days_in(p))
            if x.get("note"):
                a.note = str(x["note"])[:200]
            done.append(f"{e.name}: " + ", ".join(t for t in (f"{a.lop_days:g} din bina pagaar" if a.lop_days else "",
                                                             f"{a.ot_hours:g} ghante OT" if a.ot_hours else "",
                                                             f"+{_rs(a.extra)}" if a.extra else "", f"-{_rs(a.deduct)}" if a.deduct else "") if t))
    say = (f"{R.nice(p)} attendance: " + "; ".join(done)) if done else "Kuch save nahi hua."
    if errs:
        say += " Dhyaan: " + " ".join(errs)
    return {"saved": len(done), "errors": errs, "say": say}


_ATT_HEAD = {"employee": ("employee", "name", "emp name", "employee name", "naam", "code", "emp code", "employee code"),
             "lop_days": ("lop", "lop days", "absent", "absent days", "leave without pay", "chhutti", "lwp"),
             "present": ("present", "present days", "days present", "paid days", "payable days", "hazri"),
             "ot_hours": ("ot", "ot hours", "overtime", "overtime hours"),
             "extra": ("incentive", "bonus", "arrears", "extra"), "deduct": ("deduction", "canteen", "fine", "other deduction")}


def import_attendance(tenant: str, period: str, data: bytes, name: str = "attendance.csv") -> dict[str, Any]:
    p = R.norm_period(period)
    rows: list[list] = []
    if name.lower().endswith((".xlsx", ".xlsm")):
        import openpyxl
        try:
            wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        except Exception as e:
            raise PayError(f"Excel nahi khula: {e}")
        rows = [list(r) for r in wb.worksheets[0].iter_rows(values_only=True)]
    else:
        rows = list(csv.reader(io.StringIO(data.decode("utf-8-sig", errors="replace"))))
    keymap: dict[int, str] = {}
    start = 0
    for i, r in enumerate(rows[:10]):
        km = {}
        for j, h in enumerate(r):
            hh = re.sub(r"\s+", " ", str(h or "").strip().lower())
            for k, names in _ATT_HEAD.items():
                if hh in names and k not in km.values():
                    km[j] = k
        if "employee" in km.values() and len(km) >= 2:
            keymap, start = km, i + 1
            break
    if not keymap:
        raise PayError("Sheet mein 'Name' aur 'Present / LOP / OT' wale column chahiye.")
    entries = []
    for r in rows[start:]:
        d = {keymap[j]: r[j] for j in keymap if j < len(r) and r[j] not in (None, "")}
        if not d.get("employee"):
            continue
        if "present" in d and "lop_days" not in d:
            d["lop_days"] = max(0.0, R.days_in(p) - _f(d.pop("present")))
        d.pop("present", None)
        entries.append(d)
    return set_attendance(tenant, p, entries, "import")


def add_advance(tenant: str, employee: str, amount, emi=None, given_on: str = "", note: str = "") -> dict[str, Any]:
    amt = _f(amount)
    if amt <= 0:
        raise PayError("Advance ki rakam chahiye.")
    with session() as s:
        e = _emp(s, tenant, employee)
        g = _date(given_on) or today(tenant)
        a = PayAdvance(tenant_id=tenant, employee_id=e.id, given_on=g.isoformat(), amount=amt, emi=_f(emi) or amt, balance=amt,
                       note=str(note)[:200])
        s.add(a)
        name = e.name
    return {"say": f"{name} ko {_rs(amt)} advance note: har mahine {_rs(_f(emi) or amt)} katega."}


# ---------------------------------------------------------------- the monthly run
def _ytd(s, tenant: str, emp_id: str, period: str) -> dict:
    """Paid / approved earlier months of this financial year (TDS base)."""
    months = R.fy_months(period)
    before = months[:months.index(R.norm_period(period))]
    out = {"gross": 0.0, "tds": 0.0, "pt": 0.0, "months": 0}
    if not before:
        return out
    for r in s.scalars(select(PayRun).where(PayRun.tenant_id == tenant, PayRun.period.in_(before),
                                            PayRun.status.in_(("approved", "paid")))).all():
        for ln in _j(r.lines, []):
            if ln.get("employee_id") == emp_id:
                out["gross"] += _f(ln.get("gross"))
                out["tds"] += _f(ln.get("tds"))
                out["pt"] += _f(ln.get("pt"))
                out["months"] += 1
    return out


def _last_net(s, tenant: str, emp_id: str, period: str) -> float | None:
    y, m = R.parse_period(period)
    prev = f"{y - 1:04d}-12" if m == 1 else f"{y:04d}-{m - 1:02d}"
    r = s.scalars(select(PayRun).where(PayRun.tenant_id == tenant, PayRun.period == prev, PayRun.kind == "monthly")).first()
    if not r:
        return None
    for ln in _j(r.lines, []):
        if ln.get("employee_id") == emp_id:
            return _f(ln.get("net"))
    return None


def _line(s, tenant: str, st: PaySettings, e: PayEmployee, p: str, a: PayAttendance | None, on: date,
          last_day: date | None = None) -> tuple[dict, list[str]]:
    y, m = R.parse_period(p)
    pay = {k: _f(v) for k, v in _j(e.pay, {}).items() if k in COMPONENTS}
    mdays = R.days_in(p)
    basis = mdays if st.days_basis == "calendar" else int(st.days_basis)
    # joined / left in the month: unpaid days before joining / after leaving
    doj, dol = _date(e.doj), last_day or _date(e.dol)
    absent_extra = 0
    if doj and (doj.year, doj.month) == (y, m):
        absent_extra += doj.day - 1
    if dol and (dol.year, dol.month) == (y, m):
        absent_extra += mdays - dol.day
    lop = min(float(basis), (a.lop_days if a else 0.0) + absent_extra * basis / mdays)
    paid_days = round(basis - lop, 2)
    f = paid_days / basis if basis else 0
    earned = {k: round(v * f, 2) for k, v in pay.items()}
    full_total = sum(pay.values())
    wages_full = R.code_wages(pay.get("basic", 0) + pay.get("da", 0), full_total)
    ot = R.ot_amount(wages_full, a.ot_hours if a else 0.0)
    extra = a.extra if a else 0.0
    gross = round(sum(earned.values()) + ot + extra, 2)
    wages = R.code_wages(earned.get("basic", 0) + earned.get("da", 0), sum(earned.values()))
    problems: list[str] = []
    ln: dict[str, Any] = {"employee_id": e.id, "code": e.code, "name": e.name, "designation": e.designation, "uan": e.uan,
                          "esic_ip": e.esic_ip, "days": basis, "paid_days": paid_days, "lop_days": round(lop, 2),
                          "earned": earned, "ot_hours": a.ot_hours if a else 0.0, "ot": ot, "extra": extra, "gross": gross,
                          "wages": wages}
    # EPF
    pfx = {"ee": 0, "eps": 0, "er_epf": 0, "edli": 0, "admin": 0, "pf_wages": 0, "eps_wages": 0, "edli_wages": 0}
    if st.pf_on and e.pf_member:
        eps_ok = e.eps_member and not (e.dob and _age(e.dob, on) >= 58)
        pfx = R.pf(p, wages, eps_ok, e.pf_on_full)
        if not e.uan:
            problems.append(f"{e.name}: UAN nahi (ECR mein nahi jayega).")
    ln["pf"] = pfx
    # ESIC: coverage fixed for the contribution period
    ex = {"ee": 0, "er": 0, "esi_wages": 0}
    if st.esi_on:
        es = _j(e.esi_state, {})
        cp = R.esi_period_start(p)
        if es.get("period") != cp:
            covered = (gross - ot - extra) <= R.ESI_CEILING if paid_days else (sum(pay.values()) <= R.ESI_CEILING)
            es = {"period": cp, "covered": bool(covered)}
            e.esi_state = json.dumps(es)
        if es.get("covered"):
            ex = R.esi(gross, paid_days)
            if not e.esic_ip:
                problems.append(f"{e.name}: ESIC IP number nahi.")
    ln["esi"] = ex
    # Professional tax
    half = None
    if st.state in ("33", "32") and m in (9, 3):
        months6 = [f"{y:04d}-{mm:02d}" for mm in range(m - 5, m)]
        prev = 0.0
        for r in s.scalars(select(PayRun).where(PayRun.tenant_id == tenant, PayRun.period.in_(months6))).all():
            prev += sum(_f(x.get("gross")) for x in _j(r.lines, []) if x.get("employee_id") == e.id)
        half = prev + gross
    ln["pt"] = R.pt(st.state, gross, m, e.gender, half) if gross > 0 else 0
    # TDS: whole-year estimate spread over the months left
    yt = _ytd(s, tenant, e.id, p)
    left = R.months_left_in_fy(p)
    projected = yt["gross"] + gross + sum(pay.values()) * (left - 1) + e.prev_income
    if dol:
        projected = yt["gross"] + gross + e.prev_income
        left = 1
    t = R.monthly_tds(projected, e.regime, e.declared, yt["tds"] + e.prev_tds, left, 0.0, yt["pt"] + ln["pt"] * left)
    ln["tds"] = t["this_month"]
    ln["tax_year"] = {k: t[k] for k in ("taxable", "total")}
    if t["total"] > 0 and not e.pan:
        problems.append(f"{e.name}: TDS banta hai par PAN nahi (bina PAN 20% kaatna padega).")
    # advances + other deductions
    adv = 0.0
    for av in s.scalars(select(PayAdvance).where(PayAdvance.tenant_id == tenant, PayAdvance.employee_id == e.id,
                                                 PayAdvance.balance > 0)).all():
        adv += min(av.emi, av.balance)
    ln["advance"] = round(adv, 2)
    ln["other_deduct"] = a.deduct if a else 0.0
    ded = pfx["ee"] + ex["ee"] + ln["pt"] + ln["tds"] + ln["advance"] + ln["other_deduct"]
    ln["deductions"] = round(ded, 2)
    ln["net"] = round(gross - ded)
    ln["employer"] = {"pf": pfx["eps"] + pfx["er_epf"], "edli": pfx["edli"], "admin": pfx["admin"], "esi": ex["er"]}
    ln["ctc_month"] = round(gross + ln["employer"]["pf"] + pfx["edli"] + pfx["admin"] + ex["er"])
    # checks
    if st.min_wage and paid_days and wages_full < st.min_wage:
        problems.append(f"{e.name}: {_rs(wages_full)} minimum wage {_rs(st.min_wage)} se kam.")
    if wages and (ln["advance"] + ln["other_deduct"] + pfx["ee"] + ex["ee"]) > 0.5 * gross:
        problems.append(f"{e.name}: katauti pagaar ke 50% se zyada (Code on Wages): advance ki kist kam kijiye.")
    if not _j(_open(e.bank), {}).get("account"):
        problems.append(f"{e.name}: bank account nahi (bank file mein nahi jayega; cash / cheque ki entry rakhiye).")
    last = _last_net(s, tenant, e.id, p)
    if last and abs(ln["net"] - last) > 0.25 * last and not a:
        problems.append(f"{e.name}: pichhle mahine {_rs(last)}, is baar {_rs(ln['net'])} (25% se zyada farak).")
    ln["problems"] = problems
    return ln, problems


def prepare_run(tenant: str, period: str, conv: str = "") -> dict[str, Any]:
    p = R.norm_period(period)
    y, m = R.parse_period(p)
    first, last = date(y, m, 1), date(y, m, R.days_in(p))
    with _LOCK, session() as s:
        st = _st(s, tenant)
        r = s.scalars(select(PayRun).where(PayRun.tenant_id == tenant, PayRun.period == p, PayRun.kind == "monthly")).first()
        if r and r.status == "paid":
            raise PayError(f"{R.nice(p)} ki salary di ja chuki. Badlaav agle mahine arrears mein.")
        emps = [e for e in s.scalars(select(PayEmployee).where(PayEmployee.tenant_id == tenant)).all()
                if (not e.doj or _date(e.doj) <= last) and (not e.dol or _date(e.dol) >= first)]
        if not emps:
            raise PayError("Koi employee nahi. Pehle employee jodiye (naam, pagaar, joining date).")
        att = {a.employee_id: a for a in s.scalars(select(PayAttendance).where(PayAttendance.tenant_id == tenant,
                                                                               PayAttendance.period == p)).all()}
        on = today(tenant)
        lines, checks = [], []
        for e in sorted(emps, key=lambda x: x.code):
            ln, pr = _line(s, tenant, st, e, p, att.get(e.id), on)
            lines.append(ln)
            checks += pr
        tot = {"employees": len(lines), "gross": round(sum(x["gross"] for x in lines)), "net": round(sum(x["net"] for x in lines)),
               "pf_ee": sum(x["pf"]["ee"] for x in lines), "pf_er": sum(x["employer"]["pf"] for x in lines),
               "eps": sum(x["pf"]["eps"] for x in lines), "edli": sum(x["pf"]["edli"] for x in lines),
               "pf_admin": max(500, sum(x["pf"]["admin"] for x in lines)) if st.pf_on and any(x["pf"]["ee"] for x in lines) else 0,
               "esi_ee": sum(x["esi"]["ee"] for x in lines), "esi_er": sum(x["esi"]["er"] for x in lines),
               "pt": sum(x["pt"] for x in lines), "tds": sum(x["tds"] for x in lines), "advance": round(sum(x["advance"] for x in lines))}
        tot["pf_challan"] = tot["pf_ee"] + tot["pf_er"] + tot["edli"] + tot["pf_admin"]
        tot["esi_challan"] = tot["esi_ee"] + tot["esi_er"]
        tot["cost"] = round(tot["gross"] + tot["pf_er"] + tot["edli"] + tot["pf_admin"] + tot["esi_er"])
        ny, nm = R.next_month(p)
        pay_by = date(ny, nm, 7)
        if on > pay_by:
            checks.insert(0, f"Salary {pay_by.strftime('%d %b')} tak deni thi (Code on Wages): {(on - pay_by).days} din late.")
        if not st.min_wage:
            checks.append("Minimum wage set nahi: state ki notification wali rakam bataiye taaki Hoola check kare.")
        if len(emps) >= 20 and not st.pf_on:
            checks.append(f"{len(emps)} log hain: 20 ya zyada par EPFO registration zaroori. PF chalu kijiye.")
        if len(emps) >= 10 and not st.esi_on:
            checks.append(f"{len(emps)} log: 10 ya zyada par ESIC lagta hai (jahan ESIC notified hai).")
        if p == "2026-09" and st.pf_on:
            checks.append("September 2026: PF 1-16 ko 15,000 aur 17-30 ko 25,000 ki seema par (S.O. 5109(E)), ek hi ECR 15 Oct tak.")
        if not r:
            r = PayRun(tenant_id=tenant, period=p, kind="monthly")
            s.add(r)
        r.lines = json.dumps(lines, ensure_ascii=False)
        r.totals = json.dumps(tot)
        r.checks = json.dumps(checks, ensure_ascii=False)
        r.status = "ready"
        r.approved_at = None
        _runlog(r, f"Payroll taiyaar: {len(lines)} log, net {_rs(tot['net'])}")
        s.flush()
        v = _run_view(r)
    say = (f"{R.nice(p)} payroll taiyaar: {tot['employees']} log, kul pagaar {_rs(tot['gross'])}, haath mein {_rs(tot['net'])}. "
           f"PF challan {_rs(tot['pf_challan'])}, ESIC {_rs(tot['esi_challan'])}, PT {_rs(tot['pt'])}, TDS {_rs(tot['tds'])}. "
           f"Company ka kharcha {_rs(tot['cost'])}.")
    if checks:
        say += f" {len(checks)} cheezein dekhni hain: " + " ".join(checks[:3])
    say += " Salary dene se pehle aapki haan."
    _publish(tenant, "payroll_run", {"run_id": v["id"], "status": v["status"]}, conv)
    return {"run": v, "say": say}


def _run_view(r: PayRun, lines: bool = True) -> dict[str, Any]:
    v = {"id": r.id, "period": r.period, "period_nice": R.nice(r.period), "kind": r.kind, "status": r.status,
         "status_label": STATUS_LABEL.get(r.status, r.status), "totals": _j(r.totals, {}), "checks": _j(r.checks, []),
         "paid_on": r.paid_on, "log": _j(r.log, [])[-20:]}
    if lines:
        v["lines"] = _j(r.lines, [])
    return v


def _run(s, tenant: str, ref: str) -> PayRun:
    r = s.get(PayRun, str(ref or ""))
    if not r:
        try:
            p = R.norm_period(ref)
            r = s.scalars(select(PayRun).where(PayRun.tenant_id == tenant, PayRun.period == p, PayRun.kind == "monthly")).first()
        except ValueError:
            r = None
    if not r or r.tenant_id != tenant:
        raise PayError("Payroll nahi mila.")
    return r


def get_run(tenant: str, ref: str) -> dict[str, Any]:
    with session() as s:
        return _run_view(_run(s, tenant, ref))


def runs(tenant: str, limit: int = 24) -> list[dict[str, Any]]:
    with session() as s:
        return [_run_view(r, False) for r in s.scalars(select(PayRun).where(PayRun.tenant_id == tenant)
                                                       .order_by(PayRun.period.desc(), PayRun.created_at.desc()).limit(limit)).all()]


def decide(tenant: str, ref: str, approve: bool, note: str = "", conv: str = "") -> dict[str, Any]:
    with _LOCK, session() as s:
        r = _run(s, tenant, ref)
        if r.status not in ("ready", "rejected", "approved"):
            raise PayError(f"Abhi haan / na nahi ({STATUS_LABEL.get(r.status)}).")
        r.status = "approved" if approve else "rejected"
        r.approved_at = now() if approve else None
        _runlog(r, ("Haan" if approve else "Mana") + (f": {note[:100]}" if note else ""), "owner")
        v = _run_view(r, False)
    _publish(tenant, "payroll_run", {"run_id": v["id"], "status": v["status"]}, conv)
    if approve:
        return {"run": v, "say": "Haan mili. Bank file, payslip, ECR aur ESIC sheet bana raha hoon: bank mein upload aur payment aap karengi."}
    return {"run": v, "say": "Theek, roka. Kya badalna hai bataiye."}


# ---------------------------------------------------------------- files
def _payslip_pdf(company: str, p: str, lines: list[dict]) -> bytes:
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.units import mm
    from reportlab.pdfgen import canvas
    bio = io.BytesIO()
    c = canvas.Canvas(bio, pagesize=A4)
    W, H = A4
    for ln in lines:
        y = H - 20 * mm
        c.setFont("Helvetica-Bold", 14)
        c.drawString(20 * mm, y, company or "Payslip")
        c.setFont("Helvetica", 10)
        c.drawRightString(W - 20 * mm, y, f"Payslip: {R.nice(p)}")
        y -= 10 * mm
        for k, v in (("Employee", f"{ln['name']} ({ln['code']})"), ("Designation", ln.get("designation") or "-"),
                     ("UAN", ln.get("uan") or "-"), ("ESIC IP", ln.get("esic_ip") or "-"),
                     ("Paid days", f"{ln['paid_days']:g} of {ln['days']} (LOP {ln['lop_days']:g})")):
            c.drawString(20 * mm, y, f"{k}: {v}")
            y -= 6 * mm
        y -= 4 * mm
        c.setFont("Helvetica-Bold", 10)
        c.drawString(20 * mm, y, "Earnings")
        c.drawString(110 * mm, y, "Deductions")
        y -= 7 * mm
        c.setFont("Helvetica", 10)
        earn = [(k.title() if k != "da" else "DA", v) for k, v in ln["earned"].items() if v] + \
               ([("Overtime", ln["ot"])] if ln["ot"] else []) + ([("Incentive / arrears", ln["extra"])] if ln["extra"] else [])
        ded = [("Provident Fund", ln["pf"]["ee"]), ("ESIC", ln["esi"]["ee"]), ("Professional Tax", ln["pt"]),
               ("Income Tax (TDS)", ln["tds"]), ("Advance", ln["advance"]), ("Other", ln["other_deduct"])]
        ded = [(k, v) for k, v in ded if v]
        for i in range(max(len(earn), len(ded))):
            if i < len(earn):
                c.drawString(20 * mm, y, earn[i][0])
                c.drawRightString(95 * mm, y, f"{earn[i][1]:,.2f}")
            if i < len(ded):
                c.drawString(110 * mm, y, ded[i][0])
                c.drawRightString(W - 20 * mm, y, f"{ded[i][1]:,.2f}")
            y -= 6 * mm
        y -= 3 * mm
        c.setFont("Helvetica-Bold", 10)
        c.drawString(20 * mm, y, "Gross")
        c.drawRightString(95 * mm, y, f"{ln['gross']:,.2f}")
        c.drawString(110 * mm, y, "Total deductions")
        c.drawRightString(W - 20 * mm, y, f"{ln['deductions']:,.2f}")
        y -= 10 * mm
        c.setFont("Helvetica-Bold", 12)
        c.drawString(20 * mm, y, f"Net pay: Rs {ln['net']:,}")
        y -= 8 * mm
        c.setFont("Helvetica", 8)
        c.drawString(20 * mm, y, "Employer contributions: PF {:,} | EDLI {:,} | ESIC {:,}. Computer-generated payslip.".format(
            ln["employer"]["pf"], ln["employer"]["edli"], ln["employer"]["esi"]))
        c.showPage()
    c.save()
    return bio.getvalue()


def ecr_text(lines: list[dict]) -> str:
    """EPFO ECR 2.0: UAN#~#Name#~#Gross#~#EPF wages#~#EPS wages#~#EDLI wages#~#EE#~#EPS#~#ER diff#~#NCP days#~#Refund."""
    out = []
    for ln in lines:
        pfx = ln["pf"]
        if not pfx.get("ee") or not ln.get("uan"):
            continue
        name = re.sub(r"[^A-Za-z. ]", "", ln["name"]).upper()[:85]
        out.append("#~#".join(str(x) for x in (ln["uan"], name, R.rnd(ln["gross"]), pfx["pf_wages"], pfx["eps_wages"],
                                                pfx["edli_wages"], pfx["ee"], pfx["eps"], pfx["er_epf"], int(round(ln["lop_days"])), 0)))
    return "\n".join(out) + ("\n" if out else "")


def files(tenant: str, ref: str, conv: str = "") -> dict[str, Any]:
    """After her Haan: payslips PDF, salary register, bank bulk-payment file, ECR, ESIC sheet (she uploads / pays)."""
    import openpyxl
    from .. import chatfiles
    with session() as s:
        r = _run(s, tenant, ref)
        if r.status not in ("approved", "paid"):
            raise PayError("Pehle aapki haan chahiye: payroll dekh kar 'haan' kahiye.")
        st = _st(s, tenant)
        lines = _j(r.lines, [])
        p = r.period
        banks = {}
        for e in s.scalars(select(PayEmployee).where(PayEmployee.tenant_id == tenant)).all():
            banks[e.id] = _j(_open(e.bank), {}) if e.bank else {}
        company = st.company
        debit = _j(_open(st.pay_account), {}) if st.pay_account else {}
    out = []
    out.append(chatfiles.store(tenant, f"Payslips_{p}.pdf", _payslip_pdf(company, p, lines), "application/pdf"))
    # salary register
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Register"
    ws.append(["Code", "Name", "Paid days", "LOP", "Basic", "DA", "HRA", "Conveyance", "Special", "Other", "OT", "Incentive",
               "Gross", "PF (EE)", "ESIC (EE)", "PT", "TDS", "Advance", "Other ded.", "Net", "PF (ER)", "EDLI", "ESIC (ER)"])
    for ln in lines:
        er = ln["earned"]
        ws.append([ln["code"], ln["name"], ln["paid_days"], ln["lop_days"], er.get("basic", 0), er.get("da", 0), er.get("hra", 0),
                   er.get("conveyance", 0), er.get("special", 0), er.get("other", 0), ln["ot"], ln["extra"], ln["gross"],
                   ln["pf"]["ee"], ln["esi"]["ee"], ln["pt"], ln["tds"], ln["advance"], ln["other_deduct"], ln["net"],
                   ln["employer"]["pf"], ln["employer"]["edli"], ln["employer"]["esi"]])
    # bank bulk payment (generic NEFT upload layout most banks accept / map)
    wb2 = openpyxl.Workbook()
    wb_ = wb2.active
    wb_.title = "Bulk payment"
    wb_.append(["Beneficiary Name", "Account Number", "IFSC", "Amount", "Payment Mode", "Narration", "Debit Account"])
    missing = []
    for ln in lines:
        b = banks.get(ln["employee_id"], {})
        if not b.get("account") or ln["net"] <= 0:
            missing.append(ln["name"])
            continue
        wb_.append([ln["name"][:35], b["account"], b.get("ifsc", ""), ln["net"], "NEFT", f"Salary {R.nice(p)}", debit.get("account", "")])
    for wbk, nm in ((wb, f"Salary_Register_{p}.xlsx"), (wb2, f"Bank_Salary_{p}.xlsx")):
        bio = io.BytesIO()
        wbk.save(bio)
        out.append(chatfiles.store(tenant, nm, bio.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"))
    ecr = ecr_text(lines)
    if ecr:
        out.append(chatfiles.store(tenant, f"ECR_{p}.txt", ecr.encode(), "text/plain"))
    esi_rows = [ln for ln in lines if ln["esi"]["esi_wages"] and (ln["esi"]["ee"] or ln["esi"]["er"])]
    if esi_rows:
        wb3 = openpyxl.Workbook()
        w3 = wb3.active
        w3.append(["IP Number", "IP Name", "No of Days for which wages paid/payable during the month", "Total Monthly Wages",
                   "Reason Code for Zero workings days", "Last Working Day"])
        for ln in esi_rows:
            w3.append([ln.get("esic_ip", ""), ln["name"], int(round(ln["paid_days"])), R.rnd(ln["gross"]), "", ""])
        bio = io.BytesIO()
        wb3.save(bio)
        out.append(chatfiles.store(tenant, f"ESIC_{p}.xlsx", bio.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"))
    for f in out:
        _publish(tenant, "pc_file", {**f, "source": "payroll", "title": f.get("name", "")}, conv)
    say = (f"{R.nice(p)}: payslips (PDF), salary register, bank file, " + ("ECR, " if ecr else "") + ("ESIC sheet " if esi_rows else "") +
           "taiyaar. Bank file apne net banking ke bulk upload mein daal kar payment aap karengi; ECR EPFO portal par, ESIC sheet ESIC portal par.")
    if missing:
        say += " Bank file mein nahi (account nahi): " + ", ".join(missing[:8]) + "."
    return {"files": out, "missing_bank": missing, "say": say}


def mark_paid(tenant: str, ref: str, paid_on: str = "", by: str = "owner", conv: str = "") -> dict[str, Any]:
    with _LOCK, session() as s:
        r = _run(s, tenant, ref)
        if r.status != "approved":
            raise PayError("Pehle haan, phir payment. Abhi status: " + STATUS_LABEL.get(r.status, r.status))
        d = _date(paid_on) or today(tenant)
        r.status, r.paid_on = "paid", d.isoformat()
        for ln in _j(r.lines, []):                       # advances recovered this month
            left = _f(ln.get("advance"))
            for av in s.scalars(select(PayAdvance).where(PayAdvance.tenant_id == tenant, PayAdvance.employee_id == ln["employee_id"],
                                                         PayAdvance.balance > 0).order_by(PayAdvance.given_on)).all():
                if left <= 0:
                    break
                cut = min(left, av.balance)
                av.balance = round(av.balance - cut, 2)
                left -= cut
        _runlog(r, f"Salary di gayi {d.isoformat()}", by)
        v = _run_view(r, False)
    _publish(tenant, "payroll_run", {"run_id": v["id"], "status": v["status"]}, conv)
    return {"run": v, "say": f"{v['period_nice']} ki salary {d.isoformat()} ko di: log mein likha. Ab PF / ESIC 15 tareekh tak, TDS 7 tak."}


def payslip_drafts(tenant: str, ref: str) -> dict[str, Any]:
    """A short WhatsApp message per employee (she sends; the PDF payslip goes with it)."""
    with session() as s:
        r = _run(s, tenant, ref)
        phones = {e.id: e.phone for e in s.scalars(select(PayEmployee).where(PayEmployee.tenant_id == tenant)).all()}
        lines = _j(r.lines, [])
        st = _st(s, tenant)
    out = []
    for ln in lines:
        out.append({"employee": ln["name"], "to": phones.get(ln["employee_id"], ""),
                    "text": (f"Namaste {ln['name'].split()[0]} ji, {R.nice(r.period)} ki salary: kaam ke din {ln['paid_days']:g}/{ln['days']}, "
                             f"kul {_rs(ln['gross'])}, katauti {_rs(ln['deductions'])} (PF {_rs(ln['pf']['ee'])}, ESIC {_rs(ln['esi']['ee'])}, "
                             f"PT {_rs(ln['pt'])}, TDS {_rs(ln['tds'])}), haath mein {_rs(ln['net'])}. - {st.company or 'HR'}")})
    return {"drafts": out, "say": f"{len(out)} employee ke liye payslip message draft (aap bhejengi)."}


# ---------------------------------------------------------------- leaving: full and final
def full_and_final(tenant: str, employee: str, last_day: str, leave_days=0, notice_recovery=0, other=0,
                   death_or_disability: bool = False, conv: str = "") -> dict[str, Any]:
    """Dues of a person who leaves: salary till the last day, leave encashment, gratuity if eligible, minus recoveries.
    Code on Wages: pay within 2 working days of leaving."""
    ld = _date(last_day)
    if not ld:
        raise PayError("Aakhri din ki date chahiye.")
    p = f"{ld.year:04d}-{ld.month:02d}"
    with _LOCK, session() as s:
        st = _st(s, tenant)
        e = _emp(s, tenant, employee)
        e.dol = ld.isoformat()
        att = s.scalars(select(PayAttendance).where(PayAttendance.tenant_id == tenant, PayAttendance.employee_id == e.id,
                                                    PayAttendance.period == p)).first()
        ln, pr = _line(s, tenant, st, e, p, att, ld, ld)
        pay = {k: _f(v) for k, v in _j(e.pay, {}).items() if k in COMPONENTS}
        wages = R.code_wages(pay.get("basic", 0) + pay.get("da", 0), sum(pay.values()))
        enc = R.leave_encashment(wages, _f(leave_days))
        doj = _date(e.doj)
        gr = R.gratuity(wages, doj, ld, e.fixed_term, death_or_disability) if doj else {"eligible": False, "amount": 0, "why": "Joining date nahi."}
        adv_bal = sum(a.balance for a in s.scalars(select(PayAdvance).where(PayAdvance.tenant_id == tenant, PayAdvance.employee_id == e.id,
                                                                           PayAdvance.balance > 0)).all())
        recover = round(_f(notice_recovery) + _f(other) + max(0.0, adv_bal - ln["advance"]))
        total = round(ln["net"] + enc + gr["amount"] - recover)
        due = R.working_days_after(ld, 2)
        fnf = {"salary_line": ln, "leave_encashment": enc, "gratuity": gr, "recoveries": recover, "net_payable": total,
               "pay_by": due.isoformat()}
        r = PayRun(tenant_id=tenant, period=p, kind="fnf", status="ready", lines=json.dumps([ln | {"fnf": fnf}], ensure_ascii=False),
                   totals=json.dumps({"employees": 1, "net": total, "gross": ln["gross"]}), checks=json.dumps(pr, ensure_ascii=False))
        _runlog(r, f"Full & final: {e.name}, aakhri din {ld.isoformat()}")
        s.add(r)
        e.active = False
        s.flush()
        v = _run_view(r)
        name = e.name
    say = (f"{name} ka full & final: {ld.isoformat()} tak ki salary {_rs(ln['net'])}, chhutti ke paise {_rs(enc)}, gratuity "
           + (_rs(gr["amount"]) + f" ({gr['years']} saal)" if gr.get("eligible") else f"nahi ({gr.get('why', '')})")
           + (f", recovery {_rs(recover)}" if recover else "") + f". Kul {_rs(total)}, {due.strftime('%d %b')} tak dena hai (2 working din). Aapki haan?")
    _publish(tenant, "payroll_run", {"run_id": v["id"], "status": v["status"]}, conv)
    return {"run": v, "fnf": fnf, "say": say}


# ---------------------------------------------------------------- bonus (yearly)
def bonus_report(tenant: str, fy_period: str, rate: float = 8.33) -> dict[str, Any]:
    months = R.fy_months(fy_period)
    with session() as s:
        st = _st(s, tenant)
        wages: dict[str, list[float]] = {}
        names: dict[str, str] = {}
        for r in s.scalars(select(PayRun).where(PayRun.tenant_id == tenant, PayRun.period.in_(months), PayRun.kind == "monthly",
                                                PayRun.status.in_(("approved", "paid")))).all():
            for ln in _j(r.lines, []):
                wages.setdefault(ln["employee_id"], []).append(_f(ln.get("wages")))
                names[ln["employee_id"]] = ln["name"]
        rows = []
        for eid, ws in wages.items():
            b = R.bonus(ws, st.min_wage, rate)
            if b > 0:
                rows.append({"employee": names[eid], "months": len(ws), "bonus": b})
    tot = sum(x["bonus"] for x in rows)
    return {"fy": R.fy_of(fy_period), "rows": rows, "total": tot,
            "say": f"FY {R.fy_of(fy_period)} bonus ({rate:g}%): {len(rows)} log, kul {_rs(tot)}. 30 November tak dena (FY ke 8 mahine mein)."}


# ---------------------------------------------------------------- calendar + morning
def calendar(tenant: str, ahead: int = 12) -> dict[str, Any]:
    on = today(tenant)
    with session() as s:
        st = _st(s, tenant)
        rs_ = s.scalars(select(PayRun).where(PayRun.tenant_id == tenant, PayRun.kind == "monthly")).all()
        n_emp = len(s.scalars(select(PayEmployee.id).where(PayEmployee.tenant_id == tenant, PayEmployee.active.is_(True))).all())
        fnf = [r for r in s.scalars(select(PayRun).where(PayRun.tenant_id == tenant, PayRun.kind == "fnf",
                                                         PayRun.status.in_(("ready", "approved")))).all()]
    status = {r.period: r.status for r in rs_}
    items, late = [], []
    for back in (2, 1):
        y, m = on.year, on.month - back
        while m < 1:
            y, m = y - 1, m + 12
        p = f"{y:04d}-{m:02d}"
        for d in R.dues(p, st.pf_on, st.esi_on):
            if d["kind"] == "salary" and status.get(p) == "paid":
                continue
            days = (d["due"] - on).days
            row = {"period": p, "period_nice": R.nice(p), "kind": d["kind"], "label": d["label"], "due": d["due"].isoformat(),
                   "due_nice": d["due"].strftime("%d %b %Y"), "days": days, "run_status": status.get(p, "")}
            if days < 0 and back == 1 and d["kind"] in ("salary", "pf", "esi", "tds"):
                late.append(row)
            elif 0 <= days <= ahead:
                items.append(row)
    for r in fnf:
        fl = _j(r.lines, [{}])[0].get("fnf", {})
        d = _date(fl.get("pay_by"))
        if d:
            (late if d < on else items).append({"kind": "fnf", "label": f"Full & final: {_j(r.lines, [{}])[0].get('name', '')}",
                                                "due": d.isoformat(), "due_nice": d.strftime("%d %b %Y"), "days": (d - on).days})
    items.sort(key=lambda x: x["days"])
    parts = []
    if late:
        parts.append("Late: " + ", ".join(f"{x['label']} ({x.get('period_nice', '')})" for x in late[:4]))
    soon = [x for x in items if x["days"] <= 3]
    if soon:
        parts.append("Jaldi: " + ", ".join(f"{x['label']} {x['due_nice']}" for x in soon[:5]))
    return {"date": on.isoformat(), "employees": n_emp, "due": items, "late": late,
            "text": ". ".join(parts) or "Payroll mein abhi kuch due nahi."}


def due_mornings() -> list[str]:
    from ..life import tz_of
    out = []
    with session() as s:
        for r in s.scalars(select(PaySettings).where(PaySettings.enabled.is_(True))).all():
            if not s.scalars(select(PayEmployee.id).where(PayEmployee.tenant_id == r.tenant_id).limit(1)).first():
                continue
            local = datetime.now(tz_of(r.tenant_id))
            d = local.date().isoformat()
            try:
                hh, mm = [int(x) for x in (r.morning_time or "09:45").split(":")]
            except ValueError:
                hh, mm = 9, 45
            late = (local.hour * 60 + local.minute) - (hh * 60 + mm)
            if r.last_morning != d and 0 <= late <= 180:
                r.last_morning = d
                out.append(r.tenant_id)
    return out


def send_morning(tenant: str) -> dict[str, Any]:
    from ..events import notify
    from ..net.service import home_conv
    cal = calendar(tenant, 5)
    if not cal["due"] and not cal["late"]:
        return cal
    conv = home_conv(tenant)
    try:
        notify(conv, "Payroll: " + cal["text"])
    except Exception:
        pass
    _publish(tenant, "payroll_day", {"date": cal["date"], "text": cal["text"]}, conv)
    return cal
