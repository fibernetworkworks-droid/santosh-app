"""v0.96 Auditor's Assistant (Audit Sahayak): Hoola does an audit's leg-work end to end; the CA judges and signs.

  1. Gets the job: "Sharma Traders ka FY 2025-26 tax audit shuru karo", the client's Tally trial balance (this year and
     last), the day book / ledger export, bank statements, replies to queries, confirmation replies.
  2. Does the audit work, step by step on screen: materiality (SA 320, reason written down), analytical review of every
     ledger vs last year (SA 520), journal-entry tests (SA 240: Sundays / holidays, round amounts, just below her
     approval limit, year-end and back-dated entries, missing narration, rarely used ledgers, odd users, duplicates),
     cash rules (Sec 34(3) cash payment over Rs 10,000, Sec 221 cash receipt Rs 2 lakh+, Sec 219 / 220 cash loans),
     TDS on payments above Sec 393 thresholds, Benford first-digit test, a reproducible sample (SA 530, monetary-unit or
     random, seed recorded), balance confirmations (SA 505) from the debtors / creditors in the books.
  3. Checked against the standards and earlier work: tax audit / CARO applicability, the audit trail (Rule 11(g)),
     going-concern signs (SA 570), last year's explained items (not raised again), her own thresholds.
  4. Ready as a card, a file or a draft: findings card, the client query list, confirmation letters, a management
     representation letter draft (SA 580), the working-paper Excel (SA 230).
  5. Waits for her Haan: sending queries / confirmations to the client or third parties. Hoola never signs, never
     forms or words the audit opinion as final, never generates a UDIN, never files Form 3CD / 26.
  6. Delivers and learns: every finding is closed with her note; explained patterns per client are remembered for next
     year; the engagement log is the audit trail of the work.
"""
from __future__ import annotations

import csv
import io
import json
import logging
import re
import threading
from collections import Counter, defaultdict
from datetime import date, datetime, timedelta
from typing import Any

from sqlalchemy import Boolean, DateTime, Float, Integer, String, Text, select
from sqlalchemy.orm import Mapped, mapped_column

from ..db import Base, new_id, now, session
from . import rules as R

log = logging.getLogger("santosh.audit")
_LOCK = threading.Lock()
KINDS = {"statutory": "Statutory audit", "tax": "Tax audit", "internal": "Internal audit", "stock": "Stock audit",
         "bank": "Bank / concurrent audit", "other": "Audit"}
STATUS = ("planning", "fieldwork", "review", "done")


class AuditError(Exception):
    pass


class AuditEngagement(Base):
    __tablename__ = "audit_engagements"
    id: Mapped[str] = mapped_column(String(40), primary_key=True, default=lambda: new_id("ae"))
    tenant_id: Mapped[str] = mapped_column(String(64), index=True)
    client: Mapped[str] = mapped_column(String(160), default="")
    kind: Mapped[str] = mapped_column(String(12), default="statutory")
    entity: Mapped[str] = mapped_column(String(12), default="firm")            # company (private / public) | firm | individual | llp
    fy_start: Mapped[int] = mapped_column(Integer, default=2025)
    status: Mapped[str] = mapped_column(String(10), default="planning")
    approval_limit: Mapped[float] = mapped_column(Float, default=0.0)          # the client's own approval limit (JE test)
    materiality: Mapped[str] = mapped_column(Text, default="{}")
    tb: Mapped[str] = mapped_column(Text, default="[]")                        # this year's TB lines
    tb_prev: Mapped[str] = mapped_column(Text, default="[]")
    facts: Mapped[str] = mapped_column(Text, default="{}")                     # turnover, cash %, borrowings...
    log: Mapped[str] = mapped_column(Text, default="[]")
    created_at: Mapped[datetime] = mapped_column(DateTime(timezone=True), default=now)


class AuditEntry(Base):
    """A day-book line (voucher x ledger)."""
    __tablename__ = "audit_entries"
    id: Mapped[str] = mapped_column(String(40), primary_key=True, default=lambda: new_id("ay"))
    tenant_id: Mapped[str] = mapped_column(String(64), index=True)
    engagement_id: Mapped[str] = mapped_column(String(40), index=True)
    day: Mapped[str] = mapped_column(String(10), default="")
    voucher: Mapped[str] = mapped_column(String(40), default="")
    vtype: Mapped[str] = mapped_column(String(30), default="")
    ledger: Mapped[str] = mapped_column(String(160), default="")
    debit: Mapped[float] = mapped_column(Float, default=0.0)
    credit: Mapped[float] = mapped_column(Float, default=0.0)
    narration: Mapped[str] = mapped_column(String(300), default="")
    user: Mapped[str] = mapped_column(String(60), default="")
    entered_on: Mapped[str] = mapped_column(String(10), default="")


class AuditFinding(Base):
    __tablename__ = "audit_findings"
    id: Mapped[str] = mapped_column(String(40), primary_key=True, default=lambda: new_id("af"))
    tenant_id: Mapped[str] = mapped_column(String(64), index=True)
    engagement_id: Mapped[str] = mapped_column(String(40), index=True)
    area: Mapped[str] = mapped_column(String(30), default="")                 # analytics | je | cash | tds | benford | trail | gc | other
    rule: Mapped[str] = mapped_column(String(40), default="")
    ref: Mapped[str] = mapped_column(String(120), default="")                 # voucher / ledger
    amount: Mapped[float] = mapped_column(Float, default=0.0)
    text: Mapped[str] = mapped_column(Text, default="")
    level: Mapped[str] = mapped_column(String(6), default="warn")
    status: Mapped[str] = mapped_column(String(10), default="open")           # open | asked | explained | adjusted | reported
    note: Mapped[str] = mapped_column(Text, default="")
    key: Mapped[str] = mapped_column(String(120), default="", index=True)


class AuditSample(Base):
    __tablename__ = "audit_samples"
    id: Mapped[str] = mapped_column(String(40), primary_key=True, default=lambda: new_id("as"))
    tenant_id: Mapped[str] = mapped_column(String(64), index=True)
    engagement_id: Mapped[str] = mapped_column(String(40), index=True)
    population: Mapped[str] = mapped_column(String(80), default="")
    ref: Mapped[str] = mapped_column(String(60), default="")
    amount: Mapped[float] = mapped_column(Float, default=0.0)
    why: Mapped[str] = mapped_column(String(20), default="sample")            # key | sample
    status: Mapped[str] = mapped_column(String(10), default="pending")        # pending | ok | exception
    note: Mapped[str] = mapped_column(Text, default="")


class AuditConfirm(Base):
    __tablename__ = "audit_confirms"
    id: Mapped[str] = mapped_column(String(40), primary_key=True, default=lambda: new_id("ac"))
    tenant_id: Mapped[str] = mapped_column(String(64), index=True)
    engagement_id: Mapped[str] = mapped_column(String(40), index=True)
    party: Mapped[str] = mapped_column(String(160), default="")
    kind: Mapped[str] = mapped_column(String(10), default="debtor")
    balance: Mapped[float] = mapped_column(Float, default=0.0)
    status: Mapped[str] = mapped_column(String(10), default="draft")          # draft | approved | sent | agreed | differs
    reply: Mapped[float | None] = mapped_column(Float, nullable=True)
    note: Mapped[str] = mapped_column(Text, default="")


class AuditLearn(Base):
    """Per client: patterns she explained once (e.g. 'Sunday entries in March = stock count') so they aren't raised again."""
    __tablename__ = "audit_learn"
    id: Mapped[str] = mapped_column(String(40), primary_key=True, default=lambda: new_id("al"))
    tenant_id: Mapped[str] = mapped_column(String(64), index=True)
    client: Mapped[str] = mapped_column(String(160), index=True)
    key: Mapped[str] = mapped_column(String(120), default="")
    note: Mapped[str] = mapped_column(Text, default="")


# ---------------------------------------------------------------- helpers
def _j(s: str, default):
    try:
        v = json.loads(s or "")
        return v if isinstance(v, type(default)) else default
    except (TypeError, ValueError):
        return default


def _rs(v) -> str:
    v = round(float(v or 0))
    s = str(abs(v))
    if len(s) > 3:
        s = re.sub(r"(\d)(?=(\d\d)+$)", r"\1,", s[:-3]) + "," + s[-3:]
    return ("-" if v < 0 else "") + "₹" + s


def _date(v) -> date | None:
    from ..analyst.service import _date as d
    return d(v)


def _tz(tenant: str):
    from ..life import tz_of
    return tz_of(tenant)


def today(tenant: str) -> date:
    return datetime.now(_tz(tenant)).date()


def _log(e: AuditEngagement, what: str, by: str = "hoola") -> None:
    lg = _j(e.log, [])
    lg.append({"at": now().isoformat(timespec="seconds"), "what": what[:300], "by": by})
    e.log = json.dumps(lg[-300:], ensure_ascii=False)


def _publish(tenant: str, kind: str, payload: dict, conv: str = "") -> None:
    from ..events import bus
    try:
        if not conv:
            from ..net.service import home_conv
            conv = home_conv(tenant)
        bus.publish(conv, kind, payload)
    except Exception:
        log.exception("audit publish failed")


def _load_file(tenant: str, file_id: str) -> tuple[str, bytes]:
    from ..mailbox import service as ms
    try:
        name, _m, data = ms.load_files(tenant, [{"type": "file", "file_id": file_id}])[0]
    except Exception as e:
        raise AuditError(f"File {file_id} nahi mili: {e}")
    return name, data


def _eng(s, tenant: str, ref: str) -> AuditEngagement:
    ref = str(ref or "").strip()
    e = s.get(AuditEngagement, ref) if ref else None
    if e and e.tenant_id == tenant:
        return e
    q = select(AuditEngagement).where(AuditEngagement.tenant_id == tenant)
    rows = s.scalars(q.order_by(AuditEngagement.created_at.desc())).all()
    if not ref:
        open_ = [r for r in rows if r.status != "done"]
        if len(open_) == 1:
            return open_[0]
        raise AuditError("Kaun sa audit? Client ka naam bataiye.")
    hits = [r for r in rows if ref.lower() in r.client.lower()]
    if hits:
        return hits[0]
    raise AuditError(f"Audit '{ref}' nahi mila.")


def _view(s, e: AuditEngagement) -> dict[str, Any]:
    f = s.scalars(select(AuditFinding).where(AuditFinding.engagement_id == e.id)).all()
    by = Counter(x.status for x in f)
    return {"id": e.id, "client": e.client, "kind": e.kind, "kind_label": KINDS.get(e.kind, e.kind), "entity": e.entity,
            "fy": f"{e.fy_start}-{str(e.fy_start + 1)[2:]}", "fy_start": e.fy_start, "status": e.status,
            "materiality": _j(e.materiality, {}), "facts": _j(e.facts, {}), "has_tb": bool(_j(e.tb, [])),
            "has_prev": bool(_j(e.tb_prev, [])), "entries": len(s.scalars(select(AuditEntry.id).where(AuditEntry.engagement_id == e.id)).all()),
            "findings": {"open": by.get("open", 0) + by.get("asked", 0), "closed": sum(v for k, v in by.items() if k not in ("open", "asked")),
                         "total": len(f)},
            "log": _j(e.log, [])[-25:]}


# ---------------------------------------------------------------- engagement
def new_engagement(tenant: str, client: str, kind: str = "statutory", fy: str = "", entity: str = "firm",
                   approval_limit=None, by: str = "owner") -> dict[str, Any]:
    if not str(client).strip():
        raise AuditError("Client ka naam chahiye.")
    k = kind if kind in KINDS else "other"
    m = re.match(r"(\d{4})", str(fy or ""))
    on = today(tenant)
    fs = int(m.group(1)) if m else (on.year - 1 if on.month >= 4 else on.year - 2)
    with _LOCK, session() as s:
        e = AuditEngagement(tenant_id=tenant, client=str(client).strip()[:160], kind=k, entity=str(entity or "firm").lower()[:12],
                            fy_start=fs, approval_limit=float(approval_limit or 0))
        s.add(e)
        _log(e, f"{KINDS[k]} shuru: {e.client}, FY {fs}-{str(fs + 1)[2:]}", by)
        s.flush()
        v = _view(s, e)
    return {"engagement": v, "say": f"{v['kind_label']} {v['client']} FY {v['fy']} shuru. Ab Tally ka trial balance (is saal + pichhle saal) "
                                     "aur day book (Excel) bhejiye."}


def engagements(tenant: str) -> list[dict]:
    with session() as s:
        return [_view(s, e) for e in s.scalars(select(AuditEngagement).where(AuditEngagement.tenant_id == tenant)
                                               .order_by(AuditEngagement.created_at.desc())).all()]


def get(tenant: str, ref: str) -> dict:
    with session() as s:
        return _view(s, _eng(s, tenant, ref))


def set_facts(tenant: str, ref: str, **facts: Any) -> dict[str, Any]:
    with session() as s:
        e = _eng(s, tenant, ref)
        f = _j(e.facts, {})
        f |= {k: v for k, v in facts.items() if v not in (None, "")}
        e.facts = json.dumps(f)
        if facts.get("approval_limit") not in (None, ""):
            e.approval_limit = float(facts["approval_limit"])
        if facts.get("status") in STATUS:
            e.status = facts["status"]
        _log(e, "Jaankari update: " + ", ".join(f"{k}={v}" for k, v in facts.items() if v not in (None, "")))
        return {"engagement": _view(s, e)}


# ---------------------------------------------------------------- trial balance + day book
def import_tb(tenant: str, ref: str, file_id: str = "", data: bytes | None = None, name: str = "tb.xlsx", previous: bool = False) -> dict[str, Any]:
    from ..analyst.service import _balances, _pick, parse_books
    if file_id:
        name, data = _load_file(tenant, file_id)
    if not data:
        raise AuditError("Trial balance file chahiye.")
    try:
        lines = _pick(parse_books(name, data))
    except Exception as ex:
        raise AuditError(str(ex))
    with _LOCK, session() as s:
        e = _eng(s, tenant, ref)
        tot_dr, tot_cr = sum(x["dr"] for x in lines), sum(x["cr"] for x in lines)
        if previous:
            e.tb_prev = json.dumps(lines, ensure_ascii=False)
        else:
            e.tb = json.dumps(lines, ensure_ascii=False)
        _log(e, f"TB {'pichhla saal' if previous else 'is saal'}: {len(lines)} ledger, Dr {_rs(tot_dr)} Cr {_rs(tot_cr)}")
        cid = e.id
    say = f"Trial balance {'(pichhla saal) ' if previous else ''}liya: {len(lines)} ledger."
    if abs(tot_dr - tot_cr) > 1:
        say += f" Dhyaan: Debit {_rs(tot_dr)} aur Credit {_rs(tot_cr)} barabar nahi ({_rs(tot_dr - tot_cr)} ka farak)."
    out = {"ledgers": len(lines), "difference": round(tot_dr - tot_cr, 2), "say": say}
    if not previous:
        out["materiality"] = plan_materiality(tenant, cid)
        out["say"] += " " + out["materiality"]["say"]
    return out


_DB = {"day": ("date", "voucher date", "vch date"), "voucher": ("voucher no", "voucher no.", "vch no", "vch no.", "voucher number", "doc no"),
       "vtype": ("voucher type", "vch type", "type"), "ledger": ("ledger", "particulars", "account", "ledger name", "party"),
       "debit": ("debit", "debit amount", "dr", "debit amt"), "credit": ("credit", "credit amount", "cr", "credit amt"),
       "narration": ("narration", "remarks", "description"), "user": ("entered by", "user", "created by", "user name"),
       "entered_on": ("entry date", "created on", "altered on", "entered on", "created date")}


def import_daybook(tenant: str, ref: str, file_id: str = "", data: bytes | None = None, name: str = "daybook.xlsx") -> dict[str, Any]:
    from ..analyst.service import _rows
    if file_id:
        name, data = _load_file(tenant, file_id)
    rows = _rows(name, data or b"")
    hi, cols = None, {}
    for i, r in enumerate(rows[:25]):
        h = [re.sub(r"\s+", " ", str(x or "").strip().lower()) for x in r]
        f = {}
        for j, x in enumerate(h):
            for k, ns in _DB.items():
                if x in ns and k not in f:
                    f[k] = j
        if {"day", "ledger"} <= set(f) and ("debit" in f or "credit" in f):
            hi, cols = i, f
            break
    if hi is None:
        raise AuditError("Day book mein Date, Ledger / Particulars aur Debit / Credit column chahiye.")
    added = 0
    last = {"day": "", "voucher": "", "vtype": "", "narration": "", "user": "", "entered_on": ""}
    with _LOCK, session() as s:
        e = _eng(s, tenant, ref)
        for old in s.scalars(select(AuditEntry).where(AuditEntry.engagement_id == e.id)).all():
            s.delete(old)
        for r in rows[hi + 1:]:
            g = lambda k: r[cols[k]] if k in cols and cols[k] < len(r) else None  # noqa: E731
            led = str(g("ledger") or "").strip()
            from ..analyst.rules import clean_amount
            dr, cr = abs(clean_amount(g("debit"))), abs(clean_amount(g("credit")))
            d = _date(g("day"))
            if d:                                          # Tally prints date / voucher only on a voucher's first line
                last["day"] = d.isoformat()
                for k in ("voucher", "vtype", "narration", "user"):
                    last[k] = str(g(k) or "").strip()
                eo = _date(g("entered_on"))
                last["entered_on"] = eo.isoformat() if eo else ""
            if not led or not (dr or cr) or not last["day"] or led.lower() in ("total", "grand total"):
                continue
            s.add(AuditEntry(tenant_id=tenant, engagement_id=e.id, day=last["day"], voucher=last["voucher"][:40], vtype=last["vtype"][:30],
                             ledger=led[:160], debit=dr, credit=cr, narration=last["narration"][:300], user=last["user"][:60],
                             entered_on=last["entered_on"]))
            added += 1
        _log(e, f"Day book: {added} line")
        if e.status == "planning":
            e.status = "fieldwork"
    return {"lines": added, "say": f"Day book ki {added} line li. Ab 'jaanch chalao' kahiye: entries ki jaanch, cash niyam, TDS, Benford."}


# ---------------------------------------------------------------- materiality
def _statements(lines: list[dict]) -> dict:
    from ..analyst import rules as A
    from ..analyst.service import _balances
    return A.statements(_balances(lines))


def plan_materiality(tenant: str, ref: str, benchmark: str = "", pct=None, risk: str = "normal") -> dict[str, Any]:
    with _LOCK, session() as s:
        e = _eng(s, tenant, ref)
        lines = _j(e.tb, [])
        if not lines:
            raise AuditError("Pehle is saal ka trial balance chahiye.")
        st = _statements(lines)
        pl, bs = st["pl"], st["bs"]
        assets = bs["current_assets"] + bs["fixed_assets"]
        expenses = pl["cogs"] + pl["opex"] + pl["depreciation"] + pl["interest"]
        m = R.materiality(pl["profit_before_tax"], pl["revenue"], assets, bs["own_funds"], expenses, benchmark, pct, risk)
        e.materiality = json.dumps(m)
        f = _j(e.facts, {})
        f.setdefault("turnover", pl["revenue"])
        e.facts = json.dumps(f)
        _log(e, f"Materiality {_rs(m['overall'])} ({m['pct']}% of {m['benchmark_label']}), PM {_rs(m['performance'])}")
    m["say"] = (f"Materiality {_rs(m['overall'])} ({m['pct']:g}% {m['benchmark_label']}, kyunki {m['why']}); performance materiality "
                f"{_rs(m['performance'])}, isse chhota farak ({_rs(m['trivial'])}) nazarandaz. Yeh SA 320 ke hisaab se shuruaati andaza hai: aap badal sakti hain.")
    return m


# ---------------------------------------------------------------- tests
def _learned(s, tenant: str, client: str) -> dict[str, str]:
    return {r.key: r.note for r in s.scalars(select(AuditLearn).where(AuditLearn.tenant_id == tenant, AuditLearn.client == client)).all()}


def _add(s, out: list, tenant: str, e: AuditEngagement, learned: dict, area: str, rule: str, ref: str, amount: float, text: str,
         level: str = "warn") -> None:
    key = f"{rule}:{ref}"[:120]
    if f"{rule}:*" in learned or key in learned:
        return
    s.add(AuditFinding(tenant_id=tenant, engagement_id=e.id, area=area, rule=rule, ref=ref[:120], amount=round(amount, 2),
                       text=text[:1000], level=level, key=key))
    out.append({"area": area, "rule": rule, "ref": ref, "amount": round(amount, 2), "text": text, "level": level})


def run_tests(tenant: str, ref: str, conv: str = "") -> dict[str, Any]:
    """Analytics + journal-entry tests + cash rules + TDS + Benford + going concern. Findings replace earlier open ones."""
    with _LOCK, session() as s:
        e = _eng(s, tenant, ref)
        for old in s.scalars(select(AuditFinding).where(AuditFinding.engagement_id == e.id, AuditFinding.status == "open")).all():
            s.delete(old)
        learned = _learned(s, tenant, e.client)
        m = _j(e.materiality, {})
        pm = float(m.get("performance") or 0)
        trivial = float(m.get("trivial") or 0)
        out: list[dict] = []
        fy_end = date(e.fy_start + 1, 3, 31)
        # 1. analytics vs last year (SA 520)
        cur, prev = _j(e.tb, []), _j(e.tb_prev, [])
        if cur and prev and pm:
            pv = {x["name"].strip().lower(): x["dr"] - x["cr"] for x in prev}
            seen = set()
            for x in cur:
                k = x["name"].strip().lower()
                seen.add(k)
                v = R.variance(abs(x["dr"] - x["cr"]), abs(pv.get(k, 0.0)), pm)
                if v["flag"]:
                    _add(s, out, tenant, e, learned, "analytics", "variance", x["name"], v["diff"],
                         f"{x['name']}: {_rs(v['previous'])} se {_rs(v['current'])}" + (f" ({v['pct']:+.0f}%)" if v["pct"] is not None else " (naya ledger)")
                         + ": karan aur kaagaz poochhiye.")
            for k, amt in pv.items():
                if k not in seen and abs(amt) >= pm:
                    _add(s, out, tenant, e, learned, "analytics", "disappeared", k, amt, f"{k}: pichhle saal {_rs(abs(amt))}, is saal ledger hi nahi.")
        # 2. journal entries
        ents = s.scalars(select(AuditEntry).where(AuditEntry.engagement_id == e.id)).all()
        if ents:
            ledger_use = Counter(x.ledger.lower() for x in ents)
            users = Counter(x.user for x in ents if x.user)
            vouchers: dict[str, list[AuditEntry]] = defaultdict(list)
            for x in ents:
                vouchers[f"{x.vtype}|{x.voucher}|{x.day}"].append(x)
            dup_seen: dict[tuple, str] = {}
            cash_by_party_day: dict[tuple, float] = defaultdict(float)
            receipts_by_party_day: dict[tuple, float] = defaultdict(float)
            pay_by_party: dict[str, float] = defaultdict(float)
            for vk, xs in vouchers.items():
                vtype, vno, day = vk.split("|")
                d = _date(day)
                amt = max(sum(x.debit for x in xs), sum(x.credit for x in xs))
                vref = f"{vtype} {vno} ({day})".strip()
                if amt < trivial:
                    continue
                h = R.holiday(d) if d else ""
                if h:
                    _add(s, out, tenant, e, learned, "je", "holiday", vref, amt, f"{vref}: {h} ko entry, {_rs(amt)}.")
                if R.is_round(amt):
                    _add(s, out, tenant, e, learned, "je", "round", vref, amt, f"{vref}: gol rakam {_rs(amt)}.", "info")
                if e.approval_limit and R.just_below(amt, e.approval_limit):
                    _add(s, out, tenant, e, learned, "je", "below_limit", vref, amt,
                         f"{vref}: {_rs(amt)}, approval seema {_rs(e.approval_limit)} se thoda neeche.")
                if d and (fy_end - d).days <= 3 and (fy_end - d).days >= 0 and amt >= pm * 0.25:
                    _add(s, out, tenant, e, learned, "je", "year_end", vref, amt, f"{vref}: saal ke aakhri dino mein badi entry {_rs(amt)}.")
                eo = _date(xs[0].entered_on)
                if d and eo and (eo - d).days > 30:
                    _add(s, out, tenant, e, learned, "je", "backdated", vref, amt, f"{vref}: {eo.isoformat()} ko daali, tareekh {day} ({(eo - d).days} din peeche).", "bad")
                if d and eo and eo > fy_end and d <= fy_end:
                    _add(s, out, tenant, e, learned, "je", "post_close", vref, amt, f"{vref}: saal band hone ke baad daali gayi entry.")
                if not xs[0].narration and amt >= pm * 0.1:
                    _add(s, out, tenant, e, learned, "je", "no_narration", vref, amt, f"{vref}: narration khaali, {_rs(amt)}.", "info")
                u = xs[0].user
                if u and len(ents) >= 50 and users[u] / len(ents) < 0.02:
                    _add(s, out, tenant, e, learned, "je", "odd_user", vref, amt, f"{vref}: '{u}' ne daali (bahut kam entries wala user).")
                sig = (day, round(amt), tuple(sorted(x.ledger.lower() for x in xs)))
                if sig in dup_seen:
                    _add(s, out, tenant, e, learned, "je", "duplicate", vref, amt, f"{vref}: {dup_seen[sig]} jaisi hi entry (same din, rakam, ledger).", "bad")
                else:
                    dup_seen[sig] = vref
                ledgers = [x.ledger.lower() for x in xs]
                has_cash = any("cash" in l for l in ledgers)
                has_tds = any("tds" in l for l in ledgers)
                for x in xs:
                    l = x.ledger.lower()
                    if "cash" in l:
                        continue
                    if ledger_use[l] <= 2 and amt >= pm * 0.25 and "suspense" not in l:
                        _add(s, out, tenant, e, learned, "je", "rare_ledger", f"{x.ledger} / {vref}", amt, f"{x.ledger}: saal mein sirf {ledger_use[l]} baar, {_rs(amt)}.", "info")
                    if "suspense" in l:
                        _add(s, out, tenant, e, learned, "je", "suspense", f"{x.ledger} / {vref}", amt, f"{vref}: suspense mein {_rs(max(x.debit, x.credit))}.")
                    if has_cash and x.debit and d:
                        cash_by_party_day[(x.ledger, day)] += x.debit
                    if has_cash and x.credit and d:
                        receipts_by_party_day[(x.ledger, day)] += x.credit
                    nat = R.tds_nature(f"{x.ledger} {xs[0].narration}")
                    if nat and x.debit and not has_tds:
                        pay_by_party[f"{nat[0]}|{x.ledger}"] += x.debit
            loan_words = ("loan", "deposit", "advance from", "unsecured")
            for (party, day), amt in cash_by_party_day.items():
                lim = R.CASH_PAY_TRANSPORT if any(w in party.lower() for w in ("transport", "freight", "truck")) else R.CASH_PAY_LIMIT
                if any(w in party.lower() for w in loan_words):
                    if amt >= R.CASH_LOAN_LIMIT:
                        _add(s, out, tenant, e, learned, "cash", "cash_loan_repaid", f"{party} ({day})", amt,
                             f"{party}: {day} ko {_rs(amt)} loan / deposit cash mein lauta (Sec 220, purana 269T: Rs 20,000+ cash mein nahi).", "bad")
                elif amt > lim:
                    _add(s, out, tenant, e, learned, "cash", "cash_payment", f"{party} ({day})", amt,
                         f"{party}: {day} ko {_rs(amt)} cash mein diya; Rs {lim:,.0f} se upar ka kharcha Sec 34(3) (purana 40A(3)) mein disallow.", "bad")
            for (party, day), amt in receipts_by_party_day.items():
                if any(w in party.lower() for w in loan_words):
                    if amt >= R.CASH_LOAN_LIMIT:
                        _add(s, out, tenant, e, learned, "cash", "cash_loan_taken", f"{party} ({day})", amt,
                             f"{party}: {_rs(amt)} loan / deposit cash mein liya (Sec 219, purana 269SS).", "bad")
                elif amt >= R.CASH_RECEIPT_LIMIT:
                    _add(s, out, tenant, e, learned, "cash", "cash_receipt", f"{party} ({day})", amt,
                         f"{party}: {day} ko {_rs(amt)} cash mein liya; Rs 2 lakh ya zyada Sec 221 (purana 269ST) mein mana, penalty utni hi.", "bad")
            for k, amt in pay_by_party.items():
                nat_name, party = k.split("|", 1)
                row = next(r for r in R.TDS if r[0] == nat_name)
                if amt > row[3]:
                    _add(s, out, tenant, e, learned, "tds", f"tds_{nat_name}", party, amt,
                         f"{party}: saal mein {_rs(amt)} ({nat_name}), Rs {row[3]:,.0f} se upar, par TDS ki entry nahi dikhi "
                         f"(Sec 393, purana {row[5]}, ~{row[4]:g}%). Kaata gaya ho to dikhaiye; nahi to kharcha disallow ho sakta hai.")
            # Benford
            amounts = [max(x.debit, x.credit) for x in ents]
            bf = R.benford(amounts)
            if bf.get("ok") is False:
                _add(s, out, tenant, e, learned, "benford", "benford", "day book", 0,
                     f"Benford test: {bf['conformity']} (MAD {bf['mad']}); zyada aane wale ank: {bf['spikes'] or '-'}. In ank se shuru hone wali rakamein dekhiye.")
        else:
            bf = None
        # going concern (SA 570)
        if cur:
            st = _statements(cur)
            if st["bs"]["own_funds"] < 0:
                _add(s, out, tenant, e, learned, "gc", "net_worth_negative", "balance sheet", st["bs"]["own_funds"], "Net worth minus mein: going concern (SA 570) dekhiye.", "bad")
            if st["bs"]["current_liabilities"] and st["bs"]["current_assets"] < st["bs"]["current_liabilities"]:
                _add(s, out, tenant, e, learned, "gc", "current_ratio_below_1", "balance sheet", st["bs"]["current_assets"] - st["bs"]["current_liabilities"],
                     "Current assets current liabilities se kam: cash ki tangi, going concern dekhiye.")
            if st["pl"]["profit_before_tax"] < 0:
                _add(s, out, tenant, e, learned, "gc", "losses", "P&L", st["pl"]["profit_before_tax"], f"Ghata {_rs(-st['pl']['profit_before_tax'])}.")
        # applicability
        facts = _j(e.facts, {})
        if e.kind == "tax" or facts.get("turnover"):
            ta = R.tax_audit("profession" if facts.get("profession") else "business", float(facts.get("turnover") or 0),
                             facts.get("cash_receipts_pct"), facts.get("cash_payments_pct"), e.fy_start)
            facts["tax_audit"] = ta
        if e.entity in ("private", "public", "company"):
            kind = "private" if e.entity in ("private", "company") else "public"
            st = _statements(cur) if cur else {"bs": {"own_funds": 0, "loans": 0}, "pl": {"revenue": 0}}
            facts["caro"] = R.caro(kind, st["bs"]["own_funds"], st["bs"]["loans"], st["pl"]["revenue"], bool(facts.get("holding_or_sub")),
                                   facts.get("paid_up"), date(e.fy_start + 1, 3, 31))
            facts["audit_trail"] = R.AUDIT_TRAIL
        e.facts = json.dumps(facts)
        if e.status == "planning":
            e.status = "fieldwork"
        _log(e, f"Jaanch chali: {len(out)} baatein mili")
        eid, client = e.id, e.client
    by_area = Counter(x["area"] for x in out)
    say = f"{client}: jaanch poori, {len(out)} baatein dekhni hain (" + ", ".join(f"{k} {v}" for k, v in by_area.items()) + ")." if out else \
        f"{client}: jaanch poori, kuch ajeeb nahi mila."
    bad = [x for x in out if x["level"] == "bad"][:3]
    if bad:
        say += " Sabse zaroori: " + " ".join(x["text"] for x in bad)
    _publish(tenant, "audit_update", {"engagement": eid, "findings": len(out)}, conv)
    return {"findings": out, "benford": bf, "say": say}


def findings(tenant: str, ref: str, status: str = "") -> list[dict]:
    with session() as s:
        e = _eng(s, tenant, ref)
        q = select(AuditFinding).where(AuditFinding.engagement_id == e.id)
        if status:
            q = q.where(AuditFinding.status == status)
        order = {"bad": 0, "warn": 1, "info": 2}
        rows = s.scalars(q).all()
        return sorted([{"id": f.id, "area": f.area, "rule": f.rule, "ref": f.ref, "amount": f.amount, "text": f.text, "level": f.level,
                        "status": f.status, "note": f.note} for f in rows], key=lambda x: (order.get(x["level"], 3), -abs(x["amount"])))


def close_finding(tenant: str, finding_id: str, status: str, note: str = "", remember: bool = False, by: str = "owner") -> dict[str, Any]:
    """She closes a finding: explained | adjusted | reported (goes into the report) | asked. remember=true: the same
    pattern for this client is not raised next time."""
    if status not in ("explained", "adjusted", "reported", "asked", "open"):
        raise AuditError("status: explained | adjusted | reported | asked")
    with session() as s:
        f = s.get(AuditFinding, finding_id)
        if not f or f.tenant_id != tenant:
            raise AuditError("Baat nahi mili.")
        f.status, f.note = status, str(note or "")[:1000]
        e = s.get(AuditEngagement, f.engagement_id)
        _log(e, f"{f.rule} {f.ref}: {status}" + (f" ({note[:80]})" if note else ""), by)
        if remember and status == "explained":
            s.add(AuditLearn(tenant_id=tenant, client=e.client, key=f.key, note=note[:500]))
    return {"say": f"Note kiya: {status}." + (" Agli baar yeh baat nahi uthegi." if remember else "")}


# ---------------------------------------------------------------- sampling (SA 530)
def make_sample(tenant: str, ref: str, population: str = "", method: str = "mus", size: int = 25, confidence: int = 95,
                seed: int | None = None) -> dict[str, Any]:
    """population: a ledger name (its vouchers), 'sales' / 'purchases' / 'expenses' (vouchers of that kind), or empty = all."""
    with _LOCK, session() as s:
        e = _eng(s, tenant, ref)
        m = _j(e.materiality, {})
        ents = s.scalars(select(AuditEntry).where(AuditEntry.engagement_id == e.id)).all()
        if not ents:
            raise AuditError("Pehle day book chahiye.")
        pop = (population or "").strip().lower()
        vouchers: dict[str, float] = defaultdict(float)
        for x in ents:
            if pop in ("sales", "purchases", "purchase", "payment", "receipt", "journal") and pop[:5] not in x.vtype.lower():
                continue
            if pop and pop not in ("sales", "purchases", "purchase", "payment", "receipt", "journal") and pop not in x.ledger.lower():
                continue
            vouchers[f"{x.vtype} {x.voucher} ({x.day})"] += x.debit if not pop or pop not in x.ledger.lower() else max(x.debit, x.credit)
        items = [(k, v) for k, v in vouchers.items() if v > 0]
        if not items:
            raise AuditError(f"'{population}' mein koi voucher nahi mila.")
        sd = seed if seed is not None else (abs(hash(e.id)) % 9973) + 1
        if method == "mus" and m.get("performance"):
            r = R.mus_select(items, float(m["performance"]), confidence, sd)
            chosen = [(x, "key") for x in r["key_items"]] + [(x, "sample") for x in r["sampled"]]
            info = {"method": "monetary unit", "interval": r["interval"], "seed": sd, "confidence": confidence, "population_value": r["population_value"]}
        else:
            picks = R.random_select([k for k, _v in items], size, sd)
            chosen = [(x, "sample") for x in picks]
            info = {"method": "random", "seed": sd, "size": size}
        amt = dict(items)
        for old in s.scalars(select(AuditSample).where(AuditSample.engagement_id == e.id, AuditSample.population == (population or "all"))).all():
            s.delete(old)
        for refv, why in chosen:
            s.add(AuditSample(tenant_id=tenant, engagement_id=e.id, population=(population or "all")[:80], ref=refv[:60], amount=amt.get(refv, 0.0), why=why))
        _log(e, f"Sample ({info['method']}, seed {sd}): {len(chosen)} voucher from {population or 'sab'}")
    return {**info, "count": len(chosen), "items": [{"ref": r, "why": w, "amount": round(amt.get(r, 0.0))} for r, w in chosen],
            "say": f"{len(chosen)} voucher chune ({info['method']}, seed {sd}: dobara chalane par wahi aayenge). Inke bill / kaagaz mangwaiye."}


def sample_result(tenant: str, sample_id: str, status: str, note: str = "") -> dict:
    if status not in ("ok", "exception", "pending"):
        raise AuditError("status: ok | exception")
    with session() as s:
        x = s.get(AuditSample, sample_id)
        if not x or x.tenant_id != tenant:
            raise AuditError("Sample item nahi mila.")
        x.status, x.note = status, str(note)[:500]
        e = s.get(AuditEngagement, x.engagement_id)
        if status == "exception":
            s.add(AuditFinding(tenant_id=tenant, engagement_id=e.id, area="sample", rule="vouching", ref=x.ref, amount=x.amount,
                               text=f"Vouching mein galti: {x.ref} - {note}", level="bad", key=f"vouching:{x.ref}"[:120]))
        return {"say": f"{x.ref}: {status}."}


def samples(tenant: str, ref: str) -> list[dict]:
    with session() as s:
        e = _eng(s, tenant, ref)
        return [{"id": x.id, "population": x.population, "ref": x.ref, "amount": x.amount, "why": x.why, "status": x.status, "note": x.note}
                for x in s.scalars(select(AuditSample).where(AuditSample.engagement_id == e.id)).all()]


# ---------------------------------------------------------------- confirmations (SA 505)
def make_confirmations(tenant: str, ref: str, top: int = 10, min_amount=None) -> dict[str, Any]:
    with _LOCK, session() as s:
        e = _eng(s, tenant, ref)
        lines = _j(e.tb, [])
        m = _j(e.materiality, {})
        floor = float(min_amount) if min_amount not in (None, "") else float(m.get("performance") or 0) * 0.5
        rows = [(x["name"], "debtor" if x["category"] == "debtors" else "creditor", abs(x["dr"] - x["cr"]))
                for x in lines if x["category"] in ("debtors", "creditors") and not x.get("is_group")]
        rows = [r for r in rows if r[2] >= floor]
        rows.sort(key=lambda r: -r[2])
        for old in s.scalars(select(AuditConfirm).where(AuditConfirm.engagement_id == e.id, AuditConfirm.status == "draft")).all():
            s.delete(old)
        made = []
        for kind in ("debtor", "creditor"):
            for name, k, bal in [r for r in rows if r[1] == kind][:top]:
                c = AuditConfirm(tenant_id=tenant, engagement_id=e.id, party=name[:160], kind=k, balance=round(bal, 2))
                s.add(c)
                made.append((name, k, bal))
        _log(e, f"Confirmation draft: {len(made)}")
    return {"count": len(made), "say": f"{len(made)} party ke liye balance confirmation letter ka draft (SA 505). Bhejne se pehle aapki haan."}


def confirmation_letters(tenant: str, ref: str) -> dict[str, Any]:
    with session() as s:
        e = _eng(s, tenant, ref)
        out = []
        fy_end = f"31-03-{e.fy_start + 1}"
        for c in s.scalars(select(AuditConfirm).where(AuditConfirm.engagement_id == e.id)).all():
            side = "receivable from you" if c.kind == "debtor" else "payable to you"
            out.append({"id": c.id, "party": c.party, "kind": c.kind, "balance": c.balance, "status": c.status,
                        "letter": (f"To,\n{c.party}\n\nSub: Confirmation of balance as on {fy_end}\n\nDear Sir / Madam,\n\nOur auditors are "
                                   f"examining our accounts. As per our books, an amount of Rs {c.balance:,.2f} was {side} as on {fy_end}. "
                                   "Please confirm this balance directly to our auditors, or state the difference with details, by signing "
                                   "and returning this letter.\n\nFor " + e.client + "\n\nAuthorised Signatory\n\n---\nWe confirm the above "
                                   "balance / our balance is Rs ________ (details attached).\n\nSignature and stamp: ________")})
    return {"letters": out}


def approve_confirmations(tenant: str, ref: str, approve: bool = True) -> dict[str, Any]:
    with session() as s:
        e = _eng(s, tenant, ref)
        n = 0
        for c in s.scalars(select(AuditConfirm).where(AuditConfirm.engagement_id == e.id, AuditConfirm.status == "draft")).all():
            c.status = "approved" if approve else "draft"
            n += 1
        _log(e, f"Confirmations {'approved' if approve else 'roke'}: {n}", "owner")
    return {"approved": n if approve else 0,
            "say": f"{n} letter bhejne ki haan mili: client ke letterhead par client se bhijwaiye; jawab seedha auditor ko aaye (SA 505)." if approve else "Roka."}


def confirmation_reply(tenant: str, confirm_id: str, amount, note: str = "") -> dict:
    with session() as s:
        c = s.get(AuditConfirm, confirm_id)
        if not c or c.tenant_id != tenant:
            raise AuditError("Confirmation nahi mila.")
        from ..analyst.rules import clean_amount
        c.reply = abs(clean_amount(amount))
        diff = round((c.reply or 0) - c.balance, 2)
        c.status = "agreed" if abs(diff) <= 1 else "differs"
        c.note = str(note)[:500]
        if c.status == "differs":
            s.add(AuditFinding(tenant_id=tenant, engagement_id=c.engagement_id, area="confirm", rule="confirm_diff", ref=c.party,
                               amount=diff, text=f"{c.party}: books {_rs(c.balance)}, party kehti hai {_rs(c.reply)} ({_rs(diff)} farak). Reconciliation maangiye.",
                               level="warn", key=f"confirm_diff:{c.party}"[:120]))
        return {"say": f"{c.party}: {'mil gaya' if c.status == 'agreed' else f'{_rs(diff)} ka farak'}."}


# ---------------------------------------------------------------- queries, letters, working papers
def query_list(tenant: str, ref: str) -> dict[str, Any]:
    """Open findings as a client query list (draft; sending waits for her Haan)."""
    fs = [f for f in findings(tenant, ref) if f["status"] in ("open", "asked") and f["level"] != "info"]
    with session() as s:
        e = _eng(s, tenant, ref)
        client = e.client
        fy = f"{e.fy_start}-{str(e.fy_start + 1)[2:]}"
    lines = [f"{i}. {f['text']}" for i, f in enumerate(fs, 1)]
    text = (f"{client} - audit FY {fy}: kripya neeche ki baaton par jawab aur kaagaz bhejiye.\n\n" + "\n".join(lines)) if lines else ""
    return {"count": len(fs), "draft": text, "say": f"{len(fs)} sawaal ki list taiyaar (draft). Client ko bhejne se pehle aapki haan." if fs else "Koi khula sawaal nahi."}


def mark_asked(tenant: str, ref: str) -> dict:
    with session() as s:
        e = _eng(s, tenant, ref)
        n = 0
        for f in s.scalars(select(AuditFinding).where(AuditFinding.engagement_id == e.id, AuditFinding.status == "open",
                                                      AuditFinding.level != "info")).all():
            f.status = "asked"
            n += 1
        _log(e, f"{n} sawaal client se poochhe (owner ki haan)", "owner")
    return {"say": f"{n} sawaal 'poochha' mark kiye."}


def rep_letter(tenant: str, ref: str) -> dict[str, Any]:
    """Management representation letter draft (SA 580): the client signs; the auditor decides what it must contain."""
    with session() as s:
        e = _eng(s, tenant, ref)
        fy_end = f"31 March {e.fy_start + 1}"
    text = (f"To,\nThe Auditors\n\nSub: Management representation for the year ended {fy_end}\n\nWe confirm, to the best of our knowledge "
            "and belief:\n1. We have fulfilled our responsibility for preparing the financial statements and have given you all books, "
            "records and information.\n2. All transactions have been recorded and are reflected in the financial statements.\n3. We have "
            "disclosed all known instances of fraud or suspected fraud, and all related-party relationships and transactions.\n4. There are no "
            "unrecorded liabilities, claims or contingencies other than those disclosed.\n5. Events after the balance sheet date requiring "
            "adjustment or disclosure have been adjusted or disclosed.\n6. The audit trail (edit log) feature of our accounting software was "
            "operating throughout the year and has not been disabled or tampered with.\n\nFor " + e.client + "\n\nAuthorised Signatory\nDate:")
    return {"draft": text, "say": "Representation letter (SA 580) ka draft: client sign karega. Zaroorat ke hisaab se points badaliye."}


def workpapers(tenant: str, ref: str, conv: str = "") -> dict[str, Any]:
    import openpyxl
    from .. import chatfiles
    with session() as s:
        e = _eng(s, tenant, ref)
        v = _view(s, e)
        tb, tbp = _j(e.tb, []), _j(e.tb_prev, [])
        logs = _j(e.log, [])
    fs = findings(tenant, ref)
    sm = samples(tenant, ref)
    cf = confirmation_letters(tenant, ref)["letters"]
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Planning"
    ws.append([v["client"], v["kind_label"], f"FY {v['fy']}"])
    ws.append(["Prepared by Hoola (assistant). Reviewed and concluded by the engagement partner."])
    m = v["materiality"]
    for k in ("benchmark_label", "base", "pct", "overall", "performance_pct", "performance", "trivial", "why"):
        ws.append([k, m.get(k)])
    for k, val in v["facts"].items():
        ws.append([k, json.dumps(val, ensure_ascii=False) if isinstance(val, (dict, list)) else val])
    w = wb.create_sheet("TB")
    w.append(["Ledger", "Head", "Debit", "Credit", "Last year (net)"])
    pv = {x["name"].strip().lower(): x["dr"] - x["cr"] for x in tbp}
    for x in tb:
        w.append([x["name"], x["category"], x["dr"], x["cr"], pv.get(x["name"].strip().lower())])
    w = wb.create_sheet("Findings")
    w.append(["Area", "Rule", "Ref", "Amount", "Level", "Finding", "Status", "Her note"])
    for f in fs:
        w.append([f["area"], f["rule"], f["ref"], f["amount"], f["level"], f["text"], f["status"], f["note"]])
    w = wb.create_sheet("Sample")
    w.append(["Population", "Voucher", "Amount", "Why", "Result", "Note"])
    for x in sm:
        w.append([x["population"], x["ref"], x["amount"], x["why"], x["status"], x["note"]])
    w = wb.create_sheet("Confirmations")
    w.append(["Party", "Kind", "Books", "Status"])
    for c in cf:
        w.append([c["party"], c["kind"], c["balance"], c["status"]])
    w = wb.create_sheet("Log")
    for lg in logs:
        w.append([lg["at"], lg["by"], lg["what"]])
    bio = io.BytesIO()
    wb.save(bio)
    f = chatfiles.store(tenant, f"WP_{re.sub(r'[^A-Za-z0-9]+', '_', v['client'])}_{v['fy']}.xlsx", bio.getvalue(),
                        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    _publish(tenant, "pc_file", {**f, "source": "audit", "title": f.get("name", "")}, conv)
    return {"file": f, "say": "Working papers (Planning, TB, Findings, Sample, Confirmations, Log) Excel mein taiyaar (SA 230)."}


def status_summary(tenant: str) -> dict[str, Any]:
    on = today(tenant)
    rows = []
    for e in engagements(tenant):
        if e["status"] == "done":
            continue
        ta = (e["facts"] or {}).get("tax_audit") or {}
        due = ta.get("due")
        rows.append({"client": e["client"], "kind": e["kind_label"], "fy": e["fy"], "open": e["findings"]["open"], "due": due,
                     "days": (date.fromisoformat(due) - on).days if due else None})
    txt = "; ".join(f"{r['client']} ({r['kind']}): {r['open']} khule sawaal" + (f", due {r['due']}" if r["due"] else "") for r in rows[:6])
    return {"engagements": rows, "text": txt or "Koi chalu audit nahi."}
