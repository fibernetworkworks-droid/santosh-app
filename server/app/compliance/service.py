"""Compliance Sahayak service: Hoola as a business's compliance officer, end to end.

Her business facts (entity, state, staff, turnover, GSTIN, TAN, sector, food / factory / pollution / fire, licences) -> which statutory
duties apply and why (rules.applicability) -> a compliance calendar of items with due dates, owners, evidence and status, rebuilt every
day -> reminders before the due date and penalty andaaza after it (starters with sources) -> a monthly compliance certificate draft
she signs off.

Read, don't duplicate: when she already uses the GST Sahayak (gstp), Payroll (payroll), Daftar renewals (office) or the Company
Secretary Sahayak (cosec), their due items are SHOWN here as linked rows (status from there) and not made again.

In code: an item becomes 'done' only with an evidence file (challan / ack / certificate) or, for the owner, her own word. 'Portal par
file karo' and the certificate sign-off are cards whose hash only her app tap (`by_tap=True`, API only) decides. Hoola never logs in,
pays or submits on a portal: she does login / OTP / payment / submit herself (the Computer Operator may open forms under her control).
"""
from __future__ import annotations

import hashlib
import io
import json
import logging
import re
import threading
import urllib.parse
from datetime import date, datetime, timedelta, timezone
from typing import Any

from sqlalchemy import select

from ..db import now, session
from . import rules as R
from .models import CompCert, CompFiling, CompItem, CompOnce, CompSettings

log = logging.getLogger("santosh.compliance")
_LOCK = threading.RLock()


class CompError(Exception):
    pass


# ---------------------------------------------------------------- helpers
def _j(v, default):
    try:
        x = json.loads(v) if isinstance(v, str) and v else default
        return x if isinstance(x, type(default)) else default
    except (ValueError, TypeError):
        return default


def _d(v) -> str:
    return json.dumps(v, ensure_ascii=False, default=str)


def clock() -> datetime:
    return now()


def _tz(tenant: str):
    from ..life import tz_of
    return tz_of(tenant)


def today(tenant: str) -> date:
    return clock().astimezone(_tz(tenant)).date()


def _date(v) -> date | None:
    if isinstance(v, date):
        return v
    s = R._iso(v)
    return date.fromisoformat(s) if s else None


def _clean(text, n: int = 400) -> str:
    """Her / outside words: secrets never kept; Aadhaar / card numbers / PAN masked."""
    t = str(text or "")
    t = re.sub(r"\b(otp|pin|cvv|password|passcode|upi pin)\b\s*(?:is|hai|:|=)?\s*\S+", lambda m: f"{m.group(1)} [hataya]", t, flags=re.I)
    t = re.sub(r"\b(?:\d[ -]?){13,19}\b", "[hataya]", t)
    t = re.sub(r"\b\d{4}\s?\d{4}\s?\d{4}\b", "[hataya]", t)
    t = re.sub(r"\b[A-Z]{5}\d{4}[A-Z]\b", "[PAN hataya]", t)
    return t.strip()[:n]


def _hash(show: dict) -> str:
    return hashlib.sha256(json.dumps(show, sort_keys=True, ensure_ascii=False, default=str).encode()).hexdigest()


def _publish(tenant: str, kind: str, payload: dict, conv: str = "") -> None:
    from ..events import bus
    try:
        if not conv:
            from ..net.service import home_conv
            conv = home_conv(tenant)
        bus.publish(conv, kind, payload)
    except Exception:
        log.exception("compliance publish failed")


def is_owner(actor: str) -> bool:
    a = actor or "owner"
    return a == "owner" or a.startswith("approved:")


def _who(tenant: str, actor: str) -> str:
    if is_owner(actor):
        return "owner"
    try:
        from ..employee.service import actor_label
        return actor_label(tenant, actor)[:80]
    except Exception:
        return (actor or "")[:80]


def _st(s, tenant: str) -> CompSettings:
    r = s.get(CompSettings, tenant)
    if r is None:
        r = CompSettings(tenant_id=tenant)
        s.add(r)
        s.flush()
    return r


def _next(s, tenant: str, key: str) -> int:
    st = _st(s, tenant)
    seq = _j(st.seq, {})
    seq[key] = int(seq.get(key) or 0) + 1
    st.seq = _d(seq)
    return seq[key]


def _once(s, tenant: str, key: str) -> bool:
    k = f"{tenant}:{key}"[:200]
    if s.get(CompOnce, k) is not None:
        return False
    s.add(CompOnce(key=k, tenant_id=tenant))
    return True


def _file_meta(tenant: str, file_id: str) -> dict | None:
    try:
        from .. import chatfiles
        return chatfiles._meta(tenant, file_id)
    except Exception:
        return None


def _phone(v) -> str:
    dg = re.sub(r"\D", "", str(v or ""))
    if len(dg) == 12 and dg.startswith("91"):
        dg = dg[2:]
    return dg if len(dg) == 10 and dg[0] in "6789" else ""


def _wa(phone: str, text: str) -> str:
    p = _phone(phone)
    return f"https://wa.me/91{p}?text={urllib.parse.quote(text)}" if p else f"https://wa.me/?text={urllib.parse.quote(text)}"


def _rs(v) -> str:
    if v is None:
        return ""
    try:
        from ..india.money import inr
        return inr(v)
    except Exception:
        return f"Rs {v:,.0f}"


def _log(i: CompItem, what: str, by: str = "hoola") -> None:
    lg = _j(i.log, [])
    lg.append({"at": now().isoformat(timespec="seconds"), "by": by, "what": what[:300]})
    i.log = _d(lg[-40:])


# ---------------------------------------------------------------- settings + facts
_SET_KEYS = ("firm", "signatory", "designation", "owners", "team", "windows", "morning_time")


def _windows(st: CompSettings) -> list[int]:
    out = []
    for x in str(st.windows or "").split(","):
        try:
            v = int(x.strip())
        except ValueError:
            continue
        if 0 <= v <= 90:
            out.append(v)
    return sorted(set(out), reverse=True) or [7, 3, 1, 0]


def _set_view(r: CompSettings) -> dict:
    return {"firm": r.firm, "signatory": r.signatory, "designation": r.designation, "facts": _j(r.facts, {}),
            "owners": _j(r.owners, {}), "team": _j(r.team, []), "windows": _windows(r), "morning_time": r.morning_time,
            "entities": R.ENTITIES, "areas": R.AREAS, "licence_keys": R.LICENCE_KEYS, "pollution": list(R.POLLUTION),
            "food_kinds": list(R.FOOD_KINDS), "check": R.CHECK_LINE}


def settings(tenant: str) -> dict:
    with session() as s:
        return _set_view(_st(s, tenant))


def save_settings(tenant: str, facts: dict | None = None, **kw) -> dict:
    with _LOCK, session() as s:
        st = _st(s, tenant)
        if facts is not None:
            cur = _j(st.facts, {})
            new = dict(cur)
            for k, v in dict(facts).items():
                if k == "licences":
                    lic = dict(cur.get("licences") or {})
                    lic.update({a: b for a, b in dict(v or {}).items()})
                    new["licences"] = {a: b for a, b in lic.items() if b}
                elif k in R.FACT_KEYS:
                    new[k] = v
            try:
                st.facts = _d(R.clean_facts(new))
            except ValueError as e:
                raise CompError(str(e))
        if "firm" in kw:
            st.firm = _clean(kw["firm"], 160)
        if "signatory" in kw:
            st.signatory = _clean(kw["signatory"], 120)
        if "designation" in kw:
            st.designation = _clean(kw["designation"], 80)
        if "owners" in kw and kw["owners"] is not None:
            st.owners = _d({k: _clean(v, 80) for k, v in dict(kw["owners"]).items() if k in R.AREAS and v})
        if "team" in kw and kw["team"] is not None:
            team = []
            for m in list(kw["team"])[:40]:
                m = dict(m or {})
                if m.get("name"):
                    team.append({"name": _clean(m["name"], 80), "phone": _phone(m.get("phone"))})
            st.team = _d(team)
        if "windows" in kw and kw["windows"] is not None:
            w = kw["windows"]
            w = w if isinstance(w, str) else ",".join(str(x) for x in w)
            st.windows = ",".join(str(x) for x in sorted({int(x) for x in re.findall(r"\d+", w) if int(x) <= 90}, reverse=True))[:40] or "7,3,1,0"
        if kw.get("morning_time"):
            if not re.fullmatch(r"([01]\d|2[0-3]):[0-5]\d", str(kw["morning_time"])):
                raise CompError("morning_time HH:MM")
            st.morning_time = str(kw["morning_time"])
        st.last_build = ""
        v = _set_view(st)
    if facts is not None:
        build(tenant)
    return v


def facts_of(tenant: str) -> dict:
    with session() as s:
        return _j(_st(s, tenant).facts, {})


# ---------------------------------------------------------------- neighbours (read only)
_LIC_WORDS = {"fssai": ("fssai", "food licence", "food license"), "consent": ("consent", "pollution", "pcb", "cto"),
              "fire": ("fire",), "shop": ("shop", "establishment", "gumasta"), "trade": ("trade licence", "trade license"),
              "factory": ("factory",), "contract": ("contract labour",)}


def linked(tenant: str) -> dict:
    """Which neighbour Sahayaks she already uses (their due items are read, never copied)."""
    out: dict[str, Any] = {"gstp": False, "payroll": False, "cosec": False, "office": [], "office_licences": []}
    try:
        from ..gstp import service as G
        out["gstp"] = bool(G.clients(tenant))
    except Exception:
        log.debug("gstp not readable", exc_info=True)
    try:
        from ..payroll import service as P
        out["payroll"] = bool(P.employees(tenant))
    except Exception:
        log.debug("payroll not readable", exc_info=True)
    try:
        from ..office import service as O
        rs = O.renewals(tenant)
        out["office"] = rs
        lic = []
        for k, words in _LIC_WORDS.items():
            if any(any(w in (r.get("name") or "").lower() for w in words) for r in rs):
                lic.append(k)
        out["office_licences"] = lic
    except Exception:
        log.debug("office not readable", exc_info=True)
    try:
        from ..cosec import service as CS  # neighbour built in the same suite; optional
        fn = getattr(CS, "calendar", None)
        out["cosec"] = bool(fn and _rows_of(fn(tenant)))
    except Exception:
        out["cosec"] = False
    return out


def _rows_of(res) -> list[dict]:
    if isinstance(res, dict):
        rows = list(res.get("due") or []) + list(res.get("late") or []) + list(res.get("items") or [])
    elif isinstance(res, list):
        rows = res
    else:
        rows = []
    return [r for r in rows if isinstance(r, dict)]


def _covered(entry: dict, lk: dict) -> bool:
    link = entry.get("link")
    if link == "gstp":
        return bool(lk.get("gstp"))
    if link == "payroll":
        return bool(lk.get("payroll"))
    if link == "cosec":
        return bool(lk.get("cosec"))
    if link == "office" and entry["freq"] == "expiry":
        return entry["due"][0] in (lk.get("office_licences") or [])
    return False


def linked_rows(tenant: str, lk: dict | None = None, ahead: int = 60) -> list[dict]:
    """Due items of the neighbours, normalised; status comes from there (done there, not here)."""
    lk = lk or linked(tenant)
    on = today(tenant)
    rows: list[dict] = []

    def add(src: str, label: str, due, status: str = "", note: str = "") -> None:
        dd = _date(due)
        rows.append({"source": src, "source_label": {"gstp": "GST Sahayak", "payroll": "Payroll", "office": "Daftar renewals",
                                                     "cosec": "Company Secretary"}.get(src, src),
                     "label": str(label)[:160], "due": dd.isoformat() if dd else "", "due_nice": R.nice(dd),
                     "days": (dd - on).days if dd else None, "late": bool(dd and dd < on and status not in ("filed", "done", "paid")),
                     "status": status, "note": note})
    if lk.get("gstp"):
        try:
            from ..gstp import service as G
            cal = G.calendar(tenant, ahead)
            for x in list(cal.get("late") or []) + list(cal.get("due") or []):
                add("gstp", f"{x.get('client', '')} {x.get('label', '')} {x.get('period_nice', '')}".strip(), x.get("due"), x.get("status", ""))
        except Exception:
            log.debug("gstp calendar", exc_info=True)
    if lk.get("payroll"):
        try:
            from ..payroll import service as P
            cal = P.calendar(tenant, ahead)
            for x in list(cal.get("late") or []) + list(cal.get("due") or []):
                add("payroll", f"{x.get('label', '')} {x.get('period_nice', '')}".strip(), x.get("due"), x.get("run_status", ""))
        except Exception:
            log.debug("payroll calendar", exc_info=True)
    for r in lk.get("office") or []:
        if r.get("days_left") is not None and r["days_left"] <= ahead and r.get("kind") in ("licence", "certificate", "tax", "insurance",
                                                                                             "amc", "rent", "vehicle", "other", "domain",
                                                                                             "subscription", "bill"):
            add("office", f"{r.get('name', '')} ({r.get('kind_label', '')})", r.get("due_on"), r.get("status", ""),
                "Daftar Sahayak mein 'renewal done' kijiye")
    if lk.get("cosec"):
        try:
            from ..cosec import service as CS
            for x in _rows_of(CS.calendar(tenant)):
                add("cosec", x.get("label") or x.get("title") or x.get("form") or "", x.get("due") or x.get("due_on"), x.get("status", ""))
        except Exception:
            log.debug("cosec calendar", exc_info=True)
    rows.sort(key=lambda r: (r["due"] or "9999"))
    return rows


# ---------------------------------------------------------------- applicability + calendar build
def applicability(tenant: str) -> dict:
    f = facts_of(tenant)
    lk = linked(tenant)
    items = R.applicability(f)
    for x in items:
        e = R.BY_KEY[x["key"]]
        x["covered_by"] = (e.get("link") if _covered(e, lk) else "")
    n = sum(1 for x in items if x["applies"])
    say = (f"Aapke business par {n} statutory kaam lagte hain." if f.get("entity") else "Pehle business ki jaankari chahiye.")
    q = R.questions(f)
    if q:
        say += f" {len(q)} sawaal baaki: {q[0]}"
    return {"facts": f, "items": items, "questions": q, "linked": {k: (bool(v) if not isinstance(v, list) else len(v)) for k, v in lk.items()},
            "say": say + " " + R.CHECK_LINE, "check": R.CHECK_LINE}


def build(tenant: str, ahead: int = 92, back: int = 62) -> dict:
    """Make the calendar items for every duty that applies (idempotent; linked neighbours are read, not copied)."""
    f = facts_of(tenant)
    lk = linked(tenant)
    on = today(tenant)
    start, end = on - timedelta(days=back), on + timedelta(days=ahead)
    made = 0
    with _LOCK, session() as s:
        st = _st(s, tenant)
        owners = _j(st.owners, {})
        have = {(i.key, i.period): i for i in s.scalars(select(CompItem).where(CompItem.tenant_id == tenant)).all()}
        for e in R.CATALOGUE:
            ok, _why = (False, "")
            try:
                ok, _why = e["applies"](f)
            except Exception:
                ok = False
            if not ok or _covered(e, lk):
                continue
            inst = R.instances(e, f, start, end)
            if any(p not in ("exp:?", "AGM?") for p, _, _ in inst):
                for kk in [kk for kk in have if kk[0] == e["key"] and kk[1] in ("exp:?", "AGM?")]:
                    if have[kk].status == "open":       # she gave the date now: the 'tell me the date' item goes
                        s.delete(have.pop(kk))
            for period, due, suffix in inst:
                k = (e["key"], period)
                if k in have:
                    cur = have[k]
                    if cur.status == "open" and due and cur.due != due.isoformat() and not cur.due_override:
                        cur.due = due.isoformat()
                    continue
                if period in ("exp:?", "AGM?", "once", "her"):
                    # one 'tell me the date / status' item per duty, not per window
                    if any(kk[0] == e["key"] for kk in have):
                        continue
                i = CompItem(tenant_id=tenant, key=e["key"], area=e["area"], period=period,
                             label=f"{e['label']} ({suffix})" if suffix else e["label"], due=due.isoformat() if due else "",
                             owner=owners.get(e["area"], ""))
                i.number = f"C-{_next(s, tenant, 'item'):04d}"
                _log(i, "calendar mein joda")
                s.add(i)
                have[k] = i
                made += 1
        st.last_build = on.isoformat()
    return {"made": made, "say": f"Calendar taiyaar: {made} naye kaam jode." if made else "Calendar pehle se taiyaar hai."}


def _eff_due(i: CompItem) -> date | None:
    return _date(i.due_override or i.due)


def _state_of(i: CompItem, on: date) -> str:
    if i.status in ("done", "na"):
        return i.status
    dd = _eff_due(i)
    if dd is None:
        return "no_date"
    days = (dd - on).days
    if days < 0:
        return "late"
    return "due" if days <= 7 else "upcoming"


def _item_view(i: CompItem, on: date, facts: dict | None = None, full: bool = False) -> dict:
    e = R.BY_KEY.get(i.key) or {"label": i.label, "area": i.area, "freq": "her", "penalty": {"text": ""}}
    dd = _eff_due(i)
    st = _state_of(i, on)
    days = (dd - on).days if dd else None
    v = {"id": i.id, "number": i.number, "key": i.key, "area": i.area, "area_label": R.AREAS.get(i.area, i.area), "period": i.period,
         "label": i.label, "due": dd.isoformat() if dd else "", "due_nice": R.nice(dd), "due_own": bool(i.due_override),
         "days": days, "state": st, "status": i.status, "owner": i.owner, "done_how": i.done_how, "done_on": i.done_on,
         "done_by": i.done_by, "ack": i.ack, "evidence": _j(i.evidence, []), "note": i.note, "portal": e.get("portal", ""),
         "evidence_hint": e.get("evidence", "saboot ki file"), "source": R.source(e) if i.key in R.BY_KEY else {"title": "aapka apna kaam", "url": ""},
         "verify": bool(e.get("verify")) or (i.key in R.BY_KEY and not e.get("src")), "custom": i.key.startswith("custom:")}
    if st == "late" and i.key in R.BY_KEY:
        v["penalty"] = R.penalty(e, -days, i.period, facts)
    elif i.key in R.BY_KEY:
        v["penalty"] = {"text": (e.get("penalty") or {}).get("text", ""), "amount": None, "basis": ""}
    if full:
        v["log"] = _j(i.log, [])
        v["steps"] = R.steps(e)
    return v


def _get(s, tenant: str, ref: str) -> CompItem:
    ref = str(ref or "").strip()
    i = s.get(CompItem, ref)
    if i is None or i.tenant_id != tenant:
        i = s.scalars(select(CompItem).where(CompItem.tenant_id == tenant, CompItem.number == ref.upper())).first()
    if i is None:
        raise CompError("Kaam nahi mila (number jaise C-0003 dijiye).")
    return i


def calendar(tenant: str, ahead: int = 60, area: str = "", state: str = "", include_linked: bool = True) -> dict:
    on = today(tenant)
    f = facts_of(tenant)
    with session() as s:
        rows = s.scalars(select(CompItem).where(CompItem.tenant_id == tenant)).all()
        views = [_item_view(i, on, f) for i in rows]
    end = on + timedelta(days=ahead)
    out = []
    for v in views:
        if area and v["area"] != area:
            continue
        if state and v["state"] != state:
            continue
        if v["state"] in ("done", "na") and not state:
            dd = _date(v["due"])
            if not dd or (on - dd).days > 35:
                continue
        if v["due"] and _date(v["due"]) > end and v["state"] not in ("late",):
            continue
        out.append(v)
    order = {"late": 0, "due": 1, "no_date": 2, "upcoming": 3, "done": 4, "na": 5}
    out.sort(key=lambda v: (order.get(v["state"], 9), v["due"] or "9999"))
    lrows = linked_rows(tenant, ahead=ahead) if include_linked else []
    late = [v for v in out if v["state"] == "late"]
    soon = [v for v in out if v["state"] == "due"]
    nod = [v for v in out if v["state"] == "no_date"]
    bits = []
    if late:
        bits.append(f"{len(late)} late: " + ", ".join(f"{v['number']} {v['label']}" for v in late[:3]))
    if soon:
        bits.append(f"7 din mein {len(soon)}: " + ", ".join(f"{v['label']} {v['due_nice']}" for v in soon[:4]))
    if nod:
        bits.append(f"{len(nod)} kaam ki tareekh / status aap batayein")
    ll = [r for r in lrows if r["late"]]
    if ll:
        bits.append(f"doosre Sahayak mein {len(ll)} late")
    return {"date": on.isoformat(), "items": out, "linked": lrows, "late": len(late), "due_soon": len(soon), "no_date": len(nod),
            "say": (". ".join(bits) + ".") if bits else "Agle dino mein koi compliance kaam baaki nahi."}


def item(tenant: str, ref: str) -> dict:
    with session() as s:
        return _item_view(_get(s, tenant, ref), today(tenant), facts_of(tenant), full=True)


def item_update(tenant: str, ref: str, due: str = "", owner: str | None = None, note: str = "", actor: str = "owner") -> dict:
    with _LOCK, session() as s:
        i = _get(s, tenant, ref)
        who = _who(tenant, actor)
        if due:
            if not is_owner(actor):
                raise CompError("Tareekh sirf owner badal sakti hain.")
            dd = _date(due)
            if not dd:
                raise CompError("Tareekh samajh nahi aayi (25-10-2026).")
            i.due_override = dd.isoformat()
            _log(i, f"tareekh {dd.isoformat()} (owner ne di)", who)
        if owner is not None:
            i.owner = _clean(owner, 80)
            _log(i, f"zimmedaar: {i.owner or '-'}", who)
        if note:
            i.note = _clean(note, 600)
            _log(i, "note", who)
        v = _item_view(i, today(tenant), facts_of(tenant), full=True)
    return v | {"say": f"{v['number']} badla."}


def add_custom(tenant: str, label: str, due: str = "", area: str = "custom", owner: str = "", note: str = "", actor: str = "owner") -> dict:
    if not is_owner(actor):
        raise CompError("Naya kaam sirf owner jod sakti hain.")
    label = _clean(label, 200)
    if not label:
        raise CompError("Kaam ka naam chahiye.")
    area = area if area in R.AREAS else "custom"
    with _LOCK, session() as s:
        n = _next(s, tenant, "custom")
        i = CompItem(tenant_id=tenant, key=f"custom:{n}", area=area, period="once", label=label,
                     due=(_date(due).isoformat() if _date(due) else ""), owner=_clean(owner, 80), note=_clean(note, 600))
        i.number = f"C-{_next(s, tenant, 'item'):04d}"
        _log(i, "aapne joda", "owner")
        s.add(i)
        s.flush()
        v = _item_view(i, today(tenant))
    return v | {"say": f"{v['number']} {label} joda" + (f", due {v['due_nice']}." if v["due"] else ". Tareekh bata dijiye.")}


# ---------------------------------------------------------------- evidence, done, n/a
def add_evidence(tenant: str, ref: str, file_id: str, note: str = "", actor: str = "owner") -> dict:
    if not file_id:
        raise CompError("file_id chahiye (challan / ack / certificate ki file).")
    meta = _file_meta(tenant, file_id)
    if not meta:
        raise CompError("Yeh file nahi mili; chat mein file bhejiye phir dobara.")
    with _LOCK, session() as s:
        i = _get(s, tenant, ref)
        ev = _j(i.evidence, [])
        if any(x.get("file_id") == file_id for x in ev):
            return _item_view(i, today(tenant)) | {"say": "Yeh file pehle se judi hai."}
        who = _who(tenant, actor)
        ev.append({"file_id": file_id, "name": str(meta.get("name") or "")[:120], "note": _clean(note, 200),
                   "at": today(tenant).isoformat(), "by": who})
        i.evidence = _d(ev[-20:])
        _log(i, f"saboot joda: {meta.get('name', '')}", who)
        v = _item_view(i, today(tenant))
    return v | {"say": f"{v['number']} par saboot juda."}


def mark_done(tenant: str, ref: str, file_id: str = "", note: str = "", her_word: bool = False, ack: str = "", on: str = "",
              actor: str = "owner") -> dict:
    """'Ho gaya' needs evidence (a file on the item or given now) or, only for the owner, her own word."""
    if file_id:
        add_evidence(tenant, ref, file_id, note, actor)
    with _LOCK, session() as s:
        i = _get(s, tenant, ref)
        if i.status == "done":
            return _item_view(i, today(tenant)) | {"say": f"{i.number} pehle se ho gaya hai."}
        if i.status == "na":
            raise CompError("Yeh kaam 'lagu nahi' mein hai; pehle reopen kijiye.")
        has_ev = bool(_j(i.evidence, []))
        if not has_ev:
            if not her_word:
                raise CompError("Saboot chahiye: challan / ack / certificate ki file bhejiye (file_id), ya aap khud 'ho gaya' kahiye.")
            if not is_owner(actor):
                raise CompError("Bina saboot 'ho gaya' sirf owner keh sakti hain; staff file bheje.")
        who = _who(tenant, actor)
        i.status = "done"
        i.done_how = "evidence" if has_ev else "word"
        i.done_on = (_date(on) or today(tenant)).isoformat()
        i.done_by = who
        if ack:
            i.ack = _clean(ack, 80)
        if note and not file_id:
            i.note = _clean(note, 600)
        _log(i, "ho gaya (" + ("saboot ke saath" if has_ev else "owner ke kehne par") + ")", who)
        v = _item_view(i, today(tenant))
    return v | {"say": f"{v['number']} {v['label']}: ho gaya" + (" (saboot ke saath)." if v["done_how"] == "evidence" else
                                                                     " (aapke kehne par; saboot baad mein jod sakti hain).")}


def mark_na(tenant: str, ref: str, reason: str, actor: str = "owner") -> dict:
    if not is_owner(actor):
        raise CompError("'Lagu nahi' sirf owner keh sakti hain.")
    if not reason or len(reason.strip()) < 4:
        raise CompError("Kyun lagu nahi, ek line likhiye (certificate mein jaati hai).")
    with _LOCK, session() as s:
        i = _get(s, tenant, ref)
        i.status = "na"
        i.note = _clean(reason, 400)
        _log(i, f"lagu nahi: {i.note}", "owner")
        v = _item_view(i, today(tenant))
    return v | {"say": f"{v['number']} lagu nahi maana: {v['note']}."}


def reopen(tenant: str, ref: str, actor: str = "owner") -> dict:
    if not is_owner(actor):
        raise CompError("Reopen sirf owner.")
    with _LOCK, session() as s:
        i = _get(s, tenant, ref)
        i.status, i.done_how, i.done_on, i.done_by = "open", "", "", ""
        _log(i, "dobara khola", "owner")
        v = _item_view(i, today(tenant))
    return v | {"say": f"{v['number']} dobara khula."}


# ---------------------------------------------------------------- 'portal par file karo' card (her tap)
def _filing_show(i: CompItem) -> dict:
    e = R.BY_KEY.get(i.key) or {}
    return {"item": i.id, "number": i.number, "label": i.label, "period": i.period, "due": (i.due_override or i.due),
            "portal": e.get("portal", ""), "steps": R.steps(e) if e else [], "owner": i.owner}


def _f_view(x: CompFiling) -> dict:
    return {"id": x.id, "number": x.number, "item_id": x.item_id, "show": _j(x.show, {}), "hash": x.hash, "status": x.status,
            "reason": x.reason}


def filing_request(tenant: str, ref: str, conv: str = "") -> dict:
    """Card only: 'is kaam ko portal par file karein?' Her tap decides; she still logs in, pays and submits herself."""
    with _LOCK, session() as s:
        i = _get(s, tenant, ref)
        if i.status in ("done", "na"):
            raise CompError(f"{i.number} pehle hi '{i.status}' hai.")
        if not (i.due_override or i.due):
            raise CompError("Pehle tareekh bataiye (item_update due=...).")
        for old in s.scalars(select(CompFiling).where(CompFiling.item_id == i.id, CompFiling.status == "waiting")).all():
            old.status = "stale"
        show = _filing_show(i)
        x = CompFiling(tenant_id=tenant, item_id=i.id, number=f"F-{_next(s, tenant, 'filing'):03d}", show=_d(show), hash=_hash(show))
        s.add(x)
        s.flush()
        v = _f_view(x)
    v["say"] = f"{show['label']}: portal par file karne ka card bheja hai - aapki haan chahiye. Login / OTP / payment / submit aap khud karengi."
    _publish(tenant, "compliance_filing", {"filing": v, "say": v["say"]}, conv)
    return v


def decide_filing(tenant: str, filing_id: str, yes: bool, card_hash: str = "", reason: str = "", by_tap: bool = False) -> dict:
    if not by_tap:
        raise CompError("File karne ki haan sirf aapke app ke tap se.")
    with _LOCK, session() as s:
        x = s.get(CompFiling, filing_id)
        if x is None or x.tenant_id != tenant:
            raise CompError("Card nahi mila.")
        if x.status != "waiting":
            return {"filing": _f_view(x), "say": f"{x.number} pehle hi '{x.status}' hai."}
        i = s.get(CompItem, x.item_id)
        now_hash = _hash(_filing_show(i)) if i else ""
        stale = not card_hash or card_hash != x.hash or now_hash != x.hash or i.status in ("done", "na")
        if stale and card_hash == x.hash:
            x.status = "stale"              # the item moved under the card (a wrong hash alone does not spoil it)
    if stale:
        raise CompError("Kaam badal gaya hai; naya card banaiye.")
    with _LOCK, session() as s:
        x = s.get(CompFiling, filing_id)
        i = s.get(CompItem, x.item_id)
        x.status = "approved" if yes else "rejected"
        x.reason = _clean(reason, 300)
        x.decided_at = now()
        if yes:
            i.status = "file_ok"
            _log(i, "file karne ki haan (owner ka tap)", "owner")
        v = _f_view(x)
    if yes:
        v["say"] = (f"{v['show']['label']}: haan mili. Ab {v['show']['portal'] or 'portal'} par aap login karke file kijiye (OTP / payment / "
                    "submit aap khud). Ack / challan ki file bhejte hi 'ho gaya' kar dunga.")
    else:
        v["say"] = "Theek hai, abhi file nahi."
    return v


# ---------------------------------------------------------------- monthly compliance certificate (her sign-off)
def _month(m: str) -> tuple[int, int]:
    mm = re.fullmatch(r"(\d{4})-(\d{1,2})", str(m or "").strip())
    if not mm or not 1 <= int(mm.group(2)) <= 12:
        raise CompError("month = YYYY-MM")
    return int(mm.group(1)), int(mm.group(2))


def certificate(tenant: str, month: str = "") -> dict:
    """Draft content for a month: every own item due that month + older still-open items + neighbours' rows (read)."""
    on = today(tenant)
    if not month:
        y, m = R.add_months(on.year, on.month, -1)
        month = f"{y:04d}-{m:02d}"
    y, m = _month(month)
    first, last = date(y, m, 1), date(y, m, R.last_day(y, m))
    with session() as s:
        st = _st(s, tenant)
        firm, sign, desig = st.firm, st.signatory, st.designation
        rows = s.scalars(select(CompItem).where(CompItem.tenant_id == tenant)).all()
        items = []
        for i in rows:
            dd = _eff_due(i)
            if dd is None:
                continue
            in_month = first <= dd <= last
            carry = dd < first and i.status not in ("done", "na")
            if not (in_month or carry):
                continue
            items.append({"number": i.number, "label": i.label, "due": dd.isoformat(), "status": i.status, "how": i.done_how,
                          "done_on": i.done_on, "ack": i.ack, "evidence": [x.get("name", "") for x in _j(i.evidence, [])],
                          "carry": carry, "note": i.note})
    items.sort(key=lambda x: (x["due"], x["number"]))
    lrows = []
    for r in linked_rows(tenant, ahead=max(0, (last - on).days)):
        dd = _date(r["due"])
        if dd and (first <= dd <= last or (dd < first and r["late"])):
            lrows.append({"source": r["source_label"], "label": r["label"], "due": r["due"], "status": r["status"] or ("late" if r["late"] else "")})
    counts = {"total": len(items) + len(lrows), "done_evidence": sum(1 for x in items if x["status"] == "done" and x["how"] == "evidence"),
              "done_word": sum(1 for x in items if x["status"] == "done" and x["how"] == "word"),
              "na": sum(1 for x in items if x["status"] == "na"), "open": sum(1 for x in items if x["status"] in ("open", "file_ok")),
              "linked": len(lrows)}
    exc = [f"{x['number']} {x['label']} (due {x['due']}) baaki" for x in items if x["status"] in ("open", "file_ok")]
    exc += [f"{x['number']} {x['label']}: lagu nahi - {x['note']}" for x in items if x["status"] == "na"]
    exc += [f"{x['source']}: {x['label']} (due {x['due']}) late" for x in lrows if x["status"] == "late"]
    label = f"{R.MONTHS[m]} {y}"
    text = R.certificate_text(sign, firm, label, counts, exc)
    return {"month": f"{y:04d}-{m:02d}", "month_label": label, "firm": firm, "signatory": sign, "designation": desig, "items": items,
            "linked": lrows, "counts": counts, "exceptions": exc, "text": text, "check": R.CHECK_LINE}


def _c_view(c: CompCert) -> dict:
    return {"id": c.id, "number": c.number, "month": c.month, "content": _j(c.content, {}), "hash": c.hash, "status": c.status,
            "reason": c.reason, "file_id": c.file_id, "signed_on": c.signed_on}


def cert_request(tenant: str, month: str = "", conv: str = "") -> dict:
    content = certificate(tenant, month)
    if not content["signatory"]:
        raise CompError("Certificate par kaun sign karega? settings mein signatory (aapka naam) dijiye.")
    if not content["counts"]["total"]:
        raise CompError("Is mahine ka koi compliance kaam calendar mein nahi hai.")
    with _LOCK, session() as s:
        for old in s.scalars(select(CompCert).where(CompCert.tenant_id == tenant, CompCert.month == content["month"],
                                                     CompCert.status == "waiting")).all():
            old.status = "stale"
        c = CompCert(tenant_id=tenant, month=content["month"], number=f"CC-{_next(s, tenant, 'cert'):03d}", content=_d(content),
                     hash=_hash(content))
        s.add(c)
        s.flush()
        v = _c_view(c)
    n = content["counts"]
    v["say"] = (f"{content['month_label']} ka compliance certificate draft: {n['total']} kaam, {n['open']} baaki"
                + (f", {len(content['exceptions'])} apvaad likhe" if content["exceptions"] else "") + ". Sign-off ke liye aapka tap chahiye.")
    _publish(tenant, "compliance_cert", {"cert": v, "say": v["say"]}, conv)
    return v


def _pdf(title: str, blocks: list[tuple], watermark: str = "") -> bytes:
    from xml.sax.saxutils import escape

    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    from .. import indicpdf
    pdf = indicpdf.blocks_pdf(title, blocks, watermark)       # Devanagari names print through Chromium (shaped)
    if pdf:
        return pdf
    ss = getSampleStyleSheet()
    small = ss["BodyText"].clone("small", fontSize=8, leading=10)

    def txt(t: Any) -> str:
        return indicpdf.rl_markup(escape(str(t if t is not None else "").replace("₹", "Rs. ")))
    story: list = []
    for b in blocks:
        if b[0] == "h1":
            story.append(Paragraph(txt(b[1]), ss["Title"]))
        elif b[0] == "h":
            story.append(Paragraph(txt(b[1]), ss["Heading2"]))
        elif b[0] == "p":
            for line in str(b[1]).split("\n"):
                story.append(Paragraph(txt(line), ss["BodyText"]))
        elif b[0] == "table" and b[1]:
            t = Table([[Paragraph(txt(c), small) for c in r] for r in b[1]], repeatRows=1)
            t.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.3, colors.grey), ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#EEEEEE")),
                                   ("VALIGN", (0, 0), (-1, -1), "TOP")]))
            story.append(t)
        story.append(Spacer(1, 6))
    bio = io.BytesIO()
    SimpleDocTemplate(bio, pagesize=A4, title=title, leftMargin=36, rightMargin=36, topMargin=36, bottomMargin=36).build(story)
    return bio.getvalue()


def _cert_pdf(content: dict, signed_on: str, number: str) -> bytes:
    rows = [["No", "Kaam", "Due", "Status", "Saboot / ack"]]
    for x in content["items"]:
        st = {"done": "Ho gaya" + (" (saboot)" if x["how"] == "evidence" else " (owner ke kehne par)"), "na": "Lagu nahi",
              "open": "Baaki", "file_ok": "File ki haan, baaki"}.get(x["status"], x["status"])
        rows.append([x["number"], x["label"], x["due"], st, ", ".join(filter(None, [x["ack"]] + x["evidence"]))])
    blocks = [("h1", f"Compliance certificate - {content['month_label']}"),
              ("p", f"{content['firm'] or ''}  ·  {number}"), ("p", content["text"]), ("h", "Kaam ki list"), ("table", rows)]
    if content["linked"]:
        lr = [["Sahayak", "Kaam", "Due", "Status"]] + [[x["source"], x["label"], x["due"], x["status"] or "-"] for x in content["linked"]]
        blocks += [("h", "Doosre Sahayak ke kaam (wahan se padhe)"), ("table", lr)]
    blocks += [("p", f"Sign-off: {content['signatory']}" + (f", {content['designation']}" if content["designation"] else "")
                + f" - Hoola app mein tap se {signed_on} ko. (Yeh digital signature nahi hai.)"), ("p", R.CHECK_LINE)]
    return _pdf(f"Compliance certificate {content['month_label']}", blocks)


def decide_cert(tenant: str, cert_id: str, yes: bool, card_hash: str = "", reason: str = "", by_tap: bool = False, conv: str = "") -> dict:
    if not by_tap:
        raise CompError("Certificate ka sign-off sirf aapke app ke tap se.")
    with session() as s:
        c = s.get(CompCert, cert_id)
        if c is None or c.tenant_id != tenant:
            raise CompError("Card nahi mila.")
        if c.status != "waiting":
            return {"cert": _c_view(c), "say": f"{c.number} pehle hi '{c.status}' hai."}
        month, stored = c.month, c.hash
    content_now = certificate(tenant, month)
    if not card_hash or card_hash != stored or _hash(content_now) != stored:
        with _LOCK, session() as s:
            s.get(CompCert, cert_id).status = "stale"
        raise CompError("Card ke baad kaam badle hain; naya certificate draft banaiye.")
    signed_on = today(tenant).isoformat()
    with _LOCK, session() as s:
        c = s.get(CompCert, cert_id)
        c.decided_at = now()
        c.reason = _clean(reason, 300)
        if not yes:
            c.status = "rejected"
            return {"cert": _c_view(c), "say": "Certificate sign nahi kiya."}
        c.status, c.signed_on = "signed", signed_on
        number = c.number
    pdf = _cert_pdf(content_now, signed_on, number)
    f = {}
    try:
        from .. import chatfiles
        f = chatfiles.store(tenant, f"Compliance_certificate_{month}.pdf", pdf, "application/pdf")
        _publish(tenant, "pc_file", {**f, "source": "compliance", "title": f"Compliance certificate {month}"}, conv)
    except Exception:
        log.exception("compliance cert file")
    with _LOCK, session() as s:
        c = s.get(CompCert, cert_id)
        c.file_id = (f or {}).get("file_id", "")
        v = _c_view(c)
    return {"cert": v, "file": f, "say": f"{content_now['month_label']} ka certificate aapne sign kiya; PDF bani. {R.CHECK_LINE}"}


def certs(tenant: str) -> list[dict]:
    with session() as s:
        rows = s.scalars(select(CompCert).where(CompCert.tenant_id == tenant).order_by(CompCert.created_at.desc())).all()
        return [_c_view(c) | {"content": {}} for c in rows[:24]]


def filings(tenant: str, status: str = "waiting") -> list[dict]:
    with session() as s:
        q = select(CompFiling).where(CompFiling.tenant_id == tenant)
        if status:
            q = q.where(CompFiling.status == status)
        return [_f_view(x) for x in s.scalars(q.order_by(CompFiling.created_at.desc())).all()[:50]]


# ---------------------------------------------------------------- day note, nudge, register, forget
def day(tenant: str) -> dict:
    cal = calendar(tenant, ahead=14)
    with session() as s:
        wf = len(s.scalars(select(CompFiling.id).where(CompFiling.tenant_id == tenant, CompFiling.status == "waiting")).all())
        wc = len(s.scalars(select(CompCert.id).where(CompCert.tenant_id == tenant, CompCert.status == "waiting")).all())
    late = [v for v in cal["items"] if v["state"] == "late"]
    soon = [v for v in cal["items"] if v["state"] == "due"]
    pen = sum((v.get("penalty") or {}).get("amount") or 0 for v in late)
    bits = []
    if late:
        bits.append(f"Late {len(late)}: " + ", ".join(f"{v['label']}" for v in late[:3])
                    + (f" (penalty andaaza ~{_rs(pen)}, source ke saath)" if pen else ""))
    if soon:
        bits.append("Is hafte: " + ", ".join(f"{v['label']} {v['due_nice']}" for v in soon[:4]))
    if wf or wc:
        bits.append(f"Aapki haan baaki: {wf} file card, {wc} certificate")
    ll = [r for r in cal["linked"] if r["late"]]
    if ll:
        bits.append("Doosre Sahayak mein late: " + ", ".join(r["label"] for r in ll[:3]))
    q = R.questions(facts_of(tenant))
    return {"late": late, "soon": soon, "waiting_filings": wf, "waiting_certs": wc, "linked_late": ll, "questions": q,
            "penalty_andaaza": pen or None, "say": (". ".join(bits) + ".") if bits else "Compliance mein aaj sab theek."}


def nudge(tenant: str, ref: str) -> dict:
    """WhatsApp draft to the person responsible (she sends it)."""
    with session() as s:
        i = _get(s, tenant, ref)
        team = _j(_st(s, tenant).team, [])
        v = _item_view(i, today(tenant))
    who = next((m for m in team if m.get("name", "").lower() == (v["owner"] or "").lower()), {})
    text = (f"Namaste {v['owner'] or ''}, {v['label']} ki tareekh {v['due_nice'] or '[tareekh?]'} hai. Ho jaaye to challan / ack ki file bhej "
            "dijiye. Dhanyavaad.")
    return {"text": text, "whatsapp": _wa(who.get("phone", ""), text), "to": v["owner"],
            "say": "Message ka draft taiyaar; aap bhejiye." + ("" if who.get("phone") else " (Is vyakti ka phone team mein nahi hai.)")}


def register_file(tenant: str, conv: str = "") -> dict:
    import openpyxl
    on = today(tenant)
    f = facts_of(tenant)
    with session() as s:
        rows = [_item_view(i, on, f) for i in s.scalars(select(CompItem).where(CompItem.tenant_id == tenant)).all()]
    rows.sort(key=lambda v: v["due"] or "9999")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Calendar"
    ws.append(["No", "Area", "Kaam", "Period", "Due", "Zimmedaar", "Status", "Kaise", "Ack", "Saboot", "Penalty (starter)", "Source", "Link"])
    for v in rows:
        ws.append([v["number"], v["area_label"], v["label"], v["period"], v["due"], v["owner"], v["state"], v["done_how"], v["ack"],
                   ", ".join(x.get("name", "") for x in v["evidence"]), (v.get("penalty") or {}).get("text", ""),
                   v["source"].get("title", ""), v["source"].get("url", "")])
    ws = wb.create_sheet("Lagu kya hai")
    ws.append(["Kaam", "Area", "Lagu?", "Kyun", "Kahan (Sahayak)", "Source", "Pakka karwaiye?"])
    for a in applicability(tenant)["items"]:
        ws.append([a["label"], a["area_label"], "haan" if a["applies"] else "nahi", a["why"], a.get("covered_by", ""),
                   a["source"]["title"], "haan" if a["verify"] else ""])
    ws = wb.create_sheet("Note")
    ws.append([R.CHECK_LINE])
    bio = io.BytesIO()
    wb.save(bio)
    from .. import chatfiles
    fl = chatfiles.store(tenant, f"Compliance_register_{on.isoformat()}.xlsx", bio.getvalue(),
                         "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    _publish(tenant, "pc_file", {**fl, "source": "compliance", "title": "Compliance register"}, conv)
    return {"file": fl, "say": "Compliance register Excel bana (calendar + kya lagu hai + sources)."}


def forget_person(tenant: str, name: str) -> dict:
    """DPDP: a person who left - name removed from items and team (the items stay)."""
    name = (name or "").strip().lower()
    if not name:
        raise CompError("Naam chahiye.")
    n = 0
    with _LOCK, session() as s:
        st = _st(s, tenant)
        st.team = _d([m for m in _j(st.team, []) if (m.get("name") or "").lower() != name])
        st.owners = _d({k: v for k, v in _j(st.owners, {}).items() if (v or "").lower() != name})
        for i in s.scalars(select(CompItem).where(CompItem.tenant_id == tenant)).all():
            if (i.owner or "").lower() == name:
                i.owner = ""
                n += 1
            if (i.done_by or "").lower() == name:
                i.done_by = "purana staff"
    return {"items": n, "say": f"{n} kaam se naam hataya; team se bhi."}


# ---------------------------------------------------------------- worker
def _tenants() -> list[str]:
    with session() as s:
        return sorted(set(s.scalars(select(CompSettings.tenant_id)).all()))


def tick(when: datetime | None = None) -> dict:
    """Worker: rebuild the calendar once a day, reminders at her windows before the due date, overdue once (and weekly), a
    certificate reminder for last month on the 1st-5th, morning note at her time."""
    n = when or now()
    if n.tzinfo is None:
        n = n.replace(tzinfo=timezone.utc)
    out = {"alerts": 0, "morning": 0, "built": 0}
    for tenant in _tenants():
        try:
            loc = n.astimezone(_tz(tenant))
            on = loc.date()
            with session() as s:
                st = _st(s, tenant)
                need_build = st.last_build != on.isoformat() and bool(_j(st.facts, {}))
            if need_build:
                build(tenant)
                out["built"] += 1
            alerts = []
            f = facts_of(tenant)
            with _LOCK, session() as s:
                st = _st(s, tenant)
                wins = _windows(st)
                for i in s.scalars(select(CompItem).where(CompItem.tenant_id == tenant, CompItem.status.in_(("open", "file_ok")))).all():
                    dd = _eff_due(i)
                    if not dd:
                        continue
                    left = (dd - on).days
                    if left < 0:
                        week = (-left - 1) // 7
                        if _once(s, tenant, f"late:{i.id}:{dd.isoformat()}:{week}"):
                            v = _item_view(i, on, f)
                            p = v.get("penalty") or {}
                            alerts.append({"item": v, "say": f"{i.number} {i.label} {-left} din late." + (
                                f" Penalty andaaza ~{_rs(p['amount'])} ({p['basis']})." if p.get("amount") else f" {p.get('text', '')}")})
                        continue
                    w = next((x for x in sorted(wins) if left <= x), None)
                    if w is not None and _once(s, tenant, f"due:{i.id}:{dd.isoformat()}:{w}"):
                        alerts.append({"item": _item_view(i, on, f), "say": f"{i.number} {i.label}: " + (
                            "aaj aakhri din." if left == 0 else f"{left} din baaki ({R.nice(dd)}).")
                            + (f" Zimmedaar: {i.owner}." if i.owner else "")})
                if on.day <= 5:
                    y, m = R.add_months(on.year, on.month, -1)
                    mon = f"{y:04d}-{m:02d}"
                    signed = s.scalars(select(CompCert.id).where(CompCert.tenant_id == tenant, CompCert.month == mon,
                                                                 CompCert.status.in_(("signed", "waiting")))).first()
                    if not signed and s.scalars(select(CompItem.id).where(CompItem.tenant_id == tenant)).first() and \
                            _once(s, tenant, f"certmon:{mon}"):
                        alerts.append({"item": None, "say": f"{R.MONTHS[m]} {y} ka compliance certificate draft banaun? "
                                                            "('certificate' boliye; sign-off aapke tap se)."})
                morning = False
                if st.last_morning != on.isoformat():
                    try:
                        hh, mm = (int(x) for x in (st.morning_time or "09:00").split(":"))
                    except ValueError:
                        hh, mm = 9, 0
                    if (loc.hour, loc.minute) >= (hh, mm):
                        st.last_morning = on.isoformat()
                        morning = True
            for a in alerts:
                _publish(tenant, "compliance_alert", a)
                out["alerts"] += 1
            if morning:
                dd_ = day(tenant)
                if dd_["say"] != "Compliance mein aaj sab theek.":
                    _publish(tenant, "compliance_day", {"say": dd_["say"], "late": len(dd_["late"]), "soon": len(dd_["soon"]),
                                                        "waiting": dd_["waiting_filings"] + dd_["waiting_certs"]})
                    out["morning"] += 1
        except Exception:
            log.exception("compliance tick failed for %s", tenant)
    return out
