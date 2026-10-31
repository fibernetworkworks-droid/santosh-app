"""v0.103 SEO Sahayak: Hoola as an SEO specialist, end to end.

  1. Gets the job: "meri website ka SEO dekho", "Google par kyu nahi aate", "is hafte ki report", a review to answer,
     "naye page ke liye likho".
  2. Does the work, step by step: crawls HER site (robots.txt honoured, same host only, public addresses only), checks
     every page (title, description, H1, words, alt, viewport, https, noindex, canonical, lang, schema, city), the site
     (robots blocking, sitemap, duplicate titles), Core Web Vitals from Google's PageSpeed Insights API, and Search Console
     (official API, read-only sign-in): clicks, impressions, CTR, position vs the previous period, striking-distance
     queries (position 4-20), low-CTR queries, cannibalisation, intent mix.
  3. Checked against the rules and earlier work: Google Search Essentials + spam policies, CWV thresholds, earlier crawls
     (fixed / new issues), last period's numbers.
  4. Ready as a card, a file or a draft: a fix list ranked by impact with the exact fix, LocalBusiness JSON-LD, title /
     description suggestions, a content brief, Google Business Profile post and review-reply drafts, an Excel report.
  5. Waits for her Haan: nothing is published by Hoola. Drafts become final only on her Haan, and she (or her web person)
     puts them on the site / profile.
  6. Delivers and learns: tasks she marks done are re-checked on the next crawl; numbers are kept week by week.
"""
from __future__ import annotations

import io
import json
import logging
import re
import secrets
import threading
import time
from datetime import date, datetime, timedelta
from html.parser import HTMLParser
from typing import Any
from urllib.parse import quote, urldefrag, urlencode, urljoin, urlparse
from urllib.robotparser import RobotFileParser

import httpx
from sqlalchemy import Boolean, DateTime, Float, Integer, String, Text, select
from sqlalchemy.orm import Mapped, mapped_column

from ..db import Base, new_id, now, session
from . import rules as R

log = logging.getLogger("santosh.seo")
_LOCK = threading.Lock()
UA = "HoolaSEO/1.0 (site owner's own audit)"
GSC_SCOPE = "https://www.googleapis.com/auth/webmasters.readonly"
GA4_SCOPE = "https://www.googleapis.com/auth/analytics.readonly"   # v0.314 optional: organic sessions + leads for the SEO Employee
GSC_API = "https://www.googleapis.com/webmasters/v3"
PSI_API = "https://www.googleapis.com/pagespeedonline/v5/runPagespeed"
_http: httpx.Client | None = None          # tests inject a mock transport
_PAUSE = 0.3


class SeoError(Exception):
    pass


def set_http(c: httpx.Client | None) -> None:
    global _http
    _http = c


class SeoSite(Base):
    __tablename__ = "seo_sites"
    id: Mapped[str] = mapped_column(String(40), primary_key=True, default=lambda: new_id("ss"))
    tenant_id: Mapped[str] = mapped_column(String(64), index=True)
    url: Mapped[str] = mapped_column(String(300), default="")
    name: Mapped[str] = mapped_column(String(160), default="")
    business: Mapped[str] = mapped_column(Text, default="{}")      # city, state, street, pin, phone, schema_type, hours, price_range, same_as
    services: Mapped[str] = mapped_column(Text, default="[]")
    gsc_property: Mapped[str] = mapped_column(String(300), default="")
    score: Mapped[int] = mapped_column(Integer, default=0)
    last_crawl: Mapped[str] = mapped_column(String(20), default="")
    cwv: Mapped[str] = mapped_column(Text, default="{}")
    created_at: Mapped[datetime] = mapped_column(DateTime(timezone=True), default=now)


class SeoPage(Base):
    __tablename__ = "seo_pages"
    id: Mapped[str] = mapped_column(String(40), primary_key=True, default=lambda: new_id("sp"))
    tenant_id: Mapped[str] = mapped_column(String(64), index=True)
    site_id: Mapped[str] = mapped_column(String(40), index=True)
    url: Mapped[str] = mapped_column(String(600), default="")
    data: Mapped[str] = mapped_column(Text, default="{}")
    issues: Mapped[str] = mapped_column(Text, default="[]")
    score: Mapped[int] = mapped_column(Integer, default=0)


class SeoTask(Base):
    __tablename__ = "seo_tasks"
    id: Mapped[str] = mapped_column(String(40), primary_key=True, default=lambda: new_id("st"))
    tenant_id: Mapped[str] = mapped_column(String(64), index=True)
    site_id: Mapped[str] = mapped_column(String(40), index=True)
    key: Mapped[str] = mapped_column(String(700), default="", index=True)    # code|url
    code: Mapped[str] = mapped_column(String(30), default="")
    url: Mapped[str] = mapped_column(String(600), default="")
    level: Mapped[str] = mapped_column(String(6), default="warn")
    text: Mapped[str] = mapped_column(Text, default="")
    fix: Mapped[str] = mapped_column(Text, default="")
    status: Mapped[str] = mapped_column(String(10), default="open")         # open | done | ignored | fixed (seen fixed on re-crawl)
    seen: Mapped[str] = mapped_column(String(20), default="")


class SeoStat(Base):
    """One Search Console period (kept week by week)."""
    __tablename__ = "seo_stats"
    id: Mapped[str] = mapped_column(String(40), primary_key=True, default=lambda: new_id("sq"))
    tenant_id: Mapped[str] = mapped_column(String(64), index=True)
    site_id: Mapped[str] = mapped_column(String(40), index=True)
    start: Mapped[str] = mapped_column(String(10), default="")
    end: Mapped[str] = mapped_column(String(10), default="")
    clicks: Mapped[int] = mapped_column(Integer, default=0)
    impressions: Mapped[int] = mapped_column(Integer, default=0)
    ctr: Mapped[float] = mapped_column(Float, default=0.0)
    position: Mapped[float] = mapped_column(Float, default=0.0)
    rows: Mapped[str] = mapped_column(Text, default="[]")


class SeoDraft(Base):
    __tablename__ = "seo_drafts"
    id: Mapped[str] = mapped_column(String(40), primary_key=True, default=lambda: new_id("sd"))
    tenant_id: Mapped[str] = mapped_column(String(64), index=True)
    site_id: Mapped[str] = mapped_column(String(40), index=True)
    kind: Mapped[str] = mapped_column(String(16), default="")               # gbp_post | review_reply | brief | meta | schema
    title: Mapped[str] = mapped_column(String(200), default="")
    text: Mapped[str] = mapped_column(Text, default="")
    status: Mapped[str] = mapped_column(String(10), default="ready")        # ready | approved
    created_at: Mapped[datetime] = mapped_column(DateTime(timezone=True), default=now)


class SeoGoogle(Base):
    """Her Google account for Search Console (read-only scope). Tokens encrypted."""
    __tablename__ = "seo_google"
    tenant_id: Mapped[str] = mapped_column(String(64), primary_key=True)
    address: Mapped[str] = mapped_column(String(320), default="")
    access_token_enc: Mapped[str | None] = mapped_column(Text, nullable=True)
    refresh_token_enc: Mapped[str | None] = mapped_column(Text, nullable=True)
    expiry: Mapped[datetime | None] = mapped_column(DateTime(timezone=True), nullable=True)
    status: Mapped[str] = mapped_column(String(12), default="ok")
    last_error: Mapped[str] = mapped_column(Text, default="")
    scopes: Mapped[str | None] = mapped_column(Text, nullable=True)       # v0.314 what she granted (Analytics is optional)


class SeoSettings(Base):
    __tablename__ = "seo_settings"
    tenant_id: Mapped[str] = mapped_column(String(64), primary_key=True)
    enabled: Mapped[bool] = mapped_column(Boolean, default=True)
    weekday: Mapped[int] = mapped_column(Integer, default=0)                # Monday report
    morning_time: Mapped[str] = mapped_column(String(5), default="09:30")
    last_morning: Mapped[str] = mapped_column(String(10), default="")


# ---------------------------------------------------------------- helpers
def _j(s: str, default):
    try:
        v = json.loads(s or "")
        return v if isinstance(v, type(default)) else default
    except (TypeError, ValueError):
        return default


def _tz(tenant: str):
    from ..life import tz_of
    return tz_of(tenant)


def today(tenant: str) -> date:
    return datetime.now(_tz(tenant)).date()


def _publish(tenant: str, kind: str, payload: dict, conv: str = "") -> None:
    from ..events import bus
    try:
        if not conv:
            from ..net.service import home_conv
            conv = home_conv(tenant)
        bus.publish(conv, kind, payload)
    except Exception:
        log.exception("seo publish failed")


def _client() -> tuple[httpx.Client, bool]:
    return (_http, False) if _http is not None else (httpx.Client(timeout=20, headers={"User-Agent": UA}), True)


def _safe(url: str) -> str:
    from ..web import WebError, _check_url
    try:
        return _check_url(url)
    except WebError as e:
        raise SeoError(str(e))


def _fetch(url: str, max_bytes: int = 4_000_000) -> tuple[int, str, str, int]:
    """GET a public URL, following up to 4 redirects (each one re-checked). -> (status, final_url, text, bytes)."""
    c, own = _client()
    try:
        for _ in range(5):
            url = _safe(url)
            try:
                r = c.get(url, headers={"User-Agent": UA}, follow_redirects=False)
            except httpx.HTTPError as e:
                raise SeoError(f"{url} khula nahi ({type(e).__name__}).")
            if r.status_code in (301, 302, 303, 307, 308) and r.headers.get("location"):
                url = urljoin(url, r.headers["location"])
                continue
            body = r.content[:max_bytes]
            return r.status_code, url, body.decode(r.encoding or "utf-8", "replace"), len(r.content)
        raise SeoError("Bahut zyada redirect.")
    finally:
        if own:
            c.close()


class _Page(HTMLParser):
    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.title, self.description, self.canonical, self.lang = "", "", "", ""
        self.h1: list[str] = []
        self.h2: list[str] = []
        self.links: list[str] = []
        self.robots, self.viewport = "", False
        self.img, self.img_no_alt = 0, 0
        self.jsonld: list[str] = []
        self.text: list[str] = []
        self.mixed = 0                          # v0.314 http:// pictures / scripts / styles on an https page
        self._in, self._buf, self._skip = "", [], 0

    def handle_starttag(self, tag, attrs):
        a = {k: (v or "") for k, v in attrs}
        if tag == "html":
            self.lang = a.get("lang", "")
        elif tag == "meta":
            n = (a.get("name") or a.get("property") or "").lower()
            if n == "description":
                self.description = a.get("content", "").strip()
            elif n == "robots":
                self.robots = a.get("content", "").lower()
            elif n == "viewport":
                self.viewport = True
        elif tag == "link" and "canonical" in a.get("rel", "").lower():
            self.canonical = a.get("href", "")
        elif tag == "a" and a.get("href"):
            self.links.append(a["href"])
        elif tag == "img":
            self.mixed += a.get("src", "").startswith("http://")
            self.img += 1
            if not a.get("alt", "").strip():
                self.img_no_alt += 1
        elif tag in ("script", "iframe", "link") and (a.get("src") or (a.get("href") if tag == "link" else "") or "").startswith("http://") \
                and (tag != "link" or re.search(r"stylesheet|icon|preload", a.get("rel", ""), re.I)):
            self.mixed += 1
            if tag == "script":
                self._skip += 1
        elif tag == "script" and a.get("type", "").lower() == "application/ld+json":
            self._in, self._buf = "jsonld", []
        elif tag in ("script", "style", "noscript", "svg"):
            self._skip += 1
        elif tag in ("title", "h1", "h2"):
            self._in, self._buf = tag, []

    def handle_endtag(self, tag):
        if tag in ("script", "style", "noscript", "svg") and self._in != "jsonld" and self._skip:
            self._skip -= 1
            return
        if self._in and tag in ("title", "h1", "h2", "script"):
            t = re.sub(r"\s+", " ", "".join(self._buf)).strip()
            if self._in == "title":
                self.title = self.title or t
            elif self._in == "h1":
                self.h1.append(t)
            elif self._in == "h2":
                self.h2.append(t)
            elif self._in == "jsonld":
                self.jsonld.append("".join(self._buf))
            self._in = ""

    def handle_data(self, data):
        if self._in:
            self._buf.append(data)
        if not self._skip and self._in != "jsonld":
            self.text.append(data)


def _parse(url: str, status: int, html: str, size: int, home: str, city: str = "") -> tuple[dict, list[str]]:
    p = _Page()
    try:
        p.feed(html)
    except Exception:
        pass
    words = len(re.findall(r"\w+", " ".join(p.text)))
    canon = urljoin(url, p.canonical) if p.canonical else ""
    local = False
    for j in p.jsonld:
        if re.search(r'"@type"\s*:\s*"[^"]*(LocalBusiness|Store|Restaurant|Organization|Dentist|Physician|AutoRepair|Shop|Service|Hotel|School)',
                     j, re.I):
            local = True
    path = urlparse(url).path.rstrip("/")
    d = {"url": url, "status": status, "title": p.title[:300], "description": p.description[:400], "h1": p.h1[:5], "h2": p.h2[:15],
         "words": words, "img": p.img, "img_no_alt": p.img_no_alt, "viewport": p.viewport, "https": url.startswith("https://"),
         "noindex": "noindex" in p.robots, "lang": p.lang, "bytes": size, "schema_local": local,
         "canonical_other": canon if canon and canon.rstrip("/") != url.rstrip("/") else "",
         "is_home": url.rstrip("/") == home.rstrip("/"), "is_contact": any(w in path.lower() for w in ("contact", "sampark")),
         "mixed": p.mixed if url.startswith("https://") else 0}
    host = urlparse(home).hostname
    links = []
    for h in p.links:
        u = urldefrag(urljoin(url, h))[0]
        pu = urlparse(u)
        if pu.scheme in ("http", "https") and pu.hostname and pu.hostname.removeprefix("www.") == (host or "").removeprefix("www.") \
                and not re.search(r"\.(jpg|jpeg|png|gif|webp|pdf|zip|mp4|svg|css|js)$", pu.path, re.I):
            links.append(u)
    d["internal_links"] = len(set(links))
    return d, links


# ---------------------------------------------------------------- site
def _site(s, tenant: str, ref: str = "") -> SeoSite:
    q = select(SeoSite).where(SeoSite.tenant_id == tenant)
    rows = s.scalars(q.order_by(SeoSite.created_at)).all()
    if not rows:
        raise SeoError("Pehle website jodiye: 'meri website example.in hai'.")
    if not ref:
        return rows[0]
    for r in rows:
        if r.id == ref or ref.lower().strip("/ ") in r.url.lower() or ref.lower() == r.name.lower():
            return r
    raise SeoError(f"Website '{ref}' nahi mili.")


def _site_view(r: SeoSite) -> dict:
    return {"id": r.id, "url": r.url, "name": r.name, "business": _j(r.business, {}), "services": _j(r.services, []),
            "gsc_property": r.gsc_property, "score": r.score, "last_crawl": r.last_crawl, "cwv": _j(r.cwv, {})}


def save_site(tenant: str, url: str = "", name: str = "", site: str = "", services: list | None = None, **biz: Any) -> dict:
    with _LOCK, session() as s:
        r = None
        if site:
            r = _site(s, tenant, site)
        elif url:
            host = (urlparse(url if "://" in url else "https://" + url).hostname or "").removeprefix("www.")
            r = next((x for x in s.scalars(select(SeoSite).where(SeoSite.tenant_id == tenant)).all() if host and host in x.url), None)
        if r is None:
            if not url:
                raise SeoError("Website ka pata (URL) chahiye.")
            r = SeoSite(tenant_id=tenant)
            s.add(r)
        if url:
            u = url.strip()
            u = u if "://" in u else "https://" + u
            p = urlparse(u)
            if not p.hostname or "." not in p.hostname:
                raise SeoError(f"URL samajh nahi aaya: {url}")
            r.url = f"{p.scheme}://{p.hostname}/"
        if name:
            r.name = str(name)[:160]
        b = _j(r.business, {})
        for k in ("city", "state", "street", "pin", "phone", "schema_type", "hours", "price_range", "lat", "lng", "same_as"):
            if biz.get(k) not in (None, ""):
                b[k] = biz[k]
        r.business = json.dumps(b, ensure_ascii=False)
        if services is not None:
            r.services = json.dumps([str(x).strip()[:80] for x in services if str(x).strip()][:30], ensure_ascii=False)
        s.flush()
        v = _site_view(r)
    return {"site": v, "say": f"Website {v['url']} jodi" + (f" ({v['name']}, {v['business'].get('city')})" if v["business"].get("city") else "") + "."}


def sites(tenant: str) -> list[dict]:
    with session() as s:
        return [_site_view(r) for r in s.scalars(select(SeoSite).where(SeoSite.tenant_id == tenant).order_by(SeoSite.created_at)).all()]


# ---------------------------------------------------------------- crawl + audit
def crawl(tenant: str, site: str = "", max_pages: int = 25, conv: str = "", cap: int = 60) -> dict:
    """cap: 60 for a one-off check; the hired SEO Employee (v0.314 seo/desk.py) crawls up to 500."""
    with session() as s:
        r = _site(s, tenant, site)
        sid, home, city = r.id, r.url, _j(r.business, {}).get("city", "")
    max_pages = max(1, min(int(max_pages or 25), max(1, min(int(cap or 60), 500))))
    robots_txt, sitemap_ok, seeds = "", False, [home]
    sitemap_urls: set[str] = set()
    try:
        st_, _, robots_txt, _ = _fetch(urljoin(home, "/robots.txt"), 200_000)
        if st_ >= 400:
            robots_txt = ""
    except SeoError:
        robots_txt = ""
    rp = RobotFileParser()
    rp.parse(robots_txt.splitlines())
    maps = re.findall(r"^\s*sitemap:\s*(\S+)", robots_txt, re.I | re.M) or [urljoin(home, "/sitemap.xml")]
    for sm in maps[:3]:
        try:
            st_, _, xml, _ = _fetch(sm, 3_000_000)
        except SeoError:
            continue
        if st_ < 400 and "<loc>" in xml:
            sitemap_ok = True
            locs = re.findall(r"<loc>\s*([^<\s]+)\s*</loc>", xml)
            subs = [u for u in locs if u.endswith(".xml")]
            seeds += [u for u in locs if not u.endswith(".xml")]
            sitemap_urls.update(u for u in locs if not u.endswith(".xml"))
            for sub in subs[:3]:
                try:
                    st2, _, x2, _ = _fetch(sub, 3_000_000)
                    if st2 < 400:
                        sub_locs = [u for u in re.findall(r"<loc>\s*([^<\s]+)\s*</loc>", x2) if not u.endswith(".xml")]
                        seeds += sub_locs
                        sitemap_urls.update(sub_locs)
                except SeoError:
                    pass
    host = (urlparse(home).hostname or "").removeprefix("www.")
    queue, seen, pages, blocked = [], set(), [], []
    depth: dict[str, int] = {}
    redirects: dict[str, str] = {}                  # v0.314 every internal redirect seen, even to a page already checked
    for u in seeds:
        u = urldefrag(u)[0]
        if (urlparse(u).hostname or "").removeprefix("www.") == host and u not in seen:
            seen.add(u)
            queue.append(u)
    depth[urldefrag(home)[0]] = 0
    while queue and len(pages) < max_pages:
        u = queue.pop(0)
        if robots_txt and not rp.can_fetch(UA, u) and not rp.can_fetch("*", u):
            blocked.append(u)
            continue
        try:
            status, final, html, size = _fetch(u)
        except SeoError as e:
            pages.append({"url": u, "status": 599, "error": str(e)})
            continue
        if final != u:
            redirects[u] = final
        if final != u and any(x.get("url") == final for x in pages):
            continue                                            # a redirect to a page already checked
        d, links = _parse(final, status, html, size, home, city)
        if final != u:
            d["redirected_from"] = u
            seen.add(final)
        d["depth"] = depth.get(u, depth.get(final))           # clicks from the home page (None = only in the sitemap)
        d["out"] = sorted({l2.split("?")[0] for l2 in links})[:300]
        pages.append(d)
        for l2 in links:
            l2 = l2.split("?")[0]
            if d["depth"] is not None and (l2 not in depth or depth[l2] > d["depth"] + 1):
                depth[l2] = d["depth"] + 1
            if l2 not in seen and len(seen) < max_pages * 4:
                seen.add(l2)
                queue.append(l2)
        if _http is None:
            time.sleep(_PAUSE)
    if not pages:
        raise SeoError("Website ka ek bhi page nahi khula.")
    all_issues = []
    stamp = now().isoformat(timespec="seconds")
    with _LOCK, session() as s:
        for old in s.scalars(select(SeoPage).where(SeoPage.site_id == sid)).all():
            s.delete(old)
        for p in pages:
            iss = R.page_issues(p, city) if p.get("status", 599) != 599 else [
                {"code": "unreachable", "level": "bad", "text": f"Page nahi khula: {p.get('error', '')}", "fix": "Server / link check karein."}]
            for i in iss:
                i["url"] = p["url"]
            all_issues += iss
            s.add(SeoPage(tenant_id=tenant, site_id=sid, url=p["url"][:600], data=json.dumps(p, ensure_ascii=False),
                          issues=json.dumps(iss, ensure_ascii=False), score=R.score(iss)))
        for i in R.site_issues(pages, robots_txt, sitemap_ok):
            i["url"] = home
            all_issues.append(i)
        all_issues += R.link_issues(pages, sitemap_urls, home, redirects)   # v0.314 broken / redirected links, orphans, depth, dups
        # tasks: keep her decisions, add new, mark fixed what disappeared
        tasks = {t.key: t for t in s.scalars(select(SeoTask).where(SeoTask.site_id == sid)).all()}
        now_keys = set()
        new = 0
        for i in all_issues:
            k = f"{i['code']}|{i['url']}"[:700]
            now_keys.add(k)
            t = tasks.get(k)
            if t is None:
                s.add(SeoTask(tenant_id=tenant, site_id=sid, key=k, code=i["code"], url=i["url"][:600], level=i["level"], text=i["text"],
                              fix=i["fix"], seen=stamp))
                new += 1
            else:
                t.text, t.fix, t.seen = i["text"], i["fix"], stamp
                if t.status in ("fixed", "done"):
                    t.status = "open"          # came back / not really fixed
        fixed = 0
        for k, t in tasks.items():
            if k not in now_keys and t.status in ("open", "done"):
                t.status = "fixed"
                fixed += 1
        site_score = round(sum(R.score([x for x in all_issues if x["url"] == p["url"]]) for p in pages) / len(pages))
        site_score = max(0, site_score - 10 * sum(1 for i in all_issues if i["code"] in ("robots_block",)))
        r = s.get(SeoSite, sid)
        r.score, r.last_crawl = site_score, stamp
    bad = sum(1 for i in all_issues if i["level"] == "bad")
    warn = sum(1 for i in all_issues if i["level"] == "warn")
    say = (f"{len(pages)} page jaanche. SEO score {site_score}/100. {bad} zaroori, {warn} dhyaan wali baatein"
           + (f"; {new} nayi" if new else "") + (f"; {fixed} theek ho gayi" if fixed else "") + ".")
    if blocked:
        say += f" {len(blocked)} page robots.txt ki wajah se chhode."
    _publish(tenant, "seo_update", {"site": sid, "score": site_score, "bad": bad, "warn": warn, "pages": len(pages)}, conv)
    return {"pages": len(pages), "score": site_score, "bad": bad, "warn": warn, "new": new, "fixed": fixed, "sitemap": sitemap_ok,
            "blocked": blocked[:20], "top": sorted(all_issues, key=lambda i: {"bad": 0, "warn": 1, "info": 2}[i["level"]])[:8], "say": say}


def pages(tenant: str, site: str = "") -> list[dict]:
    with session() as s:
        sid = _site(s, tenant, site).id
        out = []
        for p in s.scalars(select(SeoPage).where(SeoPage.site_id == sid)).all():
            d = _j(p.data, {})
            out.append({"url": p.url, "title": d.get("title", ""), "words": d.get("words", 0), "status": d.get("status"), "score": p.score,
                        "issues": _j(p.issues, [])})
    return sorted(out, key=lambda x: x["score"])


def tasks(tenant: str, site: str = "", status: str = "open") -> list[dict]:
    with session() as s:
        sid = _site(s, tenant, site).id
        q = select(SeoTask).where(SeoTask.site_id == sid)
        if status:
            q = q.where(SeoTask.status == status)
        rows = s.scalars(q).all()
        out = [{"id": t.id, "code": t.code, "url": t.url, "level": t.level, "text": t.text, "fix": t.fix, "status": t.status} for t in rows]
    return sorted(out, key=lambda t: ({"bad": 0, "warn": 1, "info": 2}.get(t["level"], 3), t["url"]))


def close_task(tenant: str, task_id: str, status: str = "done") -> dict:
    if status not in ("done", "ignored", "open"):
        raise SeoError("status: done | ignored | open")
    with session() as s:
        t = s.get(SeoTask, task_id)
        if not t or t.tenant_id != tenant:
            raise SeoError("Kaam nahi mila.")
        t.status = status
    return {"say": "Note kiya; agli jaanch mein dobara dekhunga." if status == "done" else "Note kiya."}


# ---------------------------------------------------------------- PageSpeed (Core Web Vitals)
def pagespeed(tenant: str, site: str = "", url: str = "", strategy: str = "mobile") -> dict:
    from ..config import get_settings
    with session() as s:
        r = _site(s, tenant, site)
        sid, target = r.id, url or r.url
    _safe(target)
    q = {"url": target, "strategy": "desktop" if strategy == "desktop" else "mobile", "category": "performance"}
    key = getattr(get_settings(), "pagespeed_api_key", "") or ""
    if key:
        q["key"] = key
    c, own = _client()
    try:
        resp = c.get(PSI_API + "?" + urlencode(q), timeout=90)
    except httpx.HTTPError as e:
        raise SeoError(f"PageSpeed se baat nahi hui ({type(e).__name__}).")
    finally:
        if own:
            c.close()
    if resp.status_code != 200:
        raise SeoError(f"PageSpeed ne jawab nahi diya ({resp.status_code}). Thodi der baad; server par PAGESPEED_API_KEY lagane se limit badhti hai.")
    js = resp.json()
    m = (js.get("loadingExperience") or {}).get("metrics") or {}
    origin = False
    if not m:
        m = (js.get("originLoadingExperience") or {}).get("metrics") or {}
        origin = bool(m)

    def pct(k):
        v = (m.get(k) or {}).get("percentile")
        return v

    lcp, inp, cls = pct("LARGEST_CONTENTFUL_PAINT_MS"), pct("INTERACTION_TO_NEXT_PAINT"), pct("CUMULATIVE_LAYOUT_SHIFT_SCORE")
    cls = cls / 100 if cls is not None else None
    lh = js.get("lighthouseResult") or {}
    perf = round(((lh.get("categories") or {}).get("performance") or {}).get("score", 0) * 100) if lh else None
    ops = []
    for a in (lh.get("audits") or {}).values():
        det = a.get("details") or {}
        if det.get("type") == "opportunity" and (det.get("overallSavingsMs") or 0) >= 300:
            ops.append({"title": a.get("title", ""), "save_ms": round(det["overallSavingsMs"])})
    ops.sort(key=lambda o: -o["save_ms"])
    grades = {"lcp": R.cwv_grade("lcp", lcp), "inp": R.cwv_grade("inp", inp), "cls": R.cwv_grade("cls", cls)}
    field = any(v != "na" for v in grades.values())
    out = {"url": target, "strategy": q["strategy"], "lcp_ms": lcp, "inp_ms": inp, "cls": cls, "grades": grades, "field_data": field,
           "origin_level": origin, "lab_score": perf, "fixes": ops[:6], "checked": today(tenant).isoformat()}
    with session() as s:
        row = s.get(SeoSite, sid)
        cw = _j(row.cwv, {})
        cw[q["strategy"]] = out
        row.cwv = json.dumps(cw, ensure_ascii=False)
    names = {"good": "achha", "ni": "sudhaar chahiye", "poor": "kharab", "na": "data nahi"}
    if field:
        say = (f"{q['strategy']} par asli graahakon ka data: LCP {lcp / 1000 if lcp else 0:.1f}s ({names[grades['lcp']]}), "
               f"INP {inp or 0}ms ({names[grades['inp']]}), CLS {cls if cls is not None else '-'} ({names[grades['cls']]}).")
    else:
        say = "Is site par abhi kam traffic hai, Chrome ka asli data nahi; lab test se andaza."
    if perf is not None:
        say += f" Lab score {perf}/100."
    if ops:
        say += " Sabse bada sudhaar: " + ops[0]["title"] + "."
    return {**out, "say": say}


# ---------------------------------------------------------------- Search Console (official API, her read-only sign-in)
def gsc_ready() -> bool:
    from ..mailbox import accounts as mail_acc
    return mail_acc.oauth_ready("google")


def is_link_state(state: str) -> bool:
    return (state or "").startswith("gsc_")


def start_link(tenant: str) -> str:
    from .. import accounts as acc
    from ..config import get_settings
    from ..db import MailLinkState
    from ..mailbox import accounts as mail_acc
    if not gsc_ready():
        raise SeoError("Search Console jodne ke liye server par Google sign-in (GOOGLE_OAUTH_CLIENT_ID / SECRET, PUBLIC_BASE_URL) chahiye.")
    st = get_settings()
    state = "gsc_" + secrets.token_urlsafe(24)
    with session() as s:
        s.add(MailLinkState(state=state, tenant_id=tenant, provider="seo"))
    q = {"client_id": st.google_oauth_client_id, "redirect_uri": acc._redirect_uri(), "response_type": "code",
         "scope": " ".join(["openid", "email", GSC_SCOPE, GA4_SCOPE]), "access_type": "offline", "prompt": "consent", "state": state,
         "include_granted_scopes": "true"}
    return mail_acc.GOOGLE_AUTH + "?" + urlencode(q)


def finish_link(code: str, state: str) -> dict:
    from .. import accounts as acc
    from ..config import get_settings
    from ..db import MailLinkState, as_utc
    from ..mailbox import accounts as mail_acc
    st = get_settings()
    with session() as s:
        row = s.get(MailLinkState, state)
        if row is None or row.provider != "seo" or as_utc(row.created_at) < now() - timedelta(minutes=30):
            raise SeoError("Yeh link purana ho gaya. App se dobara 'Search Console jodiye' dabaiye.")
        tenant = row.tenant_id
        s.delete(row)
    try:
        tok = mail_acc._post(mail_acc.GOOGLE_TOKEN, {"code": code, "client_id": st.google_oauth_client_id,
                                                     "client_secret": st.google_oauth_client_secret,
                                                     "redirect_uri": acc._redirect_uri(), "grant_type": "authorization_code"})
        if GSC_SCOPE not in (tok.get("scope") or "").split():
            raise SeoError("Search Console dekhne ki permission nahi di gayi.")
        info = mail_acc._get("https://openidconnect.googleapis.com/v1/userinfo", tok["access_token"])
    except mail_acc.MailError as e:
        raise SeoError(str(e))
    if not tok.get("refresh_token"):
        raise SeoError("Google se refresh token nahi mila; dobara jodiye.")
    with session() as s:
        g = s.get(SeoGoogle, tenant) or SeoGoogle(tenant_id=tenant)
        s.add(g)
        g.address = (info.get("email") or "").lower()
        g.access_token_enc = acc.encrypt(tok["access_token"])
        g.refresh_token_enc = acc.encrypt(tok["refresh_token"])
        g.expiry = now() + timedelta(seconds=int(tok.get("expires_in", 3600)))
        g.status, g.last_error = "ok", ""
        g.scopes = tok.get("scope") or ""
        addr = g.address
    try:
        _auto_property(tenant)
    except SeoError:
        pass
    return {"address": addr}


def _token(tenant: str, force: bool = False) -> str:
    from .. import accounts as acc
    from ..config import get_settings
    from ..db import as_utc
    from ..mailbox import accounts as mail_acc
    with session() as s:
        g = s.get(SeoGoogle, tenant)
        if g is None:
            raise SeoError("Search Console abhi juda nahi. App mein SEO > 'Search Console jodiye' dabaiye.")
        exp = as_utc(g.expiry)
        if not force and g.access_token_enc and exp and exp > now() + timedelta(seconds=90):
            return acc.decrypt(g.access_token_enc) or ""
        refresh = acc.decrypt(g.refresh_token_enc)
    st = get_settings()
    try:
        tok = mail_acc._post(mail_acc.GOOGLE_TOKEN, {"client_id": st.google_oauth_client_id, "client_secret": st.google_oauth_client_secret,
                                                     "refresh_token": refresh, "grant_type": "refresh_token"})
    except Exception as e:
        with session() as s:
            g = s.get(SeoGoogle, tenant)
            if g:
                g.status, g.last_error = "error", str(e)[:400]
        raise SeoError("Search Console ka login expire ho gaya; dobara jodiye.")
    with session() as s:
        g = s.get(SeoGoogle, tenant)
        g.access_token_enc = acc.encrypt(tok["access_token"])
        g.expiry = now() + timedelta(seconds=int(tok.get("expires_in", 3600)))
        g.status = "ok"
    return tok["access_token"]


def _gsc(tenant: str, method: str, path: str, body: dict | None = None) -> dict:
    c, own = _client()
    try:
        for force in (False, True):
            h = {"Authorization": f"Bearer {_token(tenant, force)}"}
            try:
                r = c.request(method, GSC_API + path, headers=h, json=body, timeout=60)
            except httpx.HTTPError as e:
                raise SeoError(f"Search Console se baat nahi hui ({type(e).__name__}).")
            if r.status_code == 401 and not force:
                continue
            if r.status_code == 403:
                raise SeoError("Is Google account ke paas is website ka Search Console access nahi.")
            if r.status_code >= 400:
                raise SeoError(f"Search Console: {r.text[:200]}")
            return r.json()
    finally:
        if own:
            c.close()
    raise SeoError("Search Console login nahi chala.")


def gsc_status(tenant: str) -> dict:
    with session() as s:
        g = s.get(SeoGoogle, tenant)
        return {"connected": bool(g), "address": g.address if g else "", "status": g.status if g else "", "can_link": gsc_ready(),
                "analytics": bool(g and GA4_SCOPE in (g.scopes or "").split())}


def gsc_properties(tenant: str) -> list[dict]:
    return [{"property": x.get("siteUrl", ""), "level": x.get("permissionLevel", "")} for x in _gsc(tenant, "GET", "/sites").get("siteEntry", [])]


def _auto_property(tenant: str, site: str = "") -> str:
    props = [p for p in gsc_properties(tenant) if p["level"] != "siteUnverifiedUser"]
    with session() as s:
        r = _site(s, tenant, site)
        host = (urlparse(r.url).hostname or "").removeprefix("www.")
        pick = next((p["property"] for p in props if p["property"] == f"sc-domain:{host}"), "") or \
            next((p["property"] for p in props if host and host in p["property"]), "")
        if pick:
            r.gsc_property = pick
        return pick


def performance(tenant: str, site: str = "", days: int = 28, conv: str = "", store: bool = True) -> dict:
    """Last `days` (final data ends ~3 days ago) vs the period before: totals, top queries / pages, opportunities."""
    with session() as s:
        r = _site(s, tenant, site)
        sid, prop = r.id, r.gsc_property
    if not prop:
        prop = _auto_property(tenant, site)
        if not prop:
            raise SeoError("Search Console mein yeh website nahi mili. Pehle search.google.com/search-console par website jodkar verify kijiye.")
    days = max(7, min(int(days or 28), 90))
    end = today(tenant) - timedelta(days=3)
    start = end - timedelta(days=days - 1)
    p_end, p_start = start - timedelta(days=1), start - timedelta(days=days)
    path = f"/sites/{quote(prop, safe='')}/searchAnalytics/query"

    def q(a: date, b: date, dims: list[str], limit: int = 1000) -> list[dict]:
        rows = _gsc(tenant, "POST", path, {"startDate": a.isoformat(), "endDate": b.isoformat(), "dimensions": dims, "rowLimit": limit,
                                           "type": "web"}).get("rows", [])
        return [{**{d: k for d, k in zip(dims, x.get("keys", []))}, "clicks": int(x.get("clicks", 0)), "impressions": int(x.get("impressions", 0)),
                 "ctr": float(x.get("ctr", 0)), "position": float(x.get("position", 0))} for x in rows]

    cur_tot = q(start, end, ["date"], 100)
    prev_tot = q(p_start, p_end, ["date"], 100)
    qp = q(start, end, ["query", "page"], 2500)

    def tot(rows):
        c, i = sum(x["clicks"] for x in rows), sum(x["impressions"] for x in rows)
        pos = sum(x["position"] * x["impressions"] for x in rows) / i if i else 0
        return {"clicks": c, "impressions": i, "ctr": round(c / i, 4) if i else 0, "position": round(pos, 1)}

    cur, prev = tot(cur_tot), tot(prev_tot)
    queries: dict[str, dict] = {}
    for x in qp:
        a = queries.setdefault(x["query"], {"query": x["query"], "clicks": 0, "impressions": 0, "_pos": 0.0, "page": x["page"], "_best": 0})
        a["clicks"] += x["clicks"]
        a["impressions"] += x["impressions"]
        a["_pos"] += x["position"] * x["impressions"]
        if x["impressions"] > a["_best"]:
            a["page"], a["_best"] = x["page"], x["impressions"]
    qs = []
    for a in queries.values():
        a["position"] = round(a["_pos"] / a["impressions"], 1) if a["impressions"] else 0
        a["ctr"] = round(a["clicks"] / a["impressions"], 4) if a["impressions"] else 0
        a["intent"] = R.intent(a["query"])
        a.pop("_pos")
        a.pop("_best")
        qs.append(a)
    qs.sort(key=lambda a: -a["clicks"] * 1000 - a["impressions"])
    opps = []
    for a in qs:
        o = R.opportunity(a)
        if o:
            opps.append({**o, "query": a["query"], "page": a["page"], "position": a["position"], "impressions": a["impressions"]})
    opps.sort(key=lambda o: -o["gain"])
    cann = R.cannibal(qp)
    mix: dict[str, int] = {}
    for a in qs:
        mix[a["intent"]] = mix.get(a["intent"], 0) + a["impressions"]
    long_q = sum(1 for a in qs if len(a["query"].split()) >= 6)
    if store:          # v0.314 the SEO Employee's daily pull reads numbers only; the weekly snapshot stays the trend
        with session() as s:
            s.add(SeoStat(tenant_id=tenant, site_id=sid, start=start.isoformat(), end=end.isoformat(), clicks=cur["clicks"],
                          impressions=cur["impressions"], ctr=cur["ctr"], position=cur["position"],
                          rows=json.dumps(qs[:300], ensure_ascii=False)))

    def ch(a, b):
        return round((a - b) * 100 / b) if b else None

    delta = {"clicks": ch(cur["clicks"], prev["clicks"]), "impressions": ch(cur["impressions"], prev["impressions"])}
    say = (f"Pichhle {days} din ({start.strftime('%d %b')} - {end.strftime('%d %b')}): {cur['clicks']} click, {cur['impressions']} baar dikhe, "
           f"CTR {cur['ctr'] * 100:.1f}%, ausat position {cur['position']}.")
    if delta["clicks"] is not None:
        say += f" Click {'+' if delta['clicks'] >= 0 else ''}{delta['clicks']}% pichhle {days} din se."
    if opps:
        say += f" {len(opps)} mauke; sabse bada: {opps[0]['do']}"
    if long_q:
        say += f" {long_q} lambe sawaal-jaise search (AI Mode / voice) bhi aa rahe hain: FAQ wale jawab madad karenge."
    if delta["clicks"] is not None and delta["clicks"] <= -25:
        _publish(tenant, "seo_alert", {"site": sid, "text": f"Google se click {abs(delta['clicks'])}% gire."}, conv)
    return {"property": prop, "start": start.isoformat(), "end": end.isoformat(), "current": cur, "previous": prev, "change": delta,
            "queries": qs[:50], "opportunities": opps[:15], "cannibal": cann[:10], "intent_mix": mix, "long_queries": long_q, "say": say}


# ---------------------------------------------------------------- keywords, content, schema, GBP drafts
def keyword_plan(tenant: str, site: str = "") -> dict:
    with session() as s:
        r = _site(s, tenant, site)
        svc, city, sid = _j(r.services, []), _j(r.business, {}).get("city", ""), r.id
        last = s.scalars(select(SeoStat).where(SeoStat.site_id == sid).order_by(SeoStat.end.desc())).first()
        known = {x["query"]: x for x in _j(last.rows, [])} if last else {}
    if not svc:
        raise SeoError("Pehle bataiye aap kya kya kaam / product bechte hain (services).")
    out = []
    for sv in svc:
        base = sv.lower()
        seeds = [f"{base} in {city.lower()}" if city else base, f"{base} near me", f"best {base} in {city.lower()}" if city else f"best {base}",
                 f"{base} price in {city.lower()}" if city else f"{base} price"]
        for k in seeds:
            hit = known.get(k) or next((v for q, v in known.items() if base in q and (not city or city.lower() in q)), None)
            out.append({"keyword": k, "service": sv, "intent": R.intent(k), "position": hit["position"] if hit else None,
                        "impressions": hit["impressions"] if hit else 0})
    return {"keywords": out, "say": f"{len(out)} keyword {len(svc)} kaam ke liye. Har kaam ka alag page ho, title mein kaam + shehar. "
            "Search volume ke liye Google Ads Keyword Planner (free account) dekhiye; asli data Search Console se aata hai."}


def _draft(tenant: str, sid: str, kind: str, title: str, text: str, conv: str = "") -> dict:
    with session() as s:
        d = SeoDraft(tenant_id=tenant, site_id=sid, kind=kind, title=title[:200], text=text)
        s.add(d)
        s.flush()
        did = d.id
    _publish(tenant, "seo_draft", {"draft_id": did, "kind": kind, "title": title[:200]}, conv)
    return {"draft_id": did, "kind": kind, "title": title, "text": text}


def content_brief(tenant: str, keyword: str, site: str = "", conv: str = "") -> dict:
    if not str(keyword).strip():
        raise SeoError("Kis shabd / kaam ke liye page?")
    with session() as s:
        r = _site(s, tenant, site)
        sid, name, city = r.id, r.name or (urlparse(r.url).hostname or ""), _j(r.business, {}).get("city", "")
        ps = [(p.url, _j(p.data, {}).get("title", "")) for p in s.scalars(select(SeoPage).where(SeoPage.site_id == sid)).all()]
    kw = re.sub(r"\s+", " ", keyword).strip()
    it = R.intent(kw)
    svc = re.sub(r"\b(in|near me|best|price|cost|rate)\b.*$", "", kw, flags=re.I).strip() or kw
    title = R.title_suggestion(svc.title(), city, name)
    desc = (f"{svc.title()} {('in ' + city) if city else ''}: daam, samay aur kaam ki poori jaankari. {name} se aaj hi call / WhatsApp karein.")[:158]
    h2 = [f"{svc.title()} mein kya milta hai", f"Daam / charges", "Kaam kaise hota hai (step by step)", "Hamare kaam ki photo / pehle-baad",
          f"{city or 'Aapke shehar'} mein kahan, kitne samay mein", "Graahak kya kehte hain", "Aksar puchhe jaane wale sawaal (FAQ)"]
    faq = [f"{svc} ka kharcha kitna hai?", f"{svc} mein kitna samay lagta hai?", "Kya ghar / jagah par aakar karte hain?",
           "Payment kaise kar sakte hain?", "Guarantee / warranty milti hai?"]
    links = [u for u, t in ps if any(w in (t or "").lower() for w in svc.lower().split() if len(w) > 3)][:4]
    words = 900 if it in ("commercial", "informational") else 600
    text = "\n".join([f"Keyword: {kw} ({it})", f"Title: {title}", f"Description: {desc}", f"H1: {svc.title()}{(' in ' + city) if city else ''}",
                      "H2:", *[f"- {h}" for h in h2], "FAQ (FAQPage schema ke saath):", *[f"- {q}" for q in faq],
                      f"Lambai: ~{words} shabd, apni bhasha mein, asli daam / photo / anubhav ke saath (copy-paste ya AI ka bhara hua text nahi).",
                      "Andar ke links: " + (", ".join(links) if links else "home page aur contact page se is page par link dijiye."),
                      "Page ke neeche: phone, WhatsApp button, address + map."])
    out = _draft(tenant, sid, "brief", f"Page plan: {kw}", text, conv)
    return {**out, "say": f"'{kw}' ke liye page ka plan taiyaar (title, headings, FAQ). Aapki haan ke baad final; website par aap / web wale lagayenge."}


def schema_draft(tenant: str, site: str = "", conv: str = "") -> dict:
    with session() as s:
        r = _site(s, tenant, site)
        b = {**_j(r.business, {}), "name": r.name, "url": r.url}
        sid = r.id
    if not b.get("name"):
        raise SeoError("Business ka naam bataiye.")
    code = R.local_schema(b)
    out = _draft(tenant, sid, "schema", "LocalBusiness schema", code, conv)
    missing = [k for k in ("phone", "street", "city", "pin", "hours") if not b.get(k)]
    say = "Home page ke <head> mein lagane ka JSON-LD taiyaar."
    if missing:
        say += " Aur behtar ke liye bataiye: " + ", ".join(missing) + "."
    return {**out, "say": say + " Lagane ke baad Google Rich Results Test se jaanch lijiye."}


def gbp_post(tenant: str, topic: str, offer: str = "", site: str = "", conv: str = "") -> dict:
    with session() as s:
        r = _site(s, tenant, site)
        sid, name, b = r.id, r.name, _j(r.business, {})
    t = str(topic).strip()
    if not t:
        raise SeoError("Post kis baare mein?")
    body = f"{t}" + (f" - {offer}" if offer else "") + f". {name}{(', ' + b['city']) if b.get('city') else ''}."
    body += f" Abhi call karein: {b['phone']}." if b.get("phone") else " Abhi call karein."
    body = body[:1400]
    out = _draft(tenant, sid, "gbp_post", f"Google post: {t[:60]}", body + "\n(Button: Call now / Book; ek asli photo lagaiye.)", conv)
    return {**out, "say": "Google Business Profile post ka draft taiyaar; haan ke baad aap profile par daalein."}


def review_reply(tenant: str, review: str, stars: int = 5, reviewer: str = "", site: str = "", conv: str = "") -> dict:
    with session() as s:
        r = _site(s, tenant, site)
        sid, name, phone = r.id, r.name, _j(r.business, {}).get("phone", "")
    if not str(review).strip() and not stars:
        raise SeoError("Review ka text ya stars bataiye.")
    who = (reviewer or "").split()[0] if reviewer else ""
    hi = f"Namaste{(' ' + who) if who else ''} ji,"
    stars = int(stars or 0)
    if stars >= 4:
        text = f"{hi} itne achhe shabdon ke liye bahut shukriya! Aapko hamara kaam pasand aaya, yeh sunkar khushi hui. Phir milenge. - {name}"
    elif stars == 3:
        text = (f"{hi} review ke liye shukriya. Hum aur behtar karna chahte hain: kya kami lagi, bataiye"
                + (f" ({phone})" if phone else "") + f". - {name}")
    else:
        text = (f"{hi} hamein afsos hai ki aapka anubhav achha nahi raha. Hum ise theek karna chahte hain; kripya hamein seedha sampark karein"
                + (f" ({phone})" if phone else "") + f" taaki baat karke hal nikaal sakein. - {name}")
    out = _draft(tenant, sid, "review_reply", f"Review jawab ({stars}*)", text, conv)
    return {**out, "say": "Jawab ka draft taiyaar (bahas nahi, offline baat ka rasta). Haan ke baad aap Google par post karein."}


def approve_draft(tenant: str, draft_id: str, text: str = "") -> dict:
    with session() as s:
        d = s.get(SeoDraft, draft_id)
        if not d or d.tenant_id != tenant:
            raise SeoError("Draft nahi mila.")
        if text:
            d.text = text
        d.status = "approved"
        t, k = d.text, d.kind
    where = {"gbp_post": "Google Business Profile > Add update", "review_reply": "Google Maps / Business Profile > Reviews > Reply",
             "schema": "website ke home page <head> mein", "brief": "website ka naya page", "meta": "page ka title / description"}.get(k, "website")
    return {"text": t, "say": f"Final. Ab ise {where} par aap lagaiye (Hoola khud publish nahi karta)."}


def drafts(tenant: str, site: str = "", status: str = "") -> list[dict]:
    with session() as s:
        sid = _site(s, tenant, site).id
        q = select(SeoDraft).where(SeoDraft.site_id == sid)
        if status:
            q = q.where(SeoDraft.status == status)
        return [{"id": d.id, "kind": d.kind, "title": d.title, "text": d.text, "status": d.status}
                for d in s.scalars(q.order_by(SeoDraft.created_at.desc())).all()][:50]


# ---------------------------------------------------------------- report, morning
def summary(tenant: str, site: str = "") -> dict:
    with session() as s:
        r = _site(s, tenant, site)
        v = _site_view(r)
        stats = [{"start": x.start, "end": x.end, "clicks": x.clicks, "impressions": x.impressions, "ctr": x.ctr, "position": x.position}
                 for x in s.scalars(select(SeoStat).where(SeoStat.site_id == r.id).order_by(SeoStat.end.desc())).all()[:12]]
    t = tasks(tenant, site, "open")
    return {"site": v, "open_tasks": len(t), "bad": sum(1 for x in t if x["level"] == "bad"), "top_tasks": t[:10], "stats": stats,
            "gsc": gsc_status(tenant), "gbp_checklist": R.GBP_CHECKLIST, "never": list(R.NEVER)}


def report_file(tenant: str, site: str = "", conv: str = "") -> dict:
    import openpyxl
    from .. import chatfiles
    sm = summary(tenant, site)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Summary"
    v = sm["site"]
    ws.append(["Website", v["url"]])
    ws.append(["SEO score", v["score"]])
    ws.append(["Last crawl", v["last_crawl"]])
    for k, c in (v.get("cwv") or {}).items():
        ws.append([f"Core Web Vitals ({k})", f"LCP {c.get('lcp_ms')} ms, INP {c.get('inp_ms')} ms, CLS {c.get('cls')}, lab {c.get('lab_score')}"])
    ws2 = wb.create_sheet("Fix list")
    ws2.append(["Level", "Page", "Problem", "Fix", "Status"])
    for t in tasks(tenant, site, ""):
        ws2.append([t["level"], t["url"], t["text"], t["fix"], t["status"]])
    ws3 = wb.create_sheet("Pages")
    ws3.append(["URL", "Title", "Words", "HTTP", "Score"])
    for p in pages(tenant, site):
        ws3.append([p["url"], p["title"], p["words"], p["status"], p["score"]])
    ws4 = wb.create_sheet("Search Console")
    ws4.append(["Period start", "End", "Clicks", "Impressions", "CTR", "Position"])
    for x in sm["stats"]:
        ws4.append([x["start"], x["end"], x["clicks"], x["impressions"], x["ctr"], x["position"]])
    with session() as s:
        last = s.scalars(select(SeoStat).where(SeoStat.site_id == v["id"]).order_by(SeoStat.end.desc())).first()
        rows = _j(last.rows, []) if last else []
    if rows:
        ws5 = wb.create_sheet("Queries")
        ws5.append(["Query", "Intent", "Clicks", "Impressions", "CTR", "Position", "Page"])
        for q in rows[:300]:
            ws5.append([q["query"], q.get("intent", ""), q["clicks"], q["impressions"], q["ctr"], q["position"], q.get("page", "")])
    bio = io.BytesIO()
    wb.save(bio)
    f = chatfiles.store(tenant, "SEO_report.xlsx", bio.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    _publish(tenant, "pc_file", {**f, "source": "seo", "title": "SEO report"}, conv)
    return {"file": f, "say": "SEO report Excel mein (fix list, pages, Search Console)."}


def settings(tenant: str) -> dict:
    with session() as s:
        r = s.get(SeoSettings, tenant) or SeoSettings(tenant_id=tenant, enabled=True, weekday=0, morning_time="09:30", last_morning="")
        return {"enabled": r.enabled, "weekday": r.weekday, "morning_time": r.morning_time}


def save_settings(tenant: str, **kw: Any) -> dict:
    with session() as s:
        r = s.get(SeoSettings, tenant)
        if r is None:
            r = SeoSettings(tenant_id=tenant)
            s.add(r)
        if "enabled" in kw:
            r.enabled = bool(kw["enabled"])
        if kw.get("weekday") is not None:
            r.weekday = max(0, min(int(kw["weekday"]), 6))
        if kw.get("morning_time"):
            r.morning_time = str(kw["morning_time"]).zfill(5)[:5]
    return settings(tenant)


def due_mornings() -> list[str]:
    from ..life import tz_of
    out = []
    with session() as s:
        tenants = {x for x in s.scalars(select(SeoSite.tenant_id)).all()}
        for t in tenants:
            r = s.get(SeoSettings, t)
            if r is None:
                r = SeoSettings(tenant_id=t)
                s.add(r)
                s.flush()
            if not r.enabled:
                continue
            local = datetime.now(tz_of(t))
            if local.weekday() != r.weekday:
                continue
            try:
                hh, mm = [int(x) for x in (r.morning_time or "09:30").split(":")]
            except ValueError:
                hh, mm = 9, 30
            late = (local.hour * 60 + local.minute) - (hh * 60 + mm)
            d = local.date().isoformat()
            if r.last_morning != d and 0 <= late <= 180:
                r.last_morning = d
                out.append(t)
    return out


def send_morning(tenant: str) -> dict:
    """Weekly: fresh Search Console numbers (if connected) + the open fix list."""
    from ..events import notify
    from ..net.service import home_conv
    parts = []
    try:
        if gsc_status(tenant)["connected"]:
            p = performance(tenant, days=7)
            parts.append(f"Is hafte Google se {p['current']['clicks']} click" + (f" ({'+' if (p['change']['clicks'] or 0) >= 0 else ''}{p['change']['clicks']}%)"
                                                                                  if p["change"]["clicks"] is not None else ""))
            if p["opportunities"]:
                parts.append("Mauka: " + p["opportunities"][0]["do"])
    except SeoError as e:
        parts.append(str(e))
    try:
        t = tasks(tenant, "", "open")
        if t:
            parts.append(f"{len(t)} SEO kaam baaki ({sum(1 for x in t if x['level'] == 'bad')} zaroori)")
    except SeoError:
        return {"text": ""}
    if not parts:
        return {"text": ""}
    text = "SEO: " + "; ".join(parts) + "."
    conv = home_conv(tenant)
    try:
        notify(conv, text)
    except Exception:
        pass
    _publish(tenant, "seo_week", {"text": text}, conv)
    return {"text": text}
