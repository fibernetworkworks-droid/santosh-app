"""v0.108 Document desk: Hoola reads many documents fully, remembers them per job, finds things across them, and
updates the one that needs updating - on her PC, in code, after her Haan.

Files come from her PC (Santosh Desktop: search / send_file), from chat (file_id), or from a link she pastes (a public
Google Sheet / Drive / direct file link is downloaded here; a link that needs her login is opened in HER PC's Chrome
with pc_do and the download is picked up from the PC).

    add / read(part)   every file read in full (Word / Excel / PDF / scans via OCR / text), split into parts of 12,000
                       characters; numbered per chat for 7 days ("pehli wali file" = 1) and shown in the work diary
    search             across all of them (meaning + exact words)
    edit               Word: find / replace inside runs (bold / colours / fonts stay), new paragraphs copy their
                       neighbour's style; Excel: cell values / find-replace (formatting stays); PDF: corrected new PDF
                       (simple layout, said plainly); text / CSV: plain replace
    sheet_compare      Sheet Milaan (sheetmilaan.py): two sheets by MAC / serial / ID -> new / changed / missing
    decide             only from her tap: the NEW copy is saved into the PC's Santosh Inbox and opened there; the
                       original file is never touched (it stays as the backup). PC offline -> the copy comes to her phone.
"""
from __future__ import annotations

import asyncio
import hashlib
import io
import json
import logging
import re
from datetime import datetime, timedelta, timezone
from pathlib import Path
from typing import Any

from sqlalchemy import Integer, String, Text, DateTime, select
from sqlalchemy.orm import Mapped, mapped_column

from .db import Base, new_id, now, session, as_utc
from .events import bus

log = logging.getLogger("santosh.docdesk")
PART = 12000
MAX_TEXT = 1_500_000
MAX_BYTES = 40 * 1024 * 1024
KEEP_DAYS = 7
STATS = {"read": 0, "edits": 0, "saved": 0, "sheet": 0, "links": 0}


class DeskError(ValueError):
    pass


class DeskFile(Base):
    __tablename__ = "desk_files"
    id: Mapped[str] = mapped_column(String(40), primary_key=True, default=lambda: new_id("df"))
    tenant_id: Mapped[str] = mapped_column(String(64), index=True)
    conversation_id: Mapped[str] = mapped_column(String(64), default="", index=True)
    name: Mapped[str] = mapped_column(String(255), default="")
    source: Mapped[str] = mapped_column(String(10), default="file")      # pc | file | link | made
    pc_path: Mapped[str] = mapped_column(Text, default="")
    file_id: Mapped[str] = mapped_column(String(40), default="")         # latest server copy (PcFile)
    url: Mapped[str] = mapped_column(Text, default="")
    kind: Mapped[str] = mapped_column(String(12), default="other")      # excel | csv | word | pdf | image | text | slides | other
    size: Mapped[int] = mapped_column(Integer, default=0)
    sha: Mapped[str] = mapped_column(String(64), default="")
    chars: Mapped[int] = mapped_column(Integer, default=0)
    parts: Mapped[int] = mapped_column(Integer, default=0)
    text: Mapped[str] = mapped_column(Text, default="")
    notes: Mapped[str] = mapped_column(Text, default="[]")
    added_at: Mapped[datetime] = mapped_column(DateTime(timezone=True), default=now)
    used_at: Mapped[datetime] = mapped_column(DateTime(timezone=True), default=now)


class DeskEdit(Base):
    """One proposed change (a document edit or a Sheet Milaan) waiting for her Haan."""
    __tablename__ = "desk_edits"
    id: Mapped[str] = mapped_column(String(40), primary_key=True, default=lambda: new_id("de"))
    tenant_id: Mapped[str] = mapped_column(String(64), index=True)
    conversation_id: Mapped[str] = mapped_column(String(64), default="")
    mode: Mapped[str] = mapped_column(String(8), default="doc")           # doc | sheet
    file_ref: Mapped[str] = mapped_column(String(40), default="")         # DeskFile being updated
    other_ref: Mapped[str] = mapped_column(String(40), default="")        # Sheet Milaan: the new sheet
    name: Mapped[str] = mapped_column(String(255), default="")
    kind: Mapped[str] = mapped_column(String(12), default="")
    ops: Mapped[str] = mapped_column(Text, default="{}")
    result: Mapped[str] = mapped_column(Text, default="{}")
    sha: Mapped[str] = mapped_column(String(64), default="")              # the file the preview was made from
    status: Mapped[str] = mapped_column(String(12), default="ready")      # ready | saved | rejected | failed
    out: Mapped[str] = mapped_column(Text, default="{}")
    created_at: Mapped[datetime] = mapped_column(DateTime(timezone=True), default=now)
    decided_at: Mapped[datetime | None] = mapped_column(DateTime(timezone=True), nullable=True)


# ---------------------------------------------------------------- helpers
KINDS = {".xlsx": "excel", ".xlsm": "excel", ".xls": "excel", ".csv": "csv", ".tsv": "csv", ".docx": "word",
         ".doc": "word", ".docm": "word", ".odt": "word", ".rtf": "word", ".pdf": "pdf", ".pptx": "slides",
         ".jpg": "image", ".jpeg": "image", ".png": "image", ".webp": "image", ".txt": "text", ".md": "text",
         ".html": "text", ".htm": "text", ".json": "text", ".xml": "text"}
SHOWN = {"excel": "Excel", "csv": "CSV", "word": "Word", "pdf": "PDF", "image": "photo", "text": "text", "slides": "PPT",
         "other": "file"}
ORDINAL = {"pehli": 1, "pehla": 1, "pahli": 1, "pahla": 1, "first": 1, "1st": 1, "ek": 1, "doosri": 2, "dusri": 2,
           "doosra": 2, "dusra": 2, "second": 2, "2nd": 2, "teesri": 3, "tisri": 3, "teesra": 3, "third": 3, "3rd": 3,
           "chauthi": 4, "fourth": 4, "4th": 4, "paanchvi": 5, "fifth": 5}
LAST_WORDS = ("last", "aakhri", "akhri", "nayi", "naya", "latest", "abhi wali", "abhi wala", "new")


def kind_of(name: str) -> str:
    return KINDS.get(Path(name or "").suffix.lower(), "other")


def _sha(b: bytes) -> str:
    return hashlib.sha256(b).hexdigest()


def _lst(s: str) -> list:
    try:
        v = json.loads(s or "[]")
        return v if isinstance(v, list) else []
    except Exception:
        return []


def _obj(s: str) -> dict:
    try:
        v = json.loads(s or "{}")
        return v if isinstance(v, dict) else {}
    except Exception:
        return {}


def _since() -> datetime:
    return datetime.now(timezone.utc) - timedelta(days=KEEP_DAYS)


def job_files(tenant: str, conversation_id: str = "") -> list[DeskFile]:
    """This chat's files of the last 7 days, oldest first (her 'pehli wali' = number 1)."""
    with session() as s:
        q = select(DeskFile).where(DeskFile.tenant_id == tenant, DeskFile.added_at >= _since())
        if conversation_id:
            q = q.where(DeskFile.conversation_id == conversation_id)
        rows = s.scalars(q.order_by(DeskFile.added_at, DeskFile.id)).all()
        for r in rows:
            s.expunge(r)
        return list(rows)


def _view(df: DeskFile, n: int | None = None) -> dict[str, Any]:
    where = f"PC: {df.pc_path}" if df.pc_path else ("link" if df.source == "link" else "chat file")
    return {"n": n, "id": df.id, "name": df.name, "kind": SHOWN.get(df.kind, df.kind), "where": where,
            "parts": df.parts, "chars": df.chars, "notes": _lst(df.notes),
            "added": as_utc(df.added_at).isoformat(timespec="minutes") if df.added_at else ""}


def diary_lines(s, tenant: str, conversation_id: str, since: datetime) -> list[str]:
    q = select(DeskFile).where(DeskFile.tenant_id == tenant, DeskFile.added_at >= since)
    if conversation_id:
        q = q.where(DeskFile.conversation_id == conversation_id)
    rows = s.scalars(q.order_by(DeskFile.added_at, DeskFile.id)).all()
    out = []
    from .yaad import _when
    for i, r in enumerate(rows, start=1):
        where = f" · PC: {r.pc_path}" if r.pc_path else ""
        out.append(f"{i}. {r.name} ({SHOWN.get(r.kind, r.kind)}, {r.parts} part){where} · {_when(r.added_at, tenant)} · ref {r.id}")
    edits = s.scalars(select(DeskEdit).where(DeskEdit.tenant_id == tenant, DeskEdit.created_at >= since)
                      .order_by(DeskEdit.created_at.desc()).limit(8)).all()
    for e in edits:
        what = "Sheet Milaan" if e.mode == "sheet" else "badlav"
        st = {"ready": "uski Haan ka intezaar", "saved": "PC par nayi copy save", "rejected": "usne mana kiya",
              "failed": "nahi ho paya"}.get(e.status, e.status)
        o = _obj(e.out)
        out.append(f"✎ {what} {e.id} · {e.name} → {st}" + (f" ({o.get('saved_path') or o.get('name')})" if o else ""))
    return out


# ---------------------------------------------------------------- getting the bytes
async def _from_pc(tenant: str, path: str) -> tuple[str, bytes, str]:
    from .pc.hub import hub, PcOffline, PcError
    from .pc import service as pcs
    try:
        out = await hub.call(tenant, "send_file", {"path": path})
    except PcOffline:
        raise DeskError("PC offline hai: PC on karke Santosh Desktop kholiye (ya file chat mein bhej dijiye).")
    except PcError as e:
        raise DeskError(f"PC se file nahi aayi: {e}")
    f = pcs.get_file(out.get("file_id", ""))
    if f is None or f.tenant_id != tenant:
        raise DeskError("PC se file nahi aayi (dobara try kijiye).")
    return f.id, Path(f.stored_path).read_bytes(), f.name


def _from_file(tenant: str, file_id: str) -> tuple[str, bytes]:
    from .mailbox import service as ms
    try:
        f = ms._file_row(tenant, file_id)
    except Exception as e:
        raise DeskError(str(e))
    p = Path(f.stored_path)
    if not p.exists():
        raise DeskError("Yeh file server se hat chuki hai (24 ghante baad files hat jaati hain); dobara bhejiye.")
    return f.name, p.read_bytes()


def _link_url(url: str) -> tuple[str, str]:
    """A shared Google Sheet / Doc / Drive / Dropbox link -> its official download link, plus a default name."""
    u = url.strip()
    m = re.search(r"docs\.google\.com/spreadsheets/d/([\w-]+)", u)
    if m:
        return f"https://docs.google.com/spreadsheets/d/{m.group(1)}/export?format=xlsx", "Google sheet.xlsx"
    m = re.search(r"docs\.google\.com/document/d/([\w-]+)", u)
    if m:
        return f"https://docs.google.com/document/d/{m.group(1)}/export?format=docx", "Google doc.docx"
    m = re.search(r"drive\.google\.com/(?:file/d/|open\?id=|uc\?(?:export=download&)?id=)([\w-]+)", u)
    if m:
        return f"https://drive.google.com/uc?export=download&id={m.group(1)}", "Drive file"
    if "dropbox.com" in u:
        return re.sub(r"([?&])dl=0", r"\1dl=1", u) if "dl=0" in u else u + ("&" if "?" in u else "?") + "dl=1", "Dropbox file"
    return u, Path(re.sub(r"[?#].*$", "", u)).name or "download"


async def download(url: str) -> tuple[str, bytes]:
    """A public file link (her shared Google Sheet etc.). A login page = she must open it on her PC instead."""
    import httpx
    from urllib.parse import urljoin
    from . import web
    link, name = _link_url(url)
    try:
        link = web._check_url(link)
    except Exception as e:
        raise DeskError(str(e))
    body = b""
    disp = ""
    ctype = ""
    async with httpx.AsyncClient(timeout=40, follow_redirects=False, transport=web._transport,
                                 headers={"User-Agent": web.UA}) as c:
        for _ in range(8):
            async with c.stream("GET", link) as r:
                if r.status_code in (301, 302, 303, 307, 308) and r.headers.get("location"):
                    link = web._check_url(urljoin(link, r.headers["location"]))
                    continue
                if r.status_code >= 400:
                    raise DeskError(f"Link ne error {r.status_code} diya: shayad file sirf login karke khulti hai")
                ctype = r.headers.get("content-type", "")
                disp = r.headers.get("content-disposition", "")
                async for chunk in r.aiter_bytes():
                    body += chunk
                    if len(body) > MAX_BYTES:
                        raise DeskError("File 40 MB se badi hai")
                break
        else:
            raise DeskError("Link bahut baar ghooma (redirect)")
    if "text/html" in ctype and not name.lower().endswith((".html", ".htm")):
        raise DeskError("LOGIN_NEEDED: yeh link file nahi, ek web page / login page deta hai. Iske liye uske PC ke Chrome "
                        "mein link kholkar download karna hoga (pc_do), phir doc_desk action=recent se file uthaiye.")
    m = re.search(r"filename\*?=(?:UTF-8'')?\"?([^\";]+)", disp)
    if m:
        from urllib.parse import unquote
        name = unquote(m.group(1)).strip() or name
    if not Path(name).suffix:
        ext = {"spreadsheetml": ".xlsx", "wordprocessingml": ".docx", "pdf": ".pdf", "csv": ".csv"}
        name += next((v for k, v in ext.items() if k in ctype), "")
    STATS["links"] += 1
    return name[:200], body


# ---------------------------------------------------------------- reading
def extract(name: str, data: bytes, tenant: str = "") -> tuple[str, list[str]]:
    """The whole text of a document (tables kept as tables). Scans / photos go through OCR."""
    k = kind_of(name)
    notes: list[str] = []
    text = ""
    if k == "pdf":
        try:
            from pypdf import PdfReader
            rd = PdfReader(io.BytesIO(data))
            if rd.is_encrypted:
                raise DeskError("PDF par password laga hai: Hoola password wali files nahi kholta")
            pages = [f"--- page {i + 1} ---\n{(p.extract_text() or '').strip()}" for i, p in enumerate(rd.pages[:400])]
            text = "\n\n".join(pages)
            if len(re.sub(r"--- page \d+ ---|\s", "", text)) < 40 * max(1, len(pages)):
                from . import ocr
                r = ocr.read_pdf(data, _lang(tenant), max_pages=30)
                if r and r.get("text"):
                    text = r["text"]
                    notes.append("scan thi: OCR se padha (galti ho sakti hai)")
                else:
                    notes.append("scan wali PDF; OCR is server par nahi hai, isliye text kam mila")
        except DeskError:
            raise
        except Exception as e:
            notes.append(f"PDF theek se nahi padhi ({type(e).__name__})")
    elif k == "image":
        from . import ocr
        r = ocr.read(data, _lang(tenant))
        text = (r or {}).get("text", "")
        notes.append("photo: OCR se padha" if text else "photo mein padhne layak text nahi mila (OCR nahi hai ya photo saaf nahi)")
    elif k in ("text",):
        for enc in ("utf-8-sig", "utf-16", "cp1252"):
            try:
                text = data.decode(enc)
                break
            except Exception:
                continue
    else:
        from . import docread
        try:
            r = docread.read_text(name, "", data, limit=MAX_TEXT, max_pages=400)
            text = r.get("text", "")
        except Exception as e:
            notes.append(f"file theek se nahi padhi ({type(e).__name__})")
        if k == "excel" and not text:
            try:
                from . import sheetmilaan
                t = sheetmilaan.load(name, data)
                text = " | ".join(t.headers) + "\n" + "\n".join(" | ".join(sheetmilaan.shown(v) for v in r) for _, r in t.rows)
            except Exception:
                pass
    if k in ("excel", "word"):
        from . import sheetmilaan
        notes += sheetmilaan.risky_parts(name, data) if k == "excel" else docx_risks(name, data)
    text = (text or "").strip()
    if len(text) > MAX_TEXT:
        notes.append(f"bahut badi file: pehle {MAX_TEXT // 1000}k akshar rakhe")
        text = text[:MAX_TEXT]
    return text, notes


def _lang(tenant: str) -> str:
    try:
        from .agent.file_tools import _owner_lang
        return _owner_lang(tenant)
    except Exception:
        return ""


def docx_risks(name: str, data: bytes) -> list[str]:
    import zipfile
    if not name.lower().endswith((".docx", ".docm")):
        return []
    out = []
    try:
        z = zipfile.ZipFile(io.BytesIO(data))
        names = z.namelist()
        body = z.read("word/document.xml").decode("utf-8", "ignore")
    except Exception:
        return out
    if "<w:ins " in body or "<w:del " in body:
        out.append("is Word file mein tracked changes hain: pehle Accept / Reject kar lein to badlav saaf rahega")
    if "txbxContent" in body:
        out.append("text box ke andar ka likha Hoola nahi badal sakta (woh waisa hi rahega)")
    if any(n.startswith("word/charts/") for n in names) or any(n.startswith("word/diagrams/") for n in names):
        out.append("chart / SmartArt hai: unke andar ka text nahi badlega (baaki file theek rahegi)")
    if name.lower().endswith(".docm") or any("vbaProject" in n for n in names):
        out.append("macro wali Word file: naye copy mein macro nahi rahega")
    if "fldSimple" in body or "w:fldChar" in body:
        out.append("auto fields (page number, date, table of contents) hain: Word mein F9 se update kijiye")
    return out


def _parts(text: str) -> list[str]:
    if not text:
        return [""]
    out, i = [], 0
    while i < len(text):
        j = min(len(text), i + PART)
        if j < len(text):
            k = text.rfind("\n", i + PART // 2, j)
            j = k if k > 0 else j
        out.append(text[i:j])
        i = j
    return out


async def _resolve_new(tenant: str, conversation_id: str, ref: str) -> DeskFile:
    """Anything she can point at -> a desk file (added to this chat's list)."""
    r = (ref or "").strip()
    if not r:
        raise DeskError("Kaunsi file? Naam, PC ka rasta (path), file_id ya link dijiye.")
    if re.match(r"^https?://", r, re.I):
        name, data = await download(r)
        return await asyncio.to_thread(_store, tenant, conversation_id, name, data, "link", url=r)
    if r.startswith("file_") or r.startswith("file-"):
        name, data = await asyncio.to_thread(_from_file, tenant, r)
        return await asyncio.to_thread(_store, tenant, conversation_id, name, data, "file", file_id=r)
    if "\\" in r or "/" in r or re.match(r"^[A-Za-z]:", r):
        fid, data, name = await _from_pc(tenant, r)
        return await asyncio.to_thread(_store, tenant, conversation_id, name, data, "pc", pc_path=r, file_id=fid)
    hits = await pc_find(tenant, r)
    if not hits:
        raise DeskError(f"'{r}' naam ki file PC par nahi mili. Poora naam / folder batayein ya file chat mein bhejiye.")
    best = hits[0]
    if len(hits) > 1 and hits[1].get("_score", 0) >= best.get("_score", 0) and hits[1]["name"] != best["name"]:
        raise DeskError("CHOOSE: kai files mili, kaunsi? " + "; ".join(f"{h['name']} ({h.get('modified', '')}) {h['path']}"
                                                                      for h in hits[:5]))
    fid, data, name = await _from_pc(tenant, best["path"])
    return await asyncio.to_thread(_store, tenant, conversation_id, name, data, "pc", pc_path=best["path"], file_id=fid)


async def pc_find(tenant: str, query: str, limit: int = 8) -> list[dict[str, Any]]:
    from .pc.hub import hub, PcOffline, PcError
    try:
        out = await hub.call(tenant, "search", {"query": query, "kind": "any", "limit": 20})
    except PcOffline:
        raise DeskError("PC offline hai: PC on karke Santosh Desktop kholiye.")
    except PcError as e:
        raise DeskError(f"PC par dhoondh nahi paya: {e}")
    q = re.sub(r"[^a-z0-9ऀ-ॿ]+", " ", query.lower()).split()
    res = []
    for h in out.get("results") or []:
        n = (h.get("name") or "").lower()
        h["_score"] = sum(2 for w in q if w in n) + (3 if n.rsplit(".", 1)[0] == " ".join(q) else 0)
        res.append(h)
    res.sort(key=lambda h: (h["_score"], h.get("modified", "")), reverse=True)
    return res[:limit]


async def recent_on_pc(tenant: str, kind: str = "", limit: int = 10) -> list[dict[str, Any]]:
    """Newest files in her shared folders (e.g. what she just downloaded from WhatsApp)."""
    from .pc.hub import hub, PcOffline, PcError
    exts = {"excel": (".xlsx", ".xlsm", ".xls", ".csv"), "word": (".docx", ".doc"), "pdf": (".pdf",)}.get(kind, ())
    try:
        out = await hub.call(tenant, "recent", {"limit": 30, "days": 3})
        items = out.get("files") or []
    except PcOffline:
        raise DeskError("PC offline hai: PC on karke Santosh Desktop kholiye.")
    except PcError:        # an older Santosh Desktop: look at the top of each shared folder (Downloads etc.)
        items = []
        try:
            roots = (await hub.call(tenant, "list", {})).get("folders") or []
            for r in roots[:6]:
                lst = await hub.call(tenant, "list", {"path": r["path"]})
                items += [e for e in lst.get("entries") or [] if not e.get("dir")]
        except Exception as e:
            raise DeskError(f"PC ki nayi files nahi dekh paya: {e}")
        items.sort(key=lambda e: e.get("modified", ""), reverse=True)
    if exts:
        items = [e for e in items if str(e.get("name", "")).lower().endswith(exts)]
    return items[:limit]


def _store(tenant: str, conversation_id: str, name: str, data: bytes, source: str, pc_path: str = "",
           file_id: str = "", url: str = "") -> DeskFile:
    if len(data) > MAX_BYTES:
        raise DeskError("File 40 MB se badi hai")
    from . import filecheck
    rt = filecheck.real_type(name, data)
    if rt.get("danger"):
        raise DeskError(f"Yeh file khatarnak lag rahi hai: {rt.get('note')}. Hoola ne nahi kholi.")
    if not file_id:
        from .mailbox import service as ms
        file_id = ms.store_bytes(tenant, name, data, "to_phone").id
    text, notes = extract(name, data, tenant)
    parts = _parts(text)
    sha = _sha(data)
    with session() as s:
        old = s.scalars(select(DeskFile).where(DeskFile.tenant_id == tenant, DeskFile.conversation_id == conversation_id,
                                               DeskFile.added_at >= _since(),
                                               (DeskFile.pc_path == pc_path) if pc_path else (DeskFile.sha == sha))).first()
        row = old or DeskFile(tenant_id=tenant, conversation_id=conversation_id)
        if old is None:
            s.add(row)
        row.name, row.source, row.pc_path, row.file_id, row.url = name[:255], source, pc_path[:2000], file_id, url[:2000]
        row.kind, row.size, row.sha, row.chars, row.parts = kind_of(name), len(data), sha, len(text), len(parts)
        row.text, row.notes, row.used_at = text, json.dumps(notes, ensure_ascii=False), now()
        s.flush()
        s.expunge(row)
    STATS["read"] += 1
    try:
        from . import yaad
        yaad.add_item(tenant, "file", f"File {name} ({SHOWN.get(row.kind)}){' PC: ' + pc_path if pc_path else ''}: "
                      f"{text[:500]}", conversation_id, row.id)
    except Exception:
        log.exception("index file failed")
    return row


def find(tenant: str, conversation_id: str, ref: str) -> DeskFile | None:
    """A file already on this chat's list: by number, 'pehli / doosri / aakhri wali', id, or name."""
    r = (ref or "").strip().lower()
    files = job_files(tenant, conversation_id) or job_files(tenant)
    if not files or not r:
        return None
    raw = (ref or "").strip()
    if re.match(r"^https?://", raw, re.I):
        return next((f for f in reversed(files) if f.url == raw), None)
    if "\\" in raw or "/" in raw or re.match(r"^[A-Za-z]:", raw):
        return next((f for f in reversed(files) if f.pc_path == raw), None)
    if r.startswith("df_") or r.startswith("df-"):
        return next((f for f in files if f.id.lower() == r), None) or _by_id(tenant, ref.strip())
    m = re.fullmatch(r"#?(\d{1,2})", r)
    if m:
        i = int(m.group(1))
        return files[i - 1] if 1 <= i <= len(files) else None
    for w, i in ORDINAL.items():
        if re.search(rf"\b{w}\b", r):
            return files[i - 1] if i <= len(files) else None
    if any(w in r for w in LAST_WORDS):
        return files[-1]
    if r.startswith("file_"):
        return next((f for f in reversed(files) if f.file_id == ref.strip()), None)
    words = [w for w in re.findall(r"[a-z0-9ऀ-ॿ]+", r) if w not in ("wali", "wala", "file", "sheet", "ki", "ka", "ke")]
    best, bs = None, 0
    for f in files:
        n = f.name.lower()
        sc = sum(1 for w in words if w in n)
        if sc > bs or (sc == bs and sc and best is not None):
            best, bs = f, sc
    return best if bs else None


def _by_id(tenant: str, fid: str) -> DeskFile | None:
    with session() as s:
        f = s.get(DeskFile, fid)
        if f is None or f.tenant_id != tenant:
            return None
        s.expunge(f)
        return f


async def get(tenant: str, conversation_id: str, ref: str) -> DeskFile:
    f = find(tenant, conversation_id, ref)
    return f if f is not None else await _resolve_new(tenant, conversation_id, ref)


async def data_of(tenant: str, df: DeskFile, fresh: bool = False) -> bytes:
    """The file's bytes now: from the PC again when asked (or when the server copy expired)."""
    if df.pc_path and fresh:
        fid, data, _ = await _from_pc(tenant, df.pc_path)
        with session() as s:
            row = s.get(DeskFile, df.id)
            if row:
                row.file_id = fid
        return data
    try:
        _, data = await asyncio.to_thread(_from_file, tenant, df.file_id)
        return data
    except DeskError:
        if df.pc_path:
            return await data_of(tenant, df, fresh=True)
        if df.url:
            _, data = await download(df.url)
            return data
        raise


async def read(tenant: str, conversation_id: str, ref: str, part: int = 1) -> dict[str, Any]:
    df = await get(tenant, conversation_id, ref)
    parts = _parts(df.text)
    part = max(1, min(int(part or 1), len(parts)))
    files = job_files(tenant, conversation_id)
    n = next((i for i, f in enumerate(files, start=1) if f.id == df.id), None)
    out = {**_view(df, n), "part": part, "text": parts[part - 1]}
    if part < len(parts):
        out["next"] = f"Aage padhne ke liye part={part + 1} (kul {len(parts)})"
    return out


def search(tenant: str, conversation_id: str, query: str, refs: list[str] | None = None, top: int = 8) -> list[dict]:
    files = job_files(tenant, conversation_id) or job_files(tenant)
    if refs:
        pick = [find(tenant, conversation_id, r) for r in refs]
        files = [f for f in pick if f is not None] or files
    chunks: list[tuple[str, str]] = []
    meta: dict[str, tuple[DeskFile, int]] = {}
    for i, f in enumerate(files, start=1):
        for j, para in enumerate(re.split(r"\n\s*\n|(?<=\n)(?=\S)", f.text or "")):
            para = para.strip()
            if len(para) < 3:
                continue
            for k in range(0, len(para), 900):
                cid = f"{f.id}:{j}:{k}"
                chunks.append((cid, para[k:k + 900]))
                meta[cid] = (f, i)
            if len(chunks) > 20000:
                break
    if not chunks:
        return []
    terms = [w.lower() for w in re.findall(r"[\w@.:-]{2,}", query)]
    exact = [(c, t) for c, t in chunks if terms and all(w in t.lower() for w in terms)]
    from . import semantic
    ranked = [(c, t, 1.0) for c, t in exact[:top]]
    if len(ranked) < top:
        have = {c for c, _, _ in ranked}
        ranked += [x for x in semantic.rank(query, [(c, t) for c, t in chunks if c not in have], top=top - len(ranked),
                                            floor=0.35 if semantic.available() else 0.3)]
    out = []
    for cid, text, sc in ranked:
        f, n = meta[cid]
        pos = (f.text or "").find(text[:60])
        part = pos // PART + 1 if pos >= 0 else 1
        out.append({"n": n, "file": f.name, "id": f.id, "part": part, "score": round(float(sc), 2), "text": text[:600]})
    return out


# ---------------------------------------------------------------- edits
def _docx_paras(doc):
    seen: set[int] = set()

    def table(t):
        for row in t.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    yield p
                for t2 in cell.tables:
                    yield from table(t2)

    def walk(container):
        for p in container.paragraphs:
            yield p
        for t in container.tables:
            yield from table(t)

    for p in walk(doc):
        if id(p._p) not in seen:
            seen.add(id(p._p))
            yield p, "body"
    for sec in doc.sections:
        for hf, where in ((sec.header, "header"), (sec.footer, "footer"), (sec.first_page_header, "header"),
                          (sec.first_page_footer, "footer"), (sec.even_page_header, "header"), (sec.even_page_footer, "footer")):
            try:
                if hf.is_linked_to_previous:
                    continue
                for p in walk(hf):
                    if id(p._p) not in seen:
                        seen.add(id(p._p))
                        yield p, where
            except Exception:
                continue


def _replace_in_par(p, find: str, repl: str, limit: int) -> int:
    """Replace inside the runs of one paragraph; the first run of each match keeps its formatting."""
    n = 0
    start = 0
    while n < limit:
        runs = p.runs
        texts = [r.text for r in runs]
        full = "".join(texts)
        i = full.find(find, start)
        if i < 0 or not find:
            break
        j = i + len(find)
        pos = 0
        first = last = None
        bounds = []
        for k, t in enumerate(texts):
            bounds.append((pos, pos + len(t)))
            if first is None and pos <= i < pos + len(t):
                first = k
            if pos < j <= pos + len(t):
                last = k
            pos += len(t)
        if first is None or last is None:
            break
        b0 = bounds[first][0]
        if first == last:
            runs[first].text = texts[first][:i - b0] + repl + texts[first][j - b0:]
        else:
            runs[first].text = texts[first][:i - b0] + repl
            for k in range(first + 1, last):
                runs[k].text = ""
            bl = bounds[last][0]
            runs[last].text = texts[last][j - bl:]
        n += 1
        start = i + len(repl)
    return n


def _insert_after(p, text: str):
    """A new paragraph after p, with p's style and its first run's formatting."""
    from copy import deepcopy
    new = deepcopy(p._p)
    for child in list(new):
        if child.tag.endswith("}r") or child.tag.endswith("}hyperlink"):
            new.remove(child)
    p._p.addnext(new)
    from docx.text.paragraph import Paragraph
    np_ = Paragraph(new, p._parent)
    run = np_.add_run(text)
    if p.runs:
        src = p.runs[0]._r.rPr
        if src is not None:
            run._r.insert(0, deepcopy(src))
    return np_


def edit_docx(data: bytes, ops: list[dict]) -> tuple[bytes, list[dict], list[str]]:
    import docx
    doc = docx.Document(io.BytesIO(data))
    changes, missing = [], []
    for op in ops:
        if op.get("after") is not None and op.get("insert") is not None:
            target = str(op["after"])
            hit = next((p for p, _ in _docx_paras(doc) if target and target in p.text), None)
            if hit is None:
                missing.append(target[:120])
                continue
            _insert_after(hit, str(op["insert"]))
            changes.append({"where": "naya paragraph", "before": hit.text[:160], "after": str(op["insert"])[:300]})
            continue
        f, r = str(op.get("find") or ""), str(op.get("replace") if op.get("replace") is not None else "")
        if not f:
            continue
        limit = 10_000 if op.get("all", True) else 1
        total = 0
        for p, where in _docx_paras(doc):
            if f not in "".join(x.text for x in p.runs):
                continue
            before = p.text
            k = _replace_in_par(p, f, r, limit - total)
            if k:
                total += k
                if len(changes) < 200:
                    changes.append({"where": where, "before": before[:300], "after": p.text[:300], "count": k})
            if total >= limit:
                break
        if not total:
            missing.append(f[:120])
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue(), changes, missing


def edit_xlsx(name: str, data: bytes, ops: list[dict]) -> tuple[bytes, list[dict], list[str]]:
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(data), keep_vba=name.lower().endswith(".xlsm"))
    changes, missing = [], []
    for op in ops:
        if op.get("cell"):
            sh = op.get("sheet") or wb.sheetnames[0]
            if sh not in wb.sheetnames:
                missing.append(f"sheet {sh}")
                continue
            c = wb[sh][str(op["cell"]).upper()]
            if isinstance(c.value, str) and c.value.startswith("=") and not op.get("formula_ok"):
                missing.append(f"{sh}!{c.coordinate} mein formula hai (nahi badla)")
                continue
            old = c.value
            from .sheetmilaan import _typed
            c.value = _typed(op.get("value"), old) if op.get("value") is not None else None
            changes.append({"where": f"{sh}!{c.coordinate}", "before": "" if old is None else str(old), "after": str(c.value or "")})
            continue
        f, r = str(op.get("find") or ""), str(op.get("replace") if op.get("replace") is not None else "")
        if not f:
            continue
        n = 0
        for ws in wb.worksheets:
            for row in ws.iter_rows():
                for c in row:
                    if isinstance(c.value, str) and f in c.value and not c.value.startswith("="):
                        old = c.value
                        c.value = old.replace(f, r)
                        n += 1
                        if len(changes) < 200:
                            changes.append({"where": f"{ws.title}!{c.coordinate}", "before": old[:200], "after": c.value[:200]})
        if not n:
            missing.append(f[:120])
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue(), changes, missing


def edit_text(text: str, ops: list[dict]) -> tuple[str, list[dict], list[str]]:
    changes, missing = [], []
    for op in ops:
        f, r = str(op.get("find") or ""), str(op.get("replace") if op.get("replace") is not None else "")
        if op.get("after") is not None and op.get("insert") is not None:
            a = str(op["after"])
            i = text.find(a)
            if i < 0:
                missing.append(a[:120])
                continue
            e = text.find("\n", i + len(a))
            e = len(text) if e < 0 else e
            text = text[:e] + "\n" + str(op["insert"]) + text[e:]
            changes.append({"where": "nayi line", "before": a[:160], "after": str(op["insert"])[:300]})
            continue
        if not f:
            continue
        k = text.count(f)
        if not k:
            missing.append(f[:120])
            continue
        i = text.find(f)
        ctx = text[max(0, i - 80): i + len(f) + 80]
        text = text.replace(f, r) if op.get("all", True) else text.replace(f, r, 1)
        changes.append({"where": "text", "before": ctx[:300], "after": ctx.replace(f, r)[:300], "count": k})
    return text, changes, missing


def _out_name(name: str, ext: str | None = None) -> str:
    p = Path(name)
    stem = re.sub(r"\s*\(Hoola[^)]*\)$", "", p.stem)
    return f"{stem} (Hoola {datetime.now(timezone(timedelta(hours=5, minutes=30))).strftime('%d-%m-%Y %H%M')}){ext or p.suffix}"


def build(df: DeskFile, data: bytes, ops: list[dict], tenant: str = "") -> tuple[bytes, str, list[dict], list[str], list[str]]:
    """(new bytes, new name, changes, not found, warnings)."""
    k = df.kind
    low = df.name.lower()
    warnings: list[str] = []
    if k == "word" and low.endswith((".docx", ".docm")):
        b, ch, miss = edit_docx(data, ops)
        warnings = docx_risks(df.name, data)
        return b, _out_name(df.name, ".docx"), ch, miss, warnings
    if k == "excel" and low.endswith((".xlsx", ".xlsm")):
        from .sheetmilaan import risky_parts
        b, ch, miss = edit_xlsx(df.name, data, ops)
        return b, _out_name(df.name), ch, miss, risky_parts(df.name, data)
    if k in ("text", "csv"):
        try:
            text = data.decode("utf-8-sig")
        except Exception:
            text = data.decode("cp1252", "ignore")
        t, ch, miss = edit_text(text, ops)
        return t.encode("utf-8-sig" if k == "csv" else "utf-8"), _out_name(df.name), ch, miss, []
    if k == "pdf":
        t, ch, miss = edit_text(df.text or "", ops)
        t = re.sub(r"--- page \d+ ---\n?", "\n", t).strip()
        pdf = _text_pdf(t, tenant)
        warnings.append("PDF ka design (logo, lines, columns) naye copy mein saada ho jayega: sirf sahi kiya hua text rahega. "
                        "Asli Word / Excel file ho to woh dijiye, usme formatting ke saath badlunga.")
        if pdf is None:
            from .officegen import make_docx
            warnings.append("PDF banane wala Chromium server par nahi mila: sahi kiya hua text Word file mein diya")
            return make_docx(Path(df.name).stem, t), _out_name(df.name, ".docx"), ch, miss, warnings
        return pdf, _out_name(df.name, ".pdf"), ch, miss, warnings
    if k == "word":
        raise DeskError("Purani .doc / .odt file: Word mein kholkar .docx mein 'Save As' kar dijiye, phir badlunga")
    raise DeskError(f"{SHOWN.get(k, k)} file mein badlav Hoola nahi kar sakta (Word, Excel, PDF, CSV, text kar sakta hai)")


def _text_pdf(text: str, tenant: str = "") -> bytes | None:
    try:
        from .lekhak import service as lk
        lang = "hi" if re.search(r"[ऀ-ॿ]", text) else "en"
        return lk.make_pdf(lk.letter_html("typing", lang, {"text": text}))
    except Exception as e:
        log.info("text pdf failed: %s", e)
        return None


def _card(e: DeskEdit) -> dict[str, Any]:
    res = _obj(e.result)
    view = {"id": e.id, "mode": e.mode, "name": e.name, "kind": e.kind, "status": e.status, "out": _obj(e.out)}
    if e.mode == "sheet":
        lim = 40
        view.update({"key": res.get("key"), "counts": res.get("counts"), "compare": res.get("compare"),
                     "added": res.get("added", [])[:lim], "changed": res.get("changed", [])[:lim],
                     "missing": res.get("missing", [])[:lim], "dupes": res.get("dupes"), "new_columns": res.get("new_columns"),
                     "notes": res.get("notes", []), "warnings": res.get("warnings", []), "other": res.get("other_name", ""),
                     "options": _obj(e.ops), "more": {k: max(0, len(res.get(k, [])) - lim) for k in ("added", "changed", "missing")},
                     "say": res.get("say", "")})
    else:
        view.update({"changes": res.get("changes", [])[:60], "missing": res.get("missing", []),
                     "warnings": res.get("warnings", []), "new_name": res.get("new_name", ""), "say": res.get("say", ""),
                     "more": max(0, len(res.get("changes", [])) - 60)})
    return view


def _publish(conversation_id: str, e: DeskEdit) -> None:
    bus.publish(conversation_id or e.conversation_id, "sheet_milaan" if e.mode == "sheet" else "doc_edit", _card(e))


async def propose_edit(tenant: str, conversation_id: str, ref: str, ops: list[dict]) -> dict[str, Any]:
    if not ops:
        raise DeskError("Kya badalna hai? edits = [{find, replace}] / [{cell, value}] / [{after, insert}]")
    if len(ops) > 200:
        raise DeskError("Ek baar mein 200 se zyada badlav nahi")
    df = await get(tenant, conversation_id, ref)
    data = await data_of(tenant, df, fresh=bool(df.pc_path))
    b, new_name, ch, miss, warn = await asyncio.to_thread(build, df, data, ops, tenant)
    if not ch:
        raise DeskError("Kuch nahi badla: yeh shabd file mein nahi mile: " + "; ".join(miss[:8]))
    say = f"{df.name}: {len(ch)} jagah badlav taiyaar" + (f", {len(miss)} nahi mile" if miss else "") + \
        ". Aapki Haan par nayi copy PC ke Santosh Inbox mein save hogi aur khulegi; asli file waisi hi rahegi."
    with session() as s:
        e = DeskEdit(tenant_id=tenant, conversation_id=conversation_id, mode="doc", file_ref=df.id, name=df.name, kind=df.kind,
                     ops=json.dumps(ops, ensure_ascii=False), sha=_sha(data),
                     result=json.dumps({"changes": ch, "missing": miss, "warnings": warn, "new_name": new_name, "say": say},
                                       ensure_ascii=False))
        s.add(e)
        s.flush()
        s.expunge(e)
    STATS["edits"] += 1
    _publish(conversation_id, e)
    return {"status": "waiting_owner", "edit_id": e.id, "changes": len(ch), "not_found": miss, "warnings": warn,
            "say": say + " Card par Haan dabaiye.", "preview": ch[:6]}


async def sheet_compare(tenant: str, conversation_id: str, old_ref: str, new_ref: str, key: str = "", update=None,
                        sheet_old: str = "", sheet_new: str = "", key_new: str = "", add: bool = True,
                        missing: str = "keep") -> dict[str, Any]:
    from . import sheetmilaan as sm
    old = await get(tenant, conversation_id, old_ref)
    new = await get(tenant, conversation_id, new_ref)
    if old.id == new.id:
        raise DeskError("Dono ek hi file hain: purani aur nayi alag files batayein")
    od = await data_of(tenant, old, fresh=bool(old.pc_path))
    nd = await data_of(tenant, new)
    upd = [u for u in (update if isinstance(update, list) else [update] if update else []) if u]

    def work():
        ot = sm.load(old.name, od, sheet_old or None, upd + ([key] if key else []))
        nt = sm.load(new.name, nd, sheet_new or None, upd + ([key_new or key] if (key_new or key) else []))
        res = sm.compare(ot, nt, key or None, upd or None, key_new or None)
        res["warnings"] = sm.risky_parts(old.name, od)
        res["say"] = sm.say(res)
        res["other_name"] = new.name
        return res
    try:
        res = await asyncio.to_thread(work)
    except sm.SheetError as e:
        raise DeskError(str(e))
    opts = {"add": bool(add), "update": True, "missing": missing if missing in ("keep", "mark", "delete") else "keep",
            "highlight": False, "cols": upd}
    with session() as s:
        e = DeskEdit(tenant_id=tenant, conversation_id=conversation_id, mode="sheet", file_ref=old.id, other_ref=new.id,
                     name=old.name, kind=old.kind, ops=json.dumps(opts), sha=_sha(od),
                     result=json.dumps(res, ensure_ascii=False, default=str))
        s.add(e)
        s.flush()
        s.expunge(e)
    STATS["sheet"] += 1
    _publish(conversation_id, e)
    c = res["counts"]
    return {"status": "waiting_owner", "edit_id": e.id, "key": res["key"], "counts": c,
            "changed": res["changed"][:8], "added": res["added"][:5], "missing": res["missing"][:5],
            "dupes": res["dupes"], "notes": res["notes"], "warnings": res["warnings"],
            "say": res["say"] + f" Purani file '{old.name}' ki nayi copy banegi (asli waisi hi rahegi). Card par Haan dabaiye."}


def get_edit(tenant: str, edit_id: str) -> DeskEdit:
    with session() as s:
        e = s.get(DeskEdit, edit_id)
        if e is None or e.tenant_id != tenant:
            raise DeskError("Yeh badlav nahi mila")
        s.expunge(e)
        return e


async def save_to_pc(tenant: str, name: str, data: bytes, open_it: bool = True) -> dict[str, Any]:
    """The new copy -> PC's Santosh Inbox (+ opened there). PC offline -> the file comes to her phone instead."""
    from .mailbox import service as ms
    from .pc.hub import hub, PcOffline, PcError
    from .pc.service import public_file
    row = ms.store_bytes(tenant, name, data, "to_pc")
    view = public_file(row)
    view.pop("pc_path", None)
    try:
        got = await hub.call(tenant, "save_file", {"file_id": row.id, "name": row.name, "url": f"/api/pc/fetch/{row.id}"})
    except (PcOffline, PcError) as e:
        return {**view, "where": "phone", "saved_path": "", "opened": False,
                "say": f"PC tak nahi pahuncha ({'offline' if isinstance(e, PcOffline) else e}); nayi copy yahan phone par hai."}
    out = {**view, "where": "pc", "saved_path": got.get("saved_path", ""), "opened": False}
    if open_it and out["saved_path"]:
        try:
            await hub.call(tenant, "open_file", {"path": out["saved_path"]})
            out["opened"] = True
        except Exception as e:      # an older Santosh Desktop has no open_file: she opens it from the Inbox
            out["open_error"] = str(e)[:160]
    out["say"] = (f"Nayi copy PC ke Santosh Inbox mein save: {out['saved_path']}" +
                  (" aur khol di." if out["opened"] else " (PC par Santosh Inbox folder mein kholiye; khud kholne ke liye "
                                                       "naya Santosh Desktop chahiye)."))
    STATS["saved"] += 1
    return out


async def decide(tenant: str, edit_id: str, yes: bool, options: dict | None = None) -> dict[str, Any]:
    """Only from her tap (the API): Haan -> build the NEW copy from the file as it is NOW, save to the PC, open it."""
    e = get_edit(tenant, edit_id)
    if e.status != "ready":
        return _card(e)
    if not yes:
        with session() as s:
            row = s.get(DeskEdit, edit_id)
            row.status, row.decided_at = "rejected", now()
        e = get_edit(tenant, edit_id)
        _publish(e.conversation_id, e)
        _chat_note(e, f"Aapne mana kiya: {e.name} mein kuch nahi badla.")
        return _card(e)
    df = _by_id(tenant, e.file_ref)
    if df is None:
        raise DeskError("Asli file ki jaankari nahi mili; dobara shuru kijiye")
    data = await data_of(tenant, df, fresh=bool(df.pc_path))
    note = ""
    if _sha(data) != e.sha:
        note = "File preview ke baad PC par badli thi: naye roop par wahi badlav lagaye. "
    try:
        if e.mode == "sheet":
            from . import sheetmilaan as sm
            other = _by_id(tenant, e.other_ref)
            nd = await data_of(tenant, other) if other else b""
            opts = {**_obj(e.ops), **{k: v for k, v in (options or {}).items() if k in ("add", "update", "missing", "highlight")}}
            res = _obj(e.result)
            if note:      # her file changed: compare again so row numbers are right
                ot = sm.load(df.name, data, res["sheets"]["old"] or None)
                nt = sm.load(other.name, nd, res["sheets"]["new"] or None)
                res = sm.compare(ot, nt, res["key"]["old"], opts.get("cols") or None, res["key"]["new"])
            b, done = await asyncio.to_thread(sm.apply, df.name, data, other.name, nd, res, bool(opts.get("add", True)),
                                              bool(opts.get("update", True)), str(opts.get("missing", "keep")),
                                              bool(opts.get("highlight", False)))
            new_name = _out_name(df.name)
            summary = (f"{done['updated_cells']} cells update, {done['added_rows']} naye rows" +
                       (f", {done['missing_marked']} rows par 'naye mein nahi' likha" if done["missing_marked"] else "") +
                       (f", {done['missing_deleted']} rows hataye" if done["missing_deleted"] else "") +
                       (f"; formula wale {len(done['skipped_formula'])} cells nahi chhede" if done["skipped_formula"] else ""))
        else:
            b, new_name, ch, miss, warn = await asyncio.to_thread(build, df, data, _lst(e.ops), tenant)
            done = {"changes": len(ch), "missing": miss}
            summary = f"{len(ch)} jagah badla" + (f", {len(miss)} nahi mile" if miss else "")
    except Exception as ex:
        with session() as s:
            row = s.get(DeskEdit, edit_id)
            row.status, row.decided_at = "failed", now()
            row.out = json.dumps({"error": str(ex)[:300]}, ensure_ascii=False)
        e = get_edit(tenant, edit_id)
        _publish(e.conversation_id, e)
        _chat_note(e, f"{e.name} ki nayi copy nahi ban payi: {str(ex)[:200]}")
        return _card(e)
    saved = await save_to_pc(tenant, new_name, b)
    saved["summary"] = note + summary
    with session() as s:
        row = s.get(DeskEdit, edit_id)
        row.status, row.decided_at = "saved", now()
        row.out = json.dumps({**saved, "done": done}, ensure_ascii=False, default=str)
    try:   # the new copy joins this chat's file list (so 'nayi wali' / 'updated wali' works next)
        nd_ = await asyncio.to_thread(_store, tenant, e.conversation_id, new_name, b, "made", file_id=saved["file_id"])
        if saved.get("saved_path"):
            with session() as s:
                r = s.get(DeskFile, nd_.id)
                if r:
                    r.pc_path = saved["saved_path"]
    except Exception:
        log.exception("could not list the new copy")
    e = get_edit(tenant, edit_id)
    _publish(e.conversation_id, e)
    _chat_note(e, f"{e.name}: {saved['summary']}. {saved['say']}")
    try:
        from . import yaad
        yaad.add_item(tenant, "decision", f"{'Sheet Milaan' if e.mode == 'sheet' else 'Document badlav'} {e.name} -> "
                      f"{new_name}: {saved['summary']} ({saved.get('saved_path') or 'phone'})", e.conversation_id, e.id)
    except Exception:
        log.exception("index decision failed")
    return _card(e)


def _chat_note(e: DeskEdit, text: str) -> None:
    """A line in the chat (shown + in history) so Hoola and she both know what happened."""
    try:
        from .agent.santosh import _save
        _save(e.conversation_id, "assistant", [{"type": "text", "text": text}], display=text)
        bus.publish(e.conversation_id, "message", {"role": "assistant", "text": text, "source": "docdesk"})
    except Exception:
        log.exception("doc desk chat note failed")


def edits(tenant: str, limit: int = 30) -> list[dict[str, Any]]:
    with session() as s:
        rows = s.scalars(select(DeskEdit).where(DeskEdit.tenant_id == tenant).order_by(DeskEdit.created_at.desc())
                         .limit(limit)).all()
        return [_card(r) for r in rows]


def full_result(tenant: str, edit_id: str) -> dict[str, Any]:
    e = get_edit(tenant, edit_id)
    return {**_card(e), "result": _obj(e.result)}
