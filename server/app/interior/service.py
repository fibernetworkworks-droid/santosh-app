"""Interior Sahayak: Hoola as an interior designer's coworker, end to end.

  1. Client brief per room (style, budget, must-haves, avoid) -> site measurement sheet (L x W x H + doors / windows) -> floor,
     ceiling, net wall, skirting, perimeter worked out by code from HER numbers only.
  2. Concept / moodboard: a palette from the style (a design suggestion) or her brand kit, her own photos, and pictures from her
     words made by photogen (always labelled "AI se bana"), or a 3D render through the 3D Studio when it is on the server.
  3. Room-wise BOQ with HER rate card (modular kitchen, wardrobes, false ceiling, electrical, painting, flooring, furniture ...):
     a rate she has not given = a blank, a quantity not measured / not said = a blank. Never an invented rate, size or brand price.
     Quotation PDF with GST (CGST + SGST inside her state, IGST outside), design fee by her fee mode, payment stages, terms.
  4. Vendor / carpenter work orders, installation schedule (vendor clashes, trade order, handover date), snag list with photos
     (EXIF / GPS stripped), warranty sheet in HER words at handover.

Waits for HER tap (only the app API sets by_tap=True; tools only make the card):
  * the quotation going to the client (`interior_quote` card -> final PDF + message draft),
  * any vendor order (`interior_order`; her fingerprint once a phone key is active: it commits her money),
  * any change order that changes the price (`interior_change`),
  * any payment request to the client (`interior_payment` -> message draft with her UPI link; Hoola never collects).
Hoola never sends a message by itself: every message is a draft (wa.me / SMS) she sends."""
from __future__ import annotations

import csv
import difflib
import io
import json
import logging
import re
import secrets
import threading
from datetime import date, datetime, timedelta, timezone
from pathlib import Path
from typing import Any
from urllib.parse import quote as urlq

from sqlalchemy import select

from ..db import now, session
from . import rules as R
from .models import (
    InteriorAlert,
    InteriorChange,
    InteriorOrder,
    InteriorProject,
    InteriorQuote,
    InteriorRate,
    InteriorSettings,
    InteriorSnag,
    InteriorStage,
    InteriorTask,
    InteriorVendor,
)

log = logging.getLogger("santosh.interior")
_LOCK = threading.RLock()
DATA_ROOT = Path("data") / "interior"
STATUS = {"brief": "Brief", "measure": "Naap", "concept": "Concept", "quote": "Quote", "approved": "Pakka", "production": "Production",
          "install": "Installation", "handover": "Handover", "done": "Poora", "cancelled": "Radd"}
TRADES = ("carpenter", "modular", "electrician", "plumber", "painter", "ceiling", "flooring", "fabricator", "mason", "supplier",
          "cleaner", "other")
MAX_PHOTOS = 12


class InteriorError(Exception):
    pass


# ---------------------------------------------------------------- helpers
def _j(s: str | None, default):
    try:
        v = json.loads(s or "")
        return v if isinstance(v, type(default)) else default
    except (TypeError, ValueError):
        return default


def _dump(v) -> str:
    return json.dumps(v, ensure_ascii=False, default=str)


def _aware(d: datetime | None) -> datetime | None:
    if d is None:
        return None
    return d if d.tzinfo else d.replace(tzinfo=timezone.utc)


def clock() -> datetime:
    return now()


def _tz(tenant: str):
    from ..life import tz_of
    return tz_of(tenant)


def today(tenant: str) -> date:
    return clock().astimezone(_tz(tenant)).date()


def _date(v) -> date | None:
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    if not v:
        return None
    s = str(v).strip()[:10]
    for f in ("%Y-%m-%d", "%d-%m-%Y", "%d/%m/%Y", "%d.%m.%Y", "%d-%m-%y", "%d/%m/%y"):
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            pass
    return None


def _iso(d: date | None) -> str:
    return d.isoformat() if d else ""


def _num(v) -> float | None:
    """A number from her words, or None when she did not give one (0 is a real number only when typed)."""
    if v is None or v == "":
        return None
    if isinstance(v, bool):
        return None
    if isinstance(v, (int, float)):
        return float(v)
    s = re.sub(r"[^\d.\-]", "", str(v).replace(",", ""))
    try:
        return float(s) if s not in ("", "-", ".") else None
    except ValueError:
        return None


def _rs(v: float | None) -> str:
    if v is None:
        return "[daam?]"
    try:
        from ..india.money import inr
        return inr(v)
    except Exception:
        return f"₹{v:,.0f}"


def _pdf_rs(v: float | None) -> str:
    if v is None:
        return "[ ? ]"
    return "Rs. " + f"{v:,.2f}"


def _publish(tenant: str, kind: str, payload: dict, conv: str = "") -> None:
    from ..events import bus
    try:
        if not conv:
            from ..net.service import home_conv
            conv = home_conv(tenant)
        bus.publish(conv, kind, payload)
    except Exception:
        log.exception("interior publish failed")


def is_owner(actor: str) -> bool:
    a = actor or "owner"
    return a == "owner" or a.startswith("approved:")


def actor_name(tenant: str, actor: str) -> str:
    if is_owner(actor):
        return "Owner"
    try:
        from ..employee.service import actor_label
        return actor_label(tenant, actor)
    except Exception:
        return str(actor)[:40]


def _st(s, tenant: str) -> InteriorSettings:
    r = s.get(InteriorSettings, tenant)
    if r is None:
        r = InteriorSettings(tenant_id=tenant)
        s.add(r)
        s.flush()
    return r


def _safe(x: str) -> str:
    return re.sub(r"[^\w-]", "_", x or "")[:64] or "x"


def _dir(tenant: str, project_id: str = "") -> Path:
    p = DATA_ROOT / _safe(tenant)
    if project_id:
        p = p / _safe(project_id)
    p.mkdir(parents=True, exist_ok=True)
    return p


def _load_file(tenant: str, file_id: str) -> tuple[str, bytes]:
    from ..mailbox import service as ms
    try:
        name, _m, data = ms.load_files(tenant, [{"type": "file", "file_id": file_id}])[0]
        return name, data
    except Exception as e:
        raise InteriorError(f"File {file_id} nahi mili: {e}")


def _store(tenant: str, name: str, data: bytes, mime: str, conv: str = "", title: str = "") -> dict:
    from .. import chatfiles
    f = chatfiles.store(tenant, name, data, mime)
    _publish(tenant, "pc_file", {**f, "source": "interior", "title": title or name}, conv)
    return f


def _clean_image(data: bytes) -> bytes:
    """Re-encode from pixels: EXIF (camera, time, GPS location) and hidden chunks are dropped; orientation applied first."""
    from PIL import Image, ImageOps
    try:
        img = Image.open(io.BytesIO(data))
        img.load()
    except Exception:
        raise InteriorError("Yeh photo khul nahi rahi.")
    try:
        img = ImageOps.exif_transpose(img)
    except Exception:
        pass
    img.thumbnail((2000, 2000))
    src = img.convert("RGB")
    out = io.BytesIO()
    Image.frombytes("RGB", src.size, src.tobytes()).save(out, "JPEG", quality=86)
    return out.getvalue()


def _wa(phone: str, text: str) -> str:
    d = R.phone(phone)
    return f"https://wa.me/91{d}?text={urlq(text)}" if d else ""


def _sms(phone: str, text: str) -> str:
    d = R.phone(phone)
    return f"sms:+91{d}?body={urlq(text)}" if d else ""


def _upi(tenant: str, amount: float, note: str) -> str:
    with session() as s:
        st = _st(s, tenant)
        upi, nm = st.upi_id, st.firm
    if not upi:
        try:
            from ..employee import service as emp
            p = emp.profile(tenant) or {}
            upi = p.get("upi_id") or ""
            nm = nm or p.get("upi_name") or p.get("business_name") or ""
        except Exception:
            upi = ""
    if not upi or amount <= 0:
        return ""
    return f"upi://pay?pa={urlq(upi)}&pn={urlq(nm)}&am={amount:.2f}&cu=INR&tn={urlq(note[:40])}"


# ---------------------------------------------------------------- settings
_SET_KEYS = ("firm", "title", "coa_reg", "gstin", "state_code", "address", "phone", "email", "upi_id", "fee_mode", "fee_value",
             "gst_works", "gst_design", "stages", "terms", "validity_days", "warranty", "measure_unit", "morning_time", "rules")


def _settings_view(r: InteriorSettings) -> dict:
    stages = _j(r.stages, [])
    return {"firm": r.firm, "title": r.title, "coa_reg": r.coa_reg, "gstin": r.gstin, "state_code": r.state_code,
            "address": r.address, "phone": r.phone, "email": r.email, "upi_id": r.upi_id, "fee_mode": r.fee_mode,
            "fee_value": r.fee_value or None, "gst_works": r.gst_works, "gst_design": r.gst_design,
            "stages": stages or R.DEFAULT_STAGES, "stages_starter": not stages, "terms": r.terms, "validity_days": r.validity_days,
            "warranty": _j(r.warranty, {}), "measure_unit": r.measure_unit, "morning_time": r.morning_time, "rules": r.rules,
            "kinds": {k: v["label"] for k, v in R.KINDS.items()}}


def settings(tenant: str) -> dict:
    with session() as s:
        return _settings_view(_st(s, tenant))


def save_settings(tenant: str, **kw: Any) -> dict:
    with _LOCK, session() as s:
        r = _st(s, tenant)
        title = str(kw.get("title", r.title) or "Interior Designer")
        coa = str(kw.get("coa_reg", r.coa_reg) or "")
        why = R.title_problem(title, coa) or R.title_problem(str(kw.get("firm", r.firm) or ""), coa)
        if why:
            raise InteriorError(why)
        for k in _SET_KEYS:
            if k not in kw or kw[k] is None:
                continue
            v = kw[k]
            if k == "gstin":
                g = str(v).strip().upper()
                if g:
                    from ..india.money import gstin_ok
                    if not gstin_ok(g):
                        raise InteriorError("GSTIN sahi nahi lagta (check digit / format).")
                    r.state_code = g[:2]
                r.gstin = g
            elif k == "stages":
                st = [{"label": R.clean(x.get("label"), 120), "pct": float(_num(x.get("pct")) or 0)} for x in (v or []) if isinstance(x, dict)]
                errs = R.check_stages(st)
                if errs:
                    raise InteriorError(" ".join(errs))
                r.stages = _dump(st)
            elif k == "warranty":
                if not isinstance(v, dict):
                    raise InteriorError("warranty = {kind: aapke shabd}.")
                cur = _j(r.warranty, {})
                for kk, vv in v.items():
                    kk = kk if kk in R.KINDS else R.kind_of(kk)
                    if vv in (None, ""):
                        cur.pop(kk, None)
                    else:
                        cur[kk] = R.clean(vv, 200)
                r.warranty = _dump(cur)
            elif k == "fee_mode":
                if v not in ("none", "percent", "sqft", "flat"):
                    raise InteriorError("fee_mode = none | percent | sqft | flat")
                r.fee_mode = v
            elif k in ("fee_value", "gst_works", "gst_design"):
                n = _num(v)
                if n is None or n < 0 or (k != "fee_value" and n > 40):
                    raise InteriorError(f"{k} sahi number nahi.")
                setattr(r, k, n)
            elif k == "validity_days":
                r.validity_days = max(1, min(int(_num(v) or 15), 120))
            elif k == "measure_unit":
                r.measure_unit = v if v in ("ft", "m", "mm", "in") else "ft"
            elif k == "phone":
                r.phone = R.phone(v) or str(v)[:20]
            elif k == "state_code":
                r.state_code = str(v).strip()[:2]
            else:
                setattr(r, k, R.clean(v, 2000) if k in ("terms", "rules", "address") else R.clean(v, 160))
        r.title = title
        view = _settings_view(r)
    return {"settings": view, "say": "Niyam save ho gaye."}


# ---------------------------------------------------------------- rate card (HER numbers)
def _rate_view(x: InteriorRate) -> dict:
    return {"id": x.id, "code": x.code, "kind": x.kind, "kind_label": R.KINDS.get(x.kind, R.KINDS["other"])["label"], "name": x.name,
            "unit": x.unit, "rate": x.rate, "rate_text": _rs(x.rate), "spec": x.spec, "hsn_sac": x.hsn_sac, "gst_pct": x.gst_pct,
            "active": x.active}


def rate_save(tenant: str, name: str = "", rate=None, unit: str = "", kind: str = "", code: str = "", spec: str = "",
              hsn_sac: str = "", gst_pct=None, rate_id: str = "") -> dict:
    with _LOCK, session() as s:
        x = None
        if rate_id:
            x = s.get(InteriorRate, rate_id)
            if x is None or x.tenant_id != tenant:
                raise InteriorError("Rate nahi mila.")
        elif code:
            x = s.scalars(select(InteriorRate).where(InteriorRate.tenant_id == tenant, InteriorRate.code == code.strip())).first()
        if x is None and name:
            x = s.scalars(select(InteriorRate).where(InteriorRate.tenant_id == tenant, InteriorRate.name == R.clean(name, 200))).first()
        if x is None:
            if not name:
                raise InteriorError("Rate ka naam chahiye (jaise 'Wardrobe BWP ply + laminate').")
            x = InteriorRate(tenant_id=tenant, name=R.clean(name, 200))
            s.add(x)
        n = _num(rate)
        if n is not None:
            if n < 0 or n > 10_000_000:
                raise InteriorError("Rate sahi nahi lagta.")
            x.rate = n
        elif not rate_id and not x.rate:
            raise InteriorError("Rate (aapka, per unit, GST se pehle) chahiye; Hoola rate nahi gadhta.")
        if name:
            x.name = R.clean(name, 200)
        x.kind = kind if kind in R.KINDS else (R.kind_of(kind or x.name) if (kind or not x.kind or x.kind == "other") else x.kind)
        u = R.unit_of(unit, "")
        if u:
            x.unit = u
        elif not x.unit or x.unit == "nos":
            x.unit = R.KINDS[x.kind]["unit"]
        if code:
            x.code = code.strip()[:30]
        if spec:
            x.spec = R.clean(spec, 500)
        if hsn_sac:
            x.hsn_sac = re.sub(r"\D", "", hsn_sac)[:8]
        g = _num(gst_pct)
        if g is not None:
            x.gst_pct = g
        x.active = True
        s.flush()
        v = _rate_view(x)
    return {"rate": v, "say": f"Rate card: {v['name']} = {v['rate_text']} / {v['unit']}."}


def rate_remove(tenant: str, rate_id: str) -> dict:
    with session() as s:
        x = s.get(InteriorRate, rate_id)
        if x is None or x.tenant_id != tenant:
            raise InteriorError("Rate nahi mila.")
        x.active = False
    return {"say": "Rate card se hataya."}


def rates(tenant: str, q: str = "") -> list[dict]:
    with session() as s:
        rows = s.scalars(select(InteriorRate).where(InteriorRate.tenant_id == tenant, InteriorRate.active.is_(True))).all()
        out = [_rate_view(x) for x in rows]
    if q:
        ql = q.lower()
        out = [r for r in out if ql in r["name"].lower() or ql == r["code"].lower() or ql == r["kind"]]
    return sorted(out, key=lambda r: (R.KINDS.get(r["kind"], R.KINDS["other"])["order"], r["name"]))


def _norm(h) -> str:
    return re.sub(r"[^a-z]", "", str(h or "").lower())


def _rows(name: str, data: bytes) -> list[list]:
    if name.lower().endswith((".xlsx", ".xlsm")):
        import openpyxl
        wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        return [list(r) for r in wb.active.iter_rows(values_only=True)]
    text = data.decode("utf-8-sig", errors="replace")
    return [r for r in csv.reader(io.StringIO(text))]


def import_rates(tenant: str, file_id: str = "", data: bytes | None = None, name: str = "rates.csv") -> dict:
    if file_id:
        name, data = _load_file(tenant, file_id)
    rows = _rows(name, data or b"")
    head_i = -1
    for i, r in enumerate(rows[:15]):
        hs = [_norm(c) for c in r]
        if any(h in ("item", "name", "particulars", "description", "naam", "kaam") for h in hs) and \
                any(h in ("rate", "daam", "price", "rateperunit", "amount") for h in hs):
            head_i = i
            break
    if head_i < 0:
        raise InteriorError("Sheet mein 'Item / Name' aur 'Rate' column nahi mile.")
    hs = [_norm(c) for c in rows[head_i]]

    def col(*names):
        for n in names:
            if n in hs:
                return hs.index(n)
        return None
    ci = {"name": col("item", "name", "particulars", "description", "naam", "kaam"), "rate": col("rate", "daam", "price", "rateperunit", "amount"),
          "unit": col("unit", "uom", "ikai"), "kind": col("kind", "category", "head", "type"), "code": col("code", "itemcode", "sr", "srno"),
          "spec": col("spec", "specification", "material", "brand", "finish"), "gst": col("gst", "gstpct", "gstrate", "tax"),
          "hsn": col("hsn", "sac", "hsnsac")}
    added, bad = 0, []
    for n, r in enumerate(rows[head_i + 1:], start=head_i + 2):
        def g(k):
            i = ci[k]
            return r[i] if i is not None and i < len(r) else None
        nm = str(g("name") or "").strip()
        if not nm:
            continue
        rt = _num(g("rate"))
        if rt is None or rt <= 0:
            bad.append(f"line {n}: {nm[:40]} (rate nahi)")
            continue
        try:
            rate_save(tenant, nm, rt, str(g("unit") or ""), str(g("kind") or ""), str(g("code") or ""), str(g("spec") or ""),
                      str(g("hsn") or ""), g("gst"))
            added += 1
        except InteriorError as e:
            bad.append(f"line {n}: {e}")
    return {"added": added, "bad": bad[:20], "say": f"{added} rate card mein aaye." + (f" {len(bad)} chhode (rate nahi / galat)." if bad else "")}


def _match_rate(rows: list[InteriorRate], code: str, name: str) -> InteriorRate | None:
    if code:
        for x in rows:
            if x.code and x.code.lower() == code.lower():
                return x
    if not name:
        return None
    nl = name.lower().strip()
    for x in rows:
        if x.name.lower() == nl:
            return x
    best, score = None, 0.0
    for x in rows:
        sc = difflib.SequenceMatcher(None, nl, x.name.lower()).ratio()
        if sc > score:
            best, score = x, sc
    return best if score >= 0.82 else None


# ---------------------------------------------------------------- vendors
def _vendor_view(v: InteriorVendor) -> dict:
    return {"id": v.id, "name": v.name, "trade": v.trade, "phone": v.phone, "gstin": v.gstin, "notes": v.notes, "active": v.active}


def vendor_save(tenant: str, name: str, trade: str = "", phone: str = "", gstin: str = "", notes: str = "", vendor_id: str = "") -> dict:
    with _LOCK, session() as s:
        v = s.get(InteriorVendor, vendor_id) if vendor_id else None
        if v is not None and v.tenant_id != tenant:
            v = None
        if v is None and name:
            v = s.scalars(select(InteriorVendor).where(InteriorVendor.tenant_id == tenant, InteriorVendor.name == R.clean(name, 120))).first()
        if v is None:
            if not name:
                raise InteriorError("Vendor ka naam chahiye.")
            v = InteriorVendor(tenant_id=tenant, name=R.clean(name, 120))
            s.add(v)
        if trade:
            t = trade.lower().strip()
            v.trade = t if t in TRADES else R.KINDS.get(R.kind_of(t), R.KINDS["other"])["trade"]
        if phone:
            p = R.phone(phone)
            if not p:
                raise InteriorError("Phone number 10 ankon ka nahi lagta.")
            v.phone = p
        if gstin:
            v.gstin = gstin.strip().upper()[:15]
        if notes:
            v.notes = R.clean(notes, 400)
        v.active = True
        s.flush()
        out = _vendor_view(v)
    return {"vendor": out, "say": f"Vendor {out['name']} ({out['trade']}) save."}


def vendors(tenant: str) -> list[dict]:
    with session() as s:
        return [_vendor_view(v) for v in s.scalars(select(InteriorVendor).where(InteriorVendor.tenant_id == tenant,
                                                                             InteriorVendor.active.is_(True))).all()]


def _vendor(s, tenant: str, ref: str) -> InteriorVendor:
    v = s.get(InteriorVendor, ref) if ref else None
    if v is not None and v.tenant_id == tenant:
        return v
    rows = s.scalars(select(InteriorVendor).where(InteriorVendor.tenant_id == tenant, InteriorVendor.active.is_(True))).all()
    rl = (ref or "").lower().strip()
    hits = [v for v in rows if rl and (rl == v.name.lower() or rl in v.name.lower() or (R.phone(ref) and R.phone(ref) == v.phone))]
    if len(hits) == 1:
        return hits[0]
    if not hits:
        raise InteriorError(f"Vendor '{ref}' nahi mila; pehle vendor jodiye.")
    raise InteriorError("Kai vendor milte hain: " + ", ".join(v.name for v in hits[:5]))


# ---------------------------------------------------------------- projects + brief + measurement
def _proj(s, tenant: str, ref: str) -> InteriorProject:
    p = s.get(InteriorProject, ref) if ref else None
    if p is not None and p.tenant_id == tenant:
        return p
    rows = s.scalars(select(InteriorProject).where(InteriorProject.tenant_id == tenant)).all()
    if not ref:
        live = [p for p in rows if p.status not in ("done", "cancelled")]
        if len(live) == 1:
            return live[0]
        raise InteriorError("Kaun sa project? (number / naam / client)")
    rl = ref.lower().strip()
    hits = [p for p in rows if rl in (p.number.lower(), p.name.lower(), p.client.lower())]
    if not hits:
        hits = [p for p in rows if rl in p.name.lower() or rl in p.client.lower()]
    if len(hits) == 1:
        return hits[0]
    if not hits:
        raise InteriorError(f"Project '{ref}' nahi mila.")
    raise InteriorError("Kai project milte hain: " + ", ".join(f"{p.number} {p.name}" for p in hits[:5]))


def _log(p: InteriorProject, text: str, by: str = "") -> None:
    lg = _j(p.log, [])
    lg.append({"at": now().isoformat(timespec="minutes"), "what": text[:300], "by": by})
    p.log = _dump(lg[-200:])


def _room(rooms: list[dict], ref: str, create: bool = False) -> dict | None:
    rl = (ref or "").lower().strip()
    for r in rooms:
        if rl and (rl == r["id"] or rl == r["name"].lower()):
            return r
    if create and ref:
        r = {"id": f"r{len(rooms) + 1}", "name": R.clean(ref, 60), "type": "", "style": "", "budget": None, "must": [], "avoid": [],
             "notes": "", "dims": {}, "openings": [], "areas": {}, "measured_by": "", "measured_on": ""}
        rooms.append(r)
        return r
    return None


def _project_view(s, p: InteriorProject, full: bool = False) -> dict:
    rooms = _j(p.rooms, [])
    v = {"id": p.id, "number": p.number, "name": p.name, "client": p.client, "phone": p.phone, "email": p.email, "site": p.site,
         "city": p.city, "pin": p.pin, "home": p.home, "carpet_sqft": p.carpet_sqft or None, "budget": p.budget or None,
         "budget_text": _rs(p.budget) if p.budget else "", "style": p.style, "style_label": R.STYLES.get(p.style, {}).get("label", ""),
         "status": p.status, "status_label": STATUS.get(p.status, p.status), "rooms": [{"id": r["id"], "name": r["name"],
         "measured": bool(r.get("areas", {}).get("floor_sqft"))} for r in rooms], "target_handover": p.target_handover,
         "handed_over": p.handed_over, "contract_quote": p.contract_quote, "forgotten": p.forgotten}
    if full:
        v["brief"] = _j(p.brief, {})
        v["rooms"] = rooms
        v["concept"] = _j(p.concept, {})
        v["boq"] = boq_view_rows(s, p)
        v["log"] = _j(p.log, [])[-30:]
        v["client_state"], v["client_gstin"] = p.client_state, p.client_gstin
        v["contract_total"] = _contract_total(s, p)
    return v


_PROJ_KEYS = ("client", "phone", "email", "site", "city", "pin", "client_state", "client_gstin", "home", "carpet_sqft", "budget",
              "style", "target_handover", "status")


def project_save(tenant: str, name: str = "", project: str = "", actor: str = "owner", conv: str = "", **kw: Any) -> dict:
    with _LOCK, session() as s:
        if project:
            p = _proj(s, tenant, project)
            new = False
        else:
            if not (name or kw.get("client")):
                raise InteriorError("Project ka naam ya client ka naam chahiye.")
            st = _st(s, tenant)
            st.proj_seq += 1
            p = InteriorProject(tenant_id=tenant, number=f"IP-{st.proj_seq:04d}", conv=conv or "")
            s.add(p)
            new = True
        if name:
            p.name = R.clean(name, 160)
        for k in _PROJ_KEYS:
            if k not in kw or kw[k] is None:
                continue
            v = kw[k]
            if k == "phone":
                ph = R.phone(v)
                if v and not ph:
                    raise InteriorError("Client ka phone 10 ankon ka nahi lagta.")
                p.phone = ph
            elif k == "pin":
                p.pin = re.sub(r"\D", "", str(v))[:6]
            elif k in ("carpet_sqft", "budget"):
                n = _num(v)
                setattr(p, k, n if n and n > 0 else 0.0)
            elif k == "style":
                p.style = v if v in R.STYLES else R.style_of(str(v))
            elif k == "target_handover":
                d = _date(v)
                if v and not d:
                    raise InteriorError("Handover tareekh samajh nahi aayi (DD-MM-YYYY).")
                p.target_handover = _iso(d)
            elif k == "status":
                if v not in STATUS:
                    raise InteriorError("status = " + " | ".join(STATUS))
                if v in ("done", "handover") and not is_owner(actor):
                    raise InteriorError("Handover / poora sirf owner.")
                p.status = v
            elif k == "client_gstin":
                g = str(v).strip().upper()
                p.client_gstin = g
                if g[:2].isdigit():
                    p.client_state = g[:2]
            elif k == "client_state":
                p.client_state = str(v).strip()[:2]
            else:
                setattr(p, k, R.clean(v, 400))
        if not p.name:
            p.name = f"{p.client or 'Client'} {p.home}".strip()
        _log(p, "naya project" if new else "project badla", actor_name(tenant, actor))
        s.flush()
        view = _project_view(s, p, full=True)
    return {"project": view, "say": f"{view['number']} {view['name']} " + ("bana." if new else "save.")}


def brief(tenant: str, project: str, room: str = "", style: str = "", budget=None, must: Any = None, avoid: Any = None,
          notes: str = "", room_type: str = "", actor: str = "owner") -> dict:
    """Client's brief, per room when a room is named. Client words are data; secrets dropped."""
    def lst(v) -> list[str]:
        if v is None:
            return []
        if isinstance(v, str):
            v = [x for x in re.split(r"[,;\n]", v)]
        return [R.clean(x, 120) for x in v if str(x).strip()][:20]
    with _LOCK, session() as s:
        p = _proj(s, tenant, project)
        if room:
            rooms = _j(p.rooms, [])
            r = _room(rooms, room, create=True)
            if style:
                r["style"] = style if style in R.STYLES else R.style_of(style)
            b = _num(budget)
            if b is not None:
                r["budget"] = b
            if room_type:
                r["type"] = R.clean(room_type, 40)
            r["must"] = sorted(set(r.get("must", []) + lst(must)))
            r["avoid"] = sorted(set(r.get("avoid", []) + lst(avoid)))
            if notes:
                r["notes"] = R.clean(notes, 800)
            p.rooms = _dump(rooms)
            said = f"{r['name']} ka brief save"
        else:
            bf = _j(p.brief, {})
            bf["must"] = sorted(set(bf.get("must", []) + lst(must)))
            bf["avoid"] = sorted(set(bf.get("avoid", []) + lst(avoid)))
            if notes:
                bf["notes"] = R.clean(notes, 1500)
            p.brief = _dump(bf)
            if style:
                p.style = style if style in R.STYLES else R.style_of(style)
            b = _num(budget)
            if b is not None and b > 0:
                p.budget = b
            said = "Project ka brief save"
        if p.status == "brief" and room:
            p.status = "measure"
        _log(p, said, actor_name(tenant, actor))
        view = _project_view(s, p, full=True)
    return {"project": view, "say": said + "."}


def measure(tenant: str, project: str, room: str, l=None, w=None, h=None, unit: str = "", openings: list | None = None,
            notes: str = "", by: str = "", actor: str = "owner") -> dict:
    """Site measurement from HER / her staff's numbers. Areas by code; a missing number stays missing."""
    if not room:
        raise InteriorError("Kaun sa kamra?")
    with _LOCK, session() as s:
        st = _st(s, tenant)
        u = unit or st.measure_unit or "ft"
        p = _proj(s, tenant, project)
        rooms = _j(p.rooms, [])
        r = _room(rooms, room, create=True)
        dims = dict(r.get("dims") or {})
        for k, v in (("l", l), ("w", w), ("h", h)):
            if v not in (None, ""):
                ft = R.length(v, u)
                if ft is None:
                    raise InteriorError(f"'{v}' naap samajh nahi aaya.")
                if ft > 200:
                    raise InteriorError(f"{v} {u} bahut bada lagta hai; ikai (ft / m / mm) dekhiye.")
                dims[k] = ft
        ops = []
        for o in openings or []:
            if not isinstance(o, dict):
                continue
            ow, oh = R.length(o.get("w"), u), R.length(o.get("h"), u)
            kind = str(o.get("kind") or "other").lower()
            kind = "door" if kind.startswith(("door", "darwaza", "darwaja")) else ("window" if kind.startswith(("win", "khidki")) else "other")
            ops.append({"kind": kind, "w": ow, "h": oh, "count": int(_num(o.get("count")) or 1)})
        if openings is not None:
            r["openings"] = ops
        r["dims"] = dims
        r["areas"] = R.room_areas(dims.get("l"), dims.get("w"), dims.get("h"), r.get("openings"))
        r["measured_by"] = R.clean(by or actor_name(tenant, actor), 60)
        r["measured_on"] = today(tenant).isoformat()
        if notes:
            r["notes"] = (r.get("notes", "") + " | " + R.clean(notes, 300)).strip(" |")
        p.rooms = _dump(rooms)
        if p.status in ("brief", "measure"):
            p.status = "concept" if all(x.get("areas", {}).get("floor_sqft") for x in rooms) else "measure"
        _log(p, f"naap: {r['name']}", r["measured_by"])
        ar = r["areas"]
    say = f"{r['name']}: " + (f"floor {ar['floor_sqft']} sq ft" if ar["floor_sqft"] is not None else "floor area abhi nahi")
    if ar["wall_sqft"] is not None:
        say += f", deewar (darwaze-khidki ghata kar) {ar['wall_sqft']} sq ft"
    if ar["missing"]:
        say += ". Baaki naap: " + ", ".join(ar["missing"])
    return {"room": r, "say": say + "."}


def measure_sheet(tenant: str, project: str, conv: str = "") -> dict:
    import openpyxl
    with session() as s:
        p = _proj(s, tenant, project)
        rooms = _j(p.rooms, [])
        num, name = p.number, p.name
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Naap"
    ws.append([f"{num} {name} - site measurement (ft)"])
    ws.append(["Room", "L (ft)", "W (ft)", "H (ft)", "Openings", "Floor sqft", "Ceiling sqft", "Net wall sqft", "Skirting rft",
               "Missing", "By", "On"])
    for r in rooms:
        d, a = r.get("dims") or {}, r.get("areas") or {}
        ops = "; ".join(f"{o['kind']} {o.get('w')}x{o.get('h')} x{o.get('count')}" for o in r.get("openings") or [])
        ws.append([r["name"], d.get("l"), d.get("w"), d.get("h"), ops, a.get("floor_sqft"), a.get("ceiling_sqft"), a.get("wall_sqft"),
                   a.get("skirting_rft"), ", ".join(a.get("missing") or []), r.get("measured_by"), r.get("measured_on")])
    bio = io.BytesIO()
    wb.save(bio)
    f = _store(tenant, f"{num}_naap.xlsx", bio.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", conv,
               f"{num} naap sheet")
    return {"file": f, "say": "Naap sheet ka Excel."}


# ---------------------------------------------------------------- concept / moodboard
def _make_picture(prompt: str) -> bytes:
    """photogen (Gemini image models, labelled 'AI se bana'). Tests replace this."""
    from .. import photogen
    return photogen.make(prompt, aspect="4:3")["png"]


def concept(tenant: str, project: str, room: str = "", style: str = "", words: str = "", pictures: int = 0,
            file_ids: list | None = None, brand_colors: bool = False, conv: str = "", actor: str = "owner") -> dict:
    with session() as s:
        p = _proj(s, tenant, project)
        rooms = _j(p.rooms, [])
        r = _room(rooms, room) if room else None
        if room and r is None:
            raise InteriorError(f"Kamra '{room}' is project mein nahi.")
        stl = (style if style in R.STYLES else R.style_of(style)) or (r or {}).get("style") or p.style or "modern"
        pnum, home = p.number, p.home
        dims = (r or {}).get("dims") or {}
        must = (r or {}).get("must") or _j(p.brief, {}).get("must", [])
        avoid = (r or {}).get("avoid") or _j(p.brief, {}).get("avoid", [])
    sd = R.STYLES.get(stl, R.STYLES["modern"])
    palette = list(sd["palette"])
    source = f"{sd['label']} style ka sujhav"
    if brand_colors:
        try:
            from ..brand import service as bs
            k = bs.kit(tenant)
            cols = k.get("colors") or {}
            if cols:
                palette = [cols[x] for x in ("primary", "secondary", "accent", "dark", "light") if cols.get(x)]
                source = "aapke brand kit ke rang"
        except Exception:
            pass
    items: list[dict] = []
    for fid in (file_ids or [])[:MAX_PHOTOS]:
        items.append({"file_id": str(fid), "ai": False, "room": (r or {}).get("name", ""), "note": "aapki photo"})
    notes: list[str] = []
    n = max(0, min(int(pictures or 0), 3))
    if n:
        what = (r or {}).get("type") or (r or {}).get("name") or home or "living room"
        size = f"{dims.get('l')} x {dims.get('w')} ft" if dims.get("l") and dims.get("w") else ""
        base = (f"Interior design concept render of an Indian home {what}" + (f", about {size}" if size else "")
                + f", {sd['label']} style, colour palette {', '.join(palette)}, materials {', '.join(sd['materials'])}"
                + (f", must have: {', '.join(must)}" if must else "") + (f", avoid: {', '.join(avoid)}" if avoid else "")
                + (f". {R.clean(words, 300)}" if words else "")
                + ". Photorealistic, natural light, no people, no text, no logos, no brand names.")
        for i in range(n):
            try:
                png = _make_picture(base + ("" if i == 0 else f" Variation {i + 1}: different angle."))
                f = _store(tenant, f"{pnum}_concept_{i + 1}.png", png, "image/png", conv, f"{pnum} concept (AI se bana)")
                items.append({"file_id": f.get("file_id", ""), "ai": True, "room": (r or {}).get("name", ""), "note": "AI se bana (sujhav)",
                              "prompt": base[:400]})
            except Exception as e:
                notes.append(f"Tasveer {i + 1} nahi bani: {str(e)[:120]}")
                break
    with _LOCK, session() as s:
        p = _proj(s, tenant, project)
        c = _j(p.concept, {})
        key = (r or {}).get("id", "all")
        boards = c.get("boards", {})
        b = boards.get(key, {"items": []})
        b.update({"style": stl, "palette": palette, "palette_from": source, "materials": sd["materials"]})
        b["items"] = (b.get("items", []) + items)[-24:]
        boards[key] = b
        c["boards"] = boards
        p.concept = _dump(c)
        if p.status in ("brief", "measure"):
            p.status = "concept"
        _log(p, f"concept {key}", actor_name(tenant, actor))
    card = {"project": pnum, "room": (r or {}).get("name", "Poora ghar"), "style": sd["label"], "palette": palette,
            "palette_from": source, "materials": sd["materials"], "items": b["items"], "notes": notes,
            "ai_note": "AI se bani tasveerein sirf sujhav hain; asli material / rang site par dekh kar pakka kijiye."}
    _publish(tenant, "interior_concept", {"concept": card, "say": f"{pnum} moodboard: {card['room']}."}, conv)
    return {"concept": card, "say": f"Moodboard taiyaar ({source})." + (" " + " ".join(notes) if notes else "")}


def render3d(tenant: str, project: str, room: str, words: str = "", conv: str = "", actor: str = "owner") -> dict:
    """A 3D render of the room through the 3D Studio (v0.141) when it is on this server; its card arrives by itself."""
    with session() as s:
        p = _proj(s, tenant, project)
        r = _room(_j(p.rooms, []), room)
        if r is None:
            raise InteriorError(f"Kamra '{room}' nahi mila.")
        stl = R.STYLES.get(r.get("style") or p.style or "modern", R.STYLES["modern"])
    d = r.get("dims") or {}
    if not (d.get("l") and d.get("w") and d.get("h")):
        raise InteriorError("3D ke liye kamre ka poora naap (L, W, H) chahiye.")
    try:
        from ..blender3d import service as b3
    except Exception:
        raise InteriorError("3D Studio is server par nahi hai; moodboard ki tasveerein bana sakte hain.")
    brief_text = (f"Interior room {r['name']} {d['l']:.1f} x {d['w']:.1f} ft, height {d['h']:.1f} ft, {stl['label']} style, colours "
                  f"{', '.join(stl['palette'])}, " + ", ".join(r.get("must") or []) + (f". {R.clean(words, 300)}" if words else ""))
    out = b3.start(tenant, conv, brief_text, {"views": ["three_quarter"], "quality": "preview"}, actor)
    return {"render": out, "say": "3D render shuru; card khud aayega (AI se bana)."}


# ---------------------------------------------------------------- BOQ
def _resolve_line(s, tenant: str, p: InteriorProject, ln: dict, rate_rows: list[InteriorRate], gst_default: float) -> dict:
    rooms = _j(p.rooms, [])
    room = str(ln.get("room") or "").strip()
    r = _room(rooms, room) if room else None
    name = R.clean(ln.get("name") or ln.get("item") or "", 200)
    x = _match_rate(rate_rows, str(ln.get("code") or ""), name)
    kind = ln.get("kind") if ln.get("kind") in R.KINDS else (x.kind if x else R.kind_of(name))
    unit = R.unit_of(str(ln.get("unit") or ""), "") or (x.unit if x else R.KINDS[kind]["unit"])
    notes = []
    rate = _num(ln.get("rate"))
    rate_src = "aapka (is line ke liye)" if rate is not None else ""
    if rate is None and x is not None:
        if x.unit == unit:
            rate, rate_src = x.rate, "rate card"
        else:
            notes.append(f"rate card ki ikai {x.unit}, line ki {unit}")
    qty = _num(ln.get("qty"))
    qty_src = "aapka" if qty is not None else ""
    if qty is None:
        area_key = str(ln.get("qty_from") or R.KINDS[kind]["area"] or "")
        if area_key and r is not None:
            qty, nt = R.qty_from_area(area_key, r.get("areas") or {}, unit)
            if qty is not None:
                qty_src = f"naap ({area_key})"
            elif nt:
                notes.append(nt)
        elif area_key and room and r is None:
            notes.append(f"kamra '{room}' nahi mila")
    if x is None and not name:
        raise InteriorError("Har line mein item ka naam ya rate card ka code chahiye.")
    gst = _num(ln.get("gst_pct"))
    if gst is None:
        gst = x.gst_pct if x is not None and x.gst_pct is not None else gst_default
    blanks = []
    if rate is None:
        blanks.append("rate")
    if qty is None:
        blanks.append("qty")
    amount = R.r2(rate * qty) if rate is not None and qty is not None else None
    return {"id": ln.get("id") or "L" + secrets.token_hex(3), "room": (r or {}).get("name", room), "kind": kind,
            "kind_label": R.KINDS[kind]["label"], "name": name or (x.name if x else ""), "spec": R.clean(ln.get("spec") or (x.spec if x else ""), 400),
            "unit": unit, "qty": qty, "qty_src": qty_src, "rate": rate, "rate_src": rate_src, "rate_code": x.code if x else "",
            "rate_id": x.id if x else "", "hsn_sac": (x.hsn_sac if x and x.hsn_sac else R.KINDS[kind]["sac"]), "gst_pct": gst,
            "amount": amount, "blanks": blanks, "notes": notes, "co": str(ln.get("co") or "")}


def boq_view_rows(s, p: InteriorProject) -> list[dict]:
    return _j(p.boq, [])


def boq_set(tenant: str, project: str, lines: list | None, replace: bool = False, actor: str = "owner") -> dict:
    if not lines:
        raise InteriorError("BOQ lines chahiye: [{room, name / code, qty, unit, rate (agar card mein nahi), qty_from}].")
    with _LOCK, session() as s:
        p = _proj(s, tenant, project)
        if p.contract_quote:
            raise InteriorError("Quote pakka ho chuka hai; ab badlav = change order (action=change).")
        st = _st(s, tenant)
        rate_rows = s.scalars(select(InteriorRate).where(InteriorRate.tenant_id == tenant, InteriorRate.active.is_(True))).all()
        cur = [] if replace else _j(p.boq, [])
        for ln in lines:
            if not isinstance(ln, dict):
                continue
            row = _resolve_line(s, tenant, p, ln, rate_rows, st.gst_works)
            idx = next((i for i, c in enumerate(cur) if c["id"] == row["id"]), None)
            if idx is None:
                cur.append(row)
            else:
                cur[idx] = row
        cur.sort(key=lambda c: (c["room"], R.KINDS.get(c["kind"], R.KINDS["other"])["order"]))
        p.boq = _dump(cur)
        if p.status in ("brief", "measure", "concept"):
            p.status = "quote"
        _log(p, f"BOQ: {len(lines)} line", actor_name(tenant, actor))
        view = _boq_summary(cur)
    return {"boq": view, "say": view["say"]}


def boq_remove(tenant: str, project: str, line_id: str) -> dict:
    with _LOCK, session() as s:
        p = _proj(s, tenant, project)
        if p.contract_quote:
            raise InteriorError("Quote pakka hai; hatana = change order.")
        cur = [c for c in _j(p.boq, []) if c["id"] != line_id]
        p.boq = _dump(cur)
        view = _boq_summary(cur)
    return {"boq": view, "say": "Line hatai. " + view["say"]}


def _boq_summary(lines: list[dict]) -> dict:
    known = sum(c["amount"] for c in lines if c["amount"] is not None)
    blanks = [f"{c['room'] or '-'} / {c['name']}: " + " + ".join(c["blanks"]) for c in lines if c["blanks"]]
    by_room: dict[str, float] = {}
    for c in lines:
        if c["amount"] is not None:
            by_room[c["room"] or "Common"] = by_room.get(c["room"] or "Common", 0) + c["amount"]
    say = f"{len(lines)} line, jo pata hai unka jod {_rs(known)} (GST se pehle)."
    if blanks:
        say += f" {len(blanks)} line mein khaali jagah (rate / naap aapse chahiye)."
    return {"lines": lines, "subtotal_known": R.r2(known), "blanks": blanks, "by_room": {k: R.r2(v) for k, v in by_room.items()}, "say": say}


def boq(tenant: str, project: str) -> dict:
    with session() as s:
        return _boq_summary(_j(_proj(s, tenant, project).boq, []))


def boq_file(tenant: str, project: str, conv: str = "") -> dict:
    import openpyxl
    with session() as s:
        p = _proj(s, tenant, project)
        lines, num = _j(p.boq, []), p.number
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "BOQ"
    ws.append(["Room", "Head", "Item", "Spec", "Unit", "Qty", "Rate", "Amount", "GST %", "HSN/SAC", "Khaali"])
    for c in lines:
        ws.append([c["room"], c["kind_label"], c["name"], c["spec"], c["unit"], c["qty"], c["rate"], c["amount"], c["gst_pct"],
                   c["hsn_sac"], " + ".join(c["blanks"])])
    bio = io.BytesIO()
    wb.save(bio)
    f = _store(tenant, f"{num}_BOQ.xlsx", bio.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", conv, f"{num} BOQ")
    return {"file": f, "say": "BOQ ka Excel."}


# ---------------------------------------------------------------- quotation
def _totals(tenant: str, s, p: InteriorProject, lines: list[dict]) -> tuple[dict, list[str]]:
    st = _st(s, tenant)
    intra = not (st.state_code and p.client_state and st.state_code != p.client_state)
    blanks = [f"{c['room'] or '-'} / {c['name']}: " + " + ".join(c["blanks"]) for c in lines if c["blanks"]]
    sub = sum(c["amount"] for c in lines if c["amount"] is not None)
    taxes: dict[float, float] = {}
    for c in lines:
        if c["amount"] is not None:
            taxes[float(c["gst_pct"])] = taxes.get(float(c["gst_pct"]), 0) + c["amount"]
    fee, fee_note = 0.0, ""
    if st.fee_mode == "percent" and st.fee_value:
        fee, fee_note = R.r2(sub * st.fee_value / 100), f"{st.fee_value:g}% of works"
    elif st.fee_mode == "sqft" and st.fee_value:
        if p.carpet_sqft:
            fee, fee_note = R.r2(p.carpet_sqft * st.fee_value), f"{p.carpet_sqft:g} sq ft x Rs {st.fee_value:g}"
        else:
            blanks.append("Design fee: carpet area (sq ft) chahiye")
    elif st.fee_mode == "flat" and st.fee_value:
        fee, fee_note = R.r2(st.fee_value), "flat"
    gst_rows = []
    tax_total = 0.0
    for pct, base in sorted(taxes.items()):
        g = R.gst_split(base, pct, intra)
        gst_rows.append({"pct": pct, "base": R.r2(base), **g})
        tax_total += g["tax"]
    if fee:
        g = R.gst_split(fee, st.gst_design, intra)
        gst_rows.append({"pct": st.gst_design, "base": fee, "on": "design fee (SAC 998391)", **g})
        tax_total += g["tax"]
    total = R.r2(sub + fee + tax_total)
    return ({"works": R.r2(sub), "design_fee": fee, "fee_note": fee_note, "gst": gst_rows, "tax": R.r2(tax_total), "total": total,
             "intra": intra, "words": _words(total)}, blanks)


def _words(v: float) -> str:
    try:
        from ..india.money import rupees_words
        return rupees_words(v, "en")
    except Exception:
        return ""


def _quote_show(q: InteriorQuote, p: InteriorProject) -> dict:
    t = _j(q.totals, {})
    return {"number": q.number, "project": p.number, "client": p.client, "phone": p.phone, "total": t.get("total"),
            "lines": [[c["id"], c["name"], c["qty"], c["rate"], c["amount"], c["gst_pct"]] for c in _j(q.lines, [])],
            "fee": t.get("design_fee"), "valid_until": q.valid_until}


def _quote_view(s, q: InteriorQuote) -> dict:
    p = s.get(InteriorProject, q.project_id)
    t = _j(q.totals, {})
    return {"id": q.id, "number": q.number, "rev": q.rev, "project": p.number if p else "", "project_id": q.project_id,
            "project_name": p.name if p else "", "client": p.client if p else "", "lines": _j(q.lines, []), "totals": t,
            "total_text": _rs(t.get("total")), "blanks": _j(q.blanks, []), "status": q.status, "hash": q.hash,
            "valid_until": q.valid_until, "file_id": q.file_id, "reason": q.reason,
            "created_at": _aware(q.created_at).isoformat(timespec="minutes") if q.created_at else ""}


def quote_make(tenant: str, project: str, conv: str = "", actor: str = "owner") -> dict:
    """A new revision from the current BOQ (draft; nothing goes to the client)."""
    with _LOCK, session() as s:
        p = _proj(s, tenant, project)
        if p.contract_quote:
            raise InteriorError("Quote pehle hi pakka hai; ab change order.")
        lines = _j(p.boq, [])
        if not lines:
            raise InteriorError("Pehle BOQ banaiye (action=boq).")
        st = _st(s, tenant)
        totals, blanks = _totals(tenant, s, p, lines)
        olds = s.scalars(select(InteriorQuote).where(InteriorQuote.project_id == p.id)).all()
        for o in olds:
            if o.status in ("draft", "waiting", "approved"):
                o.status = "superseded"
        rev = len(olds) + 1
        if rev == 1:
            st.quote_seq += 1
            base = f"Q-{st.quote_seq:04d}"
        else:
            base = olds[0].number.split("/")[0]
        q = InteriorQuote(tenant_id=tenant, project_id=p.id, number=f"{base}/R{rev}", rev=rev, lines=_dump(lines), totals=_dump(totals),
                          blanks=_dump(blanks), valid_until=_iso(today(tenant) + timedelta(days=st.validity_days)))
        s.add(q)
        s.flush()
        q.hash = R.digest(_quote_show(q, p))
        view = _quote_view(s, q)
    say = f"{view['number']} draft: {view['total_text']} (GST samet)." + (f" {len(blanks)} khaali jagah bhariye, tabhi client ko jayega." if blanks else
                                                                         " 'Client ko bhejne' ke liye card maangiye.")
    return {"quote": view, "say": say}


def _quote(s, tenant: str, ref: str) -> InteriorQuote:
    q = s.get(InteriorQuote, ref)
    if q is None:
        q = s.scalars(select(InteriorQuote).where(InteriorQuote.tenant_id == tenant, InteriorQuote.number == ref)).first()
    if q is None or q.tenant_id != tenant:
        raise InteriorError("Quote nahi mila.")
    return q


def quote_view(tenant: str, ref: str) -> dict:
    with session() as s:
        return _quote_view(s, _quote(s, tenant, ref))


def quote_request(tenant: str, ref: str, conv: str = "", actor: str = "owner") -> dict:
    """Make the card for HER tap; the quotation goes to the client only after it."""
    with _LOCK, session() as s:
        q = _quote(s, tenant, ref)
        p = s.get(InteriorProject, q.project_id)
        if q.status not in ("draft", "waiting"):
            raise InteriorError(f"{q.number} '{q.status}' hai.")
        blanks = _j(q.blanks, [])
        if blanks:
            raise InteriorError("Pehle khaali jagah bhariye: " + "; ".join(blanks[:6]))
        if not p.client:
            raise InteriorError("Client ka naam chahiye.")
        st = _st(s, tenant)
        why = R.title_problem(st.title, st.coa_reg)
        if why:
            raise InteriorError(why)
        q.status = "waiting"
        q.hash = R.digest(_quote_show(q, p))
        view = _quote_view(s, q)
    say = f"{view['number']} ({view['total_text']}) client {view['client']} ko bhejne se pehle aapki haan ka card bheja hai."
    _publish(tenant, "interior_quote", {"quote": view, "say": say}, conv)
    return {"quote": view, "status": "needs_tap", "say": say}


def decide_quote(tenant: str, ref: str, yes: bool, card_hash: str = "", reason: str = "", by_tap: bool = False, conv: str = "") -> dict:
    if not by_tap:
        raise InteriorError("Quotation client ko sirf aapke tap se (app ka card).")
    with _LOCK, session() as s:
        q = _quote(s, tenant, ref)
        p = s.get(InteriorProject, q.project_id)
        if q.status != "waiting":
            return {"quote": _quote_view(s, q), "say": f"{q.number} pehle hi '{q.status}' hai."}
        if not card_hash or card_hash != q.hash or R.digest(_quote_show(q, p)) != q.hash:
            raise InteriorError("Card badal gaya hai; naya card dekhiye.")
        if not yes:
            q.status, q.reason, q.decided_at = "rejected", R.clean(reason, 300), now()
            return {"quote": _quote_view(s, q), "say": f"{q.number} nahi bheja."}
        q.status, q.decided_at = "approved", now()
        qid = q.id
    pdf = quote_pdf(tenant, qid, final=True)
    with session() as s:
        q = _quote(s, tenant, qid)
        p = s.get(InteriorProject, q.project_id)
        f = _store(tenant, f"{q.number.replace('/', '-')}_quotation.pdf", pdf, "application/pdf", conv, f"{q.number} quotation")
        q.file_id = f.get("file_id", "")
        t = _j(q.totals, {})
        st = _st(s, tenant)
        text = (f"Namaste {p.client} ji, {st.firm or 'hamari team'} ki taraf se {p.name} ka quotation {q.number} bhej rahe hain: "
                f"kul {_rs(t.get('total'))} (GST samet), {q.valid_until} tak valid. Koi sawaal ho to bataiye.")
        draft = {"text": text, "wa": _wa(p.phone, text), "sms": _sms(p.phone, text), "file_id": q.file_id}
        view = _quote_view(s, q)
    return {"quote": view, "draft": draft, "say": f"{view['number']} final PDF bana. Message draft taiyaar: aap khud bhejiye (PDF saath lagaiye)."}


def quote_mark(tenant: str, ref: str, what: str, actor: str = "owner") -> dict:
    """sent (she sent it) | accepted (client agreed: owner only) | rejected (client said no)."""
    with _LOCK, session() as s:
        q = _quote(s, tenant, ref)
        p = s.get(InteriorProject, q.project_id)
        if what == "sent":
            if q.status != "approved":
                raise InteriorError("Pehle aapki haan (card), phir bhejna.")
            q.status = "sent"
        elif what == "accepted":
            if not is_owner(actor):
                raise InteriorError("Client ki manzoori sirf owner darj karti hain.")
            if q.status not in ("approved", "sent"):
                raise InteriorError("Yeh quote aapki haan se client ko gaya hi nahi.")
            q.status = "accepted"
            p.contract_quote = q.id
            p.status = "approved"
            _make_stages(s, tenant, p, q)
        elif what == "rejected":
            q.status = "rejected"
        else:
            raise InteriorError("what = sent | accepted | rejected")
        _log(p, f"quote {q.number} {what}", actor_name(tenant, actor))
        view = _quote_view(s, q)
    return {"quote": view, "say": f"{view['number']}: {what}." + (" Payment stages bane." if what == "accepted" else "")}


def _make_stages(s, tenant: str, p: InteriorProject, q: InteriorQuote) -> None:
    if s.scalars(select(InteriorStage).where(InteriorStage.project_id == p.id)).first() is not None:
        return
    tmpl = _j(_st(s, tenant).stages, []) or R.DEFAULT_STAGES
    total = _j(q.totals, {}).get("total") or 0.0
    for i, (stg, amt) in enumerate(zip(tmpl, R.split_amount(total, tmpl))):
        s.add(InteriorStage(tenant_id=tenant, project_id=p.id, seq=i + 1, label=stg["label"], pct=stg["pct"], amount=amt))


def _contract_total(s, p: InteriorProject) -> float | None:
    if not p.contract_quote:
        return None
    q = s.get(InteriorQuote, p.contract_quote)
    base = _j(q.totals, {}).get("total") or 0.0 if q else 0.0
    ch = sum(c.delta for c in s.scalars(select(InteriorChange).where(InteriorChange.project_id == p.id,
                                                                     InteriorChange.status.in_(("approved", "applied")))).all())
    return R.r2(base + ch)


# ---------------------------------------------------------------- PDFs (reportlab, Latin text)
def _pdf(title: str, blocks: list[tuple], watermark: str = "") -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
    from xml.sax.saxutils import escape

    ss = getSampleStyleSheet()
    small = ss["BodyText"].clone("small", fontSize=8.5, leading=10.5)

    def txt(t: str) -> str:
        return escape(str(t or "").replace("₹", "Rs. "))
    story: list = []
    for b in blocks:
        if b[0] == "h":
            story.append(Paragraph(txt(b[1]), ss["Heading2"]))
        elif b[0] == "h1":
            story.append(Paragraph(txt(b[1]), ss["Title"]))
        elif b[0] == "p":
            story.append(Paragraph(txt(b[1]), ss["BodyText"]))
        elif b[0] == "table":
            rows = [[Paragraph(txt(c), small) for c in r] for r in b[1]]
            t = Table(rows, colWidths=b[2] if len(b) > 2 else None, repeatRows=1)
            t.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.3, colors.grey), ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#EEEEEE")),
                                   ("VALIGN", (0, 0), (-1, -1), "TOP")]))
            story.append(t)
        story.append(Spacer(1, 6))
    bio = io.BytesIO()

    def mark(c, _d):
        if watermark:
            c.saveState()
            c.setFont("Helvetica-Bold", 46)
            c.setFillColor(colors.Color(0.85, 0.2, 0.2, alpha=0.18))
            c.translate(300, 420)
            c.rotate(35)
            c.drawCentredString(0, 0, watermark)
            c.restoreState()
    doc = SimpleDocTemplate(bio, pagesize=A4, title=title, leftMargin=36, rightMargin=36, topMargin=36, bottomMargin=36)
    doc.build(story, onFirstPage=mark, onLaterPages=mark)
    return bio.getvalue()


def _head_blocks(st: InteriorSettings, title: str) -> list[tuple]:
    who = " · ".join(x for x in (st.firm, st.title or "Interior Designer", f"CoA {st.coa_reg}" if st.coa_reg else "",
                                 f"GSTIN {st.gstin}" if st.gstin else "") if x)
    contact = " · ".join(x for x in (st.address, st.phone, st.email) if x)
    return [("h1", title), ("p", who), ("p", contact)] if contact else [("h1", title), ("p", who)]


def quote_pdf(tenant: str, ref: str, final: bool = False) -> bytes:
    with session() as s:
        q = _quote(s, tenant, ref)
        p = s.get(InteriorProject, q.project_id)
        st = _st(s, tenant)
        lines, t = _j(q.lines, []), _j(q.totals, {})
        tmpl = _j(st.stages, []) or R.DEFAULT_STAGES
        blocks = _head_blocks(st, f"Quotation {q.number}")
        blocks += [("p", f"Client: {p.client}" + (f" (GSTIN {p.client_gstin})" if p.client_gstin else "") + f" · Site: {p.site} {p.city} {p.pin}"),
                   ("p", f"Project {p.number}: {p.name} · Date {_aware(q.created_at).date().isoformat() if q.created_at else ''} · Valid till {q.valid_until}")]
        rows = [["Room", "Item", "Spec", "HSN/SAC", "Qty", "Unit", "Rate", "Amount", "GST%"]]
        for c in lines:
            rows.append([c["room"], c["name"], c["spec"], c["hsn_sac"], "" if c["qty"] is None else f"{c['qty']:g}", c["unit"],
                         _pdf_rs(c["rate"]), _pdf_rs(c["amount"]), f"{c['gst_pct']:g}"])
        blocks.append(("table", rows, [55, 95, 110, 40, 32, 38, 55, 62, 32]))
        trows = [["", "Amount"], ["Works (before GST)", _pdf_rs(t.get("works"))]]
        if t.get("design_fee"):
            trows.append([f"Design fee ({t.get('fee_note')})", _pdf_rs(t.get("design_fee"))])
        for g in t.get("gst", []):
            on = g.get("on") or f"on {_pdf_rs(g['base'])}"
            if t.get("intra"):
                trows.append([f"CGST {g['pct'] / 2:g}% {on}", _pdf_rs(g["cgst"])])
                trows.append([f"SGST {g['pct'] / 2:g}% {on}", _pdf_rs(g["sgst"])])
            else:
                trows.append([f"IGST {g['pct']:g}% {on}", _pdf_rs(g["igst"])])
        trows.append(["TOTAL", _pdf_rs(t.get("total"))])
        blocks.append(("table", trows, [330, 120]))
        if t.get("words"):
            blocks.append(("p", f"In words: {t['words']}"))
        blocks.append(("h", "Payment schedule"))
        srows = [["Stage", "%", "Amount"]]
        for stg, amt in zip(tmpl, R.split_amount(t.get("total") or 0, tmpl)):
            srows.append([stg["label"], f"{stg['pct']:g}%", _pdf_rs(amt)])
        blocks.append(("table", srows, [300, 50, 100]))
        if _j(q.blanks, []):
            blocks.append(("p", "Blanks to fill: " + "; ".join(_j(q.blanks, []))))
        blocks.append(("h", "Terms"))
        blocks.append(("p", st.terms or "Payment as per the schedule above. Changes in scope are quoted separately (change order). "
                       "Brand / material as per the spec column; equivalent only with the client's written OK. GST as applicable."))
        blocks.append(("p", "Prepared with Hoola. Rates are the designer's own; quantities from site measurement or as stated."))
    return _pdf(f"Quotation {q.number}", blocks, "" if final else "DRAFT - NOT SENT")


# ---------------------------------------------------------------- payment stages + requests (her tap)
def _stage_view(s, x: InteriorStage) -> dict:
    p = s.get(InteriorProject, x.project_id)
    return {"id": x.id, "project": p.number if p else "", "project_id": x.project_id, "client": p.client if p else "", "seq": x.seq,
            "label": x.label, "pct": x.pct, "amount": x.amount, "amount_text": _rs(x.amount), "due_on": x.due_on, "status": x.status,
            "hash": x.hash, "paid_amount": x.paid_amount, "paid_ref": x.paid_ref, "paid_on": x.paid_on, "source": x.source}


def stages(tenant: str, project: str) -> dict:
    with session() as s:
        p = _proj(s, tenant, project)
        rows = s.scalars(select(InteriorStage).where(InteriorStage.project_id == p.id).order_by(InteriorStage.seq)).all()
        out = [_stage_view(s, x) for x in rows]
        total = _contract_total(s, p)
    paid = sum(x["paid_amount"] for x in out if x["status"] == "paid")
    return {"stages": out, "contract_total": total, "paid": R.r2(paid), "left": R.r2((total or 0) - paid) if total is not None else None,
            "say": (f"Kul {_rs(total)}, mila {_rs(paid)}." if total is not None else "Quote abhi pakka nahi (stages client ki manzoori par bante hain).")}


def stage_due(tenant: str, stage_id: str, due_on: str) -> dict:
    d = _date(due_on)
    if not d:
        raise InteriorError("Tareekh samajh nahi aayi.")
    with session() as s:
        x = s.get(InteriorStage, stage_id)
        if x is None or x.tenant_id != tenant:
            raise InteriorError("Stage nahi mila.")
        x.due_on = d.isoformat()
        v = _stage_view(s, x)
    return {"stage": v, "say": f"{v['label']}: {v['due_on']} tak."}


def _pay_show(s, x: InteriorStage, tenant: str) -> dict:
    p = s.get(InteriorProject, x.project_id)
    st = _st(s, tenant)
    return {"project": p.number, "client": p.client, "label": x.label, "amount": x.amount, "upi": st.upi_id}


def pay_request(tenant: str, stage_id: str, conv: str = "", actor: str = "owner") -> dict:
    with _LOCK, session() as s:
        x = s.get(InteriorStage, stage_id)
        if x is None or x.tenant_id != tenant:
            raise InteriorError("Stage nahi mila.")
        if x.status in ("paid", "cancelled"):
            raise InteriorError(f"Yeh stage '{x.status}' hai.")
        if x.amount <= 0:
            raise InteriorError("Is stage ki rakam 0 hai.")
        x.status = "waiting"
        x.hash = R.digest(_pay_show(s, x, tenant))
        v = _stage_view(s, x)
    say = f"{v['project']} {v['label']} {v['amount_text']}: client se maangne ka card aapki haan ke liye bheja hai."
    _publish(tenant, "interior_payment", {"stage": v, "say": say}, conv)
    return {"stage": v, "status": "needs_tap", "say": say}


def decide_payment(tenant: str, stage_id: str, yes: bool, card_hash: str = "", by_tap: bool = False) -> dict:
    if not by_tap:
        raise InteriorError("Paise maangna sirf aapke tap se (app ka card).")
    with _LOCK, session() as s:
        x = s.get(InteriorStage, stage_id)
        if x is None or x.tenant_id != tenant:
            raise InteriorError("Stage nahi mila.")
        if x.status != "waiting":
            return {"stage": _stage_view(s, x), "say": f"Yeh stage '{x.status}' hai."}
        if not card_hash or card_hash != x.hash or R.digest(_pay_show(s, x, tenant)) != x.hash:
            raise InteriorError("Card badal gaya hai; naya card dekhiye.")
        if not yes:
            x.status = "planned"
            return {"stage": _stage_view(s, x), "say": "Theek hai, abhi nahi maanga."}
        x.status = "requested"
        p = s.get(InteriorProject, x.project_id)
        st = _st(s, tenant)
        v = _stage_view(s, x)
        phone, client, firm = p.phone, p.client, st.firm
    link = _upi(tenant, v["amount"], f"{v['project']} {v['label']}")
    text = (f"Namaste {client} ji, {v['project']} ka '{v['label']}' stage: {v['amount_text']} dene ki vinti hai."
            + (f" UPI: {link}" if link else "") + f" Dhanyavaad, {firm or ''}".rstrip(", "))
    return {"stage": v, "draft": {"text": text, "wa": _wa(phone, text), "sms": _sms(phone, text), "upi": link},
            "say": "Message draft taiyaar; aap khud bhejiye. Paise mile to 'mila' darj kijiye (Hoola paise nahi leta)."}


def stage_paid(tenant: str, stage_id: str, amount=None, ref: str = "", day: str = "", actor: str = "owner") -> dict:
    if not is_owner(actor):
        raise InteriorError("Paise mile, yeh sirf owner darj karti hain.")
    with _LOCK, session() as s:
        x = s.get(InteriorStage, stage_id)
        if x is None or x.tenant_id != tenant:
            raise InteriorError("Stage nahi mila.")
        a = _num(amount)
        x.paid_amount = a if a is not None else x.amount
        x.paid_ref = R.clean(ref, 80)
        x.paid_on = (_date(day) or today(tenant)).isoformat()
        x.status = "paid"
        v = _stage_view(s, x)
    short = R.r2(v["amount"] - v["paid_amount"])
    return {"stage": v, "say": f"{v['label']}: {_rs(v['paid_amount'])} mila darj." + (f" {_rs(short)} kam aaya." if short > 0.5 else "")}


# ---------------------------------------------------------------- vendor work orders (her tap + fingerprint)
def _order_show(s, o: InteriorOrder) -> dict:
    p = s.get(InteriorProject, o.project_id)
    v = s.get(InteriorVendor, o.vendor_id)
    return {"number": o.number, "project": p.number if p else "", "vendor": v.name if v else "", "vendor_phone": v.phone if v else "",
            "amount": o.amount, "lines": [[ln.get("name"), ln.get("qty"), ln.get("unit"), ln.get("rate")] for ln in _j(o.lines, [])],
            "start": o.start, "finish": o.finish}


def _order_view(s, o: InteriorOrder) -> dict:
    show = _order_show(s, o)
    return {"id": o.id, **show, "amount_text": _rs(o.amount), "lines": _j(o.lines, []), "status": o.status, "hash": o.hash,
            "checks": _j(o.checks, []), "notes": o.notes, "file_id": o.file_id, "asked_by": o.asked_by, "reason": o.reason}


def order_request(tenant: str, project: str, vendor: str, lines: list | None = None, amount=None, start: str = "", finish: str = "",
                  notes: str = "", actor: str = "owner", conv: str = "") -> dict:
    """The vendor's price is HER word (what the vendor quoted her); never the client's rate. Only a card here."""
    with _LOCK, session() as s:
        p = _proj(s, tenant, project)
        v = _vendor(s, tenant, vendor)
        boq_rows = {c["id"]: c for c in _j(p.boq, [])}
        out, blanks, checks = [], [], []
        for ln in lines or []:
            if not isinstance(ln, dict):
                continue
            b = boq_rows.get(str(ln.get("line") or ln.get("id") or ""))
            name = R.clean(ln.get("name") or (b["name"] if b else ""), 200)
            qty = _num(ln.get("qty"))
            if qty is None and b:
                qty = b["qty"]
            unit = R.unit_of(str(ln.get("unit") or ""), "") or (b["unit"] if b else "nos")
            rate = _num(ln.get("rate"))
            row = {"name": name, "room": b["room"] if b else R.clean(ln.get("room"), 60), "qty": qty, "unit": unit, "rate": rate,
                   "amount": R.r2(qty * rate) if qty is not None and rate is not None else None, "line": b["id"] if b else "",
                   "spec": R.clean(ln.get("spec") or (b["spec"] if b else ""), 300)}
            if not name:
                blanks.append("item ka naam")
            if qty is None:
                blanks.append(f"{name}: qty")
            if rate is None and _num(amount) is None:
                blanks.append(f"{name}: vendor ka rate")
            if b and rate is not None and b.get("rate") is not None:
                if rate >= b["rate"]:
                    checks.append(f"{name}: vendor rate {_rs(rate)} client rate {_rs(b['rate'])} se kam nahi (margin nahi).")
            out.append(row)
        amt = _num(amount)
        if amt is None:
            amt = sum(r["amount"] for r in out if r["amount"] is not None) if out and not blanks else None
        if amt is None or amt <= 0:
            blanks.append("order ki kul rakam (vendor ka quote)")
        if blanks:
            raise InteriorError("Order ke liye yeh chahiye: " + "; ".join(dict.fromkeys(blanks)) + ". Hoola rate nahi gadhta.")
        sd, fd = _date(start), _date(finish)
        if start and not sd or finish and not fd:
            raise InteriorError("Tareekh samajh nahi aayi (DD-MM-YYYY).")
        if sd and fd and fd < sd:
            raise InteriorError("Khatam tareekh shuru se pehle hai.")
        if p.target_handover and fd and fd.isoformat() > p.target_handover:
            checks.append(f"Vendor ka kaam {fd} tak, handover {p.target_handover} ka hai.")
        for o in s.scalars(select(InteriorOrder).where(InteriorOrder.project_id == p.id, InteriorOrder.vendor_id == v.id,
                                                       InteriorOrder.status == "waiting")).all():
            if abs(o.amount - amt) < 0.5:
                return {"order": _order_view(s, o), "existing": True, "say": f"{o.number} ka card pehle se aapki haan ka intezaar kar raha hai."}
        st = _st(s, tenant)
        st.wo_seq += 1
        o = InteriorOrder(tenant_id=tenant, project_id=p.id, vendor_id=v.id, number=f"WO-{st.wo_seq:04d}", lines=_dump(out), amount=amt,
                          start=_iso(sd), finish=_iso(fd), notes=R.clean(notes, 500), checks=_dump(checks), asked_by=actor_name(tenant, actor))
        s.add(o)
        s.flush()
        o.hash = R.digest(_order_show(s, o))
        view = _order_view(s, o)
    say = f"{view['number']}: {view['vendor']} ko {view['amount_text']} ka order - aapki haan ka card bheja hai." + \
          (" Dhyaan: " + " ".join(checks) if checks else "")
    _publish(tenant, "interior_order", {"order": view, "say": say}, conv)
    return {"order": view, "status": "needs_tap", "say": say}


def order_card(tenant: str, order_id: str) -> dict | None:
    with session() as s:
        o = s.get(InteriorOrder, order_id)
        if o is None or o.tenant_id != tenant:
            return None
        return _order_view(s, o)


def order_challenge(tenant: str, order_id: str) -> dict:
    """approvekey's message format (kind 'interior'): her phone signs this order card as it is NOW."""
    from .. import approvekey as ak
    c = order_card(tenant, order_id)
    if not c or c["status"] != "waiting":
        raise InteriorError("Yeh card ab baaki nahi hai.")
    nonce = secrets.token_urlsafe(18)
    exp = datetime.now(timezone.utc) + ak.CHALLENGE_TTL
    msg = f"hoola-approve|v1|{tenant}|interior|{order_id}|{c['hash']}|{nonce}|{int(exp.timestamp())}"
    with session() as s:
        ch = ak.ApproveChallenge(tenant_id=tenant, kind="interior", target=order_id, digest=c["hash"], message=msg, expires_at=exp)
        s.add(ch)
        s.flush()
        chid = ch.id
    return {"challenge_id": chid, "message": msg, "show": {"title": f"Vendor order {c['number']}", "amount": c["amount_text"],
                                                             "payee": c["vendor"]}, "expires_at": exp.isoformat()}


def _verify_fingerprint(tenant: str, order_id: str, digest: str, proof: dict | None) -> None:
    from .. import approvekey as ak
    if not proof:
        raise InteriorError("fingerprint_needed")
    with session() as s:
        ak._settle(s, tenant)
        ok = ak._check(s, tenant, "interior", order_id, proof, digest)
    if not ok:
        raise InteriorError("fingerprint_invalid")


def decide_order(tenant: str, order_id: str, yes: bool, card_hash: str = "", reason: str = "", proof: dict | None = None,
                 by_tap: bool = False, conv: str = "") -> dict:
    if not by_tap:
        raise InteriorError("Vendor order sirf aapke tap se (app ka card).")
    with session() as s:
        o = s.get(InteriorOrder, order_id)
        if o is None or o.tenant_id != tenant:
            raise InteriorError("Card nahi mila.")
        if o.status != "waiting":
            return {"order": _order_view(s, o), "say": f"{o.number} pehle hi '{o.status}' hai."}
        if not card_hash or card_hash != o.hash or R.digest(_order_show(s, o)) != o.hash:
            raise InteriorError("Card badal gaya hai; naya card dekhiye.")
    if yes:
        from .. import approvekey as ak
        if ak.active_on(tenant):
            _verify_fingerprint(tenant, order_id, card_hash, proof)
    with _LOCK, session() as s:
        o = s.get(InteriorOrder, order_id)
        o.status = "approved" if yes else "rejected"
        o.reason = R.clean(reason, 300)
        p = s.get(InteriorProject, o.project_id)
        if yes and p.status in ("approved", "quote"):
            p.status = "production"
        view = _order_view(s, o)
    if not yes:
        return {"order": view, "say": f"{view['number']}: order nahi."}
    pdf = _order_pdf(tenant, order_id)
    f = _store(tenant, f"{view['number']}_work_order.pdf", pdf, "application/pdf", conv, f"{view['number']} work order")
    with session() as s:
        o = s.get(InteriorOrder, order_id)
        o.file_id = f.get("file_id", "")
        view = _order_view(s, o)
    text = (f"Namaste {view['vendor']} ji, {view['project']} ke liye work order {view['number']}: {view['amount_text']}"
            + (f", {view['start']} se {view['finish']} tak" if view["start"] and view["finish"] else "") + ". PDF bhej rahe hain.")
    return {"order": view, "draft": {"text": text, "wa": _wa(view["vendor_phone"], text), "sms": _sms(view["vendor_phone"], text)},
            "say": f"{view['number']} maana; work order PDF + vendor ke liye message draft (aap bhejiye). Paise aap khud denge."}


def _order_pdf(tenant: str, order_id: str) -> bytes:
    with session() as s:
        o = s.get(InteriorOrder, order_id)
        p = s.get(InteriorProject, o.project_id)
        v = s.get(InteriorVendor, o.vendor_id)
        st = _st(s, tenant)
        blocks = _head_blocks(st, f"Work order {o.number}")
        blocks += [("p", f"To: {v.name if v else ''} ({v.trade if v else ''}) {v.phone if v else ''}"),
                   ("p", f"Project {p.number}: {p.name} · Site: {p.site} {p.city} {p.pin}"),
                   ("p", f"Start {o.start or '-'} · Finish {o.finish or '-'}")]
        rows = [["Room", "Item", "Spec", "Qty", "Unit", "Rate", "Amount"]]
        for ln in _j(o.lines, []):
            rows.append([ln.get("room", ""), ln.get("name", ""), ln.get("spec", ""), "" if ln.get("qty") is None else f"{ln['qty']:g}",
                         ln.get("unit", ""), _pdf_rs(ln.get("rate")), _pdf_rs(ln.get("amount"))])
        blocks.append(("table", rows, [60, 110, 140, 40, 40, 60, 70]))
        blocks.append(("p", f"Order value: {_pdf_rs(o.amount)}"))
        if o.notes:
            blocks.append(("p", o.notes))
        blocks.append(("p", "Material / finish as per the spec. Any extra work only with a written change from us."))
    return _pdf(f"Work order {o.number}", blocks)


def order_done(tenant: str, order_id: str, actor: str = "owner") -> dict:
    with session() as s:
        o = s.get(InteriorOrder, order_id)
        if o is None or o.tenant_id != tenant:
            raise InteriorError("Order nahi mila.")
        if o.status != "approved":
            raise InteriorError(f"{o.number} '{o.status}' hai.")
        o.status = "done"
        v = _order_view(s, o)
    return {"order": v, "say": f"{v['number']} ka kaam poora darj."}


def orders(tenant: str, project: str = "", status: str = "") -> list[dict]:
    with session() as s:
        q = select(InteriorOrder).where(InteriorOrder.tenant_id == tenant)
        if project:
            q = q.where(InteriorOrder.project_id == _proj(s, tenant, project).id)
        if status:
            q = q.where(InteriorOrder.status == status)
        return [_order_view(s, o) for o in s.scalars(q.order_by(InteriorOrder.created_at.desc())).all()[:80]]


# ---------------------------------------------------------------- change orders
def _change_show(s, c: InteriorChange) -> dict:
    p = s.get(InteriorProject, c.project_id)
    return {"number": c.number, "project": p.number, "client": p.client, "what": c.what, "delta": c.delta,
            "add": [[a["name"], a["qty"], a["rate"], a["amount"]] for a in _j(c.add, [])], "remove": _j(c.remove, [])}


def _change_view(s, c: InteriorChange) -> dict:
    show = _change_show(s, c)
    return {"id": c.id, **show, "add": _j(c.add, []), "delta_text": ("+" if c.delta > 0 else "") + _rs(c.delta), "status": c.status,
            "hash": c.hash, "asked_by": c.asked_by, "reason": c.reason}


def change_request(tenant: str, project: str, what: str, add: list | None = None, remove: list | None = None, actor: str = "owner",
                   conv: str = "") -> dict:
    if not what:
        raise InteriorError("Badlav kya hai? (ek line)")
    with _LOCK, session() as s:
        p = _proj(s, tenant, project)
        if not p.contract_quote:
            raise InteriorError("Abhi quote pakka nahi; BOQ hi badaliye aur naya quote banaiye.")
        st = _st(s, tenant)
        rate_rows = s.scalars(select(InteriorRate).where(InteriorRate.tenant_id == tenant, InteriorRate.active.is_(True))).all()
        adds = [_resolve_line(s, tenant, p, ln, rate_rows, st.gst_works) for ln in (add or []) if isinstance(ln, dict)]
        blanks = [f"{a['name']}: " + " + ".join(a["blanks"]) for a in adds if a["blanks"]]
        if blanks:
            raise InteriorError("Change ke liye yeh chahiye: " + "; ".join(blanks) + ". Hoola rate nahi gadhta.")
        cur = {c["id"]: c for c in _j(p.boq, [])}
        rem = [str(x) for x in (remove or [])]
        missing = [x for x in rem if x not in cur]
        if missing:
            raise InteriorError("BOQ mein yeh line nahi: " + ", ".join(missing))
        delta = sum(a["amount"] * (1 + a["gst_pct"] / 100) for a in adds) - \
            sum((cur[x]["amount"] or 0) * (1 + cur[x]["gst_pct"] / 100) for x in rem)
        delta = R.r2(delta)
        st.co_seq += 1
        c = InteriorChange(tenant_id=tenant, project_id=p.id, number=f"CO-{st.co_seq:04d}", what=R.clean(what, 400), add=_dump(adds),
                           remove=_dump(rem), delta=delta, asked_by=actor_name(tenant, actor))
        s.add(c)
        s.flush()
        if abs(delta) < 0.5:
            _apply_change(s, tenant, p, c)
            c.status = "applied"
            v = _change_view(s, c)
            return {"change": v, "say": f"{v['number']}: daam nahi badla, seedha darj."}
        c.hash = R.digest(_change_show(s, c))
        v = _change_view(s, c)
    say = f"{v['number']}: {v['what']} ({v['delta_text']} GST samet) - daam badalta hai, aapki haan ka card bheja hai."
    _publish(tenant, "interior_change", {"change": v, "say": say}, conv)
    return {"change": v, "status": "needs_tap", "say": say}


def _apply_change(s, tenant: str, p: InteriorProject, c: InteriorChange) -> None:
    cur = [x for x in _j(p.boq, []) if x["id"] not in set(_j(c.remove, []))]
    for a in _j(c.add, []):
        a["co"] = c.number
        cur.append(a)
    p.boq = _dump(cur)
    if c.delta > 0.5:
        n = len(s.scalars(select(InteriorStage).where(InteriorStage.project_id == p.id)).all())
        s.add(InteriorStage(tenant_id=tenant, project_id=p.id, seq=n + 1, label=f"{c.number} (badlav): {c.what[:60]}", amount=c.delta,
                            source=c.id))
    elif c.delta < -0.5:
        left = -c.delta
        for x in sorted(s.scalars(select(InteriorStage).where(InteriorStage.project_id == p.id,
                                                              InteriorStage.status == "planned")).all(), key=lambda z: -z.seq):
            cut = min(left, x.amount)
            x.amount = R.r2(x.amount - cut)
            left = R.r2(left - cut)
            if left <= 0:
                break
    _log(p, f"change {c.number} {c.delta:+.0f}", "")


def decide_change(tenant: str, change_id: str, yes: bool, card_hash: str = "", reason: str = "", by_tap: bool = False) -> dict:
    if not by_tap:
        raise InteriorError("Daam badalne wala badlav sirf aapke tap se (app ka card).")
    with _LOCK, session() as s:
        c = s.get(InteriorChange, change_id)
        if c is None or c.tenant_id != tenant:
            raise InteriorError("Card nahi mila.")
        if c.status != "waiting":
            return {"change": _change_view(s, c), "say": f"{c.number} pehle hi '{c.status}' hai."}
        if not card_hash or card_hash != c.hash or R.digest(_change_show(s, c)) != c.hash:
            raise InteriorError("Card badal gaya hai; naya card dekhiye.")
        p = s.get(InteriorProject, c.project_id)
        if not yes:
            c.status, c.reason = "rejected", R.clean(reason, 300)
            return {"change": _change_view(s, c), "say": f"{c.number} nahi maana."}
        _apply_change(s, tenant, p, c)
        c.status = "approved"
        v = _change_view(s, c)
        phone, client = p.phone, p.client
        total = _contract_total(s, p)
    text = (f"Namaste {client} ji, {v['project']} mein badlav {v['number']}: {v['what']}. Rakam {v['delta_text']} (GST samet); "
            f"naya kul {_rs(total)}. Aapki likhit haan (reply) ke baad kaam karenge.")
    return {"change": v, "draft": {"text": text, "wa": _wa(phone, text), "sms": _sms(phone, text)},
            "say": f"{v['number']} maana; naya kul {_rs(total)}. Client ke liye message draft (aap bhejiye, unki haan likhit mein lijiye)."}


def changes(tenant: str, project: str) -> list[dict]:
    with session() as s:
        p = _proj(s, tenant, project)
        return [_change_view(s, c) for c in s.scalars(select(InteriorChange).where(InteriorChange.project_id == p.id)
                                                       .order_by(InteriorChange.created_at)).all()]


# ---------------------------------------------------------------- installation schedule
def _task_view(s, t: InteriorTask) -> dict:
    v = s.get(InteriorVendor, t.vendor_id) if t.vendor_id else None
    return {"id": t.id, "kind": t.kind, "kind_label": R.KINDS.get(t.kind, R.KINDS["other"])["label"], "room": t.room, "what": t.what,
            "vendor": v.name if v else "", "vendor_id": t.vendor_id, "start": t.start, "end": t.end, "status": t.status, "note": t.note}


def _task_checks(s, tenant: str, p: InteriorProject, t: InteriorTask) -> list[str]:
    out = []
    if t.vendor_id and t.start and t.end:
        others = s.scalars(select(InteriorTask).where(InteriorTask.tenant_id == tenant, InteriorTask.vendor_id == t.vendor_id,
                                                      InteriorTask.id != t.id, InteriorTask.status.in_(("planned", "started", "delayed")))).all()
        for o in others:
            if o.start and o.end and o.start <= t.end and t.start <= o.end:
                op = s.get(InteriorProject, o.project_id)
                out.append(f"Vendor usi samay {op.number if op else ''} {o.room} ({o.start} - {o.end}) mein laga hai.")
    if p.target_handover and t.end and t.end > p.target_handover:
        out.append(f"{t.end} handover ({p.target_handover}) ke baad.")
    d = _date(t.start)
    if d and d.weekday() == 6:
        out.append(f"{t.start} Ravivaar hai (society / building ka niyam dekhiye).")
    rows = [_task_view(s, x) for x in s.scalars(select(InteriorTask).where(InteriorTask.project_id == p.id)).all()]
    out += [x for x in R.sequence_problems(rows) if t.room in x or not t.room]
    return out


def task_add(tenant: str, project: str, what: str = "", kind: str = "", room: str = "", vendor: str = "", start: str = "",
             end: str = "", note: str = "", actor: str = "owner") -> dict:
    sd, ed = _date(start), _date(end) or _date(start)
    if not sd:
        raise InteriorError("Shuru ki tareekh chahiye.")
    if ed < sd:
        raise InteriorError("Khatam tareekh shuru se pehle.")
    with _LOCK, session() as s:
        p = _proj(s, tenant, project)
        v = _vendor(s, tenant, vendor) if vendor else None
        k = kind if kind in R.KINDS else R.kind_of(kind or what)
        t = InteriorTask(tenant_id=tenant, project_id=p.id, kind=k, room=R.clean(room, 60), what=R.clean(what or R.KINDS[k]["label"], 200),
                         vendor_id=v.id if v else "", start=sd.isoformat(), end=ed.isoformat(), note=R.clean(note, 300))
        s.add(t)
        s.flush()
        checks = _task_checks(s, tenant, p, t)
        if p.status in ("approved", "production"):
            p.status = "install"
        view = _task_view(s, t)
    return {"task": view, "checks": checks, "say": f"Schedule: {view['what']} {view['start']} - {view['end']}." +
            (" Dhyaan: " + " ".join(checks) if checks else "")}


def task_update(tenant: str, task_id: str, status: str = "", start: str = "", end: str = "", note: str = "", actor: str = "owner") -> dict:
    with _LOCK, session() as s:
        t = s.get(InteriorTask, task_id)
        if t is None or t.tenant_id != tenant:
            raise InteriorError("Kaam nahi mila.")
        if status:
            if status not in ("planned", "started", "done", "delayed", "cancelled"):
                raise InteriorError("status = planned | started | done | delayed | cancelled")
            t.status = status
        if start:
            d = _date(start)
            if not d:
                raise InteriorError("Tareekh samajh nahi aayi.")
            t.start = d.isoformat()
        if end:
            d = _date(end)
            if not d:
                raise InteriorError("Tareekh samajh nahi aayi.")
            t.end = d.isoformat()
        if note:
            t.note = R.clean(note, 300)
        p = s.get(InteriorProject, t.project_id)
        checks = _task_checks(s, tenant, p, t) if (start or end) else []
        view = _task_view(s, t)
    return {"task": view, "checks": checks, "say": f"{view['what']}: {view['status']}." + (" Dhyaan: " + " ".join(checks) if checks else "")}


def schedule(tenant: str, project: str) -> dict:
    with session() as s:
        p = _proj(s, tenant, project)
        rows = [_task_view(s, t) for t in s.scalars(select(InteriorTask).where(InteriorTask.project_id == p.id)).all()]
        target = p.target_handover
    rows.sort(key=lambda r: (r["start"], R.KINDS.get(r["kind"], R.KINDS["other"])["order"]))
    on = today(tenant).isoformat()
    late = [r for r in rows if r["status"] in ("planned", "started", "delayed") and r["end"] and r["end"] < on]
    return {"tasks": rows, "problems": R.sequence_problems(rows), "late": late, "target_handover": target,
            "say": f"{len(rows)} kaam schedule mein" + (f", {len(late)} der se" if late else "") + "."}


# ---------------------------------------------------------------- snags (photos without EXIF)
def _snag_view(s, x: InteriorSnag) -> dict:
    v = s.get(InteriorVendor, x.vendor_id) if x.vendor_id else None
    return {"id": x.id, "number": x.number, "room": x.room, "what": x.what, "vendor": v.name if v else "", "status": x.status,
            "photos": [{"id": ph["id"], "stage": ph["stage"]} for ph in _j(x.photos, [])], "by": x.by,
            "created_at": _aware(x.created_at).date().isoformat() if x.created_at else ""}


def _add_photo(tenant: str, x: InteriorSnag, data: bytes, stage: str) -> None:
    phs = _j(x.photos, [])
    if len(phs) >= MAX_PHOTOS:
        raise InteriorError(f"Ek snag par {MAX_PHOTOS} photo tak.")
    clean = _clean_image(data)
    pid = "ph" + secrets.token_hex(5)
    path = _dir(tenant, x.project_id) / f"{x.id}_{pid}.jpg"
    path.write_bytes(clean)
    phs.append({"id": pid, "path": str(path), "stage": stage if stage in ("before", "fixed") else "before"})
    x.photos = _dump(phs)


def snag_add(tenant: str, project: str, what: str, room: str = "", vendor: str = "", file_id: str = "", data: bytes | None = None,
             actor: str = "owner") -> dict:
    if not what:
        raise InteriorError("Kya kharabi hai? (ek line)")
    if file_id:
        _n, data = _load_file(tenant, file_id)
    with _LOCK, session() as s:
        p = _proj(s, tenant, project)
        v = _vendor(s, tenant, vendor) if vendor else None
        n = len(s.scalars(select(InteriorSnag).where(InteriorSnag.project_id == p.id)).all()) + 1
        x = InteriorSnag(tenant_id=tenant, project_id=p.id, number=n, room=R.clean(room, 60), what=R.clean(what, 400),
                         vendor_id=v.id if v else "", by=actor_name(tenant, actor))
        s.add(x)
        s.flush()
        if data:
            _add_photo(tenant, x, data, "before")
        view = _snag_view(s, x)
    return {"snag": view, "say": f"Snag #{view['number']} ({view['room'] or 'ghar'}): {view['what']}" + (" + photo (location / EXIF hataya)." if data else ".")}


def snag_photo(tenant: str, snag_id: str, data: bytes | None = None, file_id: str = "", stage: str = "fixed") -> dict:
    if file_id:
        _n, data = _load_file(tenant, file_id)
    if not data:
        raise InteriorError("Photo chahiye.")
    with _LOCK, session() as s:
        x = s.get(InteriorSnag, snag_id)
        if x is None or x.tenant_id != tenant:
            raise InteriorError("Snag nahi mila.")
        _add_photo(tenant, x, data, stage)
        view = _snag_view(s, x)
    return {"snag": view, "say": "Photo judi (location / EXIF hataya)."}


def snag_status(tenant: str, snag_id: str, status: str, actor: str = "owner") -> dict:
    if status not in ("open", "fixed", "closed"):
        raise InteriorError("status = open | fixed | closed")
    if status == "closed" and not is_owner(actor):
        raise InteriorError("Snag band sirf owner karti hain (theek hua dekh kar).")
    with session() as s:
        x = s.get(InteriorSnag, snag_id)
        if x is None or x.tenant_id != tenant:
            raise InteriorError("Snag nahi mila.")
        x.status = status
        view = _snag_view(s, x)
    return {"snag": view, "say": f"Snag #{view['number']}: {status}."}


def snag_file(tenant: str, snag_id: str, photo_id: str) -> Path:
    with session() as s:
        x = s.get(InteriorSnag, snag_id)
        if x is None or x.tenant_id != tenant:
            raise InteriorError("Snag nahi mila.")
        for ph in _j(x.photos, []):
            if ph["id"] == photo_id and Path(ph["path"]).exists():
                return Path(ph["path"])
    raise InteriorError("Photo nahi mili.")


def snags(tenant: str, project: str, status: str = "") -> list[dict]:
    with session() as s:
        p = _proj(s, tenant, project)
        q = select(InteriorSnag).where(InteriorSnag.project_id == p.id)
        if status:
            q = q.where(InteriorSnag.status == status)
        return [_snag_view(s, x) for x in s.scalars(q.order_by(InteriorSnag.number)).all()]


# ---------------------------------------------------------------- handover + warranty sheet
def warranty_rows(tenant: str, project: str) -> list[dict]:
    with session() as s:
        p = _proj(s, tenant, project)
        w = _j(_st(s, tenant).warranty, {})
        lines = _j(p.boq, [])
    seen: dict[str, dict] = {}
    for c in lines:
        k = c["kind"]
        row = seen.setdefault(k, {"kind": k, "label": R.KINDS.get(k, R.KINDS["other"])["label"], "items": [], "terms": w.get(k, "")})
        row["items"].append(f"{c['room']}: {c['name']}" + (f" ({c['spec']})" if c["spec"] else ""))
    return list(seen.values())


def handover(tenant: str, project: str, force: bool = False, reason: str = "", actor: str = "owner", conv: str = "") -> dict:
    if not is_owner(actor):
        raise InteriorError("Handover sirf owner.")
    with _LOCK, session() as s:
        p = _proj(s, tenant, project)
        open_snags = s.scalars(select(InteriorSnag).where(InteriorSnag.project_id == p.id, InteriorSnag.status != "closed")).all()
        if open_snags and not (force and reason):
            raise InteriorError(f"{len(open_snags)} snag abhi khule hain; band kijiye ya force + wajah dijiye.")
        unpaid = [x for x in s.scalars(select(InteriorStage).where(InteriorStage.project_id == p.id)).all()
                  if x.status not in ("paid", "cancelled") and x.amount > 0]
        p.handed_over = today(tenant).isoformat()
        p.status = "done"
        _log(p, "handover" + (f" (force: {R.clean(reason, 200)})" if force else ""), actor_name(tenant, actor))
        pnum, pname, client, site = p.number, p.name, p.client, f"{p.site} {p.city} {p.pin}"
        st = _st(s, tenant)
        head = _head_blocks(st, f"Handover & warranty - {pnum}")
        snag_rows = [[f"#{x.number}", x.room, x.what, x.status] for x in
                     s.scalars(select(InteriorSnag).where(InteriorSnag.project_id == p.id).order_by(InteriorSnag.number)).all()]
        unpaid_text = [f"{x.label}: {_rs(x.amount)}" for x in unpaid]
    rows = warranty_rows(tenant, project)
    blocks = head + [("p", f"Client: {client} · Site: {site} · Project: {pname} · Handover: {today(tenant).isoformat()}"),
                     ("h", "Warranty")]
    wt = [["Work", "Items", "Warranty (as stated by the designer)"]]
    for r in rows:
        wt.append([r["label"], "; ".join(r["items"])[:600], r["terms"] or "[not written - ask the designer]"])
    blocks.append(("table", wt, [90, 260, 170]))
    blocks.append(("p", "Brand warranties follow the brand's own card / invoice. Keep this sheet with your invoices."))
    if snag_rows:
        blocks.append(("h", "Snag list"))
        blocks.append(("table", [["No", "Room", "Snag", "Status"]] + snag_rows, [40, 80, 300, 60]))
    blocks.append(("p", "Received in good condition: ______________________ (client)      Date: __________"))
    pdf = _pdf(f"Handover {pnum}", blocks)
    f = _store(tenant, f"{pnum}_handover_warranty.pdf", pdf, "application/pdf", conv, f"{pnum} handover + warranty")
    missing = [r["label"] for r in rows if not r["terms"]]
    say = f"{pnum} handover darj; warranty sheet PDF bani."
    if missing:
        say += " Inki warranty aapne nahi likhi: " + ", ".join(missing) + " (Niyam mein likhiye)."
    if unpaid_text:
        say += " Baaki paise: " + "; ".join(unpaid_text) + "."
    return {"file": f, "missing_warranty": missing, "unpaid": unpaid_text, "say": say}


def forget(tenant: str, project: str, actor: str = "owner") -> dict:
    """DPDP: after the job ends, the client's personal details go; numbers / amounts stay for accounts."""
    if not is_owner(actor):
        raise InteriorError("Sirf owner.")
    with _LOCK, session() as s:
        p = _proj(s, tenant, project)
        if p.status not in ("done", "cancelled"):
            raise InteriorError("Project poora / radd hone ke baad hi.")
        p.client, p.phone, p.email, p.site, p.client_gstin = "Client (hataya)", "", "", "", ""
        p.pin, p.forgotten = "", True
        for x in s.scalars(select(InteriorSnag).where(InteriorSnag.project_id == p.id)).all():
            for ph in _j(x.photos, []):
                try:
                    Path(ph["path"]).unlink(missing_ok=True)
                except OSError:
                    pass
            x.photos = "[]"
        num = p.number
    return {"say": f"{num}: client ki personal details aur snag photos hata di (hisaab rakha)."}


# ---------------------------------------------------------------- views / day / worker
def projects(tenant: str, status: str = "") -> list[dict]:
    with session() as s:
        q = select(InteriorProject).where(InteriorProject.tenant_id == tenant)
        if status:
            q = q.where(InteriorProject.status == status)
        return [_project_view(s, p) for p in s.scalars(q.order_by(InteriorProject.created_at.desc())).all()[:100]]


def project_view(tenant: str, ref: str) -> dict:
    with session() as s:
        p = _proj(s, tenant, ref)
        v = _project_view(s, p, full=True)
        v["quotes"] = [_quote_view(s, q) for q in s.scalars(select(InteriorQuote).where(InteriorQuote.project_id == p.id)
                                                             .order_by(InteriorQuote.created_at)).all()]
        pid = p.id
    v["stages"] = stages(tenant, pid)["stages"]
    v["orders"] = orders(tenant, pid)
    v["changes"] = changes(tenant, pid)
    v["schedule"] = schedule(tenant, pid)
    v["snags"] = snags(tenant, pid)
    return v


def waiting(tenant: str) -> dict:
    with session() as s:
        qs = [_quote_view(s, q) for q in s.scalars(select(InteriorQuote).where(InteriorQuote.tenant_id == tenant, InteriorQuote.status == "waiting")).all()]
        os_ = [_order_view(s, o) for o in s.scalars(select(InteriorOrder).where(InteriorOrder.tenant_id == tenant, InteriorOrder.status == "waiting")).all()]
        cs = [_change_view(s, c) for c in s.scalars(select(InteriorChange).where(InteriorChange.tenant_id == tenant, InteriorChange.status == "waiting")).all()]
        ps = [_stage_view(s, x) for x in s.scalars(select(InteriorStage).where(InteriorStage.tenant_id == tenant, InteriorStage.status == "waiting")).all()]
    return {"quotes": qs, "orders": os_, "changes": cs, "payments": ps}


def day(tenant: str) -> dict:
    on = today(tenant)
    tom = (on + timedelta(days=1)).isoformat()
    with session() as s:
        ps = {p.id: p for p in s.scalars(select(InteriorProject).where(InteriorProject.tenant_id == tenant,
                                                                       InteriorProject.status.notin_(("done", "cancelled")))).all()}
        tasks = [(_task_view(s, t), ps[t.project_id].number) for t in s.scalars(select(InteriorTask).where(InteriorTask.tenant_id == tenant)).all()
                 if t.project_id in ps and t.status in ("planned", "started", "delayed")]
        due = [_stage_view(s, x) for x in s.scalars(select(InteriorStage).where(InteriorStage.tenant_id == tenant,
                                                                               InteriorStage.status == "planned")).all()
               if x.due_on and x.due_on <= on.isoformat()]
        snag_open: dict[str, int] = {}
        for x in s.scalars(select(InteriorSnag).where(InteriorSnag.tenant_id == tenant, InteriorSnag.status == "open")).all():
            if x.project_id in ps:
                snag_open[ps[x.project_id].number] = snag_open.get(ps[x.project_id].number, 0) + 1
        expiring = [_quote_view(s, q) for q in s.scalars(select(InteriorQuote).where(InteriorQuote.tenant_id == tenant,
                                                                                    InteriorQuote.status == "sent")).all()
                    if q.valid_until and q.valid_until <= (on + timedelta(days=3)).isoformat()]
        handovers = [p.number for p in ps.values() if p.target_handover and p.target_handover <= (on + timedelta(days=7)).isoformat()]
    starting = [f"{n} {t['what']} ({t['room'] or 'ghar'})" for t, n in tasks if t["start"] in (on.isoformat(), tom)]
    late = [f"{n} {t['what']}" for t, n in tasks if t["end"] and t["end"] < on.isoformat()]
    w = waiting(tenant)
    parts = []
    if starting:
        parts.append("Aaj / kal shuru: " + "; ".join(starting[:6]))
    if late:
        parts.append("Der se: " + "; ".join(late[:6]))
    if due:
        parts.append("Paise ka stage aa gaya: " + "; ".join(f"{d['project']} {d['label']} {d['amount_text']}" for d in due[:5]))
    if snag_open:
        parts.append("Khule snag: " + ", ".join(f"{k} ({v})" for k, v in snag_open.items()))
    if expiring:
        parts.append("Quote ki validity khatam ho rahi: " + ", ".join(q["number"] for q in expiring))
    if handovers:
        parts.append("Handover is hafte: " + ", ".join(handovers))
    nw = sum(len(v) for v in w.values())
    if nw:
        parts.append(f"{nw} card aapki haan ka intezaar kar rahe")
    return {"projects": len(ps), "starting": starting, "late": late, "stages_due": due, "snags": snag_open, "expiring": expiring,
            "handovers": handovers, "waiting": w, "say": (". ".join(parts) + ".") if parts else "Interior ka aaj sab theek."}


def overview(tenant: str) -> dict:
    return {"settings": settings(tenant), "projects": projects(tenant), "rates": rates(tenant), "vendors": vendors(tenant),
            "day": day(tenant), "styles": {k: {"label": v["label"], "palette": v["palette"]} for k, v in R.STYLES.items()},
            "units": R.UNITS}


def _once(s, tenant: str, key: str) -> bool:
    k = f"{tenant}:{key}"[:200]
    if s.get(InteriorAlert, k) is not None:
        return False
    s.add(InteriorAlert(key=k, tenant_id=tenant))
    return True


def _tenants() -> list[str]:
    with session() as s:
        return sorted(set(s.scalars(select(InteriorProject.tenant_id)).all()) | set(s.scalars(select(InteriorSettings.tenant_id)).all()))


def tick(when: datetime | None = None) -> dict:
    """Worker, every minute: late installation tasks and due payment stages (once each), the morning note once a day at her time
    (only when there is something)."""
    n = _aware(when) or now()
    out = {"alerts": 0, "morning": 0}
    for tenant in _tenants():
        try:
            loc = n.astimezone(_tz(tenant))
            on = loc.date().isoformat()
            alerts = []
            with _LOCK, session() as s:
                for t in s.scalars(select(InteriorTask).where(InteriorTask.tenant_id == tenant,
                                                              InteriorTask.status.in_(("planned", "started")))).all():
                    if t.end and t.end < on and _once(s, tenant, f"late:{t.id}:{t.end}"):
                        p = s.get(InteriorProject, t.project_id)
                        alerts.append({"kind": "late", "say": f"{p.number if p else ''}: {t.what} ({t.room or 'ghar'}) {t.end} tak hona tha, poora darj nahi."})
                for x in s.scalars(select(InteriorStage).where(InteriorStage.tenant_id == tenant, InteriorStage.status == "planned")).all():
                    if x.due_on and x.due_on <= on and _once(s, tenant, f"due:{x.id}:{x.due_on}"):
                        p = s.get(InteriorProject, x.project_id)
                        alerts.append({"kind": "stage", "stage": x.id, "say": f"{p.number if p else ''} '{x.label}' {_rs(x.amount)} ka samay aa gaya; "
                                                                             "maangne ka card banau?"})
            for a in alerts:
                _publish(tenant, "interior_alert", a)
                out["alerts"] += 1
            with session() as s:
                r = s.get(InteriorSettings, tenant)
                due_m = False
                if r is not None and r.last_morning != on:
                    try:
                        hh, mm = (int(z) for z in (r.morning_time or "08:30").split(":"))
                    except ValueError:
                        hh, mm = 8, 30
                    due_m = (loc.hour, loc.minute) >= (hh, mm)
                    if due_m:
                        r.last_morning = on
            if due_m:
                d = day(tenant)
                if d["say"] != "Interior ka aaj sab theek.":
                    _publish(tenant, "interior_day", {"say": d["say"], "starting": d["starting"], "late": d["late"],
                                                      "stages_due": len(d["stages_due"]), "projects": d["projects"]})
                    out["morning"] += 1
        except Exception:
            log.exception("interior tick failed for %s", tenant)
    return out
