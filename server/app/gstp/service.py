"""v0.93 GST Sahayak: Hoola does a GST practitioner's work end to end, for every client (or a business's own GST).

  1. Gets the job: she says "Sharma Traders ka October GST" or forwards sales bills / purchase bills (photos, PDF), the
     Tally / Vyapar / Excel sales and purchase registers, or the GSTR-2B JSON / Excel downloaded from the portal.
  2. Does the GST work, step by step, on screen: every invoice checked (GSTIN check digit, place of supply vs IGST or
     CGST + SGST, rate in the GST 2.0 slabs, tax maths, HSN digits by turnover, e-invoice IRN, duplicates, period), then
     GSTR-1 (B2B, B2CL, B2CS, exports, credit / debit notes, nil, HSN B2B / B2C, documents issued) as the portal's
     upload JSON, purchase register vs GSTR-2B reconciliation with IMS accept / reject / pending suggestions, and
     GSTR-3B (3.1, 3.2, 4A / 4B with 17(5) blocked credit and Rule 37 180-day reversals, legal set-off for least cash,
     interest + late fee when late).
  3. Checked against the law and earlier work: due dates per client (monthly / QRMP 22nd-24th state group /
     composition), the 3-year time bar, the 30 Nov ITC deadline, GSTR-9 / 9C need, GSTR-1 vs 3B (3.1 is hard-locked from
     GSTR-1 since July 2025: fix through GSTR-1A), DRC-01B / DRC-01C risks, notices with reply deadlines.
  4. Ready as a card, a file or a message draft: return summary card, GSTR-1 JSON + review Excel, supplier follow-up
     drafts for missing invoices, the client's "bills bhejiye" and "tax itna banta hai" drafts, notice reply drafts.
  5. Waits for her Haan: filing a return. Hoola prepares, she approves; only then the Computer Operator logs in to
     gst.gov.in (password + CAPTCHA confirm hers), uploads / fills and stops before the final submit. The EVC OTP on the
     taxpayer's phone and the tax payment (challan) are always theirs. Hoola never pays.
  6. Delivers and learns: ARN + filing log per return; item -> HSN / rate and supplier -> blocked credit are learned per
     client, so the next month's bills need fewer questions.
"""
from __future__ import annotations

import csv
import io
import json
import logging
import re
import threading
from datetime import date, datetime, timedelta
from typing import Any

from sqlalchemy import Boolean, DateTime, Float, Integer, String, Text, select
from sqlalchemy.orm import Mapped, mapped_column

from ..db import Base, new_id, now, session
from ..india.money import STATE_CODES
from . import rules as R

log = logging.getLogger("santosh.gstp")
_LOCK = threading.Lock()
SCHEMES = ("monthly", "qrmp", "composition")
DOC_KINDS = ("inv", "cn", "dn", "exp", "sez")
RETURN_KINDS = ("gstr1", "gstr3b", "cmp08", "gstr9")
RETURN_LABEL = {"gstr1": "GSTR-1", "gstr3b": "GSTR-3B", "cmp08": "CMP-08", "gstr9": "GSTR-9", "iff": "IFF", "pmt06": "PMT-06"}
STATUS_LABEL = {"draft": "Ban raha", "ready": "Aapki haan baaki", "approved": "Haan mili", "filing": "Portal par chal raha",
                "filed": "File ho gaya", "rejected": "Mana"}
PORTAL = "https://services.gst.gov.in/services/login"


class GstError(Exception):
    pass


class GstSettings(Base):
    __tablename__ = "gst_settings"
    tenant_id: Mapped[str] = mapped_column(String(64), primary_key=True)
    enabled: Mapped[bool] = mapped_column(Boolean, default=True)
    practitioner: Mapped[str] = mapped_column(String(120), default="")       # her name / firm
    enrolment: Mapped[str] = mapped_column(String(30), default="")           # GSTP enrolment id (PCT-02), if she is one
    morning_time: Mapped[str] = mapped_column(String(5), default="09:30")
    last_morning: Mapped[str] = mapped_column(String(10), default="")
    remind_days: Mapped[int] = mapped_column(Integer, default=5)             # ask clients for bills this many days before due


class GstClient(Base):
    __tablename__ = "gst_clients"
    id: Mapped[str] = mapped_column(String(40), primary_key=True, default=lambda: new_id("gc"))
    tenant_id: Mapped[str] = mapped_column(String(64), index=True)
    name: Mapped[str] = mapped_column(String(160), default="")
    gstin: Mapped[str] = mapped_column(String(15), default="", index=True)
    state: Mapped[str] = mapped_column(String(2), default="")
    scheme: Mapped[str] = mapped_column(String(12), default="monthly")
    aato: Mapped[float] = mapped_column(Float, default=0.0)                    # aggregate annual turnover (Rs), last FY
    einvoice: Mapped[bool] = mapped_column(Boolean, default=False)
    portal_user: Mapped[str] = mapped_column(String(60), default="")          # username only: the password is never kept
    login_as: Mapped[str] = mapped_column(String(12), default="taxpayer")     # taxpayer | practitioner
    contact_name: Mapped[str] = mapped_column(String(120), default="")
    contact_phone: Mapped[str] = mapped_column(String(20), default="")
    fee: Mapped[int] = mapped_column(Integer, default=0)                      # her monthly fee (paise), shown only to her
    notes: Mapped[str] = mapped_column(Text, default="")
    active: Mapped[bool] = mapped_column(Boolean, default=True)
    created_at: Mapped[datetime] = mapped_column(DateTime(timezone=True), default=now)


class GstDoc(Base):
    __tablename__ = "gst_docs"
    id: Mapped[str] = mapped_column(String(40), primary_key=True, default=lambda: new_id("gd"))
    tenant_id: Mapped[str] = mapped_column(String(64), index=True)
    client_id: Mapped[str] = mapped_column(String(40), index=True)
    period: Mapped[str] = mapped_column(String(7), index=True)
    direction: Mapped[str] = mapped_column(String(3), default="out")          # out (sales) | in (purchases)
    kind: Mapped[str] = mapped_column(String(4), default="inv")
    number: Mapped[str] = mapped_column(String(40), default="")
    doc_date: Mapped[str] = mapped_column(String(10), default="")
    party_gstin: Mapped[str] = mapped_column(String(15), default="", index=True)
    party_name: Mapped[str] = mapped_column(String(160), default="")
    pos: Mapped[str] = mapped_column(String(2), default="")
    rcm: Mapped[bool] = mapped_column(Boolean, default=False)
    items: Mapped[str] = mapped_column(Text, default="[]")
    txval: Mapped[float] = mapped_column(Float, default=0.0)
    igst: Mapped[float] = mapped_column(Float, default=0.0)
    cgst: Mapped[float] = mapped_column(Float, default=0.0)
    sgst: Mapped[float] = mapped_column(Float, default=0.0)
    cess: Mapped[float] = mapped_column(Float, default=0.0)
    value: Mapped[float] = mapped_column(Float, default=0.0)
    irn: Mapped[str] = mapped_column(String(70), default="")
    category: Mapped[str] = mapped_column(String(80), default="")            # purchases: what it was for (17(5) check)
    itc: Mapped[bool] = mapped_column(Boolean, default=True)                 # purchases: credit can be taken
    paid_on: Mapped[str] = mapped_column(String(10), default="")             # purchases: when the supplier was paid
    problems: Mapped[str] = mapped_column(Text, default="[]")
    source: Mapped[str] = mapped_column(String(20), default="manual")
    created_at: Mapped[datetime] = mapped_column(DateTime(timezone=True), default=now)


class Gst2B(Base):
    __tablename__ = "gst_2b"
    id: Mapped[str] = mapped_column(String(40), primary_key=True, default=lambda: new_id("g2"))
    tenant_id: Mapped[str] = mapped_column(String(64), index=True)
    client_id: Mapped[str] = mapped_column(String(40), index=True)
    period: Mapped[str] = mapped_column(String(7), index=True)
    typ: Mapped[str] = mapped_column(String(4), default="inv")               # inv | cn | dn | impg
    ctin: Mapped[str] = mapped_column(String(15), default="")
    trdnm: Mapped[str] = mapped_column(String(160), default="")
    inum: Mapped[str] = mapped_column(String(40), default="")
    dt: Mapped[str] = mapped_column(String(10), default="")
    val: Mapped[float] = mapped_column(Float, default=0.0)
    txval: Mapped[float] = mapped_column(Float, default=0.0)
    igst: Mapped[float] = mapped_column(Float, default=0.0)
    cgst: Mapped[float] = mapped_column(Float, default=0.0)
    sgst: Mapped[float] = mapped_column(Float, default=0.0)
    cess: Mapped[float] = mapped_column(Float, default=0.0)
    rev: Mapped[bool] = mapped_column(Boolean, default=False)
    itcavl: Mapped[bool] = mapped_column(Boolean, default=True)
    rsn: Mapped[str] = mapped_column(String(40), default="")
    supprd: Mapped[str] = mapped_column(String(7), default="")
    ims: Mapped[str] = mapped_column(String(10), default="")                 # her decision: accept | reject | pending
    created_at: Mapped[datetime] = mapped_column(DateTime(timezone=True), default=now)


class GstReturn(Base):
    __tablename__ = "gst_returns"
    id: Mapped[str] = mapped_column(String(40), primary_key=True, default=lambda: new_id("gr"))
    tenant_id: Mapped[str] = mapped_column(String(64), index=True)
    client_id: Mapped[str] = mapped_column(String(40), index=True)
    period: Mapped[str] = mapped_column(String(7), index=True)
    kind: Mapped[str] = mapped_column(String(8), default="gstr1")
    status: Mapped[str] = mapped_column(String(10), default="draft")
    data: Mapped[str] = mapped_column(Text, default="{}")                    # summary / computed tables
    portal_json: Mapped[str] = mapped_column(Text, default="")               # GSTR-1 upload JSON
    checks: Mapped[str] = mapped_column(Text, default="[]")
    approved_at: Mapped[datetime | None] = mapped_column(DateTime(timezone=True), nullable=True)
    operator_task: Mapped[str] = mapped_column(String(40), default="")
    filed_on: Mapped[str] = mapped_column(String(10), default="")
    arn: Mapped[str] = mapped_column(String(30), default="")
    log: Mapped[str] = mapped_column(Text, default="[]")
    created_at: Mapped[datetime] = mapped_column(DateTime(timezone=True), default=now)
    updated_at: Mapped[datetime] = mapped_column(DateTime(timezone=True), default=now, onupdate=now)


class GstNotice(Base):
    __tablename__ = "gst_notices"
    id: Mapped[str] = mapped_column(String(40), primary_key=True, default=lambda: new_id("gn"))
    tenant_id: Mapped[str] = mapped_column(String(64), index=True)
    client_id: Mapped[str] = mapped_column(String(40), index=True)
    form: Mapped[str] = mapped_column(String(12), default="OTHER")
    ref: Mapped[str] = mapped_column(String(60), default="")
    issued_on: Mapped[str] = mapped_column(String(10), default="")
    due_on: Mapped[str] = mapped_column(String(10), default="")
    text: Mapped[str] = mapped_column(Text, default="")
    amount: Mapped[float] = mapped_column(Float, default=0.0)
    status: Mapped[str] = mapped_column(String(10), default="open")          # open | replied | closed
    reply: Mapped[str] = mapped_column(Text, default="")
    created_at: Mapped[datetime] = mapped_column(DateTime(timezone=True), default=now)


class GstLearn(Base):
    """Per client: item words -> HSN / rate / UQC, supplier GSTIN -> blocked credit. Learned from her approved work."""
    __tablename__ = "gst_learn"
    id: Mapped[str] = mapped_column(String(40), primary_key=True, default=lambda: new_id("gl"))
    tenant_id: Mapped[str] = mapped_column(String(64), index=True)
    client_id: Mapped[str] = mapped_column(String(40), index=True)
    key: Mapped[str] = mapped_column(String(120), index=True)                # "item:<words>" | "supplier:<gstin>"
    value: Mapped[str] = mapped_column(Text, default="{}")
    seen: Mapped[int] = mapped_column(Integer, default=1)


# ---------------------------------------------------------------- helpers
def _j(s: str, default):
    try:
        v = json.loads(s or "")
        return v if isinstance(v, type(default)) else default
    except (TypeError, ValueError):
        return default


def _f(v) -> float:
    if v is None or v == "":
        return 0.0
    if isinstance(v, (int, float)):
        return float(v)
    s = re.sub(r"[^\d.\-]", "", str(v).replace(",", ""))
    try:
        return float(s) if s not in ("", "-", ".") else 0.0
    except ValueError:
        return 0.0


def _rs(v: float) -> str:
    v = round(float(v or 0), 2)
    neg = v < 0
    v = abs(v)
    whole = int(v)
    s = str(whole)
    if len(s) > 3:
        head, tail = s[:-3], s[-3:]
        head = re.sub(r"(\d)(?=(\d\d)+$)", r"\1,", head)
        s = f"{head},{tail}"
    paise = round((v - whole) * 100)
    return ("-" if neg else "") + "₹" + s + (f".{paise:02d}" if paise else "")


def _tz(tenant: str):
    from ..life import tz_of
    return tz_of(tenant)


def today(tenant: str) -> date:
    return datetime.now(_tz(tenant)).date()


MONTHS = {m: i for i, m in enumerate(("jan", "feb", "mar", "apr", "may", "jun", "jul", "aug", "sep", "oct", "nov", "dec"), 1)}


def parse_date(v) -> date | None:
    if v is None or v == "":
        return None
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    s = str(v).strip()
    m = re.fullmatch(r"(\d{4})-(\d{1,2})-(\d{1,2})(?:[ T].*)?", s)
    try:
        if m:
            return date(int(m.group(1)), int(m.group(2)), int(m.group(3)))
        m = re.fullmatch(r"(\d{1,2})[-/.](\d{1,2})[-/.](\d{2,4})", s)
        if m:
            y = int(m.group(3))
            return date(y + 2000 if y < 100 else y, int(m.group(2)), int(m.group(1)))
        m = re.fullmatch(r"(\d{1,2})[-/ ]([A-Za-z]{3})[A-Za-z]*[-/ ,]*(\d{2,4})", s)
        if m and m.group(2).lower() in MONTHS:
            y = int(m.group(3))
            return date(y + 2000 if y < 100 else y, MONTHS[m.group(2).lower()], int(m.group(1)))
    except ValueError:
        return None
    return None


def _dmy(iso: str) -> str:
    d = parse_date(iso)
    return d.strftime("%d-%m-%Y") if d else ""


def _norm_inv(n: str) -> str:
    s = re.sub(r"[^A-Z0-9]", "", str(n or "").upper())
    return re.sub(r"(?<![0-9])0+(?=\d)", "", s)


def _words(t: str) -> str:
    return " ".join(sorted(set(re.findall(r"[a-z0-9]{3,}", str(t or "").lower()))))[:110]


def _publish(tenant: str, kind: str, payload: dict, conv: str = "") -> None:
    from ..events import bus
    try:
        if not conv:
            from ..net.service import home_conv
            conv = home_conv(tenant)
        bus.publish(conv, kind, payload)
    except Exception:
        log.exception("gst publish failed")


def _log(r: GstReturn, what: str, by: str = "hoola") -> None:
    lg = _j(r.log, [])
    lg.append({"at": now().isoformat(timespec="seconds"), "what": what[:300], "by": by})
    r.log = json.dumps(lg[-150:], ensure_ascii=False)


# ---------------------------------------------------------------- settings
def _settings_row(s, tenant: str) -> GstSettings:
    r = s.get(GstSettings, tenant)
    if not r:
        r = GstSettings(tenant_id=tenant)
        s.add(r)
        s.flush()
    return r


def settings(tenant: str) -> dict[str, Any]:
    with session() as s:
        r = _settings_row(s, tenant)
        return {"enabled": r.enabled, "practitioner": r.practitioner, "enrolment": r.enrolment, "morning_time": r.morning_time,
                "remind_days": r.remind_days}


def save_settings(tenant: str, **kw: Any) -> dict[str, Any]:
    with session() as s:
        r = _settings_row(s, tenant)
        for k in ("practitioner", "enrolment"):
            if k in kw and kw[k] is not None:
                setattr(r, k, str(kw[k]).strip()[:120])
        if "enabled" in kw:
            r.enabled = bool(kw["enabled"])
        if kw.get("morning_time"):
            if not re.fullmatch(r"\d{1,2}:\d{2}", str(kw["morning_time"])):
                raise GstError("Subah ka time jaise 09:30")
            r.morning_time = str(kw["morning_time"]).zfill(5)
        if kw.get("remind_days") is not None:
            r.remind_days = max(1, min(int(kw["remind_days"]), 15))
    return settings(tenant)


# ---------------------------------------------------------------- clients
def _client_view(c: GstClient, full: bool = False) -> dict[str, Any]:
    v = {"id": c.id, "name": c.name, "gstin": c.gstin, "state": c.state, "state_name": STATE_CODES.get(c.state, ""),
         "scheme": c.scheme, "aato": c.aato, "einvoice": c.einvoice, "active": c.active, "contact_name": c.contact_name,
         "contact_phone": c.contact_phone, "login_as": c.login_as, "portal_user": c.portal_user}
    if full:
        v |= {"fee": round((c.fee or 0) / 100), "notes": c.notes}
    return v


def _client(s, tenant: str, ref: str) -> GstClient:
    ref = str(ref or "").strip()
    if not ref:
        rows = s.scalars(select(GstClient).where(GstClient.tenant_id == tenant, GstClient.active.is_(True))).all()
        if len(rows) == 1:
            return rows[0]
        raise GstError("Kis client ka? Naam ya GSTIN bataiye.")
    c = s.get(GstClient, ref)
    if c and c.tenant_id == tenant:
        return c
    up = ref.upper()
    c = s.scalars(select(GstClient).where(GstClient.tenant_id == tenant, GstClient.gstin == up)).first()
    if c:
        return c
    rows = s.scalars(select(GstClient).where(GstClient.tenant_id == tenant)).all()
    low = ref.lower()
    hits = [r for r in rows if low in r.name.lower()]
    if len(hits) == 1:
        return hits[0]
    if len(hits) > 1:
        raise GstError("Is naam ke kai client hain: " + ", ".join(f"{h.name} ({h.gstin})" for h in hits[:5]))
    raise GstError(f"Client '{ref}' nahi mila.")


def add_client(tenant: str, name: str, gstin: str, scheme: str = "", aato=None, portal_user: str = "", login_as: str = "",
               contact_name: str = "", contact_phone: str = "", fee=None, notes: str = "", einvoice: bool | None = None,
               client: str = "") -> dict[str, Any]:
    """Add or update a client. The portal password is never stored (she / the client types it on the portal)."""
    g = re.sub(r"\s", "", str(gstin or "")).upper()
    if g and not R.valid_gstin(g):
        raise GstError(f"GSTIN {g} sahi nahi (check digit / state code galat). Dobara dekhiye.")
    with _LOCK, session() as s:
        c = None
        if client:
            c = _client(s, tenant, client)
        elif g:
            c = s.scalars(select(GstClient).where(GstClient.tenant_id == tenant, GstClient.gstin == g)).first()
        new = c is None
        if new:
            if not g:
                raise GstError("GSTIN chahiye.")
            if not str(name or "").strip():
                raise GstError("Client ka naam chahiye.")
            c = GstClient(tenant_id=tenant, gstin=g, state=g[:2])
            s.add(c)
        if name:
            c.name = str(name).strip()[:160]
        if g:
            c.gstin, c.state = g, g[:2]
        if aato not in (None, ""):
            c.aato = max(0.0, _f(aato))
        if scheme:
            sc = str(scheme).lower().strip()
            sc = {"quarterly": "qrmp", "qtr": "qrmp", "comp": "composition", "month": "monthly"}.get(sc, sc)
            if sc not in SCHEMES:
                raise GstError("Scheme: monthly | qrmp | composition")
            c.scheme = sc
        elif new:
            c.scheme = "monthly" if c.aato > R.QRMP_LIMIT else "qrmp" if c.aato else "monthly"
        if c.scheme == "qrmp" and c.aato > R.QRMP_LIMIT:
            raise GstError("QRMP sirf 5 crore tak ke turnover par. Is client ko monthly rakhiye.")
        c.einvoice = bool(einvoice) if einvoice is not None else (c.aato > R.EINV_LIMIT or c.einvoice)
        for k, v in (("portal_user", portal_user), ("contact_name", contact_name), ("contact_phone", contact_phone)):
            if v:
                setattr(c, k, re.sub(r"\s+", " ", str(v)).strip()[:120])
        if login_as in ("taxpayer", "practitioner"):
            c.login_as = login_as
        if fee not in (None, ""):
            c.fee = int(round(_f(fee) * 100))
        if notes:
            c.notes = str(notes)[:2000]
        s.flush()
        v = _client_view(c, True)
    say = f"{'Naya client' if new else 'Client update'}: {v['name']} ({v['gstin']}, {v['state_name']}), {v['scheme']}."
    if v["einvoice"]:
        say += " E-invoice lagta hai: B2B bill par IRN chahiye."
    return {"client": v, "say": say}


def clients(tenant: str, active: bool = True) -> list[dict[str, Any]]:
    with session() as s:
        q = select(GstClient).where(GstClient.tenant_id == tenant)
        if active:
            q = q.where(GstClient.active.is_(True))
        return [_client_view(c, True) for c in s.scalars(q.order_by(GstClient.name)).all()]


def get_client(tenant: str, ref: str) -> dict[str, Any]:
    with session() as s:
        return _client_view(_client(s, tenant, ref), True)


# ---------------------------------------------------------------- learning
def _learned(s, tenant: str, client_id: str, key: str) -> dict:
    r = s.scalars(select(GstLearn).where(GstLearn.tenant_id == tenant, GstLearn.client_id == client_id, GstLearn.key == key)).first()
    return _j(r.value, {}) if r else {}


def _learn(s, tenant: str, client_id: str, key: str, value: dict) -> None:
    if not key.split(":", 1)[-1]:
        return
    r = s.scalars(select(GstLearn).where(GstLearn.tenant_id == tenant, GstLearn.client_id == client_id, GstLearn.key == key)).first()
    if r:
        r.value = json.dumps(_j(r.value, {}) | value, ensure_ascii=False)
        r.seen = (r.seen or 0) + 1
    else:
        s.add(GstLearn(tenant_id=tenant, client_id=client_id, key=key[:120], value=json.dumps(value, ensure_ascii=False)))


# ---------------------------------------------------------------- documents (sales / purchases)
def _items(raw, txval=None, rate=None, hsn="", desc="", qty=None, uqc="", igst=None, cgst=None, sgst=None, cess=None) -> list[dict]:
    if isinstance(raw, list) and raw:
        out = []
        for it in raw:
            if isinstance(it, dict):
                out.append({"hsn": str(it.get("hsn") or it.get("hsn_sc") or "").strip(), "desc": str(it.get("desc") or it.get("name") or "")[:120],
                            "qty": _f(it.get("qty")), "uqc": str(it.get("uqc") or "").upper()[:3], "rate": _f(it.get("rate") or it.get("rt")),
                            "txval": _f(it.get("txval") or it.get("taxable") or it.get("amount")),
                            "igst": it.get("igst"), "cgst": it.get("cgst"), "sgst": it.get("sgst"), "cess": _f(it.get("cess"))})
        return out
    return [{"hsn": str(hsn or "").strip(), "desc": str(desc or "")[:120], "qty": _f(qty), "uqc": str(uqc or "").upper()[:3],
             "rate": _f(rate), "txval": _f(txval), "igst": igst, "cgst": cgst, "sgst": sgst, "cess": _f(cess)}]


def _doc_view(d: GstDoc) -> dict[str, Any]:
    return {"id": d.id, "period": d.period, "direction": d.direction, "kind": d.kind, "number": d.number, "date": d.doc_date,
            "party_gstin": d.party_gstin, "party_name": d.party_name, "pos": d.pos, "rcm": d.rcm, "items": _j(d.items, []),
            "txval": d.txval, "igst": d.igst, "cgst": d.cgst, "sgst": d.sgst, "cess": d.cess, "value": d.value, "irn": d.irn,
            "category": d.category, "itc": d.itc, "paid_on": d.paid_on, "problems": _j(d.problems, []), "source": d.source}


def _build_doc(s, tenant: str, c: GstClient, direction: str, number: str, doc_date, party_gstin: str = "", party_name: str = "",
               pos: str = "", items=None, kind: str = "inv", rcm: bool = False, irn: str = "", category: str = "",
               paid_on: str = "", value=None, source: str = "manual", **line) -> tuple[GstDoc, list[str]]:
    direction = "in" if str(direction).lower() in ("in", "purchase", "inward", "kharid") else "out"
    kind = {"invoice": "inv", "credit": "cn", "credit_note": "cn", "debit": "dn", "debit_note": "dn", "export": "exp"}.get(
        str(kind or "inv").lower(), str(kind or "inv").lower())
    if kind not in DOC_KINDS:
        kind = "inv"
    d = parse_date(doc_date)
    if not d:
        raise GstError(f"Bill {number or ''}: date samajh nahi aayi ({doc_date!r}).")
    number = str(number or "").strip()[:40]
    if not number:
        raise GstError("Bill number chahiye.")
    pg = re.sub(r"\s", "", str(party_gstin or "")).upper()
    problems: list[str] = []
    if pg and not R.valid_gstin(pg):
        problems.append(f"Party GSTIN {pg} galat (check digit).")
    if pg and pg == c.gstin:
        problems.append("Party GSTIN client ka apna hai.")
    # place of supply: given, else the buyer's state (registered) / the client's own state
    if direction == "out":
        p = R.pos_code(pos) or (pg[:2] if pg and R.valid_gstin(pg) else "") or c.state
        sup_state = c.state
    else:
        p = R.pos_code(pos) or c.state
        sup_state = pg[:2] if pg else c.state
    sez = kind == "sez"
    exp = kind == "exp"
    is_intra = (not exp) and R.intra(sup_state, p, sez)
    its = _items(items, **{k: line.get(k) for k in ("txval", "rate", "hsn", "desc", "qty", "uqc", "igst", "cgst", "sgst", "cess")})
    tot = {"txval": 0.0, "igst": 0.0, "cgst": 0.0, "sgst": 0.0, "cess": 0.0}
    for it in its:
        learned = _learned(s, tenant, c.id, "item:" + _words(it["desc"])) if it["desc"] else {}
        if not it["hsn"] and learned.get("hsn"):
            it["hsn"] = learned["hsn"]
            it["learned"] = True
        if not it["rate"] and learned.get("rate") is not None and it["igst"] in (None, "") and it["cgst"] in (None, ""):
            it["rate"] = float(learned["rate"])
            it["learned"] = True
        if not it["uqc"]:
            it["uqc"] = learned.get("uqc") or ("NA" if R.is_service_hsn(it["hsn"]) else "NOS")
        given = any(it[k] not in (None, "") for k in ("igst", "cgst", "sgst"))
        if given:
            for k in ("igst", "cgst", "sgst"):
                it[k] = _f(it[k])
            if not it["rate"] and it["txval"]:
                it["rate"] = round((it["igst"] + it["cgst"] + it["sgst"]) * 100 / it["txval"], 2)
                it["rate"] = min(R.RATES, key=lambda r: abs(r - it["rate"]))
        else:
            sp = R.split_tax(it["txval"], it["rate"], is_intra)
            it |= {k: sp[k] for k in ("igst", "cgst", "sgst")}
        if exp and not given:
            it |= {"igst": 0.0, "cgst": 0.0, "sgst": 0.0}                 # LUT export without tax unless tax was charged
        if not exp:
            for q in R.check_line(d, it["rate"], it["hsn"], c.aato, it["txval"], it["igst"], it["cgst"], it["sgst"], is_intra):
                if direction == "out" or "IGST" in q or "CGST" in q or "Tax" in q:
                    problems.append(q)
        for k in tot:
            tot[k] += _f(it[k])
    if direction == "out" and c.einvoice and pg and kind in ("inv", "cn", "dn", "sez") and not irn:
        problems.append("E-invoice wala client: is B2B bill ka IRN nahi mila (IRP par report karna zaroori).")
    if direction == "out" and c.aato >= R.EINV_30DAY and pg and not irn and (date.today() - d).days > 30:
        problems.append("10 crore+ turnover: 30 din se purana bill IRP par report nahi hoga.")
    if direction == "in" and rcm is False and not pg:
        problems.append("Supplier ka GSTIN nahi: is par ITC nahi milega (unregistered se kharid).")
    blocked = direction == "in" and (R.blocked_hint(category) or R.blocked_hint(" ".join(i["desc"] for i in its)) or
                                     bool(_learned(s, tenant, c.id, "supplier:" + pg).get("blocked")))
    if blocked:
        problems.append("Yeh kharid Sec 17(5) blocked credit lagti hai (gaadi / khana / club / personal...): ITC nahi.")
    val = _f(value) if value not in (None, "") else round(sum(tot.values()), 2)
    doc = GstDoc(tenant_id=tenant, client_id=c.id, period=R.period_of(d), direction=direction, kind=kind, number=number,
                 doc_date=d.isoformat(), party_gstin=pg, party_name=str(party_name or "").strip()[:160], pos=p, rcm=bool(rcm),
                 items=json.dumps(its, ensure_ascii=False), value=val, irn=str(irn or "")[:70], category=str(category or "")[:80],
                 itc=not blocked and (bool(pg) or bool(rcm)), paid_on=(parse_date(paid_on).isoformat() if parse_date(paid_on) else ""),
                 source=source, **{k: round(v, 2) for k, v in tot.items()})
    if value not in (None, "") and abs(val - sum(tot.values())) > max(2.0, 0.01 * val):
        problems.append(f"Bill total {_rs(val)} hai, taxable + tax {_rs(sum(tot.values()))} banta hai.")
    doc.problems = json.dumps(problems, ensure_ascii=False)
    return doc, problems


def add_doc(tenant: str, client: str, direction: str, number: str, doc_date, party_gstin: str = "", party_name: str = "",
            pos: str = "", items=None, kind: str = "inv", rcm: bool = False, irn: str = "", category: str = "",
            paid_on: str = "", value=None, source: str = "manual", replace: bool = False, **line) -> dict[str, Any]:
    with _LOCK, session() as s:
        c = _client(s, tenant, client)
        doc, problems = _build_doc(s, tenant, c, direction, number, doc_date, party_gstin, party_name, pos, items, kind, rcm, irn,
                                   category, paid_on, value, source, **line)
        dup = s.scalars(select(GstDoc).where(GstDoc.tenant_id == tenant, GstDoc.client_id == c.id, GstDoc.direction == doc.direction,
                                             GstDoc.kind == doc.kind, GstDoc.party_gstin == doc.party_gstin)).all()
        same = [x for x in dup if _norm_inv(x.number) == _norm_inv(doc.number) and x.doc_date[:4] == doc.doc_date[:4]]
        if same and not replace:
            return {"saved": False, "duplicate": _doc_view(same[0]),
                    "say": f"Bill {doc.number} pehle se hai ({same[0].doc_date}, {_rs(same[0].value)}). Badalna ho to replace=true."}
        for x in same:
            s.delete(x)
        s.add(doc)
        s.flush()
        v = _doc_view(doc)
    say = (f"{'Sale' if v['direction'] == 'out' else 'Kharid'} bill {v['number']} ({v['date']}) {v['party_name'] or v['party_gstin'] or 'B2C'}: "
           f"taxable {_rs(v['txval'])}, tax {_rs(v['igst'] + v['cgst'] + v['sgst'])}.")
    if problems:
        say += " Dhyaan: " + " ".join(problems[:4])
    return {"saved": True, "doc": v, "problems": problems, "say": say}


def docs(tenant: str, client: str, period: str = "", direction: str = "", limit: int = 500) -> list[dict[str, Any]]:
    with session() as s:
        c = _client(s, tenant, client)
        q = select(GstDoc).where(GstDoc.tenant_id == tenant, GstDoc.client_id == c.id)
        if period:
            q = q.where(GstDoc.period.in_(_months(c, R.norm_period(period))))
        if direction:
            q = q.where(GstDoc.direction == direction)
        return [_doc_view(d) for d in s.scalars(q.order_by(GstDoc.doc_date, GstDoc.number).limit(limit)).all()]


def delete_doc(tenant: str, doc_id: str) -> dict:
    with session() as s:
        d = s.get(GstDoc, doc_id)
        if not d or d.tenant_id != tenant:
            raise GstError("Bill nahi mila.")
        s.delete(d)
    return {"deleted": True}


def mark_paid(tenant: str, client: str, number: str, paid_on: str, supplier: str = "") -> dict:
    with session() as s:
        c = _client(s, tenant, client)
        rows = s.scalars(select(GstDoc).where(GstDoc.tenant_id == tenant, GstDoc.client_id == c.id, GstDoc.direction == "in")).all()
        hits = [d for d in rows if _norm_inv(d.number) == _norm_inv(number) and (not supplier or supplier.upper() in (d.party_gstin + d.party_name.upper()))]
        if not hits:
            raise GstError(f"Kharid bill {number} nahi mila.")
        pd = parse_date(paid_on) or today(tenant)
        for d in hits:
            d.paid_on = pd.isoformat()
    return {"updated": len(hits), "say": f"Bill {number}: supplier ko payment {pd.isoformat()} note kiya (Rule 37)."}


# ---------------------------------------------------------------- import registers / GSTR-2B
_HEAD = {
    "number": ("invoice no", "invoice number", "inv no", "bill no", "voucher no", "document no", "doc no", "invoice #", "vch no",
               "note number", "note no", "inum", "bill number", "invoice"),
    "doc_date": ("invoice date", "date", "bill date", "voucher date", "doc date", "note date", "idt", "dt"),
    "party_gstin": ("gstin", "gstin/uin", "party gstin", "gst no", "customer gstin", "supplier gstin", "gstin of supplier",
                    "gstin of recipient", "ctin", "gstin/uin of recipient"),
    "party_name": ("party", "party name", "customer", "customer name", "supplier", "supplier name", "name", "particulars",
                   "receiver name", "trade/legal name", "trade name", "legal name", "trdnm"),
    "pos": ("place of supply", "pos", "state"),
    "txval": ("taxable value", "taxable", "taxable amount", "assessable value", "txval", "taxable value (₹)", "taxable value(₹)"),
    "rate": ("rate", "gst rate", "tax rate", "gst %", "rate %", "rt", "applicable % of tax rate"),
    "igst": ("igst", "integrated tax", "igst amount", "integrated tax(₹)", "integrated tax (₹)", "iamt"),
    "cgst": ("cgst", "central tax", "cgst amount", "central tax(₹)", "central tax (₹)", "camt"),
    "sgst": ("sgst", "state/ut tax", "sgst amount", "utgst", "sgst/utgst", "state/ut tax(₹)", "state/ut tax (₹)", "samt"),
    "cess": ("cess", "cess amount", "cess(₹)", "cess (₹)", "csamt"),
    "value": ("invoice value", "total", "bill amount", "grand total", "invoice amount", "val", "invoice value(₹)", "invoice value (₹)",
              "gross total"),
    "hsn": ("hsn", "hsn/sac", "hsn code", "sac", "hsn_sc"),
    "desc": ("item", "item name", "description", "product", "goods", "particular", "stock item"),
    "qty": ("qty", "quantity"),
    "uqc": ("uqc", "unit", "uom"),
    "rcm": ("reverse charge", "rcm", "supply attract reverse charge"),
    "kind": ("type", "voucher type", "note type", "invoice type", "document type"),
    "irn": ("irn",),
    "itcavl": ("itc availability", "itc available"),
    "rsn": ("reason",),
    "supprd": ("gstr-1/iff/gstr-5 period", "period"),
}


def _key_of(h: str) -> str | None:
    h = re.sub(r"\s+", " ", str(h or "").strip().lower()).replace("\n", " ")
    if not h:
        return None
    for k, names in _HEAD.items():
        if h in names:
            return k
    for k, names in _HEAD.items():
        for n in names:
            if len(n) > 4 and (h.startswith(n) or n in h):
                return k
    return None


def _table(name: str, data: bytes) -> list[dict]:
    """Rows from CSV / XLSX (the sheet with the most invoice-like header; header row found in the first 12 rows)."""
    rows_sets: list[list[list]] = []
    low = name.lower()
    if low.endswith((".xlsx", ".xlsm", ".xls")):
        try:
            import openpyxl
            wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        except Exception as e:
            raise GstError(f"Excel nahi khula: {e}")
        prefer = [ws for ws in wb.worksheets if ws.title.strip().lower() in ("b2b", "sales", "purchase", "purchases", "b2b-invoices")]
        for ws in prefer + [w for w in wb.worksheets if w not in prefer]:
            rows_sets.append([list(r) for r in ws.iter_rows(values_only=True)][:20000])
    else:
        text = data.decode("utf-8-sig", errors="replace")
        try:
            dialect = csv.Sniffer().sniff(text[:4000], delimiters=",;\t|")
        except csv.Error:
            dialect = csv.excel
        rows_sets.append(list(csv.reader(io.StringIO(text), dialect)))
    best: tuple[int, list[dict]] = (0, [])
    for rows in rows_sets:
        for hi in range(min(12, len(rows))):
            keys = [_key_of(h) for h in rows[hi]]
            score = sum(1 for k in set(keys) if k in ("number", "doc_date", "party_gstin", "txval", "igst", "cgst", "value"))
            if score >= 3 and score > best[0]:
                # 2-row headers (the portal's 2B Excel puts tax heads in the row below)
                sub = rows[hi + 1] if hi + 1 < len(rows) else []
                for i, k in enumerate(keys):
                    if k is None and i < len(sub):
                        keys[i] = _key_of(sub[i])
                out = []
                for r in rows[hi + 1:]:
                    if not any(x not in (None, "") for x in r):
                        continue
                    d = {}
                    for i, k in enumerate(keys):
                        if k and i < len(r) and k not in d and r[i] not in (None, ""):
                            d[k] = r[i]
                    if d.get("number") and (d.get("doc_date") or d.get("txval")):
                        out.append(d)
                best = (score, out)
    if not best[1]:
        raise GstError("File mein invoice wali table nahi mili (Invoice No, Date, GSTIN, Taxable jaise column chahiye).")
    return best[1]


def _load_file(tenant: str, file_id: str) -> tuple[str, bytes]:
    from ..mailbox import service as ms
    try:
        name, _mime, data = ms.load_files(tenant, [{"type": "file", "file_id": file_id}])[0]
    except Exception as e:
        raise GstError(f"File {file_id} nahi mili: {e}")
    return name, data


def import_register(tenant: str, client: str, direction: str, file_id: str = "", data: bytes | None = None,
                    name: str = "register.csv", period: str = "") -> dict[str, Any]:
    """A sales / purchase register (Tally / Vyapar / Busy / Excel / CSV export): every row checked and saved."""
    if file_id:
        name, data = _load_file(tenant, file_id)
    if not data:
        raise GstError("File chahiye.")
    rows = _table(name, data)
    want = R.norm_period(period) if period else ""
    saved, skipped, dups, problems = 0, 0, 0, []
    grouped: dict[tuple, dict] = {}
    for r in rows:                                   # one row per item line: group by invoice
        k = (str(r.get("number")).strip(), str(r.get("party_gstin") or "").strip().upper())
        g = grouped.setdefault(k, {**r, "items": []})
        g["items"].append({"hsn": r.get("hsn"), "desc": r.get("desc"), "qty": r.get("qty"), "uqc": r.get("uqc"),
                           "rate": r.get("rate"), "txval": r.get("txval"), "igst": r.get("igst"), "cgst": r.get("cgst"),
                           "sgst": r.get("sgst"), "cess": r.get("cess")})
    with _LOCK, session() as s:
        c = _client(s, tenant, client)
        existing = s.scalars(select(GstDoc).where(GstDoc.tenant_id == tenant, GstDoc.client_id == c.id)).all()
        have = {(x.direction, x.party_gstin, _norm_inv(x.number), x.doc_date[:4]) for x in existing}
        for (num, _g), r in grouped.items():
            kind = str(r.get("kind") or "inv").lower()
            kind = "cn" if "credit" in kind or kind in ("c", "cn") else "dn" if "debit" in kind or kind in ("d", "dn") else \
                "exp" if "export" in kind else "sez" if "sez" in kind else "inv"
            rcm = str(r.get("rcm") or "").strip().lower() in ("y", "yes", "true", "1")
            for it in r["items"]:
                if it.get("igst") in (None, "") and it.get("cgst") in (None, "") and not it.get("rate"):
                    it.pop("igst"), it.pop("cgst"), it.pop("sgst")
            try:
                doc, pr = _build_doc(s, tenant, c, direction, num, r.get("doc_date"), r.get("party_gstin") or "", r.get("party_name") or "",
                                     r.get("pos") or "", r["items"], kind, rcm, str(r.get("irn") or ""), "", "", r.get("value"), "import")
            except GstError as e:
                skipped += 1
                problems.append(str(e))
                continue
            if want and doc.period not in _months(c, want):
                skipped += 1
                continue
            key = (doc.direction, doc.party_gstin, _norm_inv(doc.number), doc.doc_date[:4])
            if key in have:
                dups += 1
                continue
            have.add(key)
            s.add(doc)
            saved += 1
            problems += [f"{doc.number}: {p}" for p in pr]
    say = f"{saved} bill save ({'sale' if direction == 'out' else 'kharid'})." + (f" {dups} pehle se the." if dups else "") + \
          (f" {skipped} chhode (date / period)." if skipped else "")
    if problems:
        say += f" {len(problems)} galtiyan mili, pehli: " + "; ".join(problems[:3])
    return {"saved": saved, "duplicates": dups, "skipped": skipped, "problems": problems[:200], "say": say}


def _flt2b(v) -> bool:
    return str(v or "").strip().upper() in ("Y", "YES", "TRUE", "1")


def import_2b(tenant: str, client: str, period: str, file_id: str = "", data: bytes | None = None, name: str = "gstr2b.json") -> dict[str, Any]:
    """GSTR-2B from the portal (Download JSON, or the Excel). Rows replace that period's earlier 2B."""
    if file_id:
        name, data = _load_file(tenant, file_id)
    if not data:
        raise GstError("2B file chahiye.")
    p = R.norm_period(period)
    rows: list[dict] = []
    if name.lower().endswith(".json") or data[:1] in (b"{", b"["):
        try:
            j = json.loads(data.decode("utf-8-sig"))
        except ValueError as e:
            raise GstError(f"2B JSON kharab: {e}")
        d = j.get("data", j) if isinstance(j, dict) else {}
        d = d.get("data", d) if isinstance(d, dict) and "docdata" not in d else d
        rp = str(d.get("rtnprd") or "")
        if re.fullmatch(r"\d{6}", rp) and R.norm_period(rp) != p:
            raise GstError(f"Yeh 2B {R.nice_period(rp)} ka hai, aapne {R.nice_period(p)} kaha.")
        dd = d.get("docdata") or {}
        for sec, typ in (("b2b", "inv"), ("b2ba", "inv"), ("cdnr", "note"), ("cdnra", "note"), ("ecom", "inv")):
            for sup in dd.get(sec) or []:
                docs_ = sup.get("inv") or sup.get("nt") or []
                for x in docs_:
                    t = typ if typ == "inv" else ("cn" if str(x.get("typ") or "").upper() == "C" else "dn")
                    items = x.get("items") or []
                    tx = sum(_f(i.get("txval")) for i in items) or _f(x.get("txval"))
                    rows.append({"typ": t, "ctin": sup.get("ctin", ""), "trdnm": sup.get("trdnm", ""), "inum": x.get("inum") or x.get("ntnum") or "",
                                 "dt": x.get("dt") or "", "val": _f(x.get("val")), "txval": tx,
                                 "igst": sum(_f(i.get("igst")) for i in items) or _f(x.get("igst")),
                                 "cgst": sum(_f(i.get("cgst")) for i in items) or _f(x.get("cgst")),
                                 "sgst": sum(_f(i.get("sgst")) for i in items) or _f(x.get("sgst")),
                                 "cess": sum(_f(i.get("cess")) for i in items) or _f(x.get("cess")),
                                 "rev": _flt2b(x.get("rev")), "itcavl": str(x.get("itcavl") or "Y").upper() != "N",
                                 "rsn": str(x.get("rsn") or ""), "supprd": str(sup.get("supprd") or "")})
        for x in dd.get("impg") or []:
            rows.append({"typ": "impg", "ctin": "", "trdnm": f"Import {x.get('portcode', '')}", "inum": str(x.get("boenum") or ""),
                         "dt": x.get("boedt") or "", "val": _f(x.get("txval")) + _f(x.get("igst")), "txval": _f(x.get("txval")),
                         "igst": _f(x.get("igst")), "cgst": 0.0, "sgst": 0.0, "cess": _f(x.get("cess")), "rev": False,
                         "itcavl": True, "rsn": "", "supprd": ""})
    else:
        for r in _table(name, data):
            rows.append({"typ": "cn" if "credit" in str(r.get("kind") or "").lower() else "inv", "ctin": str(r.get("party_gstin") or "").upper(),
                         "trdnm": str(r.get("party_name") or ""), "inum": str(r.get("number") or ""), "dt": r.get("doc_date"),
                         "val": _f(r.get("value")), "txval": _f(r.get("txval")), "igst": _f(r.get("igst")), "cgst": _f(r.get("cgst")),
                         "sgst": _f(r.get("sgst")), "cess": _f(r.get("cess")), "rev": _flt2b(r.get("rcm")),
                         "itcavl": str(r.get("itcavl") or "Yes").strip().lower() not in ("no", "n"), "rsn": str(r.get("rsn") or ""),
                         "supprd": str(r.get("supprd") or "")})
    if not rows:
        raise GstError("2B mein koi bill nahi mila.")
    with _LOCK, session() as s:
        c = _client(s, tenant, client)
        for old in s.scalars(select(Gst2B).where(Gst2B.tenant_id == tenant, Gst2B.client_id == c.id, Gst2B.period == p)).all():
            s.delete(old)
        for r in rows:
            d = parse_date(r["dt"])
            s.add(Gst2B(tenant_id=tenant, client_id=c.id, period=p, typ=r["typ"], ctin=str(r["ctin"])[:15], trdnm=str(r["trdnm"])[:160],
                        inum=str(r["inum"])[:40], dt=d.isoformat() if d else "", val=r["val"], txval=r["txval"], igst=r["igst"],
                        cgst=r["cgst"], sgst=r["sgst"], cess=r["cess"], rev=r["rev"], itcavl=r["itcavl"], rsn=r["rsn"][:40],
                        supprd=str(r["supprd"])[:7]))
            if r["ctin"] and r["trdnm"]:
                _learn(s, tenant, c.id, "supplier:" + r["ctin"], {"name": r["trdnm"]})
    tax = sum(r["igst"] + r["cgst"] + r["sgst"] for r in rows if r["itcavl"] and not r["rev"] and r["typ"] != "cn") - \
        sum(r["igst"] + r["cgst"] + r["sgst"] for r in rows if r["itcavl"] and r["typ"] == "cn")
    return {"rows": len(rows), "itc": round(tax, 2),
            "say": f"GSTR-2B {R.nice_period(p)}: {len(rows)} bill, ITC dikh raha {_rs(tax)}. Ab kharid register se milaata hoon."}


# ---------------------------------------------------------------- reconciliation: books vs GSTR-2B (+ IMS)
def _months(c: GstClient, p: str) -> list[str]:
    return R.quarter_months(p) if c.scheme == "qrmp" and R.is_quarter_end(p) else [p]


def reconcile(tenant: str, client: str, period: str, tolerance: float = 10.0) -> dict[str, Any]:
    p = R.norm_period(period)
    with session() as s:
        c = _client(s, tenant, client)
        months = _months(c, p)
        twob = s.scalars(select(Gst2B).where(Gst2B.tenant_id == tenant, Gst2B.client_id == c.id, Gst2B.period.in_(months))).all()
        if not twob:
            raise GstError(f"{R.nice_period(p)} ka GSTR-2B abhi nahi aaya. Portal se 2B (JSON) download karke bhejiye "
                           "(ya Hoola se kahiye: portal se 2B le aao).")
        books = s.scalars(select(GstDoc).where(GstDoc.tenant_id == tenant, GstDoc.client_id == c.id, GstDoc.direction == "in",
                                               GstDoc.period <= p)).all()
        claimed = _claimed_keys(s, tenant, c.id, p)
        by_key: dict[tuple, list[GstDoc]] = {}
        for b in books:
            by_key.setdefault((b.party_gstin, _norm_inv(b.number)), []).append(b)
        out = {"matched": [], "mismatch": [], "only_2b": [], "only_books": [], "ineligible": [], "rcm": []}
        used: set[str] = set()
        for r in twob:
            row = {"id": r.id, "supplier": r.trdnm or r.ctin, "ctin": r.ctin, "number": r.inum, "date": r.dt, "typ": r.typ,
                   "tax": round(r.igst + r.cgst + r.sgst, 2), "txval": r.txval, "ims": r.ims}
            if r.rev:
                out["rcm"].append(row | {"note": "Reverse charge: tax aap cash mein bharenge (3.1(d)), phir ITC."})
                continue
            if not r.itcavl:
                out["ineligible"].append(row | {"note": f"2B mein ITC nahi ({r.rsn or 'POS / 16(4)'})."})
                continue
            cand = by_key.get((r.ctin, _norm_inv(r.inum)), [])
            b = next((x for x in cand if x.id not in used), None)
            if not b:
                out["only_2b"].append(row | {"suggest": "pending", "note": "Books mein nahi: bill mangwaiye; asli nahi to IMS mein reject."})
                continue
            used.add(b.id)
            diff = round((r.igst + r.cgst + r.sgst) - (b.igst + b.cgst + b.sgst), 2)
            row |= {"doc_id": b.id, "books_tax": round(b.igst + b.cgst + b.sgst, 2), "diff": diff, "blocked": not b.itc}
            if abs(diff) <= tolerance:
                out["matched"].append(row | {"suggest": "accept"})
            else:
                out["mismatch"].append(row | {"suggest": "pending", "note": f"Tax mein {_rs(diff)} ka farak: supplier se sahi karwaiye."})
        twob_keys = {(r.ctin, _norm_inv(r.inum)) for r in twob}
        for b in books:
            if b.id in used or b.rcm or not b.party_gstin or (b.party_gstin, _norm_inv(b.number)) in twob_keys:
                continue
            if (b.party_gstin, _norm_inv(b.number)) in claimed:
                continue
            age = b.period < months[0]
            out["only_books"].append({"doc_id": b.id, "supplier": b.party_name or b.party_gstin, "ctin": b.party_gstin,
                                      "number": b.number, "date": b.doc_date, "tax": round(b.igst + b.cgst + b.sgst, 2), "older": bool(age),
                                      "note": "Supplier ne GSTR-1 mein nahi dala: ITC tab tak nahi (Sec 16(2)(aa)). Supplier ko yaad dilaiye."})
        name = c.name
    tot = lambda k: round(sum(x["tax"] for x in out[k]), 2)  # noqa: E731
    summary = {k: {"count": len(v), "tax": tot(k)} for k, v in out.items()}
    say = (f"{name} {R.nice_period(p)}: 2B se mile {summary['matched']['count']} ({_rs(summary['matched']['tax'])}), "
           f"farak {summary['mismatch']['count']}, sirf 2B mein {summary['only_2b']['count']}, sirf books mein "
           f"{summary['only_books']['count']} ({_rs(summary['only_books']['tax'])} ITC atka).")
    return {"client": name, "period": p, "summary": summary, **out, "say": say,
            "ims_note": "IMS: jo accept karenge woh 2B mein aata hai; reject / pending par ITC nahi. 14 tareekh tak action, "
                        "baad mein badla to 2B recompute karna padta hai. Portal par action aapki haan ke baad."}


def _claimed_keys(s, tenant: str, client_id: str, before: str) -> set:
    """Purchase bills whose ITC was already taken in an earlier filed 3B (so they don't show again as missing)."""
    out = set()
    for r in s.scalars(select(GstReturn).where(GstReturn.tenant_id == tenant, GstReturn.client_id == client_id,
                                               GstReturn.kind == "gstr3b", GstReturn.status == "filed", GstReturn.period < before)).all():
        for k in _j(r.data, {}).get("claimed", []):
            out.add((k[0], k[1]))
    return out


def ims_decide(tenant: str, client: str, period: str, decisions: dict | list, by: str = "owner") -> dict[str, Any]:
    """Her decision per 2B row: accept | reject | pending (Hoola only suggests). Pending is allowed for one period only."""
    p = R.norm_period(period)
    if isinstance(decisions, list):
        decisions = {str(x.get("id")): str(x.get("action")) for x in decisions if isinstance(x, dict)}
    n = 0
    with session() as s:
        c = _client(s, tenant, client)
        for rid, act in (decisions or {}).items():
            act = str(act).lower()
            if act not in ("accept", "reject", "pending", ""):
                raise GstError("IMS action: accept | reject | pending")
            r = s.get(Gst2B, rid)
            if r and r.tenant_id == tenant and r.client_id == c.id:
                r.ims = act
                n += 1
    return {"updated": n, "say": f"{n} bill par aapka IMS faisla note kiya. 'IMS portal par lagao' kahte hi Hoola portal par kar dega."}


def fetch_2b(tenant: str, client: str, period: str, conv: str = "", by: str = "owner") -> dict[str, Any]:
    """The Computer Operator downloads GSTR-2B (JSON) from the portal; the file comes back in chat to import."""
    from ..operator import service as op
    p = R.norm_period(period)
    with session() as s:
        c = _client(s, tenant, client)
        who = f"GSTIN {c.gstin}" + (f" (username {c.portal_user})" if c.portal_user else "")
    task = (f"On the GST portal log in for {who}; the owner types the password; read the CAPTCHA and show it to the owner to "
            f"confirm. Go to Services > Returns > Returns Dashboard, choose {R.fy_of(p)} / {R.nice_period(p)}, open Auto-drafted "
            f"ITC Statement GSTR-2B > View, then 'Download GSTR-2B details (JSON)'. Send the downloaded file to the owner. If "
            f"GSTR-2B is not generated yet (before the 14th), say so and stop. Do not change or file anything. Never pay.")
    try:
        job = op.start(tenant, task, PORTAL, [], conv, created_by=by)
    except op.OperatorJobError as e:
        raise GstError(f"Computer Operator: {e}")
    return {"operator": job, "say": f"Hoola portal se {c.gstin} ka 2B ({R.nice_period(p)}) la raha hai. File aate hi import kar dunga."}


def ims_apply(tenant: str, client: str, period: str, conv: str = "", by: str = "owner") -> dict[str, Any]:
    """Her IMS decisions put on the portal by the Computer Operator (she already decided; it stops before anything final)."""
    from ..operator import service as op
    p = R.norm_period(period)
    with session() as s:
        c = _client(s, tenant, client)
        rows = s.scalars(select(Gst2B).where(Gst2B.tenant_id == tenant, Gst2B.client_id == c.id, Gst2B.period == p,
                                             Gst2B.ims.in_(("accept", "reject", "pending")))).all()
        if not rows:
            raise GstError("Abhi koi IMS faisla nahi: pehle batayiye kaun accept / reject / pending.")
        lines = "; ".join(f"{x.ims.upper()}: supplier {x.ctin} invoice {x.inum}" for x in rows[:60])
        who = f"GSTIN {c.gstin}" + (f" (username {c.portal_user})" if c.portal_user else "")
    task = (f"On the GST portal log in for {who}; the owner types the password; read the CAPTCHA and show it to the owner to confirm. "
            f"Open Services > Returns > Invoice Management System (IMS) > Inward Supplies > 'All other ITC'. For each record set "
            f"the action exactly as listed and press SAVE: {lines}. Leave every other record as it is. After saving, if the "
            f"return period's GSTR-2B was already generated, tell the owner it needs 'Compute GSTR-2B' again and ask before "
            f"clicking it. Do not file any return. Never pay.")
    try:
        job = op.start(tenant, task, PORTAL, [], conv, created_by=by)
    except op.OperatorJobError as e:
        raise GstError(f"Computer Operator: {e}")
    return {"operator": job, "say": f"IMS par {len(rows)} faisle Hoola portal par laga raha hai (password aap daalengi)."}


def supplier_drafts(tenant: str, client: str, period: str) -> dict[str, Any]:
    """WhatsApp / email drafts to suppliers whose bills are missing or wrong in GSTR-2B (she sends them)."""
    rc = reconcile(tenant, client, period)
    by_sup: dict[str, list] = {}
    for x in rc["only_books"] + rc["mismatch"]:
        by_sup.setdefault(x["ctin"], []).append(x)
    drafts = []
    for ctin, xs in by_sup.items():
        bills = ", ".join(f"{x['number']} ({x['date']})" for x in xs[:12])
        drafts.append({"supplier": xs[0]["supplier"], "ctin": ctin,
                       "text": (f"Namaste, {rc['client']} (GST) ki taraf se. Aapke bill {bills} hamare GSTR-2B mein "
                                f"{'nahi dikh rahe' if any('doc_id' in x and 'diff' not in x for x in xs) else 'alag amount ke saath dikh rahe'}. "
                                f"Kripya GSTR-1 / IFF mein hamare GSTIN par sahi daal dijiye taaki ITC mil sake. Dhanyavaad.")})
    return {"drafts": drafts, "say": f"{len(drafts)} supplier ke liye message draft taiyaar (aap bhejengi)."}


# ---------------------------------------------------------------- GSTR-1
def _sum_items(docs_: list[GstDoc]) -> dict:
    return {k: round(sum(getattr(d, k) * (-1 if d.kind == "cn" else 1) for d in docs_), 2) for k in ("txval", "igst", "cgst", "sgst", "cess")}


def _itm_det(it: dict, inter: bool) -> dict:
    d = {"txval": round(_f(it["txval"]), 2), "rt": _f(it["rate"])}
    if inter:
        d["iamt"] = round(_f(it["igst"]), 2)
    else:
        d["camt"], d["samt"] = round(_f(it["cgst"]), 2), round(_f(it["sgst"]), 2)
    d["csamt"] = round(_f(it.get("cess")), 2)
    return d


def _by_rate(items: list[dict]) -> list[dict]:
    g: dict[float, dict] = {}
    for it in items:
        r = _f(it["rate"])
        x = g.setdefault(r, {"rate": r, "txval": 0.0, "igst": 0.0, "cgst": 0.0, "sgst": 0.0, "cess": 0.0})
        for k in ("txval", "igst", "cgst", "sgst", "cess"):
            x[k] += _f(it.get(k))
    return list(g.values())


def gstr1_json(c: GstClient, period: str, out_docs: list[GstDoc]) -> tuple[dict, dict, list[str]]:
    """The GSTR-1 upload JSON (portal offline schema keys), a readable summary and the problems still open."""
    b2b: dict[str, list] = {}
    b2cl: dict[str, list] = {}
    b2cs: dict[tuple, dict] = {}
    exp: dict[str, list] = {}
    cdnr: dict[str, list] = {}
    cdnur: list = []
    nil = {k: {"sply_ty": k, "nil_amt": 0.0, "expt_amt": 0.0, "ngsup_amt": 0.0} for k in ("INTRB2B", "INTRAB2B", "INTRB2C", "INTRAB2C")}
    hsn: dict[str, dict[tuple, dict]] = {"b2b": {}, "b2c": {}}
    docnums: dict[int, list[str]] = {1: [], 4: [], 5: []}
    problems: list[str] = []
    for d in sorted(out_docs, key=lambda x: (x.doc_date, x.number)):
        its = _j(d.items, [])
        reg = bool(d.party_gstin)
        inter = not R.intra(c.state, d.pos, d.kind == "sez") if d.kind != "exp" else True
        problems += [f"{d.number}: {p}" for p in _j(d.problems, [])]
        docnums[5 if d.kind == "cn" else 4 if d.kind == "dn" else 1].append(d.number)
        taxable = [i for i in its if _f(i["rate"]) > 0 or d.kind in ("exp", "sez")]
        zero = [i for i in its if _f(i["rate"]) == 0 and d.kind not in ("exp", "sez")]
        for i in zero:
            nk = ("INTRB2B" if inter else "INTRAB2B") if reg else ("INTRB2C" if inter else "INTRAB2C")
            nil[nk]["nil_amt"] += _f(i["txval"]) * (-1 if d.kind == "cn" else 1)
        for i in its:
            key = (re.sub(r"\D", "", i["hsn"]) or "", _f(i["rate"]), i.get("uqc") or "NOS")
            hrow = hsn["b2b" if reg else "b2c"].setdefault(key, {"hsn_sc": key[0], "desc": (i.get("desc") or "")[:30], "uqc": key[2],
                                                                 "qty": 0.0, "rt": key[1], "txval": 0.0, "iamt": 0.0, "camt": 0.0,
                                                                 "samt": 0.0, "csamt": 0.0})
            sg = -1 if d.kind == "cn" else 1
            hrow["qty"] += _f(i.get("qty")) * sg
            for a, b in (("txval", "txval"), ("iamt", "igst"), ("camt", "cgst"), ("samt", "sgst"), ("csamt", "cess")):
                hrow[a] += _f(i.get(b)) * sg
        if not taxable:
            continue
        itms = [{"num": n + 1, "itm_det": _itm_det(g, inter)} for n, g in enumerate(_by_rate(taxable))]
        if d.kind in ("cn", "dn"):
            note = {"ntty": "C" if d.kind == "cn" else "D", "nt_num": d.number, "nt_dt": _dmy(d.doc_date), "val": round(d.value, 2),
                    "pos": d.pos, "rchrg": "Y" if d.rcm else "N", "inv_typ": "R", "itms": itms}
            if reg:
                cdnr.setdefault(d.party_gstin, []).append(note)
            elif inter and d.value > R.B2CL_LIMIT:
                cdnur.append({"typ": "B2CL", **{k: v for k, v in note.items() if k not in ("rchrg", "inv_typ")}})
            else:
                for g in _by_rate(taxable):
                    k = ("INTER" if inter else "INTRA", d.pos, _f(g["rate"]))
                    x = b2cs.setdefault(k, {"sply_ty": k[0], "pos": k[1], "typ": "OE", "rt": k[2], "txval": 0.0, "iamt": 0.0,
                                            "camt": 0.0, "samt": 0.0, "csamt": 0.0})
                    sg = -1 if d.kind == "cn" else 1
                    for a, b in (("txval", "txval"), ("iamt", "igst"), ("camt", "cgst"), ("samt", "sgst"), ("csamt", "cess")):
                        x[a] += _f(g[b]) * sg
            continue
        inv = {"inum": d.number, "idt": _dmy(d.doc_date), "val": round(d.value, 2)}
        if d.kind == "exp":
            wpay = any(_f(i.get("igst")) > 0 for i in its)
            exp.setdefault("WPAY" if wpay else "WOPAY", []).append(
                inv | {"itms": [{"txval": round(_f(g["txval"]), 2), "rt": _f(g["rate"]), "iamt": round(_f(g["igst"]), 2),
                                 "csamt": round(_f(g["cess"]), 2)} for g in _by_rate(its)]})
        elif reg:
            b2b.setdefault(d.party_gstin, []).append(inv | {"pos": d.pos, "rchrg": "Y" if d.rcm else "N",
                                                            "inv_typ": ("SEWP" if any(_f(i.get("igst")) for i in its) else "SEWOP") if d.kind == "sez" else "R",
                                                            "itms": itms})
        elif inter and d.value > R.B2CL_LIMIT:
            b2cl.setdefault(d.pos, []).append(inv | {"itms": itms})
        else:
            for g in _by_rate(taxable):
                k = ("INTER" if inter else "INTRA", d.pos, _f(g["rate"]))
                x = b2cs.setdefault(k, {"sply_ty": k[0], "pos": k[1], "typ": "OE", "rt": k[2], "txval": 0.0, "iamt": 0.0,
                                        "camt": 0.0, "samt": 0.0, "csamt": 0.0})
                for a, b in (("txval", "txval"), ("iamt", "igst"), ("camt", "cgst"), ("samt", "sgst"), ("csamt", "cess")):
                    x[a] += _f(g[b])
    rnd = lambda x: {k: (round(v, 2) if isinstance(v, float) else v) for k, v in x.items()}  # noqa: E731
    j: dict[str, Any] = {"gstin": c.gstin, "fp": R.fp(period), "version": "GST2.2", "hash": "hash"}
    if b2b:
        j["b2b"] = [{"ctin": k, "inv": v} for k, v in sorted(b2b.items())]
    if b2cl:
        j["b2cl"] = [{"pos": k, "inv": v} for k, v in sorted(b2cl.items())]
    if b2cs:
        rows = []
        for x in b2cs.values():
            x = rnd(x)
            if x["sply_ty"] == "INTER":
                x.pop("camt"), x.pop("samt")
            else:
                x.pop("iamt")
            rows.append(x)
        j["b2cs"] = rows
    if exp:
        j["exp"] = [{"exp_typ": k, "inv": v} for k, v in exp.items()]
    if cdnr:
        j["cdnr"] = [{"ctin": k, "nt": v} for k, v in sorted(cdnr.items())]
    if cdnur:
        j["cdnur"] = cdnur
    nil_rows = [rnd(v) for v in nil.values() if v["nil_amt"] or v["expt_amt"] or v["ngsup_amt"]]
    if nil_rows:
        j["nil"] = {"inv": nil_rows}
    hj = {}
    for part in ("b2b", "b2c"):
        rows = []
        for n, x in enumerate(hsn[part].values(), 1):
            if not x["hsn_sc"]:
                problems.append(f"HSN khaali line ({part.upper()}): Table 12 mein nahi jayegi, HSN daaliye.")
                continue
            rows.append({"num": n, **rnd(x)})
        if rows:
            hj["hsn_" + part] = rows
    if hj:
        j["hsn"] = hj
    det = []
    for dn, nums in docnums.items():
        if nums:
            nums_sorted = sorted(nums, key=lambda s_: (len(s_), s_))
            det.append({"doc_num": dn, "docs": [{"num": 1, "from": nums_sorted[0], "to": nums_sorted[-1], "totnum": len(nums),
                                                  "cancel": 0, "net_issue": len(nums)}]})
    if det:
        j["doc_issue"] = {"doc_det": det}
    t = _sum_items([d for d in out_docs])
    summary = {"b2b_parties": len(b2b), "b2b_invoices": sum(len(v) for v in b2b.values()),
               "b2cl_invoices": sum(len(v) for v in b2cl.values()), "b2cs_rows": len(b2cs), "exports": sum(len(v) for v in exp.values()),
               "notes": sum(len(v) for v in cdnr.values()) + len(cdnur), "documents": sum(len(v) for v in docnums.values()),
               "txval": t["txval"], "igst": t["igst"], "cgst": t["cgst"], "sgst": t["sgst"], "cess": t["cess"],
               "tax": round(t["igst"] + t["cgst"] + t["sgst"] + t["cess"], 2), "nil": not out_docs}
    return j, summary, problems


def _ret(s, tenant: str, client_id: str, period: str, kind: str) -> GstReturn:
    r = s.scalars(select(GstReturn).where(GstReturn.tenant_id == tenant, GstReturn.client_id == client_id, GstReturn.period == period,
                                          GstReturn.kind == kind)).first()
    if not r:
        r = GstReturn(tenant_id=tenant, client_id=client_id, period=period, kind=kind)
        s.add(r)
        s.flush()
    return r


def _ret_view(r: GstReturn, c: GstClient | None = None) -> dict[str, Any]:
    d = _j(r.data, {})
    due = None
    if c:
        try:
            due = R.due_date(r.kind, r.period, c.scheme, c.state).isoformat()
        except ValueError:
            due = None
    return {"id": r.id, "client_id": r.client_id, "client": c.name if c else "", "gstin": c.gstin if c else "", "period": r.period,
            "period_nice": R.nice_period(r.period), "kind": r.kind, "label": RETURN_LABEL.get(r.kind, r.kind), "status": r.status,
            "status_label": STATUS_LABEL.get(r.status, r.status), "summary": d.get("summary", {}), "tables": d.get("tables", {}),
            "setoff": d.get("setoff", {}), "checks": _j(r.checks, []), "due": due, "arn": r.arn, "filed_on": r.filed_on,
            "operator_task": r.operator_task, "has_json": bool(r.portal_json), "log": _j(r.log, [])[-20:]}


def prepare_gstr1(tenant: str, client: str, period: str, conv: str = "") -> dict[str, Any]:
    p = R.norm_period(period)
    with _LOCK, session() as s:
        c = _client(s, tenant, client)
        if c.scheme == "composition":
            raise GstError("Composition client GSTR-1 nahi bharta: CMP-08 (har quarter) aur GSTR-4 (saal).")
        if c.scheme == "qrmp" and not R.is_quarter_end(p):
            label = "IFF"
        else:
            label = "GSTR-1"
        months = _months(c, p)
        out_docs = s.scalars(select(GstDoc).where(GstDoc.tenant_id == tenant, GstDoc.client_id == c.id, GstDoc.direction == "out",
                                                  GstDoc.period.in_(months))).all()
        if label == "IFF":
            out_docs = [d for d in out_docs if d.party_gstin]
        j, summary, problems = gstr1_json(c, p, out_docs)
        on = today(tenant)
        checks = list(problems)
        due = R.due_date("iff" if label == "IFF" else "gstr1", p, c.scheme, c.state)
        if R.time_barred("gstr1", p, c.scheme, c.state, on):
            checks.insert(0, "3 saal ki seema nikal chuki: yeh return portal ab nahi lega.")
        late = (on - due).days
        if late > 0 and label != "IFF":
            checks.append(f"{late} din late: late fee {_rs(R.late_fee('gstr1', late, summary['nil'], c.aato))} (GSTR-3B ke saath bharni hogi).")
        r = _ret(s, tenant, c.id, p, "gstr1")
        if r.status == "filed":
            raise GstError(f"{c.name} ka {label} {R.nice_period(p)} file ho chuka (ARN {r.arn}). Sudhaar GSTR-1A / agle mahine ke amendment mein.")
        r.data = json.dumps({"summary": summary | {"label": label, "late_days": max(0, late)}}, ensure_ascii=False)
        r.portal_json = json.dumps(j, ensure_ascii=False, separators=(",", ":"))
        r.checks = json.dumps(checks, ensure_ascii=False)
        r.status = "ready"
        _log(r, f"{label} taiyaar: {summary['b2b_invoices']} B2B, tax {_rs(summary['tax'])}, {len(checks)} dhyaan")
        v = _ret_view(r, c)
    say = (f"{v['client']} {label} {v['period_nice']} taiyaar: B2B {summary['b2b_invoices']} bill ({summary['b2b_parties']} party), "
           f"B2C {summary['b2cs_rows']} line, notes {summary['notes']}; taxable {_rs(summary['txval'])}, tax {_rs(summary['tax'])}.")
    if summary["nil"]:
        say += " Koi sale nahi: NIL return."
    if checks:
        say += f" {len(checks)} cheezein dekhni hain: " + " ".join(checks[:3])
    say += f" Due {_nice_date(due)}. File karne se pehle aapki haan."
    _publish(tenant, "gst_return", {"return_id": v["id"], "status": v["status"]}, conv)
    return {"return": v, "say": say}


def _nice_date(d: date) -> str:
    return d.strftime("%d %b %Y")


# ---------------------------------------------------------------- GSTR-3B
def prepare_gstr3b(tenant: str, client: str, period: str, other_reversal=None, conv: str = "") -> dict[str, Any]:
    p = R.norm_period(period)
    with _LOCK, session() as s:
        c = _client(s, tenant, client)
        if c.scheme == "composition":
            raise GstError("Composition client 3B nahi bharta: CMP-08.")
        if c.scheme == "qrmp" and not R.is_quarter_end(p):
            raise GstError(f"QRMP client: {R.nice_period(p)} mein 3B nahi, PMT-06 se tax (25 tareekh tak). 3B quarter ke end par.")
        months = _months(c, p)
        allq = s.scalars(select(GstDoc).where(GstDoc.tenant_id == tenant, GstDoc.client_id == c.id, GstDoc.period.in_(months))).all()
        out = [d for d in allq if d.direction == "out"]
        inn = [d for d in allq if d.direction == "in"]
        sg = lambda d: -1 if d.kind == "cn" else 1  # noqa: E731
        t = {"3.1a": {"txval": 0.0, "igst": 0.0, "cgst": 0.0, "sgst": 0.0, "cess": 0.0}, "3.1b": {"txval": 0.0, "igst": 0.0, "cess": 0.0},
             "3.1c": {"txval": 0.0}, "3.1d": {"txval": 0.0, "igst": 0.0, "cgst": 0.0, "sgst": 0.0, "cess": 0.0}, "3.1e": {"txval": 0.0}}
        unreg_inter: dict[str, dict] = {}
        for d in out:
            for i in _j(d.items, []):
                k = "3.1b" if d.kind in ("exp", "sez") else ("3.1c" if _f(i["rate"]) == 0 else "3.1a")
                for f in t[k]:
                    t[k][f] += _f(i.get("txval" if f == "txval" else f)) * sg(d)
            if not d.party_gstin and not R.intra(c.state, d.pos) and d.kind not in ("exp",):
                u = unreg_inter.setdefault(d.pos, {"pos": d.pos, "txval": 0.0, "igst": 0.0})
                u["txval"] += d.txval * sg(d)
                u["igst"] += d.igst * sg(d)
        rcm_tax = {"igst": 0.0, "cgst": 0.0, "sgst": 0.0, "cess": 0.0}
        for d in inn:
            if d.rcm:
                for f in ("txval", "igst", "cgst", "sgst", "cess"):
                    t["3.1d"][f] += getattr(d, f) * sg(d)
                for f in rcm_tax:
                    rcm_tax[f] += getattr(d, f) * sg(d) if d.itc else 0.0
        # ITC: only what GSTR-2B shows (Sec 16(2)(aa), Rule 36(4)), minus her IMS reject / pending, minus 17(5) blocked
        twob = s.scalars(select(Gst2B).where(Gst2B.tenant_id == tenant, Gst2B.client_id == c.id, Gst2B.period.in_(months))).all()
        books_idx = {(b.party_gstin, _norm_inv(b.number)): b for b in s.scalars(select(GstDoc).where(
            GstDoc.tenant_id == tenant, GstDoc.client_id == c.id, GstDoc.direction == "in")).all()}
        itc = {"igst": 0.0, "cgst": 0.0, "sgst": 0.0, "cess": 0.0}
        blocked = {"igst": 0.0, "cgst": 0.0, "sgst": 0.0, "cess": 0.0}
        claimed, skipped_ims = [], 0
        for r in twob:
            if r.rev or not r.itcavl:
                continue
            if r.ims in ("reject", "pending"):
                skipped_ims += 1
                continue
            b = books_idx.get((r.ctin, _norm_inv(r.inum)))
            sign = -1 if r.typ == "cn" else 1
            for f in itc:
                itc[f] += getattr(r, f) * sign
            if b is not None and not b.itc:
                for f in blocked:
                    blocked[f] += getattr(r, f) * sign
            claimed.append([r.ctin, _norm_inv(r.inum)])
        on = today(tenant)
        rule37 = {"igst": 0.0, "cgst": 0.0, "sgst": 0.0, "cess": 0.0}
        r37_bills = []
        for b in books_idx.values():
            bd = parse_date(b.doc_date)
            if not bd or b.rcm or not b.itc or b.paid_on:
                continue
            limit = bd + timedelta(days=R.RULE37_DAYS)
            if R.period_of(limit) in months:
                r37_bills.append(f"{b.party_name or b.party_gstin} {b.number}")
                for f in rule37:
                    rule37[f] += getattr(b, f)
        other = {f: _f((other_reversal or {}).get(f)) for f in ("igst", "cgst", "sgst", "cess")}
        itc4a = {f: round(itc[f] + rcm_tax[f], 2) for f in itc}
        rev4b1 = {f: round(blocked[f], 2) for f in itc}
        rev4b2 = {f: round(rule37[f] + other[f], 2) for f in itc}
        net = {f: round(itc4a[f] - rev4b1[f] - rev4b2[f], 2) for f in itc}
        liab_fwd = {f: round(t["3.1a"].get(f, 0.0) + t["3.1b"].get(f, 0.0), 2) for f in ("igst", "cgst", "sgst", "cess")}
        so = R.set_off(liab_fwd, net)
        cash = {f: round(so["cash"][f] + t["3.1d"][f], 2) for f in ("igst", "cgst", "sgst", "cess")}   # RCM always in cash
        cash_total = round(sum(cash.values()), 2)
        due = R.due_date("gstr3b", p, c.scheme, c.state)
        late = max(0, (on - due).days)
        nil = not out and not inn
        int_amt = R.interest(cash_total, late)
        fee = R.late_fee("gstr3b", late, nil, c.aato)
        checks = []
        if R.time_barred("gstr3b", p, c.scheme, c.state, on):
            checks.append("3 saal ki seema nikal chuki: yeh 3B portal ab nahi lega.")
        g1 = s.scalars(select(GstReturn).where(GstReturn.tenant_id == tenant, GstReturn.client_id == c.id, GstReturn.period == p,
                                               GstReturn.kind == "gstr1")).first()
        if not g1 or g1.status != "filed":
            checks.append("GSTR-1 pehle file hona chahiye: 3B ki Table 3.1 usi se lock hokar aati hai (July 2025 se).")
        else:
            g1s = _j(g1.data, {}).get("summary", {})
            d1 = round((g1s.get("tax") or 0) - (t["3.1a"]["igst"] + t["3.1a"]["cgst"] + t["3.1a"]["sgst"] + t["3.1a"]["cess"] + t["3.1b"]["igst"]), 2)
            if abs(d1) > 10:
                checks.append(f"GSTR-1 aur books mein {_rs(d1)} ka farak: 3B se pehle GSTR-1A mein sudhaariye (DRC-01B ka khatra).")
        if skipped_ims:
            checks.append(f"{skipped_ims} bill IMS mein reject / pending: unka ITC is mahine nahi liya.")
        if r37_bills:
            checks.append(f"180 din se supplier ko payment nahi ({len(r37_bills)} bill): ITC wapas (4B(2)) + byaaj; payment par wapas le sakte.")
        if any(rev4b1.values()):
            checks.append(f"Blocked credit 17(5) {_rs(sum(rev4b1.values()))} 4B(1) mein wapas.")
        if not twob and inn:
            checks.append("2B nahi aaya: ITC zero rakha (sirf 2B wala ITC milta hai). 2B bhejiye.")
        if late:
            checks.append(f"{late} din late: byaaj {_rs(int_amt)} + late fee {_rs(fee)}.")
        if p >= "2026-07":
            checks.append("July 2026 se 3B ka ITC (4A) bhi 2B se auto aata hai: portal ka figure aur yeh hisaab milaiye.")
        rr = _ret(s, tenant, c.id, p, "gstr3b")
        if rr.status == "filed":
            raise GstError(f"{c.name} ka GSTR-3B {R.nice_period(p)} file ho chuka (ARN {rr.arn}).")
        tables = {k: {f: round(v, 2) for f, v in x.items()} for k, x in t.items()}
        tables["3.2"] = [{k: (round(v, 2) if isinstance(v, float) else v) for k, v in u.items()} for u in unreg_inter.values()]
        tables["4A"] = itc4a
        tables["4A_rcm"] = {f: round(v, 2) for f, v in rcm_tax.items()}
        tables["4B1"] = rev4b1
        tables["4B2"] = rev4b2
        tables["4C"] = net
        tables["5.1"] = {"interest": int_amt, "late_fee": fee}
        summary = {"liability": round(sum(liab_fwd.values()) + sum(t["3.1d"][f] for f in ("igst", "cgst", "sgst", "cess")), 2),
                   "itc": round(sum(net.values()), 2), "cash": cash_total, "cash_split": cash, "interest": int_amt, "late_fee": fee,
                   "pay_total": round(cash_total + int_amt + fee, 2), "late_days": late, "nil": nil, "rule37": r37_bills[:20]}
        rr.data = json.dumps({"summary": summary, "tables": tables, "setoff": so, "claimed": claimed}, ensure_ascii=False)
        rr.checks = json.dumps(checks, ensure_ascii=False)
        rr.status = "ready"
        _log(rr, f"3B taiyaar: tax {_rs(summary['liability'])}, ITC {_rs(summary['itc'])}, cash {_rs(cash_total)}")
        v = _ret_view(rr, c)
    say = (f"{v['client']} GSTR-3B {v['period_nice']}: tax {_rs(summary['liability'])}, ITC {_rs(summary['itc'])}, "
           f"cash mein bharna {_rs(summary['pay_total'])}"
           + (f" (byaaj {_rs(int_amt)}, late fee {_rs(fee)})" if late else "") + f". Due {_nice_date(due)}.")
    if checks:
        say += " Dhyaan: " + " ".join(checks[:3])
    say += " File karne se pehle aapki haan; tax ka challan client / aap khud bharenge."
    _publish(tenant, "gst_return", {"return_id": v["id"], "status": v["status"]}, conv)
    return {"return": v, "say": say}


def prepare_cmp08(tenant: str, client: str, period: str, turnover=None, conv: str = "") -> dict[str, Any]:
    """Composition: tax on the quarter's turnover (1% traders / manufacturers, 5% restaurants, 6% services opt-in)."""
    p = R.norm_period(period)
    with _LOCK, session() as s:
        c = _client(s, tenant, client)
        if c.scheme != "composition":
            raise GstError("Yeh client composition mein nahi.")
        months = R.quarter_months(p)
        out = s.scalars(select(GstDoc).where(GstDoc.tenant_id == tenant, GstDoc.client_id == c.id, GstDoc.direction == "out",
                                             GstDoc.period.in_(months))).all()
        to = _f(turnover) if turnover not in (None, "") else sum(d.txval * (-1 if d.kind == "cn" else 1) for d in out)
        rate = 1.0
        nt = c.notes.lower()
        if "restaurant" in nt or "hotel" in nt:
            rate = 5.0
        elif "service" in nt:
            rate = 6.0
        tax = round(to * rate / 100, 2)
        rr = _ret(s, tenant, c.id, months[-1], "cmp08")
        rr.data = json.dumps({"summary": {"turnover": round(to, 2), "rate": rate, "tax": tax, "cgst": round(tax / 2, 2),
                                          "sgst": round(tax - round(tax / 2, 2), 2)}}, ensure_ascii=False)
        rr.status = "ready"
        _log(rr, f"CMP-08: turnover {_rs(to)}, {rate:g}% = {_rs(tax)}")
        v = _ret_view(rr, c)
    _publish(tenant, "gst_return", {"return_id": v["id"], "status": v["status"]}, conv)
    return {"return": v, "say": f"{v['client']} CMP-08 {v['period_nice']}: turnover {_rs(to)}, tax {rate:g}% = {_rs(tax)}. Due {v['due']}."}


# ---------------------------------------------------------------- her Haan, filing, ARN
def _get_ret(s, tenant: str, ref: str) -> GstReturn:
    r = s.get(GstReturn, str(ref or ""))
    if not r or r.tenant_id != tenant:
        raise GstError("Return nahi mila.")
    return r


def get_return(tenant: str, ref: str) -> dict[str, Any]:
    with session() as s:
        r = _get_ret(s, tenant, ref)
        return _ret_view(r, s.get(GstClient, r.client_id))


def find_return(tenant: str, client: str, period: str, kind: str) -> dict[str, Any]:
    with session() as s:
        c = _client(s, tenant, client)
        r = s.scalars(select(GstReturn).where(GstReturn.tenant_id == tenant, GstReturn.client_id == c.id,
                                              GstReturn.period == R.norm_period(period), GstReturn.kind == kind)).first()
        if not r:
            raise GstError(f"{c.name} ka {RETURN_LABEL.get(kind, kind)} {R.nice_period(period)} abhi bana nahi.")
        return _ret_view(r, c)


def decide(tenant: str, ref: str, approve: bool, note: str = "", conv: str = "") -> dict[str, Any]:
    with _LOCK, session() as s:
        r = _get_ret(s, tenant, ref)
        if r.status not in ("ready", "approved", "rejected"):
            raise GstError(f"Is return par abhi haan / na nahi ho sakti ({STATUS_LABEL.get(r.status)}).")
        r.status = "approved" if approve else "rejected"
        r.approved_at = now() if approve else None
        _log(r, ("Haan: file karo" if approve else "Mana") + (f" ({note[:100]})" if note else ""), "owner")
        c = s.get(GstClient, r.client_id)
        if approve:
            _learn_from(s, tenant, r)
        v = _ret_view(r, c)
    _publish(tenant, "gst_return", {"return_id": v["id"], "status": v["status"]}, conv)
    return {"return": v, "say": "Haan mili. 'file' bolte hi Hoola portal par taiyaar kar dega." if approve else "Theek, return roka."}


def _learn_from(s, tenant: str, r: GstReturn) -> None:
    c = s.get(GstClient, r.client_id)
    months = _months(c, r.period) if c else [r.period]
    for d in s.scalars(select(GstDoc).where(GstDoc.tenant_id == tenant, GstDoc.client_id == r.client_id, GstDoc.period.in_(months))).all():
        for it in _j(d.items, []):
            if it.get("desc") and it.get("hsn"):
                _learn(s, tenant, r.client_id, "item:" + _words(it["desc"]), {"hsn": it["hsn"], "rate": _f(it["rate"]), "uqc": it.get("uqc", "")})
        if d.direction == "in" and d.party_gstin and not d.itc:
            _learn(s, tenant, r.client_id, "supplier:" + d.party_gstin, {"blocked": True, "name": d.party_name})


def return_file(tenant: str, ref: str, conv: str = "") -> dict[str, Any]:
    """The GSTR-1 JSON as a chat file (to keep / upload by hand) + a readable Excel of the same data."""
    from .. import chatfiles
    with session() as s:
        r = _get_ret(s, tenant, ref)
        c = s.get(GstClient, r.client_id)
        if r.kind != "gstr1" or not r.portal_json:
            raise GstError("JSON sirf GSTR-1 / IFF ke liye banta hai.")
        name = f"{c.gstin}_GSTR1_{R.fp(r.period)}.json"
        f = chatfiles.store(tenant, name, r.portal_json.encode("utf-8"), "application/json")
        j = json.loads(r.portal_json)
    x = chatfiles.store(tenant, name.replace(".json", ".xlsx"), _gstr1_xlsx(j), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    _publish(tenant, "pc_file", {**f, "source": "gst", "title": name}, conv)
    _publish(tenant, "pc_file", {**x, "source": "gst", "title": "GSTR-1 review sheet"}, conv)
    return {"file": f, "excel": x, "say": f"GSTR-1 JSON ({name}) aur review Excel taiyaar."}


def _gstr1_xlsx(j: dict) -> bytes:
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "b2b"
    ws.append(["GSTIN of recipient", "Invoice no", "Date", "Value", "POS", "RCM", "Type", "Rate", "Taxable", "IGST", "CGST", "SGST", "Cess"])
    for p in j.get("b2b", []):
        for inv in p["inv"]:
            for it in inv["itms"]:
                d = it["itm_det"]
                ws.append([p["ctin"], inv["inum"], inv["idt"], inv["val"], inv["pos"], inv["rchrg"], inv["inv_typ"], d["rt"], d["txval"],
                           d.get("iamt", 0), d.get("camt", 0), d.get("samt", 0), d.get("csamt", 0)])
    ws = wb.create_sheet("b2cl")
    ws.append(["Invoice no", "Date", "Value", "POS", "Rate", "Taxable", "IGST", "Cess"])
    for p in j.get("b2cl", []):
        for inv in p["inv"]:
            for it in inv["itms"]:
                d = it["itm_det"]
                ws.append([inv["inum"], inv["idt"], inv["val"], p["pos"], d["rt"], d["txval"], d.get("iamt", 0), d.get("csamt", 0)])
    ws = wb.create_sheet("b2cs")
    ws.append(["Type", "POS", "Rate", "Taxable", "IGST", "CGST", "SGST", "Cess"])
    for x in j.get("b2cs", []):
        ws.append([x["sply_ty"], x["pos"], x["rt"], x["txval"], x.get("iamt", 0), x.get("camt", 0), x.get("samt", 0), x.get("csamt", 0)])
    ws = wb.create_sheet("cdnr")
    ws.append(["GSTIN", "Note no", "Date", "C/D", "Value", "POS", "Rate", "Taxable", "IGST", "CGST", "SGST"])
    for p in j.get("cdnr", []):
        for nt in p["nt"]:
            for it in nt["itms"]:
                d = it["itm_det"]
                ws.append([p["ctin"], nt["nt_num"], nt["nt_dt"], nt["ntty"], nt["val"], nt["pos"], d["rt"], d["txval"],
                           d.get("iamt", 0), d.get("camt", 0), d.get("samt", 0)])
    for part in ("hsn_b2b", "hsn_b2c"):
        ws = wb.create_sheet(part)
        ws.append(["HSN", "Description", "UQC", "Qty", "Rate", "Taxable", "IGST", "CGST", "SGST", "Cess"])
        for x in (j.get("hsn") or {}).get(part, []):
            ws.append([x["hsn_sc"], x["desc"], x["uqc"], x["qty"], x["rt"], x["txval"], x["iamt"], x["camt"], x["samt"], x["csamt"]])
    ws = wb.create_sheet("docs")
    ws.append(["Doc type", "From", "To", "Total", "Cancelled", "Net"])
    for d in (j.get("doc_issue") or {}).get("doc_det", []):
        for x in d["docs"]:
            ws.append([{1: "Invoices", 4: "Debit notes", 5: "Credit notes"}.get(d["doc_num"], d["doc_num"]), x["from"], x["to"],
                       x["totnum"], x["cancel"], x["net_issue"]])
    bio = io.BytesIO()
    wb.save(bio)
    return bio.getvalue()


def _portal_task(c: GstClient, r: GstReturn) -> tuple[str, str]:
    who = (f"GST practitioner login (practitioner username), then 'View Taxpayers / Returns' for GSTIN {c.gstin}"
           if c.login_as == "practitioner" else f"taxpayer login for GSTIN {c.gstin}" + (f" (username {c.portal_user})" if c.portal_user else ""))
    per = f"{R.fy_of(r.period)} / {R.nice_period(r.period)}"
    d = _j(r.data, {})
    s_ = d.get("summary", {})
    if r.kind == "gstr1":
        label = s_.get("label", "GSTR-1")
        task = (f"On the GST portal: {who}. The owner types the password; read the CAPTCHA and show it to the owner to confirm. "
                f"Go to Services > Returns > Returns Dashboard, choose {per}, open {label} > 'Prepare Offline' > Upload, and upload "
                f"the attached JSON file. Wait until the upload status is 'Processed' (refresh every 30 seconds, up to 10 minutes). "
                f"If it shows 'Processed with Error', download the error report and stop. Then open 'Prepare Online', click "
                f"'Generate Summary' and read the summary: it should show {s_.get('b2b_invoices', 0)} B2B invoices and total tax "
                f"about Rs {s_.get('tax', 0)}. Tell the owner the portal's totals. STOP before 'File Statement' / 'Submit': the "
                f"taxpayer files with the EVC OTP on their registered mobile (or DSC). Never pay anything.")
    elif r.kind == "gstr3b":
        tb = d.get("tables", {})
        a4, b1, b2 = tb.get("4A", {}), tb.get("4B1", {}), tb.get("4B2", {})
        task = (f"On the GST portal: {who}. The owner types the password; read the CAPTCHA and show it to the owner to confirm. "
                f"Go to Services > Returns > Returns Dashboard, choose {per}, open GSTR-3B > 'Prepare Online'. Table 3.1 is "
                f"auto-filled from GSTR-1 and locked: read it and tell the owner if it differs from Hoola's figures "
                f"(3.1(a) taxable {tb.get('3.1a', {}).get('txval', 0)}, IGST {tb.get('3.1a', {}).get('igst', 0)}, CGST "
                f"{tb.get('3.1a', {}).get('cgst', 0)}, SGST {tb.get('3.1a', {}).get('sgst', 0)}). In Table 4 check 4A(5) "
                f"against Hoola's ITC (IGST {a4.get('igst', 0)}, CGST {a4.get('cgst', 0)}, SGST {a4.get('sgst', 0)}); enter 4B(1) "
                f"IGST {b1.get('igst', 0)} CGST {b1.get('cgst', 0)} SGST {b1.get('sgst', 0)} and 4B(2) IGST {b2.get('igst', 0)} "
                f"CGST {b2.get('cgst', 0)} SGST {b2.get('sgst', 0)}. Save. Click 'Proceed to payment' and read the cash "
                f"required (Hoola expects about Rs {s_.get('pay_total', 0)}). STOP there: the owner / taxpayer pays the "
                f"challan and files with EVC OTP or DSC. Never create or pay a challan, never press 'File'.")
    elif r.kind == "cmp08":
        task = (f"On the GST portal: {who}. The owner types the password; CAPTCHA shown to the owner to confirm. Returns "
                f"Dashboard > {per} > CMP-08 > Prepare Online. Enter outward turnover {s_.get('turnover', 0)} and tax CGST "
                f"{s_.get('cgst', 0)} SGST {s_.get('sgst', 0)}. Save and STOP before payment / 'File': the taxpayer pays and files "
                f"with EVC. Never pay.")
    else:
        raise GstError("Is return ki filing abhi Hoola nahi karta.")
    return task, PORTAL


def file_return(tenant: str, ref: str, conv: str = "", by: str = "owner") -> dict[str, Any]:
    """Only after her Haan: the Computer Operator logs in, uploads / fills and stops before the final submit + payment."""
    from .. import chatfiles
    from ..operator import service as op
    with session() as s:
        r = _get_ret(s, tenant, ref)
        c = s.get(GstClient, r.client_id)
        if r.status == "filed":
            raise GstError(f"Pehle hi file ho chuka (ARN {r.arn}).")
        if r.status != "approved":
            raise GstError("Pehle aapki haan chahiye: return dekh kar 'haan, file karo' kahiye.")
        task, url = _portal_task(c, r)
        fid = []
        if r.kind == "gstr1":
            f = chatfiles.store(tenant, f"{c.gstin}_GSTR1_{R.fp(r.period)}.json", r.portal_json.encode("utf-8"), "application/json")
            fid = [f["file_id"]]
    try:
        job = op.start(tenant, task, url, fid, conv, created_by=by)
    except op.OperatorJobError as e:
        raise GstError(f"Computer Operator: {e}")
    with _LOCK, session() as s:
        r = _get_ret(s, tenant, ref)
        r.status, r.operator_task = "filing", job["id"]
        _log(r, "Portal par taiyaar karne bheja (Computer Operator)", by)
        v = _ret_view(r, s.get(GstClient, r.client_id))
    _publish(tenant, "gst_return", {"return_id": v["id"], "status": v["status"]}, conv)
    return {"return": v, "operator": job,
            "say": "Hoola portal par return taiyaar kar raha hai: password aap daalengi, CAPTCHA main padhkar dikhaunga. "
                   "Final 'File' EVC OTP se aur tax ka challan aap / client khud karenge. ARN aane par mujhe bata dijiye."}


def mark_filed(tenant: str, ref: str, arn: str = "", filed_on: str = "", by: str = "owner", conv: str = "") -> dict[str, Any]:
    arn = re.sub(r"\s", "", str(arn or "")).upper()
    if arn and not re.fullmatch(r"[A-Z]{2}\d{2}[0-9A-Z]{9,13}", arn):
        raise GstError(f"ARN {arn} sahi nahi lagta (jaise AA2710260012345).")
    with _LOCK, session() as s:
        r = _get_ret(s, tenant, ref)
        d = parse_date(filed_on) or today(tenant)
        r.status, r.arn, r.filed_on = "filed", arn, d.isoformat()
        _log(r, f"File hua {d.isoformat()}" + (f", ARN {arn}" if arn else ""), by)
        _learn_from(s, tenant, r)
        c = s.get(GstClient, r.client_id)
        v = _ret_view(r, c)
    _publish(tenant, "gst_return", {"return_id": v["id"], "status": v["status"]}, conv)
    return {"return": v, "say": f"{v['client']} {v['label']} {v['period_nice']} file: {arn or 'ARN baad mein'}. Log mein likh diya."}


def returns(tenant: str, client: str = "", period: str = "", status: str = "", limit: int = 60) -> list[dict[str, Any]]:
    with session() as s:
        q = select(GstReturn).where(GstReturn.tenant_id == tenant)
        if client:
            q = q.where(GstReturn.client_id == _client(s, tenant, client).id)
        if period:
            q = q.where(GstReturn.period == R.norm_period(period))
        if status:
            q = q.where(GstReturn.status == status)
        rows = s.scalars(q.order_by(GstReturn.updated_at.desc()).limit(limit)).all()
        cs = {c.id: c for c in s.scalars(select(GstClient).where(GstClient.tenant_id == tenant)).all()}
        return [{k: v for k, v in _ret_view(r, cs.get(r.client_id)).items() if k not in ("tables", "setoff", "log")} for r in rows]


# ---------------------------------------------------------------- notices
def add_notice(tenant: str, client: str, form: str, issued_on: str = "", ref: str = "", text: str = "", amount=None,
               due_on: str = "") -> dict[str, Any]:
    f = re.sub(r"\s", "", str(form or "").upper()).replace("GST", "").strip("-")
    f = f if f in R.NOTICES else next((k for k in R.NOTICES if k.replace("-", "") == f.replace("-", "")), "OTHER")
    what, reply, days = R.NOTICES[f]
    iss = parse_date(issued_on)
    due = parse_date(due_on) or ((iss or date.today()) + timedelta(days=days))
    with session() as s:
        c = _client(s, tenant, client)
        n = GstNotice(tenant_id=tenant, client_id=c.id, form=f, ref=str(ref or "")[:60], issued_on=iss.isoformat() if iss else "",
                      due_on=due.isoformat(), text=str(text or "")[:20000], amount=_f(amount))
        s.add(n)
        s.flush()
        v = _notice_view(n, c)
    return {"notice": v, "say": f"{v['client']}: {f} ({what}). Jawab {reply} se, aakhri din {_nice_date(due)}. Jawab ka draft bana doon?"}


def _notice_view(n: GstNotice, c: GstClient | None) -> dict[str, Any]:
    what, reply, _ = R.NOTICES.get(n.form, R.NOTICES["OTHER"])
    return {"id": n.id, "client": c.name if c else "", "client_id": n.client_id, "form": n.form, "what": what, "reply_form": reply,
            "ref": n.ref, "issued_on": n.issued_on, "due_on": n.due_on, "amount": n.amount, "status": n.status, "reply": n.reply,
            "text": n.text[:2000]}


def notice_draft(tenant: str, notice_id: str) -> dict[str, Any]:
    """A reply skeleton with the client's own figures; she (or the model) completes and checks it before filing."""
    with session() as s:
        n = s.get(GstNotice, notice_id)
        if not n or n.tenant_id != tenant:
            raise GstError("Notice nahi mila.")
        c = s.get(GstClient, n.client_id)
        rets = s.scalars(select(GstReturn).where(GstReturn.tenant_id == tenant, GstReturn.client_id == c.id)
                         .order_by(GstReturn.period.desc()).limit(12)).all()
        facts = [f"{RETURN_LABEL.get(r.kind)} {R.nice_period(r.period)}: {STATUS_LABEL.get(r.status)}"
                 + (f", ARN {r.arn}" if r.arn else "") + (f", tax {_rs(_j(r.data, {}).get('summary', {}).get('tax') or _j(r.data, {}).get('summary', {}).get('liability') or 0)}") for r in rets]
        what, reply, _ = R.NOTICES.get(n.form, R.NOTICES["OTHER"])
        body = (f"To,\nThe Proper Officer,\n\nSubject: Reply to {n.form}{' ref. ' + n.ref if n.ref else ''} dated {n.issued_on or '___'} "
                f"in the case of {c.name} (GSTIN {c.gstin})\n\nRespected Sir / Madam,\n\nWith reference to the above notice regarding "
                f"{what.lower()}, we submit as under:\n\n1. Facts: [explain the difference / issue point by point]\n"
                f"2. Reconciliation: [figures as per returns vs as per notice]\n3. Documents attached: [invoices / 2B / ledgers]\n\n"
                f"We request that the proceedings be dropped / the difference be accepted as explained. We are available for a "
                f"personal hearing if needed.\n\nFor {c.name}\n\nAuthorised Signatory")
        n.reply = body
    return {"draft": body, "reply_form": reply, "facts": facts,
            "say": f"{n.form} ka jawab draft ({reply}) taiyaar: aankde aur kaagaz bharke dekhiye. Portal par jama karna aapki haan se."}


def notice_status(tenant: str, notice_id: str, status: str) -> dict[str, Any]:
    if status not in ("open", "replied", "closed"):
        raise GstError("status: open | replied | closed")
    with session() as s:
        n = s.get(GstNotice, notice_id)
        if not n or n.tenant_id != tenant:
            raise GstError("Notice nahi mila.")
        n.status = status
        return {"notice": _notice_view(n, s.get(GstClient, n.client_id))}


def notices(tenant: str, status: str = "open") -> list[dict[str, Any]]:
    with session() as s:
        q = select(GstNotice).where(GstNotice.tenant_id == tenant)
        if status:
            q = q.where(GstNotice.status == status)
        cs = {c.id: c for c in s.scalars(select(GstClient).where(GstClient.tenant_id == tenant)).all()}
        return [_notice_view(n, cs.get(n.client_id)) for n in s.scalars(q.order_by(GstNotice.due_on)).all()]


# ---------------------------------------------------------------- the calendar (what is due, late, at risk)
def calendar(tenant: str, ahead: int = 15) -> dict[str, Any]:
    on = today(tenant)
    cur = R.period_of(on)
    py, pm = R.add_months(on.year, on.month, -1)
    prev = f"{py:04d}-{pm:02d}"
    items, late = [], []
    with session() as s:
        cs = s.scalars(select(GstClient).where(GstClient.tenant_id == tenant, GstClient.active.is_(True))).all()
        rets = s.scalars(select(GstReturn).where(GstReturn.tenant_id == tenant)).all()
        status = {(r.client_id, r.period, r.kind): r for r in rets}
        for c in cs:
            for per in sorted({prev, f"{R.add_months(on.year, on.month, -2)[0]:04d}-{R.add_months(on.year, on.month, -2)[1]:02d}", cur}):
                for ob in R.obligations(per, c.scheme, c.state):
                    kind = ob["kind"]
                    rk = "gstr1" if kind in ("gstr1", "iff") else kind
                    r = status.get((c.id, per, rk))
                    if r and r.status == "filed":
                        continue
                    days = (ob["due"] - on).days
                    row = {"client": c.name, "client_id": c.id, "gstin": c.gstin, "period": per, "period_nice": R.nice_period(per),
                           "kind": kind, "label": ob["label"], "due": ob["due"].isoformat(), "due_nice": _nice_date(ob["due"]),
                           "days": days, "status": r.status if r else "", "return_id": r.id if r else ""}
                    if days < 0 and kind not in ("iff",):
                        nil = False
                        row["late_fee"] = R.late_fee(rk, -days, nil, c.aato)
                        late.append(row)
                    elif 0 <= days <= ahead:
                        items.append(row)
            fy_end_y = on.year if on.month >= 4 else on.year - 1
            itc_last = date(fy_end_y, 11, 30)
            if 0 <= (itc_last - on).days <= 45:
                items.append({"client": c.name, "client_id": c.id, "gstin": c.gstin, "kind": "itc_last",
                              "label": f"FY {fy_end_y - 1}-{str(fy_end_y)[2:]} ka bacha ITC lene ki aakhri tareekh (Sec 16(4))",
                              "due": itc_last.isoformat(), "due_nice": _nice_date(itc_last), "days": (itc_last - on).days})
            g9 = date(fy_end_y, 12, 31)
            if c.aato > R.GSTR9_LIMIT and 0 <= (g9 - on).days <= 75 and not status.get((c.id, f"{fy_end_y}-03", "gstr9")):
                items.append({"client": c.name, "client_id": c.id, "gstin": c.gstin, "kind": "gstr9",
                              "label": "GSTR-9" + (" + 9C" if c.aato > R.GSTR9C_LIMIT else "") + f" FY {fy_end_y - 1}-{str(fy_end_y)[2:]}",
                              "due": g9.isoformat(), "due_nice": _nice_date(g9), "days": (g9 - on).days})
        nts = [n for n in s.scalars(select(GstNotice).where(GstNotice.tenant_id == tenant, GstNotice.status == "open")).all()]
        cmap = {c.id: c for c in cs}
        for n in nts:
            dd = parse_date(n.due_on)
            if dd:
                row = {"client": cmap[n.client_id].name if n.client_id in cmap else "", "client_id": n.client_id, "kind": "notice",
                       "label": f"{n.form} ka jawab", "due": n.due_on, "due_nice": _nice_date(dd), "days": (dd - on).days, "notice_id": n.id}
                (late if row["days"] < 0 else items).append(row)
    items.sort(key=lambda x: x["days"])
    late.sort(key=lambda x: x["days"])
    lines = []
    if late:
        lines.append(f"Late: {len(late)} (" + ", ".join(f"{x['client']} {x['label']}" for x in late[:4]) + ")")
    soon = [x for x in items if x["days"] <= 3]
    if soon:
        lines.append("3 din mein: " + ", ".join(f"{x['client']} {x['label']} {x['due_nice']}" for x in soon[:6]))
    text = ". ".join(lines) or "Agle dino mein koi GST kaam due nahi."
    return {"date": on.isoformat(), "due": items, "late": late, "clients": len(cs), "text": text}


def client_request_draft(tenant: str, client: str, period: str = "") -> dict[str, Any]:
    """WhatsApp draft to the client: send this month's sales + purchase bills (she sends it)."""
    with session() as s:
        c = _client(s, tenant, client)
        st = _settings_row(s, tenant)
        on = today(tenant)
        p = R.norm_period(period) if period else f"{R.add_months(on.year, on.month, -1)[0]:04d}-{R.add_months(on.year, on.month, -1)[1]:02d}"
        obs = R.obligations(p, c.scheme, c.state)
        due = min((o["due"] for o in obs), default=None)
        txt = (f"Namaste {c.contact_name or c.name} ji, {R.nice_period(p)} ka GST taiyaar karna hai. Kripya sale bills, kharid bills, "
               f"credit / debit notes aur bank se supplier payment ki details bhej dijiye"
               + (f" {_nice_date(due - timedelta(days=4))} tak" if due else "") + f". - {st.practitioner or 'aapke GST consultant'}")
        return {"draft": txt, "to": c.contact_phone, "say": "Client ke liye message draft (aap bhejengi)."}


def tax_message_draft(tenant: str, ref: str) -> dict[str, Any]:
    """After 3B is ready: tell the client how much tax to pay (they pay the challan themselves)."""
    with session() as s:
        r = _get_ret(s, tenant, ref)
        c = s.get(GstClient, r.client_id)
        sm = _j(r.data, {}).get("summary", {})
        if r.kind not in ("gstr3b", "cmp08"):
            raise GstError("Yeh 3B / CMP-08 ke liye hai.")
        amt = sm.get("pay_total", sm.get("tax", 0))
        due = R.due_date(r.kind, r.period, c.scheme, c.state)
        txt = (f"Namaste {c.contact_name or c.name} ji, {R.nice_period(r.period)} ka {RETURN_LABEL[r.kind]} taiyaar hai. Cash mein "
               f"tax {_rs(amt)} banta hai" + (f" (byaaj {_rs(sm.get('interest', 0))}, late fee {_rs(sm.get('late_fee', 0))} shaamil)" if sm.get("late_days") else "")
               + f". GST portal par challan (PMT-06) bana kar {_nice_date(due)} se pehle bhar dijiye, phir OTP aayega return file karne ke liye.")
        return {"draft": txt, "to": c.contact_phone, "say": "Client ko tax batane ka draft (aap bhejengi; Hoola paise nahi bharta)."}


# ---------------------------------------------------------------- morning
def due_mornings() -> list[str]:
    from ..life import tz_of
    out = []
    with session() as s:
        for r in s.scalars(select(GstSettings).where(GstSettings.enabled.is_(True))).all():
            if not s.scalars(select(GstClient.id).where(GstClient.tenant_id == r.tenant_id, GstClient.active.is_(True)).limit(1)).first():
                continue
            local = datetime.now(tz_of(r.tenant_id))
            d = local.date().isoformat()
            try:
                hh, mm = [int(x) for x in (r.morning_time or "09:30").split(":")]
            except ValueError:
                hh, mm = 9, 30
            late = (local.hour * 60 + local.minute) - (hh * 60 + mm)
            if r.last_morning != d and 0 <= late <= 180:
                r.last_morning = d
                out.append(r.tenant_id)
    return out


def send_morning(tenant: str) -> dict[str, Any]:
    from ..events import notify
    from ..net.service import home_conv
    cal = calendar(tenant, 7)
    if not cal["due"] and not cal["late"]:
        return cal
    conv = home_conv(tenant)
    try:
        notify(conv, "GST: " + cal["text"])
    except Exception:
        pass
    _publish(tenant, "gst_day", {"date": cal["date"], "text": cal["text"], "due": cal["due"][:20], "late": cal["late"][:20]}, conv)
    return cal


def prompt_block(tenant: str) -> str:
    with session() as s:
        n = s.scalars(select(GstClient).where(GstClient.tenant_id == tenant, GstClient.active.is_(True))).all()
    if not n:
        return ""
    return "GST clients: " + ", ".join(f"{c.name} ({c.gstin}, {c.scheme})" for c in n[:25])
