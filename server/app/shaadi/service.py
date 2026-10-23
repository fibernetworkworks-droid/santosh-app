"""Shaadi Sahayak: Hoola as a wedding planner's coworker (Indian weddings, every community), end to end.

  1. Wedding file: community-aware functions (sagai, haldi, mehendi, sangeet, shaadi, reception ... editable), dates ONLY from her,
     panchang FACTS for those dates from app.india.days (tithi, nakshatra, vaar, holidays); the muhurat is the family pandit's word,
     stored only as she says it ("muhurat_by"), never worked out by Hoola.
  2. Guests per side with RSVP per function and headcount for the caterer; room blocks with allocation (overbooking refused in code);
     guest travel grouped by city, handed to the Travel Sahayak on her word.
  3. Vendors (decor, caterer, photographer, mehendi, DJ, pandit, makeup, transport ...) with HER quotes, a payment schedule with
     advances, budget by function and head (planned / committed / paid), day-wise timeline (noise rule after 10 pm), checklist
     counted back from the first function, shagun / gift register (owner only, never shared), invite text drafts (design via
     the Brand Sahayak), wedding registration checklist by law (checklist only).

Waits for HER tap (only the app API sets by_tap=True; tools only make the card):
  * vendor booking + its advance (`vendor_book`), every vendor payment (`vendor_pay`), any other spend (`spend`) -> money cards,
    her phone's fingerprint once a phone key is active (approvekey kind "shaadi");
  * the client quote (`client_quote`) and client payment requests (`client_pay`) when she plans for clients;
  * every message / invite to guests or vendors (`message`) -> wa.me / SMS drafts she sends herself.
Hoola never pays, never collects, never sends a message, never invents a price, date or muhurat."""
from __future__ import annotations

import csv
import io
import json
import logging
import re
import secrets
import threading
from datetime import date, datetime, timedelta, timezone
from pathlib import Path
from typing import Any
from urllib.parse import quote as urlq

from sqlalchemy import select

from ..db import now, session
from . import rules as R
from .models import ShaadiAlert, ShaadiCard, ShaadiGift, ShaadiGuest, ShaadiRoom, ShaadiSettings, ShaadiVendor, ShaadiWedding

log = logging.getLogger("santosh.shaadi")
_LOCK = threading.RLock()
CARD_KINDS = ("vendor_book", "vendor_pay", "spend", "client_quote", "client_pay", "message")
MONEY_KINDS = ("vendor_book", "vendor_pay", "spend")
DEFAULT_STAGES = [{"label": "Booking par", "pct": 30}, {"label": "Shaadi se 30 din pehle", "pct": 50},
                  {"label": "Shaadi ke baad", "pct": 20}]


class ShaadiError(Exception):
    pass


# ---------------------------------------------------------------- helpers
def _j(s: str | None, default):
    try:
        v = json.loads(s or "")
        return v if isinstance(v, type(default)) else default
    except (TypeError, ValueError):
        return default


def _dump(v) -> str:
    return json.dumps(v, ensure_ascii=False, default=str)


def _aware(d: datetime | None) -> datetime | None:
    if d is None:
        return None
    return d if d.tzinfo else d.replace(tzinfo=timezone.utc)


def clock() -> datetime:
    return now()


def _tz(tenant: str):
    from ..life import tz_of
    return tz_of(tenant)


def today(tenant: str) -> date:
    return clock().astimezone(_tz(tenant)).date()


def _date(v) -> date | None:
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    if not v:
        return None
    s = str(v).strip()[:10]
    for f in ("%Y-%m-%d", "%d-%m-%Y", "%d/%m/%Y", "%d.%m.%Y", "%d-%m-%y", "%d/%m/%y"):
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            pass
    return None


def _iso(d: date | None) -> str:
    return d.isoformat() if d else ""


def _num(v) -> float | None:
    if v is None or v == "" or isinstance(v, bool):
        return None
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).lower().replace(",", "").strip()
    mult = 1.0
    if re.search(r"\b(lakh|lac|l)\b", s):
        mult = 100000.0
    elif re.search(r"\b(k|hazar|hazaar|thousand)\b", s):
        mult = 1000.0
    elif re.search(r"\b(cr|crore)\b", s):
        mult = 10000000.0
    m = re.search(r"-?\d+(?:\.\d+)?", s)
    return float(m.group(0)) * mult if m else None


def _int(v, default: int = 0) -> int:
    n = _num(v)
    return int(n) if n is not None else default


def _rs(v: float | None) -> str:
    if v is None:
        return "[daam?]"
    try:
        from ..india.money import inr
        return inr(v)
    except Exception:
        return f"₹{v:,.0f}"


def _pdf_rs(v: float | None) -> str:
    return "[ ? ]" if v is None else "Rs. " + f"{v:,.2f}"


def r2(v: float) -> float:
    return round(float(v) + 1e-9, 2)


def _publish(tenant: str, kind: str, payload: dict, conv: str = "") -> None:
    from ..events import bus
    try:
        if not conv:
            from ..net.service import home_conv
            conv = home_conv(tenant)
        bus.publish(conv, kind, payload)
    except Exception:
        log.exception("shaadi publish failed")


def is_owner(actor: str) -> bool:
    a = actor or "owner"
    return a == "owner" or a.startswith("approved:")


def actor_name(tenant: str, actor: str) -> str:
    if is_owner(actor):
        return "Owner"
    try:
        from ..employee.service import actor_label
        return actor_label(tenant, actor)
    except Exception:
        return str(actor)[:40]


def _st(s, tenant: str) -> ShaadiSettings:
    r = s.get(ShaadiSettings, tenant)
    if r is None:
        r = ShaadiSettings(tenant_id=tenant)
        s.add(r)
        s.flush()
    return r


def _load_file(tenant: str, file_id: str) -> tuple[str, bytes]:
    from ..mailbox import service as ms
    try:
        name, _m, data = ms.load_files(tenant, [{"type": "file", "file_id": file_id}])[0]
        return name, data
    except Exception as e:
        raise ShaadiError(f"File {file_id} nahi mili: {e}")


def _store(tenant: str, name: str, data: bytes, mime: str, conv: str = "", title: str = "") -> dict:
    from .. import chatfiles
    f = chatfiles.store(tenant, name, data, mime)
    _publish(tenant, "pc_file", {**f, "source": "shaadi", "title": title or name}, conv)
    return f


def _wa(phone: str, text: str) -> str:
    d = R.phone(phone)
    return f"https://wa.me/91{d}?text={urlq(text)}" if d else ""


def _sms(phone: str, text: str) -> str:
    d = R.phone(phone)
    return f"sms:+91{d}?body={urlq(text)}" if d else ""


def _upi(tenant: str, amount: float, note: str) -> str:
    with session() as s:
        st = _st(s, tenant)
        upi, nm = st.upi_id, st.firm
    if not upi:
        try:
            from ..employee import service as emp
            p = emp.profile(tenant) or {}
            upi = p.get("upi_id") or ""
            nm = nm or p.get("upi_name") or p.get("business_name") or ""
        except Exception:
            upi = ""
    if not upi or amount <= 0:
        return ""
    return f"upi://pay?pa={urlq(upi)}&pn={urlq(nm)}&am={amount:.2f}&cu=INR&tn={urlq(note[:40])}"


def _pdf(title: str, blocks: list[tuple], watermark: str = "") -> bytes:
    from xml.sax.saxutils import escape

    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    from .. import indicpdf
    pdf = indicpdf.blocks_pdf(title, blocks, watermark)       # Hindi / Marathi names: printed by Chromium (shaped)
    if pdf:
        return pdf
    ss = getSampleStyleSheet()
    small = ss["BodyText"].clone("small", fontSize=8.5, leading=10.5)

    def txt(t: str) -> str:
        return indicpdf.rl_markup(escape(str(t or "").replace("₹", "Rs. ")))
    story: list = []
    for b in blocks:
        if b[0] == "h":
            story.append(Paragraph(txt(b[1]), ss["Heading2"]))
        elif b[0] == "h1":
            story.append(Paragraph(txt(b[1]), ss["Title"]))
        elif b[0] == "p":
            story.append(Paragraph(txt(b[1]), ss["BodyText"]))
        elif b[0] == "table":
            rows = [[Paragraph(txt(c), small) for c in r] for r in b[1]]
            t = Table(rows, colWidths=b[2] if len(b) > 2 else None, repeatRows=1)
            t.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.3, colors.grey), ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#EEEEEE")),
                                   ("VALIGN", (0, 0), (-1, -1), "TOP")]))
            story.append(t)
        story.append(Spacer(1, 6))
    bio = io.BytesIO()

    def mark(c, _d):
        if watermark:
            c.saveState()
            c.setFont("Helvetica-Bold", 46)
            c.setFillColor(colors.Color(0.85, 0.2, 0.2, alpha=0.18))
            c.translate(300, 420)
            c.rotate(35)
            c.drawCentredString(0, 0, watermark)
            c.restoreState()
    doc = SimpleDocTemplate(bio, pagesize=A4, title=title, leftMargin=36, rightMargin=36, topMargin=36, bottomMargin=36)
    doc.build(story, onFirstPage=mark, onLaterPages=mark)
    return bio.getvalue()


# ---------------------------------------------------------------- settings
_SET_KEYS = ("firm", "mode", "gstin", "phone", "upi_id", "fee_mode", "fee_value", "gst_pct", "stages", "terms", "validity_days",
             "morning_time", "rules")


def _settings_view(r: ShaadiSettings) -> dict:
    return {"firm": r.firm, "mode": r.mode, "gstin": r.gstin, "phone": r.phone, "upi_id": r.upi_id, "fee_mode": r.fee_mode,
            "fee_value": r.fee_value, "gst_pct": r.gst_pct, "stages": _j(r.stages, []) or DEFAULT_STAGES,
            "stages_custom": bool(_j(r.stages, [])), "terms": r.terms, "validity_days": r.validity_days,
            "morning_time": r.morning_time, "rules": r.rules}


def settings(tenant: str) -> dict:
    with session() as s:
        return _settings_view(_st(s, tenant))


def save_settings(tenant: str, **kw: Any) -> dict:
    with _LOCK, session() as s:
        r = _st(s, tenant)
        for k in ("firm", "terms", "rules"):
            if k in kw and kw[k] is not None:
                setattr(r, k, R.clean(kw[k], 2000 if k != "firm" else 160))
        if kw.get("mode"):
            if kw["mode"] not in ("family", "agency"):
                raise ShaadiError("mode: family (apni family ki shaadi) ya agency (clients ke liye).")
            r.mode = kw["mode"]
        if "gstin" in kw and kw["gstin"] is not None:
            g = str(kw["gstin"]).strip().upper()
            if g:
                from ..india.money import gstin_ok
                if not gstin_ok(g):
                    raise ShaadiError("GSTIN sahi nahi lag raha (check digit).")
            r.gstin = g
        if "phone" in kw and kw["phone"] is not None:
            p = R.phone(kw["phone"])
            if kw["phone"] and not p:
                raise ShaadiError("Phone 10 ankon ka chahiye.")
            r.phone = p
        if "upi_id" in kw and kw["upi_id"] is not None:
            u = str(kw["upi_id"]).strip()
            if u and not re.fullmatch(r"[\w.\-]{2,}@[a-zA-Z]{2,}", u):
                raise ShaadiError("UPI ID jaisa nahi lag raha (naam@bank).")
            r.upi_id = u
        if kw.get("fee_mode"):
            if kw["fee_mode"] not in ("none", "percent", "flat"):
                raise ShaadiError("fee_mode: none | percent | flat")
            r.fee_mode = kw["fee_mode"]
        for k in ("fee_value", "gst_pct"):
            if k in kw and kw[k] is not None and kw[k] != "":
                v = _num(kw[k])
                if v is None or v < 0:
                    raise ShaadiError(f"{k} sahi number nahi.")
                setattr(r, k, v)
        if "validity_days" in kw and kw["validity_days"]:
            r.validity_days = max(1, min(90, _int(kw["validity_days"], 15)))
        if "morning_time" in kw and kw["morning_time"]:
            m = R.hhmm(kw["morning_time"])
            if m is None:
                raise ShaadiError("Subah ka samay HH:MM mein.")
            r.morning_time = R.fmt(m)
        if "stages" in kw and kw["stages"] is not None:
            st = kw["stages"]
            if isinstance(st, str):
                st = _j(st, [])
            st = [{"label": R.clean(x.get("label"), 80), "pct": float(_num(x.get("pct")) or 0)} for x in st if isinstance(x, dict)]
            if st and abs(sum(x["pct"] for x in st) - 100) > 0.01:
                raise ShaadiError("Payment stages ka jod 100% hona chahiye.")
            r.stages = _dump(st)
        return _settings_view(r)


# ---------------------------------------------------------------- wedding file
def _wed(s, tenant: str, ref: str) -> ShaadiWedding:
    ref = (ref or "").strip()
    rows = s.scalars(select(ShaadiWedding).where(ShaadiWedding.tenant_id == tenant)).all()
    if not ref:
        live = [w for w in rows if w.status == "planning"]
        if len(live) == 1:
            return live[0]
        raise ShaadiError("Kaunsi shaadi? (number / naam batayiye)")
    for w in rows:
        if ref in (w.id, w.number) or ref.lower() == (w.title or "").lower():
            return w
    hits = [w for w in rows if ref.lower() in (w.title or "").lower() or ref.lower() in f"{w.bride} {w.groom}".lower()]
    if len(hits) == 1:
        return hits[0]
    raise ShaadiError(f"Shaadi '{ref}' nahi mili.")


def _log(w: ShaadiWedding, text: str, by: str = "") -> None:
    lg = _j(w.log, [])
    lg.append({"at": clock().isoformat(timespec="minutes"), "text": text[:300], "by": by})
    w.log = _dump(lg[-200:])


def _new_fn(key: str, name: str = "") -> dict:
    return {"id": "f" + secrets.token_hex(3), "key": key, "name": name or R.F.get(key, key.title()), "date": "", "start": "",
            "end": "", "venue": "", "side": "both", "muhurat": "", "muhurat_by": "", "notes": "", "timeline": []}


def _fn(fns: list[dict], ref: str) -> dict | None:
    r = (ref or "").strip().lower()
    if not r:
        return None
    for f in fns:
        if r in (f["id"].lower(), f["key"], f["name"].lower()):
            return f
    k = R.fn_key(r)
    hits = [f for f in fns if k and f["key"] == k] or [f for f in fns if r in f["name"].lower()]
    return hits[0] if len(hits) == 1 else None


def _dates(w: ShaadiWedding) -> list[date]:
    return sorted(d for d in (_date(f.get("date")) for f in _j(w.functions, [])) if d)


def _checklist_items(community: str) -> list[dict]:
    rows = list(R.CHECKLIST) + list(R.EXTRA_CHECKLIST.get(community, []))
    rows.sort(key=lambda x: -x[0])
    return [{"id": f"c{i + 1}", "days": d, "what": what, "kind": k, "done": False, "done_on": ""} for i, (d, what, k) in enumerate(rows)]


def wedding_new(tenant: str, title: str = "", community: str = "", bride: str = "", groom: str = "", client: str = "",
                client_phone: str = "", city: str = "", state: str = "", venue: str = "", guests_estimate=None, budget=None,
                age_bride=None, age_groom=None, functions: list | None = None, dates: dict | None = None,
                actor: str = "owner", conv: str = "") -> dict:
    p = R.age_problem(age_bride, age_groom)
    if p:
        raise ShaadiError(p)
    p = R.dowry_problem(f"{title} {venue}")
    if p:
        raise ShaadiError(p)
    comm = R.community_of(community) or (community if community in R.COMMUNITIES else "other")
    keys = []
    for x in functions or []:
        k = R.fn_key(str(x))
        keys.append((k or "custom", str(x)))
    if not keys:
        keys = [(k, "") for k in R.COMMUNITIES[comm]["fns"]]
    fns = [_new_fn(k, "" if k != "custom" else R.clean(nm, 60)) for k, nm in keys]
    for k, v in (dates or {}).items():
        f = _fn(fns, str(k))
        d = _date(v)
        if f and d:
            f["date"] = d.isoformat()
    cp = R.phone(client_phone)
    if client_phone and not cp:
        raise ShaadiError("Client ka phone 10 ankon ka chahiye.")
    with _LOCK, session() as s:
        st = _st(s, tenant)
        st.wed_seq += 1
        b = _num(budget)
        w = ShaadiWedding(tenant_id=tenant, number=f"W-{st.wed_seq:03d}", community=comm,
                          title=R.clean(title, 160) or f"{R.clean(bride, 60) or 'Dulhan'} weds {R.clean(groom, 60) or 'Dulha'}",
                          bride=R.clean(bride, 120), groom=R.clean(groom, 120), client=R.clean(client, 120), client_phone=cp,
                          city=R.clean(city, 60), state=R.clean(state, 40), venue=R.clean(venue, 400),
                          guests_estimate=_int(guests_estimate), budget=_dump({"total": b, "heads": {}}), functions=_dump(fns),
                          checklist=_dump(_checklist_items(comm)), conv=conv)
        _log(w, "Shaadi ki file bani", actor_name(tenant, actor))
        s.add(w)
        s.flush()
        view = _wedding_view(s, w)
    say = (f"{view['number']} {view['title']}: {len(fns)} functions ({R.COMMUNITIES[comm]['label']} ke hisaab se, badal sakti hain). "
           "Dates aap batayiye; muhurat panditji tay karenge, Hoola nahi.")
    return {"wedding": view, "say": say}


_W_KEYS = ("title", "bride", "groom", "client", "client_phone", "city", "state", "venue", "guests_estimate", "status", "community")


def wedding_update(tenant: str, wedding: str, actor: str = "owner", **kw: Any) -> dict:
    p = R.age_problem(kw.get("age_bride"), kw.get("age_groom"))
    if p:
        raise ShaadiError(p)
    with _LOCK, session() as s:
        w = _wed(s, tenant, wedding)
        for k in _W_KEYS:
            if k not in kw or kw[k] is None:
                continue
            v = kw[k]
            if k == "client_phone":
                ph = R.phone(v)
                if v and not ph:
                    raise ShaadiError("Phone 10 ankon ka chahiye.")
                w.client_phone = ph
            elif k == "guests_estimate":
                w.guests_estimate = _int(v)
            elif k == "status":
                if v not in ("planning", "done", "cancelled"):
                    raise ShaadiError("status: planning | done | cancelled")
                if v != w.status and not is_owner(actor):
                    raise ShaadiError("Status sirf owner badal sakti hain.")
                w.status = v
            elif k == "community":
                c = R.community_of(v) or (v if v in R.COMMUNITIES else "")
                if c:
                    w.community = c
            else:
                setattr(w, k, R.clean(v, 400))
        if "budget" in kw and kw["budget"] not in (None, ""):
            b = _j(w.budget, {})
            b["total"] = _num(kw["budget"])
            w.budget = _dump(b)
        _log(w, "Badlav: " + ", ".join(k for k in kw if kw[k] is not None)[:200], actor_name(tenant, actor))
        return {"wedding": _wedding_view(s, w), "say": "Theek, badal diya."}


def fn_set(tenant: str, wedding: str, fn: str = "", name: str = "", date_: str = "", start: str = "", end: str = "", venue: str = "",
           side: str = "", muhurat: str = "", muhurat_by: str = "", notes: str = "", remove: bool = False, add: bool = False,
           actor: str = "owner") -> dict:
    """Add / change / remove a function. The muhurat is stored ONLY as she gives it, with who gave it (pandit's word)."""
    warns = []
    with _LOCK, session() as s:
        w = _wed(s, tenant, wedding)
        fns = _j(w.functions, [])
        f = None if add else _fn(fns, fn or name)
        if f is None:
            if not (add or name or fn):
                raise ShaadiError("Kaunsa function?")
            k = R.fn_key(fn or name) or "custom"
            f = _new_fn(k, R.clean(name or fn, 60) if k == "custom" or name else "")
            fns.append(f)
            warns.append(f"Naya function: {f['name']}")
        if remove:
            fns = [x for x in fns if x["id"] != f["id"]]
            w.functions = _dump(fns)
            _log(w, f"Function hataya: {f['name']}", actor_name(tenant, actor))
            return {"wedding": _wedding_view(s, w), "say": f"{f['name']} hata diya."}
        if name:
            f["name"] = R.clean(name, 60)
        if date_:
            d = _date(date_)
            if not d:
                raise ShaadiError("Tareekh samajh nahi aayi (DD-MM-YYYY).")
            if d < today(tenant):
                warns.append("Yeh tareekh beet chuki hai.")
            f["date"] = d.isoformat()
        for key, val in (("start", start), ("end", end)):
            if val:
                m = R.hhmm(val)
                if m is None:
                    raise ShaadiError(f"{key} ka samay samajh nahi aaya (jaise 7 pm / 19:30).")
                f[key] = R.fmt(m)
        if venue:
            f["venue"] = R.clean(venue, 200)
        if side:
            if side not in ("both", "bride", "groom"):
                raise ShaadiError("side: both | bride | groom")
            f["side"] = side
        if muhurat:
            f["muhurat"] = R.clean(muhurat, 120)
            f["muhurat_by"] = R.clean(muhurat_by, 80) or "aapne bataya"
        if notes:
            p = R.dowry_problem(notes)
            if p:
                raise ShaadiError(p)
            f["notes"] = R.clean(notes, 500)
        # same date + same venue + overlapping hours
        for o in fns:
            if o["id"] != f["id"] and o.get("date") and o.get("date") == f.get("date") and f.get("start") and o.get("start"):
                a1, a2 = R.hhmm(f["start"]), R.hhmm(f.get("end") or "") or (R.hhmm(f["start"]) or 0) + 180
                b1, b2 = R.hhmm(o["start"]), R.hhmm(o.get("end") or "") or (R.hhmm(o["start"]) or 0) + 180
                if a1 is not None and b1 is not None and a1 < b2 and b1 < a2:
                    warns.append(f"{o['name']} se samay takra raha hai ({o['start']}).")
        w.functions = _dump(fns)
        _log(w, f"{f['name']}: {f.get('date')} {f.get('start')}", actor_name(tenant, actor))
        view = _wedding_view(s, w)
    say = f"{f['name']}: {f.get('date') or 'tareekh baaki'} {f.get('start') or ''}".strip()
    if f["key"] in R.RITUAL_FNS and not f.get("muhurat"):
        say += ". Muhurat panditji batayenge; jab bataayein mujhe bata dijiye."
    if warns:
        say += " Dhyaan: " + " ".join(warns)
    return {"wedding": view, "function": f, "warnings": warns, "say": say}


def panchang(tenant: str, wedding: str) -> dict:
    """Panchang FACTS for the dates she chose (india.days) + public holidays. No muhurat: that is the pandit's decision."""
    from ..india import days
    with session() as s:
        w = _wed(s, tenant, wedding)
        fns, city, state = _j(w.functions, []), w.city, w.state
    out = []
    for f in fns:
        d = _date(f.get("date"))
        if not d:
            out.append({"function": f["name"], "date": "", "say": "Tareekh abhi tay nahi."})
            continue
        try:
            p = days.panchang(d, city)
            info = days.day_info(d, state, city)
            hol = [h["name"] for h in info.get("holidays", [])]
        except Exception as e:
            log.warning("panchang failed: %s", e)
            p, hol = {}, []
        row = {"function": f["name"], "date": d.isoformat(), "vaar": p.get("vaar"), "tithi": p.get("tithi"), "paksha": p.get("paksha"),
               "nakshatra": p.get("nakshatra"), "month_amanta": p.get("month_amanta"), "sunrise": p.get("sunrise"),
               "holidays": hol, "muhurat": f.get("muhurat") or "", "muhurat_by": f.get("muhurat_by") or ""}
        if f["key"] in R.RITUAL_FNS and not f.get("muhurat"):
            row["muhurat_note"] = "Muhurat panditji se poochhiye (Hoola nahi nikalta)."
        if hol:
            row["holiday_note"] = "Chhutti ka din: venue / hotel / rail mein bheed aur daam zyada ho sakte hain."
        out.append(row)
    return {"days": out, "note": "Panchang hisaab se nikala (±5 min, amanta mahina). Shubh muhurat aapke panditji ka faisla hai.",
            "say": "; ".join(f"{r['function']} {r['date']}: {r.get('tithi') or ''} {r.get('paksha') or ''}, {r.get('nakshatra') or ''}"
                              for r in out if r.get("date")) or "Abhi kisi function ki tareekh tay nahi."}


def timeline_set(tenant: str, wedding: str, fn: str, rows: list | None = None, starter: bool = False, actor: str = "owner") -> dict:
    """Her rows [{time, what, who}] or a starter from her start time. Ritual rows get her muhurat or a blank, never a guess."""
    with _LOCK, session() as s:
        w = _wed(s, tenant, wedding)
        fns = _j(w.functions, [])
        f = _fn(fns, fn)
        if f is None:
            raise ShaadiError(f"Function '{fn}' nahi mila.")
        out = []
        if starter:
            base = R.hhmm(f.get("start") or "")
            if base is None:
                raise ShaadiError(f"{f['name']} ka shuru hone ka samay pehle batayiye; uske hisaab se timeline banegi.")
            tmpl = R.TIMELINE.get(f["key"]) or R.TIMELINE["default"]
            for off, what, who in tmpl:
                out.append({"time": R.fmt(base + off), "what": what, "who": who})
            if f["key"] in R.RITUAL_FNS:
                out.append({"time": f.get("muhurat") or "", "what": "Rasm / phere / nikah (muhurat " +
                            ("panditji ka: " + f["muhurat"] if f.get("muhurat") else "panditji se poochhna baaki") + ")", "who": "pandit"})
        for r in rows or []:
            if not isinstance(r, dict):
                continue
            t = R.hhmm(r.get("time")) if r.get("time") else None
            if r.get("time") and t is None:
                raise ShaadiError(f"Samay samajh nahi aaya: {r.get('time')}")
            out.append({"time": R.fmt(t) if t is not None else "", "what": R.clean(r.get("what"), 200),
                        "who": R.vendor_kind(str(r.get("who") or "")) if r.get("who") else ""})
        out.sort(key=lambda x: (x["time"] == "", x["time"]))
        f["timeline"] = out if (starter or rows) else f.get("timeline", [])
        w.functions = _dump(fns)
        vendors = s.scalars(select(ShaadiVendor).where(ShaadiVendor.wedding_id == w.id)).all()
        warns = _fn_warnings(f, vendors)
        view = _wedding_view(s, w)
    return {"function": f, "warnings": warns, "wedding": view,
            "say": f"{f['name']} ki timeline: {len(f['timeline'])} kaam." + (" Dhyaan: " + " ".join(warns) if warns else "")}


def _fn_warnings(f: dict, vendors: list[ShaadiVendor]) -> list[str]:
    loud_kinds = {"dj", "band"}
    loud = any(v.kind in loud_kinds and (not _j(v.functions, []) or f["id"] in _j(v.functions, []))
               and v.status in ("booked", "waiting", "quoted") for v in vendors)
    loud = loud or any(r.get("who") in loud_kinds for r in f.get("timeline", []))
    out = []
    last_loud = max([R.hhmm(r["time"]) for r in f.get("timeline", []) if r.get("who") in loud_kinds and r.get("time")] or [-1])
    st, en = R.hhmm(f.get("start") or ""), R.hhmm(f.get("end") or "")
    if last_loud >= 0 and (en is None or last_loud > en):
        en = last_loud + 60
    p = R.noise_problem(st, en, loud)
    if p:
        out.append(p)
    return out


def day_plan(tenant: str, wedding: str) -> dict:
    with session() as s:
        w = _wed(s, tenant, wedding)
        fns = _j(w.functions, [])
        vendors = s.scalars(select(ShaadiVendor).where(ShaadiVendor.wedding_id == w.id)).all()
        vname = {}
        for v in vendors:
            if v.status in ("booked", "done"):
                vname.setdefault(v.kind, []).append(v.name)
    days: dict[str, list] = {}
    for f in fns:
        rows = [{**r, "vendor": ", ".join(vname.get(r.get("who") or "", []))} for r in f.get("timeline", [])]
        days.setdefault(f.get("date") or "tay nahi", []).append({"function": f["name"], "start": f.get("start"), "end": f.get("end"),
                                                               "venue": f.get("venue") or w.venue, "muhurat": f.get("muhurat"),
                                                               "timeline": rows, "warnings": _fn_warnings(f, vendors)})
    order = sorted(days, key=lambda d: (d == "tay nahi", d))
    return {"days": [{"date": d, "functions": sorted(days[d], key=lambda x: x["start"] or "99")} for d in order]}


# ---------------------------------------------------------------- checklist
def _checklist_view(w: ShaadiWedding, on: date) -> list[dict]:
    ds = _dates(w)
    first, last = (ds[0], ds[-1]) if ds else (None, None)
    out = []
    for c in _j(w.checklist, []):
        days = c.get("days", 0)
        base = first if days >= 0 else last
        due = base - timedelta(days=days) if base else None
        state = "done" if c.get("done") else ("" if not due else "late" if due < on else "today" if due == on else "soon"
                                               if due <= on + timedelta(days=7) else "later")
        out.append({**c, "due": _iso(due), "state": state})
    out.sort(key=lambda x: (x["done"], x["due"] or "9999"))
    return out


def checklist(tenant: str, wedding: str) -> dict:
    with session() as s:
        w = _wed(s, tenant, wedding)
        items = _checklist_view(w, today(tenant))
        has = bool(_dates(w))
    return {"items": items, "say": "" if has else "Function ki tareekh tay hote hi har kaam ki last date nikal jayegi."}


def checklist_mark(tenant: str, wedding: str, item: str, done: bool = True, what: str = "", days=None, actor: str = "owner") -> dict:
    with _LOCK, session() as s:
        w = _wed(s, tenant, wedding)
        cl = _j(w.checklist, [])
        if what and not item:
            d = _int(days, 30)
            cl.append({"id": f"c{len(cl) + 1}x{secrets.token_hex(2)}", "days": d, "what": R.clean(what, 200), "kind": "custom",
                       "done": False, "done_on": ""})
            w.checklist = _dump(cl)
            return {"items": _checklist_view(w, today(tenant)), "say": "Checklist mein joda."}
        hit = [c for c in cl if c["id"] == item] or [c for c in cl if item and item.lower() in c["what"].lower()]
        if not hit:
            raise ShaadiError("Checklist mein yeh kaam nahi mila.")
        hit[0]["done"] = bool(done)
        hit[0]["done_on"] = today(tenant).isoformat() if done else ""
        w.checklist = _dump(cl)
        _log(w, f"Checklist: {hit[0]['what'][:60]} {'ho gaya' if done else 'wapas khula'}", actor_name(tenant, actor))
        return {"items": _checklist_view(w, today(tenant)), "say": f"'{hit[0]['what']}' {'ho gaya' if done else 'phir se baaki'}."}


# ---------------------------------------------------------------- guests + RSVP
def _guest_view(g: ShaadiGuest, fns: list[dict] | None = None, full: bool = True) -> dict:
    inv = _j(g.invited, [])
    names = {f["id"]: f["name"] for f in fns or []}
    return {"id": g.id, "name": g.name, "side": g.side, "group": g.grp, "phone": g.phone if full else (g.phone[-4:] and "xxxxxx" + g.phone[-4:]),
            "adults": g.adults, "kids": g.kids, "heads": g.adults + g.kids, "invited": [names.get(i, i) for i in inv] or ["sab"],
            "invited_ids": inv, "rsvp": {names.get(k, k): v for k, v in _j(g.rsvp, {}).items()}, "rsvp_ids": _j(g.rsvp, {}),
            "diet": g.diet, "city": g.city, "needs_room": g.needs_room, "room_id": g.room_id, "arrive": g.arrive, "depart": g.depart,
            "travel": _j(g.travel, {}), "notes": g.notes, "no_messages": g.no_messages}


def _invited(g: ShaadiGuest, fn_id: str) -> bool:
    inv = _j(g.invited, [])
    return not inv or fn_id in inv


def guest_add(tenant: str, wedding: str, name: str, side: str = "", phone: str = "", adults=None, kids=None, group: str = "",
              invited: list | None = None, diet: str = "", city: str = "", needs_room=None, arrive: str = "", depart: str = "",
              travel_mode: str = "", travel_from: str = "", notes: str = "", guest: str = "", no_messages=None,
              actor: str = "owner") -> dict:
    nm = R.clean(name, 120)
    ph = R.phone(phone)
    if phone and not ph:
        raise ShaadiError(f"{nm}: phone 10 ankon ka chahiye.")
    if side and side not in R.GUEST_SIDES:
        side = {"ladki": "bride", "dulhan": "bride", "ladka": "groom", "dulha": "groom"}.get(side.lower(), "common")
    with _LOCK, session() as s:
        w = _wed(s, tenant, wedding)
        fns = _j(w.functions, [])
        g = None
        if guest:
            g = s.get(ShaadiGuest, guest)
            if g is None or g.wedding_id != w.id:
                raise ShaadiError("Mehmaan nahi mila.")
        elif ph:
            g = s.scalars(select(ShaadiGuest).where(ShaadiGuest.wedding_id == w.id, ShaadiGuest.phone == ph)).first()
        if g is None:
            if not nm:
                raise ShaadiError("Mehmaan ka naam chahiye.")
            g = ShaadiGuest(tenant_id=tenant, wedding_id=w.id, name=nm)
            s.add(g)
            new = True
        else:
            new = False
        if nm:
            g.name = nm
        if ph:
            g.phone = ph
        if side:
            g.side = side
        if adults not in (None, ""):
            g.adults = max(0, _int(adults, 1))
        if kids not in (None, ""):
            g.kids = max(0, _int(kids))
        for k, v in (("grp", group), ("diet", diet), ("city", city)):
            if v:
                setattr(g, k, R.clean(v, 60))
        if notes:
            g.notes = R.clean(notes, 300)
        if invited is not None:
            ids = []
            for x in invited:
                f = _fn(fns, str(x))
                if f is None and str(x).lower() not in ("sab", "all"):
                    raise ShaadiError(f"Function '{x}' nahi mila.")
                if f:
                    ids.append(f["id"])
            g.invited = _dump(ids)
        if needs_room is not None:
            g.needs_room = bool(needs_room)
        for k, v in (("arrive", arrive), ("depart", depart)):
            if v:
                d = _date(v)
                if not d:
                    raise ShaadiError(f"{k} tareekh samajh nahi aayi.")
                setattr(g, k, d.isoformat())
        if travel_mode or travel_from:
            t = _j(g.travel, {})
            if travel_mode:
                t["mode"] = R.clean(travel_mode, 20)
            if travel_from:
                t["from"] = R.clean(travel_from, 60)
            g.travel = _dump(t)
        if no_messages is not None:
            g.no_messages = bool(no_messages)
        s.flush()
        view = _guest_view(g, fns)
    return {"guest": view, "say": f"{view['name']} ({view['heads']} log, {view['side']}) {'jode' if new else 'badle'}."}


def guests_import(tenant: str, wedding: str, file_id: str = "", data: bytes | None = None, name: str = "guests.csv",
                  side: str = "") -> dict:
    if data is None:
        name, data = _load_file(tenant, file_id)
    rows: list[list] = []
    if name.lower().endswith((".xlsx", ".xlsm")):
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        rows = [list(r) for r in wb.active.iter_rows(values_only=True)]
    else:
        rows = list(csv.reader(io.StringIO(data.decode("utf-8-sig", "ignore"))))
    if not rows:
        raise ShaadiError("File khali hai.")
    head = [re.sub(r"[^a-z]", "", str(h or "").lower()) for h in rows[0]]

    def col(*words):
        for i, h in enumerate(head):
            if any(h.startswith(w) or w in h for w in words):
                return i
        return None
    ci = {"name": col("name", "naam", "mehmaan", "guest"), "phone": col("phone", "mobile", "number", "whatsapp"),
          "side": col("side", "taraf", "paksh"), "adults": col("adult", "log", "count", "members", "heads"),
          "kids": col("kid", "bachch", "child"), "city": col("city", "shahar", "gaon"), "group": col("group", "relation", "rishta")}
    if ci["name"] is None:
        raise ShaadiError("Sheet mein 'Name / Naam' column chahiye.")
    added, bad = 0, []
    for r in rows[1:]:
        def cell(k):
            i = ci[k]
            return "" if i is None or i >= len(r) or r[i] is None else str(r[i]).strip()
        nm = cell("name")
        if not nm:
            continue
        try:
            guest_add(tenant, wedding, nm, side=cell("side").lower() or side, phone=cell("phone"), adults=cell("adults") or None,
                      kids=cell("kids") or None, group=cell("group"), city=cell("city"))
            added += 1
        except ShaadiError as e:
            bad.append(str(e))
    return {"added": added, "problems": bad[:20], "say": f"{added} mehmaan jode" + (f"; {len(bad)} line mein dikkat." if bad else ".")}


def _guest(s, tenant: str, w: ShaadiWedding, ref: str) -> ShaadiGuest:
    rows = s.scalars(select(ShaadiGuest).where(ShaadiGuest.wedding_id == w.id)).all()
    for g in rows:
        if ref in (g.id, g.phone):
            return g
    hits = [g for g in rows if ref and ref.lower() in g.name.lower()]
    if len(hits) == 1:
        return hits[0]
    raise ShaadiError(f"Mehmaan '{ref}' nahi mila" + (" (ek se zyada naam milte hain)" if hits else "") + ".")


def rsvp(tenant: str, wedding: str, guest: str, answer: str, fn: str = "", adults=None, kids=None, actor: str = "owner") -> dict:
    """Her word (or what the guest told her) for one function or all the functions the guest is invited to."""
    a = {"haan": "yes", "ha": "yes", "aayenge": "yes", "nahi": "no", "na": "no", "shayad": "maybe", "pata nahi": "maybe"}.get(
        (answer or "").lower().strip(), (answer or "").lower().strip())
    if a not in R.RSVP:
        raise ShaadiError("Jawab: yes | no | maybe | pending")
    with _LOCK, session() as s:
        w = _wed(s, tenant, wedding)
        fns = _j(w.functions, [])
        g = _guest(s, tenant, w, guest)
        rs = _j(g.rsvp, {})
        targets = [f for f in fns if _invited(g, f["id"])]
        if fn:
            f = _fn(fns, fn)
            if f is None:
                raise ShaadiError(f"Function '{fn}' nahi mila.")
            if not _invited(g, f["id"]):
                raise ShaadiError(f"{g.name} {f['name']} mein bulaye nahi gaye.")
            targets = [f]
        for f in targets:
            rs[f["id"]] = a
        g.rsvp = _dump(rs)
        if adults not in (None, ""):
            g.adults = max(0, _int(adults, g.adults))
        if kids not in (None, ""):
            g.kids = max(0, _int(kids, g.kids))
        view = _guest_view(g, fns)
    return {"guest": view, "say": f"{view['name']}: {', '.join(f['name'] for f in targets)} -> {a}."}


def headcount(tenant: str, wedding: str) -> dict:
    with session() as s:
        w = _wed(s, tenant, wedding)
        fns = _j(w.functions, [])
        gs = s.scalars(select(ShaadiGuest).where(ShaadiGuest.wedding_id == w.id)).all()
        out = []
        for f in fns:
            c = {"yes": 0, "no": 0, "maybe": 0, "pending": 0}
            diet: dict[str, int] = {}
            sides = {"bride": 0, "groom": 0, "common": 0}
            kids = 0
            for g in gs:
                if not _invited(g, f["id"]):
                    continue
                a = _j(g.rsvp, {}).get(f["id"], "pending")
                c[a] = c.get(a, 0) + g.adults + g.kids
                if a == "yes":
                    kids += g.kids
                    sides[g.side] = sides.get(g.side, 0) + g.adults + g.kids
                    if g.diet:
                        diet[g.diet] = diet.get(g.diet, 0) + g.adults + g.kids
            out.append({"function": f["name"], "id": f["id"], "date": f.get("date"), **c, "kids_yes": kids, "diet_yes": diet,
                        "sides_yes": sides, "invited": sum(c.values()),
                        "plates_range": [c["yes"], c["yes"] + c["maybe"] + c["pending"]]})
        total_guests = sum(g.adults + g.kids for g in gs)
    say = "; ".join(f"{r['function']}: haan {r['yes']}, shayad {r['maybe']}, jawab baaki {r['pending']}" for r in out)
    return {"functions": out, "guests": total_guests, "say": say or "Abhi mehmaan nahi jode.",
            "note": "Plate count caterer ko 'haan' se 'haan + shayad + baaki' ke beech; final ginti aap tay karein."}


def guests(tenant: str, wedding: str, q: str = "", side: str = "", fn: str = "", answer: str = "") -> dict:
    with session() as s:
        w = _wed(s, tenant, wedding)
        fns = _j(w.functions, [])
        f = _fn(fns, fn) if fn else None
        rows = []
        for g in s.scalars(select(ShaadiGuest).where(ShaadiGuest.wedding_id == w.id)).all():
            if q and q.lower() not in f"{g.name} {g.grp} {g.city}".lower():
                continue
            if side and g.side != side:
                continue
            if f and not _invited(g, f["id"]):
                continue
            if answer and f and _j(g.rsvp, {}).get(f["id"], "pending") != answer:
                continue
            rows.append(_guest_view(g, fns))
    rows.sort(key=lambda x: (x["side"], x["name"].lower()))
    return {"guests": rows, "count": len(rows), "heads": sum(r["heads"] for r in rows)}


# ---------------------------------------------------------------- rooms
def _room_view(r: ShaadiRoom, occupants: list[ShaadiGuest]) -> dict:
    heads = sum(g.adults + g.kids for g in occupants)
    return {"id": r.id, "place": r.place, "label": r.label, "capacity": r.capacity, "check_in": r.check_in, "check_out": r.check_out,
            "rate": r.rate, "notes": r.notes, "guests": [{"id": g.id, "name": g.name, "heads": g.adults + g.kids} for g in occupants],
            "heads": heads, "free": max(0, r.capacity - heads)}


def room_add(tenant: str, wedding: str, place: str, label: str = "", capacity=None, check_in: str = "", check_out: str = "",
             rate=None, count=None, notes: str = "") -> dict:
    ci, co = _date(check_in), _date(check_out)
    if not ci or not co:
        raise ShaadiError("Room ki check-in aur check-out tareekh chahiye.")
    if co <= ci:
        raise ShaadiError("Check-out check-in ke baad hona chahiye.")
    cap = _int(capacity, 2)
    if cap < 1 or cap > 20:
        raise ShaadiError("Ek room mein 1-20 log.")
    n = max(1, min(60, _int(count, 1)))
    out = []
    with _LOCK, session() as s:
        w = _wed(s, tenant, wedding)
        have = len(s.scalars(select(ShaadiRoom).where(ShaadiRoom.wedding_id == w.id)).all())
        for i in range(n):
            r = ShaadiRoom(tenant_id=tenant, wedding_id=w.id, place=R.clean(place, 120) or "Room",
                           label=(R.clean(label, 30) + (f"-{i + 1}" if n > 1 else "")) if label else f"R{have + i + 1}",
                           capacity=cap, check_in=ci.isoformat(), check_out=co.isoformat(), rate=_num(rate), notes=R.clean(notes, 300))
            s.add(r)
            s.flush()
            out.append(_room_view(r, []))
    return {"rooms": out, "say": f"{n} room ({place}) {ci} se {co} tak block mein jode. Hotel ki booking / advance alag card se."}


def _stay(g: ShaadiGuest, r: ShaadiRoom) -> tuple[date, date]:
    a = _date(g.arrive) or _date(r.check_in)
    b = _date(g.depart) or _date(r.check_out)
    return a, b


def room_assign(tenant: str, wedding: str, guest: str, room: str = "", remove: bool = False) -> dict:
    """Overbooking refused in code: heads staying on any night never above the room's capacity; stay inside the room's dates."""
    with _LOCK, session() as s:
        w = _wed(s, tenant, wedding)
        g = _guest(s, tenant, w, guest)
        if remove:
            g.room_id = ""
            return {"say": f"{g.name} ka room hata diya."}
        rooms = s.scalars(select(ShaadiRoom).where(ShaadiRoom.wedding_id == w.id)).all()
        r = next((x for x in rooms if room in (x.id, x.label)), None)
        if r is None:
            raise ShaadiError(f"Room '{room}' nahi mila.")
        a, b = _stay(g, r)
        ci, co = _date(r.check_in), _date(r.check_out)
        if a < ci or b > co or b <= a:
            raise ShaadiError(f"{g.name} {a} se {b} rukenge; room {r.label} sirf {ci} se {co} tak block hai.")
        others = [x for x in s.scalars(select(ShaadiGuest).where(ShaadiGuest.wedding_id == w.id, ShaadiGuest.room_id == r.id)).all()
                  if x.id != g.id]
        d = a
        while d < b:
            heads = g.adults + g.kids + sum(o.adults + o.kids for o in others if _stay(o, r)[0] <= d < _stay(o, r)[1])
            if heads > r.capacity:
                raise ShaadiError(f"Room {r.label} ({r.place}) mein {d} ki raat {heads} log ho jaayenge, jagah {r.capacity} ki hai. "
                                  "Overbooking nahi karenge: doosra room chuniye.")
            d += timedelta(days=1)
        g.room_id = r.id
        g.needs_room = True
        view = _room_view(r, others + [g])
    return {"room": view, "say": f"{g.name} -> {r.place} {r.label} ({view['heads']}/{r.capacity})."}


def rooms(tenant: str, wedding: str) -> dict:
    with session() as s:
        w = _wed(s, tenant, wedding)
        gs = s.scalars(select(ShaadiGuest).where(ShaadiGuest.wedding_id == w.id)).all()
        rs = s.scalars(select(ShaadiRoom).where(ShaadiRoom.wedding_id == w.id)).all()
        out = [_room_view(r, [g for g in gs if g.room_id == r.id]) for r in rs]
        need = [{"id": g.id, "name": g.name, "heads": g.adults + g.kids, "arrive": g.arrive, "depart": g.depart}
                for g in gs if g.needs_room and not g.room_id]
    return {"rooms": out, "waiting": need, "free_beds": sum(r["free"] for r in out),
            "say": f"{len(out)} room, {sum(r['free'] for r in out)} jagah khali; {len(need)} mehmaan ko abhi room nahi mila.",
            "note": "Hotel / homestay ki asli booking unke front office se (Hoola ke Front office / Homestay Sahayak jodte hain)."}


# ---------------------------------------------------------------- guest travel (hands over to the Travel Sahayak)
def travel_list(tenant: str, wedding: str) -> dict:
    with session() as s:
        w = _wed(s, tenant, wedding)
        city = (w.city or "").lower()
        groups: dict[str, list] = {}
        for g in s.scalars(select(ShaadiGuest).where(ShaadiGuest.wedding_id == w.id)).all():
            t = _j(g.travel, {})
            frm = (t.get("from") or g.city or "").strip()
            if not frm or frm.lower() == city:
                continue
            key = f"{frm}|{g.arrive or ''}"
            groups.setdefault(key, []).append({"id": g.id, "name": g.name, "heads": g.adults + g.kids, "mode": t.get("mode", ""),
                                               "trip": t.get("trip", "")})
    out = [{"from": k.split("|")[0], "arrive": k.split("|")[1], "guests": v, "heads": sum(x["heads"] for x in v)}
           for k, v in sorted(groups.items())]
    return {"groups": out, "say": f"{len(out)} shahar se mehmaan aa rahe hain." if out else "Bahar se aane wale mehmaan nahi jode.",
            "note": "Train / flight / bus ke liye Travel Sahayak (IRCTC 60 din pehle 8 baje khulta hai). Booking aur paise aapke tap se."}


def to_trip(tenant: str, wedding: str, from_city: str, actor: str = "owner", conv: str = "") -> dict:
    """Make a DRAFT trip in the Travel Sahayak for one city's guests (names only, never IDs)."""
    if not is_owner(actor):
        raise ShaadiError("Trip sirf owner ke kehne par.")
    groups = [g for g in travel_list(tenant, wedding)["groups"] if g["from"].lower() == (from_city or "").lower()]
    if not groups:
        raise ShaadiError(f"{from_city} se koi mehmaan nahi.")
    with session() as s:
        w = _wed(s, tenant, wedding)
        ds = _dates(w)
        dest, title = w.city, w.title
    names = [x["name"] for g in groups for x in g["guests"]]
    arrive = min((g["arrive"] for g in groups if g["arrive"]), default="") or ((ds[0] - timedelta(days=1)).isoformat() if ds else "")
    from ..travel import service as TS
    trip = TS.new_trip(tenant, text=f"{title}: {from_city} se {dest} shaadi ke liye {len(names)} mehmaan", origin=from_city,
                       destination=dest, start_date=arrive, travellers=[{"name": n} for n in names], kind="family",
                       actor=actor, conv=conv)
    tid = (trip.get("trip") or {}).get("id") or (trip.get("trip") or {}).get("number") or ""
    with _LOCK, session() as s:
        w = _wed(s, tenant, wedding)
        for g in groups:
            for x in g["guests"]:
                row = s.get(ShaadiGuest, x["id"])
                t = _j(row.travel, {})
                t["trip"] = tid
                row.travel = _dump(t)
    return {"trip": trip.get("trip"), "say": f"Travel Sahayak mein {from_city} -> {dest} ki trip ({len(names)} log) bani; "
                                             "options aur booking wahan, aapke tap se."}


# ---------------------------------------------------------------- vendors
def _vendor_view(s, v: ShaadiVendor, fns: list[dict]) -> dict:
    names = {f["id"]: f["name"] for f in fns}
    sch = _j(v.schedule, [])
    paid = sum(x.get("paid_amount") or 0 for x in sch if x.get("status") == "paid")
    return {"id": v.id, "name": v.name, "kind": v.kind, "kind_label": R.VENDOR_KINDS.get(v.kind, v.kind), "phone": v.phone,
            "functions": [names.get(i, i) for i in _j(v.functions, [])] or ["sab"], "function_ids": _j(v.functions, []),
            "quote": v.quote, "quote_text": _rs(v.quote), "quote_note": v.quote_note, "per_plate": v.per_plate, "plates": v.plates,
            "schedule": sch, "paid": paid, "due": None if v.quote is None else r2(v.quote - paid), "status": v.status, "notes": v.notes}


def _vendor(s, tenant: str, w: ShaadiWedding, ref: str) -> ShaadiVendor:
    rows = s.scalars(select(ShaadiVendor).where(ShaadiVendor.wedding_id == w.id)).all()
    for v in rows:
        if ref in (v.id, v.phone):
            return v
    hits = [v for v in rows if ref and ref.lower() in v.name.lower()]
    if len(hits) == 1:
        return hits[0]
    hits = [v for v in rows if R.vendor_kind(ref) == v.kind and ref.lower() in R.VENDOR_KINDS.get(v.kind, "").lower()]
    if len(hits) == 1:
        return hits[0]
    raise ShaadiError(f"Vendor '{ref}' nahi mila.")


def _clean_schedule(rows: list | None, quote: float | None, old: list | None = None) -> list[dict]:
    old_by = {x.get("id"): x for x in old or []}
    out = []
    for i, r in enumerate(rows or []):
        if not isinstance(r, dict):
            continue
        amt = _num(r.get("amount"))
        if amt is None and _num(r.get("pct")) is not None and quote is not None:
            amt = r2(quote * _num(r.get("pct")) / 100)
        if amt is None or amt <= 0:
            raise ShaadiError(f"Payment '{r.get('label') or i + 1}' ki rakam chahiye (vendor ne jo kaha). Hoola rakam nahi gadhta.")
        due = _date(r.get("due"))
        if r.get("due") and not due:
            raise ShaadiError(f"Tareekh samajh nahi aayi: {r.get('due')}")
        prev = old_by.get(r.get("id")) or {}
        if prev.get("status") in ("approved", "paid") and abs((prev.get("amount") or 0) - amt) > 0.5:
            raise ShaadiError(f"'{prev.get('label')}' pe aapki haan ho chuki hai; uski rakam ab nahi badlegi.")
        out.append({"id": r.get("id") or "p" + secrets.token_hex(3), "label": R.clean(r.get("label"), 60) or f"Kist {i + 1}",
                    "amount": r2(amt), "due": _iso(due), "status": prev.get("status", "planned"), "paid_amount": prev.get("paid_amount"),
                    "ref": prev.get("ref", ""), "paid_on": prev.get("paid_on", "")})
    if quote is not None and out and sum(x["amount"] for x in out) > quote + 0.5:
        raise ShaadiError(f"Kiston ka jod {_rs(sum(x['amount'] for x in out))} quote {_rs(quote)} se zyada hai.")
    return out


def vendor_add(tenant: str, wedding: str, name: str = "", kind: str = "", phone: str = "", functions: list | None = None,
               quote=None, quote_note: str = "", per_plate=None, plates=None, schedule: list | None = None, notes: str = "",
               vendor: str = "", actor: str = "owner") -> dict:
    """The quote is HER word (what the vendor quoted). Per-plate x plates is worked out only from her two numbers."""
    ph = R.phone(phone)
    if phone and not ph:
        raise ShaadiError("Vendor ka phone 10 ankon ka chahiye.")
    p = R.dowry_problem(f"{name} {quote_note} {notes}")
    if p:
        raise ShaadiError(p)
    warns = []
    with _LOCK, session() as s:
        w = _wed(s, tenant, wedding)
        fns = _j(w.functions, [])
        v = _vendor(s, tenant, w, vendor) if vendor else None
        if v is None:
            if not name:
                raise ShaadiError("Vendor ka naam chahiye.")
            v = ShaadiVendor(tenant_id=tenant, wedding_id=w.id, name=R.clean(name, 120), kind=R.vendor_kind(kind or name))
            s.add(v)
        elif v.status in ("booked", "done") and (quote is not None or per_plate is not None or plates is not None):
            raise ShaadiError(f"{v.name} book ho chuka; daam badalna hai to naya quote card (vendor ka likhit quote) chahiye: "
                              "pehle status 'quoted' karne ke liye owner se baat karein.")
        if name and v.name != R.clean(name, 120):
            v.name = R.clean(name, 120)
        if kind:
            v.kind = R.vendor_kind(kind)
        if ph:
            v.phone = ph
        if functions is not None:
            ids = []
            for x in functions:
                f = _fn(fns, str(x))
                if f is None and str(x).lower() not in ("sab", "all"):
                    raise ShaadiError(f"Function '{x}' nahi mila.")
                if f:
                    ids.append(f["id"])
            v.functions = _dump(ids)
        if per_plate not in (None, ""):
            v.per_plate = _num(per_plate)
        if plates not in (None, ""):
            v.plates = _int(plates)
        if quote not in (None, ""):
            q = _num(quote)
            if q is None or q <= 0:
                raise ShaadiError("Quote ki rakam sahi nahi.")
            v.quote = r2(q)
        elif v.quote is None and v.per_plate and v.plates:
            v.quote = r2(v.per_plate * v.plates)
            warns.append(f"Quote = {_rs(v.per_plate)} x {v.plates} plate = {_rs(v.quote)} (aapke diye number se). GST alag hai ya shamil, vendor se likhit lijiye.")
        if quote_note:
            v.quote_note = R.clean(quote_note, 500)
        if notes:
            v.notes = R.clean(notes, 500)
        if schedule is not None:
            v.schedule = _dump(_clean_schedule(schedule, v.quote, _j(v.schedule, [])))
        s.flush()
        # clashes: same kind already booked for the same function; same vendor (phone) booked elsewhere on the same dates
        mine = set(_j(v.functions, [])) or {f["id"] for f in fns}
        for o in s.scalars(select(ShaadiVendor).where(ShaadiVendor.wedding_id == w.id, ShaadiVendor.kind == v.kind,
                                                     ShaadiVendor.status.in_(("booked", "waiting")))).all():
            if o.id != v.id and (set(_j(o.functions, [])) or {f["id"] for f in fns}) & mine and v.kind not in ("other", "transport"):
                warns.append(f"{R.VENDOR_KINDS.get(v.kind)} ke liye {o.name} pehle se {o.status} hai.")
        if v.phone:
            my_dates = {f.get("date") for f in fns if f["id"] in mine and f.get("date")}
            for o in s.scalars(select(ShaadiVendor).where(ShaadiVendor.tenant_id == tenant, ShaadiVendor.phone == v.phone,
                                                         ShaadiVendor.status == "booked")).all():
                if o.wedding_id == w.id:
                    continue
                ow = s.get(ShaadiWedding, o.wedding_id)
                if ow is None or ow.status != "planning":
                    continue
                ofn = _j(ow.functions, [])
                odates = {f.get("date") for f in ofn if (not _j(o.functions, []) or f["id"] in _j(o.functions, [])) and f.get("date")}
                both = sorted(my_dates & odates)
                if both:
                    warns.append(f"{v.name} {ow.number} mein bhi {', '.join(both)} ko book hai.")
        view = _vendor_view(s, v, fns)
    return {"vendor": view, "warnings": warns, "say": f"{view['name']} ({view['kind_label']}): {view['quote_text']}."
            + (" Dhyaan: " + " ".join(warns) if warns else "")}


def vendors(tenant: str, wedding: str) -> dict:
    with session() as s:
        w = _wed(s, tenant, wedding)
        fns = _j(w.functions, [])
        rows = [_vendor_view(s, v, fns) for v in s.scalars(select(ShaadiVendor).where(ShaadiVendor.wedding_id == w.id)).all()]
    by_kind: dict[str, list] = {}
    for r in rows:
        by_kind.setdefault(r["kind"], []).append(r)
    compare = {k: sorted(v, key=lambda x: (x["quote"] is None, x["quote"] or 0)) for k, v in by_kind.items() if len(v) > 1}
    return {"vendors": rows, "compare": {k: [{"name": x["name"], "quote": x["quote_text"], "note": x["quote_note"], "status": x["status"]}
                                            for x in v] for k, v in compare.items()}}


# ---------------------------------------------------------------- cards (her tap)
def _card_view(c: ShaadiCard) -> dict:
    return {"id": c.id, "number": c.number, "kind": c.kind, "show": _j(c.show, {}), "amount": c.amount, "amount_text": _rs(c.amount) if c.amount else "",
            "money": c.money, "status": c.status, "hash": c.hash, "asked_by": c.asked_by, "result": _j(c.result, {}), "reason": c.reason,
            "created_at": _aware(c.created_at).isoformat() if c.created_at else ""}


def _fn_line(fns: list[dict], ids: list[str]) -> list[str]:
    pick = [f for f in fns if not ids or f["id"] in ids]
    return [f"{f['name']} {f.get('date') or '(tareekh baaki)'}" for f in pick]


def _show(s, tenant: str, kind: str, params: dict) -> tuple[dict, float]:
    """Exactly what the card shows, rebuilt from the data as it is NOW (so a change after the card makes it stale)."""
    w = s.get(ShaadiWedding, params["wedding"])
    if w is None:
        raise ShaadiError("Shaadi nahi mili.")
    fns = _j(w.functions, [])
    base = {"wedding": w.number, "title": w.title}
    if kind in ("vendor_book", "vendor_pay"):
        v = s.get(ShaadiVendor, params["vendor"])
        if v is None or v.wedding_id != w.id:
            raise ShaadiError("Vendor nahi mila.")
        sch = _j(v.schedule, [])
        if kind == "vendor_book":
            adv = next((x for x in sch if "advance" in x["label"].lower() or "booking" in x["label"].lower()), sch[0] if sch else None)
            show = {**base, "vendor": v.name, "kind": R.VENDOR_KINDS.get(v.kind, v.kind), "phone": v.phone,
                    "functions": _fn_line(fns, _j(v.functions, [])), "quote": v.quote, "quote_note": v.quote_note,
                    "per_plate": v.per_plate, "plates": v.plates,
                    "schedule": [{"label": x["label"], "amount": x["amount"], "due": x["due"]} for x in sch],
                    "advance": {"label": adv["label"], "amount": adv["amount"], "due": adv["due"]} if adv else None}
            return show, (adv["amount"] if adv else 0.0)
        item = next((x for x in sch if x["id"] == params.get("item")), None)
        if item is None:
            raise ShaadiError("Yeh kist nahi mili.")
        show = {**base, "vendor": v.name, "kind": R.VENDOR_KINDS.get(v.kind, v.kind), "phone": v.phone, "item": item["label"],
                "amount": item["amount"], "due": item["due"], "quote": v.quote,
                "paid_before": sum(x.get("paid_amount") or 0 for x in sch if x.get("status") == "paid")}
        return show, item["amount"]
    if kind == "spend":
        return {**base, "what": params["what"], "head": params.get("head", ""), "function": params.get("function", ""),
                "payee": params.get("payee", ""), "amount": params["amount"]}, params["amount"]
    if kind in ("client_quote", "client_pay"):
        st = _st(s, tenant)
        q = _quote_numbers(st, w, params)
        if kind == "client_quote":
            return {**base, "client": w.client, "client_phone": w.client_phone, **q}, 0.0
        stg = q["stages"][params["stage"]] if 0 <= params.get("stage", -1) < len(q["stages"]) else None
        if stg is None:
            raise ShaadiError("Yeh stage nahi mila.")
        return {**base, "client": w.client, "client_phone": w.client_phone, "stage": stg["label"], "amount": stg["amount"],
                "total": q["total"]}, 0.0
    if kind == "message":
        rec, skipped = _recipients(s, w, params)
        text = params["text"]
        return {**base, "audience": params.get("audience_say", ""), "text": text, "count": len(rec),
                "recipients": [{"id": r["id"], "name": r["name"], "phone": r["phone"]} for r in rec], "skipped": skipped}, 0.0
    raise ShaadiError("Card ka kind nahi pehchana.")


def _make_card(tenant: str, wedding_id: str, kind: str, params: dict, actor: str, conv: str, key: str = "") -> dict:
    with _LOCK, session() as s:
        show, amt = _show(s, tenant, kind, {**params, "wedding": wedding_id})
        h = R.digest({"kind": kind, "show": show})
        for c in s.scalars(select(ShaadiCard).where(ShaadiCard.wedding_id == wedding_id, ShaadiCard.kind == kind,
                                                   ShaadiCard.status == "waiting")).all():
            p = _j(c.ref, {})
            if key and p.get("_key") == key:
                if c.hash == h:
                    return {"card": _card_view(c), "existing": True, "status": "needs_tap",
                            "say": f"{c.number} ka card pehle se aapki haan ka intezaar kar raha hai."}
                c.status = "stale"
        st = _st(s, tenant)
        st.card_seq += 1
        c = ShaadiCard(tenant_id=tenant, wedding_id=wedding_id, number=f"SC-{st.card_seq:04d}", kind=kind,
                       ref=_dump({**params, "wedding": wedding_id, "_key": key}), show=_dump(show), amount=r2(amt or 0),
                       money=kind in MONEY_KINDS and (amt or 0) > 0, hash=h, asked_by=actor_name(tenant, actor))
        s.add(c)
        s.flush()
        view = _card_view(c)
    return {"card": view, "status": "needs_tap"}


def _say_card(view: dict) -> str:
    sh = view["show"]
    k = view["kind"]
    if k == "vendor_book":
        adv = sh.get("advance")
        return (f"{view['number']}: {sh['vendor']} ({sh['kind']}) {_rs(sh.get('quote'))} book karna" +
                (f", advance {_rs(adv['amount'])}" if adv else "") + " - aapki haan ka card bheja hai.")
    if k == "vendor_pay":
        return f"{view['number']}: {sh['vendor']} ko '{sh['item']}' {_rs(sh['amount'])} dena - aapki haan ka card bheja hai."
    if k == "spend":
        return f"{view['number']}: {sh['what']} {_rs(sh['amount'])} kharch - aapki haan ka card bheja hai."
    if k == "client_quote":
        return f"{view['number']}: client {sh.get('client')} ko quote {_rs(sh.get('total'))} - aapki haan ka card bheja hai."
    if k == "client_pay":
        return f"{view['number']}: client se '{sh['stage']}' {_rs(sh['amount'])} maangna - aapki haan ka card bheja hai."
    return f"{view['number']}: {sh.get('count')} logon ko message - aapki haan ka card bheja hai (bhejengi aap)."


def _card_out(tenant: str, out: dict, conv: str) -> dict:
    if not out.get("existing"):
        _publish(tenant, "shaadi_card", {"card": out["card"], "say": _say_card(out["card"])}, conv)
    out.setdefault("say", _say_card(out["card"]))
    return out


def book_request(tenant: str, wedding: str, vendor: str, actor: str = "owner", conv: str = "") -> dict:
    with session() as s:
        w = _wed(s, tenant, wedding)
        v = _vendor(s, tenant, w, vendor)
        if v.quote is None:
            raise ShaadiError(f"{v.name} ka quote (vendor ne jo daam bataya) pehle likhiye. Hoola daam nahi gadhta.")
        if v.status in ("booked", "done"):
            raise ShaadiError(f"{v.name} pehle se book hai.")
        wid, vid = w.id, v.id
    out = _make_card(tenant, wid, "vendor_book", {"vendor": vid}, actor, conv, key=f"book:{vid}")
    with _LOCK, session() as s:
        v = s.get(ShaadiVendor, vid)
        v.status = "waiting"
    return _card_out(tenant, out, conv)


def pay_request(tenant: str, wedding: str, vendor: str, item: str = "", actor: str = "owner", conv: str = "") -> dict:
    with session() as s:
        w = _wed(s, tenant, wedding)
        v = _vendor(s, tenant, w, vendor)
        if v.status not in ("booked", "done"):
            raise ShaadiError(f"{v.name} abhi book nahi hua; pehle booking ka card.")
        sch = _j(v.schedule, [])
        it = next((x for x in sch if item and (x["id"] == item or item.lower() in x["label"].lower())), None)
        if it is None:
            it = next((x for x in sch if x["status"] == "planned"), None) if not item else None
        if it is None:
            raise ShaadiError("Kaunsi kist? (vendor ki payment schedule mein jodiye)")
        if it["status"] in ("approved", "paid"):
            raise ShaadiError(f"'{it['label']}' par pehle hi haan / payment ho chuki.")
        wid, vid, iid = w.id, v.id, it["id"]
    return _card_out(tenant, _make_card(tenant, wid, "vendor_pay", {"vendor": vid, "item": iid}, actor, conv, key=f"pay:{vid}:{iid}"), conv)


def spend_request(tenant: str, wedding: str, what: str, amount, head: str = "", fn: str = "", payee: str = "",
                  actor: str = "owner", conv: str = "") -> dict:
    amt = _num(amount)
    if amt is None or amt <= 0:
        raise ShaadiError("Kharche ki rakam chahiye (bill / quote se).")
    p = R.dowry_problem(f"{what} {payee}")
    if p:
        raise ShaadiError(p)
    with session() as s:
        w = _wed(s, tenant, wedding)
        f = _fn(_j(w.functions, []), fn) if fn else None
        if fn and f is None:
            raise ShaadiError(f"Function '{fn}' nahi mila.")
        wid = w.id
    params = {"what": R.clean(what, 200), "amount": r2(amt), "head": head if head in R.HEADS else ("other" if head else ""),
              "function": f["id"] if f else "", "payee": R.clean(payee, 80)}
    return _card_out(tenant, _make_card(tenant, wid, "spend", params, actor, conv, key=f"spend:{params['what']}:{params['amount']}"), conv)


def _quote_numbers(st: ShaadiSettings, w: ShaadiWedding, params: dict) -> dict:
    lines = []
    blanks = []
    for x in params.get("lines") or []:
        amt = _num(x.get("amount")) if isinstance(x, dict) else None
        nm = R.clean(x.get("name"), 120) if isinstance(x, dict) else ""
        if not nm or amt is None:
            blanks.append(f"'{nm or '?'}' ki rakam")
            continue
        lines.append({"name": nm, "amount": r2(amt)})
    budget = _j(w.budget, {}).get("total")
    fee = 0.0
    fee_note = ""
    if st.fee_mode == "percent":
        if not budget:
            blanks.append("shaadi ka budget (fee % ke liye)")
        else:
            fee = r2(budget * st.fee_value / 100)
            fee_note = f"{st.fee_value:g}% of budget {_rs(budget)}"
    elif st.fee_mode == "flat":
        fee = r2(st.fee_value)
        fee_note = "flat planning fee"
    if st.fee_mode == "none" and not lines:
        blanks.append("planning fee (Niyam mein fee_mode) ya quote ki lines")
    sub = r2(fee + sum(x["amount"] for x in lines))
    gst = r2(sub * st.gst_pct / 100) if st.gstin else 0.0
    total = r2(sub + gst)
    tmpl = _j(st.stages, []) or DEFAULT_STAGES
    amounts, acc = [], 0.0
    for i, x in enumerate(tmpl):
        a = r2(total - acc) if i == len(tmpl) - 1 else r2(total * x["pct"] / 100)
        acc += a
        amounts.append({"label": x["label"], "pct": x["pct"], "amount": a})
    valid = (clock().date() + timedelta(days=st.validity_days)).isoformat()
    return {"fee": fee, "fee_note": fee_note, "lines": lines, "subtotal": sub, "gst_pct": st.gst_pct if st.gstin else 0, "gst": gst,
            "total": total, "stages": amounts, "valid_until": valid if not params.get("valid_until") else params["valid_until"],
            "blanks": blanks, "gst_note": "" if st.gstin else "GSTIN nahi diya: GST nahi lagaya."}


def quote_request(tenant: str, wedding: str, lines: list | None = None, actor: str = "owner", conv: str = "") -> dict:
    with session() as s:
        w = _wed(s, tenant, wedding)
        st = _st(s, tenant)
        if st.mode != "agency":
            raise ShaadiError("Client quote sirf agency mode mein (Niyam: mode agency).")
        if not w.client:
            raise ShaadiError("Client ka naam pehle likhiye.")
        valid = (today(tenant) + timedelta(days=st.validity_days)).isoformat()
        q = _quote_numbers(st, w, {"lines": lines or [], "valid_until": valid})
        if q["blanks"]:
            raise ShaadiError("Quote ke liye yeh chahiye: " + "; ".join(q["blanks"]) + ". Hoola daam nahi gadhta.")
        wid = w.id
    return _card_out(tenant, _make_card(tenant, wid, "client_quote", {"lines": lines or [], "valid_until": valid}, actor, conv,
                                        key="client_quote"), conv)


def client_pay_request(tenant: str, wedding: str, stage: int = 0, actor: str = "owner", conv: str = "") -> dict:
    with session() as s:
        w = _wed(s, tenant, wedding)
        acc = [c for c in s.scalars(select(ShaadiCard).where(ShaadiCard.wedding_id == w.id, ShaadiCard.kind == "client_quote",
                                                            ShaadiCard.status == "approved")).all()]
        if not acc:
            raise ShaadiError("Pehle client quote par aapki haan chahiye.")
        params = {**_j(acc[-1].ref, {}), "stage": int(stage)}
        params.pop("_key", None)
        wid = w.id
    return _card_out(tenant, _make_card(tenant, wid, "client_pay", params, actor, conv, key=f"client_pay:{stage}"), conv)


# ---- messages to guests / vendors (always a card; she sends)
def _recipients(s, w: ShaadiWedding, params: dict) -> tuple[list[dict], list[str]]:
    fns = _j(w.functions, [])
    out, skipped = [], []
    if params.get("to") == "vendors":
        for v in s.scalars(select(ShaadiVendor).where(ShaadiVendor.wedding_id == w.id, ShaadiVendor.status.in_(("booked", "done")))).all():
            if params.get("kind") and v.kind != params["kind"]:
                continue
            if not v.phone:
                skipped.append(f"{v.name} (phone nahi)")
                continue
            out.append({"id": v.id, "name": v.name, "phone": v.phone})
        return out, skipped
    f = _fn(fns, params.get("fn") or "") if params.get("fn") else None
    ids = set(params.get("guests") or [])
    for g in s.scalars(select(ShaadiGuest).where(ShaadiGuest.wedding_id == w.id)).all():
        if ids and g.id not in ids:
            continue
        if params.get("side") and g.side != params["side"]:
            continue
        if f and not _invited(g, f["id"]):
            continue
        if params.get("pending") and f and _j(g.rsvp, {}).get(f["id"], "pending") != "pending":
            continue
        if params.get("pending") and not f and all(v != "pending" for v in _j(g.rsvp, {}).values()) and _j(g.rsvp, {}):
            continue
        if g.no_messages:
            skipped.append(f"{g.name} (message nahi chahte)")
            continue
        if not g.phone:
            skipped.append(f"{g.name} (phone nahi)")
            continue
        out.append({"id": g.id, "name": g.name, "phone": g.phone})
    out.sort(key=lambda x: x["name"].lower())
    return out, skipped


INVITE = {
    "hinglish": "Namaste {name} ji,\n{bride} aur {groom} ki shaadi mein aap saparivaar saadar aamantrit hain.\n{lines}\nAapka aana humare liye "
                "aashirwad hoga. Kripya bataiye aap aa payenge ya nahi.\n- {host}",
    "hindi": "आदरणीय {name} जी,\n{bride} और {groom} के शुभ विवाह में आप सपरिवार सादर आमंत्रित हैं।\n{lines}\nकृपया अपने आने की सूचना दें।\n- {host}",
    "marathi": "आदरणीय {name},\n{bride} आणि {groom} यांच्या शुभविवाहास आपण सहकुटुंब अवश्य यावे.\n{lines}\nकृपया येण्याविषयी कळवावे.\n- {host}",
    "english": "Dear {name},\nYou and your family are warmly invited to the wedding of {bride} and {groom}.\n{lines}\nKindly let us know if you "
               "can join.\n- {host}",
}
RSVP_ASK = {"hinglish": "Namaste {name} ji, {fn} ({when}) mein aap aa payenge? Kitne log? Bas 'haan / nahi' likh dijiye. - {host}",
            "english": "Dear {name}, will you be able to join us for {fn} ({when})? How many of you? Just reply yes / no. - {host}"}


def invite_text(tenant: str, wedding: str, kind: str = "invite", fn: str = "", lang: str = "hinglish", host: str = "") -> dict:
    """Text from the wedding's REAL facts only. Anything missing stays a [blank] (and blocks sending)."""
    with session() as s:
        w = _wed(s, tenant, wedding)
        fns = _j(w.functions, [])
        st = _st(s, tenant)
        pick = [f for f in fns if not fn or f is _fn(fns, fn)]
        if fn and not pick:
            raise ShaadiError(f"Function '{fn}' nahi mila.")
        hostname = R.clean(host, 80) or (w.client if st.mode == "agency" else "") or "[aapka naam]"

        def when(f):
            parts = [f.get("date") or "[tareekh]", f.get("start") or "[samay]", f.get("venue") or w.venue or "[jagah]"]
            return ", ".join(parts)
        lines = "\n".join(f"• {f['name']}: {when(f)}" for f in pick)
        if kind == "rsvp":
            tmpl = RSVP_ASK.get(lang, RSVP_ASK["hinglish"])
            f = pick[0] if pick else None
            text = tmpl.format(name="{name}", fn=f["name"] if f else "shaadi", when=when(f) if f else "[tareekh]", host=hostname)
        else:
            tmpl = INVITE.get(lang, INVITE["hinglish"])
            text = tmpl.format(name="{name}", bride=w.bride or "[dulhan]", groom=w.groom or "[dulha]", lines=lines, host=hostname)
    blanks = re.findall(r"\[[^\]]+\]", text)
    return {"text": text, "blanks": blanks, "say": "Draft taiyaar" + (f"; khaali: {', '.join(sorted(set(blanks)))}" if blanks else "") + ".",
            "note": "Card ka design chahiye to Brand Sahayak (action invite_design)."}


def message_request(tenant: str, wedding: str, text: str = "", kind: str = "invite", to: str = "guests", fn: str = "", side: str = "",
                    pending: bool = False, guests_: list | None = None, lang: str = "hinglish", vendor_kind: str = "",
                    actor: str = "owner", conv: str = "") -> dict:
    """Every message to guests / vendors is a card; on her tap she gets one WhatsApp / SMS draft link per person and sends herself."""
    if not text:
        text = invite_text(tenant, wedding, "rsvp" if kind == "rsvp" else "invite", fn, lang)["text"]
    text = R.clean(text, 2000)
    p = R.dowry_problem(text)
    if p:
        raise ShaadiError(p)
    blanks = [b for b in re.findall(r"\[[^\]]+\]", text) if b != "[hataya]"]
    if blanks:
        raise ShaadiError("Message mein khaali jagah hai: " + ", ".join(sorted(set(blanks))) + ". Pehle bhariye.")
    with session() as s:
        w = _wed(s, tenant, wedding)
        wid = w.id
        aud = []
        if to == "vendors":
            aud.append("booked vendors" + (f" ({vendor_kind})" if vendor_kind else ""))
        else:
            if side:
                aud.append({"bride": "dulhan ki taraf", "groom": "dulhe ki taraf"}.get(side, side))
            if fn:
                aud.append(f"{fn} ke mehmaan")
            if pending:
                aud.append("jinka jawab baaki")
            if guests_:
                aud.append(f"{len(guests_)} chune hue")
    params = {"text": text, "to": to, "fn": fn, "side": side, "pending": bool(pending), "guests": list(guests_ or []),
              "kind": R.vendor_kind(vendor_kind) if vendor_kind else "", "audience_say": ", ".join(aud) or "saare mehmaan"}
    with session() as s:
        rec, _sk = _recipients(s, s.get(ShaadiWedding, wid), params)
    if not rec:
        raise ShaadiError("Kisi ko bhejne layak phone nahi mila.")
    return _card_out(tenant, _make_card(tenant, wid, "message", params, actor, conv, key=f"msg:{R.digest(params)[:16]}"), conv)


def challenge(tenant: str, card_id: str) -> dict:
    """approvekey's message format (kind 'shaadi'): her phone signs this money card as it is NOW."""
    from .. import approvekey as ak
    with session() as s:
        c = s.get(ShaadiCard, card_id)
        if c is None or c.tenant_id != tenant or c.status != "waiting":
            raise ShaadiError("Yeh card ab baaki nahi hai.")
        view = _card_view(c)
    nonce = secrets.token_urlsafe(18)
    exp = datetime.now(timezone.utc) + ak.CHALLENGE_TTL
    msg = f"hoola-approve|v1|{tenant}|shaadi|{card_id}|{view['hash']}|{nonce}|{int(exp.timestamp())}"
    with session() as s:
        ch = ak.ApproveChallenge(tenant_id=tenant, kind="shaadi", target=card_id, digest=view["hash"], message=msg, expires_at=exp)
        s.add(ch)
        s.flush()
        chid = ch.id
    sh = view["show"]
    return {"challenge_id": chid, "message": msg, "expires_at": exp.isoformat(),
            "show": {"title": f"Shaadi {view['number']}", "amount": view["amount_text"], "payee": sh.get("vendor") or sh.get("payee") or ""}}


def _verify_fingerprint(tenant: str, card_id: str, digest: str, proof: dict | None) -> None:
    from .. import approvekey as ak
    if not proof:
        raise ShaadiError("fingerprint_needed")
    with session() as s:
        ak._settle(s, tenant)
        ok = ak._check(s, tenant, "shaadi", card_id, proof, digest)
    if not ok:
        raise ShaadiError("fingerprint_invalid")


def decide(tenant: str, card_id: str, yes: bool, card_hash: str = "", reason: str = "", proof: dict | None = None,
           by_tap: bool = False, conv: str = "") -> dict:
    if not by_tap:
        raise ShaadiError("Yeh faisla sirf aapke tap se (app ka card).")
    with session() as s:
        c = s.get(ShaadiCard, card_id)
        if c is None or c.tenant_id != tenant:
            raise ShaadiError("Card nahi mila.")
        if c.status != "waiting":
            return {"card": _card_view(c), "say": f"{c.number} pehle hi '{c.status}' hai."}
        params = _j(c.ref, {})
        kind = c.kind
        if not card_hash or card_hash != c.hash:
            raise ShaadiError("Card badal gaya hai; naya card dekhiye.")
        try:
            show, _amt = _show(s, tenant, kind, params)
        except ShaadiError:
            show = None
        stale = show is None or R.digest({"kind": kind, "show": show}) != c.hash
        if stale:
            c.status = "stale"
        money = c.money
    if stale:
        raise ShaadiError("Card banne ke baad cheezein badal gayi (daam / kist / mehmaan); naya card banwaiye.")
    if yes and money:
        from .. import approvekey as ak
        if ak.active_on(tenant):
            _verify_fingerprint(tenant, card_id, card_hash, proof)
    result: dict[str, Any] = {}
    with _LOCK, session() as s:
        c = s.get(ShaadiCard, card_id)
        c.status = "approved" if yes else "rejected"
        c.reason = R.clean(reason, 300)
        c.decided_at = clock()
        w = s.get(ShaadiWedding, params["wedding"])
        if kind in ("vendor_book", "vendor_pay"):
            v = s.get(ShaadiVendor, params["vendor"])
            if kind == "vendor_book":
                v.status = "booked" if yes else "quoted"
                if yes:
                    sh = show or {}
                    adv = sh.get("advance")
                    text = (f"Namaste {v.name} ji, {w.title} ke liye aapki booking pakki: " + ", ".join(sh.get("functions", []))
                            + f". Tay rakam {_rs(v.quote)}." + (f" Advance {_rs(adv['amount'])} hum bhej rahe hain." if adv else ""))
                    result["draft"] = {"text": text, "wa": _wa(v.phone, text), "sms": _sms(v.phone, text)}
                    if adv:
                        sch = _j(v.schedule, [])
                        for x in sch:
                            if x["label"] == adv["label"] and x["status"] == "planned":
                                x["status"] = "approved"
                                break
                        v.schedule = _dump(sch)
            else:
                sch = _j(v.schedule, [])
                for x in sch:
                    if x["id"] == params["item"]:
                        x["status"] = "approved" if yes else "planned"
                v.schedule = _dump(sch)
            if yes:
                result["pay_note"] = "Hoola paise nahi bhejta: aap apne UPI / bank se dijiye, phir 'Maine de diya' (UTR) dabaiye."
            _log(w, f"{c.number} {'haan' if yes else 'nahi'}: {v.name}", "Owner")
        elif kind == "spend":
            _log(w, f"{c.number} {'haan' if yes else 'nahi'}: {params['what']} {_rs(params['amount'])}", "Owner")
        elif kind == "message" and yes:
            sh = show or {}
            links = []
            for r in sh.get("recipients", []):
                t = sh["text"].replace("{name}", r["name"].split()[0] if r["name"] else "")
                links.append({"id": r["id"], "name": r["name"], "wa": _wa(r["phone"], t), "sms": _sms(r["phone"], t)})
            result["links"] = links
            result["note"] = "Har naam par tap karke WhatsApp / SMS se aap khud bhejiye; Hoola apne aap nahi bhejta."
        result_kind = kind
        c.result = _dump(result)
        view = _card_view(c)
        wnum, wtitle, cphone = w.number, w.title, w.client_phone
    if yes and result_kind == "client_quote":
        pdf = quote_pdf(tenant, card_id)
        f = _store(tenant, f"{view['number']}_wedding_quote.pdf", pdf, "application/pdf", conv, f"{wtitle} quote")
        sh = view["show"]
        text = f"Namaste {sh.get('client')} ji, {wtitle} ki planning ka quote {_rs(sh.get('total'))} (PDF saath mein), {sh.get('valid_until')} tak."
        result.update({"file_id": f.get("file_id", ""), "draft": {"text": text, "wa": _wa(cphone, text), "sms": _sms(cphone, text)}})
    if yes and result_kind == "client_pay":
        sh = view["show"]
        up = _upi(tenant, sh["amount"], f"{wnum} {sh['stage']}")
        text = f"Namaste {sh.get('client')} ji, {wtitle}: '{sh['stage']}' ki rakam {_rs(sh['amount'])}." + (f" UPI: {up}" if up else "")
        result.update({"upi": up, "draft": {"text": text, "wa": _wa(cphone, text), "sms": _sms(cphone, text)},
                       "note": "Client apne app se aapko bhejega; Hoola paise nahi leta."})
    if result_kind in ("client_quote", "client_pay"):
        with session() as s:
            c = s.get(ShaadiCard, card_id)
            c.result = _dump(result)
            view = _card_view(c)
    if result_kind == "message" and yes:
        with _LOCK, session() as s:
            for r in (show or {}).get("recipients", []):
                g = s.get(ShaadiGuest, r["id"])
                if g is not None:
                    sent = _j(g.invite_sent, [])
                    sent.append(card_id)
                    g.invite_sent = _dump(sent[-20:])
    say = f"{view['number']}: " + ("haan, " + {"vendor_book": "booking pakki; vendor ko message draft aur advance aap bhejiye.",
                                               "vendor_pay": "theek; payment aap apne UPI / bank se kijiye, phir 'Maine de diya'.",
                                               "spend": "kharch maana.", "client_quote": "quote PDF + client ke liye message draft taiyaar.",
                                               "client_pay": "client ke liye message draft + UPI link (aap bhejiye).",
                                               "message": f"{len(result.get('links', []))} drafts taiyaar, aap bhejiye."}[result_kind]
                                    if yes else "nahi.")
    return {"card": view, "say": say, **result}


def mark_paid(tenant: str, wedding: str, vendor: str, item: str = "", amount=None, ref: str = "", actor: str = "owner") -> dict:
    """Her word that SHE paid (after her Haan). Hoola records it; never moves money."""
    if not is_owner(actor):
        raise ShaadiError("Payment ka hisaab sirf owner likhti hain.")
    with _LOCK, session() as s:
        w = _wed(s, tenant, wedding)
        v = _vendor(s, tenant, w, vendor)
        sch = _j(v.schedule, [])
        it = next((x for x in sch if item and (x["id"] == item or item.lower() in x["label"].lower())), None) or \
            next((x for x in sch if x["status"] == "approved"), None)
        if it is None:
            raise ShaadiError("Is vendor ki koi kist aapki haan ke baad baaki nahi.")
        if it["status"] != "approved":
            raise ShaadiError(f"'{it['label']}' par pehle aapki haan (card) chahiye.")
        amt = _num(amount) if amount not in (None, "") else it["amount"]
        it.update({"status": "paid", "paid_amount": r2(amt), "ref": R.clean(ref, 40)[-12:], "paid_on": today(tenant).isoformat()})
        v.schedule = _dump(sch)
        if v.quote is not None and sum(x.get("paid_amount") or 0 for x in sch if x["status"] == "paid") >= v.quote - 0.5:
            v.status = "done"
        _log(w, f"{v.name}: {it['label']} {_rs(amt)} diya", "Owner")
        view = _vendor_view(s, v, _j(w.functions, []))
    short = " (kist se kam)" if amt < it["amount"] - 0.5 else ""
    return {"vendor": view, "say": f"{view['name']}: {it['label']} {_rs(amt)} diya darj{short}. Baaki {_rs(view['due'])}."}


def spend_paid(tenant: str, card_id: str, ref: str = "", actor: str = "owner") -> dict:
    if not is_owner(actor):
        raise ShaadiError("Sirf owner.")
    with _LOCK, session() as s:
        c = s.get(ShaadiCard, card_id)
        if c is None or c.tenant_id != tenant or c.kind != "spend":
            raise ShaadiError("Kharche ka card nahi mila.")
        if c.status != "approved":
            raise ShaadiError("Pehle is kharche par aapki haan chahiye.")
        c.status = "done"
        r = _j(c.result, {})
        r["ref"] = R.clean(ref, 40)[-12:]
        r["paid_on"] = today(tenant).isoformat()
        c.result = _dump(r)
        return {"card": _card_view(c), "say": f"{c.number} diya darj."}


def cards(tenant: str, wedding: str = "", status: str = "") -> list[dict]:
    with session() as s:
        q = select(ShaadiCard).where(ShaadiCard.tenant_id == tenant)
        if wedding:
            q = q.where(ShaadiCard.wedding_id == _wed(s, tenant, wedding).id)
        if status:
            q = q.where(ShaadiCard.status == status)
        return [_card_view(c) for c in s.scalars(q.order_by(ShaadiCard.created_at.desc())).all()][:100]


def quote_pdf(tenant: str, card_id: str) -> bytes:
    with session() as s:
        c = s.get(ShaadiCard, card_id)
        if c is None or c.tenant_id != tenant or c.kind != "client_quote":
            raise ShaadiError("Quote nahi mila.")
        st = _st(s, tenant)
        sh = _j(c.show, {})
        final = c.status == "approved"
        blocks = [("h1", st.firm or "Wedding Planner"), ("p", (f"GSTIN {st.gstin} · " if st.gstin else "") + (st.phone or "")),
                  ("h", f"Quotation {c.number} - {sh.get('title')}"), ("p", f"Client: {sh.get('client')} · Valid till {sh.get('valid_until')}")]
        rows = [["Item", "Amount"]]
        if sh.get("fee"):
            rows.append([f"Planning fee ({sh.get('fee_note')})", _pdf_rs(sh.get("fee"))])
        for x in sh.get("lines", []):
            rows.append([x["name"], _pdf_rs(x["amount"])])
        rows.append(["Subtotal", _pdf_rs(sh.get("subtotal"))])
        if sh.get("gst"):
            rows.append([f"GST {sh.get('gst_pct'):g}% (SAC 998596, as applicable)", _pdf_rs(sh.get("gst"))])
        rows.append(["TOTAL", _pdf_rs(sh.get("total"))])
        blocks.append(("table", rows, [330, 120]))
        srows = [["Stage", "%", "Amount"]] + [[x["label"], f"{x['pct']:g}%", _pdf_rs(x["amount"])] for x in sh.get("stages", [])]
        blocks += [("h", "Payment schedule"), ("table", srows, [300, 50, 100]), ("h", "Terms"),
                   ("p", st.terms or "Vendor charges are paid by the client directly to vendors unless listed above. Cancellation and refund "
                                     "terms as agreed in writing before booking. GST as applicable.")]
    return _pdf(f"Quotation {c.number}", blocks, "" if final else "DRAFT - NOT SENT")


# ---------------------------------------------------------------- budget
def budget_set(tenant: str, wedding: str, amount, head: str = "", fn: str = "", actor: str = "owner") -> dict:
    v = _num(amount)
    if v is None or v < 0:
        raise ShaadiError("Budget ki rakam chahiye.")
    with _LOCK, session() as s:
        w = _wed(s, tenant, wedding)
        b = _j(w.budget, {})
        if not head and not fn:
            b["total"] = r2(v)
        else:
            h = head if head in R.HEADS else R.vendor_kind(head) if head else "other"
            if h not in R.HEADS:
                h = "other"
            key = "all"
            if fn:
                f = _fn(_j(w.functions, []), fn)
                if f is None:
                    raise ShaadiError(f"Function '{fn}' nahi mila.")
                key = f["id"]
            b.setdefault("heads", {}).setdefault(key, {})[h] = r2(v)
        w.budget = _dump(b)
    return budget(tenant, wedding)


def budget(tenant: str, wedding: str) -> dict:
    """Planned (her numbers) vs committed (booked vendors' quotes + approved spends) vs paid. Vendors on several functions
    are shared out equally across them (said)."""
    with session() as s:
        w = _wed(s, tenant, wedding)
        fns = _j(w.functions, [])
        b = _j(w.budget, {})
        names = {f["id"]: f["name"] for f in fns}
        names["all"] = "Poori shaadi"
        table: dict[str, dict[str, dict]] = {}

        def cell(k, h):
            return table.setdefault(k, {}).setdefault(h, {"planned": 0.0, "committed": 0.0, "paid": 0.0})
        for k, heads in (b.get("heads") or {}).items():
            for h, amt in heads.items():
                cell(k, h)["planned"] += amt or 0
        vendors = s.scalars(select(ShaadiVendor).where(ShaadiVendor.wedding_id == w.id)).all()
        for v in vendors:
            if v.status not in ("booked", "done") or v.quote is None:
                continue
            ids = _j(v.functions, []) or ["all"]
            paid = sum(x.get("paid_amount") or 0 for x in _j(v.schedule, []) if x.get("status") == "paid")
            h = v.kind if v.kind in R.HEADS else "other"
            for k in ids:
                cell(k, h)["committed"] += v.quote / len(ids)
                cell(k, h)["paid"] += paid / len(ids)
        for c in s.scalars(select(ShaadiCard).where(ShaadiCard.wedding_id == w.id, ShaadiCard.kind == "spend",
                                                   ShaadiCard.status.in_(("approved", "done")))).all():
            p = _j(c.ref, {})
            x = cell(p.get("function") or "all", p.get("head") or "other")
            x["committed"] += c.amount
            if c.status == "done":
                x["paid"] += c.amount
        rows, warns = [], []
        for k, heads in table.items():
            for h, x in heads.items():
                over = x["planned"] and x["committed"] > x["planned"] + 0.5
                rows.append({"function": names.get(k, k), "head": h, "head_label": R.VENDOR_KINDS.get(h, h.title()),
                             "planned": r2(x["planned"]), "committed": r2(x["committed"]), "paid": r2(x["paid"]), "over": bool(over)})
                if over:
                    warns.append(f"{names.get(k, k)} / {R.VENDOR_KINDS.get(h, h)}: {_rs(x['committed'])} pakka, budget {_rs(x['planned'])}.")
        tot = {"planned_total": b.get("total"), "planned_heads": r2(sum(r["planned"] for r in rows)),
               "committed": r2(sum(r["committed"] for r in rows)), "paid": r2(sum(r["paid"] for r in rows))}
        if b.get("total") and tot["committed"] > b["total"]:
            warns.append(f"Pakka kharch {_rs(tot['committed'])} kul budget {_rs(b['total'])} se zyada.")
    rows.sort(key=lambda r: (r["function"], r["head"]))
    return {"rows": rows, **tot, "warnings": warns,
            "say": f"Budget {_rs(b.get('total'))}; pakka {_rs(tot['committed'])}; diya {_rs(tot['paid'])}." + (" " + " ".join(warns[:3]) if warns else "")}


# ---------------------------------------------------------------- shagun / gift register (owner only, never shared)
def gift_add(tenant: str, wedding: str, giver: str, value=None, kind: str = "cash", what: str = "", relation: str = "",
             side: str = "", for_whom: str = "both", fn: str = "", day: str = "", note: str = "", actor: str = "owner") -> dict:
    if not is_owner(actor):
        raise ShaadiError("Shagun register sirf owner ke liye hai.")
    p = R.dowry_problem(f"{what} {note} {relation}")
    if p:
        raise ShaadiError(p)
    if kind not in ("cash", "upi", "cheque", "gold", "item", "other"):
        kind = "item" if what else "cash"
    v = _num(value)
    if kind in ("cash", "upi", "cheque") and v is None:
        raise ShaadiError("Shagun ki rakam chahiye.")
    warns = []
    with _LOCK, session() as s:
        w = _wed(s, tenant, wedding)
        f = _fn(_j(w.functions, []), fn) if fn else None
        d = _date(day) or today(tenant)
        g = ShaadiGift(tenant_id=tenant, wedding_id=w.id, giver=R.clean(giver, 120), relation=R.clean(relation, 60),
                       side=side if side in R.GUEST_SIDES else "common", for_whom=for_whom if for_whom in ("bride", "groom", "both") else "both",
                       function_id=f["id"] if f else "", kind=kind, what=R.clean(what, 300), value=r2(v) if v is not None else None,
                       day=d.isoformat(), note=R.clean(note, 300))
        s.add(g)
        s.flush()
        if kind == "cash":
            same = sum(x.value or 0 for x in s.scalars(select(ShaadiGift).where(ShaadiGift.wedding_id == w.id, ShaadiGift.kind == "cash",
                                                                                 ShaadiGift.day == g.day)).all()
                       if x.giver.lower() == g.giver.lower())
            if same >= R.CASH_LIMIT:
                warns.append(f"{g.giver} se ek din mein {_rs(same)} nakad: Rs 2 lakh ya zyada nakad lena income-tax mein mana hai "
                             "(purana s.269ST); bank / UPI se lijiye. CA se poochhiye.")
        gid = g.id
    return {"gift": {"id": gid, "giver": R.mask_name(R.clean(giver, 120)), "kind": kind, "value": v}, "warnings": warns,
            "say": "Shagun register mein likha." + (" Dhyaan: " + " ".join(warns) if warns else "")}


def gifts(tenant: str, wedding: str, full: bool = False) -> dict:
    """full=True only for HER app screen; the chat model gets masked names. Never shared with anyone else."""
    with session() as s:
        w = _wed(s, tenant, wedding)
        names = {f["id"]: f["name"] for f in _j(w.functions, [])}
        rows = s.scalars(select(ShaadiGift).where(ShaadiGift.wedding_id == w.id).order_by(ShaadiGift.day)).all()
        out = [{"id": g.id, "giver": g.giver if full else R.mask_name(g.giver), "relation": g.relation, "side": g.side,
                "for_whom": g.for_whom, "function": names.get(g.function_id, ""), "kind": g.kind, "what": g.what, "value": g.value,
                "day": g.day, "note": g.note if full else ""} for g in rows]
    by_kind: dict[str, float] = {}
    by_side: dict[str, float] = {}
    for x in out:
        by_kind[x["kind"]] = r2(by_kind.get(x["kind"], 0) + (x["value"] or 0))
        by_side[x["side"]] = r2(by_side.get(x["side"], 0) + (x["value"] or 0))
    return {"gifts": out, "count": len(out), "by_kind": by_kind, "by_side": by_side,
            "total": r2(sum(x["value"] or 0 for x in out)), "private": True,
            "say": f"{len(out)} shagun / tohfe, kul {_rs(sum(x['value'] or 0 for x in out))}. Yeh list sirf aapke liye hai."}


def gift_remove(tenant: str, gift_id: str, actor: str = "owner") -> dict:
    if not is_owner(actor):
        raise ShaadiError("Sirf owner.")
    with _LOCK, session() as s:
        g = s.get(ShaadiGift, gift_id)
        if g is None or g.tenant_id != tenant:
            raise ShaadiError("Entry nahi mili.")
        s.delete(g)
    return {"say": "Hata diya."}


def gift_list_pdf(tenant: str, wedding: str, for_whom: str = "bride") -> bytes:
    """List of presents in the Dowry Prohibition (Maintenance of Lists of Presents) Rules 1985 shape: description, approximate
    value, giver, relationship; signed by bride and bridegroom. Owner's own file."""
    with session() as s:
        w = _wed(s, tenant, wedding)
        rows = [g for g in s.scalars(select(ShaadiGift).where(ShaadiGift.wedding_id == w.id).order_by(ShaadiGift.day)).all()
                if g.for_whom in (for_whom, "both")]
        title = w.title
        bride, groom = w.bride, w.groom
    who = {"bride": f"the bride ({bride})", "groom": f"the bridegroom ({groom})"}.get(for_whom, "the couple")
    t = [["#", "Present (short description)", "Approx. value", "Given by", "Relationship"]]
    for i, g in enumerate(rows, 1):
        t.append([str(i), g.what or {"cash": "Cash", "upi": "UPI transfer", "cheque": "Cheque"}.get(g.kind, g.kind),
                  _pdf_rs(g.value), g.giver, g.relation])
    blocks = [("h1", "List of presents"), ("p", f"{title} - presents received by {who} at the time of marriage"),
              ("p", "Kept as per the Dowry Prohibition (Maintenance of Lists of Presents to the Bride and Bridegroom) Rules, 1985."),
              ("table", t, [20, 170, 70, 110, 80]),
              ("p", "Signature of the bride: ____________________     Signature of the bridegroom: ____________________"),
              ("p", "Signatures of relatives / persons present (optional): ____________________")]
    return _pdf("List of presents", blocks)


def gift_list_file(tenant: str, wedding: str, for_whom: str = "bride", actor: str = "owner", conv: str = "") -> dict:
    """The presents list as HER file (only she gets it; it is never sent anywhere by Hoola)."""
    if not is_owner(actor):
        raise ShaadiError("Shagun register sirf owner ke liye hai.")
    pdf = gift_list_pdf(tenant, wedding, for_whom)
    f = _store(tenant, f"presents_list_{for_whom}.pdf", pdf, "application/pdf", conv, "List of presents")
    return {"file_id": f.get("file_id", ""), "say": "Tohfon ki list PDF aapki file mein; dulha-dulhan dono sign karein (Dowry Prohibition Rules 1985)."}


# ---------------------------------------------------------------- registration checklist + invite design
def registration(tenant: str, wedding: str) -> dict:
    with session() as s:
        w = _wed(s, tenant, wedding)
        law = R.COMMUNITIES.get(w.community, {}).get("law", "") or "hma"
        state = w.state
        ds = _dates(w)
    from ..india import days
    code = days.state_code(state) if state else None
    r = R.REGISTRATION[law]
    last = ds[-1] if ds else None
    out = {"law": r["law"], "items": r["items"], "state_note": R.STATE_NOTES.get(code or "", ""),
           "by": (last + timedelta(days=90)).isoformat() if last and code == "MH" else "",
           "note": "Sirf checklist: asli niyam registrar office / vakil batayenge. Hoola koi application file nahi karta.",
           "say": f"{r['law']}: {len(r['items'])} kaagaz ki list." + (" " + R.STATE_NOTES[code] if code in R.STATE_NOTES else "")}
    if law == "sma":
        out["say"] += " Special Marriage Act mein 30 din pehle notice deni hoti hai."
    return out


def invite_design(tenant: str, wedding: str, fn: str = "", words: str = "", conv: str = "") -> dict:
    """Card design through the Brand Sahayak (its own final-artwork approval). Text = the real facts only."""
    t = invite_text(tenant, wedding, "invite", fn)
    with session() as s:
        w = _wed(s, tenant, wedding)
        brief = f"Wedding invitation card: {w.bride} weds {w.groom}. " + t["text"].split("\n", 2)[-1][:600] + (" " + words if words else "")
    if t["blanks"]:
        brief += " (khaali: " + ", ".join(sorted(set(t["blanks"]))) + ")"
    from ..brand import service as BS
    try:
        out = BS.make(tenant, brief=R.clean(brief, 1500), kind="", sizes="whatsapp", conv=conv)
    except Exception as e:
        raise ShaadiError(f"Brand Sahayak se design nahi bana: {e}")
    return {"design": out.get("design") if isinstance(out, dict) else out,
            "say": "Card ka design Brand Sahayak mein ban raha hai; final artwork wahan aapki haan se."}


# ---------------------------------------------------------------- views
def _wedding_view(s, w: ShaadiWedding, full: bool = False) -> dict:
    fns = _j(w.functions, [])
    on = clock().date()
    ds = _dates(w)
    v = {"id": w.id, "number": w.number, "title": w.title, "community": w.community,
         "community_label": R.COMMUNITIES.get(w.community, {}).get("label", w.community), "bride": w.bride, "groom": w.groom,
         "client": w.client, "client_phone": w.client_phone, "city": w.city, "state": w.state, "venue": w.venue,
         "guests_estimate": w.guests_estimate, "budget_total": _j(w.budget, {}).get("total"), "status": w.status,
         "functions": fns, "first_date": _iso(ds[0]) if ds else "", "last_date": _iso(ds[-1]) if ds else "",
         "days_left": (ds[0] - on).days if ds else None, "forgotten": w.forgotten}
    if full:
        v["checklist"] = _checklist_view(w, on)
        v["log"] = _j(w.log, [])[-30:]
    return v


def weddings(tenant: str, status: str = "") -> list[dict]:
    with session() as s:
        q = select(ShaadiWedding).where(ShaadiWedding.tenant_id == tenant)
        if status:
            q = q.where(ShaadiWedding.status == status)
        return [_wedding_view(s, w) for w in s.scalars(q.order_by(ShaadiWedding.created_at.desc())).all()]


def wedding_view(tenant: str, ref: str) -> dict:
    with session() as s:
        w = _wed(s, tenant, ref)
        out = _wedding_view(s, w, full=True)
        wid = w.id
    out["headcount"] = headcount(tenant, wid)["functions"]
    out["vendors"] = vendors(tenant, wid)["vendors"]
    out["rooms"] = rooms(tenant, wid)
    out["budget"] = budget(tenant, wid)
    out["cards"] = cards(tenant, wid)
    out["gift_total"] = gifts(tenant, wid)["total"]
    return out


def waiting(tenant: str) -> list[dict]:
    return cards(tenant, status="waiting")


def day(tenant: str) -> dict:
    on = today(tenant)
    parts, soon, due_pay, todo = [], [], [], []
    with session() as s:
        ws = s.scalars(select(ShaadiWedding).where(ShaadiWedding.tenant_id == tenant, ShaadiWedding.status == "planning")).all()
        for w in ws:
            for f in _j(w.functions, []):
                d = _date(f.get("date"))
                if d and on <= d <= on + timedelta(days=7):
                    soon.append(f"{w.number} {f['name']} {d.strftime('%d %b')}" + (f" {f['start']}" if f.get("start") else ""))
            for c in _checklist_view(w, on):
                if not c["done"] and c["state"] in ("late", "today"):
                    todo.append(f"{w.number}: {c['what']}" + (" (der)" if c["state"] == "late" else ""))
            for v in s.scalars(select(ShaadiVendor).where(ShaadiVendor.wedding_id == w.id, ShaadiVendor.status == "booked")).all():
                for x in _j(v.schedule, []):
                    d = _date(x.get("due"))
                    if x["status"] in ("planned", "approved") and d and d <= on + timedelta(days=3):
                        due_pay.append(f"{w.number} {v.name} {x['label']} {_rs(x['amount'])} ({d.strftime('%d %b')})")
    w_cards = waiting(tenant)
    if soon:
        parts.append("Is hafte: " + "; ".join(soon[:6]))
    if todo:
        parts.append("Kaam: " + "; ".join(todo[:6]))
    if due_pay:
        parts.append("Payment aane wale: " + "; ".join(due_pay[:5]))
    if w_cards:
        parts.append(f"{len(w_cards)} card aapki haan ka intezaar kar rahe")
    return {"soon": soon, "todo": todo, "payments": due_pay, "waiting": w_cards, "weddings": len(ws),
            "say": (". ".join(parts) + ".") if parts else "Shaadi ka aaj sab theek."}


def overview(tenant: str) -> dict:
    return {"settings": settings(tenant), "weddings": weddings(tenant), "day": day(tenant),
            "communities": {k: {"label": v["label"], "functions": [R.F[f] for f in v["fns"]]} for k, v in R.COMMUNITIES.items()},
            "vendor_kinds": R.VENDOR_KINDS, "heads": R.HEADS}


# ---------------------------------------------------------------- DPDP forget
def forget(tenant: str, wedding: str, include_gifts: bool = False, actor: str = "owner") -> dict:
    if not is_owner(actor):
        raise ShaadiError("Sirf owner.")
    with _LOCK, session() as s:
        w = _wed(s, tenant, wedding)
        if w.status not in ("done", "cancelled"):
            raise ShaadiError("Shaadi poori / radd hone ke baad hi mehmaanon ki details mitayi ja sakti hain.")
        n = 0
        for i, g in enumerate(s.scalars(select(ShaadiGuest).where(ShaadiGuest.wedding_id == w.id)).all(), 1):
            g.name, g.phone, g.notes, g.travel, g.city, g.grp = f"Mehmaan {i}", "", "", "{}", "", ""
            n += 1
        w.client, w.client_phone = ("Client" if w.client else ""), ""
        w.forgotten = True
        gn = 0
        if include_gifts:
            for g in s.scalars(select(ShaadiGift).where(ShaadiGift.wedding_id == w.id)).all():
                s.delete(g)
                gn += 1
        _log(w, f"DPDP: {n} mehmaanon ki details mitayi" + (f", {gn} shagun entry" if gn else ""), "Owner")
    return {"guests": n, "gifts": gn, "say": f"{n} mehmaanon ke naam / phone / safar mita diye; ginti aur hisaab rahe."}


# ---------------------------------------------------------------- worker
def _once(s, tenant: str, key: str) -> bool:
    k = f"{tenant}:{key}"[:200]
    if s.get(ShaadiAlert, k) is not None:
        return False
    s.add(ShaadiAlert(key=k, tenant_id=tenant))
    return True


def _tenants() -> list[str]:
    with session() as s:
        return sorted(set(s.scalars(select(ShaadiWedding.tenant_id).where(ShaadiWedding.status == "planning")).all()))


def tick(when: datetime | None = None) -> dict:
    """Worker, every minute: function tomorrow, checklist item due today, vendor payment due (once each); morning note at her time."""
    n = _aware(when) or now()
    out = {"alerts": 0, "morning": 0}
    for tenant in _tenants():
        try:
            loc = n.astimezone(_tz(tenant))
            on = loc.date()
            alerts = []
            with _LOCK, session() as s:
                for w in s.scalars(select(ShaadiWedding).where(ShaadiWedding.tenant_id == tenant, ShaadiWedding.status == "planning")).all():
                    for f in _j(w.functions, []):
                        d = _date(f.get("date"))
                        if d and d == on + timedelta(days=1) and _once(s, tenant, f"fn:{f['id']}:{d}"):
                            tl = "; ".join(f"{r['time']} {r['what']}" for r in f.get("timeline", [])[:6])
                            alerts.append({"kind": "tomorrow", "wedding": w.number,
                                           "say": f"Kal {w.number} {f['name']}" + (f" {f['start']}" if f.get("start") else "")
                                                  + (f". Timeline: {tl}" if tl else ". Timeline abhi nahi bani.")})
                    for c in _checklist_view(w, on):
                        if not c["done"] and c["state"] == "today" and _once(s, tenant, f"cl:{w.id}:{c['id']}:{c['due']}"):
                            alerts.append({"kind": "checklist", "wedding": w.number, "say": f"{w.number}: aaj tak '{c['what']}'."})
                    for v in s.scalars(select(ShaadiVendor).where(ShaadiVendor.wedding_id == w.id, ShaadiVendor.status == "booked")).all():
                        for x in _j(v.schedule, []):
                            d = _date(x.get("due"))
                            if x["status"] == "planned" and d and d <= on and _once(s, tenant, f"pay:{v.id}:{x['id']}:{x['due']}"):
                                alerts.append({"kind": "payment", "wedding": w.number, "vendor": v.id, "item": x["id"],
                                               "say": f"{w.number}: {v.name} ki '{x['label']}' {_rs(x['amount'])} ka din aa gaya; "
                                                      "card banau?"})
            for a in alerts:
                _publish(tenant, "shaadi_alert", a)
                out["alerts"] += 1
            due_m = False
            with session() as s:
                r = s.get(ShaadiSettings, tenant)
                mt = (r.morning_time if r else "") or "08:30"
                last = r.last_morning if r else ""
                if last != on.isoformat():
                    m = R.hhmm(mt) or 510
                    due_m = loc.hour * 60 + loc.minute >= m
                    if due_m:
                        st = _st(s, tenant)
                        st.last_morning = on.isoformat()
            if due_m:
                d = day(tenant)
                if d["say"] != "Shaadi ka aaj sab theek.":
                    _publish(tenant, "shaadi_day", {"say": d["say"], "soon": d["soon"], "todo": d["todo"][:6], "payments": d["payments"][:5]})
                    out["morning"] += 1
        except Exception:
            log.exception("shaadi tick failed for %s", tenant)
    return out
