"""Cosec Sahayak service: the company secretary's desk, end to end.

Company / LLP master (CIN / LLPIN, FY, directors by DIN only, auditor, class) -> MCA21 V3 forms calendar computed by code from
HER FY / AGM / event dates (AOC-4, MGT-7 / 7A, ADT-1, DIR-3 KYC, DPT-3, MSME-1, INC-20A, BEN-2, LLP Form 8 / 11 ...) with
additional-fee starters -> Board meeting / AGM: notice drafts with the SS-1 / SS-2 notice period checked, agenda, attendance +
quorum, minutes drafts, resolutions register -> statutory registers (members MGT-1 style, directors & KMP, charges) as Excel ->
event filings from her words (DIR-12, PAS-3, INC-22, CHG-1, MGT-14 ...).

HER tap (enforced here, `by_tap=True` only from the API): sending a meeting notice, signing off the minutes, and "file karo" on a
filing card (her fingerprint too once a phone key is active). Hoola never logs in to MCA V3, never uses a DSC, never pays a fee:
cards, drafts and steps only. Everything is "draft / jaankari - CS / vakil / CA se pakka karwaiye"."""
from __future__ import annotations

import hashlib
import io
import json
import logging
import re
import secrets
import threading
import urllib.parse
from datetime import date, datetime, timedelta, timezone
from typing import Any

from sqlalchemy import select

from ..db import now, session
from . import rules as R
from .models import CosecCharge, CosecEntity, CosecFiling, CosecMeeting, CosecOnce, CosecPerson, CosecResolution, CosecSettings

log = logging.getLogger("santosh.cosec")
_LOCK = threading.RLock()


class CosecError(Exception):
    pass


# ---------------------------------------------------------------- small helpers
def _j(v, default):
    try:
        x = json.loads(v) if isinstance(v, str) and v else default
        return x if isinstance(x, type(default)) else default
    except (ValueError, TypeError):
        return default


def _d(v) -> str:
    return json.dumps(v, ensure_ascii=False, default=str)


def clock() -> datetime:
    return now()


def _tz(tenant: str):
    from ..life import tz_of
    return tz_of(tenant)


def today(tenant: str) -> date:
    return clock().astimezone(_tz(tenant)).date()


def _date(v) -> date | None:
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    if not v:
        return None
    s = str(v).strip()[:10]
    for f in ("%Y-%m-%d", "%d-%m-%Y", "%d/%m/%Y", "%d.%m.%Y", "%d-%m-%y", "%d/%m/%y"):
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            pass
    return None


def _iso(v) -> str:
    d = _date(v)
    return d.isoformat() if d else ""


def _clean(text, n: int = 500) -> str:
    """Her / outside words: secrets never kept; Aadhaar / PAN only by the locker reference."""
    t = str(text or "")
    t = re.sub(r"\b(otp|pin|cvv|password|upi pin)\b\s*(?:is|hai|:|=)?\s*\S+", lambda m: f"{m.group(1)} [hataya]", t, flags=re.I)
    t = re.sub(r"\b(?:\d[ -]?){13,19}\b", "[hataya]", t)
    t = re.sub(r"\b\d{4}\s?\d{4}\s?\d{4}\b", "[hataya]", t)
    t = re.sub(r"\b[A-Z]{5}\d{4}[A-Z]\b", "[PAN hataya]", t)
    return t.strip()[:n]


def _phone(v) -> str:
    d = re.sub(r"\D", "", str(v or ""))
    if len(d) == 12 and d.startswith("91"):
        d = d[2:]
    return d if len(d) == 10 and d[0] in "6789" else ""


def _num(v) -> float | None:
    if v is None or v == "":
        return None
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).lower().replace(",", "").replace("rs", "").replace("₹", "").strip()
    m = re.match(r"^([\d.]+)\s*(crore|cr|lakh|lac|l)?$", s)
    if not m:
        raise CosecError(f"'{v}' number nahi samjha.")
    x = float(m.group(1))
    mult = {"crore": 1e7, "cr": 1e7, "lakh": 1e5, "lac": 1e5, "l": 1e5}.get(m.group(2) or "", 1)
    return x * mult


def _hash(show: dict) -> str:
    return hashlib.sha256(json.dumps(show, sort_keys=True, ensure_ascii=False, default=str).encode()).hexdigest()


def _publish(tenant: str, kind: str, payload: dict, conv: str = "") -> None:
    from ..events import bus
    try:
        if not conv:
            from ..net.service import home_conv
            conv = home_conv(tenant)
        bus.publish(conv, kind, payload)
    except Exception:
        log.exception("cosec publish failed")


def is_owner(actor: str) -> bool:
    a = actor or "owner"
    return a == "owner" or a.startswith("approved:")


def _st(s, tenant: str) -> CosecSettings:
    r = s.get(CosecSettings, tenant)
    if r is None:
        r = CosecSettings(tenant_id=tenant)
        s.add(r)
        s.flush()
    return r


def _next(s, tenant: str, key: str) -> int:
    st = _st(s, tenant)
    seq = _j(st.seq, {})
    seq[key] = int(seq.get(key) or 0) + 1
    st.seq = _d(seq)
    return seq[key]


def _once(s, tenant: str, key: str) -> bool:
    k = f"{tenant}:{key}"[:200]
    if s.get(CosecOnce, k) is not None:
        return False
    s.add(CosecOnce(key=k, tenant_id=tenant))
    return True


def _wa(phone: str, text: str) -> str:
    p = _phone(phone)
    return f"https://wa.me/91{p}?text={urllib.parse.quote(text)}" if p else f"https://wa.me/?text={urllib.parse.quote(text)}"


def _mail(to: str, subject: str, text: str) -> str:
    return f"mailto:{urllib.parse.quote(to or '')}?subject={urllib.parse.quote(subject)}&body={urllib.parse.quote(text)}"


def _store(tenant: str, name: str, data: bytes, ctype: str, conv: str = "", title: str = "") -> dict:
    from .. import chatfiles
    f = chatfiles.store(tenant, name, data, ctype)
    _publish(tenant, "pc_file", {**f, "source": "cosec", "title": title or name}, conv)
    return f


DOCX = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


# ---------------------------------------------------------------- settings
_SET_KEYS = ("firm", "cs_name", "membership_no", "email", "phone", "windows", "morning_time", "rules")


def _set_view(r: CosecSettings) -> dict:
    return {"firm": r.firm, "cs_name": r.cs_name, "membership_no": r.membership_no, "email": r.email, "phone": r.phone,
            "windows": r.windows, "morning_time": r.morning_time, "rules": r.rules, "kinds": R.KINDS}


def settings(tenant: str) -> dict:
    with session() as s:
        return _set_view(_st(s, tenant))


def _windows(st: CosecSettings) -> list[int]:
    out = []
    for x in str(st.windows or "").split(","):
        try:
            out.append(max(0, min(120, int(x.strip()))))
        except ValueError:
            pass
    return sorted(set(out or [30, 15, 7, 1, 0]))


def save_settings(tenant: str, **kw) -> dict:
    with _LOCK, session() as s:
        r = _st(s, tenant)
        for k in _SET_KEYS:
            if k not in kw or kw[k] is None:
                continue
            v = str(kw[k]).strip()
            if k == "phone" and v and not _phone(v):
                raise CosecError("Mobile 10 ankon ka chahiye.")
            if k == "morning_time" and v and not re.match(r"^\d{1,2}:\d{2}$", v):
                raise CosecError("Samay HH:MM mein.")
            if k == "membership_no" and v and not re.match(r"^[AF]CS\s*\d{1,6}$", v.upper()):
                raise CosecError("Membership no. jaise ACS 12345 / FCS 1234.")
            setattr(r, k, _phone(v) if k == "phone" else (v.upper() if k == "membership_no" else _clean(v, 500)))
        v = _set_view(r)
    v["say"] = "Settings save."
    return v


# ---------------------------------------------------------------- entities (company / LLP master)
def _ent(s, tenant: str, ref: str) -> CosecEntity:
    ref = str(ref or "").strip()
    if not ref:
        rows = s.scalars(select(CosecEntity).where(CosecEntity.tenant_id == tenant, CosecEntity.status == "active")).all()
        if len(rows) == 1:
            return rows[0]
        raise CosecError("Kaunsi company / LLP? Naam ya E-number batao.")
    e = s.get(CosecEntity, ref)
    if e is not None and e.tenant_id == tenant:
        return e
    rows = s.scalars(select(CosecEntity).where(CosecEntity.tenant_id == tenant)).all()
    up = ref.upper()
    for e in rows:
        if up in (e.number.upper(), e.cin, e.llpin):
            return e
    hits = [e for e in rows if ref.lower() in e.name.lower()]
    if len(hits) == 1:
        return hits[0]
    if len(hits) > 1:
        raise CosecError("Ek se zyada mili: " + ", ".join(f"{e.number} {e.name}" for e in hits[:5]))
    raise CosecError(f"'{ref}' naam ki company / LLP nahi mili.")


def _light(e: CosecEntity) -> bool:
    f = _j(e.flags, {})
    return e.kind == "opc" or bool(e.small) or bool(f.get("dormant")) or bool(f.get("startup"))


def _ent_view(e: CosecEntity) -> dict:
    sc = R.small_check(e.kind, e.paid_up, e.turnover, e.holding_or_sub)
    return {"id": e.id, "number": e.number, "name": e.name, "kind": e.kind, "kind_label": R.KINDS.get(e.kind, e.kind), "small": e.small,
            "small_check": sc, "holding_or_sub": e.holding_or_sub, "cin": e.cin, "llpin": e.llpin, "roc": e.roc,
            "incorporated_on": e.incorporated_on, "paid_up": e.paid_up, "turnover": e.turnover, "registered_office": e.registered_office,
            "email": e.email, "auditor": _j(e.auditor, {}), "flags": _j(e.flags, {}), "last_agm": e.last_agm, "status": e.status,
            "notes": e.notes, "light_board": _light(e), "draft_line": R.DRAFT_LINE}


_ENT_KEYS = ("kind", "small", "holding_or_sub", "cin", "llpin", "roc", "incorporated_on", "paid_up", "turnover", "registered_office",
             "email", "auditor_name", "auditor_frn", "auditor_appointed_on", "auditor_term_to_agm", "last_agm", "status", "notes",
             "deposits", "msme_dues", "dormant", "startup", "first_fy")


def save_entity(tenant: str, name: str = "", ref: str = "", **kw) -> dict:
    with _LOCK, session() as s:
        if ref:
            e = _ent(s, tenant, ref)
        else:
            if not str(name or "").strip():
                raise CosecError("Company / LLP ka naam chahiye.")
            e = CosecEntity(tenant_id=tenant, number=f"E-{_next(s, tenant, 'entity'):03d}")
            s.add(e)
        if name:
            e.name = _clean(name, 200)
        if "kind" in kw and kw["kind"]:
            k = str(kw["kind"]).lower().strip()
            k = {"pvt": "private", "pvt ltd": "private", "public limited": "public", "ltd": "public", "one person": "opc"}.get(k, k)
            if k not in R.KINDS:
                raise CosecError("kind = " + " | ".join(R.KINDS))
            e.kind = k
        try:
            if kw.get("cin") is not None:
                e.cin = R.check_cin(kw["cin"])
            if kw.get("llpin") is not None:
                e.llpin = R.check_llpin(kw["llpin"])
        except ValueError as ex:
            raise CosecError(str(ex))
        for k in ("roc", "registered_office", "notes"):
            if kw.get(k) is not None:
                setattr(e, k, _clean(kw[k], 400))
        if kw.get("email") is not None:
            v = str(kw["email"]).strip()
            if v and "@" not in v:
                raise CosecError("Email sahi nahi.")
            e.email = v[:120]
        for k in ("incorporated_on", "last_agm"):
            if kw.get(k) is not None:
                v = _iso(kw[k])
                if kw[k] and not v:
                    raise CosecError(f"{k}: tareekh YYYY-MM-DD ya DD-MM-YYYY mein.")
                setattr(e, k, v)
        for k in ("paid_up", "turnover"):
            if k in kw and kw[k] is not None:
                setattr(e, k, _num(kw[k]))
        if "holding_or_sub" in kw and kw["holding_or_sub"] is not None:
            e.holding_or_sub = bool(kw["holding_or_sub"])
        if "small" in kw and kw["small"] is not None:
            e.small = bool(kw["small"])
        else:
            sc = R.small_check(e.kind, e.paid_up, e.turnover, e.holding_or_sub)
            if sc["small"] is not None:
                e.small = sc["small"]
        if e.kind != "private":
            e.small = False
        aud = _j(e.auditor, {})
        for k, kk in (("auditor_name", "name"), ("auditor_frn", "frn")):
            if kw.get(k) is not None:
                aud[kk] = _clean(kw[k], 120)
        if kw.get("auditor_appointed_on") is not None:
            aud["appointed_on"] = _iso(kw["auditor_appointed_on"])
        if kw.get("auditor_term_to_agm") not in (None, ""):
            y = R.fy_end_year(kw["auditor_term_to_agm"])
            if not y:
                raise CosecError("Auditor term: kis FY ke AGM tak (jaise 2029-30).")
            aud["term_to_agm"] = y
        e.auditor = _d(aud)
        fl = _j(e.flags, {})
        for k in ("deposits", "msme_dues", "dormant", "startup", "first_fy"):
            if k in kw and kw[k] is not None:
                fl[k] = bool(kw[k])
        e.flags = _d(fl)
        if kw.get("status"):
            st = str(kw["status"]).lower()
            if st not in ("active", "closed", "struck_off"):
                raise CosecError("status = active | closed | struck_off")
            e.status = st
        if e.kind == "llp" and e.cin:
            raise CosecError("LLP ka LLPIN hota hai, CIN nahi.")
        s.flush()
        v = _ent_view(e)
    warn = []
    if v["kind"] == "llp" and not v["llpin"]:
        warn.append("LLPIN baaki")
    if v["kind"] != "llp" and not v["cin"]:
        warn.append("CIN baaki")
    if v["kind"] == "private" and v["small"] is None:
        warn.append(v["small_check"]["why"])
    v["say"] = f"{v['number']} {v['name']} ({v['kind_label']}) save." + (f" Baaki: {'; '.join(warn)}." if warn else "")
    return v


def entities(tenant: str, status: str = "") -> list[dict]:
    with session() as s:
        q = select(CosecEntity).where(CosecEntity.tenant_id == tenant)
        if status:
            q = q.where(CosecEntity.status == status)
        out = []
        for e in s.scalars(q.order_by(CosecEntity.number)).all():
            v = _ent_view(e)
            v["directors"] = len(s.scalars(select(CosecPerson.id).where(CosecPerson.entity_id == e.id, CosecPerson.role.in_(("director", "partner")),
                                                                        CosecPerson.ceased_on == "")).all())
            out.append(v)
        return out


# ---------------------------------------------------------------- people: directors / partners / KMP / members
def _person_view(p: CosecPerson) -> dict:
    return {"id": p.id, "role": p.role, "name": p.name, "din": p.din, "designation": p.designation, "appointed_on": p.appointed_on,
            "ceased_on": p.ceased_on, "active": not p.ceased_on, "din_on": p.din_on, "kyc_on": p.kyc_on, "email": p.email, "phone": p.phone,
            "address": p.address, "id_ref": p.id_ref, "folio": p.folio, "shares": p.shares, "share_class": p.share_class, "nominal": p.nominal,
            "member_since": p.member_since, "history": _j(p.history, [])}


def _find_person(s, e: CosecEntity, ref: str, role: str = "") -> CosecPerson | None:
    ref = str(ref or "").strip()
    if not ref:
        return None
    p = s.get(CosecPerson, ref)
    if p is not None and p.entity_id == e.id:
        return p
    rows = s.scalars(select(CosecPerson).where(CosecPerson.entity_id == e.id)).all()
    for p in rows:
        if (p.din and p.din == re.sub(r"\D", "", ref)) or (p.folio and p.folio.lower() == ref.lower()):
            return p
    hits = [p for p in rows if p.name.lower() == ref.lower() and (not role or p.role == role)]
    return hits[0] if hits else None


def person(tenant: str, entity_ref: str, name: str = "", role: str = "director", ref: str = "", **kw) -> dict:
    role = (role or "director").lower()
    if role not in ("director", "partner", "kmp", "member"):
        raise CosecError("role = director | partner | kmp | member")
    with _LOCK, session() as s:
        e = _ent(s, tenant, entity_ref)
        p = _find_person(s, e, ref or name, role) if (ref or name) else None
        if p is None:
            if not str(name or "").strip():
                raise CosecError("Naam chahiye.")
            p = CosecPerson(tenant_id=tenant, entity_id=e.id, role=role)
            s.add(p)
        if name:
            p.name = _clean(name, 160)
        if kw.get("role_change"):
            p.role = role
        try:
            if kw.get("din") is not None:
                p.din = R.check_din(kw["din"])
        except ValueError as ex:
            raise CosecError(str(ex))
        if e.kind == "llp" and role == "director":
            raise CosecError("LLP mein director nahi, designated partner (role=partner).")
        for k in ("designation", "address", "share_class", "folio"):
            if kw.get(k) is not None:
                setattr(p, k, _clean(kw[k], 300))
        for k in ("appointed_on", "ceased_on", "din_on", "kyc_on", "member_since"):
            if kw.get(k) is not None:
                v = _iso(kw[k])
                if kw[k] and not v:
                    raise CosecError(f"{k}: tareekh samajh nahi aayi.")
                setattr(p, k, v)
        if kw.get("email") is not None:
            v = str(kw["email"]).strip()
            if v and "@" not in v:
                raise CosecError("Email sahi nahi.")
            p.email = v[:120]
        if kw.get("phone") is not None:
            v = _phone(kw["phone"])
            if kw["phone"] and not v:
                raise CosecError("Mobile 10 ankon ka chahiye.")
            p.phone = v
        if kw.get("id_ref") is not None:
            v = str(kw["id_ref"]).strip()
            if v and not re.fullmatch(r"\{id:[A-Za-z0-9_-]{3,40}\}", v):
                raise CosecError("Aadhaar / PAN yahan nahi: Pehchaan Locker mein daaliye, yahan sirf {id:...} reference.")
            p.id_ref = v
        for k in ("shares", "nominal"):
            if kw.get(k) is not None:
                x = _num(kw[k])
                if x is not None and x < 0:
                    raise CosecError(f"{k} minus nahi ho sakta.")
                setattr(p, k, (x or 0.0) if k == "shares" else x)
        if p.role == "director" and not p.din and kw.get("need_din", False):
            raise CosecError("Director ka DIN chahiye.")
        s.flush()
        v = _person_view(p)
    v["say"] = f"{v['name']} ({v['role']}) save" + (f", DIN {v['din']}" if v["din"] else "") + "."
    return v


# ---------------------------------------------------------------- filings calendar
def _filing_view(f: CosecFiling, on: date | None = None) -> dict:
    meta = R.FORMS.get(f.form, {})
    due = _date(f.due)
    left = (due - on).days if (due and on) else None
    fee = R.additional_fee(f.form, -left) if (left is not None and left < 0 and f.status not in ("filed", "not_needed")) else None
    if f.status == "filed" and due and _date(f.filed_on) and _date(f.filed_on) > due:
        fee = R.additional_fee(f.form, (_date(f.filed_on) - due).days)
    return {"id": f.id, "form": f.form, "label": meta.get("label", f.form), "law": meta.get("law", ""), "period": f.period, "source": f.source,
            "event": _j(f.event, {}), "due": f.due, "basis": f.basis, "est": f.est, "note": f.note, "status": f.status, "days_left": left,
            "late": left is not None and left < 0 and f.status not in ("filed", "not_needed"), "fee": fee, "hash": f.hash,
            "card": _j(f.card, {}), "srn": f.srn, "filed_on": f.filed_on, "fee_paid": f.fee_paid, "reason": f.reason,
            "is_meeting": f.form in ("AGM", "FIRST-BOARD", "FIRST-AUDITOR")}


def _upsert(s, e: CosecEntity, form: str, period: str, due: date | None, basis: str, est: bool, note: str, source: str,
            event: dict | None = None) -> CosecFiling:
    key = f"{form}|{period}"[:160]
    f = s.scalars(select(CosecFiling).where(CosecFiling.entity_id == e.id, CosecFiling.key == key)).first()
    if f is None:
        f = CosecFiling(tenant_id=e.tenant_id, entity_id=e.id, key=key, form=form, period=period, source=source)
        s.add(f)
    if f.status in ("todo", "stale") or not f.due:
        f.due = due.isoformat() if due else ""
        f.basis, f.est, f.note = basis, bool(est), note
        if event is not None:
            f.event = _d(event)
    return f


def _agm_for(s, e: CosecEntity, fy_end: int) -> date | None:
    lab = R.fy_label(fy_end)
    for m in s.scalars(select(CosecMeeting).where(CosecMeeting.entity_id == e.id, CosecMeeting.kind == "agm", CosecMeeting.fy == lab)).all():
        if m.status in ("held", "minutes_waiting", "minuted") and _date(m.on):
            return _date(m.on)
    return None


def _calendar(s, e: CosecEntity, fy_end: int) -> None:
    fl = _j(e.flags, {})
    aud = _j(e.auditor, {})
    inc = _date(e.incorporated_on)
    first = bool(fl.get("first_fy")) or (inc is not None and R.fy_end_year(inc) == fy_end)
    agm_on = _agm_for(s, e, fy_end)
    prev_agm = _agm_for(s, e, fy_end - 1) or _date(e.last_agm)
    if prev_agm and agm_on and prev_agm >= agm_on:
        prev_agm = None
    term = aud.get("term_to_agm")
    rows = R.annual(e.kind, bool(e.small), fy_end, agm_on=agm_on, last_agm=prev_agm, first_fy=first,
                    auditor_term_ends=(not aud.get("name")) or term == fy_end, deposits=fl.get("deposits", True) is not False,
                    msme_dues=bool(fl.get("msme_dues")), paid_up=e.paid_up, turnover=e.turnover)
    lab = R.fy_label(fy_end)
    keys = set()
    for r in rows:
        f = _upsert(s, e, r["form"], r["period"], r["due"], r["basis"], r["est"], r["note"], "annual", {"fy": lab})
        keys.add(f.key)
        if r["form"] == "AGM" and agm_on and f.status in ("todo", "stale"):
            f.status, f.filed_on, f.reason = "filed", agm_on.isoformat(), "AGM ho gaya"
    for f in s.scalars(select(CosecFiling).where(CosecFiling.entity_id == e.id, CosecFiling.source == "annual")).all():
        if _j(f.event, {}).get("fy") == lab and f.key not in keys and f.status == "todo":
            f.status, f.reason = "not_needed", "Ab lagoo nahi (company ki jaankari badli)"
    if e.kind != "llp":
        for p in s.scalars(select(CosecPerson).where(CosecPerson.entity_id == e.id, CosecPerson.role == "director", CosecPerson.ceased_on == "")).all():
            if not p.din:
                continue
            k = R.kyc_due(_date(p.kyc_on), _date(p.din_on))
            if k["due"]:
                _upsert(s, e, "DIR-3 KYC", f"{p.name} DIN {p.din} · {k['due'].isoformat()}", k["due"], k["basis"], k["est"],
                        "DIN deactivate na ho - har director khud apne DSC / OTP se bharta hai", "kyc", {"person": p.id})


def calendar(tenant: str, ref: str, fy: str = "") -> dict:
    """Upserts this FY's annual filings (and directors' KYC) from her dates; returns the entity's filing list."""
    with _LOCK, session() as s:
        e = _ent(s, tenant, ref)
        on = today(tenant)
        y = R.fy_end_year(fy) if fy else (on.year if on.month >= 4 else on.year - 1)
        if not y:
            raise CosecError("FY jaise 2025-26.")
        _calendar(s, e, y)
        s.flush()
        rows = [_filing_view(f, on) for f in s.scalars(select(CosecFiling).where(CosecFiling.entity_id == e.id)).all()]
        name = f"{e.number} {e.name}"
    rows.sort(key=lambda x: (x["status"] in ("filed", "not_needed"), x["due"] or "9999"))
    late = [r for r in rows if r["late"]]
    soon = [r for r in rows if not r["late"] and r["days_left"] is not None and r["days_left"] <= 30 and r["status"] not in ("filed", "not_needed")]
    say = f"{name} FY {R.fy_label(y)}: {len(rows)} kaam." + (f" Der: {', '.join(r['form'] for r in late[:4])}." if late else "") + (
        f" 30 din mein: {', '.join(r['form'] + ' ' + r['due'] for r in soon[:4])}." if soon else "") + f" {R.DRAFT_LINE}"
    return {"entity": name, "fy": R.fy_label(y), "filings": rows, "say": say}


def ensure(tenant: str) -> None:
    on = today(tenant)
    y0 = on.year if on.month >= 4 else on.year - 1
    with _LOCK, session() as s:
        for e in s.scalars(select(CosecEntity).where(CosecEntity.tenant_id == tenant, CosecEntity.status == "active")).all():
            for y in (y0, y0 + 1):
                try:
                    _calendar(s, e, y)
                except Exception:
                    log.exception("cosec calendar failed for %s", e.id)


def filings(tenant: str, ref: str = "", status: str = "", within_days: int | None = None) -> list[dict]:
    on = today(tenant)
    with session() as s:
        q = select(CosecFiling).where(CosecFiling.tenant_id == tenant)
        if ref:
            q = q.where(CosecFiling.entity_id == _ent(s, tenant, ref).id)
        if status:
            q = q.where(CosecFiling.status == status)
        ents = {e.id: f"{e.number} {e.name}" for e in s.scalars(select(CosecEntity).where(CosecEntity.tenant_id == tenant)).all()}
        out = []
        for f in s.scalars(q).all():
            v = _filing_view(f, on) | {"entity": ents.get(f.entity_id, "")}
            if within_days is not None and (v["days_left"] is None or v["days_left"] > within_days or v["status"] in ("filed", "not_needed")):
                continue
            out.append(v)
    out.sort(key=lambda x: x["due"] or "9999")
    return out[:300]


# ---------------------------------------------------------------- events from her words -> filings + registers
def event(tenant: str, ref: str, kind: str = "", on: str = "", words: str = "", details: dict | None = None, conv: str = "") -> dict:
    kind = (kind or R.guess_event(words)).strip()
    if kind not in R.EVENTS:
        raise CosecError("Kaunsa badlav? " + " | ".join(f"{k} ({v['label']})" for k, v in R.EVENTS.items()))
    d0 = _date(on) or today(tenant)
    det = {k: v for k, v in (details or {}).items() if v not in (None, "", [], {})}
    notes: list[str] = []
    with _LOCK, session() as s:
        e = _ent(s, tenant, ref)
        if kind in ("director_appoint", "director_resign", "kyc_change") and e.kind == "llp":
            raise CosecError("LLP mein partner badlav = partner_change.")
        if kind == "partner_change" and e.kind != "llp":
            raise CosecError("partner_change sirf LLP ke liye.")
        # register effects
        if kind in ("director_appoint", "partner_change") and det.get("name"):
            role = "partner" if e.kind == "llp" else ("kmp" if str(det.get("role", "")).lower() == "kmp" else "director")
            p = _find_person(s, e, det.get("din") or det["name"], role) or CosecPerson(tenant_id=tenant, entity_id=e.id, role=role)
            if p.id is None:
                s.add(p)
            p.name = _clean(det["name"], 160)
            if det.get("din"):
                try:
                    p.din = R.check_din(det["din"])
                except ValueError as ex:
                    raise CosecError(str(ex))
            p.designation = _clean(det.get("designation") or p.designation or ("Designated Partner" if role == "partner" else "Director"), 80)
            if det.get("ceased"):
                p.ceased_on = d0.isoformat()
            else:
                p.appointed_on, p.ceased_on = d0.isoformat(), ""
            if role == "director" and not p.din:
                notes.append("Naye director ka DIN chahiye (DIN nahi to pehle DIR-3 / SPICe+ se DIN).")
        if kind == "director_resign":
            p = _find_person(s, e, det.get("din") or det.get("name") or "")
            if p is None:
                raise CosecError("Kis director ne istifa diya? naam / DIN batao.")
            p.ceased_on = d0.isoformat()
            det["name"] = p.name
            left = s.scalars(select(CosecPerson).where(CosecPerson.entity_id == e.id, CosecPerson.role == "director", CosecPerson.ceased_on == "")).all()
            need = 1 if e.kind == "opc" else 3 if e.kind in ("public", "listed") else 2
            if len(left) < need:
                notes.append(f"Ab sirf {len(left)} director: {R.KINDS[e.kind]} mein kam se kam {need} chahiye - jaldi naya director (CS se pakka).")
        if kind == "allotment":
            allottees = det.get("allottees") or ([{"name": det["name"], "shares": det.get("shares")}] if det.get("name") else [])
            if not allottees:
                raise CosecError("Allotment: kisko kitne shares (allottees [{name, shares}]).")
            total = 0.0
            for a in allottees:
                n = _num(a.get("shares"))
                if not a.get("name") or not n or n <= 0:
                    raise CosecError("Har allottee ka naam aur shares chahiye.")
                p = _find_person(s, e, a.get("folio") or a["name"], "member") or CosecPerson(tenant_id=tenant, entity_id=e.id, role="member")
                if p.id is None:
                    s.add(p)
                    p.folio = f"F-{_next(s, tenant, 'folio:' + e.id):04d}"
                    p.member_since = d0.isoformat()
                p.name = _clean(a["name"], 160)
                p.shares = (p.shares or 0) + n
                p.share_class = _clean(det.get("share_class") or p.share_class or "Equity", 30)
                if a.get("nominal") or det.get("nominal"):
                    p.nominal = _num(a.get("nominal") or det.get("nominal"))
                h = _j(p.history, [])
                h.append({"on": d0.isoformat(), "what": "allotment", "shares": n})
                p.history = _d(h[-50:])
                total += n
            det["total_shares"] = total
            if e.kind in ("private", "opc"):
                notes.append("Private company mein Rule 9B (demat) lagoo ho to naye shares demat mein - CS se pakka.")
        if kind == "transfer":
            src = _find_person(s, e, det.get("from") or "", "member")
            n = _num(det.get("shares"))
            if src is None or not n or n <= 0 or not det.get("to"):
                raise CosecError("Transfer: from (member), to (naam), shares chahiye.")
            if (src.shares or 0) < n:
                raise CosecError(f"{src.name} ke paas sirf {src.shares:g} shares.")
            dst = _find_person(s, e, det["to"], "member") or CosecPerson(tenant_id=tenant, entity_id=e.id, role="member")
            if dst.id is None:
                s.add(dst)
                dst.folio = f"F-{_next(s, tenant, 'folio:' + e.id):04d}"
                dst.member_since = d0.isoformat()
                dst.share_class = src.share_class
                dst.nominal = src.nominal
            dst.name = dst.name or _clean(det["to"], 160)
            src.shares -= n
            dst.shares = (dst.shares or 0) + n
            for p, w, q in ((src, "transfer out", -n), (dst, "transfer in", n)):
                h = _j(p.history, [])
                h.append({"on": d0.isoformat(), "what": w, "shares": q})
                p.history = _d(h[-50:])
            if src.shares == 0 and not src.ceased_on:
                src.ceased_on = d0.isoformat()
            notes.append(R.EVENTS["transfer"]["note"])
        if kind.startswith("ro_change") and det.get("address"):
            det["old_address"] = e.registered_office
            e.registered_office = _clean(det["address"], 400)
        if kind == "charge_create":
            c = CosecCharge(tenant_id=tenant, entity_id=e.id, holder=_clean(det.get("holder"), 160), amount=_num(det.get("amount")),
                            property=_clean(det.get("property"), 300), created_on=d0.isoformat(), charge_id=_clean(det.get("charge_id"), 30))
            s.add(c)
            if not c.holder:
                notes.append("Charge holder (bank) ka naam batao.")
        if kind == "charge_satisfy":
            cs = [c for c in s.scalars(select(CosecCharge).where(CosecCharge.entity_id == e.id, CosecCharge.satisfied_on == "")).all()
                  if not det.get("holder") or str(det["holder"]).lower() in c.holder.lower() or det.get("charge_id") == c.charge_id]
            if len(cs) == 1:
                cs[0].satisfied_on = d0.isoformat()
            else:
                notes.append("Kaunsa charge chuka? holder / charge id batao (register mein mark karna hai).")
        if kind == "auditor_appoint" and det.get("name"):
            aud = _j(e.auditor, {})
            aud |= {"name": _clean(det["name"], 120), "frn": _clean(det.get("frn"), 30), "appointed_on": d0.isoformat()}
            if det.get("term_to_agm"):
                aud["term_to_agm"] = R.fy_end_year(det["term_to_agm"])
            e.auditor = _d(aud)
        if kind == "incorporation":
            e.incorporated_on = d0.isoformat()
        # filings
        made = []
        for r in R.event_filings(kind, d0, e.kind):
            tag = det.get("name") or det.get("address") or det.get("holder") or ""
            period = f"{R.EVENTS[kind]['label']} {d0.isoformat()}" + (f" · {str(tag)[:30]}" if tag else "")
            f = _upsert(s, e, r["form"], period, r["due"], r["basis"], False, r["note"], "event", {"kind": kind, "on": d0.isoformat(), **det})
            s.flush()
            made.append(_filing_view(f, today(tenant)))
        name = f"{e.number} {e.name}"
    if any(m["days_left"] is not None and m["days_left"] < 0 for m in made):
        notes.append("Tareekh pehle hi nikal chuki: extra fee lagegi, jaldi file karwaiye.")
    say = (f"{name}: {R.EVENTS[kind]['label']} ({d0.isoformat()}) note kiya. "
           + (("Filing: " + ", ".join(f"{m['form']} {m['due'] or '(RD / vakil)'}" for m in made) + ". ") if made else "")
           + (" ".join(notes) + " " if notes else "") + R.DRAFT_LINE)
    out = {"event": kind, "label": R.EVENTS[kind]["label"], "on": d0.isoformat(), "filings": made, "notes": notes, "say": say}
    _publish(tenant, "cosec_event", {"entity": name, **out}, conv)
    return out


# ---------------------------------------------------------------- filing card (HER tap, fingerprint once a key is active)
_STEPS = ["mca.gov.in par apne MCA V3 business user ID se login (aap / director khud)",
          "MCA Services -> E-Filing -> {group} Forms -> {form}",
          "CIN / LLPIN daal kar prefill dekhiye; neeche ki jaankari milaiye",
          "Attachments lagaiye; jahan zaroori ho PCS / CA certification",
          "Director / designated partner ka DSC lagaiye (Hoola DSC nahi chhoota)",
          "Submit -> fee portal par khud bhariye -> SRN aur challan sambhaliye",
          "SRN yahan 'Ho gaya' mein bataiye"]


def _filing_card(s, f: CosecFiling, tenant: str) -> dict:
    e = s.get(CosecEntity, f.entity_id)
    on = today(tenant)
    v = _filing_view(f, on)
    ev = _j(f.event, {})
    blockers, checks = [], []
    if e.kind == "llp":
        if not e.llpin:
            blockers.append("LLPIN entity mein daaliye.")
    elif not e.cin and f.form not in ("INC-20A",):
        blockers.append("CIN entity mein daaliye.")
    signers = s.scalars(select(CosecPerson).where(CosecPerson.entity_id == e.id, CosecPerson.role.in_(("director", "partner")),
                                                  CosecPerson.ceased_on == "")).all()
    if not [p for p in signers if p.din]:
        blockers.append("Sign karne wale director / designated partner ka DIN chahiye.")
    if f.form in ("AOC-4", "MGT-7", "MGT-7A") and v["est"]:
        blockers.append("AGM abhi hua nahi (ya meeting mein darj nahi): AGM ke baad hi yeh form.")
    if f.form in ("AOC-4", "AOC-4 OPC", "ADT-1") and not _j(e.auditor, {}).get("name"):
        blockers.append("Auditor ka naam entity mein daaliye.")
    if f.form == "ADT-1" and not _j(e.auditor, {}).get("frn"):
        checks.append("Auditor ka FRN (firm registration no.) bhi rakhiye.")
    if f.form in ("MGT-7",) and R.mgt8_needed(e.kind, e.paid_up, e.turnover):
        checks.append("MGT-8 (PCS ka certificate) chahiye.")
    if f.form == "DIR-12" and ev.get("kind") == "director_appoint" and not ev.get("din"):
        blockers.append("Naye director ka DIN chahiye.")
    if f.form == "PAS-3" and not ev.get("allottees") and not ev.get("total_shares"):
        blockers.append("Allottees ki list chahiye.")
    if f.form in ("AGM", "FIRST-BOARD", "FIRST-AUDITOR"):
        blockers.append("Yeh MCA form nahi, meeting / Board ka kaam hai: Meetings mein darj kijiye.")
    if f.form == "INC-23":
        blockers.append("Doosre state ka shift Regional Director ke paas jaata hai: vakil / CS ka kaam.")
    if v["late"]:
        checks.append(f"{-v['days_left']} din der: {v['fee']['text']}")
        if f.form in ("AOC-4", "MGT-7", "MGT-7A", "AOC-4 OPC") and -v["days_left"] > 365:
            checks.append("Lambi der: s.164(2) (3 saal non-filing par directors disqualify) ka khatra - CS se baat kijiye.")
    if v["est"]:
        checks.append("Tareekh andaaza hai: " + f.basis)
    group = "LLP" if e.kind == "llp" else "Company"
    show = {"entity": f"{e.number} {e.name}", "cin": e.cin or e.llpin, "form": f.form, "label": v["label"], "law": v["law"],
            "period": f.period, "due": f.due, "basis": f.basis, "event": ev, "blockers": blockers, "checks": checks,
            "signers": [f"{p.name} (DIN {p.din})" for p in signers if p.din][:4],
            "steps": [x.format(group=group, form=f.form) for x in _STEPS], "draft_line": R.DRAFT_LINE}
    return show


def filing_request(tenant: str, fid: str, conv: str = "") -> dict:
    """Prepares the filing card. 'File karo' becomes true ONLY on her tap (decide_filing from the API)."""
    with _LOCK, session() as s:
        f = s.get(CosecFiling, fid)
        if f is None or f.tenant_id != tenant:
            raise CosecError("Filing nahi mili.")
        if f.status in ("filed", "not_needed"):
            raise CosecError(f"{f.form} pehle hi '{f.status}' hai.")
        show = _filing_card(s, f, tenant)
        f.card, f.hash, f.status = _d(show), _hash(show), "waiting"
        s.flush()
        v = _filing_view(f, today(tenant))
    v["say"] = (f"{v['form']} ({v['period']}): " + (f"{len(show['blockers'])} rukawat - pehle: {show['blockers'][0]}" if show["blockers"]
                                                       else "card taiyaar, aapki haan chahiye. MCA par login / DSC / fee aap khud.")) + f" {R.DRAFT_LINE}"
    _publish(tenant, "cosec_filing", {"filing": v, "say": v["say"]}, conv)
    return v


def filing_challenge(tenant: str, fid: str) -> dict:
    """approvekey's message format (kind 'cosec'): her phone signs the filing card as it is NOW."""
    from .. import approvekey as ak
    with session() as s:
        f = s.get(CosecFiling, fid)
        if f is None or f.tenant_id != tenant or f.status != "waiting":
            raise CosecError("Yeh card ab baaki nahi hai.")
        digest, form, period = f.hash, f.form, f.period
        ent = _j(f.card, {}).get("entity", "")
    nonce = secrets.token_urlsafe(18)
    exp = datetime.now(timezone.utc) + ak.CHALLENGE_TTL
    msg = f"hoola-approve|v1|{tenant}|cosec|{fid}|{digest}|{nonce}|{int(exp.timestamp())}"
    with session() as s:
        ch = ak.ApproveChallenge(tenant_id=tenant, kind="cosec", target=fid, digest=digest, message=msg, expires_at=exp)
        s.add(ch)
        s.flush()
        chid = ch.id
    return {"challenge_id": chid, "message": msg, "show": {"title": f"{form} file karo", "subtitle": f"{ent} · {period}"[:120]},
            "expires_at": exp.isoformat()}


def _verify_fingerprint(tenant: str, fid: str, digest: str, proof: dict | None) -> None:
    from .. import approvekey as ak
    if not proof:
        raise CosecError("fingerprint_needed")
    with session() as s:
        ak._settle(s, tenant)
        ok = ak._check(s, tenant, "cosec", fid, proof, digest)
    if not ok:
        raise CosecError("fingerprint_invalid")


def decide_filing(tenant: str, fid: str, yes: bool, card_hash: str = "", reason: str = "", proof: dict | None = None,
                  by_tap: bool = False) -> dict:
    if not by_tap:
        raise CosecError("Filing ki haan sirf aapke tap se (app ka card).")
    with _LOCK, session() as s:
        f = s.get(CosecFiling, fid)
        if f is None or f.tenant_id != tenant:
            raise CosecError("Card nahi mila.")
        if f.status != "waiting":
            return {"filing": _filing_view(f, today(tenant)), "say": f"{f.form} pehle hi '{f.status}' hai."}
        now_show = _filing_card(s, f, tenant)
        stale = not card_hash or card_hash != f.hash or _hash(now_show) != f.hash
        if stale:
            f.status, f.hash = "todo", ""
    if stale:
        raise CosecError("Jaankari badal gayi hai; naya card banaiye.")
    if yes and now_show["blockers"]:
        raise CosecError("Rukawat baaki hai: " + now_show["blockers"][0])
    if yes:
        from .. import approvekey as ak
        if ak.active_on(tenant):
            _verify_fingerprint(tenant, fid, card_hash, proof)
    with _LOCK, session() as s:
        f = s.get(CosecFiling, fid)
        f.status = "approved" if yes else "todo"
        f.reason = _clean(reason, 300)
        f.decided_at = now()
        v = _filing_view(f, today(tenant))
    if yes:
        v["say"] = f"{v['form']} maana. Ab MCA V3 par aap login karke: " + " -> ".join(v["card"].get("steps", [])[1:6]) + f". {R.DRAFT_LINE}"
    else:
        v["say"] = f"{v['form']}: abhi nahi; list mein wapas."
    return v


def filing_done(tenant: str, fid: str, srn: str = "", filed_on: str = "", fee=None) -> dict:
    with _LOCK, session() as s:
        f = s.get(CosecFiling, fid)
        if f is None or f.tenant_id != tenant:
            raise CosecError("Filing nahi mili.")
        if f.status != "approved" and f.form not in ("AGM", "FIRST-BOARD", "FIRST-AUDITOR", "DIR-3 KYC", "DIR-3 KYC WEB", "ADT-3"):
            raise CosecError(f"{f.form}: pehle card par aapki haan (tap), phir 'ho gaya'.")
        f.status = "filed"
        f.srn = _clean(srn, 40)
        f.filed_on = _iso(filed_on) or today(tenant).isoformat()
        f.fee_paid = _num(fee) if fee not in (None, "") else None
        v = _filing_view(f, today(tenant))
    v["say"] = f"{v['form']} file hua" + (f" (SRN {v['srn']})" if v["srn"] else "") + "." + (
        f" Der hui: {v['fee']['text']}" if v.get("fee") and v["fee"].get("days") else "")
    return v


def filing_skip(tenant: str, fid: str, reason: str) -> dict:
    if not str(reason or "").strip():
        raise CosecError("Kyun zaroori nahi? wajah likhiye (record ke liye).")
    with _LOCK, session() as s:
        f = s.get(CosecFiling, fid)
        if f is None or f.tenant_id != tenant:
            raise CosecError("Filing nahi mili.")
        if f.status == "filed":
            raise CosecError("File ho chuki filing ko 'zaroori nahi' nahi kar sakte.")
        f.status, f.reason = "not_needed", _clean(reason, 300)
        v = _filing_view(f, today(tenant))
    v["say"] = f"{v['form']} ({v['period']}): zaroori nahi - '{v['reason']}'."
    return v


# ---------------------------------------------------------------- meetings (notice -> attendance -> resolutions -> minutes)
def _meet_view(m: CosecMeeting, on: date | None = None) -> dict:
    v = {"id": m.id, "number": m.number, "kind": m.kind, "kind_label": R.MEETING_KINDS.get(m.kind, m.kind), "fy": m.fy, "on": m.on,
         "time": m.time, "place": m.place, "mode": m.mode, "agenda": _j(m.agenda, []), "notice_mode": m.notice_mode,
         "notice_sent_on": m.notice_sent_on, "shorter_consent": m.shorter_consent, "notice_text": m.notice_text, "notice_hash": m.notice_hash,
         "notice_file": m.notice_file, "attendance": _j(m.attendance, []), "quorum": _j(m.quorum, {}), "chair": m.chair, "minutes": m.minutes,
         "minutes_hash": m.minutes_hash, "minutes_file": m.minutes_file, "minutes_signed_on": m.minutes_signed_on, "status": m.status}
    d = _date(m.on)
    if d:
        v["deadlines"] = {k: x.isoformat() for k, x in R.minutes_deadlines(m.kind, d).items()}
        if not m.notice_sent_on:
            v["latest_notice"] = R.notice_check(m.kind, d, on or d, m.notice_mode)["latest_send"].isoformat()
    return v


def _meeting(s, tenant: str, mid: str) -> CosecMeeting:
    m = s.get(CosecMeeting, mid)
    if m is None or m.tenant_id != tenant:
        rows = s.scalars(select(CosecMeeting).where(CosecMeeting.tenant_id == tenant, CosecMeeting.number == str(mid or "").upper())).all()
        if len(rows) == 1:
            return rows[0]
        raise CosecError("Meeting nahi mili.")
    return m


def meeting_new(tenant: str, ref: str, kind: str = "board", on: str = "", time: str = "", place: str = "", mode: str = "physical",
                agenda: list | None = None, fy: str = "", chair: str = "", conv: str = "") -> dict:
    kind = (kind or "board").lower()
    if kind not in R.MEETING_KINDS:
        raise CosecError("kind = board | agm | egm | committee")
    d = _date(on)
    if not d:
        raise CosecError("Meeting ki tareekh chahiye.")
    if mode not in ("physical", "vc", "hybrid"):
        raise CosecError("mode = physical | vc | hybrid")
    if time and not re.match(r"^\d{1,2}:\d{2}(\s?[ap]m)?$", str(time).strip(), re.I):
        raise CosecError("Samay jaise 11:00 ya 4:30 pm.")
    with _LOCK, session() as s:
        e = _ent(s, tenant, ref)
        if e.kind == "llp":
            raise CosecError("LLP mein Board meeting / AGM ka kanooni dhaancha nahi (LLP agreement dekhiye).")
        if e.kind == "opc" and kind in ("agm",):
            raise CosecError("OPC ko AGM nahi karna hota (s.96).")
        items = [_clean(x, 400) for x in (agenda or []) if str(x).strip()] or list(R.AGENDA[kind])
        fyl = ""
        if kind == "agm":
            y = R.fy_end_year(fy) if fy else R.fy_end_year(d) - 1
            if not y:
                raise CosecError("AGM kis FY ke accounts ke liye? (jaise 2025-26)")
            fyl = R.fy_label(y)
        pre = {"board": "BM", "agm": "AGM", "egm": "EGM", "committee": "CM"}[kind]
        m = CosecMeeting(tenant_id=tenant, entity_id=e.id, number=f"{pre}-{_next(s, tenant, pre + ':' + e.id):03d}", kind=kind, fy=fyl,
                         on=d.isoformat(), time=str(time or "").strip(), place=_clean(place, 300), mode=mode, agenda=_d(items),
                         chair=_clean(chair, 120))
        s.add(m)
        s.flush()
        v = _meet_view(m, today(tenant))
        ename = f"{e.number} {e.name}"
    checks = []
    if kind == "agm" and fyl:
        last = R.add_months(date(R.fy_end_year(fyl), 3, 31), 6)
        if d > last:
            checks.append(f"AGM ki aam aakhri tareekh {last.isoformat()} (pehle AGM ko 9 mahine): yeh baad mein hai - "
                          "ROC se extension ya der ka asar, CS se pakka.")
    v["checks"] = checks
    v["say"] = (f"{ename}: {v['number']} {v['kind_label']} {v['on']} banayi. Notice {v.get('latest_notice', '')} tak jaana chahiye "
                f"({R.notice_need(kind)} clear din)." + (" " + checks[0] if checks else ""))
    return v


def meeting_update(tenant: str, mid: str, **kw) -> dict:
    with _LOCK, session() as s:
        m = _meeting(s, tenant, mid)
        if m.status not in ("draft", "notice_waiting"):
            raise CosecError(f"{m.number} ka notice ja chuka: tareekh / agenda badalna = naya notice (nayi meeting banaiye).")
        if kw.get("on"):
            d = _date(kw["on"])
            if not d:
                raise CosecError("Tareekh samajh nahi aayi.")
            m.on = d.isoformat()
        for k in ("time", "place", "chair"):
            if kw.get(k) is not None:
                setattr(m, k, _clean(kw[k], 300))
        if kw.get("mode") in ("physical", "vc", "hybrid"):
            m.mode = kw["mode"]
        if kw.get("agenda"):
            m.agenda = _d([_clean(x, 400) for x in kw["agenda"] if str(x).strip()])
        if m.status == "notice_waiting":
            m.status = "draft"
            m.notice_hash = ""
        v = _meet_view(m, today(tenant))
    v["say"] = f"{v['number']} badli." + (" Purana notice card ab bekaar: naya card banaiye." if kw else "")
    return v


def _recipients(s, e: CosecEntity, kind: str) -> list[dict]:
    rows = s.scalars(select(CosecPerson).where(CosecPerson.entity_id == e.id, CosecPerson.ceased_on == "")).all()
    roles = ("director",) if kind in ("board", "committee") else ("director", "member")
    out = [{"name": p.name, "role": p.role, "email": p.email, "phone": p.phone} for p in rows if p.role in roles]
    aud = _j(e.auditor, {})
    if kind in ("agm", "egm") and aud.get("name"):
        out.append({"name": aud["name"], "role": "auditor", "email": "", "phone": ""})
    return out


def _notice_text(e: CosecEntity, m: CosecMeeting, st: CosecSettings) -> str:
    d = _date(m.on)
    when = d.strftime("%A, %d %B %Y") if d else "[tareekh]"
    place = ("through Video Conferencing (VC) / Other Audio Visual Means; deemed venue: Registered Office" if m.mode == "vc"
             else (m.place or "[jagah]"))
    head = {"board": "Meeting of the Board of Directors", "committee": "Committee Meeting", "agm": "Annual General Meeting",
            "egm": "Extra-ordinary General Meeting"}[m.kind]
    lines = [f"{e.name}", f"CIN: {e.cin or '[CIN]'}", f"Regd. Office: {e.registered_office or '[registered office]'}", "",
             f"NOTICE of {head} ({m.number})", "",
             f"NOTICE is hereby given that the {head} of {e.name} will be held on {when} at {m.time or '[samay]'} {place}, "
             "to transact the following business:", ""]
    for i, a in enumerate(_j(m.agenda, []), 1):
        lines.append(f"{i}. {a}")
    if m.kind in ("agm", "egm"):
        lines += ["", "Notes:",
                  "1. A member entitled to attend and vote is entitled to appoint a proxy to attend and vote instead of himself / herself; "
                  "a proxy need not be a member. [Check: proxy rules for this company / VC meeting]",
                  "2. Explanatory statement under Section 102 for special business: [chahiye, agar special business hai]"]
    if m.shorter_consent:
        lines += ["", "This meeting is being convened at a shorter notice. [Consent on record]"]
    lines += ["", "By order of the Board", f"For {e.name}", "", f"{st.cs_name or '[Company Secretary / Director]'}"
              + (f" ({st.membership_no})" if st.membership_no else ""), f"Date: [notice ki tareekh]", "",
              f"[{R.DRAFT_LINE}]"]
    return "\n".join(lines)


def _notice_show(s, tenant: str, m: CosecMeeting, send_on: date) -> dict:
    e = s.get(CosecEntity, m.entity_id)
    st = _st(s, tenant)
    d = _date(m.on)
    chk = R.notice_check(m.kind, d, send_on, m.notice_mode, m.shorter_consent)
    text = _notice_text(e, m, st)
    blockers = []
    for ph in ("[samay]", "[jagah]", "[CIN]", "[registered office]"):
        if ph in text:
            blockers.append(f"Notice mein {ph} khaali hai.")
    if not chk["ok"]:
        blockers.append(chk["note"])
    if send_on > d:
        blockers.append("Meeting ki tareekh nikal chuki.")
    rec = _recipients(s, e, m.kind)
    if not rec:
        blockers.append("Directors / members ki list khaali: pehle log jodiye.")
    no_contact = [r["name"] for r in rec if not r["email"] and not r["phone"] and r["role"] != "auditor"]
    checks = [chk["note"]] + ([f"Contact nahi: {', '.join(no_contact[:5])} (haath se notice)"] if no_contact else [])
    return {"meeting": m.number, "entity": f"{e.number} {e.name}", "kind": m.kind, "on": m.on, "send_on": send_on.isoformat(),
            "mode": m.notice_mode, "text": text, "recipients": rec, "blockers": blockers, "checks": checks, "draft_line": R.DRAFT_LINE}


def notice_request(tenant: str, mid: str, send_on: str = "", mode: str = "", shorter_consent: bool | None = None, conv: str = "") -> dict:
    with _LOCK, session() as s:
        m = _meeting(s, tenant, mid)
        if m.status not in ("draft", "notice_waiting"):
            raise CosecError(f"{m.number} ka notice pehle hi '{m.status}'.")
        if mode:
            if mode not in ("email", "hand", "post", "courier"):
                raise CosecError("mode = email | hand | post | courier")
            m.notice_mode = mode
        if shorter_consent is not None:
            m.shorter_consent = bool(shorter_consent)
        sd = _date(send_on) or today(tenant)
        show = _notice_show(s, tenant, m, sd)
        m.notice_text, m.notice_hash, m.status = show["text"], _hash(show), "notice_waiting"
        m.notice_plan_on = sd.isoformat()
        s.flush()
        v = _meet_view(m, today(tenant)) | {"card": show}
    v["say"] = (f"{v['number']} ka notice: " + (f"rukawat - {show['blockers'][0]}" if show["blockers"] else
                                                  f"taiyaar ({len(show['recipients'])} logon ko). Aapki haan ka card bheja.") + f" {R.DRAFT_LINE}")
    _publish(tenant, "cosec_notice", {"meeting": v, "say": v["say"]}, conv)
    return v


def decide_notice(tenant: str, mid: str, yes: bool, card_hash: str = "", send_on: str = "", by_tap: bool = False, conv: str = "") -> dict:
    if not by_tap:
        raise CosecError("Notice bhejne ki haan sirf aapke tap se (app ka card).")
    with _LOCK, session() as s:
        m = _meeting(s, tenant, mid)
        if m.status != "notice_waiting":
            return {"meeting": _meet_view(m), "say": f"{m.number} pehle hi '{m.status}'."}
        sd = _date(send_on) or _date(m.notice_plan_on) or today(tenant)
        show = _notice_show(s, tenant, m, sd)
        stale = not card_hash or card_hash != m.notice_hash or _hash(show) != m.notice_hash
        if stale and _hash(show) != m.notice_hash:
            m.status = "draft"
        if not stale and not yes:
            m.status = "draft"
            return {"meeting": _meet_view(m), "say": f"{m.number}: notice abhi nahi."}
        text, number, ent = m.notice_text, m.number, show["entity"]
    if stale:
        raise CosecError("Notice badal gaya (ya tareekh) - naya card banaiye.")
    if show["blockers"]:
        raise CosecError("Rukawat: " + show["blockers"][0])
    from .. import officegen
    f = _store(tenant, f"Notice_{number}.docx", officegen.make_docx(f"Notice {number}", text), DOCX, conv, f"Notice {number}")
    subj = f"Notice: {ent.split(' ', 1)[-1]} {number} on {show['on']}"
    drafts = [{"name": r["name"], "role": r["role"], "email": _mail(r["email"], subj, text) if r["email"] else "",
               "whatsapp": _wa(r["phone"], text) if r["phone"] else ""} for r in show["recipients"]]
    with _LOCK, session() as s:
        m = _meeting(s, tenant, mid)
        m.notice_sent_on, m.notice_file, m.status = sd.isoformat(), f.get("file_id", ""), "notice_sent"
        v = _meet_view(m)
    return {"meeting": v, "file": f, "drafts": drafts, "text": text,
            "say": f"{number} notice taiyaar (Word file). Har director / member ko aap bhejiye (email / WhatsApp drafts); Hoola khud nahi bhejta."}


def attendance(tenant: str, mid: str, present: list | None = None, interested: list | None = None, chair: str = "",
               members_count: int | None = None, held: bool = True) -> dict:
    present = [str(x).strip() for x in (present or []) if str(x).strip()]
    if not present:
        raise CosecError("Kaun maujood tha? naam / DIN ki list.")
    interested_set = {str(x).strip().lower() for x in (interested or [])}
    with _LOCK, session() as s:
        m = _meeting(s, tenant, mid)
        e = s.get(CosecEntity, m.entity_id)
        if m.status in ("minutes_waiting", "minuted", "cancelled"):
            raise CosecError(f"{m.number} '{m.status}': attendance ab nahi badlegi.")
        if m.status in ("draft", "notice_waiting") and not m.shorter_consent:
            raise CosecError("Notice abhi gaya nahi (card par aapki haan). Shorter notice ho to notice card mein bataiye.")
        dirs = s.scalars(select(CosecPerson).where(CosecPerson.entity_id == e.id, CosecPerson.role == "director", CosecPerson.ceased_on == "")).all()
        rows, unknown = [], []
        for name in present:
            p = _find_person(s, e, name)
            rows.append({"person_id": p.id if p else "", "name": p.name if p else _clean(name, 120), "present": True,
                         "interested": (p.name.lower() if p else name.lower()) in interested_set or name.lower() in interested_set})
            if p is None:
                unknown.append(name)
        names = {r["name"].lower() for r in rows}
        for p in dirs:
            if p.name.lower() not in names and m.kind in ("board", "committee"):
                rows.append({"person_id": p.id, "name": p.name, "present": False, "interested": False})
        here = [r for r in rows if r["present"]]
        q = R.quorum(m.kind, strength=len(dirs), members=int(members_count or 0), entity_kind=e.kind, present=len(here),
                     interested=sum(1 for r in here if r["interested"]))
        m.attendance, m.quorum = _d(rows), _d(q)
        if chair:
            m.chair = _clean(chair, 120)
        if held:
            m.status = "held"
            if m.kind == "agm":
                e.last_agm = m.on
        s.flush()
        v = _meet_view(m)
        fyl = m.fy
        eref = e.id
        kind = m.kind
    if kind == "agm" and fyl and held:
        calendar(tenant, eref, fyl)
    v["say"] = (f"{v['number']}: {len([r for r in rows if r['present']])} maujood, quorum {q['need']} - " + ("poora." if q["ok"] else
                "POORA NAHI: meeting adjourn / dobara (s.103 / s.174), resolutions mat likhiye.") + (f" Anjaan naam: {', '.join(unknown)}." if unknown else ""))
    return v


def resolution(tenant: str, mid: str = "", ref: str = "", kind: str = "board", text: str = "", passed_on: str = "", votes_for=None,
               votes_against=None, mgt14: bool = False) -> dict:
    kind = (kind or "board").lower()
    if kind not in ("board", "ordinary", "special", "circular"):
        raise CosecError("kind = board | ordinary | special | circular")
    if not str(text or "").strip():
        raise CosecError("Resolution ka text chahiye.")
    with _LOCK, session() as s:
        m = None
        if mid:
            m = _meeting(s, tenant, mid)
            e = s.get(CosecEntity, m.entity_id)
            if m.status != "held":
                raise CosecError(f"{m.number} abhi '{m.status}': resolution meeting hone ke baad (attendance + quorum) aur minutes se pehle.")
            if not _j(m.quorum, {}).get("ok"):
                raise CosecError("Quorum poora nahi tha: resolution nahi.")
            if kind in ("ordinary", "special") and m.kind not in ("agm", "egm"):
                raise CosecError("Ordinary / special resolution general meeting mein; Board mein 'board'.")
            if kind == "board" and m.kind not in ("board", "committee"):
                raise CosecError("General meeting mein 'ordinary' ya 'special'.")
        else:
            e = _ent(s, tenant, ref)
            if kind != "circular":
                raise CosecError("Meeting ke bina sirf circular resolution (Board).")
        d0 = _date(passed_on) or (_date(m.on) if m else today(tenant))
        vf, va = (_num(votes_for) if votes_for not in (None, "") else None), (_num(votes_against) if votes_against not in (None, "") else None)
        result = "passed"
        note = ""
        if vf is not None and va is not None:
            if kind == "special" and vf < 3 * va:
                result, note = "failed", "Special resolution ko 'against' ke kam se kam teen guna 'for' chahiye (s.114(2))."
            elif kind in ("ordinary", "board", "circular") and vf <= va:
                result, note = "failed", "'For' 'against' se zyada nahi."
        y = d0.year
        r = CosecResolution(tenant_id=tenant, entity_id=e.id, meeting_id=m.id if m else "", kind=kind, text=_clean(text, 3000),
                            passed_on=d0.isoformat(), votes_for=vf, votes_against=va, result=result,
                            number=f"R-{y}/{_next(s, tenant, f'res:{e.id}:{y}'):03d}")
        s.add(r)
        s.flush()
        rid, rnum, eid = r.id, r.number, e.id
    filing = None
    if result == "passed" and (kind == "special" or mgt14):
        out = event(tenant, eid, "special_resolution", d0.isoformat(), details={"name": rnum})
        filing = out["filings"][0] if out["filings"] else None
        if filing:
            with session() as s:
                s.get(CosecResolution, rid).filing_id = filing["id"]
    v = {"id": rid, "number": rnum, "kind": kind, "result": result, "passed_on": d0.isoformat(), "filing": filing}
    v["say"] = (f"{rnum} ({kind}) {'pass' if result == 'passed' else 'FAIL'}." + (f" {note}" if note else "")
                + (f" MGT-14 {filing['due']} tak." if filing else "")
                + (" Private company ki kuch Board resolutions par MGT-14 nahi lagta - CS se pakka." if kind == "board" and not mgt14 else ""))
    return v


def resolutions(tenant: str, ref: str) -> list[dict]:
    with session() as s:
        e = _ent(s, tenant, ref)
        return [{"id": r.id, "number": r.number, "kind": r.kind, "text": r.text, "passed_on": r.passed_on, "result": r.result,
                 "votes_for": r.votes_for, "votes_against": r.votes_against, "meeting_id": r.meeting_id, "filing_id": r.filing_id}
                for r in s.scalars(select(CosecResolution).where(CosecResolution.entity_id == e.id).order_by(CosecResolution.created_at)).all()]


def _minutes_text(s, m: CosecMeeting) -> str:
    e = s.get(CosecEntity, m.entity_id)
    att = _j(m.attendance, [])
    q = _j(m.quorum, {})
    head = R.MEETING_KINDS.get(m.kind, m.kind)
    lines = [f"MINUTES of the {head} ({m.number}) of {e.name} held on {m.on} at {m.time or '[samay]'} "
             f"{'through VC' if m.mode == 'vc' else ('at ' + (m.place or '[jagah]'))}", "",
             "Present: " + (", ".join(r["name"] for r in att if r.get("present")) or "[ ]")]
    absent = [r["name"] for r in att if not r.get("present")]
    if absent:
        lines.append("Leave of absence: " + ", ".join(absent))
    lines += [f"Chairperson: {m.chair or '[Chairperson]'}", "",
              f"Quorum: {q.get('present', '?')} present against {q.get('need', '?')} required - the quorum was "
              + ("present." if q.get("ok") else "NOT present."), ""]
    res = s.scalars(select(CosecResolution).where(CosecResolution.meeting_id == m.id).order_by(CosecResolution.created_at)).all()
    for i, a in enumerate(_j(m.agenda, []), 1):
        lines.append(f"{i}. {a}")
        lines.append("   [Discussion: kya baat hui]")
    if res:
        lines += ["", "Resolutions:"]
        for r in res:
            lines.append(f"{r.number} ({r.kind} resolution, {r.result}): RESOLVED THAT {r.text}"
                         + (f" [For: {r.votes_for:g}, Against: {r.votes_against:g}]" if r.votes_for is not None and r.votes_against is not None else ""))
    lines += ["", "There being no other business, the meeting concluded with a vote of thanks to the Chair.", "",
              "Chairperson's signature: ____________   Date of signing: ________   Place: ________", "", f"[{R.DRAFT_LINE}]"]
    return "\n".join(lines)


def minutes_draft(tenant: str, mid: str) -> dict:
    with _LOCK, session() as s:
        m = _meeting(s, tenant, mid)
        if m.status not in ("held", "minutes_waiting"):
            raise CosecError(f"{m.number} '{m.status}': minutes meeting hone ke baad (attendance).")
        if not m.minutes:
            m.minutes = _minutes_text(s, m)
        v = _meet_view(m, today(tenant))
    dl = v.get("deadlines", {})
    v["say"] = (f"{v['number']} minutes ka draft. " + (f"Draft directors ko {dl.get('draft_by')} tak, comments {dl.get('comments_by')} tak, "
                                                       if "draft_by" in dl else "") + f"sign {dl.get('sign_by')} tak. [..] wali jagah bhariye.")
    return v


def minutes_edit(tenant: str, mid: str, text: str) -> dict:
    if not str(text or "").strip():
        raise CosecError("Minutes ka text chahiye.")
    with _LOCK, session() as s:
        m = _meeting(s, tenant, mid)
        if m.status not in ("held", "minutes_waiting"):
            raise CosecError(f"{m.number} '{m.status}': signed minutes badle nahi jaate.")
        m.minutes = _clean(text, 20000)
        if m.status == "minutes_waiting":
            m.status = "held"
            m.minutes_hash = ""
        v = _meet_view(m)
    v["say"] = f"{v['number']} minutes badle." + (" Purana sign card bekaar." if v["status"] == "held" else "")
    return v


def _minutes_show(s, m: CosecMeeting) -> dict:
    e = s.get(CosecEntity, m.entity_id)
    blockers = [f"Minutes mein '{ph}' khaali" for ph in ("[samay]", "[jagah]", "[Chairperson]", "[Discussion: kya baat hui]") if ph in (m.minutes or "")]
    if not m.minutes:
        blockers.append("Minutes ka draft nahi bana.")
    return {"meeting": m.number, "entity": f"{e.number} {e.name}", "on": m.on, "text": m.minutes, "blockers": blockers, "draft_line": R.DRAFT_LINE}


def minutes_request(tenant: str, mid: str, conv: str = "") -> dict:
    with _LOCK, session() as s:
        m = _meeting(s, tenant, mid)
        if m.status not in ("held", "minutes_waiting"):
            raise CosecError(f"{m.number} '{m.status}'.")
        if not m.minutes:
            m.minutes = _minutes_text(s, m)
        show = _minutes_show(s, m)
        m.minutes_hash, m.status = _hash(show), "minutes_waiting"
        v = _meet_view(m, today(tenant)) | {"card": show}
    v["say"] = f"{v['number']} minutes: " + (f"rukawat - {show['blockers'][0]}" if show["blockers"] else "chairperson ke sign ke liye taiyaar, aapki haan ka card.")
    _publish(tenant, "cosec_minutes", {"meeting": v, "say": v["say"]}, conv)
    return v


def decide_minutes(tenant: str, mid: str, yes: bool, card_hash: str = "", signed_on: str = "", by_tap: bool = False, conv: str = "") -> dict:
    if not by_tap:
        raise CosecError("Minutes pakke sirf aapke tap se (app ka card).")
    with _LOCK, session() as s:
        m = _meeting(s, tenant, mid)
        if m.status != "minutes_waiting":
            return {"meeting": _meet_view(m), "say": f"{m.number} '{m.status}'."}
        show = _minutes_show(s, m)
        stale = not card_hash or card_hash != m.minutes_hash or _hash(show) != m.minutes_hash
        if stale or not yes:
            m.status = "held"
        if not stale and not yes:
            return {"meeting": _meet_view(m), "say": f"{m.number}: minutes abhi pakke nahi."}
        text, number = m.minutes, m.number
        held = _date(m.on)
    if stale:
        raise CosecError("Minutes badal gaye - naya card banaiye.")
    if show["blockers"]:
        raise CosecError("Rukawat: " + show["blockers"][0])
    sd = _date(signed_on) or today(tenant)
    late = (sd - held).days > 30 if held else False
    from .. import officegen
    f = _store(tenant, f"Minutes_{number}.docx", officegen.make_docx(f"Minutes {number}", text), DOCX, conv, f"Minutes {number}")
    with _LOCK, session() as s:
        m = _meeting(s, tenant, mid)
        m.minutes_signed_on, m.minutes_file, m.status = sd.isoformat(), f.get("file_id", ""), "minuted"
        v = _meet_view(m)
    return {"meeting": v, "file": f, "say": f"{number} minutes pakke (Word file). Chairperson har page initial + aakhri page sign kare; minutes book mein lagaiye."
            + (" Dhyaan: meeting ke 30 din baad sign ho rahe hain (s.118 / SS)." if late else "")}


def meetings(tenant: str, ref: str = "", status: str = "") -> list[dict]:
    on = today(tenant)
    with session() as s:
        q = select(CosecMeeting).where(CosecMeeting.tenant_id == tenant)
        if ref:
            q = q.where(CosecMeeting.entity_id == _ent(s, tenant, ref).id)
        if status:
            q = q.where(CosecMeeting.status == status)
        return [_meet_view(m, on) for m in s.scalars(q.order_by(CosecMeeting.on.desc())).all()[:200]]


# ---------------------------------------------------------------- entity view, registers
def entity_view(tenant: str, ref: str) -> dict:
    on = today(tenant)
    with session() as s:
        e = _ent(s, tenant, ref)
        v = _ent_view(e)
        ppl = [_person_view(p) for p in s.scalars(select(CosecPerson).where(CosecPerson.entity_id == e.id).order_by(CosecPerson.created_at)).all()]
        v["directors"] = [p for p in ppl if p["role"] in ("director", "partner")]
        v["kmp"] = [p for p in ppl if p["role"] == "kmp"]
        v["members"] = [p for p in ppl if p["role"] == "member"]
        v["charges"] = [{"id": c.id, "holder": c.holder, "amount": c.amount, "property": c.property, "created_on": c.created_on,
                         "charge_id": c.charge_id, "satisfied_on": c.satisfied_on}
                        for c in s.scalars(select(CosecCharge).where(CosecCharge.entity_id == e.id)).all()]
        boards = [_date(m.on) for m in s.scalars(select(CosecMeeting).where(CosecMeeting.entity_id == e.id, CosecMeeting.kind == "board",
                                                                         CosecMeeting.status.in_(("held", "minutes_waiting", "minuted")))).all()]
        light = _light(e)
        single = e.kind == "opc" and len([p for p in v["directors"] if p["active"]]) <= 1
        eid = e.id
    v["meetings"] = meetings(tenant, eid)
    v["filings"] = filings(tenant, eid)
    v["resolutions"] = resolutions(tenant, eid)
    v["board_issues"] = R.board_frequency([d for d in boards if d], on.year, light, single) if e.kind != "llp" else []
    last = max([d for d in boards if d], default=None)
    nb = R.next_board_by(last, light)
    v["next_board_by"] = nb.isoformat() if nb else ""
    return v


def registers_file(tenant: str, ref: str, conv: str = "") -> dict:
    import openpyxl
    v = entity_view(tenant, ref)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Members (MGT-1)"
    ws.append(["Folio", "Name", "Class", "Shares", "Nominal value", "Member since", "Ceased", "Movements"])
    for p in v["members"]:
        ws.append([p["folio"], p["name"], p["share_class"], p["shares"], p["nominal"], p["member_since"], p["ceased_on"],
                   "; ".join(f"{h['on']} {h['what']} {h['shares']:g}" for h in p["history"])])
    ws = wb.create_sheet("Directors & KMP")
    ws.append(["Name", "Role", "DIN / DPIN", "Designation", "Appointed", "Ceased", "KYC on"])
    for p in v["directors"] + v["kmp"]:
        ws.append([p["name"], p["role"], p["din"], p["designation"], p["appointed_on"], p["ceased_on"], p["kyc_on"]])
    ws = wb.create_sheet("Charges")
    ws.append(["Charge ID", "Holder", "Amount", "Property", "Created", "Satisfied"])
    for c in v["charges"]:
        ws.append([c["charge_id"], c["holder"], c["amount"], c["property"], c["created_on"], c["satisfied_on"]])
    ws = wb.create_sheet("Resolutions")
    ws.append(["No", "Kind", "Passed on", "Result", "For", "Against", "Text"])
    for r in v["resolutions"]:
        ws.append([r["number"], r["kind"], r["passed_on"], r["result"], r["votes_for"], r["votes_against"], r["text"]])
    ws = wb.create_sheet("Meetings")
    ws.append(["No", "Kind", "On", "Notice sent", "Quorum", "Minutes signed", "Status"])
    for m in v["meetings"]:
        ws.append([m["number"], m["kind_label"], m["on"], m["notice_sent_on"], "ok" if m["quorum"].get("ok") else "", m["minutes_signed_on"], m["status"]])
    ws = wb.create_sheet("Filings")
    ws.append(["Form", "Period", "Due", "Basis", "Andaaza", "Status", "SRN", "Filed on"])
    for f in v["filings"]:
        ws.append([f["form"], f["period"], f["due"], f["basis"], "haan" if f["est"] else "", f["status"], f["srn"], f["filed_on"]])
    ws = wb.create_sheet("Note")
    ws.append([R.DRAFT_LINE])
    ws.append(["Statutory registers ka format Companies (Management and Administration) Rules 2014 ke hisaab se CS milaye; yeh working copy hai."])
    bio = io.BytesIO()
    wb.save(bio)
    f = _store(tenant, f"Registers_{v['number']}.xlsx", bio.getvalue(), XLSX, conv, f"Registers {v['number']}")
    return {"file": f, "say": f"{v['number']} {v['name']}: members / directors / charges / resolutions / meetings / filings register Excel."}


def forget(tenant: str, ref: str) -> dict:
    """DPDP: after the client company / LLP is closed (left her practice / struck off), personal details are blanked."""
    with _LOCK, session() as s:
        e = _ent(s, tenant, ref)
        if e.status == "active":
            raise CosecError("Pehle entity ka status 'closed' / 'struck_off' kijiye; chalu client ka data nahi mitta.")
        n = 0
        for i, p in enumerate(s.scalars(select(CosecPerson).where(CosecPerson.entity_id == e.id)).all(), 1):
            p.name, p.email, p.phone, p.address, p.id_ref, p.din = f"Purane {p.role} {i}", "", "", "", "", ""
            n += 1
        e.email, e.notes, e.forgotten = "", "", True
        for m in s.scalars(select(CosecMeeting).where(CosecMeeting.entity_id == e.id)).all():
            m.attendance, m.notice_text, m.minutes = "[]", "", ""
        name = f"{e.number} {e.name}"
    return {"forgotten": n, "say": f"{name}: {n} logon ki niji jaankari mita di (filing / tareekh ka record bacha)."}


# ---------------------------------------------------------------- day note + tick
def day(tenant: str) -> dict:
    on = today(tenant)
    out: dict[str, Any] = {"overdue": [], "soon": [], "waiting": [], "meetings": [], "minutes": [], "board": []}
    with session() as s:
        ents = {e.id: e for e in s.scalars(select(CosecEntity).where(CosecEntity.tenant_id == tenant, CosecEntity.status == "active")).all()}
        for f in s.scalars(select(CosecFiling).where(CosecFiling.tenant_id == tenant, CosecFiling.status.in_(("todo", "waiting", "approved")))).all():
            e = ents.get(f.entity_id)
            due = _date(f.due)
            if not e or not due:
                continue
            tag = f"{e.number} {f.form} ({due.isoformat()}{', andaaza' if f.est else ''})"
            if f.status == "waiting":
                out["waiting"].append(f"{e.number} {f.form} card")
            if due < on:
                out["overdue"].append(tag + f" {(on - due).days} din der")
            elif (due - on).days <= 30:
                out["soon"].append(tag)
        for m in s.scalars(select(CosecMeeting).where(CosecMeeting.tenant_id == tenant)).all():
            e = ents.get(m.entity_id)
            d = _date(m.on)
            if not e or not d:
                continue
            if m.status in ("notice_waiting", "minutes_waiting"):
                out["waiting"].append(f"{e.number} {m.number} {'notice' if m.status == 'notice_waiting' else 'minutes'}")
            if m.status in ("draft", "notice_waiting") and d >= on:
                ls = R.notice_check(m.kind, d, on, m.notice_mode)["latest_send"]
                if (ls - on).days <= 3:
                    out["meetings"].append(f"{e.number} {m.number} {m.on}: notice {ls.isoformat()} tak bhejna" + (" (nikal gaya)" if ls < on else ""))
            if m.status in ("held", "minutes_waiting"):
                sign = R.minutes_deadlines(m.kind, d)["sign_by"]
                if (sign - on).days <= 10:
                    out["minutes"].append(f"{e.number} {m.number} minutes {sign.isoformat()} tak sign")
        for e in ents.values():
            if e.kind == "llp":
                continue
            boards = [_date(m.on) for m in s.scalars(select(CosecMeeting).where(CosecMeeting.entity_id == e.id, CosecMeeting.kind == "board",
                                                                             CosecMeeting.status.in_(("held", "minutes_waiting", "minuted")))).all()]
            nb = R.next_board_by(max([d for d in boards if d], default=None), _light(e))
            if nb and (nb - on).days <= 15:
                out["board"].append(f"{e.number} agli Board meeting {nb.isoformat()} se pehle (120 din)")
    bits = []
    for key, title in (("overdue", "Der ho gayi"), ("waiting", "Aapki haan baaki"), ("soon", "30 din mein"), ("meetings", "Notice"),
                       ("minutes", "Minutes"), ("board", "Board")):
        if out[key]:
            bits.append(f"{title}: {', '.join(out[key][:4])}" + (f" +{len(out[key]) - 4}" if len(out[key]) > 4 else ""))
    out["say"] = ". ".join(bits) + "." if bits else "Company secretary desk par aaj sab theek."
    return out


def _tenants() -> list[str]:
    with session() as s:
        return sorted(set(s.scalars(select(CosecEntity.tenant_id)).all()) | set(s.scalars(select(CosecSettings.tenant_id)).all()))


def tick(when: datetime | None = None) -> dict:
    """Worker (every minute is fine): calendars kept up to date, reminders once per window (her windows, default 30/15/7/1/0 days)
    + once a week while overdue, morning note at her time."""
    n = when or now()
    if n.tzinfo is None:
        n = n.replace(tzinfo=timezone.utc)
    out = {"alerts": 0, "morning": 0}
    for tenant in _tenants():
        try:
            loc = n.astimezone(_tz(tenant))
            on = loc.date()
            alerts = []
            morning = False
            with _LOCK, session() as s:
                st = _st(s, tenant)
                wins = _windows(st)
                ents = {e.id: e for e in s.scalars(select(CosecEntity).where(CosecEntity.tenant_id == tenant, CosecEntity.status == "active")).all()}
                for f in s.scalars(select(CosecFiling).where(CosecFiling.tenant_id == tenant, CosecFiling.status.in_(("todo", "waiting", "approved")))).all():
                    e, due = ents.get(f.entity_id), _date(f.due)
                    if not e or not due:
                        continue
                    left = (due - on).days
                    if left < 0:
                        if _once(s, tenant, f"late:{f.id}:{(-left) // 7}"):
                            fee = R.additional_fee(f.form, -left)
                            alerts.append(f"{e.number} {e.name}: {f.form} ({f.period}) {-left} din der. {fee['text']}")
                        continue
                    win = next((w for w in wins if left <= w), None)
                    if win is not None and _once(s, tenant, f"due:{f.id}:{f.due}:{win}"):
                        alerts.append(f"{e.number} {e.name}: {f.form} ({f.period}) " + ("aaj aakhri din" if left == 0 else f"{left} din mein ({f.due})")
                                      + (" - tareekh andaaza" if f.est else "") + ".")
                try:
                    hh, mm = (int(x) for x in (st.morning_time or "09:00").split(":"))
                except ValueError:
                    hh, mm = 9, 0
                if st.last_morning != on.isoformat() and (loc.hour, loc.minute) >= (hh, mm):
                    st.last_morning = on.isoformat()
                    morning = True
            if morning:
                ensure(tenant)
            for a in alerts:
                _publish(tenant, "cosec_alert", {"say": a})
                out["alerts"] += 1
            if morning:
                d = day(tenant)
                if d["say"] != "Company secretary desk par aaj sab theek.":
                    _publish(tenant, "cosec_day", {"say": d["say"], "overdue": len(d["overdue"]), "soon": len(d["soon"]), "waiting": len(d["waiting"])})
                    out["morning"] += 1
        except Exception:
            log.exception("cosec tick failed for %s", tenant)
    return out
