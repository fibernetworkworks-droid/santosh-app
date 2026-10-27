"""App Review Sahayak: review HER app (APK / AAB / source ZIP / GitHub repo) or website -> risks in plain Hinglish with evidence and
a fix + code snippet each -> report PDF / Excel + a fix list for her developer or the Developer Sahayak -> re-review shows fixed /
still open / new. Static reading only; nothing is run or uploaded to anyone. Sending the fix list to the Developer Sahayak waits
for her tap (card + hash, POST /api/appreview/handoffs/{id}/decide)."""
from __future__ import annotations

import hashlib
import io
import json
import logging
import os
import re
import secrets
import shutil
import tempfile
from datetime import date, datetime, timedelta, timezone
from pathlib import Path
from typing import Any
from urllib.parse import urlparse

from sqlalchemy import select

from ..db import now, session
from . import rules as R
from . import scan as S
from .models import AppReview, AppReviewDomain, AppReviewFile, AppReviewHandoff, AppReviewSettings

log = logging.getLogger("santosh.appreview")

GH_API = "https://api.github.com"
GH_TRANSPORT = None                 # tests: httpx.MockTransport
MAX_REPO_BYTES = 150 * 1024 * 1024
KEEP_UPLOAD_DAYS = 7


class ReviewError(Exception):
    pass


# ---------------------------------------------------------------- basics
def _dir(tenant: str) -> Path:
    safe = re.sub(r"[^A-Za-z0-9_-]", "_", tenant)[:64] or "t"
    p = Path(os.environ.get("APPREVIEW_DIR") or "data/appreview") / safe
    p.mkdir(parents=True, exist_ok=True)
    return p


def _publish(tenant: str, kind: str, payload: dict, conv: str = "") -> None:
    from ..events import bus
    try:
        if not conv:
            from ..net.service import home_conv
            conv = home_conv(tenant)
        bus.publish(conv, kind, payload)
    except Exception:
        log.exception("appreview publish failed")


def _settings_row(s, tenant: str) -> AppReviewSettings:
    st = s.get(AppReviewSettings, tenant)
    if st is None:
        st = AppReviewSettings(tenant_id=tenant)
        s.add(st)
        s.flush()
    return st


def settings(tenant: str) -> dict[str, Any]:
    with session() as s:
        st = _settings_row(s, tenant)
        return {"stack": st.stack, "kids": bool(st.kids), "ios": bool(st.ios), "developer": st.developer,
                "github": {"joined": bool(st.gh_sealed), "label": st.gh_label},
                "play_target": {"new_and_updates": R.PLAY_TARGET["new_and_updates"], "existing_min": R.PLAY_TARGET["existing_min"],
                                "since": R.PLAY_TARGET["since"].isoformat()}, "checked_on": R.CHECKED_ON}


def save_settings(tenant: str, **kw: Any) -> dict[str, Any]:
    with session() as s:
        st = _settings_row(s, tenant)
        if "stack" in kw:
            v = str(kw["stack"] or "")
            if v not in ("", "android", "flutter", "react_native", "web"):
                raise ReviewError("stack: android | flutter | react_native | web")
            st.stack = v
        for k in ("kids", "ios"):
            if k in kw:
                setattr(st, k, bool(kw[k]))
        if "developer" in kw:
            st.developer = str(kw["developer"] or "")[:120]
    return settings(tenant)


def _opts(tenant: str) -> dict[str, Any]:
    st = settings(tenant)
    return {"kids": st["kids"], "ios": st["ios"], "stack": st["stack"]}


# ---------------------------------------------------------------- her files
def upload(tenant: str, name: str, data: bytes) -> dict[str, Any]:
    from ..config import get_settings
    cap = int(get_settings().appcheck_max_mb) * 1024 * 1024
    if not data:
        raise ReviewError("File khaali hai.")
    if len(data) > cap:
        raise ReviewError(f"File {cap // 1024 // 1024} MB se badi hai.")
    if data[:2] != b"PK":
        raise ReviewError("APK / AAB / source code ki ZIP bhejiye (yeh ZIP jaisi file nahi hai).")
    safe = re.sub(r"[^A-Za-z0-9._ -]", "_", Path(name or "app.zip").name)[:120] or "app.zip"
    d = _dir(tenant) / "uploads"
    d.mkdir(exist_ok=True)
    with session() as s:
        row = AppReviewFile(tenant_id=tenant, name=safe, size=len(data), sha256=hashlib.sha256(data).hexdigest())
        s.add(row)
        s.flush()
        p = d / f"{row.id}_{safe}"
        p.write_bytes(data)
        row.path = str(p)
        return {"file": row.id, "name": safe, "size": len(data), "sha256": row.sha256}


def _resolve_file(tenant: str, ref: str) -> tuple[str, Path]:
    ref = str(ref or "").strip()
    if not ref:
        raise ReviewError("Kaunsi file? App screen se APK / AAB / ZIP upload kijiye ya chat mein bhejiye (file_id).")
    with session() as s:
        row = s.get(AppReviewFile, ref)
        if row is not None and row.tenant_id == tenant:
            p = Path(row.path)
            if not p.is_file():
                raise ReviewError("Yeh upload ab server par nahi hai (7 din baad hat jaata hai); dobara bhejiye.")
            return row.name, p
    from ..mailbox import service as ms
    try:
        f = ms._file_row(tenant, ref)
    except ms.MailError as e:
        raise ReviewError(str(e))
    p = Path(f.stored_path)
    if not p.is_file():
        raise ReviewError("Yeh file ab server par nahi hai (24 ghante baad hat jaati hai); dobara bhejiye.")
    return f.name, p


# ---------------------------------------------------------------- risks
def build_risk(f: dict[str, Any], stack: str) -> dict[str, Any]:
    rl = R.RULES[f["code"]]
    detail = f.get("detail") or ""
    explain = rl["explain"].replace("{detail}", detail) if "{detail}" in rl["explain"] else (
        rl["explain"] + (f" ({detail})" if detail else ""))
    return {"key": f["key"], "code": f["code"], "title": rl["title"], "area": rl["area"], "area_label": R.AREAS[rl["area"]],
            "sev": f.get("sev") or rl["sev"], "masvs": rl["masvs"], "explain": explain, "evidence": f.get("evidence") or [],
            "fix": rl["fix"], "snippet": R.stack_snippet(f["code"], stack),
            "sources": [{"id": r, **R.SOURCES[r]} for r in rl.get("refs", []) if r in R.SOURCES], "status": "open"}


def _sort(risks: list[dict[str, Any]]) -> list[dict[str, Any]]:
    order = {s: i for i, s in enumerate(R.SEVERITIES)}
    return sorted(risks, key=lambda r: (r.get("status") == "ignored", order.get(r["sev"], 9), r["area"], r["key"]))


def _counts(risks: list[dict[str, Any]]) -> dict[str, int]:
    c = {s: 0 for s in R.SEVERITIES}
    for r in risks:
        if r.get("status") != "ignored":
            c[r["sev"]] = c.get(r["sev"], 0) + 1
    return c


def _compare(prev: AppReview | None, risks: list[dict[str, Any]]) -> dict[str, Any]:
    if prev is None:
        for r in risks:
            r["change"] = "new"
        return {}
    old = {r["key"]: r for r in prev.risks or []}
    ignored = {k for k, r in old.items() if r.get("status") == "ignored"}
    cur = set()
    for r in risks:
        cur.add(r["key"])
        r["change"] = "still_open" if r["key"] in old else "new"
        if r["key"] in ignored:
            r["status"] = "ignored"
            r["note"] = old[r["key"]].get("note", "")
    fixed = [{"key": k, "title": o["title"], "sev": o["sev"]} for k, o in old.items() if k not in cur and o["sev"] != "info"]
    return {"previous": prev.number, "previous_id": prev.id, "previous_score": prev.score,
            "fixed": fixed, "still_open": sum(1 for r in risks if r["change"] == "still_open" and r["sev"] != "info"),
            "new": sum(1 for r in risks if r["change"] == "new" and r["sev"] != "info")}


def _target(kind: str, facts: dict[str, Any], name: str) -> str:
    if kind in ("apk", "aab") and facts.get("package"):
        return f"app:{facts['package']}"
    if kind == "web":
        return f"web:{facts.get('host', name)}"
    if kind == "repo":
        return f"repo:{name.lower()}"
    if facts.get("package"):
        return f"app:{facts['package']}"
    return f"zip:{re.sub(r'[-_ ]?v?\d[\d.]*', '', Path(name).stem).lower() or name.lower()}"


def _finish(tenant: str, kind: str, name: str, out: dict[str, Any], conv: str) -> dict[str, Any]:
    st = _opts(tenant)
    stack = st.get("stack") or out.get("stack") or "android"
    risks = _sort([build_risk(f, stack) for f in out["findings"]])
    facts = out.get("facts") or {}
    target = _target(kind, facts, name)
    with session() as s:
        prev = s.scalars(select(AppReview).where(AppReview.tenant_id == tenant, AppReview.target == target)
                         .order_by(AppReview.created_at.desc())).first()
        cmp = _compare(prev, risks)
        cfg = _settings_row(s, tenant)
        cfg.seq += 1
        title = facts.get("app_label") or facts.get("package") or facts.get("host") or name
        row = AppReview(tenant_id=tenant, number=f"AR-{cfg.seq:04d}", kind=kind, target=target, title=str(title)[:200], stack=stack,
                        score=R.score(risks), facts=facts, risks=risks, compare=cmp, limits=out.get("limits") or [], conv=conv or "")
        s.add(row)
        s.flush()
        v = view(row)
    _publish(tenant, "appreview_report", card(v), conv)
    return v


def view(r: AppReview, full: bool = True) -> dict[str, Any]:
    risks = r.risks or []
    c = _counts(risks)
    v = {"id": r.id, "number": r.number, "kind": r.kind, "target": r.target, "title": r.title, "stack": r.stack, "score": r.score,
         "counts": c, "compare": r.compare or {}, "created_at": r.created_at.isoformat() if r.created_at else "",
         "files": r.files or []}
    v["say"] = _say(v, risks)
    if full:
        v.update({"facts": r.facts or {}, "risks": risks, "limits": r.limits or [], "checked_on": R.CHECKED_ON})
    return v


def _say(v: dict[str, Any], risks: list[dict[str, Any]]) -> str:
    c = v["counts"]
    serious = c.get("critical", 0) + c.get("high", 0)
    head = f"{v['number']} {v['title']}: score {v['score']}/100, {serious} gambhir, {c.get('medium', 0)} dhyaan wale."
    top = [r["title"] for r in risks if r["sev"] in ("critical", "high") and r.get("status") != "ignored"][:3]
    if top:
        head += " Sabse pehle: " + "; ".join(top) + "."
    cmp = v.get("compare") or {}
    if cmp.get("previous"):
        head += f" Pichhli jaanch {cmp['previous']} se: {len(cmp.get('fixed') or [])} theek hue, {cmp.get('still_open', 0)} abhi baaki, {cmp.get('new', 0)} naye."
    return head


def card(v: dict[str, Any]) -> dict[str, Any]:
    risks = v.get("risks") or []
    top = [{"key": r["key"], "title": r["title"], "sev": r["sev"], "sev_label": R.SEV_LABEL[r["sev"]], "area": r["area_label"],
            "explain": r["explain"][:240], "change": r.get("change", "")} for r in risks if r.get("status") != "ignored"
           and r["sev"] != "info"][:6]
    return {"id": v["id"], "number": v["number"], "title": v["title"], "kind": v["kind"], "score": v["score"], "counts": v["counts"],
            "compare": v.get("compare") or {}, "top": top, "say": v["say"]}


# ---------------------------------------------------------------- reviews
def review_file(tenant: str, ref: str, conv: str = "", today: date | None = None) -> dict[str, Any]:
    name, path = _resolve_file(tenant, ref)
    opts = _opts(tenant)
    kind = S.classify_zip(path, name)
    tmp = None
    try:
        if kind == "apks":
            from ..appcheck import analyze as an
            tmp = Path(tempfile.mkdtemp(prefix="arv_"))
            base, _parts = an.base_apk(path, tmp)
            if base is None:
                raise ReviewError("Bundle ke andar main APK nahi mila.")
            path, kind = base, "apk"
        if kind in ("apk", "aab"):
            out = S.scan_android(path, name, opts, today=today)
        else:
            out = S.scan_source(path, name, opts, today=today, strip_top=_single_top(path))
    except S.ScanError as e:
        raise ReviewError(str(e))
    finally:
        if tmp is not None:
            shutil.rmtree(tmp, ignore_errors=True)
    return _finish(tenant, out["kind"], name, out, conv)


def _single_top(path: Path) -> bool:
    import zipfile
    try:
        with zipfile.ZipFile(path) as z:
            tops = {n.split("/", 1)[0] for n in z.namelist() if n}
            return len(tops) == 1 and all("/" in n for n in z.namelist() if n and not n.endswith("/"))
    except zipfile.BadZipFile:
        return False


# ---- GitHub (read-only)
def _gh(tenant: str) -> str:
    from ..bazar.security import vault
    with session() as s:
        st = _settings_row(s, tenant)
        sealed = st.gh_sealed
    if not sealed:
        return ""
    try:
        return vault().open(sealed)
    except ValueError:
        return ""


def _gh_req(path: str, token: str, stream_to: Path | None = None) -> Any:
    import httpx
    h = {"Accept": "application/vnd.github+json", "X-GitHub-Api-Version": "2022-11-28", "User-Agent": "Hoola-AppReview"}
    if token:
        h["Authorization"] = f"Bearer {token}"
    try:
        with httpx.Client(timeout=60, transport=GH_TRANSPORT, follow_redirects=True) as c:
            if stream_to is None:
                r = c.get(GH_API + path, headers=h)
                if r.status_code >= 400:
                    raise ReviewError(_gh_err(r))
                return r
            with c.stream("GET", GH_API + path, headers=h) as r:
                if r.status_code >= 400:
                    r.read()
                    raise ReviewError(_gh_err(r))
                n = 0
                with stream_to.open("wb") as f:
                    for chunk in r.iter_bytes():
                        n += len(chunk)
                        if n > MAX_REPO_BYTES:
                            raise ReviewError(f"Repo {MAX_REPO_BYTES // 1024 // 1024} MB se bada hai; app folder ki ZIP bhejiye.")
                        f.write(chunk)
                return n
    except httpx.HTTPError as e:
        raise ReviewError(f"GitHub tak nahi pahunch paya: {e.__class__.__name__}")


def _gh_err(r) -> str:
    try:
        msg = r.json().get("message", "")
    except ValueError:
        msg = r.text[:150]
    if r.status_code in (401, 403, 404):
        msg += " (private repo ho to App Review > Niyam mein read-only GitHub token jodiye: Contents = Read-only)"
    return f"GitHub: {r.status_code} {msg}"


def parse_repo(repo: str) -> tuple[str, str]:
    t = (repo or "").strip().rstrip("/")
    t = re.sub(r"^(?:https?://)?(?:www\.)?github\.com/", "", t)
    t = re.sub(r"\.git$", "", t)
    m = re.match(r"^([A-Za-z0-9-]{1,39})/([A-Za-z0-9._-]{1,100})", t)
    if not m:
        raise ReviewError("Repo 'owner/naam' ya github.com link ki tarah dijiye.")
    return m.group(1), m.group(2)


def review_repo(tenant: str, repo: str, ref: str = "", conv: str = "", today: date | None = None) -> dict[str, Any]:
    owner, name = parse_repo(repo)
    token = _gh(tenant)
    info = _gh_req(f"/repos/{owner}/{name}", token).json()
    branch = re.sub(r"[^A-Za-z0-9._/-]", "", ref or info.get("default_branch") or "main")[:100]
    tmp = Path(tempfile.mkdtemp(prefix="arv_repo_"))
    try:
        z = tmp / "repo.zip"
        _gh_req(f"/repos/{owner}/{name}/zipball/{branch}", token, stream_to=z)
        try:
            out = S.scan_source(z, f"{owner}/{name}", _opts(tenant), today=today, strip_top=True)
        except S.ScanError as e:
            raise ReviewError(str(e))
        out["kind"] = "repo"
        out["facts"].update({"repo": f"{owner}/{name}", "branch": branch, "private": bool(info.get("private"))})
        return _finish(tenant, "repo", f"{owner}/{name}", out, conv)
    finally:
        shutil.rmtree(tmp, ignore_errors=True)


def set_github(tenant: str, token: str) -> dict[str, Any]:
    token = (token or "").strip()
    if not re.match(r"^(?:github_pat_[A-Za-z0-9_]{20,}|gh[pousr]_[A-Za-z0-9]{30,})$", token):
        raise ReviewError("Yeh GitHub token jaisa nahi lagta (github_pat_... chahiye).")
    r = _gh_req("/user", token)
    scopes = r.headers.get("x-oauth-scopes", "")
    if scopes and any(sc.strip() in ("repo", "public_repo", "admin:org", "workflow", "delete_repo", "write:packages")
                      for sc in scopes.split(",")):
        raise ReviewError("Is token se repo mein likha ja sakta hai. Review ke liye fine-grained token banaiye: sirf "
                          "'Contents: Read-only' (aur 'Metadata: Read-only').")
    login = (r.json() or {}).get("login", "")
    from ..bazar.security import vault
    with session() as s:
        st = _settings_row(s, tenant)
        st.gh_sealed = vault().seal(token)
        st.gh_label = f"@{login}"[:80] if login else "joda gaya"
    return {"github": {"joined": True, "label": f"@{login}" if login else "joda gaya"},
            "say": "GitHub (sirf padhne wala) token jud gaya; Hoola use kabhi chat mein nahi dikhayega."}


def clear_github(tenant: str) -> dict[str, Any]:
    with session() as s:
        st = _settings_row(s, tenant)
        st.gh_sealed, st.gh_label = "", ""
    return {"github": {"joined": False, "label": ""}, "say": "GitHub token hata diya. GitHub par bhi use revoke kar dijiye."}


# ---- website
def norm_host(url: str) -> str:
    t = (url or "").strip()
    if not re.match(r"^https?://", t, re.I):
        t = "https://" + t
    h = (urlparse(t).hostname or "").lower().rstrip(".")
    if not re.match(r"^(?=.{4,200}$)([a-z0-9-]+\.)+[a-z]{2,}$", h):
        raise ReviewError("Website ka sahi address dijiye (jaise aapki-dukaan.in).")
    return h


def _verified(tenant: str, host: str) -> bool:
    with session() as s:
        rows = s.scalars(select(AppReviewDomain).where(AppReviewDomain.tenant_id == tenant, AppReviewDomain.verified_at.is_not(None))).all()
        return any(host == r.host or host.endswith("." + r.host) for r in rows)


def review_web(tenant: str, url: str, conv: str = "", fetch=None, certs=None, pause: float = 0.3) -> dict[str, Any]:
    host = norm_host(url)
    try:
        out = S.scan_web(host, _verified(tenant, host), fetch=fetch, certs=certs, pause=pause)
    except Exception as e:                     # web._check_url refusal etc.
        raise ReviewError(f"Website jaanch nahi ho paayi: {str(e)[:160]}")
    return _finish(tenant, "web", host, out, conv)


def domain_add(tenant: str, url: str) -> dict[str, Any]:
    host = norm_host(url)
    with session() as s:
        row = s.scalars(select(AppReviewDomain).where(AppReviewDomain.tenant_id == tenant, AppReviewDomain.host == host)).first()
        if row is None:
            row = AppReviewDomain(tenant_id=tenant, host=host, token=secrets.token_hex(12))
            s.add(row)
            s.flush()
        return _domain_view(row)


def _domain_view(r: AppReviewDomain) -> dict[str, Any]:
    val = f"hoola-verify={r.token}"
    return {"id": r.id, "host": r.host, "verified": r.verified_at is not None, "method": r.method,
            "verified_at": r.verified_at.isoformat() if r.verified_at else "",
            "how": {"dns": {"type": "TXT", "name": f"_hoola.{r.host}", "value": val},
                    "file": {"url": f"https://{r.host}/.well-known/hoola-verify.txt", "content": val}},
            "say": ("Verify ho gaya." if r.verified_at else
                    f"Domain aapka hai yeh dikhane ke liye: DNS mein TXT record _hoola.{r.host} = {val} daaliye, ya "
                    f"https://{r.host}/.well-known/hoola-verify.txt file mein yahi likhiye. Phir 'Verify karo'.")}


def domain_verify(tenant: str, url: str, fetch=None) -> dict[str, Any]:
    host = norm_host(url)
    fetch = fetch or S.http_fetch
    with session() as s:
        row = s.scalars(select(AppReviewDomain).where(AppReviewDomain.tenant_id == tenant, AppReviewDomain.host == host)).first()
        if row is None:
            raise ReviewError("Pehle domain jodiye (domain_add).")
        val = f"hoola-verify={row.token}"
        method = ""
        try:
            d = fetch(f"https://cloudflare-dns.com/dns-query?name=_hoola.{host}&type=TXT", headers={"Accept": "application/dns-json"},
                      limit=65536)
            ans = json.loads((d.get("body") or b"{}").decode("utf-8", "ignore") or "{}").get("Answer") or []
            if any(val in str(a.get("data", "")) for a in ans):
                method = "dns"
        except Exception:
            pass
        if not method:
            try:
                f = fetch(f"https://{host}/.well-known/hoola-verify.txt", follow=False, limit=4096)
                if f.get("status") == 200 and val.encode() in (f.get("body") or b""):
                    method = "file"
            except Exception:
                pass
        if not method:
            v = _domain_view(row)
            v["say"] = "Abhi TXT record / file nahi mili (DNS badalne mein kuch der lag sakti hai). " + v["say"]
            return v
        row.verified_at, row.method = now(), method
        return _domain_view(row)


def domains(tenant: str) -> list[dict[str, Any]]:
    with session() as s:
        return [_domain_view(r) for r in s.scalars(select(AppReviewDomain).where(AppReviewDomain.tenant_id == tenant)).all()]


# ---------------------------------------------------------------- reading / marking
def _row(s, tenant: str, ref: str) -> AppReview:
    ref = str(ref or "").strip()
    r = s.get(AppReview, ref) if ref.startswith("arv") else None
    if r is None and ref:
        r = s.scalars(select(AppReview).where(AppReview.tenant_id == tenant, AppReview.number == ref.upper())).first()
    if r is None and not ref:
        r = s.scalars(select(AppReview).where(AppReview.tenant_id == tenant).order_by(AppReview.created_at.desc())).first()
    if r is None or r.tenant_id != tenant:
        raise ReviewError("Yeh review nahi mila.")
    return r


def get(tenant: str, ref: str = "") -> dict[str, Any]:
    with session() as s:
        return view(_row(s, tenant, ref))


def reviews(tenant: str, limit: int = 30) -> list[dict[str, Any]]:
    with session() as s:
        rows = s.scalars(select(AppReview).where(AppReview.tenant_id == tenant).order_by(AppReview.created_at.desc()).limit(limit)).all()
        return [view(r, full=False) for r in rows]


def mark(tenant: str, ref: str, key: str, status: str, note: str = "") -> dict[str, Any]:
    """Her decision on one risk: ignored (it's intentional / accepted) | open. The next review keeps it."""
    if status not in ("ignored", "open"):
        raise ReviewError("status: ignored | open")
    with session() as s:
        r = _row(s, tenant, ref)
        risks = [dict(x) for x in r.risks or []]
        hit = next((x for x in risks if x["key"] == key), None)
        if hit is None:
            raise ReviewError("Yeh risk is review mein nahi hai.")
        if status == "ignored" and hit["sev"] == "critical" and not note.strip():
            raise ReviewError("Bahut gambhir risk ko chhodne ke liye wajah likhiye.")
        hit["status"], hit["note"] = status, note[:300]
        r.risks = _sort(risks)
        r.score = R.score(r.risks)
        return view(r)


def summary(tenant: str) -> dict[str, Any]:
    rs = reviews(tenant, 50)
    latest: dict[str, dict] = {}
    for r in rs:
        latest.setdefault(r["target"], r)
    return {"settings": settings(tenant), "targets": list(latest.values()), "reviews": rs[:20], "domains": domains(tenant)}


# ---------------------------------------------------------------- fix list + report files
def fix_list(v: dict[str, Any]) -> list[dict[str, Any]]:
    out = []
    for i, r in enumerate([x for x in v["risks"] if x.get("status") != "ignored" and x["sev"] != "info"], 1):
        out.append({"n": i, "key": r["key"], "title": r["title"], "severity": r["sev"], "area": r["area_label"], "masvs": r["masvs"],
                    "why": r["explain"], "where": [f"{e['where']}" + (f":{e['line']}" if e.get("line") else "") for e in r["evidence"]],
                    "do": r["fix"], "snippet": (r.get("snippet") or {}).get("code", ""),
                    "done_when": f"Agli App Review mein '{r['title']}' na dikhe."})
    return out


def fix_list_md(v: dict[str, Any], developer: str = "") -> str:
    lines = [f"# Fix list: {v['title']} ({v['number']})", "",
             f"Score {v['score']}/100 · {v['created_at'][:10]} · Hoola App Review (static review, policies checked {R.CHECKED_ON})"]
    if developer:
        lines.append(f"For: {developer}")
    lines.append("")
    for t in fix_list(v):
        lines += [f"## {t['n']}. [{t['severity'].upper()}] {t['title']}", f"- Kyon: {t['why']}"]
        if t["where"]:
            lines.append("- Kahan: " + ", ".join(t["where"]))
        lines.append(f"- Kya karein: {t['do']}")
        if t["snippet"]:
            lines += ["", "```", t["snippet"], "```"]
        lines += [f"- Ho gaya kab maanein: {t['done_when']}", ""]
    lines += ["Secrets: koi bhi key mili ho to use badalna (rotate) zaroori hai, sirf code se hatana kaafi nahi."]
    return "\n".join(lines)


def _latin(t: str) -> str:
    return (t or "").replace("…", "...").replace("→", "->").replace("–", "-").replace("—", "-").replace("₹", "Rs ").encode(
        "latin-1", "replace").decode("latin-1")


def _pdf(v: dict[str, Any]) -> bytes:
    from xml.sax.saxutils import escape

    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import Paragraph, Preformatted, SimpleDocTemplate, Spacer, Table

    ss = getSampleStyleSheet()
    bio = io.BytesIO()
    doc = SimpleDocTemplate(bio, pagesize=A4, leftMargin=16 * mm, rightMargin=16 * mm, topMargin=14 * mm, bottomMargin=14 * mm,
                            title=_latin(f"App Review {v['number']}"))

    def P(t: str, st: str = "BodyText") -> Paragraph:
        return Paragraph(escape(_latin(t)), ss[st])
    c = v["counts"]
    el = [P(f"App Review: {v['title']}", "Title"),
          P(f"{v['number']} · {v['kind'].upper()} · {v['created_at'][:10]} · Score {v['score']}/100"),
          Table([["Bahut gambhir", "Gambhir", "Dhyaan", "Chhota", "Jaankari"],
                 [c.get("critical", 0), c.get("high", 0), c.get("medium", 0), c.get("low", 0), c.get("info", 0)]],
                style=[("GRID", (0, 0), (-1, -1), 0.4, colors.grey), ("BACKGROUND", (0, 0), (-1, 0), colors.whitesmoke)]),
          Spacer(1, 6), P(v["say"])]
    cmp = v.get("compare") or {}
    if cmp.get("previous"):
        el.append(P(f"Pichhli jaanch {cmp['previous']} (score {cmp.get('previous_score')}): theek hue {len(cmp.get('fixed') or [])}, "
                    f"baaki {cmp.get('still_open', 0)}, naye {cmp.get('new', 0)}."))
        for f in (cmp.get("fixed") or [])[:15]:
            el.append(P(f"  ✓ theek: {f['title']}".replace("✓", "+")))
    facts = v.get("facts") or {}
    keep = [(k, facts.get(k)) for k in ("package", "version_name", "target_sdk", "min_sdk", "built_with", "signed_by", "repo", "host")
            if facts.get(k) not in (None, "", [])]
    if keep:
        el += [Spacer(1, 6), Table([[k, _latin(str(x))[:90]] for k, x in keep], style=[("GRID", (0, 0), (-1, -1), 0.3, colors.lightgrey)])]
    for r in v["risks"]:
        if r.get("status") == "ignored":
            continue
        el += [Spacer(1, 8), P(f"[{R.SEV_LABEL[r['sev']]}] {r['title']}  ({r['area_label']} · {r['masvs']})", "Heading4"),
               P(r["explain"])]
        if r["evidence"]:
            el.append(P("Kahan: " + "; ".join(e["where"] + (f":{e['line']}" if e.get("line") else "") +
                                              (f" ({e['text']})" if e.get("text") else "") for e in r["evidence"][:4])))
        el.append(P("Kya karein: " + r["fix"]))
        if r.get("snippet"):
            el.append(Preformatted(_latin(r["snippet"]["code"])[:1500], ss["Code"]))
        if r.get("sources"):
            el.append(P("Source: " + "; ".join(f"{s['title']} ({s['url']}, {s['date']})" for s in r["sources"][:2])))
    el += [Spacer(1, 10)] + [P("Seema: " + x) for x in v.get("limits") or []]
    el.append(P("Yeh static review hai, legal salah nahi. Store policies badalti rehti hain; launch se pehle Play Console / App Store "
                "Connect ke warnings bhi dekhiye."))
    doc.build(el)
    return bio.getvalue()


def _xlsx(v: dict[str, Any]) -> bytes:
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Risks"
    ws.append(["#", "Severity", "Area", "MASVS", "Title", "Explain", "Where", "Fix", "Change", "Status", "Source"])
    for i, r in enumerate(v["risks"], 1):
        ws.append([i, r["sev"], r["area_label"], r["masvs"], r["title"], r["explain"],
                   "; ".join(e["where"] + (f":{e['line']}" if e.get("line") else "") for e in r["evidence"]), r["fix"],
                   r.get("change", ""), r.get("status", ""), "; ".join(s["url"] for s in r.get("sources") or [])])
    fl = wb.create_sheet("Fix list")
    fl.append(["#", "Severity", "Title", "Where", "Do", "Snippet", "Done when"])
    for t in fix_list(v):
        fl.append([t["n"], t["severity"], t["title"], ", ".join(t["where"]), t["do"], t["snippet"], t["done_when"]])
    ds = wb.create_sheet("Data safety draft")
    ds.append(["Data type (Play Data safety)", "Kahan se mila", "Collected?", "Shared?", "Kyon (purpose)"])
    facts = v.get("facts") or {}
    by_sdk = {}
    for sd in facts.get("sdks") or []:
        for d in (R.SDK_BY_NAME.get(sd["name"]) or {}).get("data", []):
            by_sdk.setdefault(d, []).append(sd["name"])
    for d in facts.get("data_safety") or []:
        src = by_sdk.get(d, []) + (["permission"] if d in (facts.get("data_from_permissions") or []) else [])
        ds.append([d, ", ".join(src), "Haan (jaanchiye)", "SDK ho to shayad", ""])
    cm = wb.create_sheet("Compare")
    cmp = v.get("compare") or {}
    cm.append(["Pichhli jaanch", cmp.get("previous", "-")])
    cm.append(["Theek hue", len(cmp.get("fixed") or [])])
    for f in cmp.get("fixed") or []:
        cm.append(["  theek", f["title"], f["sev"]])
    cm.append(["Abhi baaki", cmp.get("still_open", 0)])
    cm.append(["Naye", cmp.get("new", 0)])
    so = wb.create_sheet("Sources")
    so.append(["Id", "Title", "URL", "Checked"])
    for k, s_ in R.SOURCES.items():
        so.append([k, s_["title"], s_["url"], s_["date"]])
    bio = io.BytesIO()
    wb.save(bio)
    return bio.getvalue()


def report_files(tenant: str, ref: str = "", conv: str = "") -> dict[str, Any]:
    from .. import chatfiles
    v = get(tenant, ref)
    dev = settings(tenant)["developer"]
    slug = re.sub(r"[^A-Za-z0-9]+", "_", v["title"])[:40].strip("_") or "app"
    made = []
    for name, data, ct in (
            (f"AppReview_{v['number']}_{slug}.pdf", _pdf(v), "application/pdf"),
            (f"AppReview_{v['number']}_{slug}.xlsx", _xlsx(v), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"),
            (f"FixList_{v['number']}_{slug}.md", fix_list_md(v, dev).encode(), "text/markdown"),
            (f"FixList_{v['number']}_{slug}.json", json.dumps({"review": v["number"], "title": v["title"], "tasks": fix_list(v)},
                                                               ensure_ascii=False, indent=1).encode(), "application/json")):
        f = chatfiles.store(tenant, name, data, ct)
        made.append({"file_id": f.get("file_id") or f.get("id"), "name": name})
        _publish(tenant, "pc_file", {**f, "source": "appreview", "title": name}, conv)
    with session() as s:
        r = _row(s, tenant, v["id"])
        r.files = made
    return {"files": made, "say": "Report (PDF + Excel) aur developer ke liye fix list taiyaar hai."}


# ---------------------------------------------------------------- hand-off to the Developer Sahayak (her tap)
def _handoff_hash(review_id: str, to: str, tasks: list[dict[str, Any]]) -> str:
    return hashlib.sha256(json.dumps([review_id, to, [(t["key"], t["do"]) for t in tasks]], sort_keys=True).encode()).hexdigest()


def handoff_request(tenant: str, ref: str = "", keys: list[str] | None = None, to: str = "devsahayak", conv: str = "") -> dict[str, Any]:
    if to not in ("devsahayak", "developer"):
        raise ReviewError("to: devsahayak | developer")
    v = get(tenant, ref)
    tasks = fix_list(v)
    if keys:
        tasks = [t for t in tasks if t["key"] in set(keys)]
    if not tasks:
        raise ReviewError("Fix list khaali hai (koi khula risk nahi).")
    h = _handoff_hash(v["id"], to, tasks)
    with session() as s:
        old = s.scalars(select(AppReviewHandoff).where(AppReviewHandoff.tenant_id == tenant, AppReviewHandoff.review_id == v["id"],
                                                       AppReviewHandoff.status == "card", AppReviewHandoff.hash == h)).first()
        row = old or AppReviewHandoff(tenant_id=tenant, review_id=v["id"], to=to, keys=[t["key"] for t in tasks], hash=h)
        if old is None:
            s.add(row)
            s.flush()
        hid = row.id
    payload = {"id": hid, "review": v["number"], "title": v["title"], "to": to, "hash": h,
               "tasks": [{"n": t["n"], "title": t["title"], "severity": t["severity"]} for t in tasks][:20], "count": len(tasks),
               "say": (f"{len(tasks)} kaam Developer Sahayak ko dene ke liye aapki haan chahiye (woh aapke code mein badlav "
                       f"taiyaar karega)." if to == "devsahayak" else
                       f"{len(tasks)} kaam ki fix list developer ke liye; haan par message taiyaar hoga, bhejna aap karengi.")}
    _publish(tenant, "appreview_handoff", payload, conv or v.get("conv", ""))
    return {**payload, "status": "needs_tap", "say": payload["say"] + " Card bheja hai."}


def handoff_decide(tenant: str, hid: str, yes: bool, hash_: str, *, by_tap: bool = False) -> dict[str, Any]:
    if not by_tap:
        raise ReviewError("Yeh sirf aapke app ke card par tap se hota hai.")
    with session() as s:
        row = s.get(AppReviewHandoff, hid)
        if row is None or row.tenant_id != tenant:
            raise ReviewError("Yeh card nahi mila.")
        if row.status != "card":
            raise ReviewError("Is card par faisla ho chuka hai.")
        if hash_ != row.hash:
            raise ReviewError("Card purana hai (fix list badal gayi); naya card banaiye.")
        rid, to, keys = row.review_id, row.to, list(row.keys or [])
        if not yes:
            row.status, row.decided_at = "rejected", now()
            return {"status": "rejected", "say": "Theek hai, fix list nahi bheji."}
    v = get(tenant, rid)
    tasks = [t for t in fix_list(v) if t["key"] in set(keys)]
    if _handoff_hash(rid, to, tasks) != hash_:
        raise ReviewError("Review badal gaya hai; naya card banaiye.")
    result: dict[str, Any]
    status = "sent"
    if to == "devsahayak":
        try:
            from ..devsahayak import service as dev          # neighbour module (another builder); optional
            fn = getattr(dev, "tasks_from_review", None)
            if fn is None:
                raise ImportError
            result = {"devsahayak": fn(tenant, {"review": v["number"], "title": v["title"], "target": v["target"], "tasks": tasks})}
        except ImportError:
            status = "failed"
            result = {"error": "Developer Sahayak abhi juda nahi hai; fix list file developer ko bhejiye.",
                      "fix_list": fix_list_md(v, settings(tenant)["developer"])}
    else:
        text = (f"Namaste, hamari app '{v['title']}' ki review ({v['number']}) mein {len(tasks)} kaam nikle. Sabse pehle: "
                + "; ".join(t["title"] for t in tasks[:3]) + ". Poori fix list file mein bhej rahi hoon.")
        from urllib.parse import quote
        result = {"message": text, "wa": f"https://wa.me/?text={quote(text)}", "fix_list": fix_list_md(v, settings(tenant)["developer"])}
    with session() as s:
        row = s.get(AppReviewHandoff, hid)
        row.status, row.result, row.decided_at = status, result, now()
    return {"status": status, **result, "say": "Developer Sahayak ko kaam de diye." if status == "sent" and to == "devsahayak"
            else result.get("error") or "Message taiyaar hai; aap bhejiye."}


# ---------------------------------------------------------------- worker
def tick(when: datetime | None = None) -> dict[str, Any]:
    """Once a day: remove uploads older than 7 days; one alert per app whose last review targets below Play's current API
    requirement (at most every 30 days per app)."""
    when = when or datetime.now(timezone.utc)
    day = when.date().isoformat()
    alerts = cleaned = 0
    with session() as s:
        cut = when - timedelta(days=KEEP_UPLOAD_DAYS)
        for f in s.scalars(select(AppReviewFile)).all():
            ca = f.created_at if f.created_at.tzinfo else f.created_at.replace(tzinfo=timezone.utc)
            if ca < cut:
                try:
                    Path(f.path).unlink(missing_ok=True)
                except OSError:
                    pass
                s.delete(f)
                cleaned += 1
        tenants = [st for st in s.scalars(select(AppReviewSettings)).all() if st.last_tick != day]
        for st in tenants:
            st.last_tick = day
            rows = s.scalars(select(AppReview).where(AppReview.tenant_id == st.tenant_id, AppReview.kind.in_(("apk", "aab", "source", "repo")))
                             .order_by(AppReview.created_at.desc())).all()
            seen = set()
            for r in rows:
                if r.target in seen:
                    continue
                seen.add(r.target)
                tg = (r.facts or {}).get("target_sdk")
                if tg is None or tg >= R.PLAY_TARGET["new_and_updates"] or when.date() < R.PLAY_TARGET["since"]:
                    continue
                told = (r.compare or {}).get("_alerted", "")
                if told and (when.date() - date.fromisoformat(told)).days < 30:
                    continue
                cmp = dict(r.compare or {})
                cmp["_alerted"] = day
                r.compare = cmp
                alerts += 1
                _publish(st.tenant_id, "appreview_alert", {
                    "review": r.number, "id": r.id, "title": r.title, "target_sdk": tg,
                    "say": f"{r.title}: targetSdk {tg} hai. Google Play par naya update tabhi jayega jab app API "
                           f"{R.PLAY_TARGET['new_and_updates']} target kare. Developer ko bataiye, phir nayi file se dobara review."},
                    r.conv)
    return {"alerts": alerts, "cleaned": cleaned}
