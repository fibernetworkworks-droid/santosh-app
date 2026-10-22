"""v0.102 Hoola Deals, phase 1: one avatar buys from another avatar with typed messages and one shared state machine.

The owner's example: the buyer tells her Hoola "make my list and send it to Sharma Traders" -> the buyer's Hoola sends
a `deal_ask` (items, quantities) -> the wholesaler's Hoola checks its OWN stock list in code and builds a quote (price,
GST, total, valid until) -> the quote goes out by itself only when the wholesaler's rules allow it, else the wholesaler
taps "Bhejo" -> the buyer's Hoola re-checks every number of the quote in code -> the buyer TAPS Haan (a card; never the
model) -> `deal_confirm` -> the wholesaler's Hoola checks the stock again and accepts by itself only inside the
wholesaler's signed rule, else the wholesaler taps Accept -> order number + pick list for packing -> `deal_ready` with a
pickup code -> at the counter the code is checked -> handed over.

Why typed messages and code checks instead of free chat: AI negotiators take the first offer and can be talked into
things by the other side's words (Microsoft Magentic Marketplace 2025, NIST agent red-teaming 2026). Here the model
only fills fields and words; prices, totals, stock, expiry, versions and state changes are decided by code, and the two
commitments (the buyer's confirm, the seller's accept) are the owners' own taps or the seller's written rule.

Messages (EXTRA_KINDS, need an active connection; a public business counter makes one with business.contact):
  deal_ask      buyer -> seller   {items:[{item, qty, unit, note}], mode pickup|delivery, area, need_by, note, title}
  deal_quote    seller -> buyer   {v, lines:[...], subtotal, gst, total, valid_until, note, place}
  deal_confirm  buyer -> seller   {v, quote_hash, mode}
  deal_accept   seller -> buyer   {number, ready_by, note}
  deal_reject   seller -> buyer   {reason}
  deal_status   seller -> buyer   {status packing|delayed|handed_over, note}
  deal_ready    seller -> buyer   {number, code, place, until}
  deal_cancel   either side      {reason}
Every body carries `step` (1, 2, 3...) and `prev` (hash of the step before), so a message out of order, replayed or
rewritten is refused and both sides hold the same history. Each side keeps its own `hoola_deals` row (shared `dkey`).
"""
from __future__ import annotations

import difflib
import hashlib
import json
import re
import secrets
from datetime import datetime, timedelta, timezone
from typing import Any

from sqlalchemy import JSON, Boolean, DateTime, Float, Integer, String, Text, UniqueConstraint, select
from sqlalchemy.orm import Mapped, mapped_column

from ..db import Base, new_id, now, session
from . import identity as idn
from . import policy
from .service import EXTRA_KINDS, NetError, _link, _post, _thread, resolve, tell

DEAL_KINDS = ("deal_ask", "deal_quote", "deal_confirm", "deal_accept", "deal_reject", "deal_status", "deal_ready",
              "deal_cancel", "deal_catalog", "deal_pay", "deal_paid", "deal_receipt", "deal_refund", "deal_cancel_req",
              "deal_cancel_no", "deal_ok", "deal_dispute", "deal_resolve", "deal_dispute_close", "deal_escalate")
OPEN_STATES = ("asked", "quoted", "confirmed", "accepted", "packing", "ready", "handed_over", "disputed")
END_STATES = ("closed", "rejected", "cancelled", "expired")
STATES = OPEN_STATES + END_STATES
# (state before, message) -> state after. Anything not listed is refused, on both sides.
TRANSITIONS: dict[tuple[str, str], str] = {
    ("", "deal_ask"): "asked",
    ("asked", "deal_quote"): "quoted",
    ("quoted", "deal_quote"): "quoted",               # a revised quote (new version) replaces the old one
    ("quoted", "deal_confirm"): "confirmed",
    ("confirmed", "deal_accept"): "accepted",
    ("asked", "deal_reject"): "rejected",
    ("quoted", "deal_reject"): "rejected",
    ("confirmed", "deal_reject"): "rejected",
    ("accepted", "deal_status:packing"): "packing",
    ("accepted", "deal_status:delayed"): "accepted",
    ("packing", "deal_status:delayed"): "packing",
    ("accepted", "deal_ready"): "ready",
    ("packing", "deal_ready"): "ready",
    ("ready", "deal_ready"): "ready",                 # v0.105 a new pickup code / more time (old code stops working)
    ("ready", "deal_status:handed_over"): "handed_over",
}
for _s in ("asked", "quoted", "confirmed", "accepted", "packing", "ready"):
    TRANSITIONS[(_s, "deal_cancel")] = "cancelled"
# v0.106 money, cancel requests, the buyer's check after handover, disputes (payment steps never change the state)
for _s in ("accepted", "packing", "ready", "handed_over", "disputed", "closed"):
    for _k in ("deal_pay", "deal_paid", "deal_receipt"):
        TRANSITIONS[(_s, _k)] = _s
for _s in ("cancelled", "disputed", "closed", "handed_over"):
    TRANSITIONS[(_s, "deal_refund")] = _s
for _s in ("accepted", "packing"):
    TRANSITIONS[(_s, "deal_cancel_req")] = _s
    TRANSITIONS[(_s, "deal_cancel_no")] = _s
TRANSITIONS.update({
    ("handed_over", "deal_ok"): "closed",               # the buyer: everything came right
    ("handed_over", "deal_dispute"): "disputed",        # the buyer: short / damaged / wrong, inside the seller's window
    ("disputed", "deal_resolve"): "disputed",           # the seller's answer (refund / credit / replace / reject)
    ("disputed", "deal_dispute_close"): "closed",       # the buyer accepts the answer
    ("disputed", "deal_escalate"): "disputed",          # the buyer doesn't: official routes are named, the record stays
})
FROM_SELLER = {"deal_quote", "deal_accept", "deal_reject", "deal_status", "deal_ready", "deal_pay", "deal_receipt",
               "deal_refund", "deal_resolve", "deal_cancel_no"}
FROM_BUYER = {"deal_ask", "deal_confirm", "deal_paid", "deal_ok", "deal_dispute", "deal_dispute_close", "deal_escalate",
              "deal_cancel_req"}
GST_RATES = (0.0, 0.1, 0.25, 1.5, 3.0, 5.0, 12.0, 18.0, 28.0, 40.0)
MAX_LINES = 100
ASKS_PER_DAY = 30                  # lists one peer may send one seller in a day
OPEN_PER_PEER = 10                 # unanswered lists from one peer at a time
MAX_TOTAL = 1_00_00_000.0          # Rs 1 crore per deal (phase 1 ceiling)
STATE_WORDS = {"asked": "list bheji gayi", "quoted": "quote aaya", "confirmed": "confirm hua", "accepted": "order pakka",
               "packing": "packing chal rahi", "ready": "parcel taiyaar", "handed_over": "saamaan de diya",
               "disputed": "shikayat khuli hai", "closed": "poora hua",
               "rejected": "mana ho gaya", "cancelled": "cancel hua", "expired": "quote ka samay khatam"}
STATS = {"asks": 0, "quotes": 0, "auto_quotes": 0, "confirms": 0, "accepts": 0, "auto_accepts": 0, "ready": 0,
         "handed_over": 0, "refused": 0}


class DealError(NetError):
    pass


# ---------------------------------------------------------------- tables
class Deal(Base):
    """One side's record of a deal (buyer and seller each keep one; `dkey` is shared)."""
    __tablename__ = "hoola_deals"
    __table_args__ = (UniqueConstraint("tenant_id", "dkey", name="uq_hoola_deal"),)
    id: Mapped[str] = mapped_column(String(40), primary_key=True, default=lambda: new_id("dl"))
    tenant_id: Mapped[str] = mapped_column(String(64), index=True)
    dkey: Mapped[str] = mapped_column(String(60), index=True)
    role: Mapped[str] = mapped_column(String(8))                                # buyer | seller
    peer: Mapped[str] = mapped_column(String(200), index=True)
    thread_id: Mapped[str] = mapped_column(String(40), default="")
    title: Mapped[str] = mapped_column(String(200), default="")
    state: Mapped[str] = mapped_column(String(14), default="asked")
    number: Mapped[str] = mapped_column(String(30), default="")                # the seller's order number
    ask: Mapped[dict] = mapped_column(JSON, default=dict)                       # what was asked (items, mode, need_by...)
    quote: Mapped[dict] = mapped_column(JSON, default=dict)                     # the quote in force (sent / received)
    draft: Mapped[dict] = mapped_column(JSON, default=dict)                     # seller: quote built, not sent yet
    accept: Mapped[dict] = mapped_column(JSON, default=dict)                    # number, ready_by, note
    ready: Mapped[dict] = mapped_column(JSON, default=dict)                     # code, place, until
    reserved: Mapped[list] = mapped_column(JSON, default=list)                  # seller: [{stock_id, qty}] held for this order
    smap: Mapped[list] = mapped_column(JSON, default=list)                      # seller: stock id per quote line ("" = typed line)
    pick: Mapped[dict] = mapped_column(JSON, default=dict)                      # v0.105 seller: pick task {to, status, by, short, at}
    pay: Mapped[dict] = mapped_column(JSON, default=dict)                       # v0.106 {choice, status, due, paid, payee, requests, claims, receipts, refunds}
    cancel_req: Mapped[dict] = mapped_column(JSON, default=dict)               # v0.106 {by, reason, at, status}
    dispute: Mapped[dict] = mapped_column(JSON, default=dict)                   # v0.106 {reason, lines, note, at, answer, escalated}
    accepted_at: Mapped[datetime | None] = mapped_column(DateTime(timezone=True), nullable=True)
    handed_at: Mapped[datetime | None] = mapped_column(DateTime(timezone=True), nullable=True)
    timers: Mapped[dict] = mapped_column(JSON, default=dict)                    # v0.106 reminders already given (once each)
    step: Mapped[int] = mapped_column(Integer, default=0)
    last_hash: Mapped[str] = mapped_column(String(64), default="")
    history: Mapped[list] = mapped_column(JSON, default=list)                   # [{step, kind, by, at, hash, say}]
    reason: Mapped[str] = mapped_column(Text, default="")
    waiting: Mapped[str] = mapped_column(String(16), default="")               # what this side's owner must tap: quote|confirm|accept|""
    created_at: Mapped[datetime] = mapped_column(DateTime(timezone=True), default=now)
    updated_at: Mapped[datetime] = mapped_column(DateTime(timezone=True), default=now)


class StockItem(Base):
    """The seller's own stock list: what the avatar quotes from (price, GST, quantity in hand)."""
    __tablename__ = "hoola_stock"
    __table_args__ = (UniqueConstraint("tenant_id", "key", name="uq_hoola_stock"),)
    id: Mapped[str] = mapped_column(String(40), primary_key=True, default=lambda: new_id("sk"))
    tenant_id: Mapped[str] = mapped_column(String(64), index=True)
    key: Mapped[str] = mapped_column(String(140))                               # normalised name (unique per seller)
    name: Mapped[str] = mapped_column(String(140))
    aliases: Mapped[list] = mapped_column(JSON, default=list)
    unit: Mapped[str] = mapped_column(String(30), default="")
    price: Mapped[float] = mapped_column(Float, default=0.0)                    # per unit, before GST
    gst_pct: Mapped[float] = mapped_column(Float, default=0.0)
    qty: Mapped[float | None] = mapped_column(Float, nullable=True)            # None = not counted (ask the owner)
    reserved: Mapped[float] = mapped_column(Float, default=0.0)                # held by accepted, not yet handed-over orders
    low_at: Mapped[float | None] = mapped_column(Float, nullable=True)         # v0.105 tell her when free stock falls to this
    tiers: Mapped[list] = mapped_column(JSON, default=list)                     # v0.105 [{min, price}] bulk prices
    subs: Mapped[list] = mapped_column(JSON, default=list)                      # v0.105 names that can stand in when this runs out
    public: Mapped[bool] = mapped_column(Boolean, default=True)                 # v0.105 shown in the price list others see
    updated_at: Mapped[datetime] = mapped_column(DateTime(timezone=True), default=now)


class DealRules(Base):
    """The seller's counter settings. What the avatar may do ALONE (auto quote / accept, catalog, packers) lives in her
    Mere Niyam (`avatar.rules["selling"]`, v0.105), read back to her with the other rules and set from her own words."""
    __tablename__ = "hoola_deal_rules"
    tenant_id: Mapped[str] = mapped_column(String(64), primary_key=True)
    place: Mapped[str] = mapped_column(String(200), default="")                # pickup counter / address shown on ready
    pickup_minutes: Mapped[int] = mapped_column(Integer, default=180)          # how long a ready parcel waits
    prefix: Mapped[str] = mapped_column(String(8), default="D")
    next_no: Mapped[int] = mapped_column(Integer, default=1001)
    upi_id: Mapped[str] = mapped_column(String(80), default="")               # v0.106 her UPI ID for deal payments (set by her)
    payee_name: Mapped[str] = mapped_column(String(80), default="")
    next_receipt: Mapped[int] = mapped_column(Integer, default=1)
    updated_at: Mapped[datetime] = mapped_column(DateTime(timezone=True), default=now)


# ---------------------------------------------------------------- small helpers
def _utc(d: datetime | None) -> datetime | None:
    if d is None:
        return None
    return d if d.tzinfo else d.replace(tzinfo=timezone.utc)


def _when(iso: Any) -> datetime | None:
    try:
        d = datetime.fromisoformat(str(iso))
    except (TypeError, ValueError):
        return None
    return d if d.tzinfo else d.replace(tzinfo=timezone.utc)


def _num(x: Any, lo: float = 0.0, hi: float = 1e9) -> float | None:
    try:
        v = float(x)
    except (TypeError, ValueError):
        return None
    if v != v or v < lo or v > hi:                 # NaN / out of range
        return None
    return v


def _qty(x: Any) -> str:
    try:
        v = float(x)
    except (TypeError, ValueError):
        return str(x)
    return str(int(v)) if v == int(v) else f"{v:.2f}".rstrip("0").rstrip(".")


def _r2(x: float) -> float:
    return round(float(x) + 0.0, 2)


def inr(x: Any) -> str:
    try:
        from ..india.money import inr as _inr
        return _inr(float(x or 0))
    except Exception:
        return f"Rs {float(x or 0):,.2f}"


def norm(name: Any) -> str:
    t = re.sub(r"[^\w\s]", " ", str(name or "").lower(), flags=re.UNICODE)
    return re.sub(r"\s+", " ", t).strip()[:140]


def canon(body: dict[str, Any]) -> str:
    """The exact text a step's hash is made from (both sides compute the same)."""
    b = {k: v for k, v in body.items() if k not in ("step", "prev")}
    return json.dumps(b, sort_keys=True, ensure_ascii=False, separators=(",", ":"))


def step_hash(prev: str, kind: str, step: int, body: dict[str, Any]) -> str:
    return hashlib.sha256(f"{prev}|{kind}|{step}|{canon(body)}".encode()).hexdigest()


def quote_hash(q: dict[str, Any]) -> str:
    keep = {k: q.get(k) for k in ("v", "lines", "subtotal", "gst", "total", "valid_until", "terms")}
    return hashlib.sha256(json.dumps(keep, sort_keys=True, ensure_ascii=False, separators=(",", ":")).encode()).hexdigest()


def _event_key(kind: str, body: dict[str, Any]) -> str:
    return f"deal_status:{body.get('status')}" if kind == "deal_status" else kind


def _say_who(peer: str) -> str:
    return idn.label(idn.peer_card(peer)) if peer else "Hoola"


# ---------------------------------------------------------------- rules + stock (seller side)
def rules(tenant: str) -> DealRules:
    with session() as s:
        r = s.get(DealRules, tenant)
        if r is None:
            r = DealRules(tenant_id=tenant)
            s.add(r)
            s.flush()
        s.expunge(r)
        return r


def selling(tenant: str) -> dict[str, Any]:
    """Her Mere Niyam for selling (avatar rules, cleaned with the fixed limits)."""
    from ..avatar.service import get as avatar_get
    return dict(avatar_get(tenant)["rules"]["selling"])


SELL_KEYS = ("auto_quote", "auto_quote_max", "auto_accept", "auto_accept_max", "trusted", "valid_minutes", "catalog",
             "show_stock", "auto_ready", "packers", "pay_before", "counter_ok", "khata", "cancel_minutes", "dispute_hours")
COUNTER_KEYS = ("place", "pickup_minutes", "prefix", "upi_id", "payee_name")


def rules_view(tenant: str) -> dict[str, Any]:
    r, so = rules(tenant), selling(tenant)
    from . import dealpay
    pe = dealpay.payee(tenant)
    return so | {"place": r.place, "pickup_minutes": r.pickup_minutes, "prefix": r.prefix, "say": _rules_say(so),
                 "packers_found": [m["name"] for m in packers(tenant)], "upi_id": r.upi_id or "",
                 "payee_name": r.payee_name or "", "business_verified": bool(pe.get("business_verified"))}


def _rules_say(so: dict[str, Any]) -> str:
    q = {"off": "har quote aapke tap se jaata hai", "trusted": f"bharose wale graahakon ko {inr(so['auto_quote_max'])} tak ka "
         "poora quote khud jaata hai", "connected": f"jude hue sabhi graahakon ko {inr(so['auto_quote_max'])} tak ka poora quote "
         "khud jaata hai"}.get(so["auto_quote"], "")
    a = (f"stock ho to {inr(so['auto_accept_max'])} tak ka order khud pakka" if so["auto_accept"] else "har order aapke tap se pakka")
    p = "packing hote hi parcel khud 'taiyaar'" if so.get("auto_ready", True) else "packing ke baad 'taiyaar' aapke tap se"
    m = ("saamaan se pehle paisa (paid / khata / counter)" if so.get("pay_before", "handover") == "handover" else "paisa baad mein bhi")
    return (f"{q}; {a}; {p}; {m}; {so.get('cancel_minutes', 30)} min tak free cancel; shikayat "
            f"{so.get('dispute_hours', 48)} ghante tak" + ("; khata chalu" if so.get("khata") else "") + ".")


def set_rules(tenant: str, **kw) -> dict[str, Any]:
    """Selling rules go into her Mere Niyam; the counter (place, pickup time, order prefix) stays here."""
    sell = {k: kw[k] for k in SELL_KEYS if k in kw and kw[k] is not None}
    if "auto_quote" in sell and sell["auto_quote"] not in ("off", "trusted", "connected"):
        raise DealError("auto_quote: off, trusted, connected")
    if "catalog" in sell and sell["catalog"] not in ("off", "connected", "public"):
        raise DealError("catalog: off, connected, public")
    for k in ("auto_quote_max", "auto_accept_max"):
        if k in sell and _num(sell[k], 0, MAX_TOTAL) is None:
            raise DealError(f"{k}: 0 se 1 crore ke beech")
    if "pay_before" in sell and sell["pay_before"] not in ("handover", "none"):
        raise DealError("pay_before: handover, none")
    if "cancel_minutes" in sell and _num(sell["cancel_minutes"], 0, 1440) is None:
        raise DealError("cancel_minutes: 0 se 1440")
    if "dispute_hours" in sell and _num(sell["dispute_hours"], 1, 720) is None:
        raise DealError("dispute_hours: 1 se 720")
    if kw.get("upi_id") is not None or kw.get("payee_name") is not None:
        from . import dealpay
        cur_r = rules(tenant)
        dealpay.set_payee(tenant, kw["upi_id"] if kw.get("upi_id") is not None else (cur_r.upi_id or ""),
                          kw["payee_name"] if kw.get("payee_name") is not None else (cur_r.payee_name or ""))
    if sell:
        from ..avatar import service as av
        cur = av.get(tenant)["rules"]
        av.save(tenant, rules=dict(cur, selling=dict(cur.get("selling") or {}, **sell)))
    with session() as s:
        r = s.get(DealRules, tenant) or DealRules(tenant_id=tenant)
        s.add(r)
        if kw.get("pickup_minutes") is not None:
            v = _num(kw["pickup_minutes"], 15, 7 * 24 * 60)
            if v is None:
                raise DealError("pickup_minutes: 15 minute se 7 din")
            r.pickup_minutes = int(v)
        if isinstance(kw.get("place"), str):
            if policy.secret_found(kw["place"]):
                raise DealError("Jagah mein mana ki gayi jaankari hai.")
            r.place = kw["place"].strip()[:200]
        if isinstance(kw.get("prefix"), str) and re.fullmatch(r"[A-Za-z]{1,6}", kw["prefix"].strip()):
            r.prefix = kw["prefix"].strip().upper()
        r.updated_at = now()
    return rules_view(tenant)


def _named(names: list[str], link, card: dict[str, Any] | None, address: str = "") -> bool:
    """Is this peer one of the names / Hoola IDs she wrote (same matching as VIP)."""
    want = {str(x).lower().lstrip("@") for x in names or [] if str(x).strip()}
    if not want:
        return False
    cand = {address.lower()}
    if link is not None:
        cand |= {(link.nick or "").lower(), (link.relation or "").lower()}
    if card:
        cand |= {str(card.get("owner_name", "")).lower(), str(card.get("handle", "")).lower(),
                 str(card.get("owner_name", "")).lower().split(" ")[0], str(card.get("address", "")).lower()}
    return bool(want & {c for c in cand if c})


def packers(tenant: str) -> list[dict[str, Any]]:
    """Her team members who pack orders: on the Hoola Network, active, manager or given the 'packing' duty; if she
    named packers in Mere Niyam, only those."""
    try:
        from ..employee import service as emp
        team = emp.team(tenant)
    except Exception:
        return []
    names = {str(x).lower().lstrip("@") for x in selling(tenant).get("packers") or []}
    out = []
    for m in team:
        addr = str(m.get("address") or "")
        if not m.get("active", True) or not addr:
            continue
        if not (m.get("role") == "manager" or "packing" in (m.get("duties") or [])):
            continue
        if names and not ({str(m.get("name", "")).lower(), addr.lower()} & names
                          or any(n in addr.lower() for n in names)):
            continue
        l = _link(tenant, addr)
        if l is None or l.status != "active":
            continue
        out.append({"id": m.get("id"), "name": m.get("name"), "address": addr})
    return out


def stock_view(it: StockItem) -> dict[str, Any]:
    free = None if it.qty is None else _r2(max(0.0, it.qty - (it.reserved or 0.0)))
    low = it.low_at is not None and free is not None and free <= it.low_at
    return {"id": it.id, "name": it.name, "aliases": it.aliases or [], "unit": it.unit, "price": it.price,
            "gst_pct": it.gst_pct, "qty": it.qty, "reserved": it.reserved or 0.0, "free": free, "low_at": it.low_at,
            "low": low, "tiers": it.tiers or [], "subs": it.subs or [], "public": it.public is not False,
            "updated": it.updated_at.isoformat() if it.updated_at else ""}


def stock(tenant: str, q: str = "") -> list[dict[str, Any]]:
    w = norm(q)
    with session() as s:
        rows = s.scalars(select(StockItem).where(StockItem.tenant_id == tenant).order_by(StockItem.name)).all()
        out = [stock_view(r) for r in rows if not w or w in r.key or any(w in norm(a) for a in r.aliases or [])]
    return out[:500]


def set_stock(tenant: str, items: list[dict[str, Any]]) -> dict[str, Any]:
    """Add or update items by name. Only the fields given change (qty can be set to null = not counted)."""
    if not isinstance(items, list) or not items:
        raise DealError("items: [{name, price, qty, unit, gst_pct, aliases}] chahiye")
    done = []
    with session() as s:
        for raw in items[:500]:
            if not isinstance(raw, dict):
                continue
            name = str(raw.get("name") or raw.get("item") or "").strip()[:140]
            key = norm(name)
            if not key:
                raise DealError("Har item ka naam chahiye.")
            if policy.secret_found(policy.all_text(raw)):
                raise DealError("Stock mein mana ki gayi jaankari hai.")
            it = s.scalars(select(StockItem).where(StockItem.tenant_id == tenant, StockItem.key == key)).first()
            if it is None:
                it = StockItem(tenant_id=tenant, key=key, name=name)
                s.add(it)
            if "price" in raw:
                p = _num(raw["price"], 0, MAX_TOTAL)
                if p is None:
                    raise DealError(f"{name}: daam galat")
                it.price = _r2(p)
            if "gst_pct" in raw:
                g = _num(raw["gst_pct"], 0, 40)
                if g is None or g not in GST_RATES:
                    raise DealError(f"{name}: GST % in mein se: {', '.join(str(x).rstrip('0').rstrip('.') for x in GST_RATES)}")
                it.gst_pct = g
            if "qty" in raw:
                if raw["qty"] is None:
                    it.qty = None
                else:
                    qv = _num(raw["qty"], 0, 1e7)
                    if qv is None:
                        raise DealError(f"{name}: maatra galat")
                    it.qty = qv
            if "add" in raw:                         # stock came in: add to what is there
                av = _num(raw["add"], -1e7, 1e7)
                if av is None:
                    raise DealError(f"{name}: add galat")
                it.qty = max(0.0, (it.qty or 0.0) + av)
            if isinstance(raw.get("unit"), str):
                it.unit = raw["unit"].strip()[:30]
            if isinstance(raw.get("aliases"), list):
                it.aliases = [str(a).strip()[:80] for a in raw["aliases"][:12] if str(a).strip()]
            if "low_at" in raw:
                it.low_at = None if raw["low_at"] in (None, "") else _num(raw["low_at"], 0, 1e7)
            if isinstance(raw.get("tiers"), list):
                tiers = []
                for t in raw["tiers"][:6]:
                    mn, pr = _num((t or {}).get("min"), 1, 1e7), _num((t or {}).get("price"), 0, MAX_TOTAL)
                    if mn is None or pr is None:
                        raise DealError(f"{name}: bulk daam [{{min, price}}] galat")
                    tiers.append({"min": mn, "price": _r2(pr)})
                it.tiers = sorted(tiers, key=lambda t: t["min"])
            if isinstance(raw.get("subs"), list):
                it.subs = [str(a).strip()[:140] for a in raw["subs"][:5] if str(a).strip()]
            if isinstance(raw.get("public"), bool):
                it.public = raw["public"]
            it.name, it.updated_at = name, now()
            s.flush()
            done.append(stock_view(it))
    return {"saved": len(done), "items": done}


def remove_stock(tenant: str, item: str) -> dict[str, Any]:
    with session() as s:
        it = s.get(StockItem, item)
        if it is None or it.tenant_id != tenant:
            it = s.scalars(select(StockItem).where(StockItem.tenant_id == tenant, StockItem.key == norm(item))).first()
        if it is None or it.tenant_id != tenant:
            raise DealError("Yeh item stock mein nahi mila.")
        if (it.reserved or 0) > 0:
            raise DealError("Is item ka kuch maal pakke order mein rakha hai; pehle wo order poora / cancel kijiye.")
        s.delete(it)
    return {"removed": True}


def _catalog_items(tenant: str) -> list[dict[str, Any]]:
    """Stock list first; a business with only the v0.33 price list quotes from it (quantity unknown = ask her)."""
    with session() as s:
        rows = s.scalars(select(StockItem).where(StockItem.tenant_id == tenant)).all()
        if rows:
            return [{"id": r.id, "name": r.name, "key": r.key, "aliases": [norm(a) for a in r.aliases or []], "unit": r.unit,
                     "price": r.price, "gst_pct": r.gst_pct, "tiers": r.tiers or [], "subs": r.subs or [],
                     "public": r.public is not False,
                     "free": None if r.qty is None else max(0.0, r.qty - (r.reserved or 0.0))} for r in rows]
    from .business import Business
    with session() as s:
        b = s.get(Business, tenant)
        cat = list(b.catalog or []) if b is not None else []
    out = []
    for c in cat:
        if c.get("price") is None:
            continue
        out.append({"id": "", "name": str(c.get("item") or ""), "key": norm(c.get("item")), "aliases": [],
                    "unit": str(c.get("unit") or ""), "price": float(c["price"]), "gst_pct": 0.0, "free": None,
                    "tiers": [], "subs": [], "public": True})
    return out


def match(items: list[dict[str, Any]], asked: str) -> dict[str, Any] | None:
    """The stock item a buyer's words mean: same name / alias, all asked words inside the name, or a very close spelling."""
    w = norm(asked)
    if not w:
        return None
    for it in items:
        if w == it["key"] or w in it["aliases"]:
            return it
    words = set(w.split())
    inside = [it for it in items if words and words <= set(it["key"].split())]
    if len(inside) == 1:
        return inside[0]
    best, score = None, 0.0
    for it in items:
        for k in [it["key"], *it["aliases"]]:
            r = difflib.SequenceMatcher(None, w, k).ratio()
            if r > score:
                best, score = it, r
    return best if score >= 0.86 else None


def build_quote(tenant: str, asked: list[dict[str, Any]]) -> dict[str, Any]:
    """The quote, made by code from the seller's own list (never by the model): one line per asked item."""
    items = _catalog_items(tenant)
    lines = []
    for a in asked:
        want = float(a["qty"])
        m = match(items, a["item"])
        if m is None:
            lines.append({"item": a["item"], "asked": want, "qty": 0.0, "unit": a.get("unit", ""), "price": None,
                          "gst_pct": 0.0, "amount": 0.0, "gst": 0.0, "status": "not_found"})
            continue
        free = m["free"]
        if free is None:
            qty, status = want, "unknown_stock"
        elif free >= want:
            qty, status = want, "ok"
        elif free > 0:
            qty, status = float(int(free) if free == int(free) else free), "partial"
        else:
            qty, status = 0.0, "out"
        lines.append(_line(m["name"], want, qty, m["unit"] or a.get("unit", ""), price_for(m, qty), m["gst_pct"], status,
                           stock_id=m["id"], asked_as=a["item"]))
        gap = want - qty if status in ("partial", "out") else 0.0
        for sub_name in m.get("subs") or []:                 # her own "this can stand in" list, only from free stock
            if gap <= 0:
                break
            sm = next((it for it in items if it["key"] == norm(sub_name) and it["id"] != m["id"]), None)
            if sm is None or sm["free"] is None or sm["free"] <= 0:
                continue
            sq = min(gap, sm["free"])
            sq = float(int(sq)) if sq == int(sq) else sq
            lines.append(_line(sm["name"], 0.0, sq, sm["unit"], price_for(sm, sq), sm["gst_pct"], "substitute",
                               stock_id=sm["id"], asked_as=a["item"]))
            gap -= sq
    return totals({"lines": lines})


def price_for(it: dict[str, Any], qty: float) -> float:
    """Her bulk price: the lowest tier whose minimum this quantity reaches, else the list price."""
    p = float(it["price"])
    for t in it.get("tiers") or []:
        if qty >= float(t["min"]):
            p = min(p, float(t["price"]))
    return p


def _line(item: str, asked: float, qty: float, unit: str, price: float | None, gst_pct: float, status: str,
          stock_id: str = "", asked_as: str = "") -> dict[str, Any]:
    amount = _r2(qty * (price or 0.0))
    ln = {"item": item, "asked": asked, "qty": qty, "unit": unit, "price": price, "gst_pct": gst_pct, "amount": amount,
          "gst": _r2(amount * gst_pct / 100.0), "status": status}
    if stock_id:
        ln["stock_id"] = stock_id
    if asked_as and norm(asked_as) != norm(item):
        ln["asked_as"] = asked_as[:120]
    return ln


def totals(q: dict[str, Any]) -> dict[str, Any]:
    sub = _r2(sum(l["amount"] for l in q["lines"]))
    gst = _r2(sum(l["gst"] for l in q["lines"]))
    q.update(subtotal=sub, gst=gst, total=_r2(sub + gst),
             complete=all(l["status"] == "ok" for l in q["lines"]),
             missing=[l["item"] for l in q["lines"] if l["status"] in ("not_found", "out")],
             short=[l["item"] for l in q["lines"] if l["status"] == "partial"],
             unknown=[l["item"] for l in q["lines"] if l["status"] == "unknown_stock" or l.get("price") is None and l["qty"] > 0])
    return q


# ---------------------------------------------------------------- validators (every message is checked on arrival)
def clean_items(raw: Any) -> list[dict[str, Any]]:
    if not isinstance(raw, list) or not raw:
        raise DealError("Kya chahiye? items: [{item, qty, unit}] (kam se kam ek)")
    if len(raw) > MAX_LINES:
        raise DealError(f"Ek baar mein {MAX_LINES} se zyada item nahi.")
    out = []
    for x in raw:
        if isinstance(x, str):
            x = {"item": x, "qty": 1}
        if not isinstance(x, dict):
            raise DealError("item galat")
        item = policy.clean_name(x.get("item") or x.get("name"), 120)
        qty = _num(x.get("qty", 1), 0.001, 1e6)
        if not item or qty is None:
            raise DealError(f"'{x.get('item') or x.get('name') or '?'}': naam aur maatra (0 se zyada) chahiye")
        out.append({"item": item, "qty": qty, "unit": policy.clean_name(x.get("unit"), 20),
                    "note": policy.quote(x.get("note"), 200)})
    return out


def check_ask(body: dict[str, Any]) -> dict[str, Any]:
    items = clean_items(body.get("items"))
    mode = body.get("mode") if body.get("mode") in ("pickup", "delivery") else "pickup"
    need_by = _when(body.get("need_by")) if body.get("need_by") else None
    if body.get("need_by") and (need_by is None or need_by < now() - timedelta(hours=1)
                                or need_by > now() + timedelta(days=120)):
        raise DealError("need_by galat (aaj se 120 din ke andar)")
    return {"items": items, "mode": mode, "area": policy.quote(body.get("area"), 160), "note": policy.quote(body.get("note"), 600),
            "title": policy.quote(body.get("title"), 120), "need_by": need_by.isoformat() if need_by else ""}


def check_quote(body: dict[str, Any]) -> dict[str, Any]:
    """The buyer's verifier: every number is recomputed here; a quote whose sums don't add up is refused."""
    v = _num(body.get("v"), 1, 999)
    if v is None or v != int(v):
        raise DealError("quote version galat")
    raw = body.get("lines")
    if not isinstance(raw, list) or not raw or len(raw) > MAX_LINES:
        raise DealError("quote lines galat")
    lines = []
    for x in raw:
        if not isinstance(x, dict):
            raise DealError("quote line galat")
        qty = _num(x.get("qty"), 0, 1e6)
        asked = _num(x.get("asked", x.get("qty")), 0, 1e6)
        price = None if x.get("price") is None else _num(x.get("price"), 0, MAX_TOTAL)
        gpct = _num(x.get("gst_pct", 0), 0, 40)
        status = x.get("status") if x.get("status") in ("ok", "partial", "out", "not_found", "unknown_stock", "substitute") else ""
        if qty is None or asked is None or gpct is None or gpct not in GST_RATES or not status:
            raise DealError("quote line mein maatra / GST / status galat")
        if price is None and qty > 0:
            raise DealError("daam ke bina maatra nahi ho sakti")
        ln = _line(policy.clean_name(x.get("item"), 120) or "?", asked, qty, policy.clean_name(x.get("unit"), 20), price,
                   gpct, status, asked_as=policy.clean_name(x.get("asked_as"), 120))
        for k in ("amount", "gst"):
            if abs(float(x.get(k) or 0) - ln[k]) > 0.011:
                raise DealError(f"hisaab galat: {ln['item']} ka {k} {x.get(k)} likha hai, sahi {ln[k]}")
        lines.append(ln)
    q = totals({"v": int(v), "lines": lines})
    for k in ("subtotal", "gst", "total"):
        if abs(float(body.get(k) or 0) - q[k]) > 0.011:
            raise DealError(f"hisaab galat: {k} {body.get(k)} likha hai, joda to {q[k]}")
    if q["total"] > MAX_TOTAL:
        raise DealError("rakam bahut badi")
    vu = _when(body.get("valid_until"))
    if vu is None or vu < now() or vu > now() + timedelta(days=7, minutes=5):
        raise DealError("quote ki aakhri tareekh galat ya beet chuki")
    q["valid_until"] = vu.isoformat()
    q["note"] = policy.quote(body.get("note"), 600)
    q["place"] = policy.quote(body.get("place"), 200)
    from . import dealpay
    q["terms"] = dealpay.check_terms(body.get("terms"), q["total"])
    return q


def check_ready(body: dict[str, Any]) -> dict[str, Any]:
    code = str(body.get("code") or "")
    if not re.fullmatch(r"\d{4}", code):
        raise DealError("pickup code galat")
    until = _when(body.get("until"))
    if until is None or until < now() or until > now() + timedelta(days=8):
        raise DealError("kab tak le jaana hai, galat")
    return {"code": code, "place": policy.quote(body.get("place"), 200), "until": until.isoformat(),
            "number": policy.clean_name(body.get("number"), 30), "note": policy.quote(body.get("note"), 300)}


# ---------------------------------------------------------------- records
def _get(tenant: str, deal_id: str) -> Deal:
    with session() as s:
        d = s.get(Deal, deal_id)
        if d is None or d.tenant_id != tenant:
            d = s.scalars(select(Deal).where(Deal.tenant_id == tenant, Deal.number == deal_id)).first()
        if d is None:
            raise DealError("Yeh deal nahi mili.")
        s.expunge(d)
        return d


def _by_key(tenant: str, dkey: str, peer: str) -> Deal | None:
    with session() as s:
        d = s.scalars(select(Deal).where(Deal.tenant_id == tenant, Deal.dkey == dkey, Deal.peer == peer)).first()
        if d is not None:
            s.expunge(d)
        return d


def _expired(d: Deal) -> bool:
    vu = _when((d.quote or {}).get("valid_until"))
    return d.state == "quoted" and vu is not None and vu < now()


def actions(d: Deal) -> list[str]:
    st, buyer = d.state, d.role == "buyer"
    from . import dealpay
    extra = dealpay.actions(d)
    if st in END_STATES:
        return extra
    if buyer:
        return {"asked": ["cancel"], "quoted": ["confirm", "cancel"], "confirmed": ["cancel"],
                "accepted": ["cancel"], "packing": ["cancel"], "ready": [], "handed_over": ["ok", "dispute"],
                "disputed": []}.get(st, []) + extra
    return {"asked": ["send_quote", "reject"], "quoted": ["send_quote", "reject"], "confirmed": ["accept", "reject"],
            "accepted": ["packing", "ready", "cancel"], "packing": ["ready", "cancel"],
            "ready": ["handed_over", "cancel"], "handed_over": [], "disputed": []}.get(st, []) + extra


def view(d: Deal) -> dict[str, Any]:
    if _expired(d):
        d.state = "expired"
    c = idn.peer_card(d.peer) or {}
    q = d.quote or {}
    rdy = dict(d.ready or {})
    out = {"id": d.id, "deal_id": d.id, "role": d.role, "state": d.state, "state_say": STATE_WORDS.get(d.state, d.state),
           "number": d.number, "title": d.title, "peer": d.peer, "with": idn.label(c), "peer_avatar": c.get("avatar_name", ""),
           "business": bool(c.get("business")), "verified": bool(c.get("verified")),
           "ask": d.ask or {}, "quote": q, "draft": (d.draft or {}) if d.role == "seller" else {},
           "total": q.get("total"), "valid_until": q.get("valid_until", ""), "accept": d.accept or {}, "ready": rdy,
           "reason": d.reason, "waiting": d.waiting if d.state in OPEN_STATES else "", "actions": actions(d),
           "thread_id": d.thread_id, "history": (d.history or [])[-20:], "updated": d.updated_at.isoformat() if d.updated_at else ""}
    if d.role == "seller" and rdy.get("code"):
        out["ready"] = {k: v for k, v in rdy.items() if k != "code"}           # her staff checks the code; it isn't shown here
        out["ready"]["code_set"] = True
    from . import dealpay
    out |= {"pay": dealpay.pay_view(d), "dispute": d.dispute or {}, "cancel_req": d.cancel_req or {},
            "terms": q.get("terms") or {}, "handed_at": d.handed_at.isoformat() if d.handed_at else ""}
    return out


def deals(tenant: str, state: str = "", role: str = "", limit: int = 60) -> list[dict[str, Any]]:
    with session() as s:
        q = select(Deal).where(Deal.tenant_id == tenant)
        if role in ("buyer", "seller"):
            q = q.where(Deal.role == role)
        if state == "open":
            q = q.where(Deal.state.in_(OPEN_STATES))
        elif state:
            q = q.where(Deal.state == state)
        rows = s.scalars(q.order_by(Deal.updated_at.desc()).limit(max(1, min(int(limit or 60), 200)))).all()
        for r in rows:
            s.expunge(r)
    return [view(r) for r in rows]


def show(tenant: str, deal_id: str) -> dict[str, Any]:
    return view(_get(tenant, deal_id))


def _record(d: Deal, kind: str, body: dict[str, Any], by: str, say: str) -> str:
    """Advance this side's chain: the next step's hash, stored with a short line for the history."""
    h = step_hash(d.last_hash, kind, d.step + 1, body)
    d.step += 1
    d.last_hash = h
    d.history = (d.history or []) + [{"step": d.step, "kind": kind, "by": by, "at": now().isoformat(), "hash": h[:12],
                                      "say": say[:200]}]
    d.updated_at = now()
    return h


def _save(d: Deal) -> None:
    with session() as s:
        s.merge(d)


def _publish(tenant: str, d: Deal, card: str, say: str, note: str = "") -> None:
    v = view(d)
    tell(tenant, "hoola_deal", v | {"card": card, "say": say},
         note or f"Deal ({d.role}) {d.number or d.id} {v['with']}: {say} deal_id={d.id}.")


def _send(tenant: str, d: Deal, kind: str, body: dict[str, Any], say: str) -> dict[str, Any]:
    """Advance the state, sign and send one step. The record is saved BEFORE delivery: the answer may arrive while this
    call is still delivering (both avatars can live on this server)."""
    key = (d.state if not (kind == "deal_ask") else "", _event_key(kind, body))
    if key not in TRANSITIONS:
        raise DealError(f"Abhi '{STATE_WORDS.get(d.state, d.state)}' hai; ab yeh nahi ho sakta.")
    if policy.secret_found(policy.all_text(body)):
        raise DealError("Isme PIN / OTP / password / card / Aadhaar jaisi jaankari hai; Hoola yeh kabhi nahi bhejta.")
    before = (d.state, d.step, d.last_hash, list(d.history or []))
    body = dict(body, step=d.step + 1, prev=d.last_hash)
    _record(d, kind, body, d.role, say)
    d.state = TRANSITIONS[key]
    _save(d)
    res = _post(tenant, d.peer, kind, body, d.dkey, d.thread_id)
    if not res.get("ok") and res.get("status") not in ("duplicate",):
        with session() as s:                      # the other side refused / didn't get it: undo this step here
            row = s.get(Deal, d.id)
            if row is not None and row.step == before[1] + 1:
                row.state, row.step, row.last_hash, row.history = before[0], before[1], before[2], before[3]
        raise DealError(f"{_say_who(d.peer)} tak nahi pahuncha: {res.get('reason') or res.get('status')}")
    return res


# ---------------------------------------------------------------- buyer side
def ask(tenant: str, to: str, items: Any, mode: str = "pickup", need_by: str = "", note: str = "", title: str = "",
        area: str = "") -> dict[str, Any]:
    """The buyer's list goes to the seller's avatar (a connected person, or a public business counter)."""
    body = check_ask({"items": items, "mode": mode, "need_by": need_by, "note": note, "title": title, "area": area})
    if policy.secret_found(policy.all_text(body)):
        raise DealError("List mein PIN / OTP / password / card / Aadhaar jaisi jaankari hai; yeh nahi jaayega.")
    peer = resolve(tenant, to)
    l = _link(tenant, peer)
    if l is None or l.status != "active":
        from . import business
        try:
            peer = business.contact(tenant, to)        # a public business counter opens a customer connection
        except NetError:
            raise DealError(f"{_say_who(peer)} se juda nahi hai. Pehle connect kijiye (ya woh public business ho).")
    if peer == idn.my_card(tenant)["address"]:
        raise DealError("Yeh aapka apna Hoola hai.")
    dkey = new_id("dk")
    title = body["title"] or (", ".join(f"{i['item']} {i['qty']:g}" for i in body["items"][:3])
                              + (f" +{len(body['items']) - 3}" if len(body["items"]) > 3 else ""))
    tid = _thread(tenant, peer, dkey, "Deal: " + title, "deal", True)
    d = Deal(id=new_id("dl"), tenant_id=tenant, dkey=dkey, role="buyer", peer=peer, thread_id=tid, title=title[:200],
             state="", ask=body, step=0, last_hash="", history=[], quote={}, draft={}, accept={}, ready={}, reserved=[],
             smap=[], reason="", waiting="", number="", created_at=now(), updated_at=now())
    with session() as s:
        s.add(d)
        s.flush()
        s.expunge(d)
    STATS["asks"] += 1
    try:
        res = _send(tenant, d, "deal_ask", body, f"{len(body['items'])} item ki list bheji")
    except DealError:
        with session() as s:
            row = s.get(Deal, d.id)
            if row is not None and row.step == 0:
                s.delete(row)
        raise
    d = _get(tenant, d.id)
    if d.state == "asked":                 # (a quote may already have come back while sending: its card says it all)
        _publish(tenant, d, "info", f"List {_say_who(peer)} ko bheji ({len(body['items'])} item).")
    return {"status": res.get("status"), "deal": view(d), "say": f"List bhej di. {_say_who(peer)} ka Hoola stock dekh kar quote bhejega."}


def confirm_card(tenant: str, deal_id: str) -> dict[str, Any] | None:
    """What the buyer's confirm tap says (also hashed by the fingerprint approval)."""
    try:
        d = _get(tenant, deal_id)
    except DealError:
        return None
    if d.role != "buyer" or d.state != "quoted":
        return None
    q = d.quote or {}
    return {"title": f"Order confirm: {d.title}", "payee": _say_who(d.peer), "amount": inr(q.get("total")),
            "items": [{"name": f"{l['item']} x {l['qty']:g}", "price": inr(l['amount'] + l['gst'])}
                      for l in q.get("lines") or [] if l.get("qty")], "v": q.get("v"), "hash": quote_hash(q)}


def confirm(tenant: str, deal_id: str, v: int | None = None, mode: str = "", by_tap: bool = False,
            pay: str = "") -> dict[str, Any]:
    """The buyer commits. ONLY from her tap on the card (the API route); the model can never call this with by_tap."""
    if not by_tap:
        raise DealError("Order pakka karna sirf aapke tap se hota hai: card par 'Haan, confirm' dabaiye.")
    d = _get(tenant, deal_id)
    if d.role != "buyer":
        raise DealError("Yeh aapki khareed nahi hai.")
    if _expired(d):
        d.state, d.waiting = "expired", ""
        _save(d)
        raise DealError("Quote ka samay khatam ho gaya. Naya quote maangiye.")
    if d.state != "quoted":
        raise DealError(f"Abhi '{STATE_WORDS.get(d.state, d.state)}' hai; confirm nahi ho sakta.")
    q = d.quote or {}
    if v is not None and int(v) != int(q.get("v") or 0):
        raise DealError("Quote badal gaya hai; naya quote dekh kar confirm kijiye.")
    if not any(l.get("qty") for l in q.get("lines") or []):
        raise DealError("Is quote mein koi saamaan nahi hai.")
    opts = list(((q.get("terms") or {}).get("pay_options")) or ["upi", "counter"])
    if pay and pay not in opts:
        raise DealError(f"Is quote mein '{pay}' nahi hai (khata sirf dukaan ke kholne par). Chuniye: " + ", ".join(opts))
    pay = pay or ("upi" if "upi" in opts else opts[0])
    body = {"v": q.get("v"), "quote_hash": quote_hash(q), "mode": mode if mode in ("pickup", "delivery") else
            (d.ask or {}).get("mode", "pickup"), "pay": pay}
    d.waiting = ""
    d.pay = dict(d.pay or {}, choice=pay)
    _send(tenant, d, "deal_confirm", body, f"Quote v{q.get('v')} confirm, {inr(q.get('total'))}")
    STATS["confirms"] += 1
    d = _get(tenant, deal_id)
    if d.state == "confirmed":
        _publish(tenant, d, "info", f"Order confirm bheja ({inr(q.get('total'))}). Unka Hoola stock dekh kar pakka karega.")
    return {"status": d.state, "deal": view(d)}


# ---------------------------------------------------------------- seller side
def _may_auto_quote(tenant: str, peer: str, q: dict[str, Any]) -> bool:
    """Her Mere Niyam: a quote goes by itself only when EVERY line is in stock, no line was swapped or guessed, the total
    is inside her limit and this peer is someone her rule covers."""
    so = selling(tenant)
    if so["auto_quote"] == "off" or not q.get("complete") or q["total"] <= 0 or q["total"] > so["auto_quote_max"]:
        return False
    if any(l["status"] != "ok" for l in q.get("lines") or []):
        return False
    l = _link(tenant, peer)
    if l is None or l.status != "active" or l.key_warning:
        return False
    if so["auto_quote"] == "connected":
        return True
    return bool(l.vip or l.relation in ("family", "team") or _named(so.get("trusted") or [], l, idn.peer_card(peer), peer))


def send_quote(tenant: str, deal_id: str, lines: Any = None, note: str = "", valid_minutes: Any = None,
               by: str = "owner") -> dict[str, Any]:
    """The seller's quote goes out: the built draft, or her changed lines (qty / price per item). Code does the sums."""
    d = _get(tenant, deal_id)
    if d.role != "seller":
        raise DealError("Quote sirf bechne wala bhejta hai.")
    if d.state not in ("asked", "quoted"):
        raise DealError(f"Abhi '{STATE_WORDS.get(d.state, d.state)}' hai; quote nahi bhej sakte.")
    base = dict(d.draft or {}) or build_quote(tenant, (d.ask or {}).get("items") or [])
    if lines:
        base = _edit_lines(base, lines)
    if not any(l["qty"] for l in base["lines"]):
        raise DealError("Quote mein ek bhi item nahi hai (sab khatam / nahi mila). Mana (reject) kar dijiye ya daam / maatra dijiye.")
    for l in base["lines"]:
        if l["qty"] and l.get("price") is None:
            raise DealError(f"'{l['item']}' ka daam bataiye.")
    r = rules(tenant)
    mins = int(_num(valid_minutes, 5, 7 * 24 * 60) or selling(tenant)["valid_minutes"])
    v = int((d.quote or {}).get("v") or 0) + 1
    q = {"v": v, "lines": [{k: l[k] for k in ("item", "asked", "qty", "unit", "price", "gst_pct", "amount", "gst", "status")}
                           | ({"asked_as": l["asked_as"]} if l.get("asked_as") else {}) for l in base["lines"]],
         "subtotal": base["subtotal"], "gst": base["gst"], "total": base["total"],
         "valid_until": (now() + timedelta(minutes=mins)).replace(microsecond=0).isoformat(),
         "note": policy.quote(note, 600), "place": r.place}
    if q["total"] > MAX_TOTAL:
        raise DealError("Rakam bahut badi (1 crore tak).")
    from . import dealpay
    q["terms"] = dealpay.terms_for(tenant, d.peer, q["total"])          # v0.106 pay options, payee, windows
    q = check_quote(q)                  # exactly what the buyer's side will compute, so both hash the same quote
    d.quote, d.draft, d.waiting = q, {}, ""
    d.smap = [str(l.get("stock_id") or "") for l in base["lines"]]
    _send(tenant, d, "deal_quote", q, f"Quote v{v}: {inr(q['total'])}, {mins} min ke liye" + (" (khud)" if by == "rules" else ""))
    STATS["quotes"] += 1
    d = _get(tenant, deal_id)
    _publish(tenant, d, "info", ("Aapke niyam se quote khud bheja: " if by == "rules" else "Quote bheja: ") + f"{inr(q['total'])}.")
    return {"status": d.state, "deal": view(d)}


def _edit_lines(base: dict[str, Any], lines: Any) -> dict[str, Any]:
    if not isinstance(lines, list):
        raise DealError("lines: [{item, qty, price, gst_pct}]")
    by_item = {norm(l["item"]): l for l in base["lines"]}
    by_item.update({norm(l.get("asked_as")): l for l in base["lines"] if l.get("asked_as")})
    out = [dict(l) for l in base["lines"]]
    for ch in lines[:MAX_LINES]:
        if not isinstance(ch, dict):
            continue
        k = norm(ch.get("item"))
        cur = by_item.get(k)
        if cur is None:                           # an extra line / a substitute she offers
            item = policy.clean_name(ch.get("item"), 120)
            qty, price = _num(ch.get("qty"), 0, 1e6), _num(ch.get("price"), 0, MAX_TOTAL)
            g = _num(ch.get("gst_pct", 0), 0, 40)
            if not item or qty is None or price is None or g is None or g not in GST_RATES:
                raise DealError(f"'{ch.get('item')}': naam, maatra, daam aur sahi GST % chahiye")
            out.append(_line(item, 0.0, qty, policy.clean_name(ch.get("unit"), 20), price, g, "substitute"))
            continue
        i = out.index(next(x for x in out if x["item"] == cur["item"]))
        ln = dict(out[i])
        if ch.get("qty") is not None:
            qv = _num(ch["qty"], 0, 1e6)
            if qv is None:
                raise DealError(f"{ln['item']}: maatra galat")
            ln["qty"] = qv
        if ch.get("price") is not None:
            pv = _num(ch["price"], 0, MAX_TOTAL)
            if pv is None:
                raise DealError(f"{ln['item']}: daam galat")
            ln["price"] = pv
        if ch.get("gst_pct") is not None:
            gv = _num(ch["gst_pct"], 0, 40)
            if gv is None or gv not in GST_RATES:
                raise DealError(f"{ln['item']}: GST % galat")
            ln["gst_pct"] = gv
        status = "substitute" if ln["status"] == "substitute" else \
            ("ok" if ln["qty"] >= ln["asked"] else ("partial" if ln["qty"] > 0 else "out"))
        out[i] = _line(ln["item"], ln["asked"], ln["qty"], ln["unit"], ln["price"], ln["gst_pct"], status,
                       stock_id=ln.get("stock_id", ""), asked_as=ln.get("asked_as", ""))
    return totals({"lines": out})


def _stock_holds(tenant: str, q: dict[str, Any], smap: list[str] | None = None) -> list[str]:
    """Items the stock can no longer cover (checked again at accept: someone else may have bought them)."""
    items = {it["id"]: it for it in _catalog_items(tenant) if it["id"]}
    short = []
    for i, l in enumerate(q.get("lines") or []):
        if not l.get("qty"):
            continue
        it = items.get((smap or [])[i] if i < len(smap or []) else "")
        if it is None:
            continue                                # an extra / substitute line she typed herself: her word
        if it["free"] is not None and it["free"] + 1e-9 < l["qty"]:
            short.append(f"{l['item']} ({it['free']:g} bacha)")
    return short


def accept(tenant: str, deal_id: str, ready_by: str = "", note: str = "", by: str = "owner") -> dict[str, Any]:
    """The seller commits: order number, stock held for it, pick list for packing."""
    d = _get(tenant, deal_id)
    if d.role != "seller" or d.state != "confirmed":
        raise DealError(f"Abhi '{STATE_WORDS.get(d.state, d.state)}' hai; order pakka nahi ho sakta.")
    short = _stock_holds(tenant, d.quote or {}, d.smap or [])
    if short:
        raise DealError("Stock kam pad gaya: " + ", ".join(short) + ". Mana kijiye ya naya quote bhejiye.")
    from . import dealpay
    pay_part = dealpay.at_accept(tenant, d)            # v0.106 khata limit checked again; the pay record starts
    rules(tenant)                                  # makes her rules row if missing
    with session() as s:
        rr = s.get(DealRules, tenant)
        number = f"{rr.prefix}-{rr.next_no}"
        rr.next_no += 1
        held = []
        smap = d.smap or []
        for i, l in enumerate((d.quote or {}).get("lines") or []):
            if not l.get("qty") or i >= len(smap) or not smap[i]:
                continue
            it = s.get(StockItem, smap[i])
            if it is not None and it.tenant_id == tenant and it.qty is not None:
                it.reserved = _r2((it.reserved or 0.0) + l["qty"])
                held.append({"stock_id": it.id, "qty": l["qty"]})
    rb = _when(ready_by) if ready_by else None
    if rb is None and ready_by:
        try:
            from ..life import parse_local, tz_of
            rb = parse_local(ready_by, tz_of(tenant)).astimezone(timezone.utc)
        except Exception:
            rb = None
    body = {"number": number, "ready_by": rb.isoformat() if rb else "", "note": policy.quote(note, 300)} | pay_part["body"]
    d.number, d.accept, d.reserved, d.waiting = number, body, held, ""
    d.pay, d.accepted_at = pay_part["pay"], now()
    try:
        _send(tenant, d, "deal_accept", body, f"Order {number} pakka" + (" (niyam se)" if by == "rules" else ""))
    except DealError:
        _release(tenant, held)
        raise
    dealpay.after_accept(tenant, _get(tenant, deal_id))
    STATS["accepts"] += 1
    pick = [f"{l['item']} x {l['qty']:g} {l.get('unit') or ''}".strip() for l in (d.quote or {}).get("lines") or [] if l.get("qty")]
    from . import pick as pk
    sent = pk.send_tasks(tenant, deal_id)          # v0.105 the pick list goes to her packers' own Hoolas
    d = _get(tenant, deal_id)
    _publish(tenant, d, "pack", ("Aapke niyam se order khud pakka hua. " if by == "rules" else "") +
             f"Order {number}: {len(pick)} item pack karne hain." +
             (f" Pick list {', '.join(sent)} ko bheji." if sent else ""),
             f"Order {number} pakka ({_say_who(d.peer)}). Pick list: " + "; ".join(pick)[:900] + f". deal_id={d.id}.")
    return {"status": d.state, "number": number, "pick_list": pick, "packers": sent, "deal": view(d)}


def _release(tenant: str, held: list[dict[str, Any]], used: bool = False) -> None:
    with session() as s:
        for h in held or []:
            it = s.get(StockItem, h.get("stock_id"))
            if it is None or it.tenant_id != tenant:
                continue
            q = float(h.get("qty") or 0)
            it.reserved = _r2(max(0.0, (it.reserved or 0.0) - q))
            if used and it.qty is not None:
                it.qty = _r2(max(0.0, it.qty - q))
            it.updated_at = now()


def reject(tenant: str, deal_id: str, reason: str = "") -> dict[str, Any]:
    d = _get(tenant, deal_id)
    if d.role != "seller":
        raise DealError("Khareedne wala cancel karta hai, mana nahi.")
    reason = policy.quote(reason, 300) or "Abhi yeh order nahi le sakte."
    d.reason, d.waiting = reason, ""
    _send(tenant, d, "deal_reject", {"reason": reason}, "Mana: " + reason)
    d = _get(tenant, deal_id)
    _publish(tenant, d, "info", "Mana kar diya.")
    return {"status": d.state, "deal": view(d)}


def packing(tenant: str, deal_id: str, note: str = "") -> dict[str, Any]:
    d = _get(tenant, deal_id)
    if d.role != "seller":
        raise DealError("Packing bechne wala karta hai.")
    _send(tenant, d, "deal_status", {"status": "packing", "note": policy.quote(note, 300)}, "Packing shuru")
    return {"status": "packing", "deal": view(_get(tenant, deal_id))}


def mark_ready(tenant: str, deal_id: str, place: str = "", minutes: Any = None, note: str = "") -> dict[str, Any]:
    """Packed: a 4-digit pickup code goes to the buyer; at the counter the code is checked before handing over."""
    d = _get(tenant, deal_id)
    if d.role != "seller":
        raise DealError("Parcel bechne wala taiyaar karta hai.")
    again = d.state == "ready"
    r = rules(tenant)
    mins = int(_num(minutes, 15, 7 * 24 * 60) or r.pickup_minutes)
    place = policy.quote(place, 200) or r.place or "Dukaan ke counter par"
    code = f"{secrets.randbelow(9000) + 1000}"
    body = {"number": d.number, "code": code, "place": place,
            "until": (now() + timedelta(minutes=mins)).replace(microsecond=0).isoformat(), "note": policy.quote(note, 300)}
    d.ready = body
    _send(tenant, d, "deal_ready", body, f"Parcel taiyaar: {place}")
    STATS["ready"] += 1
    from . import pick as pk
    pk.notify(tenant, deal_id, "ready")
    d = _get(tenant, deal_id)
    _publish(tenant, d, "info", (f"Order {d.number}: naya pickup code aur samay graahak ko bheja (purana code band)." if again else
                                 f"Order {d.number} taiyaar; graahak ko pickup code bheja. Counter par code milaiye."))
    return {"status": d.state, "deal": view(d)}


def handed_over(tenant: str, deal_id: str, code: str = "", note: str = "", collect: dict | None = None) -> dict[str, Any]:
    """At the counter: the buyer's code must match, and the money must be settled (paid, on khata, or taken now at the
    counter: `collect` = {mode cash|upi, amount, utr}). Stock held for the order is taken out of the count."""
    d = _get(tenant, deal_id)
    if d.role != "seller" or d.state != "ready":
        raise DealError("Pehle parcel taiyaar (ready) hona chahiye.")
    want = str((d.ready or {}).get("code") or "")
    if not want or str(code or "").strip() != want:
        STATS["refused"] += 1
        raise DealError("Pickup code nahi mila. Graahak se unke Hoola mein dikha code poochhiye.")
    from . import dealpay
    dealpay.check_handover(tenant, d, collect)           # v0.106 raises when money is open and nothing was collected
    d = _get(tenant, deal_id)
    d.handed_at = now()
    _send(tenant, d, "deal_status", {"status": "handed_over", "note": policy.quote(note, 300)}, "Saamaan de diya")
    _release(tenant, d.reserved or [], used=True)
    STATS["handed_over"] += 1
    from . import pick as pk
    pk.notify(tenant, deal_id, "handed_over")
    low = low_stock(tenant, [h.get("stock_id") for h in d.reserved or []])
    d = _get(tenant, deal_id)
    _publish(tenant, d, "info", f"Order {d.number} graahak ko de diya." +
             (" Stock kam: " + ", ".join(f"{x['name']} ({_qty(x['free'])} bacha)" for x in low) + "." if low else ""))
    return {"status": d.state, "deal": view(d)}


def cancel(tenant: str, deal_id: str, reason: str = "") -> dict[str, Any]:
    d = _get(tenant, deal_id)
    if d.state in END_STATES:
        raise DealError(f"Yeh deal pehle hi '{STATE_WORDS.get(d.state, d.state)}' hai.")
    if d.role == "buyer" and d.state == "ready":
        raise DealError("Parcel taiyaar ho chuka hai; cancel ke liye dukaan se baat kijiye.")
    if d.state in ("handed_over", "disputed"):
        raise DealError("Saamaan de diya gaya hai; ab cancel nahi, kuch galat ho to shikayat (dispute) kijiye.")
    from . import dealpay
    if d.role == "buyer" and dealpay.cancel_needs_ok(d):  # v0.106 after the seller's free-cancel window: a request
        return dealpay.cancel_request(tenant, d.id, reason)
    reason = policy.quote(reason, 300) or "Cancel"
    held = list(d.reserved or [])
    d.reason, d.waiting = reason, ""
    _send(tenant, d, "deal_cancel", {"reason": reason}, "Cancel: " + reason)
    if d.role == "seller":
        _release(tenant, held)
        from . import pick as pk
        pk.notify(tenant, deal_id, "cancelled")
    dealpay.after_cancel(tenant, deal_id)
    d = _get(tenant, deal_id)
    _publish(tenant, d, "info", "Deal cancel ho gayi." + (f" {d.pay.get('refund_say')}" if (d.pay or {}).get("refund_say") else ""))
    return {"status": d.state, "deal": view(d)}


# ---------------------------------------------------------------- receiving
def receive(tenant: str, env: dict[str, Any], card: dict[str, Any], link, rules_: dict[str, Any]) -> dict[str, Any]:
    kind, raw, peer = env["kind"], env.get("body") or {}, env["from"]
    if kind == "deal_catalog":
        return _on_catalog(tenant, peer, card, link)
    try:
        return _receive(tenant, kind, raw, peer, env, card, link)
    except DealError as e:
        STATS["refused"] += 1
        return {"ok": False, "status": "refused", "reason": str(e)[:300]}


def _receive(tenant: str, kind: str, raw: dict[str, Any], peer: str, env: dict, card: dict, link) -> dict[str, Any]:
    if policy.secret_found(policy.all_text(raw)):
        raise DealError("mana ki gayi jaankari")
    d = _by_key(tenant, env.get("thread", ""), peer)
    try:
        step = int(raw.get("step"))
    except (TypeError, ValueError):
        raise DealError("step galat")
    prev = str(raw.get("prev") or "")
    if kind == "deal_ask":
        if d is not None:
            raise DealError("yeh deal pehle se hai")
        if step != 1 or prev:
            raise DealError("pehla step galat")
        with session() as s:
            from sqlalchemy import func
            n = s.scalar(select(func.count()).select_from(Deal).where(
                Deal.tenant_id == tenant, Deal.peer == peer, Deal.created_at >= now() - timedelta(days=1))) or 0
            open_n = s.scalar(select(func.count()).select_from(Deal).where(
                Deal.tenant_id == tenant, Deal.peer == peer, Deal.state.in_(("asked", "quoted")))) or 0
        if n >= ASKS_PER_DAY or open_n >= OPEN_PER_PEER:
            raise DealError("aaj is graahak se bahut saari list aa chuki hain")
        body = check_ask(raw)
        title = body["title"] or ", ".join(f"{i['item']} {i['qty']:g}" for i in body["items"][:3])
        tid = _thread(tenant, peer, env["thread"], "Deal: " + title, "deal", False)
        d = Deal(id=new_id("dl"), tenant_id=tenant, dkey=env["thread"], role="seller", peer=peer, thread_id=tid,
                 title=title[:200], state="asked", ask=body, step=0, last_hash="", history=[], quote={}, draft={},
                 accept={}, ready={}, reserved=[], smap=[], reason="", waiting="", number="", created_at=now(),
                 updated_at=now())
        _record(d, kind, raw, "buyer", f"{len(body['items'])} item ki list aayi")
        d.draft = build_quote(tenant, body["items"])
        with session() as s:
            s.add(d)
        STATS["asks"] += 1
        return _after_ask(tenant, d, card)
    if d is None:
        raise DealError("yeh deal yahan nahi hai")
    if (kind in FROM_SELLER and d.role == "seller") or (kind in FROM_BUYER and d.role == "buyer"):
        raise DealError("yeh step aapki taraf se nahi aa sakta")
    if step != d.step + 1 or prev != d.last_hash:
        raise DealError("step kram se nahi (purana / dohraya hua / badla hua)")
    if _expired(d) and kind != "deal_cancel":
        d.state = "expired"
        _save(d)
        raise DealError("quote ka samay khatam")
    key = (d.state, _event_key(kind, raw))
    if key not in TRANSITIONS:
        raise DealError(f"abhi '{d.state}' hai; '{kind}' nahi ho sakta")
    who = idn.label(card)
    from . import dealpay
    if kind in dealpay.STEP_KINDS:                       # v0.106 money / cancel request / check / dispute steps
        return dealpay.receive_step(tenant, d, kind, raw, who, key)
    if kind == "deal_quote":
        q = check_quote(raw)
        if q["v"] <= int((d.quote or {}).get("v") or 0):
            raise DealError("purana quote")
        _record(d, kind, raw, "seller", f"Quote v{q['v']}: {inr(q['total'])}")
        d.quote, d.state, d.waiting = q, "quoted", "confirm"
        _save(d)
        parts = []
        if q["short"]:
            parts.append("kam maal: " + ", ".join(q["short"]))
        if q["missing"]:
            parts.append("nahi hai: " + ", ".join(q["missing"]))
        subs = [l["item"] for l in q["lines"] if l["status"] == "substitute"]
        if subs:
            parts.append("badle mein: " + ", ".join(subs))
        say = f"{who} ka quote: {inr(q['total'])} (GST {inr(q['gst'])} milakar)" + (f"; {'; '.join(parts)}" if parts else "") + \
              f". Haan dabaane par hi order jaayega."
        _publish(tenant, d, "confirm", say, f"Deal quote aaya {who} se: {inr(q['total'])}. " + "; ".join(parts) +
                 f" Owner card par 'Haan, confirm' dabaye tabhi order jaata hai (tum confirm nahi kar sakte). deal_id={d.id}.")
        return {"ok": True, "status": "quoted"}
    if kind == "deal_confirm":
        q = d.quote or {}
        try:
            v = int(raw.get("v"))
        except (TypeError, ValueError):
            raise DealError("version galat")
        if v != int(q.get("v") or 0) or str(raw.get("quote_hash") or "") != quote_hash(q):
            raise DealError("yeh quote ab lagoo nahi (badal gaya)")
        opts = list(((q.get("terms") or {}).get("pay_options")) or ["upi", "counter"])
        choice = raw.get("pay") if raw.get("pay") in opts else ("upi" if "upi" in opts else opts[0])
        _record(d, kind, raw, "buyer", f"Quote v{v} confirm, {inr(q.get('total'))}, " +
                {"upi": "UPI se", "counter": "counter par", "khata": "khata par"}.get(choice, choice))
        d.state, d.waiting = "confirmed", "accept"
        d.pay = dict(d.pay or {}, choice=choice)
        d.ask = dict(d.ask or {}, mode=raw.get("mode") if raw.get("mode") in ("pickup", "delivery") else (d.ask or {}).get("mode"))
        _save(d)
        STATS["confirms"] += 1
        return _after_confirm(tenant, d, who)
    if kind == "deal_accept":
        body = {"number": policy.clean_name(raw.get("number"), 30), "ready_by": "", "note": policy.quote(raw.get("note"), 300)}
        rb = _when(raw.get("ready_by"))
        if rb and now() - timedelta(hours=1) < rb < now() + timedelta(days=60):
            body["ready_by"] = rb.isoformat()
        _record(d, kind, raw, "seller", f"Order {body['number']} pakka")
        d.state, d.number, d.accept, d.waiting = "accepted", body["number"], body, ""
        from . import dealpay
        d.pay, d.accepted_at = dealpay.buyer_at_accept(tenant, d, raw), now()
        _save(d)
        when = ""
        if body["ready_by"]:
            try:
                from ..life import fmt_local, tz_of
                when = " · taiyaar hoga " + fmt_local(datetime.fromisoformat(body["ready_by"]), tz_of(tenant))
            except Exception:
                when = ""
        _publish(tenant, d, "info", f"{who} ne order pakka kiya: {body['number']}{when}.")
        return {"ok": True, "status": "accepted"}
    if kind == "deal_reject":
        reason = policy.quote(raw.get("reason"), 300)
        _record(d, kind, raw, "seller", "Mana: " + reason)
        d.state, d.reason, d.waiting = "rejected", reason, ""
        _save(d)
        _publish(tenant, d, "info", f"{who} ne mana kiya: {reason}")
        return {"ok": True, "status": "rejected"}
    if kind == "deal_status":
        st = raw.get("status")
        note = policy.quote(raw.get("note"), 300)
        _record(d, kind, raw, "seller", {"packing": "Packing shuru", "delayed": "Der hogi", "handed_over": "Saamaan mil gaya"}[st]
                + (f": {note}" if note else ""))
        d.state = TRANSITIONS[key]
        if st == "handed_over":
            d.handed_at, d.waiting = now(), "check"
        _save(d)
        if st == "handed_over":
            hrs = int(((d.quote or {}).get("terms") or {}).get("dispute_hours") or 48)
            _publish(tenant, d, "check", f"{who}: order {d.number} aapko de diya gaya. Sab sahi mila? Kuch kam / kharab ho to "
                     f"{hrs} ghante ke andar bataiye." + (f" {note}" if note else ""))
            return {"ok": True, "status": d.state}
        say = {"packing": f"{who}: aapka order pack ho raha hai.", "delayed": f"{who}: order mein thodi der hogi.",
               "handed_over": f"{who}: order {d.number} aapko de diya gaya. ✓"}[st] + (f" {note}" if note else "")
        _publish(tenant, d, "info", say)
        return {"ok": True, "status": d.state}
    if kind == "deal_ready":
        rdy = check_ready(raw)
        _record(d, kind, raw, "seller", "Parcel taiyaar")
        d.state, d.ready = "ready", rdy
        _save(d)
        try:
            from ..life import fmt_local, tz_of
            until = fmt_local(datetime.fromisoformat(rdy["until"]), tz_of(tenant))
        except Exception:
            until = rdy["until"]
        _publish(tenant, d, "ready", f"Parcel taiyaar: {rdy['place']}. Counter par code {rdy['code']} dikhaiye ({until} tak).",
                 f"{who}: order {d.number} ka parcel taiyaar hai, {rdy['place']}, {until} tak. Pickup code card par hai "
                 f"(code sirf owner ko batao, kisi aur Hoola ko kabhi nahi). deal_id={d.id}.")
        return {"ok": True, "status": "ready"}
    if kind == "deal_cancel":
        reason = policy.quote(raw.get("reason"), 300)
        held = list(d.reserved or [])
        _record(d, kind, raw, "seller" if d.role == "buyer" else "buyer", "Cancel: " + reason)
        d.state, d.reason, d.waiting = "cancelled", reason, ""
        _save(d)
        if d.role == "seller":
            _release(tenant, held)
            from . import pick as pk
            pk.notify(tenant, d.id, "cancelled")
        from . import dealpay
        dealpay.after_cancel(tenant, d.id)
        d = _get(tenant, d.id)
        _publish(tenant, d, "info", f"{who} ne deal cancel ki: {reason}" +
                 (f" {d.pay.get('refund_say')}" if (d.pay or {}).get("refund_say") else ""))
        return {"ok": True, "status": "cancelled"}
    raise DealError("unknown")


def _after_ask(tenant: str, d: Deal, card: dict[str, Any]) -> dict[str, Any]:
    """A list arrived: quote by itself only inside her written rule, else her card with the quote already worked out."""
    q = d.draft or {}
    who = idn.label(card)
    if _may_auto_quote(tenant, d.peer, q):
        STATS["auto_quotes"] += 1
        try:
            send_quote(tenant, d.id, by="rules")
            return {"ok": True, "status": "quoted"}
        except DealError:
            pass                                       # fall through to asking her
    d = _get(tenant, d.id)
    d.waiting = "quote"
    _save(d)
    gaps = []
    if q.get("missing"):
        gaps.append("nahi hai: " + ", ".join(q["missing"]))
    if q.get("short"):
        gaps.append("kam hai: " + ", ".join(q["short"]))
    if q.get("unknown"):
        gaps.append("stock / daam pata nahi: " + ", ".join(q["unknown"]))
    say = f"{who} ne {len((d.ask or {}).get('items') or [])} item ki list bheji. Quote taiyaar: {inr(q.get('total'))}" + \
          (f" ({'; '.join(gaps)})" if gaps else "") + ". Bhejun?"
    _publish(tenant, d, "quote", say, f"Deal list aayi {who} se. Code ne aapke stock se quote banaya: {inr(q.get('total'))}. "
             + "; ".join(gaps) + f" Owner se poochho: quote bhejein / daam badlein / mana. deal_id={d.id}.")
    return {"ok": True, "status": "asked"}


def _after_confirm(tenant: str, d: Deal, who: str) -> dict[str, Any]:
    so = selling(tenant)
    q = d.quote or {}
    if so["auto_accept"] and 0 < float(q.get("total") or 0) <= so["auto_accept_max"] and not _stock_holds(tenant, q, d.smap or []):
        try:
            accept(tenant, d.id, by="rules")
            STATS["auto_accepts"] += 1
            return {"ok": True, "status": "accepted"}
        except DealError:
            pass
    _publish(tenant, d, "accept", f"{who} ne {inr(q.get('total'))} ka order confirm kiya. Pakka karein?",
             f"Deal confirm aaya {who} se ({inr(q.get('total'))}). Owner 'Pakka' dabaye to order number + pick list banegi. deal_id={d.id}.")
    return {"ok": True, "status": "confirmed"}


# ---------------------------------------------------------------- background
def tick() -> int:
    """Quotes past their time become 'expired' on this side (both sides run this)."""
    n = 0
    with session() as s:
        rows = s.scalars(select(Deal).where(Deal.state == "quoted")).all()
        due = []
        for d in rows:
            vu = _when((d.quote or {}).get("valid_until"))
            if vu is not None and vu < now():
                d.state, d.waiting, d.updated_at = "expired", "", now()
                s.flush()
                s.expunge(d)
                due.append(d)
    for d in due:
        try:
            _publish(d.tenant_id, d, "info", f"{_say_who(d.peer)} wale quote ka samay khatam ho gaya." +
                     (" Zaroorat ho to dobara list bhejiye." if d.role == "buyer" else ""))
        except Exception:
            pass
        n += 1
    # v0.105 a ready parcel nobody collected in time: one reminder each side (the seller can send a new code / time)
    late = []
    with session() as s:
        for d in s.scalars(select(Deal).where(Deal.state == "ready")).all():
            rdy = dict(d.ready or {})
            until = _when(rdy.get("until"))
            if until is not None and until < now() and not rdy.get("late_told"):
                rdy["late_told"] = True
                d.ready, d.updated_at = rdy, now()
                s.flush()
                s.expunge(d)
                late.append(d)
    for d in late:
        try:
            if d.role == "seller":
                _publish(d.tenant_id, d, "pack", f"Order {d.number} ka parcel samay par nahi liya gaya. Naya samay / code bhejun "
                         "(Parcel taiyaar dobara) ya cancel karein?")
            else:
                _publish(d.tenant_id, d, "info", f"{_say_who(d.peer)}: order {d.number} lene ka samay nikal gaya. Dukaan se "
                         "naya samay maangiye.")
        except Exception:
            pass
        n += 1
    return n


# ---------------------------------------------------------------- v0.105 price list for other avatars, stock import, low stock
_CAT_ASKS: dict[str, list[float]] = {}


def catalog(tenant: str, show_stock: bool | None = None) -> list[dict[str, Any]]:
    """What another avatar may see of her list: name, unit, price (+ bulk prices, GST); stock only if she allows it."""
    so = selling(tenant)
    show = so.get("show_stock") if show_stock is None else show_stock
    out = []
    for it in _catalog_items(tenant):
        if not it.get("public", True):
            continue
        row = {"item": it["name"], "unit": it["unit"], "price": it["price"], "gst_pct": it["gst_pct"],
               "tiers": it.get("tiers") or []}
        if show:
            row["in_stock"] = None if it["free"] is None else it["free"] > 0
        out.append(row)
    return out[:300]


def _on_catalog(tenant: str, peer: str, card: dict[str, Any], link) -> dict[str, Any]:
    import time
    so = selling(tenant)
    if so.get("catalog") == "off":
        return {"ok": False, "status": "refused", "reason": "Yeh dukaan apni price list share nahi karti."}
    asks = [t for t in _CAT_ASKS.get(f"{tenant}|{peer}", []) if time.time() - t < 3600]
    if len(asks) >= 20:
        return {"ok": False, "status": "refused", "reason": "bahut baar maangi gayi"}
    _CAT_ASKS[f"{tenant}|{peer}"] = asks + [time.time()]
    return {"ok": True, "status": "catalog", "items": catalog(tenant)}


def catalog_of(tenant: str, to: str) -> dict[str, Any]:
    """The buyer's avatar asks a seller's avatar for its price list (to make a list from it)."""
    peer = resolve(tenant, to)
    l = _link(tenant, peer)
    if l is None or l.status != "active":
        from . import business
        try:
            peer = business.contact(tenant, to)
        except NetError:
            raise DealError(f"{_say_who(peer)} se juda nahi hai.")
    res = _post(tenant, peer, "deal_catalog", {}, new_id("tk"))
    if not res.get("ok"):
        raise DealError(res.get("reason") or "Price list nahi mili.")
    items = []
    for x in (res.get("items") or [])[:300]:
        if isinstance(x, dict) and x.get("item"):
            p = _num(x.get("price"), 0, MAX_TOTAL)
            items.append({"item": policy.clean_name(x.get("item"), 120), "unit": policy.clean_name(x.get("unit"), 20),
                          "price": p, "gst_pct": _num(x.get("gst_pct"), 0, 40) or 0.0,
                          **({"in_stock": bool(x["in_stock"])} if isinstance(x.get("in_stock"), bool) else {})})
    return {"from": _say_who(peer), "items": items,
            "say": "Yeh unki price list hai (jaankari); quote list bhejne par unke stock se banta hai."}


_COLS = {"name": ("name", "item", "naam", "product", "cheez", "particulars", "description", "item name"),
         "price": ("price", "rate", "daam", "mrp", "selling price", "sale price", "unit price"),
         "qty": ("qty", "quantity", "stock", "maatra", "closing", "closing stock", "balance"),
         "unit": ("unit", "uom", "pack"), "gst_pct": ("gst", "gst %", "gst%", "tax", "tax %", "gst rate")}


def import_stock(tenant: str, name: str, data: bytes) -> dict[str, Any]:
    """Her stock sheet (CSV / Excel: name, price, qty, unit, GST in any common header words) into her stock list."""
    import csv
    import io
    rows: list[list[Any]] = []
    if name.lower().endswith((".xlsx", ".xlsm")):
        from openpyxl import load_workbook
        ws = load_workbook(io.BytesIO(data), read_only=True, data_only=True).active
        rows = [list(r) for r in ws.iter_rows(values_only=True)][:5001]
    else:
        text = data.decode("utf-8-sig", errors="replace")
        rows = list(csv.reader(io.StringIO(text)))[:5001]
    head_i, cols = -1, {}
    for i, r in enumerate(rows[:15]):
        cells = [norm(c) for c in r]
        found = {k: cells.index(w) for k, ws_ in _COLS.items() for w in ws_ if w in cells}
        if "name" in found and ("price" in found or "qty" in found):
            head_i, cols = i, found
            break
    if head_i < 0:
        raise DealError("Sheet mein 'naam' aur 'daam' ya 'stock' wale column nahi mile (pehli line mein naam likhiye).")
    items, bad = [], 0
    for r in rows[head_i + 1:]:
        def cell(k):
            j = cols.get(k)
            return r[j] if j is not None and j < len(r) else None
        nm = str(cell("name") or "").strip()
        if not nm:
            continue
        it: dict[str, Any] = {"name": nm}
        for k in ("price", "qty"):
            v = cell(k)
            if v not in (None, ""):
                n = _num(str(v).replace(",", "").replace("₹", "").strip(), 0, MAX_TOTAL)
                if n is None:
                    bad += 1
                    continue
                it[k] = n
        g = cell("gst_pct")
        if g not in (None, ""):
            gv = _num(str(g).replace("%", "").strip(), 0, 40)
            if gv is not None and gv in GST_RATES:
                it["gst_pct"] = gv
        if cell("unit"):
            it["unit"] = str(cell("unit"))[:30]
        items.append(it)
    if not items:
        raise DealError("Sheet mein koi item nahi mila.")
    out = set_stock(tenant, items)
    return {"saved": out["saved"], "skipped_cells": bad, "say": f"{out['saved']} item stock mein aa gaye."}


def low_stock(tenant: str, ids: list[str] | None = None) -> list[dict[str, Any]]:
    """Items at or below her 'tell me when' level (only the given ones, or all)."""
    with session() as s:
        q = select(StockItem).where(StockItem.tenant_id == tenant, StockItem.low_at.is_not(None))
        rows = s.scalars(q).all()
        out = []
        for it in rows:
            if ids is not None and it.id not in ids:
                continue
            v = stock_view(it)
            if v["low"]:
                out.append(v)
    return out


def sync_catalog(tenant: str, catalog_rows: list[dict[str, Any]]) -> None:
    """v0.33 price list -> stock list (price / unit only; counted quantities stay as they are)."""
    items = [{"name": c["item"], "price": c["price"], **({"unit": c["unit"]} if c.get("unit") else {})}
             for c in catalog_rows if c.get("item") and c.get("price") is not None]
    if items:
        set_stock(tenant, items)


for _k in DEAL_KINDS:
    EXTRA_KINDS[_k] = receive
_ = (Text,)
