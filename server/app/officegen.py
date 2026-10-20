"""v0.51 The avatar makes REAL Office files (python-docx / openpyxl / python-pptx, all MIT/BSD): quotations, reports,
letters, price lists, attendance sheets, decks. Content comes from the avatar as simple structure:

  docx  text: '# Heading', '## Sub', '- bullet', '1. numbered', '| a | b |' table rows, blank line = new paragraph
  xlsx  sheets: {"Sheet name": [[header...], [row...], ...]}; a cell starting with '=' is a formula
  pptx  slides: [{"title": "...", "bullets": ["...", ...], "notes": "..."}]
Returns bytes; the caller stores them as the owner's file.
"""
from __future__ import annotations

import io
import re
from typing import Any

MAX_ROWS = 5000
MAX_SLIDES = 40


def make_docx(title: str, text: str) -> bytes:
    from docx import Document
    from docx.shared import Pt
    d = Document()
    st = d.styles["Normal"]
    st.font.name = "Arial"
    st.font.size = Pt(11)
    if title:
        d.add_heading(title[:200], 0)
    table_rows: list[list[str]] = []

    def flush_table():
        if not table_rows:
            return
        cols = max(len(r) for r in table_rows)
        t = d.add_table(rows=len(table_rows), cols=cols)
        t.style = "Table Grid"
        for i, r in enumerate(table_rows):
            for j in range(cols):
                t.cell(i, j).text = r[j] if j < len(r) else ""
                if i == 0:
                    for p in t.cell(i, j).paragraphs:
                        for run in p.runs:
                            run.bold = True
        table_rows.clear()

    for raw in (text or "").splitlines():
        line = raw.rstrip()
        if line.strip().startswith("|") and line.strip().endswith("|"):
            cells = [c.strip() for c in line.strip().strip("|").split("|")]
            if all(re.fullmatch(r":?-{2,}:?", c) for c in cells if c):
                continue                          # markdown separator row
            table_rows.append(cells)
            continue
        flush_table()
        s = line.strip()
        if not s:
            continue
        if s.startswith("### "):
            d.add_heading(s[4:], 3)
        elif s.startswith("## "):
            d.add_heading(s[3:], 2)
        elif s.startswith("# "):
            d.add_heading(s[2:], 1)
        elif re.match(r"^[-*•] ", s):
            d.add_paragraph(s[2:], style="List Bullet")
        elif re.match(r"^\d+[.)] ", s):
            d.add_paragraph(re.sub(r"^\d+[.)] ", "", s), style="List Number")
        else:
            p = d.add_paragraph()
            for part in re.split(r"(\*\*[^*]+\*\*)", s):
                if part.startswith("**") and part.endswith("**"):
                    p.add_run(part[2:-2]).bold = True
                elif part:
                    p.add_run(part)
    flush_table()
    b = io.BytesIO()
    d.save(b)
    return b.getvalue()


def make_xlsx(title: str, sheets: dict[str, Any]) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    wb = Workbook()
    wb.remove(wb.active)
    if not isinstance(sheets, dict) or not sheets:
        raise ValueError("sheets chahiye: {'Sheet1': [[...], ...]}")
    for name, rows in list(sheets.items())[:20]:
        ws = wb.create_sheet(re.sub(r"[\[\]:*?/\\\\]", " ", str(name))[:31] or "Sheet")
        rows = [r if isinstance(r, list) else [r] for r in (rows or [])][:MAX_ROWS]
        for r in rows:
            ws.append([c if not isinstance(c, (dict, list)) else str(c) for c in r[:200]])
        if rows:
            for c in ws[1]:
                c.font = Font(bold=True, color="FFFFFF")
                c.fill = PatternFill("solid", fgColor="5B3FD6")
                c.alignment = Alignment(vertical="center")
            ws.freeze_panes = "A2"
            for i in range(1, max(len(r) for r in rows) + 1):
                width = max((len(str(r[i - 1])) for r in rows if len(r) >= i), default=8)
                ws.column_dimensions[get_column_letter(i)].width = min(60, max(8, width + 2))
            ws.auto_filter.ref = ws.dimensions
    b = io.BytesIO()
    wb.properties.title = title[:200]
    wb.save(b)
    return b.getvalue()


def make_pptx(title: str, slides: list[dict[str, Any]]) -> bytes:
    from pptx import Presentation
    from pptx.util import Pt
    p = Presentation()
    p.slide_width, p.slide_height = 12192000, 6858000          # 16:9
    if title:
        s = p.slides.add_slide(p.slide_layouts[0])
        s.shapes.title.text = title[:200]
        if len(s.placeholders) > 1:
            s.placeholders[1].text = ""
    for sl in (slides or [])[:MAX_SLIDES]:
        if not isinstance(sl, dict):
            sl = {"title": str(sl)}
        s = p.slides.add_slide(p.slide_layouts[1])
        s.shapes.title.text = str(sl.get("title") or "")[:200]
        body = s.placeholders[1].text_frame
        body.clear()
        for i, bullet in enumerate([str(x) for x in (sl.get("bullets") or [])][:12]):
            para = body.paragraphs[0] if i == 0 else body.add_paragraph()
            para.text = bullet[:300]
            para.font.size = Pt(22)
        if sl.get("notes"):
            s.notes_slide.notes_text_frame.text = str(sl["notes"])[:3000]
    b = io.BytesIO()
    p.save(b)
    return b.getvalue()
