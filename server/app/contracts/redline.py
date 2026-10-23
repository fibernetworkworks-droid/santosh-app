"""Redline = a NEW Word copy with tracked changes (w:ins / w:del, author 'Hoola (draft)'): the other side's Word shows them as
suggestions to accept / reject. The original file is only read, never written.

Edits (from the risk review or her own words):
  {"op": "replace", "find": "ninety (90) days", "replace": "forty five (45) days", "anchor": "<start of the clause>"}
  {"op": "delete", "find": "<sentence>"}
  {"op": "insert", "text": "<new clause>", "anchor": "<clause it goes after>"}   (no anchor = at the end)
  {"op": "replace_para", "text": "<new wording>", "anchor": "<the clause>"}         (old paragraph struck, new one inserted)
"""
from __future__ import annotations

import copy
import io
import re
from datetime import datetime, timezone
from typing import Any

AUTHOR = "Hoola (draft)"


def _norm(s: str) -> str:
    return re.sub(r"\s+", " ", (s or "")).strip().lower()


class _Ids:
    def __init__(self):
        self.n = 100

    def next(self) -> str:
        self.n += 1
        return str(self.n)


def _all_paragraphs(doc) -> list:
    out = list(doc.paragraphs)
    for t in doc.tables:
        for row in t.rows:
            for cell in row.cells:
                out.extend(cell.paragraphs)
    return out


def _rpr_of(p):
    from docx.oxml.ns import qn
    for r in p._p.iter(qn("w:r")):
        rpr = r.find(qn("w:rPr"))
        return copy.deepcopy(rpr) if rpr is not None else None
    return None


def _run(text: str, rpr, deleted: bool = False):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    r = OxmlElement("w:r")
    if rpr is not None:
        r.append(copy.deepcopy(rpr))
    t = OxmlElement("w:delText" if deleted else "w:t")
    t.set(qn("xml:space"), "preserve")
    t.text = text
    r.append(t)
    return r


def _mark(tag: str, ids: _Ids, when: str):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    e = OxmlElement(tag)
    e.set(qn("w:id"), ids.next())
    e.set(qn("w:author"), AUTHOR)
    e.set(qn("w:date"), when)
    return e


def _clear_runs(p) -> None:
    from docx.oxml.ns import qn
    for child in list(p._p):
        if child.tag in (qn("w:r"), qn("w:ins"), qn("w:del"), qn("w:hyperlink"), qn("w:smartTag"), qn("w:proofErr")):
            p._p.remove(child)


def _rebuild(p, spans: list[tuple[int, int, str, str]], ids: _Ids, when: str) -> None:
    """spans: (start, end, op, new) on p.text, non-overlapping, sorted."""
    text = p.text
    rpr = _rpr_of(p)
    _clear_runs(p)
    pos = 0
    for start, end, op, new in spans:
        if start > pos:
            p._p.append(_run(text[pos:start], rpr))
        if end > start:
            d = _mark("w:del", ids, when)
            d.append(_run(text[start:end], rpr, deleted=True))
            p._p.append(d)
        if op == "replace" and new:
            i = _mark("w:ins", ids, when)
            i.append(_run(new, rpr))
            p._p.append(i)
        pos = end
    if pos < len(text):
        p._p.append(_run(text[pos:], rpr))


def _insert_after(p_el, text: str, rpr, ids: _Ids, when: str, ppr=None):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    np_ = OxmlElement("w:p")
    pp = copy.deepcopy(ppr) if ppr is not None else OxmlElement("w:pPr")
    # the paragraph mark itself is an insertion too, so "reject" removes the whole new paragraph
    prpr = pp.find(qn("w:rPr"))
    if prpr is None:
        prpr = OxmlElement("w:rPr")
        pp.append(prpr)
    prpr.append(_mark("w:ins", ids, when))
    np_.append(pp)
    ins = _mark("w:ins", ids, when)
    ins.append(_run(text, rpr))
    np_.append(ins)
    p_el.addnext(np_)
    return np_


def _find_para(paras, anchor: str, find: str = "") -> int | None:
    a = _norm(anchor)[:60]
    f = _norm(find)
    for i, p in enumerate(paras):
        t = _norm(p.text)
        if not t:
            continue
        if f and f in t and (not a or a[:30] in t):
            return i
    if f:
        for i, p in enumerate(paras):
            if f in _norm(p.text):
                return i
    if a:
        for i, p in enumerate(paras):
            t = _norm(p.text)
            if t and (a[:40] in t or (len(t) > 20 and t[:40] in a)):
                return i
    return None


def build(source_docx: bytes | None, text: str, edits: list[dict[str, Any]], title: str = "") -> tuple[bytes, dict[str, Any]]:
    from docx import Document
    from docx.shared import Pt
    when = datetime.now(timezone.utc).replace(microsecond=0).isoformat().replace("+00:00", "Z")
    ids = _Ids()
    if source_docx:
        doc = Document(io.BytesIO(source_docx))
    else:
        from .rules import paragraphs
        doc = Document()
        doc.styles["Normal"].font.name = "Arial"
        doc.styles["Normal"].font.size = Pt(11)
        for para in paragraphs(text):
            doc.add_paragraph(para)
    paras = _all_paragraphs(doc)
    applied, skipped = [], []
    per_para: dict[int, list[tuple[int, int, str, str, dict]]] = {}
    inserts: list[tuple[int | None, str, dict]] = []
    for e in edits or []:
        op = e.get("op")
        if op in ("replace", "delete"):
            find = e.get("find") or ""
            if not find.strip():
                skipped.append({**e, "why": "kya badalna hai woh text nahi"})
                continue
            i = _find_para(paras, e.get("anchor", ""), find)
            if i is None:
                skipped.append({**e, "why": "contract mein yeh shabd nahi mile"})
                continue
            ptext = paras[i].text
            taken = per_para.setdefault(i, [])
            start = -1
            for m in re.finditer(re.escape(find), ptext):
                if not any(m.start() < en and m.end() > st for st, en, *_ in taken):
                    start = m.start()
                    break
            if start < 0:
                low = ptext.lower().find(find.lower())
                if low >= 0 and not any(low < en and low + len(find) > st for st, en, *_ in taken):
                    start = low
            if start < 0:
                skipped.append({**e, "why": "yeh hissa pehle hi badla ja chuka"})
                continue
            taken.append((start, start + len(find), op, e.get("replace", "") if op == "replace" else "", e))
        elif op == "replace_para":
            i = _find_para(paras, e.get("anchor", ""))
            if i is None or not (e.get("text") or "").strip():
                skipped.append({**e, "why": "clause nahi mili"})
                continue
            per_para.setdefault(i, []).append((0, len(paras[i].text), "delete", "", e))
            inserts.append((i, e["text"], e))
        elif op == "insert":
            if not (e.get("text") or "").strip():
                skipped.append({**e, "why": "nayi clause ka text nahi"})
                continue
            i = _find_para(paras, e.get("anchor", "")) if e.get("anchor") else None
            inserts.append((i, e["text"], e))
    for i, spans in per_para.items():
        full = any(op == "delete" and st == 0 and en == len(paras[i].text) for st, en, op, _, _ in spans)
        if full:
            spans = [s for s in spans if s[0] == 0 and s[1] == len(paras[i].text)][:1]
        spans.sort(key=lambda x: x[0])
        _rebuild(paras[i], [(a, b, c, d) for a, b, c, d, _ in spans], ids, when)
        applied.extend(x[4] for x in spans)
    last = None
    for i, txt, e in inserts:
        if i is not None:
            anchor = paras[i]
        else:
            body = [p for p in doc.paragraphs if p.text.strip()]
            anchor = body[-1] if body else doc.add_paragraph("")
        rpr = _rpr_of(anchor)
        from docx.oxml.ns import qn
        ppr = anchor._p.find(qn("w:pPr"))
        el = _insert_after(anchor._p if (i is not None or last is None) else last, txt, rpr, ids, when, ppr)
        if i is None:
            last = el
        applied.append(e)
    bio = io.BytesIO()
    doc.save(bio)
    return bio.getvalue(), {"applied": len(applied), "skipped": skipped, "changes": ids.n - 100}


def tracked(docx_bytes: bytes) -> dict[str, Any]:
    """What a Word reader would show: inserted / deleted text (used by tests and the 'kya badla' list)."""
    import zipfile
    xml = zipfile.ZipFile(io.BytesIO(docx_bytes)).read("word/document.xml").decode("utf-8")
    ins = re.findall(r"<w:ins [^>]*>(.*?)</w:ins>", xml, re.S)
    dels = re.findall(r"<w:del [^>]*>(.*?)</w:del>", xml, re.S)
    txt = lambda frag: "".join(re.findall(r"<w:(?:t|delText)[^>]*>([^<]*)</w:(?:t|delText)>", frag))  # noqa: E731
    return {"inserted": [t for t in (txt(x) for x in ins) if t], "deleted": [t for t in (txt(x) for x in dels) if t],
            "authors": sorted(set(re.findall(r'w:author="([^"]+)"', xml)))}


_NUMS = re.compile(r"\d[\d,]*(?:\.\d+)?")


def check_edit(e: dict, text: str, allowed_numbers: set[str]) -> str:
    """A model-proposed edit is used only if its 'find' is really in the contract and it adds no number from nowhere."""
    op = e.get("op")
    if op not in ("replace", "delete", "insert", "replace_para"):
        return "op galat"
    if op in ("replace", "delete") and (not e.get("find") or _norm(e["find"]) not in _norm(text)):
        return "find contract mein nahi"
    new = (e.get("replace") or e.get("text") or "")
    for n in _NUMS.findall(new):
        n2 = n.replace(",", "")
        if n2 not in allowed_numbers and n not in text:
            return f"naya number {n} kahin se nahi aaya"
    return ""
