"""Fleet Sahayak: Hoola as a fleet manager (trucks, tempos, cars, buses), end to end.

  1. Plans routes and loads: trip TR-0001 (origin, stops, destination, day), the load (weight / volume / pieces) checked against
     each vehicle's payload, body cbm and body size; the smallest free vehicle with valid papers that takes it is suggested, or a
     split across vehicles; distance from Google Maps when the server has it, else HER number; fuel / toll estimate ONLY from her
     rates and this vehicle's own km/l; long drives get a driver-hours note; inter-state goods trips without a national permit warn.
  2. Tracks fuel and service: fuel log (litres, rupees, odometer, full tank) -> km/l tank-to-tank per vehicle; outliers (km/l far
     below usual, more litres than the tank, odometer backwards, price far from usual) = possible theft / leak, a signal not a
     verdict; fuel-card CSV / Excel import (deduped); service schedule by km / months (starter lines, maker's manual decides);
     tyre log with km run; breakdowns with repeat detection; cost per km; idle days.
  3. Keeps permits up to date: RC, insurance, fitness (2 years up to 8 years, then yearly; ATS), PUC, national / state permit,
     road tax, AIS-140 GPS, speed governor, FASTag note; reminder windows per document, once each; Vahan / mParivahan / eChallan
     only as official steps she follows (no scraping); challans she found are logged.
  4. Waits for HER tap on ANY spend: service, tyres, permit fee, tax, insurance, challan, repair -> `fleet_spend` card with a hash;
     only `decide_spend(by_tap=True)` from the app approves (her fingerprint once a phone key is active). Hoola never pays or buys;
     she pays and taps "Maine kar diya".
"""
from __future__ import annotations

import asyncio
import csv
import hashlib
import io
import json
import logging
import re
import secrets
import statistics
import threading
from datetime import date, datetime, timedelta, timezone
from typing import Any

from sqlalchemy import select

from ..db import now, session
from . import rules as R
from .models import (
    FleetAlert,
    FleetDoc,
    FleetEvent,
    FleetFuel,
    FleetServiceItem,
    FleetSettings,
    FleetSpend,
    FleetTrip,
    FleetVehicle,
)

log = logging.getLogger("santosh.fleet")
_LOCK = threading.RLock()
SPEND_KINDS = {"service": "Service", "tyre": "Tyre", "permit": "Permit fees", "tax": "Road tax", "insurance": "Insurance",
               "fitness": "Fitness / ATS fees", "puc": "PUC", "challan": "Challan", "repair": "Repair", "fastag": "FASTag recharge",
               "gps": "GPS / AIS-140", "other": "Doosra kharcha"}
TRIP_STATUS = {"planned": "Plan", "assigned": "Gaadi tay", "started": "Raste mein", "done": "Poori", "cancelled": "Radd"}
BLOCK_DOCS = ("insurance", "fitness", "national_permit", "state_permit", "rc")   # expired = the vehicle can't legally run
WARN_DOCS = ("puc", "road_tax", "ais140", "speed_governor")


class FleetError(Exception):
    pass


# ---------------------------------------------------------------- helpers
def _j(s: str, default):
    try:
        v = json.loads(s or "")
        return v if isinstance(v, type(default)) else default
    except (TypeError, ValueError):
        return default


def _dump(v) -> str:
    return json.dumps(v, ensure_ascii=False, default=str)


def _aware(d: datetime | None) -> datetime | None:
    if d is None:
        return None
    return d if d.tzinfo else d.replace(tzinfo=timezone.utc)


def clock() -> datetime:
    return now()


def _tz(tenant: str):
    from ..life import tz_of
    return tz_of(tenant)


def today(tenant: str) -> date:
    return clock().astimezone(_tz(tenant)).date()


def _date(v) -> date | None:
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    if not v:
        return None
    s = str(v).strip()[:10]
    for f in ("%Y-%m-%d", "%d-%m-%Y", "%d/%m/%Y", "%d.%m.%Y", "%d-%m-%y", "%d/%m/%y", "%d %b %Y", "%d-%b-%Y", "%d-%b-%y"):
        try:
            return datetime.strptime(s if "%b" not in f else str(v).strip()[:11], f).date()
        except ValueError:
            pass
    return None


def _iso(d: date | None) -> str:
    return d.isoformat() if d else ""


def _num(v) -> float:
    if v is None or v == "":
        return 0.0
    if isinstance(v, (int, float)):
        return float(v)
    s = re.sub(r"[^\d.\-]", "", str(v).replace(",", ""))
    try:
        return float(s) if s not in ("", "-", ".") else 0.0
    except ValueError:
        return 0.0


def _rs(v: float | None) -> str:
    if v is None:
        return ""
    try:
        from ..india.money import inr
        return inr(v)
    except Exception:
        return f"₹{v:,.0f}"


def _publish(tenant: str, kind: str, payload: dict, conv: str = "") -> None:
    from ..events import bus
    try:
        if not conv:
            from ..net.service import home_conv
            conv = home_conv(tenant)
        bus.publish(conv, kind, payload)
    except Exception:
        log.exception("fleet publish failed")


def is_owner(actor: str) -> bool:
    a = actor or "owner"
    return a == "owner" or a.startswith("approved:")


def actor_name(tenant: str, actor: str) -> str:
    if is_owner(actor):
        return "Owner"
    try:
        from ..employee.service import actor_label
        return actor_label(tenant, actor)
    except Exception:
        return actor[:40]


def _st(s, tenant: str) -> FleetSettings:
    r = s.get(FleetSettings, tenant)
    if r is None:
        r = FleetSettings(tenant_id=tenant)
        s.add(r)
        s.flush()
    return r


def _clean(text: str, n: int = 500) -> str:
    """Her / staff words: secrets never kept (card / OTP / long numbers)."""
    t = str(text or "")
    t = re.sub(r"\b(otp|pin|cvv|password|upi pin)\b\s*(?:is|hai|:|=)?\s*\S+", lambda m: f"{m.group(1)} [hataya]", t, flags=re.I)
    t = re.sub(r"\b(?:\d[ -]?){13,19}\b", "[hataya]", t)
    return t.strip()[:n]


# ---------------------------------------------------------------- settings
_SET_KEYS = ("business", "home_state", "diesel_rate", "petrol_rate", "cng_rate", "toll_per_km", "fill_pct", "low_kmpl_pct",
             "idle_days", "avg_speed", "morning_time", "rules")


def _settings_view(r: FleetSettings) -> dict:
    return {"business": r.business, "home_state": r.home_state, "diesel_rate": r.diesel_rate or None, "petrol_rate": r.petrol_rate or None,
            "cng_rate": r.cng_rate or None, "toll_per_km": _j(r.toll_per_km, {}), "fill_pct": r.fill_pct, "low_kmpl_pct": r.low_kmpl_pct,
            "idle_days": r.idle_days, "avg_speed": r.avg_speed or None, "morning_time": r.morning_time, "rules": r.rules}


def settings(tenant: str) -> dict:
    with session() as s:
        return _settings_view(_st(s, tenant))


def save_settings(tenant: str, **kw: Any) -> dict:
    with _LOCK, session() as s:
        r = _st(s, tenant)
        for k in _SET_KEYS:
            if k not in kw or kw[k] is None:
                continue
            v = kw[k]
            if k in ("diesel_rate", "petrol_rate", "cng_rate", "avg_speed"):
                x = _num(v)
                if x < 0 or (k.endswith("_rate") and x > 500) or (k == "avg_speed" and x > 100):
                    raise FleetError(f"{k}: {v} sahi nahi lagta.")
                setattr(r, k, x)
            elif k in ("fill_pct", "low_kmpl_pct"):
                x = _num(v)
                if not 5 <= x <= 100:
                    raise FleetError(f"{k} 5-100 ke beech.")
                setattr(r, k, x)
            elif k == "idle_days":
                r.idle_days = max(1, min(60, int(_num(v))))
            elif k == "toll_per_km":
                d = v if isinstance(v, dict) else {}
                r.toll_per_km = _dump({str(a): _num(b) for a, b in d.items() if _num(b) > 0})
            elif k == "morning_time":
                if not re.fullmatch(r"\d{1,2}:\d{2}", str(v)):
                    raise FleetError("Subah ka samay HH:MM mein.")
                r.morning_time = str(v)
            elif k == "home_state":
                r.home_state = str(v).upper()[:4]
            else:
                setattr(r, k, str(v)[:2000])
        return {"settings": _settings_view(r), "say": "Fleet ki settings save."}


# ---------------------------------------------------------------- vehicles
def _veh(s, tenant: str, ref: str) -> FleetVehicle:
    ref = str(ref or "").strip()
    if not ref:
        rows = s.scalars(select(FleetVehicle).where(FleetVehicle.tenant_id == tenant, FleetVehicle.status != "sold")).all()
        if len(rows) == 1:
            return rows[0]
        raise FleetError("Kaun si gaadi? Number bataiye.")
    v = s.get(FleetVehicle, ref)
    if v is not None and v.tenant_id == tenant:
        return v
    key = re.sub(r"[\s\-./]", "", ref).upper()
    rows = s.scalars(select(FleetVehicle).where(FleetVehicle.tenant_id == tenant)).all()
    for v in rows:
        if v.reg_key == key:
            return v
    tail = [v for v in rows if key and v.reg_key.endswith(key) and len(key) >= 3]
    if len(tail) == 1:
        return tail[0]
    if len(tail) > 1:
        raise FleetError(f"'{ref}' se kai gaadiyan milti hain: " + ", ".join(v.reg for v in tail[:5]))
    raise FleetError(f"Gaadi '{ref}' nahi mili.")


def _odo_now(s, v: FleetVehicle) -> float:
    return v.odo or 0.0


def _kmpl_hist(s, v: FleetVehicle, limit: int = 10) -> list[float]:
    rows = s.scalars(select(FleetFuel).where(FleetFuel.vehicle_id == v.id, FleetFuel.kmpl.is_not(None))
                     .order_by(FleetFuel.odo.desc())).all()
    return [r.kmpl for r in rows[:limit] if r.kmpl]


def _kmpl_of(s, v: FleetVehicle) -> float | None:
    h = _kmpl_hist(s, v)
    return round(statistics.median(h), 2) if h else None


def _doc_view(d: FleetDoc, on: date) -> dict:
    k = R.DOC_KINDS.get(d.kind, R.DOC_KINDS["other"])
    win = _j(d.windows, []) or k.get("windows", [])
    st = R.doc_state(_date(d.expires), on, win) if d.kind != "fastag" else {"state": "note", "days_left": None, "window": None}
    if d.not_needed:
        st = {"state": "not_needed", "days_left": None, "window": None}
    return {"id": d.id, "kind": d.kind, "label": k["label"], "number": d.number, "issued": d.issued, "expires": d.expires,
            "guess": d.expires_guess, "issuer": d.issuer, "file_id": d.file_id, "notes": d.notes, "windows": win, **st,
            "where": k.get("where", ""), "rule": k.get("note", "")}


def _vehicle_view(s, v: FleetVehicle, on: date, full: bool = False) -> dict:
    docs = [_doc_view(d, on) for d in s.scalars(select(FleetDoc).where(FleetDoc.vehicle_id == v.id)).all()]
    bad = [d for d in docs if d["state"] in ("expired", "missing") and d["kind"] in BLOCK_DOCS]
    due = [d for d in docs if d["state"] in ("due", "expired")]
    out = {"id": v.id, "reg": v.reg, "type": v.vtype, "type_label": R.VEHICLE_TYPES.get(v.vtype, {}).get("label", v.vtype),
           "make": v.make, "model": v.model, "fuel": v.fuel, "bs": v.bs, "reg_date": v.reg_date, "capacity_t": v.capacity_t or None,
           "capacity_cbm": v.capacity_cbm or None, "body": _j(v.body, {}), "axles": v.axles or None, "tank_l": v.tank_l or None,
           "owner_name": v.owner_name, "ownership": v.ownership, "national": v.national, "driver": v.driver, "odo": v.odo or None,
           "odo_at": v.odo_at, "status": v.status, "fastag_note": v.fastag_note, "kmpl": _kmpl_of(s, v),
           "papers_ok": not bad, "papers_bad": [d["label"] for d in bad], "papers_due": [f"{d['label']} ({d['expires'] or '?'})" for d in due]}
    if full:
        out["docs"] = sorted(docs, key=lambda d: (d["days_left"] if d["days_left"] is not None else -9999))
        out["service"] = _service_view(s, v, on)
        out["tyres"] = _tyres(s, v)
        out["events"] = [_event_view(e) for e in s.scalars(select(FleetEvent).where(FleetEvent.vehicle_id == v.id)
                                                            .order_by(FleetEvent.created_at.desc())).all()[:30]]
        out["fuel"] = [_fuel_view(f) for f in s.scalars(select(FleetFuel).where(FleetFuel.vehicle_id == v.id)
                                                        .order_by(FleetFuel.day.desc(), FleetFuel.odo.desc())).all()[:20]]
    return out


def save_vehicle(tenant: str, reg: str = "", vehicle: str = "", **kw: Any) -> dict:
    """Add (reg) or change (vehicle) one vehicle. New ones get their usual document list (dates blank until she gives them) and the
    starter service schedule."""
    on = today(tenant)
    with _LOCK, session() as s:
        if vehicle:
            v = _veh(s, tenant, vehicle)
            new = False
        else:
            chk = R.reg_check(reg)
            if not chk["ok"]:
                raise FleetError(chk["problem"])
            same = s.scalars(select(FleetVehicle).where(FleetVehicle.tenant_id == tenant, FleetVehicle.reg_key == chk["key"])).first()
            v, new = (same, False) if same else (FleetVehicle(tenant_id=tenant, reg=chk["reg"], reg_key=chk["key"]), True)
            if new:
                s.add(v)
        vt = kw.get("type") or kw.get("vtype")
        if vt:
            t = vt if vt in R.VEHICLE_TYPES else R.type_of(vt)
            if not t:
                raise FleetError(f"Gaadi ka type '{vt}' samajh nahi aaya (truck / tempo / pickup / bus / taxi / car ...).")
            v.vtype = t
        elif new:
            v.vtype = R.type_of(" ".join(str(kw.get(k) or "") for k in ("make", "model"))) or "truck"
        for k in ("make", "model", "owner_name", "driver", "notes", "fastag_note"):
            if kw.get(k) is not None:
                setattr(v, k, _clean(kw[k], 160 if k != "notes" else 1000))
        if kw.get("fuel"):
            f = str(kw["fuel"]).lower()
            if f not in R.FUELS:
                raise FleetError("Fuel: diesel / petrol / cng / lng / electric / lpg.")
            v.fuel = f
        if kw.get("bs"):
            v.bs = R.bs_of(kw["bs"]) or v.bs
        if kw.get("reg_date"):
            d = _date(kw["reg_date"])
            if not d or d > on:
                raise FleetError("Registration ki taarikh sahi nahi.")
            v.reg_date = d.isoformat()
        for k in ("capacity_t", "capacity_cbm", "tank_l"):
            if kw.get(k) is not None:
                x = _num(kw[k])
                if x < 0 or (k == "capacity_t" and x > 60) or (k == "capacity_cbm" and x > 200) or (k == "tank_l" and x > 1500):
                    raise FleetError(f"{k}: {kw[k]} sahi nahi lagta.")
                setattr(v, k, x)
        if kw.get("axles") is not None:
            v.axles = max(0, min(10, int(_num(kw["axles"]))))
        if kw.get("body") is not None:
            b = kw["body"] if isinstance(kw["body"], dict) else {}
            v.body = _dump({k: _num(b.get(k)) for k in ("l", "w", "h") if _num(b.get(k)) > 0})
        if kw.get("ownership"):
            if kw["ownership"] not in ("own", "attached", "leased"):
                raise FleetError("Ownership: own / attached / leased.")
            v.ownership = kw["ownership"]
        if kw.get("national") is not None:
            v.national = bool(kw["national"])
        if kw.get("status"):
            if kw["status"] not in ("active", "workshop", "sold"):
                raise FleetError("Status: active / workshop / sold.")
            v.status = kw["status"]
        if kw.get("odo") is not None:
            o = _num(kw["odo"])
            if o and v.odo and o < v.odo - 1 and not kw.get("odo_reset"):
                raise FleetError(f"Odometer {o:g} pichhle ({v.odo:g}) se kam: meter badla ho to odo_reset=true ke saath bhejiye.")
            v.odo, v.odo_at = o, on.isoformat()
        s.flush()
        if new:
            for k in R.required_docs(v.vtype, v.national):
                s.add(FleetDoc(tenant_id=tenant, vehicle_id=v.id, kind=k))
            for it in R.SERVICE_STARTERS.get(v.vtype, []):
                s.add(FleetServiceItem(tenant_id=tenant, vehicle_id=v.id, item=it["item"], every_km=it["km"], every_months=it["months"]))
            s.flush()
        elif kw.get("national") and v.vtype in R.VEHICLE_TYPES and R.VEHICLE_TYPES[v.vtype].get("goods"):
            if not s.scalars(select(FleetDoc).where(FleetDoc.vehicle_id == v.id, FleetDoc.kind == "national_permit")).first():
                s.add(FleetDoc(tenant_id=tenant, vehicle_id=v.id, kind="national_permit"))
        view = _vehicle_view(s, v, on, full=True)
    missing = [d["label"] for d in view["docs"] if d["state"] == "missing"]
    say = f"{view['reg']} ({view['type_label']}) " + ("jod di." if new else "update.")
    if missing:
        say += " In kaagazon ki expiry bataiye (ya photo bhejiye): " + ", ".join(missing[:8]) + "."
    return {"vehicle": view, "new": new, "say": say}


def vehicles(tenant: str, status: str = "") -> list[dict]:
    on = today(tenant)
    with session() as s:
        q = select(FleetVehicle).where(FleetVehicle.tenant_id == tenant)
        q = q.where(FleetVehicle.status == status) if status else q.where(FleetVehicle.status != "sold")
        return [_vehicle_view(s, v, on) for v in s.scalars(q.order_by(FleetVehicle.reg)).all()]


def vehicle_view(tenant: str, ref: str) -> dict:
    with session() as s:
        return _vehicle_view(s, _veh(s, tenant, ref), today(tenant), full=True)


def find_vehicle(tenant: str, ref: str) -> dict | None:
    """For the Driver Sahayak: {id, number, kind, kmpl, gvw_kg, capacity_t, papers_ok, papers_bad} of one vehicle by id / number,
    or None. Fleet keeps payload (capacity_t), not gross weight, so gvw_kg is None here (drivers uses her own number for it)."""
    if not str(ref or "").strip():
        return None
    try:
        with session() as s:
            v = _vehicle_view(s, _veh(s, tenant, ref), today(tenant))
    except FleetError:
        return None
    return {"id": v["id"], "number": v["reg"], "kind": v["type"], "kmpl": v.get("kmpl"), "gvw_kg": None,
            "capacity_t": v.get("capacity_t"), "papers_ok": v.get("papers_ok"), "papers_bad": v.get("papers_bad") or []}


# ---------------------------------------------------------------- documents
def save_doc(tenant: str, vehicle: str, kind: str, expires: str = "", issued: str = "", number: str = "", issuer: str = "",
             file_id: str = "", windows: list | None = None, not_needed: bool | None = None, notes: str = "") -> dict:
    on = today(tenant)
    k = R.doc_kind(kind)
    with _LOCK, session() as s:
        v = _veh(s, tenant, vehicle)
        d = s.scalars(select(FleetDoc).where(FleetDoc.vehicle_id == v.id, FleetDoc.kind == k)).first() if k != "other" else None
        if d is None:
            d = FleetDoc(tenant_id=tenant, vehicle_id=v.id, kind=k)
            s.add(d)
        exp, iss = _date(expires), _date(issued)
        if expires and not exp:
            raise FleetError(f"Expiry '{expires}' taarikh nahi lagti (DD-MM-YYYY).")
        if issued and not iss:
            raise FleetError(f"Issue ki taarikh '{issued}' samajh nahi aayi.")
        if exp and iss and exp <= iss:
            raise FleetError("Expiry issue ki taarikh ke baad honi chahiye.")
        if iss:
            d.issued = iss.isoformat()
        if exp:
            d.expires, d.expires_guess = exp.isoformat(), False
        elif iss and not d.expires:
            g = R.suggest_expiry(k, iss, _date(v.reg_date), v.bs)
            if g:
                d.expires, d.expires_guess = g.isoformat(), True
        if number:
            d.number = _clean(number, 60)
        if issuer:
            d.issuer = _clean(issuer, 120)
        if file_id:
            d.file_id = str(file_id)[:60]
        if windows is not None:
            w = sorted({int(_num(x)) for x in windows if 0 <= _num(x) <= 365}, reverse=True)
            d.windows = _dump(w) if w else ""
        if not_needed is not None:
            d.not_needed = bool(not_needed)
        if notes:
            d.notes = _clean(notes, 500)
        if k == "fastag" and notes:
            v.fastag_note = _clean(notes, 160)
        d.reminded = "{}" if exp else d.reminded
        s.flush()
        view = _doc_view(d, on)
        reg = v.reg
    say = f"{reg}: {view['label']}"
    if view["expires"]:
        say += f" {view['expires']} tak" + (" (andaaza: kaagaz ki asli taarikh bataiye)" if view["guess"] else "") + "."
        if view["state"] == "expired":
            say += " Yeh khatam ho chuka hai: renewal ka kharcha card se."
    else:
        say += " darj (expiry abhi khaali)."
    return {"doc": view, "say": say}


def docs_due(tenant: str, days: int = 45) -> list[dict]:
    on = today(tenant)
    out = []
    with session() as s:
        vs = {v.id: v for v in s.scalars(select(FleetVehicle).where(FleetVehicle.tenant_id == tenant, FleetVehicle.status != "sold")).all()}
        for d in s.scalars(select(FleetDoc).where(FleetDoc.tenant_id == tenant)).all():
            v = vs.get(d.vehicle_id)
            if v is None:
                continue
            x = _doc_view(d, on)
            if x["state"] == "expired" or (x["state"] in ("due", "ok") and x["days_left"] is not None and x["days_left"] <= days) \
                    or (x["state"] == "missing" and d.kind in BLOCK_DOCS):
                out.append({**x, "vehicle": v.reg, "vehicle_id": v.id})
    order = {"expired": 0, "missing": 1}
    return sorted(out, key=lambda x: (order.get(x["state"], 2), x["days_left"] if x["days_left"] is not None else 0))


def official_steps(what: str = "echallan") -> dict:
    if what == "vahan":
        return {"steps": R.VAHAN_STEPS, "say": "Vahan / mParivahan se taarikhein aise dekhiye (Hoola site nahi padhta)."}
    return {"steps": R.ECHALLAN_STEPS, "say": "eChallan aise dekhiye; challan mile to yahan bataiye. Bharna aapka kaam, card se."}


# ---------------------------------------------------------------- fuel
def _fuel_view(f: FleetFuel) -> dict:
    return {"id": f.id, "day": f.day, "litres": f.litres, "amount": f.amount or None, "odo": f.odo or None, "full": f.full,
            "station": f.station, "source": f.source, "kmpl": f.kmpl, "flags": _j(f.flags, []), "by": f.by,
            "rate": round(f.amount / f.litres, 2) if f.litres and f.amount else None}


def _recompute(s, v: FleetVehicle) -> None:
    fills = s.scalars(select(FleetFuel).where(FleetFuel.vehicle_id == v.id)).all()
    ser = {x["id"]: x["kmpl"] for x in R.kmpl_series([{"id": f.id, "odo": f.odo, "litres": f.litres, "full": f.full} for f in fills])}
    for f in fills:
        f.kmpl = ser.get(f.id)


def add_fuel(tenant: str, vehicle: str, litres, amount=None, odo=None, day: str = "", full: bool = True, station: str = "",
             source: str = "manual", ref: str = "", actor: str = "owner", conv: str = "", quiet: bool = False) -> dict:
    lit = _num(litres)
    if lit <= 0 or lit > 2000:
        raise FleetError("Litre sahi nahi (0-2000).")
    amt, o = _num(amount), _num(odo)
    d = _date(day) if day else today(tenant)
    if day and not d:
        raise FleetError(f"Taarikh '{day}' samajh nahi aayi.")
    if d and d > today(tenant):
        raise FleetError("Aage ki taarikh ka diesel nahi.")
    with _LOCK, session() as s:
        v = _veh(s, tenant, vehicle)
        if ref:
            dup = s.scalars(select(FleetFuel).where(FleetFuel.tenant_id == tenant, FleetFuel.ref == str(ref)[:80])).first()
            if dup:
                return {"fuel": _fuel_view(dup), "duplicate": True, "say": f"Yeh entry ({ref}) pehle se hai."}
        same = s.scalars(select(FleetFuel).where(FleetFuel.vehicle_id == v.id, FleetFuel.day == d.isoformat(), FleetFuel.litres == lit,
                                                 FleetFuel.amount == amt)).first()
        if same and (not o or same.odo == o):
            return {"fuel": _fuel_view(same), "duplicate": True, "say": "Yeh diesel entry pehle se hai."}
        prev = s.scalars(select(FleetFuel).where(FleetFuel.vehicle_id == v.id, FleetFuel.odo > 0).order_by(FleetFuel.odo.desc())).first()
        prev_odo = prev.odo if prev else (v.odo or None)
        hist = _kmpl_hist(s, v)
        prices = [r.amount / r.litres for r in s.scalars(select(FleetFuel).where(FleetFuel.tenant_id == tenant, FleetFuel.litres > 0,
                                                                                  FleetFuel.amount > 0)).all()[-30:]]
        f = FleetFuel(tenant_id=tenant, vehicle_id=v.id, day=d.isoformat(), litres=lit, amount=amt, odo=o, full=bool(full),
                      station=_clean(station, 120), source=source if source in ("manual", "card") else "manual", ref=str(ref or "")[:80],
                      by=actor_name(tenant, actor))
        s.add(f)
        s.flush()
        _recompute(s, v)
        flags = R.fuel_flags({"odo": o, "litres": lit, "amount": amt}, prev_odo, hist, f.kmpl, v.tank_l, prices,
                             _st(s, tenant).low_kmpl_pct)
        f.flags = _dump(flags)
        if o and (not v.odo or o >= v.odo):
            v.odo, v.odo_at = o, d.isoformat()
        view, reg = _fuel_view(f), v.reg
    say = f"{reg}: {lit:g} L" + (f" {_rs(amt)}" if amt else "") + (f", odo {o:,.0f}" if o else "")
    if view["kmpl"]:
        say += f", average {view['kmpl']:g} km/l"
    elif not o:
        say += " (odometer bina average nahi banta)"
    say += "."
    if flags:
        say += " Dhyaan: " + " ".join(x["say"] for x in flags)
        if not quiet:
            _publish(tenant, "fleet_alert", {"kind": "fuel", "vehicle": reg, "fuel": view, "say": say}, conv)
    return {"fuel": view, "flags": flags, "say": say}


_FUEL_COLS = {
    "vehicle": ("vehicle", "vehicle no", "vehicle number", "registration", "reg no", "veh no", "truck no", "gaadi", "vehicle reg no"),
    "day": ("date", "txn date", "transaction date", "date time", "trans date", "bill date"),
    "litres": ("litres", "liters", "ltr", "qty", "quantity", "volume", "litre", "qty (ltr)", "quantity (ltrs)"),
    "amount": ("amount", "amt", "value", "txn amount", "transaction amount", "total", "amount (rs)", "rs"),
    "odo": ("odometer", "odo", "odometer reading", "km reading", "kms", "meter reading"),
    "station": ("station", "outlet", "ro name", "pump", "merchant", "location", "retail outlet"),
    "ref": ("txn id", "transaction id", "txn no", "reference", "ref no", "bill no", "invoice no", "receipt no"),
    "product": ("product", "fuel", "product name", "fuel type"),
}


def _norm(x) -> str:
    return re.sub(r"\s+", " ", str(x or "").strip().lower().replace("_", " ").replace(".", ""))


def _rows(name: str, data: bytes) -> list[list]:
    if name.lower().endswith((".xlsx", ".xlsm")):
        import openpyxl
        wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        return [list(r) for r in wb.worksheets[0].iter_rows(values_only=True)]
    text = data.decode("utf-8-sig", errors="replace")
    return [r for r in csv.reader(io.StringIO(text))]


def _load_file(tenant: str, file_id: str) -> tuple[str, bytes]:
    from ..mailbox import service as ms
    try:
        name, _m, data = ms.load_files(tenant, [{"type": "file", "file_id": file_id}])[0]
    except Exception as e:
        raise FleetError(f"File {file_id} nahi mili: {e}")
    return name, data


def import_fuel(tenant: str, file_id: str = "", data: bytes | None = None, name: str = "fuel.csv", vehicle: str = "",
                actor: str = "owner", conv: str = "") -> dict:
    """Fuel card statement (IOCL XTRAPOWER / BPCL SmartFleet / HPCL DriveTrack style exports, or any sheet): header row found in
    the first 15 rows by column families; non-fuel products (lubes, AdBlue / DEF) skipped; repeats skipped."""
    if file_id:
        name, data = _load_file(tenant, file_id)
    rows = _rows(name, data or b"")
    hi, cols = None, {}
    for i, r in enumerate(rows[:15]):
        h = [_norm(x) for x in r]
        found = {}
        for key, words in _FUEL_COLS.items():
            for j, x in enumerate(h):
                if x in words and key not in found:
                    found[key] = j
        if "litres" in found and "day" in found:
            hi, cols = i, found
            break
    if hi is None:
        raise FleetError("Sheet mein 'Date' aur 'Litres / Quantity' wale column nahi mile.")
    if "vehicle" not in cols and not vehicle:
        raise FleetError("Sheet mein gaadi number ka column nahi: kaun si gaadi ki statement hai, bataiye.")
    added, dup, bad, flagged = 0, 0, [], []
    for n, r in enumerate(rows[hi + 1:], start=hi + 2):
        def g(k):
            j = cols.get(k)
            return r[j] if j is not None and j < len(r) else None
        if not any(x not in (None, "") for x in r):
            continue
        prod = _norm(g("product"))
        if prod and not any(w in prod for w in ("diesel", "hsd", "petrol", "ms", "cng", "speed", "power", "xtra", "fuel")):
            continue
        lit = _num(g("litres"))
        if lit <= 0:
            continue
        veh = str(g("vehicle") or vehicle or "")
        d = _date(g("day"))
        if not d:
            bad.append(f"line {n}: taarikh")
            continue
        try:
            out = add_fuel(tenant, veh, lit, g("amount"), g("odo"), d.isoformat(), True, str(g("station") or ""), "card",
                           str(g("ref") or ""), actor, conv, quiet=True)
        except FleetError as e:
            bad.append(f"line {n}: {e}")
            continue
        if out.get("duplicate"):
            dup += 1
            continue
        added += 1
        if out["flags"]:
            flagged.append({"line": n, "vehicle": veh, "say": " ".join(x["say"] for x in out["flags"])})
    say = f"Fuel statement: {added} entry jodi" + (f", {dup} pehle se thi" if dup else "") + (f", {len(bad)} nahi padh paya" if bad else "") + "."
    if flagged:
        say += f" {len(flagged)} entry par dhyaan dijiye."
        _publish(tenant, "fleet_alert", {"kind": "fuel_import", "flagged": flagged[:10], "say": say}, conv)
    return {"added": added, "duplicates": dup, "problems": bad[:20], "flagged": flagged[:20], "say": say}


def fuel_report(tenant: str, vehicle: str = "", days: int = 30) -> dict:
    on = today(tenant)
    since = (on - timedelta(days=days)).isoformat()
    out = []
    with session() as s:
        q = select(FleetVehicle).where(FleetVehicle.tenant_id == tenant, FleetVehicle.status != "sold")
        vs = [_veh(s, tenant, vehicle)] if vehicle else s.scalars(q).all()
        for v in vs:
            rows = s.scalars(select(FleetFuel).where(FleetFuel.vehicle_id == v.id, FleetFuel.day >= since)).all()
            if not rows:
                continue
            k = [r.kmpl for r in rows if r.kmpl]
            out.append({"vehicle": v.reg, "litres": round(sum(r.litres for r in rows), 1), "amount": round(sum(r.amount for r in rows)),
                        "fills": len(rows), "kmpl": round(statistics.median(k), 2) if k else None, "usual_kmpl": _kmpl_of(s, v),
                        "flags": [f for r in rows for f in _j(r.flags, [])][:5]})
    return {"days": days, "vehicles": out}


# ---------------------------------------------------------------- service schedule / events
def _service_view(s, v: FleetVehicle, on: date) -> list[dict]:
    out = []
    for it in s.scalars(select(FleetServiceItem).where(FleetServiceItem.vehicle_id == v.id)).all():
        st = R.service_due({"km": it.every_km, "months": it.every_months, "last_km": it.last_km, "last_on": _date(it.last_on)},
                           v.odo or None, on)
        out.append({"id": it.id, "item": it.item, "every_km": it.every_km or None, "every_months": it.every_months or None,
                    "last_km": it.last_km, "last_on": it.last_on, "starter": it.starter, **st})
    order = {"late": 0, "due": 1, "unknown": 2, "ok": 3}
    return sorted(out, key=lambda x: order.get(x["state"], 9))


def service_set(tenant: str, vehicle: str, item: str, every_km=None, every_months=None, last_km=None, last_on: str = "",
                remove: bool = False) -> dict:
    if not item:
        raise FleetError("Kaun sa service item?")
    with _LOCK, session() as s:
        v = _veh(s, tenant, vehicle)
        it = next((x for x in s.scalars(select(FleetServiceItem).where(FleetServiceItem.vehicle_id == v.id)).all()
                   if x.item.lower() == str(item).lower().strip()), None)
        if remove:
            if it:
                s.delete(it)
            return {"say": f"{v.reg}: '{item}' schedule se hataya."}
        if it is None:
            it = FleetServiceItem(tenant_id=tenant, vehicle_id=v.id, item=_clean(item, 80))
            s.add(it)
        if every_km is not None:
            it.every_km = max(0.0, _num(every_km))
        if every_months is not None:
            it.every_months = max(0, min(60, int(_num(every_months))))
        if last_km is not None:
            it.last_km = _num(last_km)
        if last_on:
            d = _date(last_on)
            if not d:
                raise FleetError("Pichhli service ki taarikh samajh nahi aayi.")
            it.last_on = d.isoformat()
        it.starter, it.reminded = False, ""
        s.flush()
        view = [x for x in _service_view(s, v, today(tenant)) if x["id"] == it.id][0]
        return {"item": view, "say": f"{v.reg}: {view['item']} har " + " / ".join(
            p for p in ((f"{view['every_km']:,.0f} km" if view["every_km"] else ""), (f"{view['every_months']} mahine" if view["every_months"] else "")) if p)
            + (f"; abhi: {view['why']}" if view["why"] else "") + "."}


def _event_view(e: FleetEvent) -> dict:
    d = _j(e.data, {})
    return {"id": e.id, "kind": e.kind, "day": e.day, "what": e.what, "amount": e.amount or None, "odo": e.odo, "status": e.status,
            "data": d, "by": e.by, "spend_id": e.spend_id}


def service_done(tenant: str, vehicle: str, items: list | str | None = None, odo=None, day: str = "", amount=None, what: str = "",
                 workshop: str = "", actor: str = "owner", spend: str = "") -> dict:
    """She (or staff) says a service was done: schedule lines reset; the amount is written down as she says it (a spend she
    already approved, or what she paid herself)."""
    d = _date(day) if day else today(tenant)
    if not d:
        raise FleetError("Taarikh samajh nahi aayi.")
    names = [items] if isinstance(items, str) else list(items or [])
    with _LOCK, session() as s:
        v = _veh(s, tenant, vehicle)
        o = _num(odo) or (v.odo or 0)
        done = []
        for it in s.scalars(select(FleetServiceItem).where(FleetServiceItem.vehicle_id == v.id)).all():
            low = it.item.lower()
            if any(str(n).lower().strip() and (str(n).lower() in low or low in str(n).lower()) for n in names) or "all" in [str(n).lower() for n in names]:
                it.last_km, it.last_on, it.reminded = (o or None), d.isoformat(), ""
                done.append(it.item)
        if o and (not v.odo or o >= v.odo):
            v.odo, v.odo_at = o, d.isoformat()
        e = FleetEvent(tenant_id=tenant, vehicle_id=v.id, kind="service", day=d.isoformat(), what=_clean(what or ", ".join(done) or "Service", 500),
                       amount=_num(amount), odo=o or None, data=_dump({"workshop": _clean(workshop, 120), "items": done}),
                       by=actor_name(tenant, actor), spend_id=spend)
        s.add(e)
        s.flush()
        reg = v.reg
        ev = _event_view(e)
    say = f"{reg}: service darj" + (f" ({', '.join(done)})" if done else "") + (f", {_rs(_num(amount))}" if _num(amount) else "") + "."
    if names and not done:
        say += " Schedule mein in naamon ka item nahi mila; schedule badalna ho to service_set."
    return {"event": ev, "reset": done, "say": say}


def _tyres(s, v: FleetVehicle) -> list[dict]:
    out = []
    for e in s.scalars(select(FleetEvent).where(FleetEvent.vehicle_id == v.id, FleetEvent.kind == "tyre")).all():
        d = _j(e.data, {})
        run = None
        if e.odo is not None:
            end = d.get("removed_km") or (v.odo or None)
            run = (end - e.odo) if end else None
        life = d.get("life_km") or None
        out.append({**_event_view(e), "position": d.get("position", ""), "brand": d.get("brand", ""), "serial": d.get("serial", ""),
                    "km_run": run, "life_km": life, "due": bool(e.status == "fitted" and life and run and run >= life * 0.9)})
    return sorted(out, key=lambda x: (x["status"] != "fitted", x["position"]))


def tyre(tenant: str, vehicle: str, op: str = "fit", position: str = "", brand: str = "", serial: str = "", odo=None, day: str = "",
         life_km=None, amount=None, tyre_id: str = "", reason: str = "", actor: str = "owner", spend: str = "") -> dict:
    d = _date(day) if day else today(tenant)
    with _LOCK, session() as s:
        v = _veh(s, tenant, vehicle)
        o = _num(odo) or (v.odo or 0)
        if op == "remove":
            e = s.get(FleetEvent, tyre_id) if tyre_id else next((x for x in s.scalars(select(FleetEvent).where(
                FleetEvent.vehicle_id == v.id, FleetEvent.kind == "tyre", FleetEvent.status == "fitted")).all()
                if position and _j(x.data, {}).get("position", "").lower() == position.lower()), None)
            if e is None or e.vehicle_id != v.id:
                raise FleetError("Kaun sa tyre? (position ya tyre id)")
            data = _j(e.data, {})
            data.update({"removed_km": o or None, "removed_on": d.isoformat(), "reason": _clean(reason, 200)})
            e.data, e.status = _dump(data), "removed"
            run = (o - e.odo) if (o and e.odo) else None
            return {"tyre": _event_view(e), "say": f"{v.reg}: {data.get('position', '')} tyre utara" + (f", {run:,.0f} km chala" if run else "") + "."}
        if not position:
            raise FleetError("Tyre ki position bataiye (jaise front-left, rear-2-outer, stepney).")
        old = [x for x in s.scalars(select(FleetEvent).where(FleetEvent.vehicle_id == v.id, FleetEvent.kind == "tyre",
                                                            FleetEvent.status == "fitted")).all()
               if _j(x.data, {}).get("position", "").lower() == position.lower()]
        for x in old:
            data = _j(x.data, {})
            data.update({"removed_km": o or None, "removed_on": d.isoformat(), "reason": "naya lagaya"})
            x.data, x.status = _dump(data), "removed"
        e = FleetEvent(tenant_id=tenant, vehicle_id=v.id, kind="tyre", day=d.isoformat(), what=f"Tyre {position}", amount=_num(amount),
                       odo=o or None, status="fitted", by=actor_name(tenant, actor), spend_id=spend,
                       data=_dump({"position": _clean(position, 40), "brand": _clean(brand, 60), "serial": _clean(serial, 40),
                                   "life_km": _num(life_km) or None}))
        s.add(e)
        s.flush()
        return {"tyre": _event_view(e), "say": f"{v.reg}: {position} par naya tyre" + (f" ({brand})" if brand else "") + " darj."}


def breakdown(tenant: str, vehicle: str, what: str, place: str = "", day: str = "", hours=None, amount=None, odo=None,
              actor: str = "owner", conv: str = "") -> dict:
    if not what:
        raise FleetError("Kya kharabi hui?")
    d = _date(day) if day else today(tenant)
    with _LOCK, session() as s:
        v = _veh(s, tenant, vehicle)
        e = FleetEvent(tenant_id=tenant, vehicle_id=v.id, kind="breakdown", day=d.isoformat(), what=_clean(what, 500), amount=_num(amount),
                       odo=_num(odo) or None, status="open", by=actor_name(tenant, actor),
                       data=_dump({"place": _clean(place, 160), "hours": _num(hours) or None}))
        s.add(e)
        since = (d - timedelta(days=90)).isoformat()
        n = len(s.scalars(select(FleetEvent).where(FleetEvent.vehicle_id == v.id, FleetEvent.kind == "breakdown", FleetEvent.day >= since)).all())
        if v.status == "active":
            v.status = "workshop"
        s.flush()
        ev, reg = _event_view(e), v.reg
    say = f"{reg}: breakdown darj ({what[:60]}). Gaadi 'workshop' mein maani; theek hone par batana."
    if n >= 3:
        say += f" 90 din mein {n} breakdown: badi marammat / purani gaadi ka hisaab dekhiye."
    if not is_owner(actor):
        _publish(tenant, "fleet_alert", {"kind": "breakdown", "vehicle": reg, "say": say}, conv)
    return {"event": ev, "repeat": n, "say": say}


def event_close(tenant: str, event: str, amount=None, note: str = "") -> dict:
    with _LOCK, session() as s:
        e = s.get(FleetEvent, event)
        if e is None or e.tenant_id != tenant:
            raise FleetError("Entry nahi mili.")
        e.status = "done"
        if amount is not None:
            e.amount = _num(amount)
        if note:
            e.what = (e.what + " | " + _clean(note, 300))[:1000]
        v = s.get(FleetVehicle, e.vehicle_id)
        if e.kind == "breakdown" and v and v.status == "workshop":
            v.status = "active"
        return {"event": _event_view(e), "say": "Band kiya" + (f"; {v.reg} phir active." if v and e.kind == "breakdown" else ".")}


def expense(tenant: str, vehicle: str, what: str, amount, day: str = "", actor: str = "owner", spend: str = "") -> dict:
    a = _num(amount)
    if a <= 0:
        raise FleetError("Rakam chahiye.")
    d = _date(day) if day else today(tenant)
    with _LOCK, session() as s:
        v = _veh(s, tenant, vehicle)
        e = FleetEvent(tenant_id=tenant, vehicle_id=v.id, kind="expense", day=d.isoformat(), what=_clean(what, 500), amount=a,
                       by=actor_name(tenant, actor), spend_id=spend)
        s.add(e)
        s.flush()
        return {"event": _event_view(e), "say": f"{v.reg}: {what[:60]} {_rs(a)} darj."}


def challan(tenant: str, vehicle: str, number: str = "", amount=None, day: str = "", offence: str = "", place: str = "",
            actor: str = "owner") -> dict:
    """A challan she found on eChallan (her steps) or got on the road. Hoola never pays it."""
    d = _date(day) if day else today(tenant)
    with _LOCK, session() as s:
        v = _veh(s, tenant, vehicle)
        num = _clean(number, 40)
        if num:
            for x in s.scalars(select(FleetEvent).where(FleetEvent.tenant_id == tenant, FleetEvent.kind == "challan")).all():
                if _j(x.data, {}).get("number") == num:
                    return {"event": _event_view(x), "duplicate": True, "say": f"Challan {num} pehle se darj hai."}
        e = FleetEvent(tenant_id=tenant, vehicle_id=v.id, kind="challan", day=d.isoformat(), what=_clean(offence or "Challan", 300),
                       amount=_num(amount), status="open", by=actor_name(tenant, actor),
                       data=_dump({"number": num, "offence": _clean(offence, 200), "place": _clean(place, 120)}))
        s.add(e)
        s.flush()
        reg = v.reg
        ev = _event_view(e)
    say = f"{reg}: challan" + (f" {num}" if num else "") + (f" {_rs(_num(amount))}" if _num(amount) else "") + " darj (khula)."
    if re.search(r"overload|over load|194", offence or "", re.I):
        say += " Overload: aage se trip plan mein capacity se zyada maal nahi."
    say += " Bharna ho to 'spend' card banaiye; paise aap portal par khud bharengi."
    return {"event": ev, "say": say}


# ---------------------------------------------------------------- spends: HER tap only
def _spend_show(x: FleetSpend, reg: str) -> dict:
    return {"number": x.number, "vehicle": reg, "kind": x.kind, "kind_label": SPEND_KINDS.get(x.kind, x.kind), "what": x.what,
            "amount": round(x.amount, 2), "vendor": x.vendor, "where": x.where}


def _spend_hash(show: dict) -> str:
    return hashlib.sha256(_dump(show).encode()).hexdigest()


def _spend_view(s, x: FleetSpend) -> dict:
    v = s.get(FleetVehicle, x.vehicle_id) if x.vehicle_id else None
    show = _spend_show(x, v.reg if v else "")
    return {"id": x.id, **show, "amount_text": _rs(x.amount), "checks": _j(x.checks, []), "status": x.status, "hash": x.hash,
            "asked_by": x.asked_by, "reason": x.reason, "paid_amount": x.paid_amount, "paid_ref": x.paid_ref,
            "created_at": _aware(x.created_at).isoformat(timespec="minutes") if x.created_at else ""}


_PAY_WHERE = {"permit": "vahan.parivahan.gov.in/npermit (national) ya Vahan permit services: aap khud bharti hain.",
              "tax": "vahan.parivahan.gov.in > Pay your tax: aap khud bharti hain.",
              "challan": "echallan.parivahan.gov.in > Pay now (OTP / payment aapka).",
              "fitness": "Vahan par fitness appointment fees / ATS counter.",
              "insurance": "Apna insurer / agent (policy ki copy yahan bhejiye).",
              "fastag": "Apne FASTag bank ka app."}


def spend_request(tenant: str, vehicle: str, kind: str, what: str, amount, vendor: str = "", link_id: str = "",
                  actor: str = "owner", conv: str = "") -> dict:
    """Make the card. Nothing is approved here: only her tap in the app (decide_spend by_tap=True)."""
    k = kind if kind in SPEND_KINDS else "other"
    a = _num(amount)
    if a <= 0:
        raise FleetError("Kharche ki rakam (aapke workshop / portal / bill se) chahiye; Hoola rakam nahi gadhta.")
    if a > 5_000_000:
        raise FleetError("Rakam bahut badi lagti hai; dobara jaanchiye.")
    if not what:
        raise FleetError("Kis cheez ka kharcha?")
    on = today(tenant)
    with _LOCK, session() as s:
        v = _veh(s, tenant, vehicle) if vehicle else None
        w = _clean(what, 300)
        same = s.scalars(select(FleetSpend).where(FleetSpend.tenant_id == tenant, FleetSpend.status == "waiting",
                                                  FleetSpend.kind == k, FleetSpend.vehicle_id == (v.id if v else ""))).all()
        for x in same:
            if x.what.lower() == w.lower() and abs(x.amount - a) < 0.5:
                return {"spend": _spend_view(s, x), "existing": True, "say": f"{x.number} ka card pehle se aapki haan ka intezaar kar raha hai."}
        checks = []
        if v is not None:
            if v.status == "sold":
                checks.append("Yeh gaadi 'sold' darj hai.")
            past = [e.amount for e in s.scalars(select(FleetEvent).where(FleetEvent.vehicle_id == v.id, FleetEvent.amount > 0)).all()
                    if (k == "service" and e.kind == "service") or (k == "tyre" and e.kind == "tyre")]
            if len(past) >= 2:
                med = statistics.median(past)
                if a > med * 1.3:
                    checks.append(f"Pichhle {SPEND_KINDS[k].lower()} ka aam kharcha {_rs(med)}; yeh {round((a / med - 1) * 100)}% zyada.")
            if k in ("permit", "insurance", "fitness", "puc", "tax"):
                dk = {"permit": "national_permit" if v.national else "state_permit", "insurance": "insurance", "fitness": "fitness",
                      "puc": "puc", "tax": "road_tax"}[k]
                dd = s.scalars(select(FleetDoc).where(FleetDoc.vehicle_id == v.id, FleetDoc.kind == dk)).first()
                if dd and dd.expires:
                    left = (_date(dd.expires) - on).days
                    if left > 60:
                        checks.append(f"{R.DOC_KINDS[dk]['label']} abhi {left} din aur valid ({dd.expires}); itni jaldi renewal?")
            if k == "challan" and link_id:
                ev = s.get(FleetEvent, link_id)
                if ev and ev.amount and abs(ev.amount - a) > 0.5:
                    checks.append(f"Challan par rakam {_rs(ev.amount)} likhi thi, card par {_rs(a)}.")
        st = _st(s, tenant)
        st.spend_seq += 1
        x = FleetSpend(tenant_id=tenant, number=f"FS-{st.spend_seq:04d}", vehicle_id=v.id if v else "", kind=k, what=w, amount=a,
                       vendor=_clean(vendor, 120), where=_PAY_WHERE.get(k, "Aap khud dijiye (workshop / dukaan / UPI)."),
                       checks=_dump(checks), link_id=str(link_id or "")[:40], asked_by=actor_name(tenant, actor), conv=conv)
        s.add(x)
        s.flush()
        x.hash = _spend_hash(_spend_show(x, v.reg if v else ""))
        view = _spend_view(s, x)
    say = (f"{view['number']}: {view['kind_label']} {view['amount_text']}" + (f" ({view['vehicle']})" if view["vehicle"] else "")
           + " ka card aapki haan ke liye bheja hai. Hoola khud paise nahi deta.")
    if checks:
        say += " Dhyaan: " + " ".join(checks)
    _publish(tenant, "fleet_spend", {"spend": view, "say": say}, conv)
    return {"spend": view, "status": "needs_tap", "say": say}


def spend_card(tenant: str, spend_id: str) -> dict | None:
    with session() as s:
        x = s.get(FleetSpend, spend_id)
        if x is None or x.tenant_id != tenant:
            return None
        return _spend_view(s, x)


def spend_needs_fingerprint(tenant: str, spend_id: str) -> bool:
    from .. import approvekey as ak
    c = spend_card(tenant, spend_id)
    return bool(c and c["amount"] > 0 and ak.active_on(tenant))


def spend_challenge(tenant: str, spend_id: str) -> dict:
    """approvekey's message format (kind 'fleet'): the phone signs this card as it is NOW."""
    from .. import approvekey as ak
    c = spend_card(tenant, spend_id)
    if not c or c["status"] != "waiting":
        raise FleetError("Yeh card ab baaki nahi hai.")
    digest = c["hash"]
    nonce = secrets.token_urlsafe(18)
    exp = datetime.now(timezone.utc) + ak.CHALLENGE_TTL
    msg = f"hoola-approve|v1|{tenant}|fleet|{spend_id}|{digest}|{nonce}|{int(exp.timestamp())}"
    with session() as s:
        ch = ak.ApproveChallenge(tenant_id=tenant, kind="fleet", target=spend_id, digest=digest, message=msg, expires_at=exp)
        s.add(ch)
        s.flush()
        chid = ch.id
    show = {"title": f"{c['kind_label']}: {c['what'][:60]}", "amount": c["amount_text"], "payee": c["vendor"] or c["vehicle"]}
    return {"challenge_id": chid, "message": msg, "show": show, "expires_at": exp.isoformat()}


def _verify_fingerprint(tenant: str, spend_id: str, digest: str, proof: dict | None) -> None:
    from .. import approvekey as ak
    if not proof:
        raise FleetError("fingerprint_needed")
    with session() as s:
        ak._settle(s, tenant)
        ok = ak._check(s, tenant, "fleet", spend_id, proof, digest)
    if not ok:
        raise FleetError("fingerprint_invalid")


def decide_spend(tenant: str, spend_id: str, yes: bool, card_hash: str = "", reason: str = "", proof: dict | None = None,
                 by_tap: bool = False) -> dict:
    """ONLY her tap in the app (API). The tool can never reach this with by_tap."""
    if not by_tap:
        raise FleetError("Kharche ki haan sirf aapke tap se (app ka card).")
    with session() as s:
        x = s.get(FleetSpend, spend_id)
        if x is None or x.tenant_id != tenant:
            raise FleetError("Card nahi mila.")
        if x.status != "waiting":
            return {"spend": _spend_view(s, x), "say": f"{x.number} pehle hi '{x.status}' hai."}
        v = s.get(FleetVehicle, x.vehicle_id) if x.vehicle_id else None
        now_hash = _spend_hash(_spend_show(x, v.reg if v else ""))
        if not card_hash or card_hash != x.hash or now_hash != x.hash:
            raise FleetError("Card badal gaya hai; naya card dekhiye.")
    if yes:
        from .. import approvekey as ak
        if ak.active_on(tenant):
            _verify_fingerprint(tenant, spend_id, card_hash, proof)
    with _LOCK, session() as s:
        x = s.get(FleetSpend, spend_id)
        x.status = "approved" if yes else "rejected"
        x.reason = _clean(reason, 300)
        x.decided_at = now()
        view = _spend_view(s, x)
    if yes:
        say = f"{view['number']} maana: {view['amount_text']}. Paise aap khud dijiye ({view['where']}) aur phir 'Maine kar diya' dabaiye."
    else:
        say = f"{view['number']}: kharcha nahi" + (f" ({view['reason']})" if view["reason"] else "") + "."
    return {"spend": view, "say": say}


def spend_done(tenant: str, spend_id: str, amount=None, ref: str = "", day: str = "", actor: str = "owner") -> dict:
    """She paid herself (UPI / cash / portal). Only after her approval; the spend becomes a cost entry."""
    with _LOCK, session() as s:
        x = s.get(FleetSpend, spend_id)
        if x is None or x.tenant_id != tenant:
            raise FleetError("Card nahi mila.")
        if x.status != "approved":
            raise FleetError(f"{x.number} abhi '{x.status}' hai: pehle aapki haan (card par tap).")
        a = _num(amount) or x.amount
        x.status, x.paid_amount, x.paid_ref = "done", a, _clean(ref, 80)
        d = (_date(day) or today(tenant)).isoformat()
        ev_kind = {"service": "service", "tyre": "tyre", "challan": "challan"}.get(x.kind, "expense")
        link = s.get(FleetEvent, x.link_id) if x.link_id else None
        if link is not None and link.tenant_id == tenant:
            link.amount = a
            link.spend_id = x.id
            if link.kind == "challan":
                link.status = "paid"
        elif x.vehicle_id and ev_kind == "expense":
            s.add(FleetEvent(tenant_id=tenant, vehicle_id=x.vehicle_id, kind="expense", day=d, what=f"{SPEND_KINDS.get(x.kind, '')}: {x.what}",
                             amount=a, spend_id=x.id, by=actor_name(tenant, actor)))
        view = _spend_view(s, x)
    say = f"{view['number']} poora: {_rs(a)} aapne diya." + (" Service / tyre ho jaye to 'service_done' / 'tyre' se darj kijiye."
                                                             if x.kind in ("service", "tyre") and not x.link_id else "")
    return {"spend": view, "say": say}


def spends(tenant: str, status: str = "") -> list[dict]:
    with session() as s:
        q = select(FleetSpend).where(FleetSpend.tenant_id == tenant)
        if status:
            q = q.where(FleetSpend.status == status)
        return [_spend_view(s, x) for x in s.scalars(q.order_by(FleetSpend.created_at.desc())).all()[:80]]


# ---------------------------------------------------------------- trips + loads
def _maps_distance(tenant: str, origin: str, destination: str, stops: list[str]) -> tuple[float, int] | None:
    """Google Maps Routes (only when the server has the key). Stops are summed leg by leg. None = not available."""
    from .. import maps
    if not maps.configured():
        return None
    try:
        asyncio.get_running_loop()
        return None                                      # called on the event loop thread: never block it
    except RuntimeError:
        pass

    async def go():
        pts = [origin, *stops, destination]
        km, mins = 0.0, 0
        prev = await maps.resolve(pts[0], tenant)
        for p in pts[1:]:
            nxt = await maps.resolve(p, tenant)
            if not prev or not nxt:
                return None
            r = await maps.route(prev, nxt, "DRIVE")
            km += r["distance_km"]
            mins += r["duration_min"]
            prev = nxt
        return round(km, 1), mins
    try:
        return asyncio.run(go())
    except Exception as e:
        log.info("fleet maps distance failed: %s", e)
        return None


def _items(items) -> list[dict]:
    out = []
    for i in items or []:
        if not isinstance(i, dict):
            continue
        dims = [x for x in (i.get("dims") or [i.get("l"), i.get("w"), i.get("h")]) if _num(x) > 0]
        out.append({"name": _clean(i.get("name") or i.get("item") or "saamaan", 80), "weight_kg": _num(i.get("weight_kg") or i.get("kg")) or None,
                    "cbm": _num(i.get("cbm")) or None, "count": max(1, int(_num(i.get("count") or i.get("qty") or 1))),
                    "dims": [_num(x) for x in dims][:3]})
    return out


def _veh_dict(s, v: FleetVehicle) -> dict:
    b = _j(v.body, {})
    return {"id": v.id, "reg": v.reg, "capacity_t": v.capacity_t, "capacity_cbm": v.capacity_cbm, "body_l": b.get("l"),
            "body_w": b.get("w"), "body_h": b.get("h"), "type": v.vtype}


def _paper_check(s, v: FleetVehicle, on: date, interstate: bool) -> tuple[list[str], list[str]]:
    block, warn = [], []
    docs = {d.kind: d for d in s.scalars(select(FleetDoc).where(FleetDoc.vehicle_id == v.id)).all()}
    commercial = R.VEHICLE_TYPES.get(v.vtype, {}).get("commercial", True)
    for k in BLOCK_DOCS + WARN_DOCS:
        d = docs.get(k)
        if d is None or d.not_needed:
            continue
        exp = _date(d.expires)
        label = R.DOC_KINDS[k]["label"]
        if exp is None:
            if k in ("insurance", "fitness") and (k != "fitness" or commercial):
                warn.append(f"{label} ki expiry darj nahi.")
            continue
        if exp < on:
            (block if k in BLOCK_DOCS else warn).append(f"{label} {d.expires} ko khatam.")
    if interstate and R.VEHICLE_TYPES.get(v.vtype, {}).get("goods") and not v.national:
        warn.append("Doosre rajya ka goods trip: national permit darj nahi (ya us rajya ka permit chahiye).")
    if v.status != "active":
        block.append(f"Gaadi abhi '{v.status}' hai.")
    return block, warn


def _busy(s, tenant: str, vid: str, day: str, exclude: str = "") -> bool:
    q = select(FleetTrip).where(FleetTrip.tenant_id == tenant, FleetTrip.vehicle_id == vid, FleetTrip.day == day,
                                FleetTrip.status.in_(("assigned", "started")))
    return any(t.id != exclude for t in s.scalars(q).all())


def _trip_view(s, t: FleetTrip) -> dict:
    v = s.get(FleetVehicle, t.vehicle_id) if t.vehicle_id else None
    return {"id": t.id, "number": t.number, "origin": t.origin, "destination": t.destination, "stops": _j(t.stops, []), "day": t.day,
            "origin_state": t.origin_state, "dest_state": t.dest_state, "items": _j(t.items, []), "vehicle": v.reg if v else "",
            "vehicle_id": t.vehicle_id, "driver": t.driver, "driver_phone": t.driver_phone, "km": t.km, "km_source": t.km_source,
            "duration_min": t.duration_min, "toll": t.toll, "estimate": _j(t.estimate, {}), "checks": _j(t.checks, {}),
            "status": t.status, "status_label": TRIP_STATUS.get(t.status, t.status), "odo_start": t.odo_start, "odo_end": t.odo_end,
            "ref": t.ref, "notes": t.notes}


def _evaluate(s, tenant: str, t: FleetTrip) -> dict:
    """Load fit, papers, estimate, suggestion: written into t.checks / t.estimate. Returns the checks."""
    st = _st(s, tenant)
    on = _date(t.day) or today(tenant)
    items = _j(t.items, [])
    inter = R.interstate(t.origin_state or st.home_state, t.dest_state)
    checks: dict[str, Any] = {"block": [], "warn": [], "notes": [], "fit": None, "suggest": [], "split": None}
    v = s.get(FleetVehicle, t.vehicle_id) if t.vehicle_id else None
    free = [x for x in s.scalars(select(FleetVehicle).where(FleetVehicle.tenant_id == tenant, FleetVehicle.status == "active")).all()
            if not _busy(s, tenant, x.id, t.day, t.id)]
    if items:
        if v is not None:
            fit = R.load_fit(items, _veh_dict(s, v), st.fill_pct)
            checks["fit"] = fit
            if not fit["fits"]:
                checks["block"] += fit["problems"]
            checks["notes"] += fit["notes"]
        cands = []
        for x in free:
            if x.capacity_t <= 0:
                continue
            f = R.load_fit(items, _veh_dict(s, x), st.fill_pct)
            b, _w = _paper_check(s, x, on, inter)
            if f["fits"] and not b:
                cands.append((x.capacity_t, x.reg, f["weight_pct"]))
        checks["suggest"] = [{"vehicle": r, "weight_pct": p} for _c, r, p in sorted(cands)[:3]]
        if not cands and (v is None or not (checks["fit"] or {}).get("fits")):
            sp = R.split_load(items, [_veh_dict(s, x) for x in free if not _paper_check(s, x, on, inter)[0]], st.fill_pct)
            if sp["plan"]:
                checks["split"] = {"plan": sp["plan"], "left": [p.get("name") for p in sp["left"]]}
    if v is not None:
        b, w = _paper_check(s, v, on, inter)
        checks["block"] += b
        checks["warn"] += w
        if _busy(s, tenant, v.id, t.day, t.id):
            checks["block"].append(f"{v.reg} {t.day} ko doosri trip par hai.")
        rate = {"diesel": st.diesel_rate, "petrol": st.petrol_rate, "cng": st.cng_rate}.get(v.fuel) or None
        tpk = _j(st.toll_per_km, {}).get(v.vtype) or None
        est = R.trip_estimate(t.km, _kmpl_of(s, v), rate, t.toll, tpk)
        t.estimate = _dump(est)
    else:
        t.estimate = _dump(R.trip_estimate(t.km, None, None, t.toll))
    hours = round(t.duration_min / 60, 1) if t.duration_min else (round(t.km / st.avg_speed, 1) if (t.km and st.avg_speed) else None)
    note = R.drive_hours_note(hours)
    if note:
        checks["warn"].append(note)
    if not t.km:
        checks["notes"].append("Distance pata nahi: aap km bataiye (Google Maps server par juda ho to Hoola khud nikalega).")
    t.checks = _dump(checks)
    return checks


def plan_trip(tenant: str, origin: str, destination: str, day: str = "", stops: list | None = None, items: list | None = None,
              vehicle: str = "", driver: str = "", driver_phone: str = "", km=None, toll=None, origin_state: str = "",
              dest_state: str = "", ref: str = "", notes: str = "", actor: str = "owner", conv: str = "") -> dict:
    if not origin or not destination:
        raise FleetError("Kahan se kahan? (origin aur destination)")
    d = _date(day) if day else today(tenant)
    if not d:
        raise FleetError("Trip ki taarikh samajh nahi aayi.")
    if d < today(tenant):
        raise FleetError("Beeti taarikh ki trip plan nahi hoti.")
    phone = re.sub(r"\D", "", driver_phone or "")[-10:]
    if driver_phone and len(phone) != 10:
        raise FleetError("Driver ka phone 10 ank ka.")
    with _LOCK, session() as s:
        st = _st(s, tenant)
        st.trip_seq += 1
        t = FleetTrip(tenant_id=tenant, number=f"TR-{st.trip_seq:04d}", origin=_clean(origin, 160), destination=_clean(destination, 160),
                      stops=_dump([_clean(x, 160) for x in (stops or []) if str(x).strip()][:10]), day=d.isoformat(),
                      items=_dump(_items(items)), driver=_clean(driver, 80), driver_phone=phone,
                      origin_state=(origin_state or st.home_state or "").upper()[:4], dest_state=(dest_state or "").upper()[:4],
                      ref=_clean(ref, 80), notes=_clean(notes, 1000), conv=conv)
        if vehicle:
            t.vehicle_id = _veh(s, tenant, vehicle).id
        if _num(km) > 0:
            t.km, t.km_source = _num(km), "her"
        if _num(toll) > 0:
            t.toll = _num(toll)
        s.add(t)
        s.flush()
        tid = t.id
    if not _num(km):
        md = _maps_distance(tenant, origin, destination, list(stops or []))
        if md:
            with session() as s:
                t = s.get(FleetTrip, tid)
                t.km, t.duration_min, t.km_source = md[0], md[1], "maps"
    return _finish_plan(tenant, tid, conv)


def _finish_plan(tenant: str, tid: str, conv: str = "", publish: bool = True) -> dict:
    with _LOCK, session() as s:
        t = s.get(FleetTrip, tid)
        ch = _evaluate(s, tenant, t)
        if t.vehicle_id and not ch["block"]:
            t.status = "assigned" if t.status == "planned" else t.status
        elif t.status == "assigned" and ch["block"]:
            t.status = "planned"
        view = _trip_view(s, t)
    say = f"{view['number']}: {view['origin']} -> {view['destination']} ({view['day']})"
    if view["km"]:
        say += f", {view['km']:g} km" + (" (Google Maps)" if view["km_source"] == "maps" else "")
    say += "."
    if view["vehicle"]:
        say += f" Gaadi {view['vehicle']}" + (" tay." if view["status"] == "assigned" else " abhi tay nahi:")
    if ch["block"]:
        say += " Rukawat: " + " ".join(ch["block"])
    elif ch["suggest"] and not view["vehicle"]:
        say += " Sahi gaadi: " + ", ".join(f"{x['vehicle']} ({x['weight_pct']:g}% bhari)" for x in ch["suggest"]) + "."
    if ch.get("split"):
        say += " Ek gaadi mein nahi aata; baantne ka plan: " + "; ".join(
            f"{p['vehicle']} {p['weight_kg']:,.0f} kg" for p in ch["split"]["plan"]) + (
            f"; phir bhi baaki: {', '.join(x for x in ch['split']['left'] if x)}" if ch["split"]["left"] else "") + "."
    if ch["warn"]:
        say += " Dhyaan: " + " ".join(ch["warn"])
    est = view["estimate"]
    if est.get("total_rs"):
        say += f" Andaaza (aapke rate se): diesel {_rs(est['fuel_rs'])} + toll {_rs(est['toll_rs'])} = {_rs(est['total_rs'])}."
    elif est.get("missing") and view["vehicle"]:
        say += " Kharche ka andaaza ke liye chahiye: " + ", ".join(est["missing"]) + "."
    if publish:
        _publish(tenant, "fleet_trip", {"trip": view, "say": say}, conv)
    return {"trip": view, "say": say}


def _trip(s, tenant: str, ref: str) -> FleetTrip:
    t = s.get(FleetTrip, ref)
    if t is not None and t.tenant_id == tenant:
        return t
    r = str(ref or "").strip().upper()
    if r.isdigit():
        r = f"TR-{int(r):04d}"
    t = s.scalars(select(FleetTrip).where(FleetTrip.tenant_id == tenant, FleetTrip.number == r)).first()
    if t is None:
        raise FleetError(f"Trip '{ref}' nahi mili.")
    return t


def assign_trip(tenant: str, trip: str, vehicle: str = "", driver: str = "", driver_phone: str = "", km=None, toll=None,
                items: list | None = None, conv: str = "") -> dict:
    with _LOCK, session() as s:
        t = _trip(s, tenant, trip)
        if t.status in ("done", "cancelled"):
            raise FleetError(f"{t.number} '{t.status}' hai.")
        if vehicle:
            t.vehicle_id = _veh(s, tenant, vehicle).id
        if driver:
            t.driver = _clean(driver, 80)
        if driver_phone:
            p = re.sub(r"\D", "", driver_phone)[-10:]
            if len(p) != 10:
                raise FleetError("Driver ka phone 10 ank ka.")
            t.driver_phone = p
        if _num(km) > 0:
            t.km, t.km_source = _num(km), "her"
        if _num(toll) > 0:
            t.toll = _num(toll)
        if items is not None:
            t.items = _dump(_items(items))
        tid = t.id
    return _finish_plan(tenant, tid, conv)


def trip_status(tenant: str, trip: str, status: str, odo=None, note: str = "", actor: str = "owner") -> dict:
    if status not in ("started", "done", "cancelled"):
        raise FleetError("Status: started / done / cancelled.")
    with _LOCK, session() as s:
        t = _trip(s, tenant, trip)
        if t.status in ("done", "cancelled"):
            raise FleetError(f"{t.number} pehle hi '{t.status}' hai.")
        o = _num(odo)
        v = s.get(FleetVehicle, t.vehicle_id) if t.vehicle_id else None
        if status == "started":
            if v is None:
                raise FleetError("Pehle gaadi tay kijiye.")
            ch = _evaluate(s, tenant, t)
            if ch["block"]:
                raise FleetError("Trip shuru nahi: " + " ".join(ch["block"]))
            t.odo_start = o or (v.odo or None)
        if status == "done":
            if o:
                if t.odo_start and o < t.odo_start:
                    raise FleetError("Aakhri odometer shuru se kam.")
                t.odo_end = o
                if v and (not v.odo or o >= v.odo):
                    v.odo, v.odo_at = o, today(tenant).isoformat()
        t.status = status
        if note:
            t.notes = (t.notes + "\n" + _clean(note, 500)).strip()[:2000]
        view = _trip_view(s, t)
    say = f"{view['number']}: {view['status_label']}."
    if status == "done" and view["odo_start"] and view["odo_end"]:
        run = view["odo_end"] - view["odo_start"]
        say += f" {run:,.0f} km chali" + (f" (plan {view['km']:g} km)" if view["km"] else "") + "."
        if view["km"] and run > view["km"] * 1.2 + 10:
            say += " Plan se kaafi zyada: rasta badla ya galat odometer?"
    return {"trip": view, "say": say}


def trips(tenant: str, status: str = "", day: str = "") -> list[dict]:
    with session() as s:
        q = select(FleetTrip).where(FleetTrip.tenant_id == tenant)
        if status:
            q = q.where(FleetTrip.status == status)
        if day:
            q = q.where(FleetTrip.day == day)
        return [_trip_view(s, t) for t in s.scalars(q.order_by(FleetTrip.day.desc(), FleetTrip.number.desc())).all()[:80]]


def trip_view(tenant: str, trip: str) -> dict:
    with session() as s:
        return _trip_view(s, _trip(s, tenant, trip))


def load_check(tenant: str, items: list, vehicle: str = "") -> dict:
    """Which vehicle takes this load (no trip saved)."""
    its = _items(items)
    st = settings(tenant)
    with session() as s:
        vs = [_veh(s, tenant, vehicle)] if vehicle else s.scalars(select(FleetVehicle).where(
            FleetVehicle.tenant_id == tenant, FleetVehicle.status == "active")).all()
        res = [{"vehicle": v.reg, **R.load_fit(its, _veh_dict(s, v), st["fill_pct"])} for v in vs]
        split = R.split_load(its, [_veh_dict(s, v) for v in vs], st["fill_pct"]) if not any(r["fits"] for r in res) else None
    ok = sorted([r for r in res if r["fits"]], key=lambda r: r["weight_pct"] or 0, reverse=True)
    say = (f"Yeh maal {ok[0]['vehicle']} mein ({ok[0]['weight_pct']:g}% wazan)." if ok else "Kisi ek gaadi mein nahi aata.")
    if split and split["plan"]:
        say += " Baantkar: " + "; ".join(f"{p['vehicle']} {p['weight_kg']:,.0f} kg" for p in split["plan"]) + "."
    return {"results": res, "split": split, "say": say}


# ---------------------------------------------------------------- reports
def cost_report(tenant: str, vehicle: str = "", days: int = 30) -> dict:
    """Cost per km over the period: fuel + service + tyres + repairs + expenses + challans paid, divided by km run (odometer
    span from fuel / trip readings). Idle days = days without a trip or a fill."""
    on = today(tenant)
    since = on - timedelta(days=days)
    out = []
    with session() as s:
        vs = [_veh(s, tenant, vehicle)] if vehicle else s.scalars(select(FleetVehicle).where(FleetVehicle.tenant_id == tenant,
                                                                                            FleetVehicle.status != "sold")).all()
        for v in vs:
            fuel = s.scalars(select(FleetFuel).where(FleetFuel.vehicle_id == v.id, FleetFuel.day >= since.isoformat())).all()
            evs = s.scalars(select(FleetEvent).where(FleetEvent.vehicle_id == v.id, FleetEvent.day >= since.isoformat())).all()
            tr = s.scalars(select(FleetTrip).where(FleetTrip.vehicle_id == v.id, FleetTrip.day >= since.isoformat(),
                                                   FleetTrip.status != "cancelled")).all()
            odos = [f.odo for f in fuel if f.odo] + [x for t in tr for x in (t.odo_start, t.odo_end) if x]
            km = (max(odos) - min(odos)) if len(odos) >= 2 else sum((t.odo_end - t.odo_start) for t in tr if t.odo_end and t.odo_start)
            cost = {"fuel": round(sum(f.amount for f in fuel)), "service": 0, "tyre": 0, "breakdown": 0, "expense": 0, "challan": 0}
            for e in evs:
                if e.kind in cost and (e.kind != "challan" or e.status == "paid"):
                    cost[e.kind] += round(e.amount)
            total = sum(cost.values())
            active_days = {f.day for f in fuel} | {t.day for t in tr}
            out.append({"vehicle": v.reg, "km": round(km) if km else None, "cost": cost, "total": total,
                        "per_km": round(total / km, 2) if km else None, "trips": len(tr), "idle_days": days - len(active_days),
                        "kmpl": _kmpl_of(s, v)})
    out.sort(key=lambda x: -(x["per_km"] or 0))
    return {"days": days, "vehicles": out}


def idle(tenant: str) -> list[dict]:
    on = today(tenant)
    with session() as s:
        st = _st(s, tenant)
        since = (on - timedelta(days=st.idle_days)).isoformat()
        out = []
        for v in s.scalars(select(FleetVehicle).where(FleetVehicle.tenant_id == tenant, FleetVehicle.status == "active")).all():
            last_f = s.scalars(select(FleetFuel.day).where(FleetFuel.vehicle_id == v.id).order_by(FleetFuel.day.desc())).first()
            last_t = s.scalars(select(FleetTrip.day).where(FleetTrip.vehicle_id == v.id, FleetTrip.status != "cancelled",
                                                           FleetTrip.day <= on.isoformat()).order_by(FleetTrip.day.desc())).first()
            last = max([x for x in (last_f, last_t) if x], default="")
            if not last:
                created = _aware(v.created_at).date() if v.created_at else on
                if (on - created).days < st.idle_days:
                    continue
            if not last or last < since:
                out.append({"vehicle": v.reg, "last": last, "days": (on - _date(last)).days if last else None})
        return out


def day(tenant: str) -> dict:
    on = today(tenant)
    docs = docs_due(tenant, 30)
    expired = [d for d in docs if d["state"] == "expired"]
    due = [d for d in docs if d["state"] == "due"]
    with session() as s:
        vs = s.scalars(select(FleetVehicle).where(FleetVehicle.tenant_id == tenant, FleetVehicle.status != "sold")).all()
        svc, tyres = [], []
        for v in vs:
            for it in _service_view(s, v, on):
                if it["state"] in ("due", "late"):
                    svc.append({"vehicle": v.reg, **it})
            tyres += [{"vehicle": v.reg, **t} for t in _tyres(s, v) if t["due"]]
        since = (on - timedelta(days=3)).isoformat()
        fuel_flags = [{"vehicle": next((v.reg for v in vs if v.id == f.vehicle_id), ""), "day": f.day, "flags": _j(f.flags, [])}
                      for f in s.scalars(select(FleetFuel).where(FleetFuel.tenant_id == tenant, FleetFuel.day >= since)).all()
                      if _j(f.flags, [])]
        challans = [_event_view(e) | {"vehicle": next((v.reg for v in vs if v.id == e.vehicle_id), "")}
                    for e in s.scalars(select(FleetEvent).where(FleetEvent.tenant_id == tenant, FleetEvent.kind == "challan",
                                                                FleetEvent.status == "open")).all()]
        workshop = [v.reg for v in vs if v.status == "workshop"]
    waiting = spends(tenant, "waiting")
    approved = spends(tenant, "approved")
    today_trips = [t for t in trips(tenant, day=on.isoformat()) if t["status"] in ("planned", "assigned", "started")]
    blocked = [t for t in today_trips if t["checks"].get("block")]
    idle_v = idle(tenant)
    parts = []
    if today_trips:
        parts.append(f"Aaj {len(today_trips)} trip" + (f" ({len(blocked)} mein rukawat)" if blocked else "") + ".")
    if expired:
        parts.append("Khatam kaagaz: " + ", ".join(f"{d['vehicle']} {d['label']}" for d in expired[:5]) + ".")
    if due:
        parts.append("Jaldi khatam: " + ", ".join(f"{d['vehicle']} {d['label']} ({d['days_left']} din)" for d in due[:5]) + ".")
    if svc:
        parts.append("Service: " + ", ".join(f"{x['vehicle']} {x['item']} ({x['why']})" for x in svc[:5]) + ".")
    if tyres:
        parts.append(f"{len(tyres)} tyre apni life ke paas.")
    if fuel_flags:
        parts.append(f"Diesel par {len(fuel_flags)} ishaare (pichhle 3 din).")
    if challans:
        parts.append(f"{len(challans)} challan khule.")
    if waiting:
        parts.append(f"{len(waiting)} kharche aapki haan ka intezaar.")
    if approved:
        parts.append(f"{len(approved)} maane hue kharche abhi dene baaki.")
    if idle_v:
        parts.append("Khadi gaadiyan: " + ", ".join(x["vehicle"] for x in idle_v[:5]) + ".")
    if workshop:
        parts.append("Workshop mein: " + ", ".join(workshop[:5]) + ".")
    return {"say": " ".join(parts) or "Fleet mein aaj sab theek.", "trips": today_trips, "blocked": blocked, "expired": expired,
            "due": due, "service": svc, "tyres": tyres, "fuel_flags": fuel_flags, "challans": challans, "waiting": waiting,
            "approved": approved, "idle": idle_v, "workshop": workshop}


def overview(tenant: str) -> dict:
    return {"settings": settings(tenant), "vehicles": vehicles(tenant), "day": day(tenant), "trips": trips(tenant)[:30],
            "spends": spends(tenant)[:30], "never": list(R.NEVER), "echallan_steps": R.ECHALLAN_STEPS, "vahan_steps": R.VAHAN_STEPS,
            "doc_kinds": {k: v["label"] for k, v in R.DOC_KINDS.items()}, "spend_kinds": SPEND_KINDS,
            "types": {k: v["label"] for k, v in R.VEHICLE_TYPES.items()}}


# ---------------------------------------------------------------- worker
def _once(s, tenant: str, key: str) -> bool:
    k = f"{tenant}:{key}"[:200]
    if s.get(FleetAlert, k) is not None:
        return False
    s.add(FleetAlert(key=k, tenant_id=tenant))
    return True


def _tenants() -> list[str]:
    with session() as s:
        return sorted(set(s.scalars(select(FleetVehicle.tenant_id)).all()) | set(s.scalars(select(FleetSettings.tenant_id)).all()))


def tick(when: datetime | None = None) -> dict:
    """Worker, every minute: document reminders (once per window), service due / late (once per state), the morning note
    (once a day at her time, only when there is something)."""
    n = _aware(when) or now()
    out = {"docs": 0, "service": 0, "morning": 0}
    for tenant in _tenants():
        try:
            loc = n.astimezone(_tz(tenant))
            on = loc.date()
            alerts = []
            with _LOCK, session() as s:
                vs = {v.id: v for v in s.scalars(select(FleetVehicle).where(FleetVehicle.tenant_id == tenant, FleetVehicle.status != "sold")).all()}
                for d in s.scalars(select(FleetDoc).where(FleetDoc.tenant_id == tenant, FleetDoc.expires != "")).all():
                    v = vs.get(d.vehicle_id)
                    if v is None or d.not_needed:
                        continue
                    x = _doc_view(d, on)
                    if x["window"] is None:
                        continue
                    if _once(s, tenant, f"doc:{d.id}:{d.expires}:{x['window']}"):
                        if x["state"] == "expired":
                            say = f"{v.reg}: {x['label']} {d.expires} ko khatam ho gaya. {x['where']}"
                        else:
                            say = f"{v.reg}: {x['label']} {x['days_left']} din mein khatam ({d.expires}). {x['where']}"
                        alerts.append({"kind": "doc", "vehicle": v.reg, "doc": x, "say": say})
                for v in vs.values():
                    for it in _service_view(s, v, on):
                        if it["state"] in ("due", "late") and _once(s, tenant, f"svc:{it['id']}:{it['last_on']}:{it['last_km']}:{it['state']}"):
                            alerts.append({"kind": "service", "vehicle": v.reg, "item": it,
                                           "say": f"{v.reg}: {it['item']} {'ka samay nikal gaya' if it['state'] == 'late' else 'jaldi'} ({it['why']})."})
            for a in alerts:
                out["docs" if a["kind"] == "doc" else "service"] += 1
                _publish(tenant, "fleet_alert", a)
            if _morning_due(tenant, loc):
                d = day(tenant)
                with session() as s:
                    _st(s, tenant).last_morning = on.isoformat()
                if d["say"] != "Fleet mein aaj sab theek.":
                    _publish(tenant, "fleet_day", _day_card(d))
                    out["morning"] += 1
        except Exception:
            log.exception("fleet tick failed for %s", tenant)
    return out


def _day_card(d: dict) -> dict:
    return {"say": d["say"], "trips": len(d["trips"]), "blocked": len(d["blocked"]), "expired": len(d["expired"]), "due": len(d["due"]),
            "service": len(d["service"]), "waiting": len(d["waiting"]), "idle": [x["vehicle"] for x in d["idle"]][:5]}


def _morning_due(tenant: str, loc: datetime) -> bool:
    with session() as s:
        r = s.get(FleetSettings, tenant)
        if r is None or r.last_morning == loc.date().isoformat():
            return False
        try:
            hh, mm = (int(x) for x in (r.morning_time or "08:30").split(":"))
        except ValueError:
            hh, mm = 8, 30
        return (loc.hour, loc.minute) >= (hh, mm)


def send_morning(tenant: str) -> dict:
    d = day(tenant)
    _publish(tenant, "fleet_day", _day_card(d))
    return {"text": d["say"]}


def register_file(tenant: str, conv: str = "") -> dict:
    """Vehicles + papers + cost per km + trips as an Excel file (her file, pc_file card)."""
    import openpyxl

    from .. import chatfiles
    on = today(tenant)
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    ws = wb.create_sheet("Vehicles")
    ws.append(["Number", "Type", "Make / model", "Capacity (t)", "Body (cbm)", "Fuel", "Odometer", "km/l", "Driver", "Status"])
    for v in vehicles(tenant):
        ws.append([v["reg"], v["type_label"], f"{v['make']} {v['model']}".strip(), v["capacity_t"], v["capacity_cbm"], v["fuel"], v["odo"],
                   v["kmpl"], v["driver"], v["status"]])
    ws = wb.create_sheet("Papers")
    ws.append(["Number", "Paper", "Doc no", "Expires", "Days left", "State"])
    with session() as s:
        for v in s.scalars(select(FleetVehicle).where(FleetVehicle.tenant_id == tenant, FleetVehicle.status != "sold")).all():
            for d in s.scalars(select(FleetDoc).where(FleetDoc.vehicle_id == v.id)).all():
                x = _doc_view(d, on)
                ws.append([v.reg, x["label"], x["number"], x["expires"] + (" (andaaza)" if x["guess"] else ""), x["days_left"], x["state"]])
    ws = wb.create_sheet("Cost per km (30 d)")
    ws.append(["Number", "km", "Fuel", "Service", "Tyre", "Breakdown", "Expense", "Challan", "Total", "Per km", "Idle days"])
    for r in cost_report(tenant)["vehicles"]:
        c = r["cost"]
        ws.append([r["vehicle"], r["km"], c["fuel"], c["service"], c["tyre"], c["breakdown"], c["expense"], c["challan"], r["total"],
                   r["per_km"], r["idle_days"]])
    ws = wb.create_sheet("Trips")
    ws.append(["Trip", "Day", "From", "To", "Vehicle", "Driver", "km", "Status", "Ref"])
    for t in trips(tenant):
        ws.append([t["number"], t["day"], t["origin"], t["destination"], t["vehicle"], t["driver"], t["km"], t["status_label"], t["ref"]])
    bio = io.BytesIO()
    wb.save(bio)
    f = chatfiles.store(tenant, f"Fleet_register_{on.isoformat()}.xlsx", bio.getvalue(),
                        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    _publish(tenant, "pc_file", {**f, "source": "fleet", "title": "Fleet register"}, conv)
    return {"file": f, "say": "Gaadiyan, kaagaz, cost per km aur trips ka Excel."}
