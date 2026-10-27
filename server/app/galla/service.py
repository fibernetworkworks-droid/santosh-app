"""v0.84 Galla: the day's sales of a small shop, without anyone writing them down.

Where the money comes from (official routes only, nothing scraped, no SMS permission):
  * upi_alert  - her phone reads ONLY the 'money received' notifications of the payment apps she picked (PhonePe
                 Business, Paytm for Business, Google Pay for Business, BharatPe...) after she turns on notification
                 access herself; the phone sends just amount, time, app, payer's first name, last 4 of the reference.
  * statement  - a CSV / Excel statement she downloads from her payment app or bank and forwards.
  * cash / manual - what she tells Hoola ("aaj 1,800 cash aaya", "subah 7 se 9 mein 40 chai").
Reused by any shop agent (chai stall first)."""
from __future__ import annotations

import csv
import hashlib
import io
import json
import re
import statistics
from datetime import date, datetime, time, timedelta, timezone
from typing import Any

from sqlalchemy import Boolean, DateTime, Integer, String, Text, select
from sqlalchemy.orm import Mapped, mapped_column

from ..db import Base, new_id, now, session
from .parse import APPS, parse_alert

DAYS = ("Som", "Mangal", "Budh", "Guru", "Shukra", "Shani", "Ravi")


class GallaError(Exception):
    pass


class GallaEntry(Base):
    __tablename__ = "galla_entries"
    id: Mapped[str] = mapped_column(String(40), primary_key=True, default=lambda: new_id("ge"))
    tenant_id: Mapped[str] = mapped_column(String(64), index=True)
    amount_paise: Mapped[int] = mapped_column(Integer, default=0)
    at: Mapped[datetime] = mapped_column(DateTime(timezone=True), index=True)        # UTC
    day: Mapped[str] = mapped_column(String(10), index=True)                         # her local date
    hour: Mapped[int] = mapped_column(Integer, default=-1)                           # her local hour, -1 = not known
    source: Mapped[str] = mapped_column(String(16), default="upi_alert")             # upi_alert | statement | cash | manual
    app: Mapped[str] = mapped_column(String(60), default="")
    payer: Mapped[str] = mapped_column(String(40), default="")
    ref_tail: Mapped[str] = mapped_column(String(8), default="")
    key: Mapped[str] = mapped_column(String(64), default="", index=True)
    note: Mapped[str] = mapped_column(String(200), default="")
    void: Mapped[bool] = mapped_column(Boolean, default=False)
    created_at: Mapped[datetime] = mapped_column(DateTime(timezone=True), default=now)


class GallaSettings(Base):
    __tablename__ = "galla_settings"
    tenant_id: Mapped[str] = mapped_column(String(64), primary_key=True)
    apps: Mapped[str] = mapped_column(Text, default="")                              # package names she picked
    last_alert_at: Mapped[datetime | None] = mapped_column(DateTime(timezone=True), nullable=True)
    alerts_total: Mapped[int] = mapped_column(Integer, default=0)
    open_hour: Mapped[int] = mapped_column(Integer, default=6)
    close_hour: Mapped[int] = mapped_column(Integer, default=22)


# ---------------------------------------------------------------- helpers
def _tz(tenant: str):
    from ..life import tz_of
    return tz_of(tenant)


def _utc(d: datetime | None) -> datetime | None:
    if d is None:
        return None
    return d if d.tzinfo else d.replace(tzinfo=timezone.utc)


def _rs(paise: int) -> str:
    r = int(paise or 0) / 100
    return f"₹{int(r):,}" if r == int(r) else f"₹{r:,.2f}"


def _paise(v) -> int:
    try:
        return int(round(float(str(v).replace(",", "").replace("₹", "").replace("Rs", "").replace("rs", "").strip()) * 100))
    except (TypeError, ValueError):
        raise GallaError(f"'{v}' rupaye ki rakam nahi lagti.")


def today(tenant: str) -> date:
    return datetime.now(_tz(tenant)).date()


def _day(tenant: str, v: str | None) -> date:
    s = str(v or "").strip().lower()
    t = today(tenant)
    if s in ("", "aaj", "today"):
        return t
    if s in ("kal", "yesterday", "kal wala", "beeta kal"):
        return t - timedelta(days=1)
    try:
        return date.fromisoformat(s)
    except ValueError:
        m = re.fullmatch(r"(\d{1,2})[/.-](\d{1,2})(?:[/.-](\d{2,4}))?", s)
        if m:
            y = int(m.group(3)) if m.group(3) else t.year
            return date(y + 2000 if y < 100 else y, int(m.group(2)), int(m.group(1)))
    raise GallaError(f"Din '{v}' samajh nahi aaya (aaj, kal, 2026-10-20, 20/10).")


def _nice(d: date) -> str:
    return f"{DAYS[d.weekday()]} {d.day:02d}/{d.month:02d}"


def _hour_label(h: int) -> str:
    def one(x: int) -> str:
        x %= 24
        return f"{12 if x % 12 == 0 else x % 12} {'subah' if x < 12 else ('dopahar' if x < 16 else ('shaam' if x < 20 else 'raat'))}"
    return f"{one(h)} - {one(h + 1)}"


def _settings(s, tenant: str) -> GallaSettings:
    r = s.get(GallaSettings, tenant)
    if r is None:
        r = GallaSettings(tenant_id=tenant, apps=json.dumps([k for k, v in APPS.items() if v["default"]]))
        s.add(r)
        s.flush()
    return r


def settings(tenant: str) -> dict[str, Any]:
    with session() as s:
        r = _settings(s, tenant)
        picked = json.loads(r.apps or "[]")
        return {"apps": [{"package": k, "label": v["label"], "on": k in picked, "business": v["default"]} for k, v in APPS.items()],
                "picked": picked, "last_alert_at": _utc(r.last_alert_at).isoformat() if r.last_alert_at else None,
                "alerts_total": r.alerts_total, "open_hour": r.open_hour, "close_hour": r.close_hour}


def save_settings(tenant: str, apps: list[str] | None = None, open_hour: int | None = None, close_hour: int | None = None) -> dict:
    with session() as s:
        r = _settings(s, tenant)
        if apps is not None:
            r.apps = json.dumps([a for a in apps if a in APPS])
        for k, v in (("open_hour", open_hour), ("close_hour", close_hour)):
            if v is not None:
                if not 0 <= int(v) <= 24:
                    raise GallaError("Ghanta 0 se 24 ke beech.")
                setattr(r, k, int(v))
    return settings(tenant)


def _add(s, tenant: str, paise: int, at: datetime, source: str, app: str = "", payer: str = "", ref_tail: str = "",
         key: str = "", note: str = "", hour_known: bool = True) -> GallaEntry | None:
    at = _utc(at) or now()
    local = at.astimezone(_tz(tenant))
    if key and s.scalars(select(GallaEntry.id).where(GallaEntry.tenant_id == tenant, GallaEntry.key == key)).first():
        return None
    day = local.date().isoformat()
    if ref_tail and source in ("upi_alert", "statement"):
        dup = s.scalars(select(GallaEntry).where(GallaEntry.tenant_id == tenant, GallaEntry.day == day,
                                                 GallaEntry.ref_tail == ref_tail, GallaEntry.amount_paise == paise,
                                                 GallaEntry.void.is_(False))).first()
        if dup is not None:
            return None
    if source == "upi_alert" and not ref_tail:       # the same alert shown twice (app updates its notification)
        lo, hi = at - timedelta(seconds=90), at + timedelta(seconds=90)
        for e in s.scalars(select(GallaEntry).where(GallaEntry.tenant_id == tenant, GallaEntry.day == day,
                                                    GallaEntry.amount_paise == paise, GallaEntry.app == app,
                                                    GallaEntry.payer == payer, GallaEntry.void.is_(False))).all():
            if lo <= _utc(e.at) <= hi:
                return None
    e = GallaEntry(tenant_id=tenant, amount_paise=paise, at=at, day=day, hour=local.hour if hour_known else -1,
                   source=source, app=app[:60], payer=payer[:40], ref_tail=ref_tail[-4:], key=key[:64], note=note[:200])
    s.add(e)
    return e


# ---------------------------------------------------------------- in: phone alerts, statements, cash
def ingest_alerts(tenant: str, items: list[dict]) -> dict[str, Any]:
    """From her phone (already parsed there). Also accepts raw {title, text} (forwarded alert text) and parses here."""
    added = skipped = 0
    with session() as s:
        st = _settings(s, tenant)
        picked = set(json.loads(st.apps or "[]"))
        for it in (items or [])[:300]:
            app = str(it.get("app") or "")[:60]
            if app and app not in APPS:
                skipped += 1
                continue
            if app and picked and app not in picked:
                skipped += 1
                continue
            if it.get("amount") is None and (it.get("text") or it.get("title")):
                p = parse_alert(str(it.get("title") or ""), str(it.get("text") or ""))
                if not p:
                    skipped += 1
                    continue
                it = {**it, **p}
            try:
                paise = _paise(it.get("amount"))
            except GallaError:
                skipped += 1
                continue
            if not 0 < paise <= 200000 * 100:
                skipped += 1
                continue
            try:
                ms = it.get("at_ms")
                at = datetime.fromtimestamp(int(ms) / 1000, tz=timezone.utc) if ms else \
                    (datetime.fromisoformat(str(it["at"])) if it.get("at") else now())
            except (TypeError, ValueError, OverflowError):
                at = now()
            if _utc(at) > now() + timedelta(minutes=10) or _utc(at) < now() - timedelta(days=8):
                skipped += 1
                continue
            payer = re.sub(r"[^A-Za-z .'-]", "", str(it.get("payer") or ""))[:20]
            ref = re.sub(r"\D", "", str(it.get("ref_tail") or ""))[-4:]
            key = str(it.get("key") or "")[:64] or hashlib.sha1(f"{app}|{at.isoformat()}|{paise}|{ref}".encode()).hexdigest()
            if _add(s, tenant, paise, at, "upi_alert", app, payer, ref, key) is None:
                skipped += 1
                continue
            added += 1
        if added:
            st.last_alert_at = now()
            st.alerts_total = (st.alerts_total or 0) + added
    if added:
        _publish(tenant, "galla", {"added": added})
    return {"added": added, "skipped": skipped}


def add_cash(tenant: str, amount, day: str = "", hour: int | None = None, note: str = "", source: str = "cash") -> dict:
    p = _paise(amount)
    if p <= 0:
        raise GallaError("Kitne rupaye?")
    d = _day(tenant, day)
    tz = _tz(tenant)
    known = hour is not None and 0 <= int(hour) <= 23
    local = datetime.combine(d, time(int(hour) if known else 12, 0), tz)
    with session() as s:
        e = _add(s, tenant, p, local.astimezone(timezone.utc), source if source in ("cash", "manual") else "cash",
                 note=note, hour_known=known)
        s.flush()
        eid = e.id if e else ""
    _publish(tenant, "galla", {"added": 1})
    return {"id": eid, "say": f"{_rs(p)} {'cash' if source == 'cash' else ''} {_nice(d)} mein likh liya."}


def void(tenant: str, entry_id: str) -> dict:
    with session() as s:
        e = s.get(GallaEntry, entry_id)
        if e is None or e.tenant_id != tenant:
            raise GallaError("Entry nahi mili.")
        e.void = True
        return {"ok": True, "say": f"{_rs(e.amount_paise)} hata diya."}


_COLS = {
    "date": ("date", "txn date", "transaction date", "date & time", "date/time", "datetime", "value date", "tarikh"),
    "time": ("time", "txn time", "transaction time"),
    "amount": ("amount", "amount (rs)", "amount(inr)", "amount (inr)", "txn amount", "transaction amount", "credit",
               "credit amount", "deposit", "deposits", "rakam"),
    "type": ("type", "txn type", "transaction type", "cr/dr", "dr/cr", "status", "debit/credit"),
    "name": ("name", "payer", "payer name", "from", "customer name", "sender", "description", "narration", "particulars", "details"),
    "ref": ("utr", "rrn", "upi ref", "reference", "ref no", "reference no", "transaction id", "txn id", "upi transaction id"),
    "debit": ("debit", "withdrawal", "withdrawals", "debit amount"),
}


def _col(headers: list[str], kind: str) -> int:
    hs = [re.sub(r"\s+", " ", h.strip().lower()) for h in headers]
    for i, h in enumerate(hs):
        if h in _COLS[kind]:
            return i
    for i, h in enumerate(hs):
        if any(w in h for w in _COLS[kind] if len(w) > 3):
            return i
    return -1


def _parse_dt(raw: str, tz) -> datetime | None:
    s = str(raw or "").strip()
    for fmt in ("%d/%m/%Y %H:%M:%S", "%d/%m/%Y %H:%M", "%d-%m-%Y %H:%M:%S", "%d-%m-%Y %H:%M", "%Y-%m-%d %H:%M:%S",
                "%Y-%m-%d %H:%M", "%d %b %Y %I:%M %p", "%d %b %Y, %I:%M %p", "%b %d, %Y %I:%M %p", "%d/%m/%Y %I:%M %p",
                "%d-%b-%Y %H:%M:%S", "%d-%b-%Y %I:%M %p", "%Y-%m-%dT%H:%M:%S", "%d/%m/%Y", "%d-%m-%Y", "%Y-%m-%d",
                "%d %b %Y", "%d-%b-%Y", "%d/%m/%y"):
        try:
            d = datetime.strptime(s[:30], fmt)
            return d.replace(tzinfo=tz)
        except ValueError:
            continue
    return None


def import_statement(tenant: str, name: str, data: bytes) -> dict[str, Any]:
    """A CSV / Excel statement from her payment app or bank: only credits are added; repeats are skipped."""
    rows: list[list[str]] = []
    low = (name or "").lower()
    if low.endswith((".xlsx", ".xlsm")):
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        for r in wb.active.iter_rows(values_only=True):
            rows.append(["" if c is None else (c.strftime("%d/%m/%Y %H:%M:%S") if isinstance(c, datetime) else str(c)) for c in r])
    else:
        text = data.decode("utf-8-sig", errors="replace")
        rows = [r for r in csv.reader(io.StringIO(text))]
    hdr_at = next((i for i, r in enumerate(rows[:30]) if _col(r, "amount") >= 0 and (_col(r, "date") >= 0)), -1)
    if hdr_at < 0:
        raise GallaError("Is file mein Date aur Amount wale column nahi mile. CSV / Excel statement bhejiye.")
    h = rows[hdr_at]
    ci = {k: _col(h, k) for k in _COLS}
    tz = _tz(tenant)
    added = skipped = 0
    with session() as s:
        for r in rows[hdr_at + 1:]:
            def g(k):
                i = ci[k]
                return r[i].strip() if 0 <= i < len(r) and r[i] is not None else ""
            typ = g("type").lower()
            if typ and re.search(r"\b(dr|debit|paid|sent|failed|pending|refund)", typ):
                skipped += 1
                continue
            if ci["debit"] >= 0 and ci["debit"] != ci["amount"] and g("debit") not in ("", "0", "0.0", "0.00", "-"):
                skipped += 1
                continue
            raw = g("amount").replace("₹", "").replace(",", "").replace("+", "").strip()
            if not raw or raw.startswith("-"):
                skipped += 1
                continue
            try:
                paise = int(round(float(raw) * 100))
            except ValueError:
                skipped += 1
                continue
            if paise <= 0:
                skipped += 1
                continue
            when = g("date") + (" " + g("time") if g("time") and ":" not in g("date") else "")
            dt = _parse_dt(when, tz)
            if dt is None:
                skipped += 1
                continue
            has_time = bool(re.search(r"\d{1,2}:\d{2}", when))
            ref = re.sub(r"\D", "", g("ref"))[-4:]
            payer = ""
            nm = g("name")
            if nm:
                m = re.search(r"(?:from|by)\s+([A-Za-z][A-Za-z .]+)", nm, re.I)
                payer = re.sub(r"[^A-Za-z]", " ", (m.group(1) if m else nm)).strip().split(" ")[0][:20].title()
            key = hashlib.sha1(f"stmt|{dt.isoformat()}|{paise}|{ref}|{nm[:40]}".encode()).hexdigest()
            if _add(s, tenant, paise, dt.astimezone(timezone.utc), "statement", "statement", payer, ref, key,
                    hour_known=has_time) is None:
                skipped += 1
                continue
            added += 1
    if added:
        _publish(tenant, "galla", {"added": added})
    return {"added": added, "skipped": skipped, "say": f"Statement se {added} payment jode ({skipped} chhode: debit / dobara / galat line)."}


# ---------------------------------------------------------------- out: the day, best hours, trend
def _rows(s, tenant: str, d0: date, d1: date) -> list[GallaEntry]:
    return s.scalars(select(GallaEntry).where(GallaEntry.tenant_id == tenant, GallaEntry.void.is_(False),
                                              GallaEntry.day >= d0.isoformat(), GallaEntry.day <= d1.isoformat())).all()


def totals(tenant: str, days: int = 28, until: date | None = None) -> dict[str, dict]:
    """{YYYY-MM-DD: {upi, cash, total (paise), count}} for the last `days` days (only days with sales)."""
    end = until or today(tenant)
    out: dict[str, dict] = {}
    with session() as s:
        for e in _rows(s, tenant, end - timedelta(days=days - 1), end):
            d = out.setdefault(e.day, {"upi": 0, "cash": 0, "total": 0, "count": 0})
            k = "cash" if e.source in ("cash", "manual") else "upi"
            d[k] += e.amount_paise
            d["total"] += e.amount_paise
            d["count"] += 1 if k == "upi" else 0
    return out


def day_summary(tenant: str, on: str = "") -> dict[str, Any]:
    d = _day(tenant, on)
    with session() as s:
        rows = [e for e in _rows(s, tenant, d, d)]
        st = _settings(s, tenant)
        last_alert = _utc(st.last_alert_at)
    upi = [e for e in rows if e.source not in ("cash", "manual")]
    cash = [e for e in rows if e.source in ("cash", "manual")]
    by_hour: dict[int, dict] = {}
    for e in rows:
        if e.hour >= 0:
            b = by_hour.setdefault(e.hour, {"hour": e.hour, "amount": 0, "count": 0})
            b["amount"] += e.amount_paise
            b["count"] += 1
    hours = [dict(v, label=_hour_label(v["hour"]), rupees=_rs(v["amount"])) for _, v in sorted(by_hour.items())]
    best = sorted(hours, key=lambda x: (x["amount"], x["count"]), reverse=True)[:3]
    tot = sum(e.amount_paise for e in rows)
    # earlier work: the same weekday in the last 4 weeks, and yesterday
    past = totals(tenant, 35, d - timedelta(days=1))
    same = [v["total"] for k, v in past.items() if date.fromisoformat(k).weekday() == d.weekday()][-4:]
    avg_same = int(statistics.mean(same)) if same else 0
    yday = past.get((d - timedelta(days=1)).isoformat(), {}).get("total", 0)
    lines = [f"{_nice(d)}: kul {_rs(tot)} (UPI {_rs(sum(e.amount_paise for e in upi))} · {len(upi)} payment"
             + (f", cash {_rs(sum(e.amount_paise for e in cash))}" if cash else "") + ")."]
    if best:
        lines.append(f"Sabse zyada bikri: {best[0]['label']} ({best[0]['rupees']}, {best[0]['count']} payment).")
    if avg_same:
        diff = round((tot - avg_same) * 100 / avg_same)
        lines.append(f"Pichhle {len(same)} {DAYS[d.weekday()]} ka ausat {_rs(avg_same)}: aaj {'+' if diff >= 0 else ''}{diff}%.")
    if d == today(tenant) and not upi and not last_alert:
        lines.append("UPI alerts abhi phone se nahi aa rahe: Chai screen > Galla mein 'UPI alerts' chalu kijiye.")
    if d == today(tenant) and not cash:
        lines.append("Cash ka hisaab nahi mila: shaam ko bata dijiye ('aaj 1200 cash').")
    return {"date": d.isoformat(), "nice": _nice(d), "total": _rs(tot), "total_paise": tot,
            "upi": _rs(sum(e.amount_paise for e in upi)), "upi_count": len(upi), "cash": _rs(sum(e.amount_paise for e in cash)),
            "cash_paise": sum(e.amount_paise for e in cash), "hours": hours, "best": best,
            "same_weekday_avg": _rs(avg_same) if avg_same else "", "yesterday": _rs(yday) if yday else "",
            "first": min((e.at for e in upi), default=None) and _utc(min(e.at for e in upi)).astimezone(_tz(tenant)).strftime("%H:%M"),
            "last": max((e.at for e in upi), default=None) and _utc(max(e.at for e in upi)).astimezone(_tz(tenant)).strftime("%H:%M"),
            "text": " ".join(lines), "alerts_live": bool(last_alert and now() - last_alert < timedelta(hours=6))}


def best_hours(tenant: str, days: int = 28) -> dict[str, Any]:
    """Which hours sell most over the last weeks (UPI + cash with a known hour), overall and per weekday."""
    end = today(tenant)
    with session() as s:
        rows = [e for e in _rows(s, tenant, end - timedelta(days=days - 1), end) if e.hour >= 0]
    if not rows:
        return {"days": days, "hours": [], "text": "Abhi ghante ke hisaab ka data nahi (UPI alerts chalu kijiye)."}
    n_days = len({e.day for e in rows}) or 1
    agg: dict[int, list[int]] = {}
    wk: dict[int, dict[int, int]] = {}
    for e in rows:
        agg.setdefault(e.hour, [0, 0])
        agg[e.hour][0] += e.amount_paise
        agg[e.hour][1] += 1
        wd = date.fromisoformat(e.day).weekday()
        wk.setdefault(wd, {}).setdefault(e.hour, 0)
        wk[wd][e.hour] += e.amount_paise
    hours = sorted(({"hour": h, "label": _hour_label(h), "avg": _rs(v[0] // n_days), "avg_paise": v[0] // n_days,
                     "payments_per_day": round(v[1] / n_days, 1)} for h, v in agg.items()), key=lambda x: x["hour"])
    top = sorted(hours, key=lambda x: x["avg_paise"], reverse=True)[:3]
    per_day = {DAYS[w]: _hour_label(max(v, key=v.get)) for w, v in sorted(wk.items())}
    text = (f"Pichhle {n_days} din: sabse zyada bikri {top[0]['label']} (roz ~{top[0]['avg']}, ~{top[0]['payments_per_day']} payment)"
            + (f", phir {top[1]['label']}" if len(top) > 1 else "") + ".")
    return {"days": n_days, "hours": hours, "top": top, "per_weekday": per_day, "text": text}


def recent(tenant: str, limit: int = 50) -> list[dict]:
    with session() as s:
        rows = s.scalars(select(GallaEntry).where(GallaEntry.tenant_id == tenant, GallaEntry.void.is_(False))
                         .order_by(GallaEntry.at.desc()).limit(limit)).all()
        tz = _tz(tenant)
        return [{"id": e.id, "amount": _rs(e.amount_paise), "at": _utc(e.at).astimezone(tz).strftime("%d/%m %H:%M"),
                 "source": e.source, "app": APPS.get(e.app, {}).get("label", e.app), "payer": e.payer, "ref_tail": e.ref_tail,
                 "note": e.note} for e in rows]


_LAST_PUB: dict[str, float] = {}


def _publish(tenant: str, kind: str, payload: dict) -> None:
    """A 'refresh' nudge for the open app; at most one a minute per owner (a busy stall gets 100s of payments)."""
    import time as _t
    from ..events import bus
    if kind == "galla" and _t.time() - _LAST_PUB.get(tenant, 0) < 60:
        return
    _LAST_PUB[tenant] = _t.time()
    try:
        from ..net.service import home_conv
        bus.publish(home_conv(tenant), kind, payload)
    except Exception:
        pass
