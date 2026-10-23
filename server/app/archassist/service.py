"""Archassist Sahayak service: an architect's assistant, end to end.

Project file per client (plot + title papers) -> area statement + zoning checks from HER numbers against the rule set she picks
(starter tables, always "rule book se pakka kijiye") -> drawing register with revisions -> approval / NOC checklist by authority ->
consultants, RFIs, meeting minutes -> fee stages and GST invoices.

HER tap (enforced here, `by_tap=True` only from the API): issuing any transmittal (drawings / certificates to a client, authority,
consultant or contractor), turning a submission into "submit karo", numbering / sending an invoice (her fingerprint too once a phone
key is active). Hoola never logs in to BPMS / AutoDCR, never pays a fee, never sends anything itself: drafts and steps only."""
from __future__ import annotations

import hashlib
import io
import json
import logging
import re
import secrets
import threading
import urllib.parse
from datetime import date, datetime, timedelta, timezone
from typing import Any

from sqlalchemy import select

from ..db import now, session
from . import rules as R
from .models import (ArchConsultant, ArchDrawing, ArchInvoice, ArchMeeting, ArchOnce, ArchProject, ArchRfi, ArchSettings,
                     ArchSubmission, ArchTransmittal)

log = logging.getLogger("santosh.archassist")
_LOCK = threading.RLock()


class ArchError(Exception):
    pass


# ---------------------------------------------------------------- small helpers
def _j(v, default):
    try:
        x = json.loads(v) if isinstance(v, str) and v else default
        return x if isinstance(x, type(default)) else default
    except (ValueError, TypeError):
        return default


def _d(v) -> str:
    return json.dumps(v, ensure_ascii=False, default=str)


def clock() -> datetime:
    return now()


def _tz(tenant: str):
    from ..life import tz_of
    return tz_of(tenant)


def today(tenant: str) -> date:
    return clock().astimezone(_tz(tenant)).date()


def _date(v) -> date | None:
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    if not v:
        return None
    s = str(v).strip()[:10]
    for f in ("%Y-%m-%d", "%d-%m-%Y", "%d/%m/%Y", "%d.%m.%Y", "%d-%m-%y", "%d/%m/%y"):
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            pass
    return None


def _iso(v) -> str:
    d = _date(v)
    return d.isoformat() if d else ""


def _clean(text, n: int = 500) -> str:
    """Her / outside words: secrets never kept; Aadhaar / PAN go only by the locker reference."""
    t = str(text or "")
    t = re.sub(r"\b(otp|pin|cvv|password|upi pin)\b\s*(?:is|hai|:|=)?\s*\S+", lambda m: f"{m.group(1)} [hataya]", t, flags=re.I)
    t = re.sub(r"\b(?:\d[ -]?){13,19}\b", "[hataya]", t)
    t = re.sub(r"\b\d{4}\s?\d{4}\s?\d{4}\b", "[hataya]", t)
    t = re.sub(r"\b[A-Z]{5}\d{4}[A-Z]\b", "[PAN hataya]", t)
    return t.strip()[:n]


def _phone(v) -> str:
    d = re.sub(r"\D", "", str(v or ""))
    if len(d) == 12 and d.startswith("91"):
        d = d[2:]
    return d if len(d) == 10 and d[0] in "6789" else ""


def _rs(v) -> str:
    if v is None:
        return "[daam?]"
    try:
        from ..india.money import inr
        return inr(v)
    except Exception:
        return f"Rs {v:,.0f}"


def _hash(show: dict) -> str:
    return hashlib.sha256(json.dumps(show, sort_keys=True, ensure_ascii=False, default=str).encode()).hexdigest()


def _publish(tenant: str, kind: str, payload: dict, conv: str = "") -> None:
    from ..events import bus
    try:
        if not conv:
            from ..net.service import home_conv
            conv = home_conv(tenant)
        bus.publish(conv, kind, payload)
    except Exception:
        log.exception("archassist publish failed")


def is_owner(actor: str) -> bool:
    a = actor or "owner"
    return a == "owner" or a.startswith("approved:")


def _st(s, tenant: str) -> ArchSettings:
    r = s.get(ArchSettings, tenant)
    if r is None:
        r = ArchSettings(tenant_id=tenant)
        s.add(r)
        s.flush()
    return r


def _next(s, tenant: str, key: str) -> int:
    st = _st(s, tenant)
    seq = _j(st.seq, {})
    seq[key] = int(seq.get(key) or 0) + 1
    st.seq = _d(seq)
    return seq[key]


def _once(s, tenant: str, key: str) -> bool:
    k = f"{tenant}:{key}"[:200]
    if s.get(ArchOnce, k) is not None:
        return False
    s.add(ArchOnce(key=k, tenant_id=tenant))
    return True


def _wa(phone: str, text: str) -> str:
    p = _phone(phone)
    return f"https://wa.me/91{p}?text={urllib.parse.quote(text)}" if p else f"https://wa.me/?text={urllib.parse.quote(text)}"


def _mail(to: str, subject: str, text: str) -> str:
    return f"mailto:{urllib.parse.quote(to or '')}?subject={urllib.parse.quote(subject)}&body={urllib.parse.quote(text)}"


def _store(tenant: str, name: str, data: bytes, ctype: str, conv: str = "", title: str = "") -> dict:
    from .. import chatfiles
    f = chatfiles.store(tenant, name, data, ctype)
    _publish(tenant, "pc_file", {**f, "source": "archassist", "title": title or name}, conv)
    return f


def _file_ok(tenant: str, file_id: str) -> bool:
    if not file_id:
        return True
    try:
        from .. import chatfiles
        return chatfiles._meta(tenant, file_id) is not None
    except Exception:
        return False


# ---------------------------------------------------------------- settings
_SET_KEYS = ("firm", "architect", "coa_no", "coa_valid_to", "gstin", "address", "phone", "email", "upi_id", "ruleset", "terms",
             "rules", "morning_time", "inv_prefix", "fee_stages")


def _set_view(r: ArchSettings) -> dict:
    stages = _j(r.fee_stages, [])
    return {"firm": r.firm, "architect": r.architect, "coa_no": r.coa_no, "coa_valid_to": r.coa_valid_to, "gstin": r.gstin,
            "address": r.address, "phone": r.phone, "email": r.email, "upi_id": r.upi_id, "ruleset": r.ruleset, "terms": r.terms,
            "rules": r.rules, "morning_time": r.morning_time, "inv_prefix": r.inv_prefix,
            "fee_stages": stages or [{"key": k, "label": l, "pct": p} for k, l, p in R.FEE_STAGES], "fee_starter": not stages,
            "rulesets": {k: v["label"] for k, v in R.RULESETS.items()},
            "authorities": {k: v["label"] for k, v in R.AUTHORITIES.items()}}


def settings(tenant: str) -> dict:
    with session() as s:
        return _set_view(_st(s, tenant))


def save_settings(tenant: str, **kw) -> dict:
    with _LOCK, session() as s:
        r = _st(s, tenant)
        for k in _SET_KEYS:
            if k not in kw or kw[k] is None:
                continue
            v = kw[k]
            if k == "gstin":
                g = str(v).strip().upper()
                if g:
                    from ..india.money import gstin_ok
                    if not gstin_ok(g):
                        raise ArchError("GSTIN sahi nahi lag raha (check digit).")
                v = g
            elif k == "ruleset":
                if v not in R.RULESETS:
                    raise ArchError("ruleset = " + " | ".join(R.RULESETS))
            elif k == "coa_valid_to":
                v = _iso(v)
            elif k == "phone":
                v = _phone(v) or ""
            elif k == "fee_stages":
                st = v if isinstance(v, list) else _j(v, [])
                clean = []
                for x in st:
                    pct = R.num(x.get("pct")) if isinstance(x, dict) else None
                    if pct is None:
                        continue
                    lab = str(x.get("label") or x.get("key") or "")[:80]
                    clean.append({"key": re.sub(r"\W+", "_", str(x.get("key") or lab).lower())[:20] or f"s{len(clean) + 1}",
                                  "label": lab, "pct": pct})
                if clean and abs(sum(c["pct"] for c in clean) - 100) > 0.01:
                    raise ArchError(f"Stage % ka jod {sum(c['pct'] for c in clean):g} hai; 100 hona chahiye.")
                v = _d(clean) if clean else ""
            else:
                v = _clean(v, 2000) if k in ("terms", "rules", "address") else str(v).strip()[:160]
            setattr(r, k, v)
        out = _set_view(r)
    out["say"] = "Settings save."
    return out


# ---------------------------------------------------------------- projects
_PLOT_KEYS = ("survey_no", "cts_no", "plot_no", "village", "taluka", "area_sqm", "road_width_m", "road_widening_sqm",
              "reservation_sqm", "fsi_on_gross", "zone", "trees", "near_airport", "heritage", "coastal", "railway", "highway",
              "promoter", "units", "congested")


def _proj(s, tenant: str, ref: str) -> ArchProject:
    ref = str(ref or "").strip()
    q = select(ArchProject).where(ArchProject.tenant_id == tenant)
    rows = s.scalars(q.order_by(ArchProject.created_at.desc())).all()
    if not ref:
        act = [p for p in rows if p.status == "active"]
        if len(act) == 1:
            return act[0]
        raise ArchError("Kaun sa project? (number / naam)")
    for p in rows:
        if p.id == ref or p.number.lower() == ref.lower():
            return p
    hits = [p for p in rows if ref.lower() in (p.name + " " + p.client).lower()]
    if len(hits) == 1:
        return hits[0]
    if hits:
        raise ArchError("Kai project milte hain: " + ", ".join(f"{p.number} {p.name}" for p in hits[:5]))
    raise ArchError(f"Project '{ref}' nahi mila.")


def _docs_view(p: ArchProject) -> list[dict]:
    d = _j(p.docs, {})
    out = []
    for k, label in R.PLOT_DOCS:
        x = d.get(k) or {}
        out.append({"key": k, "label": label, "state": x.get("state", "missing"), "file_id": x.get("file_id", ""), "note": x.get("note", "")})
    for k, x in d.items():
        if k not in dict(R.PLOT_DOCS):
            out.append({"key": k, "label": x.get("label") or k, "state": x.get("state", "missing"), "file_id": x.get("file_id", ""),
                        "note": x.get("note", "")})
    return out


def _proj_view(p: ArchProject, full: bool = False) -> dict:
    plot = _j(p.plot, {})
    v = {"id": p.id, "number": p.number, "name": p.name, "client": p.client, "city": p.city, "pin": p.pin, "status": p.status,
         "authority": p.authority, "authority_label": R.AUTHORITIES.get(p.authority, {}).get("label", p.authority),
         "ruleset": p.ruleset, "ruleset_label": R.RULESETS.get(p.ruleset, {}).get("label", p.ruleset), "use": p.use,
         "plot_area_sqm": plot.get("area_sqm"), "road_width_m": plot.get("road_width_m"), "rera_no": p.rera_no}
    docs = _docs_view(p)
    v["docs_missing"] = [x["label"] for x in docs if x["state"] == "missing"]
    if full:
        v |= {"client_phone": p.client_phone, "client_email": p.client_email, "client_gstin": p.client_gstin,
              "client_pan_ref": p.client_pan_ref, "site": p.site, "plot": plot, "proposal": _j(p.proposal, {}),
              "overrides": _j(p.overrides, {}), "docs": docs, "area": _j(p.area, {}), "fee": _j(p.fee, {}), "notes": p.notes}
    return v


def save_project(tenant: str, name: str = "", ref: str = "", **kw) -> dict:
    with _LOCK, session() as s:
        if ref:
            p = _proj(s, tenant, ref)
            made = False
        else:
            if not str(name or "").strip():
                raise ArchError("Project ka naam chahiye.")
            st = _st(s, tenant)
            p = ArchProject(tenant_id=tenant, number=f"P-{_next(s, tenant, 'project'):03d}", name=_clean(name, 160),
                            ruleset=st.ruleset or "udcpr")
            p.docs = _d({k: {"state": "missing"} for k, _ in R.PLOT_DOCS})
            s.add(p)
            made = True
        if name and not made:
            p.name = _clean(name, 160)
        for k in ("client", "client_email", "site", "city", "notes", "rera_no"):
            if kw.get(k) is not None:
                setattr(p, k, _clean(kw[k], 1000 if k in ("site", "notes") else 160))
        if kw.get("client_phone") is not None:
            ph = _phone(kw["client_phone"])
            if kw["client_phone"] and not ph:
                raise ArchError("Client ka phone 10 ank ka mobile hona chahiye.")
            p.client_phone = ph
        if kw.get("client_gstin"):
            from ..india.money import gstin_ok
            g = str(kw["client_gstin"]).strip().upper()
            if not gstin_ok(g):
                raise ArchError("Client ka GSTIN sahi nahi lag raha.")
            p.client_gstin = g
        if kw.get("client_pan_ref") is not None:
            ref_s = str(kw["client_pan_ref"]).strip()
            if ref_s and not re.fullmatch(r"\{id:[A-Za-z0-9_\-]{3,40}\}", ref_s):
                raise ArchError("PAN sirf Pehchaan Locker ke {id:...} reference se.")
            p.client_pan_ref = ref_s
        if kw.get("pin") is not None:
            pin = re.sub(r"\D", "", str(kw["pin"]))
            if pin:
                from ..india import pincode
                if not pincode.valid(pin):
                    raise ArchError("Yeh PIN code India Post ki list mein nahi hai.")
            p.pin = pin
        if kw.get("authority"):
            if kw["authority"] not in R.AUTHORITIES:
                raise ArchError("authority = " + " | ".join(R.AUTHORITIES))
            p.authority = kw["authority"]
        if kw.get("ruleset"):
            if kw["ruleset"] not in R.RULESETS:
                raise ArchError("ruleset = " + " | ".join(R.RULESETS))
            p.ruleset = kw["ruleset"]
        if kw.get("use"):
            p.use = str(kw["use"])[:20]
        if kw.get("status") in ("active", "on_hold", "done", "dropped"):
            p.status = kw["status"]
        plot = _j(p.plot, {})
        src = dict(kw.get("plot") or {})
        src |= {k: kw[k] for k in _PLOT_KEYS if k in kw and kw[k] is not None}
        if "area" in kw and kw["area"] not in (None, ""):
            src["area_sqm"] = R.area_sqm(kw["area"], str(kw.get("area_unit") or ""))
        for k, v in src.items():
            if k not in _PLOT_KEYS:
                continue
            if k in ("area_sqm", "road_width_m", "road_widening_sqm", "reservation_sqm"):
                n = R.area_sqm(v) if k.endswith("sqm") else R.num(v)
                if n is not None and n < 0:
                    raise ArchError(f"{k} minus nahi ho sakta.")
                plot[k] = n
            elif k in ("trees", "near_airport", "heritage", "coastal", "railway", "highway", "promoter", "fsi_on_gross", "congested"):
                plot[k] = bool(v)
            elif k == "units":
                plot[k] = int(R.num(v) or 0)
            else:
                plot[k] = _clean(v, 80)
        p.plot = _d(plot)
        s.flush()
        v = _proj_view(p, True)
    v["say"] = (f"{v['number']} {'bana' if made else 'update'}: {v['name']}." +
                (f" Kaagaz baaki: {', '.join(v['docs_missing'][:4])}." if v["docs_missing"] else ""))
    return v


def projects(tenant: str, status: str = "") -> list[dict]:
    with session() as s:
        q = select(ArchProject).where(ArchProject.tenant_id == tenant)
        if status:
            q = q.where(ArchProject.status == status)
        return [_proj_view(p) for p in s.scalars(q.order_by(ArchProject.created_at.desc())).all()[:100]]


def project_view(tenant: str, ref: str) -> dict:
    with session() as s:
        p = _proj(s, tenant, ref)
        v = _proj_view(p, True)
        pid = p.id
    v["drawings"] = drawings(tenant, pid)
    v["approvals"] = approvals(tenant, pid)
    v["rfis"] = rfis(tenant, pid)
    v["invoices"] = invoices(tenant, pid)
    with session() as s:
        v["transmittals"] = [_tr_view(x) for x in s.scalars(select(ArchTransmittal).where(
            ArchTransmittal.project_id == pid).order_by(ArchTransmittal.created_at.desc())).all()[:30]]
        v["meetings"] = [_mt_view(x) for x in s.scalars(select(ArchMeeting).where(
            ArchMeeting.project_id == pid).order_by(ArchMeeting.created_at.desc())).all()[:20]]
    return v


def doc(tenant: str, ref: str, key: str, state: str = "have", file_id: str = "", note: str = "", label: str = "") -> dict:
    if state not in ("have", "missing", "na"):
        raise ArchError("state = have | missing | na")
    if file_id and not _file_ok(tenant, file_id):
        raise ArchError("Yeh file nahi mili (file_id).")
    with _LOCK, session() as s:
        p = _proj(s, tenant, ref)
        d = _j(p.docs, {})
        k = str(key or "").strip()
        known = {lab.lower(): kk for kk, lab in R.PLOT_DOCS}
        if k not in dict(R.PLOT_DOCS) and k not in d:
            hit = next((kk for lab, kk in known.items() if k.lower() in lab), None)
            if hit:
                k = hit
            else:
                k = re.sub(r"\W+", "_", k.lower())[:30] or "other"
                d.setdefault(k, {})["label"] = _clean(label or key, 80)
        d.setdefault(k, {}).update({"state": state, "file_id": file_id or d.get(k, {}).get("file_id", ""), "note": _clean(note, 200)})
        p.docs = _d(d)
        v = _proj_view(p, True)
    return {"docs": v["docs"], "say": f"{k}: {state}." + (f" Baaki: {', '.join(v['docs_missing'][:4])}." if v["docs_missing"] else " Saare kaagaz aa gaye.")}


def forget(tenant: str, ref: str) -> dict:
    """DPDP: after the project ends, the client's personal contact details are erased (drawings / numbers stay)."""
    with _LOCK, session() as s:
        p = _proj(s, tenant, ref)
        if p.status not in ("done", "dropped"):
            raise ArchError("Pehle project 'done' ya 'dropped' kijiye.")
        p.client_phone = p.client_email = p.client_pan_ref = ""
        p.forgotten = True
        return {"say": f"{p.number}: client ka phone / email / PAN reference mita diya."}


# ---------------------------------------------------------------- area statement + zoning check
_PROP_KEYS = ("floors", "bua_sqm", "height_m", "front_m", "side_m", "rear_m", "units", "parking", "footprint_sqm", "open_space_sqm",
              "amenity_sqm", "use")


def area_check(tenant: str, ref: str, proposal: dict | None = None, overrides: dict | None = None, conv: str = "") -> dict:
    with _LOCK, session() as s:
        p = _proj(s, tenant, ref)
        prop = _j(p.proposal, {})
        for k, v in (proposal or {}).items():
            if k in _PROP_KEYS and v is not None:
                prop[k] = v
        prop.setdefault("use", p.use)
        ov = _j(p.overrides, {})
        for k, v in (overrides or {}).items():
            if v is None:
                continue
            if v == "":
                ov.pop(k, None)
            elif R.num(v) is not None:
                ov[k] = R.num(v)
        p.proposal, p.overrides = _d(prop), _d(ov)
        plot = _j(p.plot, {})
        res = R.area_statement(plot, prop, R.ruleset(p.ruleset, ov))
        res["at"] = today(tenant).isoformat()
        p.area = _d(res)
        # NOCs the numbers now point to (added as 'needed'; her own statuses are kept)
        ap = _j(p.approvals, {})
        need = R.nocs({**plot, "use": prop.get("use")}, R.num(prop.get("height_m")), res.get("statement", {}).get("proposed_bua"))
        for n in need:
            cur = ap.get(n["key"]) or {}
            if cur.get("status") in (None, "", "not_needed_auto"):
                ap[n["key"]] = {**cur, "status": "needed", "why": n["why"]}
            else:
                cur["why"] = n["why"]
                ap[n["key"]] = cur
        p.approvals = _d(ap)
        res["project"] = p.number
        res["nocs"] = need
    _publish(tenant, "archassist_area", {"project": res["project"], "ok": res["ok"], "say": res["say"],
                                         "checks": res["checks"][:12], "statement": res.get("statement", {})}, conv)
    return res


# ---------------------------------------------------------------- approvals + submissions
def approvals(tenant: str, ref: str) -> dict:
    with session() as s:
        p = _proj(s, tenant, ref)
        ap = _j(p.approvals, {})
        auth = R.AUTHORITIES.get(p.authority) or R.AUTHORITIES["custom"]
        items = []
        for k in R.STAGES + [k for k in ap if k not in R.STAGES]:
            x = ap.get(k) or {}
            items.append({"key": k, "label": R.STAGE_LABEL.get(k, k), "status": x.get("status", "pending"), "ref": x.get("ref", ""),
                          "on": x.get("on", ""), "valid_to": x.get("valid_to", ""), "why": x.get("why", ""), "file_id": x.get("file_id", "")})
        subs = [_sub_view(x) for x in s.scalars(select(ArchSubmission).where(ArchSubmission.project_id == p.id)
                                                 .order_by(ArchSubmission.created_at.desc())).all()[:30]]
    return {"authority": auth["label"], "portal": auth["portal"], "steps": auth["steps"], "items": items, "submissions": subs,
            "note": "Portal par login, scrutiny fee aur OTP aap khud. Hoola sirf kaagaz aur steps taiyaar karta hai."}


_APPROVAL_STATES = ("pending", "needed", "applied", "query", "granted", "refused", "not_needed")


def approval_set(tenant: str, ref: str, key: str, status: str, ref_no: str = "", on: str = "", valid_to: str = "",
                 file_id: str = "", note: str = "") -> dict:
    if status not in _APPROVAL_STATES:
        raise ArchError("status = " + " | ".join(_APPROVAL_STATES))
    if file_id and not _file_ok(tenant, file_id):
        raise ArchError("Yeh file nahi mili (file_id).")
    k = key if key in R.STAGE_LABEL else re.sub(r"\W+", "_", str(key).lower())[:30]
    with _LOCK, session() as s:
        p = _proj(s, tenant, ref)
        ap = _j(p.approvals, {})
        cur = ap.get(k) or {}
        cur |= {"status": status, "ref": _clean(ref_no, 80) or cur.get("ref", ""), "on": _iso(on) or cur.get("on", ""),
                "valid_to": _iso(valid_to) or cur.get("valid_to", ""), "file_id": file_id or cur.get("file_id", "")}
        if note:
            cur["note"] = _clean(note, 300)
        ap[k] = cur
        p.approvals = _d(ap)
    return {"say": f"{R.STAGE_LABEL.get(k, k)}: {status}" + (f" ({cur['ref']})" if cur.get("ref") else "") + "."}


_PREREQ = {"cc": ["building_permission"], "plinth": ["cc"], "oc": ["cc"]}


def _sub_view(x: ArchSubmission) -> dict:
    snap = _j(x.snapshot, {})
    return {"id": x.id, "number": x.number, "stage": x.stage, "stage_label": R.STAGE_LABEL.get(x.stage, x.stage),
            "authority": x.authority, "status": x.status, "ref": x.ref, "on": x.on, "valid_to": x.valid_to, "hash": x.hash,
            "blockers": snap.get("blockers", []), "warnings": snap.get("warnings", []), "docs": snap.get("docs", []),
            "drawings": snap.get("drawings", []), "steps": snap.get("steps", []), "project": snap.get("project", ""),
            "note": x.note}


def _sub_snapshot(s, tenant: str, p: ArchProject, stage: str) -> dict:
    docs = _docs_view(p)
    ap = _j(p.approvals, {})
    area = _j(p.area, {})
    blockers, warnings = [], []
    for d in docs:
        if d["state"] == "missing" and d["key"] in ("title", "property_card", "measurement", "zone", "deed"):
            blockers.append(f"Kaagaz baaki: {d['label']}")
        elif d["state"] == "missing":
            warnings.append(f"Kaagaz: {d['label']} (zaroori ho to)")
    for pre in _PREREQ.get(stage, []):
        if (ap.get(pre) or {}).get("status") != "granted":
            blockers.append(f"Pehle {R.STAGE_LABEL[pre]} milna chahiye.")
    if stage == "building_permission":
        if not area:
            blockers.append("Area statement abhi nahi bana (archassist action=area).")
        else:
            for c in area.get("checks", []):
                if c["status"] == "fail":
                    blockers.append(f"Niyam se bahar: {c['item']} ({c.get('got')} vs {c.get('need')})")
                elif c["status"] == "check":
                    warnings.append(f"Dekhna hai: {c['item']}")
    if stage.startswith("noc_") and (ap.get(stage) or {}).get("status") == "not_needed":
        warnings.append("Aapne yeh NOC 'zaroori nahi' kaha tha.")
    if stage == "oc":
        for k, x in ap.items():
            if k.startswith("noc_") and x.get("status") not in ("granted", "not_needed"):
                blockers.append(f"{R.STAGE_LABEL.get(k, k)} abhi '{x.get('status')}' hai.")
    st = _st(s, tenant)
    if not st.coa_no:
        blockers.append("Aapka CoA registration number settings mein daaliye (har submission par lagta hai).")
    elif st.coa_valid_to and _date(st.coa_valid_to) and _date(st.coa_valid_to) < today(tenant):
        blockers.append("CoA registration ki tareekh nikal gayi: renew kijiye.")
    drs = s.scalars(select(ArchDrawing).where(ArchDrawing.project_id == p.id, ArchDrawing.status != "superseded")).all()
    dview = [{"sheet": d.sheet, "rev": d.rev, "title": d.title, "has_file": bool((_j(d.revisions, []) or [{}])[-1].get("file_id"))}
             for d in sorted(drs, key=lambda d: d.sheet)]
    if not dview and stage in ("building_permission", "oc"):
        blockers.append("Drawing register khaali: submission ki sheets daaliye.")
    for d in dview:
        if not d["has_file"]:
            warnings.append(f"Sheet {d['sheet']} {d['rev']} ki file nahi judi.")
    auth = R.AUTHORITIES.get(p.authority) or R.AUTHORITIES["custom"]
    return {"project": f"{p.number} {p.name}", "stage": stage, "authority": auth["label"], "docs": [d["label"] for d in docs if d["state"] == "have"],
            "drawings": dview, "area": area.get("statement", {}), "blockers": blockers, "warnings": warnings, "steps": auth["steps"]}


def submission_request(tenant: str, ref: str, stage: str, conv: str = "", note: str = "") -> dict:
    """Prepares the submission card. 'Submit karo' becomes true ONLY on her tap (decide_submission from the API)."""
    if stage not in R.STAGE_LABEL:
        raise ArchError("stage = " + " | ".join(R.STAGE_LABEL))
    with _LOCK, session() as s:
        p = _proj(s, tenant, ref)
        snap = _sub_snapshot(s, tenant, p, stage)
        for old in s.scalars(select(ArchSubmission).where(ArchSubmission.project_id == p.id, ArchSubmission.stage == stage,
                                                          ArchSubmission.status == "waiting")).all():
            old.status = "stale"
        x = ArchSubmission(tenant_id=tenant, project_id=p.id, number=f"S-{_next(s, tenant, 'sub'):03d}", stage=stage,
                           authority=p.authority, snapshot=_d(snap), note=_clean(note, 300))
        x.hash = _hash({"stage": stage, "project": p.id, "snap": snap})
        s.add(x)
        s.flush()
        v = _sub_view(x)
    v["say"] = (f"{v['number']} {v['stage_label']}: " + (f"{len(v['blockers'])} rukawat - pehle yeh: {v['blockers'][0]}."
                                                          if v["blockers"] else "kaagaz taiyaar. Aapki haan ka card bheja hai."))
    _publish(tenant, "archassist_submission", {"submission": v, "say": v["say"]}, conv)
    return v


def decide_submission(tenant: str, sub_id: str, yes: bool, card_hash: str = "", reason: str = "", by_tap: bool = False) -> dict:
    if not by_tap:
        raise ArchError("Submission ki haan sirf aapke tap se (app ka card).")
    with _LOCK, session() as s:
        x = s.get(ArchSubmission, sub_id)
        if x is None or x.tenant_id != tenant:
            raise ArchError("Card nahi mila.")
        if x.status != "waiting":
            return {"submission": _sub_view(x), "say": f"{x.number} pehle hi '{x.status}' hai."}
        p = s.get(ArchProject, x.project_id)
        snap_now = _sub_snapshot(s, tenant, p, x.stage)
        now_hash = _hash({"stage": x.stage, "project": p.id, "snap": snap_now})
        if not card_hash or card_hash != x.hash or now_hash != x.hash:
            x.status = "stale"
            raise ArchError("Project badal gaya hai; naya card banaiye.")
        if yes and snap_now["blockers"]:
            raise ArchError("Rukawat baaki hai: " + snap_now["blockers"][0])
        x.status = "approved" if yes else "rejected"
        x.note = _clean(reason, 300) or x.note
        x.decided_at = now()
        v = _sub_view(x)
    if yes:
        v["say"] = (f"{v['number']} maana. Ab {v['authority']} portal par aap login karke file kijiye (fee / OTP aap khud): "
                    + " -> ".join(v["steps"][:4]) + ". Inward / application number milte hi batana.")
    else:
        v["say"] = f"{v['number']}: abhi submit nahi."
    return v


def submission_update(tenant: str, sub_id: str, status: str, ref_no: str = "", on: str = "", valid_to: str = "", note: str = "") -> dict:
    if status not in ("submitted", "query", "granted", "refused"):
        raise ArchError("status = submitted | query | granted | refused")
    with _LOCK, session() as s:
        x = s.get(ArchSubmission, sub_id)
        if x is None or x.tenant_id != tenant:
            raise ArchError("Submission nahi mila.")
        if x.status in ("waiting", "rejected", "stale"):
            raise ArchError(f"{x.number} abhi '{x.status}' hai: pehle aapki haan (card par tap).")
        x.status, x.ref = status, _clean(ref_no, 80) or x.ref
        x.on = _iso(on) or x.on or today(tenant).isoformat()
        x.valid_to = _iso(valid_to) or x.valid_to
        if note:
            x.note = _clean(note, 300)
        p = s.get(ArchProject, x.project_id)
        ap = _j(p.approvals, {})
        cur = ap.get(x.stage) or {}
        cur |= {"status": {"submitted": "applied"}.get(status, status), "ref": x.ref, "on": x.on, "valid_to": x.valid_to}
        ap[x.stage] = cur
        p.approvals = _d(ap)
        v = _sub_view(x)
    v["say"] = f"{v['number']} {v['stage_label']}: {status}" + (f" ({v['ref']})" if v["ref"] else "") + "."
    return v


# ---------------------------------------------------------------- drawing register
def _dr_view(d: ArchDrawing) -> dict:
    revs = _j(d.revisions, [])
    return {"id": d.id, "sheet": d.sheet, "title": d.title, "discipline": d.discipline,
            "discipline_label": R.DISCIPLINES.get(d.discipline, d.discipline), "scale": d.scale, "rev": d.rev, "status": d.status,
            "revisions": revs, "issued": _j(d.issued, []), "file_id": (revs[-1].get("file_id", "") if revs else "")}


def drawing(tenant: str, ref: str, sheet: str, title: str = "", discipline: str = "", scale: str = "", status: str = "",
            rev: str = "", file_id: str = "", note: str = "", by: str = "Owner", revise: bool = False) -> dict:
    sheet = str(sheet or "").strip().upper()
    if not R.sheet_ok(sheet):
        raise ArchError("Sheet number jaise A-101 / S-02 (24 akshar tak).")
    if file_id and not _file_ok(tenant, file_id):
        raise ArchError("Yeh file nahi mili (file_id).")
    if status and status not in ("draft", "for_approval", "for_construction", "superseded"):
        raise ArchError("status = draft | for_approval | for_construction | superseded")
    with _LOCK, session() as s:
        p = _proj(s, tenant, ref)
        d = s.scalars(select(ArchDrawing).where(ArchDrawing.project_id == p.id, ArchDrawing.sheet == sheet)).first()
        on = today(tenant).isoformat()
        if d is None:
            disc = (discipline or sheet[:1]).upper()[:1]
            d = ArchDrawing(tenant_id=tenant, project_id=p.id, sheet=sheet, title=_clean(title, 160), scale=_clean(scale, 20),
                            discipline=disc if disc in R.DISCIPLINES else "A", rev=(rev or "R0")[:10], status=status or "draft")
            d.revisions = _d([{"rev": d.rev, "on": on, "note": _clean(note, 200) or "Pehli baar", "file_id": file_id, "by": by}])
            s.add(d)
            made = True
        else:
            made = False
            if title:
                d.title = _clean(title, 160)
            if scale:
                d.scale = _clean(scale, 20)
            if status:
                d.status = status
            revs = _j(d.revisions, [])
            if revise or (rev and rev != d.rev):
                new = (rev or R.next_rev(d.rev))[:10]
                if any(r.get("rev") == new for r in revs):
                    raise ArchError(f"Revision {new} pehle se hai.")
                if not note:
                    raise ArchError("Revision mein kya badla, likhiye (note).")
                revs.append({"rev": new, "on": on, "note": _clean(note, 200), "file_id": file_id, "by": by})
                d.rev = new
                # the waiting transmittals that carry this sheet are now out of date
                for t in s.scalars(select(ArchTransmittal).where(ArchTransmittal.project_id == p.id,
                                                                 ArchTransmittal.status == "waiting")).all():
                    if any(i.get("sheet") == sheet for i in _j(t.items, [])):
                        t.status = "stale"
            elif file_id:
                revs[-1]["file_id"] = file_id
            d.revisions = _d(revs)
        s.flush()
        v = _dr_view(d)
    v["say"] = f"Sheet {v['sheet']} {'register mein' if made else 'update'}: rev {v['rev']}."
    return v


def drawings(tenant: str, ref: str) -> list[dict]:
    with session() as s:
        p = _proj(s, tenant, ref)
        return [_dr_view(d) for d in sorted(s.scalars(select(ArchDrawing).where(ArchDrawing.project_id == p.id)).all(),
                                            key=lambda d: d.sheet)]


# ---------------------------------------------------------------- transmittals (HER tap issues them)
def _tr_view(t: ArchTransmittal) -> dict:
    return {"id": t.id, "number": t.number, "to_kind": t.to_kind, "to_label": R.TO_KINDS.get(t.to_kind, t.to_kind), "to_name": t.to_name,
            "to_contact": t.to_contact, "purpose": t.purpose, "purpose_label": R.PURPOSES.get(t.purpose, t.purpose),
            "items": _j(t.items, []), "note": t.note, "hash": t.hash, "status": t.status, "file_id": t.file_id, "reason": t.reason}


def _tr_show(t: ArchTransmittal) -> dict:
    return {"to": [t.to_kind, t.to_name, t.to_contact], "purpose": t.purpose, "note": t.note,
            "items": [[i.get("sheet"), i.get("rev"), i.get("file_id"), i.get("copies")] for i in _j(t.items, [])]}


def transmittal_request(tenant: str, ref: str, to_kind: str, to_name: str, sheets: list | None = None, purpose: str = "for_approval",
                        to_contact: str = "", note: str = "", copies: int = 1, conv: str = "") -> dict:
    if to_kind not in R.TO_KINDS:
        raise ArchError("to_kind = " + " | ".join(R.TO_KINDS))
    if purpose not in R.PURPOSES:
        raise ArchError("purpose = " + " | ".join(R.PURPOSES))
    if not str(to_name or "").strip():
        raise ArchError("Kise bhejna hai (naam)?")
    with _LOCK, session() as s:
        p = _proj(s, tenant, ref)
        drs = {d.sheet: d for d in s.scalars(select(ArchDrawing).where(ArchDrawing.project_id == p.id)).all()}
        want = [str(x).strip().upper() for x in (sheets or []) if str(x).strip()]
        if not want:
            raise ArchError("Kaun si sheets? (sheet numbers)")
        items, warn = [], []
        for sh in want:
            d = drs.get(sh)
            if d is None:
                raise ArchError(f"Sheet {sh} register mein nahi.")
            if d.status == "superseded":
                raise ArchError(f"Sheet {sh} superseded hai.")
            revs = _j(d.revisions, [])
            fid = revs[-1].get("file_id", "") if revs else ""
            if not fid:
                warn.append(f"{sh} {d.rev} ki file nahi judi")
            if purpose == "for_construction" and d.status != "for_construction":
                warn.append(f"{sh} abhi '{d.status}' hai, construction ke liye nahi")
            items.append({"drawing_id": d.id, "sheet": d.sheet, "rev": d.rev, "title": d.title, "file_id": fid,
                          "copies": max(1, int(copies or 1)), "format": "PDF"})
        if to_kind == "authority" and not _st(s, tenant).coa_no:
            warn.append("CoA number settings mein nahi (authority ko jaati sheet par chahiye)")
        t = ArchTransmittal(tenant_id=tenant, project_id=p.id, number=f"T-{_next(s, tenant, 'transmittal'):04d}", to_kind=to_kind,
                            to_name=_clean(to_name, 160), to_contact=_clean(to_contact, 120), purpose=purpose, items=_d(items),
                            note=_clean(note, 500), status="waiting")
        t.hash = _hash(_tr_show(t))
        s.add(t)
        s.flush()
        v = _tr_view(t)
        v["project"] = f"{p.number} {p.name}"
    v["warnings"] = warn
    v["say"] = (f"{v['number']}: {len(items)} sheet {v['to_label']} {v['to_name']} ko ({v['purpose_label']}). Aapki haan ka card bheja hai"
                + (f"; dhyaan: {'; '.join(warn[:3])}" if warn else "") + ".")
    _publish(tenant, "archassist_transmittal", {"transmittal": v, "say": v["say"]}, conv)
    return v


def _tr_pdf(st: ArchSettings, p: ArchProject, t: ArchTransmittal, on: str) -> bytes:
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.units import mm
    from reportlab.pdfgen import canvas
    bio = io.BytesIO()
    c = canvas.Canvas(bio, pagesize=A4)
    w, h = A4
    y = h - 20 * mm
    c.setFont("Helvetica-Bold", 14)
    c.drawString(20 * mm, y, (st.firm or "Architect")[:70])
    c.setFont("Helvetica", 9)
    y -= 6 * mm
    c.drawString(20 * mm, y, f"{st.architect}  CoA Reg. No. {st.coa_no or '[  ]'}"[:110])
    y -= 10 * mm
    c.setFont("Helvetica-Bold", 12)
    c.drawString(20 * mm, y, f"DRAWING TRANSMITTAL  {t.number}")
    c.setFont("Helvetica", 10)
    for line in (f"Date: {on}", f"Project: {p.number} {p.name}", f"To: {t.to_name} ({R.TO_KINDS.get(t.to_kind, '')}) {t.to_contact}",
                 f"Purpose: {R.PURPOSES.get(t.purpose, t.purpose)}"):
        y -= 6 * mm
        c.drawString(20 * mm, y, line[:110])
    y -= 10 * mm
    c.setFont("Helvetica-Bold", 10)
    for x, txt in ((20, "Sheet"), (50, "Rev"), (65, "Title"), (160, "Copies"), (178, "Format")):
        c.drawString(x * mm, y, txt)
    c.setFont("Helvetica", 10)
    for i in _j(t.items, []):
        y -= 6 * mm
        if y < 30 * mm:
            c.showPage()
            y = h - 20 * mm
        for x, txt in ((20, i["sheet"]), (50, i["rev"]), (65, (i.get("title") or "")[:52]), (160, str(i.get("copies", 1))), (178, i.get("format", "PDF"))):
            c.drawString(x * mm, y, str(txt))
    if t.note:
        y -= 10 * mm
        c.drawString(20 * mm, y, ("Note: " + t.note)[:110])
    y -= 20 * mm
    c.drawString(20 * mm, y, "Received by: ____________________   Date: __________")
    c.showPage()
    c.save()
    return bio.getvalue()


def decide_transmittal(tenant: str, tr_id: str, yes: bool, card_hash: str = "", reason: str = "", by_tap: bool = False,
                       conv: str = "") -> dict:
    if not by_tap:
        raise ArchError("Drawing issue karne ki haan sirf aapke tap se (app ka card).")
    with _LOCK, session() as s:
        t = s.get(ArchTransmittal, tr_id)
        if t is None or t.tenant_id != tenant:
            raise ArchError("Card nahi mila.")
        if t.status != "waiting":
            return {"transmittal": _tr_view(t), "say": f"{t.number} pehle hi '{t.status}' hai."}
        if not card_hash or card_hash != t.hash or _hash(_tr_show(t)) != t.hash:
            t.status = "stale"
            raise ArchError("Card badal gaya hai; naya card banaiye.")
        items = _j(t.items, [])
        for i in items:                                   # a sheet revised after the card = old card
            d = s.get(ArchDrawing, i["drawing_id"])
            if d is None or d.rev != i["rev"]:
                t.status = "stale"
                raise ArchError(f"Sheet {i['sheet']} ki revision badal gayi; naya card banaiye.")
        t.decided_at = now()
        t.reason = _clean(reason, 300)
        if not yes:
            t.status = "rejected"
            return {"transmittal": _tr_view(t), "say": f"{t.number}: issue nahi kiya."}
        on = today(tenant).isoformat()
        st, p = _st(s, tenant), s.get(ArchProject, t.project_id)
        pdf = _tr_pdf(st, p, t, on)
        t.status = "issued"
        for i in items:
            d = s.get(ArchDrawing, i["drawing_id"])
            iss = _j(d.issued, [])
            iss.append({"transmittal": t.number, "rev": i["rev"], "to": t.to_name, "purpose": t.purpose, "on": on})
            d.issued = _d(iss)
        num, to_name, to_contact = t.number, t.to_name, t.to_contact
        pname = f"{p.number} {p.name}"
        sheets = ", ".join(f"{i['sheet']} rev {i['rev']}" for i in items)
    f = _store(tenant, f"Transmittal_{num}.pdf", pdf, "application/pdf", conv, f"Transmittal {num}")
    with session() as s:
        s.get(ArchTransmittal, tr_id).file_id = f.get("file_id", "")
    text = f"Namaste {to_name}, {pname} ki drawings ({sheets}) transmittal {num} ke saath bhej rahe hain. Kripya received confirm kijiye."
    links = {"whatsapp": _wa(to_contact, text)}
    if "@" in to_contact:
        links["email"] = _mail(to_contact, f"{pname}: drawings {num}", text)
    return {"transmittal_id": tr_id, "file": f, "text": text, "links": links,
            "say": f"{num} issue maana. Transmittal PDF bana; drawings ki files aur yeh message aap khud bhejiye (Hoola nahi bhejta)."}


# ---------------------------------------------------------------- RERA architect certificate (draft; issued via transmittal)
def rera_certificate(tenant: str, ref: str, building: str, pct_done, as_on: str = "", stages: dict | None = None,
                     conv: str = "") -> dict:
    """MahaRERA Form 1 style architect's certificate DRAFT from HER site-visit numbers; it becomes a register sheet that only a
    transmittal (her tap) issues. Hoola never makes up a percentage."""
    pct = R.num(pct_done)
    if pct is None or not 0 <= pct <= 100:
        raise ArchError("Kaam kitna % hua (aapki site visit se) chahiye.")
    from .. import officegen
    with session() as s:
        p = _proj(s, tenant, ref)
        st = _st(s, tenant)
        if not p.rera_no:
            raise ArchError("Project ka RERA registration number pehle daaliye.")
        lines = [f"# Architect's Certificate (Form 1 style) - DRAFT", f"Project: {p.name}", f"RERA registration no.: {p.rera_no}",
                 f"Building / wing: {_clean(building, 80) or '[building]'}", f"As on: {_iso(as_on) or today(tenant).isoformat()}",
                 f"Work done (as seen by me on site): {pct:g}%"]
        for k, v in (stages or {}).items():
            lines.append(f"- {_clean(k, 60)}: {_clean(v, 40)}")
        lines += ["", f"{st.architect or '[architect]'}  CoA Reg. No. {st.coa_no or '[ ]'}",
                  "Yeh draft hai: MahaRERA ke current format se milaiye; sign aap khud kijiye."]
        pnum = p.number
    data = officegen.make_docx("Architect certificate", "\n".join(lines))
    bslug = re.sub(r"\W+", "_", building or "bldg")
    f = _store(tenant, f"RERA_Form1_{pnum}_{bslug}.docx", data,
               "application/vnd.openxmlformats-officedocument.wordprocessingml.document", conv, "RERA certificate draft")
    sheet = ("RERA-F1-" + re.sub(r"[^A-Za-z0-9]", "", building or "B")[:8]).upper()
    try:
        drawing(tenant, ref, sheet, title=f"RERA Form 1 certificate {building} {pct:g}%", discipline="A", file_id=f.get("file_id", ""),
                note=f"{pct:g}% as on {as_on or 'aaj'}", revise=bool(_sheet_exists(tenant, ref, sheet)))
    except ArchError:
        pass
    return {"file": f, "sheet": sheet, "say": f"Form 1 certificate draft bana ({pct:g}%). Promoter ko dena = transmittal (sheet {sheet}) + aapka tap."}


def _sheet_exists(tenant: str, ref: str, sheet: str) -> bool:
    with session() as s:
        p = _proj(s, tenant, ref)
        return s.scalars(select(ArchDrawing).where(ArchDrawing.project_id == p.id, ArchDrawing.sheet == sheet)).first() is not None


# ---------------------------------------------------------------- consultants, RFIs, minutes
def consultant(tenant: str, name: str, kind: str = "structural", phone: str = "", email: str = "", firm: str = "", licence: str = "") -> dict:
    if kind not in R.CONSULTANT_KINDS:
        raise ArchError("kind = " + " | ".join(R.CONSULTANT_KINDS))
    if not str(name or "").strip():
        raise ArchError("Naam chahiye.")
    with _LOCK, session() as s:
        c = s.scalars(select(ArchConsultant).where(ArchConsultant.tenant_id == tenant, ArchConsultant.name == name.strip())).first()
        if c is None:
            c = ArchConsultant(tenant_id=tenant, name=_clean(name, 120))
            s.add(c)
        c.kind = kind
        if phone:
            c.phone = _phone(phone)
        for k, v in (("email", email), ("firm", firm), ("licence", licence)):
            if v:
                setattr(c, k, _clean(v, 160))
        s.flush()
        return {"id": c.id, "name": c.name, "kind": c.kind, "phone": c.phone, "email": c.email, "firm": c.firm, "licence": c.licence,
                "say": f"{R.CONSULTANT_KINDS[kind]} consultant {c.name} save."}


def consultants(tenant: str) -> list[dict]:
    with session() as s:
        return [{"id": c.id, "name": c.name, "kind": c.kind, "kind_label": R.CONSULTANT_KINDS.get(c.kind, c.kind), "phone": c.phone,
                 "email": c.email, "firm": c.firm, "licence": c.licence}
                for c in s.scalars(select(ArchConsultant).where(ArchConsultant.tenant_id == tenant)).all()]


def _rfi_view(r: ArchRfi, on: date | None = None) -> dict:
    late = bool(r.status == "open" and r.due and on and _date(r.due) and _date(r.due) < on)
    return {"id": r.id, "number": r.number, "direction": r.direction, "party": r.party, "subject": r.subject, "question": r.question,
            "sheets": _j(r.sheets, []), "due": r.due, "status": r.status, "answer": r.answer, "late": late}


def rfi(tenant: str, ref: str, subject: str, question: str, party: str = "", sheets: list | None = None, due: str = "",
        direction: str = "out") -> dict:
    if not str(question or "").strip():
        raise ArchError("Sawaal kya hai?")
    with _LOCK, session() as s:
        p = _proj(s, tenant, ref)
        cons = None
        if party:
            cons = s.scalars(select(ArchConsultant).where(ArchConsultant.tenant_id == tenant,
                                                          ArchConsultant.name.ilike(f"%{party.strip()}%"))).first()
        r = ArchRfi(tenant_id=tenant, project_id=p.id, number=f"RFI-{_next(s, tenant, 'rfi'):03d}", direction="in" if direction == "in" else "out",
                    party=_clean(cons.name if cons else party, 160), consultant_id=cons.id if cons else "", subject=_clean(subject, 200),
                    question=_clean(question, 2000), sheets=_d([str(x).upper() for x in sheets or []]), due=_iso(due))
        s.add(r)
        s.flush()
        v = _rfi_view(r, today(tenant))
        pname, phone, email = f"{p.number} {p.name}", (cons.phone if cons else ""), (cons.email if cons else "")
    text = (f"{v['number']} ({pname}): {v['subject']}\n{v['question']}" + (f"\nSheets: {', '.join(v['sheets'])}" if v["sheets"] else "")
            + (f"\nJawab {v['due']} tak chahiye." if v["due"] else ""))
    v["text"] = text
    v["links"] = {"whatsapp": _wa(phone, text)} | ({"email": _mail(email, f"{v['number']} {pname}", text)} if email else {})
    v["say"] = f"{v['number']} bana. Message draft taiyaar - aap bhejiye."
    return v


def rfi_answer(tenant: str, rfi_id: str, answer: str = "", close: bool = False) -> dict:
    with _LOCK, session() as s:
        r = s.get(ArchRfi, rfi_id)
        if r is None or r.tenant_id != tenant:
            r = s.scalars(select(ArchRfi).where(ArchRfi.tenant_id == tenant, ArchRfi.number == str(rfi_id).upper())).first()
        if r is None:
            raise ArchError("RFI nahi mila.")
        if answer:
            r.answer = _clean(answer, 2000)
            r.status = "answered"
        if close:
            r.status = "closed"
        return _rfi_view(r, today(tenant)) | {"say": f"{r.number}: {r.status}."}


def rfis(tenant: str, ref: str = "", status: str = "") -> list[dict]:
    with session() as s:
        q = select(ArchRfi).where(ArchRfi.tenant_id == tenant)
        if ref:
            q = q.where(ArchRfi.project_id == _proj(s, tenant, ref).id)
        if status:
            q = q.where(ArchRfi.status == status)
        on = today(tenant)
        return [_rfi_view(r, on) for r in s.scalars(q.order_by(ArchRfi.created_at.desc())).all()[:100]]


def _mt_view(m: ArchMeeting) -> dict:
    return {"id": m.id, "number": m.number, "on": m.on, "place": m.place, "attendees": _j(m.attendees, []), "points": _j(m.points, []),
            "file_id": m.file_id}


def meeting(tenant: str, ref: str, points: list, attendees: list | None = None, on: str = "", place: str = "", conv: str = "") -> dict:
    """Minutes from HER points (nothing added): Word file + a draft she sends."""
    clean = []
    for i, x in enumerate(points or [], 1):
        x = x if isinstance(x, dict) else {"point": x}
        if str(x.get("point") or "").strip():
            clean.append({"n": i, "point": _clean(x["point"], 400), "by": _clean(x.get("by"), 80), "due": _iso(x.get("due")), "done": False})
    if not clean:
        raise ArchError("Meeting ke points chahiye.")
    from .. import officegen
    with _LOCK, session() as s:
        p = _proj(s, tenant, ref)
        m = ArchMeeting(tenant_id=tenant, project_id=p.id, number=f"M-{_next(s, tenant, 'meeting'):03d}", on=_iso(on) or today(tenant).isoformat(),
                        place=_clean(place, 160), attendees=_d([_clean(a, 80) for a in attendees or []]), points=_d(clean))
        s.add(m)
        s.flush()
        pname, mid, mnum, mon = f"{p.number} {p.name}", m.id, m.number, m.on
    body = [f"# Minutes of meeting {mnum}", f"Project: {pname}", f"Date: {mon}" + (f"  Place: {place}" if place else ""),
            "Present: " + (", ".join(attendees or []) or "[ ]"), "", "| # | Point | Action by | By date |", "|---|---|---|---|"]
    body += [f"| {x['n']} | {x['point']} | {x['by']} | {x['due']} |" for x in clean]
    f = _store(tenant, f"MoM_{mnum}.docx", officegen.make_docx(f"MoM {mnum}", "\n".join(body)),
               "application/vnd.openxmlformats-officedocument.wordprocessingml.document", conv, f"Minutes {mnum}")
    with session() as s:
        s.get(ArchMeeting, mid).file_id = f.get("file_id", "")
    text = f"{pname} meeting {mon} ke minutes ({mnum}):\n" + "\n".join(f"{x['n']}. {x['point']}" + (f" - {x['by']}" if x["by"] else "")
                                                               + (f" ({x['due']})" if x["due"] else "") for x in clean)
    return {"id": mid, "number": mnum, "file": f, "text": text, "links": {"whatsapp": _wa("", text)},
            "say": f"{mnum} minutes bane ({len(clean)} points). Bhejna aapke haath mein."}


def point_done(tenant: str, meeting_id: str, n: int) -> dict:
    with _LOCK, session() as s:
        m = s.get(ArchMeeting, meeting_id)
        if m is None or m.tenant_id != tenant:
            raise ArchError("Meeting nahi mili.")
        pts = _j(m.points, [])
        for x in pts:
            if int(x.get("n") or 0) == int(n):
                x["done"] = True
        m.points = _d(pts)
        return _mt_view(m) | {"say": f"Point {n} ho gaya."}


# ---------------------------------------------------------------- fees + invoices (HER tap, fingerprint once a key is active)
def fee_set(tenant: str, ref: str, basis: str = "", percent=None, project_cost=None, lump=None, stages: list | None = None) -> dict:
    with _LOCK, session() as s:
        p = _proj(s, tenant, ref)
        fee = _j(p.fee, {})
        if basis:
            if basis not in ("percent", "lump"):
                raise ArchError("basis = percent | lump")
            fee["basis"] = basis
        for k, v in (("percent", percent), ("project_cost", project_cost), ("lump", lump)):
            if v is not None and v != "":
                n = R.num(v)
                if n is None or n < 0:
                    raise ArchError(f"{k} sahi number nahi.")
                fee[k] = n
        st = _st(s, tenant)
        base = stages or fee.get("stages") or _j(st.fee_stages, []) or [{"key": k, "label": l, "pct": pc} for k, l, pc in R.FEE_STAGES]
        if abs(sum(R.num(x.get("pct")) or 0 for x in base) - 100) > 0.01:
            raise ArchError("Stage % ka jod 100 hona chahiye.")
        old = {x.get("key"): x for x in fee.get("stages") or []}
        fee["stages"] = [{"key": x["key"], "label": x.get("label", x["key"]), "pct": R.num(x["pct"]),
                          "invoice": (old.get(x["key"]) or {}).get("invoice", "")} for x in base]
        fee["starter_stages"] = not (stages or _j(st.fee_stages, []))
        p.fee = _d(fee)
        out = _fee_view(fee)
    out["say"] = (f"Fee: {_rs(out['total'])}" if out["total"] else "Fee basis adhoora (percent + project cost, ya lump sum).") + (
        " Stages CoA ke reference wale starter hain - aapke agreement ke hisaab se badaliye." if fee["starter_stages"] else "")
    return out


def _fee_view(fee: dict) -> dict:
    total = R.stage_amount(fee, 100.0)
    return {"basis": fee.get("basis", ""), "percent": fee.get("percent"), "project_cost": fee.get("project_cost"), "lump": fee.get("lump"),
            "total": total, "stages": [x | {"amount": R.stage_amount(fee, x["pct"])} for x in fee.get("stages") or []],
            "starter_stages": fee.get("starter_stages", True)}


def _inv_view(i: ArchInvoice) -> dict:
    tax = _j(i.tax, {})
    return {"id": i.id, "number": i.number, "stage": i.stage, "label": i.label, "amount": i.amount, "tax": tax, "tds": i.tds,
            "total": i.total, "amount_text": _rs(i.total), "status": i.status, "hash": i.hash, "issued_on": i.issued_on,
            "file_id": i.file_id, "paid_amount": i.paid_amount, "tds_deducted": i.tds_deducted, "paid_on": i.paid_on, "reason": i.reason}


def _inv_show(i: ArchInvoice, client: str) -> dict:
    return {"project": i.project_id, "stage": i.stage, "label": i.label, "amount": i.amount, "tax": _j(i.tax, {}), "client": client}


def _fy_billed(s, tenant: str, project_id: str, on: date) -> float:
    fy = R.fy_of(on)
    tot = 0.0
    for x in s.scalars(select(ArchInvoice).where(ArchInvoice.tenant_id == tenant, ArchInvoice.project_id == project_id,
                                                 ArchInvoice.status.in_(("issued", "paid")))).all():
        if _date(x.issued_on) and R.fy_of(_date(x.issued_on)) == fy:
            tot += x.amount
    return tot


def invoice_request(tenant: str, ref: str, stage: str = "", amount=None, label: str = "", conv: str = "") -> dict:
    with _LOCK, session() as s:
        p = _proj(s, tenant, ref)
        st = _st(s, tenant)
        fee = _j(p.fee, {})
        stg = next((x for x in fee.get("stages") or [] if x["key"] == stage), None)
        if stg:
            if stg.get("invoice"):
                old = s.get(ArchInvoice, stg["invoice"])
                if old is not None and old.status in ("waiting", "issued", "paid"):
                    raise ArchError(f"'{stg['label']}' ka invoice pehle se hai ({old.number or 'card baaki'}).")
            amt = R.num(amount) if amount not in (None, "") else R.stage_amount(fee, stg["pct"])
            lab = label or f"Professional fee: {stg['label']} ({stg['pct']:g}%)"
        else:
            if stage:
                raise ArchError("Yeh stage fee mein nahi: " + ", ".join(x["key"] for x in fee.get("stages") or []))
            amt = R.num(amount)
            lab = label or "Professional fee"
        if not amt or amt <= 0:
            raise ArchError("Raqam pata nahi: fee basis (percent + project cost / lump) ya amount daaliye.")
        tax = R.gst_split(amt, st.gstin, p.client_gstin)
        tds = R.tds_expected(_fy_billed(s, tenant, p.id, today(tenant)), amt)
        for old in s.scalars(select(ArchInvoice).where(ArchInvoice.project_id == p.id, ArchInvoice.stage == (stage or ""),
                                                       ArchInvoice.status == "waiting")).all():
            old.status = "stale"
        i = ArchInvoice(tenant_id=tenant, project_id=p.id, stage=stage or "", label=_clean(lab, 160), amount=round(amt, 2),
                        tax=_d(tax | {"tds_note": tds["note"]}), tds=tds["tds"], total=tax["total"])
        i.hash = _hash(_inv_show(i, p.client))
        s.add(i)
        s.flush()
        if stg:
            stg["invoice"] = i.id
            p.fee = _d(fee)
        v = _inv_view(i)
        v["project"], v["client"] = f"{p.number} {p.name}", p.client
    v["say"] = (f"Invoice card: {v['label']} {_rs(v['amount'])} + GST = {v['amount_text']}" + (f" (client ~{_rs(v['tds'])} TDS kaatega)" if v["tds"] else "")
                + ". Aapki haan ke baad hi number lagega.")
    _publish(tenant, "archassist_invoice", {"invoice": v, "say": v["say"]}, conv)
    return v


def approval_card(tenant: str, inv_id: str) -> dict | None:
    with session() as s:
        i = s.get(ArchInvoice, inv_id)
        if i is None or i.tenant_id != tenant:
            return None
        p = s.get(ArchProject, i.project_id)
        return _inv_view(i) | {"client": p.client if p else "", "project": f"{p.number} {p.name}" if p else ""}


def invoice_challenge(tenant: str, inv_id: str) -> dict:
    """approvekey's message format (kind 'archassist'): her phone signs the invoice card as it is NOW."""
    from .. import approvekey as ak
    c = approval_card(tenant, inv_id)
    if not c or c["status"] != "waiting":
        raise ArchError("Yeh card ab baaki nahi hai.")
    nonce = secrets.token_urlsafe(18)
    exp = datetime.now(timezone.utc) + ak.CHALLENGE_TTL
    msg = f"hoola-approve|v1|{tenant}|archassist|{inv_id}|{c['hash']}|{nonce}|{int(exp.timestamp())}"
    with session() as s:
        ch = ak.ApproveChallenge(tenant_id=tenant, kind="archassist", target=inv_id, digest=c["hash"], message=msg, expires_at=exp)
        s.add(ch)
        s.flush()
        chid = ch.id
    return {"challenge_id": chid, "message": msg, "show": {"title": f"Invoice: {c['label'][:60]}", "amount": c["amount_text"],
                                                           "payee": c["client"]}, "expires_at": exp.isoformat()}


def _verify_fingerprint(tenant: str, inv_id: str, digest: str, proof: dict | None) -> None:
    from .. import approvekey as ak
    if not proof:
        raise ArchError("fingerprint_needed")
    with session() as s:
        ak._settle(s, tenant)
        ok = ak._check(s, tenant, "archassist", inv_id, proof, digest)
    if not ok:
        raise ArchError("fingerprint_invalid")


def _inv_pdf(st: ArchSettings, p: ArchProject, i: ArchInvoice) -> bytes:
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.units import mm
    from reportlab.pdfgen import canvas
    tax = _j(i.tax, {})
    bio = io.BytesIO()
    c = canvas.Canvas(bio, pagesize=A4)
    w, h = A4
    y = h - 20 * mm
    c.setFont("Helvetica-Bold", 14)
    c.drawString(20 * mm, y, (st.firm or "Architect")[:70])
    c.setFont("Helvetica", 9)
    for line in (st.address[:110], f"CoA Reg. No. {st.coa_no or '[ ]'}   GSTIN {st.gstin or '-'}   {st.phone} {st.email}"):
        y -= 5 * mm
        c.drawString(20 * mm, y, line)
    y -= 12 * mm
    c.setFont("Helvetica-Bold", 12)
    c.drawString(20 * mm, y, "TAX INVOICE" if tax.get("gst") else "INVOICE")
    c.setFont("Helvetica", 10)
    try:
        from ..india.money import rupees_words
        words = rupees_words(i.total, "en")
    except Exception:
        words = ""
    rows = [f"Invoice No: {i.number}    Date: {i.issued_on}", f"Bill to: {p.client}" + (f"  GSTIN {p.client_gstin}" if p.client_gstin else ""),
            f"Project: {p.name}  ({p.site[:60]})", "", f"{i.label}   SAC {R.SAC}", f"Taxable value: Rs {i.amount:,.2f}"]
    if tax.get("igst"):
        rows.append(f"IGST 18%: Rs {tax['igst']:,.2f}")
    elif tax.get("gst"):
        rows += [f"CGST 9%: Rs {tax['cgst']:,.2f}", f"SGST 9%: Rs {tax['sgst']:,.2f}"]
    rows += [f"Total: Rs {i.total:,.2f}", f"({words})" if words else "", ""]
    if i.tds:
        rows.append("TDS u/s 393(1) Income-tax Act 2025 (if applicable) to be deducted by the client on the fee.")
    if st.upi_id:
        rows.append(f"UPI: {st.upi_id}")
    for line in st.terms.splitlines()[:6]:
        rows.append(line[:110])
    for line in rows:
        y -= 6 * mm
        c.drawString(20 * mm, y, line[:110])
    y -= 20 * mm
    c.drawString(130 * mm, y, f"For {st.firm or ''}"[:40])
    c.showPage()
    c.save()
    return bio.getvalue()


def decide_invoice(tenant: str, inv_id: str, yes: bool, card_hash: str = "", reason: str = "", proof: dict | None = None,
                   by_tap: bool = False, conv: str = "") -> dict:
    if not by_tap:
        raise ArchError("Invoice ki haan sirf aapke tap se (app ka card).")
    with session() as s:
        i = s.get(ArchInvoice, inv_id)
        if i is None or i.tenant_id != tenant:
            raise ArchError("Card nahi mila.")
        if i.status != "waiting":
            return {"invoice": _inv_view(i), "say": f"Invoice pehle hi '{i.status}' hai."}
        p = s.get(ArchProject, i.project_id)
        if not card_hash or card_hash != i.hash or _hash(_inv_show(i, p.client)) != i.hash:
            raise ArchError("Card badal gaya hai; naya card dekhiye.")
    if yes:
        from .. import approvekey as ak
        if ak.active_on(tenant):
            _verify_fingerprint(tenant, inv_id, card_hash, proof)
    with _LOCK, session() as s:
        i = s.get(ArchInvoice, inv_id)
        i.reason = _clean(reason, 300)
        if not yes:
            i.status = "rejected"
            p = s.get(ArchProject, i.project_id)
            fee = _j(p.fee, {})
            for x in fee.get("stages") or []:
                if x.get("invoice") == inv_id:
                    x["invoice"] = ""
            p.fee = _d(fee)
            return {"invoice": _inv_view(i), "say": "Invoice nahi banaya."}
        st, p = _st(s, tenant), s.get(ArchProject, i.project_id)
        on = today(tenant)
        n = _next(s, tenant, f"inv:{R.fy_of(on)}")
        i.number = f"{(st.inv_prefix or 'AR')[:6]}/{R.fy_of(on)}/{n:03d}"
        i.issued_on, i.status = on.isoformat(), "issued"
        pdf = _inv_pdf(st, p, i)
        view = _inv_view(i)
        upi = st.upi_id
        client, phone, pname = p.client, p.client_phone, p.name
    f = _store(tenant, f"Invoice_{view['number'].replace('/', '-')}.pdf", pdf, "application/pdf", conv, f"Invoice {view['number']}")
    with session() as s:
        s.get(ArchInvoice, inv_id).file_id = f.get("file_id", "")
    text = f"Namaste {client}, {pname} ka invoice {view['number']}: {view['amount_text']}." + (
        f" UPI: {upi}." if upi else "") + " Dhanyavaad."
    links = {"whatsapp": _wa(phone, text)}
    if upi:
        links["upi"] = f"upi://pay?pa={urllib.parse.quote(upi)}&am={view['total']:.2f}&cu=INR&tn={urllib.parse.quote(view['number'])}"
    return {"invoice": view | {"file_id": f.get("file_id", "")}, "file": f, "text": text, "links": links,
            "say": f"Invoice {view['number']} bana ({view['amount_text']}). PDF aur message aap client ko bhejiye; paise aane par 'paid' bataiye."}


def invoice_paid(tenant: str, inv_id: str, amount=None, tds=None, on: str = "", ref: str = "") -> dict:
    with _LOCK, session() as s:
        i = s.get(ArchInvoice, inv_id)
        if i is None or i.tenant_id != tenant:
            i = s.scalars(select(ArchInvoice).where(ArchInvoice.tenant_id == tenant, ArchInvoice.number == str(inv_id))).first()
        if i is None:
            raise ArchError("Invoice nahi mila.")
        if i.status not in ("issued", "paid"):
            raise ArchError(f"Invoice abhi '{i.status}' hai.")
        a = R.num(amount)
        t = R.num(tds) or 0.0
        if a is None:
            raise ArchError("Kitne paise aaye (aapke bank / UPI se)?")
        i.paid_amount, i.tds_deducted = round(i.paid_amount + a, 2), round(i.tds_deducted + t, 2)
        i.paid_on, i.paid_ref = _iso(on) or today(tenant).isoformat(), _clean(ref, 60)
        left = round(i.total - i.paid_amount - i.tds_deducted, 2)
        if left <= 1:
            i.status = "paid"
        v = _inv_view(i)
    v["say"] = f"{v['number']}: {_rs(a)} aaye" + (f" + TDS {_rs(t)} (26AS mein dekhiye)" if t else "") + (
        "; poora." if v["status"] == "paid" else f"; baaki {_rs(left)}.")
    return v


def invoices(tenant: str, ref: str = "", status: str = "") -> list[dict]:
    with session() as s:
        q = select(ArchInvoice).where(ArchInvoice.tenant_id == tenant)
        if ref:
            q = q.where(ArchInvoice.project_id == _proj(s, tenant, ref).id)
        if status:
            q = q.where(ArchInvoice.status == status)
        return [_inv_view(i) for i in s.scalars(q.order_by(ArchInvoice.created_at.desc())).all()[:100]]


# ---------------------------------------------------------------- day note, register, tick
def day(tenant: str) -> dict:
    on = today(tenant)
    out: dict[str, Any] = {"waiting": [], "rfi_late": [], "expiring": [], "unpaid": [], "coa": ""}
    with session() as s:
        st = _st(s, tenant)
        if st.coa_valid_to and _date(st.coa_valid_to):
            left = (_date(st.coa_valid_to) - on).days
            if left <= 60:
                out["coa"] = f"CoA registration {st.coa_valid_to} ({'nikal gayi' if left < 0 else f'{left} din'}): renew kijiye."
        for t in s.scalars(select(ArchTransmittal).where(ArchTransmittal.tenant_id == tenant, ArchTransmittal.status == "waiting")).all():
            out["waiting"].append(f"{t.number} transmittal ({t.to_name})")
        for x in s.scalars(select(ArchSubmission).where(ArchSubmission.tenant_id == tenant, ArchSubmission.status == "waiting")).all():
            out["waiting"].append(f"{x.number} {R.STAGE_LABEL.get(x.stage, x.stage)}")
        for i in s.scalars(select(ArchInvoice).where(ArchInvoice.tenant_id == tenant, ArchInvoice.status.in_(("waiting", "issued")))).all():
            if i.status == "waiting":
                out["waiting"].append(f"Invoice {_rs(i.total)}")
            elif _date(i.issued_on) and (on - _date(i.issued_on)).days >= 30:
                out["unpaid"].append(f"{i.number} {_rs(i.total - i.paid_amount)} ({(on - _date(i.issued_on)).days} din)")
        for r in s.scalars(select(ArchRfi).where(ArchRfi.tenant_id == tenant, ArchRfi.status == "open")).all():
            if r.due and _date(r.due) and _date(r.due) < on:
                out["rfi_late"].append(f"{r.number} {r.party}: {r.subject}")
        for p in s.scalars(select(ArchProject).where(ArchProject.tenant_id == tenant, ArchProject.status == "active")).all():
            for k, x in _j(p.approvals, {}).items():
                vt = _date(x.get("valid_to"))
                if x.get("status") == "granted" and vt and (vt - on).days <= 60:
                    out["expiring"].append(f"{p.number} {R.STAGE_LABEL.get(k, k)} valid {x['valid_to']}")
    bits = []
    if out["waiting"]:
        bits.append(f"Aapki haan baaki: {', '.join(out['waiting'][:4])}")
    if out["rfi_late"]:
        bits.append(f"RFI late: {', '.join(out['rfi_late'][:3])}")
    if out["expiring"]:
        bits.append(f"Approval ki tareekh paas: {', '.join(out['expiring'][:3])}")
    if out["unpaid"]:
        bits.append(f"Payment baaki: {', '.join(out['unpaid'][:3])}")
    if out["coa"]:
        bits.append(out["coa"])
    out["say"] = ". ".join(bits) + "." if bits else "Architecture office mein aaj sab theek."
    return out


def register_file(tenant: str, ref: str, conv: str = "") -> dict:
    import openpyxl
    with session() as s:
        pnum = _proj(s, tenant, ref).number
    pv = project_view(tenant, ref)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Drawings"
    ws.append(["Sheet", "Title", "Discipline", "Scale", "Rev", "Status", "Revisions", "Issued"])
    for d in pv["drawings"]:
        ws.append([d["sheet"], d["title"], d["discipline_label"], d["scale"], d["rev"], d["status"],
                   "; ".join(f"{r['rev']} {r['on']} {r['note']}" for r in d["revisions"]),
                   "; ".join(f"{x['transmittal']} {x['rev']} {x['to']} {x['on']}" for x in d["issued"])])
    ws = wb.create_sheet("Transmittals")
    ws.append(["No", "To", "Purpose", "Sheets", "Status"])
    for t in pv["transmittals"]:
        ws.append([t["number"], t["to_name"], t["purpose_label"], ", ".join(f"{i['sheet']} {i['rev']}" for i in t["items"]), t["status"]])
    ws = wb.create_sheet("Approvals")
    ws.append(["Item", "Status", "Ref", "On", "Valid to", "Why"])
    for a in pv["approvals"]["items"]:
        ws.append([a["label"], a["status"], a["ref"], a["on"], a["valid_to"], a["why"]])
    ws = wb.create_sheet("RFI")
    ws.append(["No", "Party", "Subject", "Due", "Status", "Answer"])
    for r in pv["rfis"]:
        ws.append([r["number"], r["party"], r["subject"], r["due"], r["status"], r["answer"]])
    ws = wb.create_sheet("Fees")
    ws.append(["Invoice", "Stage", "Amount", "Total", "Status", "Paid", "TDS"])
    for i in pv["invoices"]:
        ws.append([i["number"], i["label"], i["amount"], i["total"], i["status"], i["paid_amount"], i["tds_deducted"]])
    bio = io.BytesIO()
    wb.save(bio)
    f = _store(tenant, f"Register_{pnum}.xlsx", bio.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
               conv, f"Register {pnum}")
    return {"file": f, "say": f"{pnum} ka drawing / transmittal / approval / RFI / fee register Excel."}


def _tenants() -> list[str]:
    with session() as s:
        return sorted(set(s.scalars(select(ArchProject.tenant_id)).all()) | set(s.scalars(select(ArchSettings.tenant_id)).all()))


def tick(when: datetime | None = None) -> dict:
    """Worker: once per item - RFI past due, approval validity in 60 / 30 / 7 days, invoice unpaid 30 days; morning note at her time."""
    n = when or now()
    if n.tzinfo is None:
        n = n.replace(tzinfo=timezone.utc)
    out = {"alerts": 0, "morning": 0}
    for tenant in _tenants():
        try:
            loc = n.astimezone(_tz(tenant))
            on = loc.date()
            alerts = []
            with _LOCK, session() as s:
                for r in s.scalars(select(ArchRfi).where(ArchRfi.tenant_id == tenant, ArchRfi.status == "open")).all():
                    if r.due and _date(r.due) and _date(r.due) < on and _once(s, tenant, f"rfi:{r.id}:{r.due}"):
                        alerts.append(f"{r.number} ({r.party}) ka jawab {r.due} tak aana tha: yaad dilaiye.")
                for p in s.scalars(select(ArchProject).where(ArchProject.tenant_id == tenant, ArchProject.status == "active")).all():
                    for k, x in _j(p.approvals, {}).items():
                        vt = _date(x.get("valid_to"))
                        if x.get("status") != "granted" or not vt:
                            continue
                        left = (vt - on).days
                        win = next((w for w in (0, 7, 30, 60) if left <= w), None)
                        if win is not None and _once(s, tenant, f"apv:{p.id}:{k}:{x['valid_to']}:{win}"):
                            alerts.append(f"{p.number} {R.STAGE_LABEL.get(k, k)} {x['valid_to']} ko khatam" + (
                                " ho gaya" if left < 0 else f" ({left} din)") + ": renewal / revalidation dekhiye.")
                for i in s.scalars(select(ArchInvoice).where(ArchInvoice.tenant_id == tenant, ArchInvoice.status == "issued")).all():
                    d0 = _date(i.issued_on)
                    if d0 and (on - d0).days >= 30 and _once(s, tenant, f"inv:{i.id}:30"):
                        alerts.append(f"Invoice {i.number} ({_rs(i.total - i.paid_amount)}) 30 din se baaki: client ko yaad dilaiye.")
                st = s.get(ArchSettings, tenant)
                morning = False
                if st is not None and st.last_morning != on.isoformat():
                    try:
                        hh, mm = (int(x) for x in (st.morning_time or "09:00").split(":"))
                    except ValueError:
                        hh, mm = 9, 0
                    if (loc.hour, loc.minute) >= (hh, mm):
                        st.last_morning = on.isoformat()
                        morning = True
            for a in alerts:
                _publish(tenant, "archassist_alert", {"say": a})
                out["alerts"] += 1
            if morning:
                d = day(tenant)
                if d["say"] != "Architecture office mein aaj sab theek.":
                    _publish(tenant, "archassist_day", {"say": d["say"], "waiting": len(d["waiting"]), "rfi_late": len(d["rfi_late"]),
                                                        "expiring": len(d["expiring"]), "unpaid": len(d["unpaid"])})
                    out["morning"] += 1
        except Exception:
            log.exception("archassist tick failed for %s", tenant)
    return out
