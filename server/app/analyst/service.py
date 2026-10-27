"""v0.95 Financial Analyst: Hoola reads a business's books and bank and tells the owner where the money stands.

  1. Gets the job: "September ka P&L dikhao", "cash kitne mahine chalega", "20 lakh ka loan le sakte hain?", a Tally
     trial balance / P&L / balance sheet export, a bank statement (Excel / CSV), the outstanding list.
  2. Does the analysis, step by step: books -> P&L + balance sheet (Tally groups understood) -> ratios (margins, current
     / quick, DSCR, TOL/TNW, debtor / creditor / stock days, interest cover) against bank benchmarks; bank -> money in /
     out by category, burn, runway, recurring payments, 13-week cash forecast; receivables aging; MSME payables at risk.
  3. Checked against the rules and earlier work: last period's numbers (what moved and why), Sec 37(2)(g) (old 43B(h))
     15 / 45-day MSME payments + MSMED interest at 3x bank rate, advance tax instalments, working-capital limit (Nayak),
     DSCR after a new loan's EMI, payroll and GST figures Hoola already holds.
  4. Ready as a card, a file or a draft: MIS card, MIS Excel + one-page PDF, collection reminder drafts, a note for the
     bank / CA.
  5. Waits for her Haan: sharing the report outside (bank, CA, investor) and sending reminders. Hoola never moves money,
     never invests, never applies for a loan; its numbers are analysis, not financial advice.
  6. Delivers and learns: every period is kept for trends; ledger -> head and bank narration -> category corrections she
     makes are remembered for next time.
"""
from __future__ import annotations

import csv
import hashlib
import io
import json
import logging
import re
import threading
from datetime import date, datetime, timedelta
from typing import Any

from sqlalchemy import Boolean, DateTime, Float, Integer, String, Text, select
from sqlalchemy.orm import Mapped, mapped_column

from ..db import Base, new_id, now, session
from . import rules as R

log = logging.getLogger("santosh.analyst")
_LOCK = threading.Lock()
CATS = ("revenue", "other_income_direct", "other_income", "purchases", "direct_exp", "indirect_exp", "opening_stock", "closing_stock",
        "depreciation", "interest", "debtors", "cash", "bank", "other_ca", "fixed_assets", "investments", "creditors", "duties",
        "provisions", "other_cl", "bank_od", "secured_loans", "unsecured_loans", "capital", "reserves", "suspense")
CAT_LABEL = {"revenue": "Bikri (sales)", "purchases": "Kharid", "direct_exp": "Seedhe kharche", "indirect_exp": "Office / chalane ke kharche",
             "depreciation": "Depreciation", "interest": "Byaaj", "other_income": "Doosri kamai", "debtors": "Lene hain (debtors)",
             "creditors": "Dene hain (creditors)", "closing_stock": "Stock", "cash": "Cash", "bank": "Bank", "bank_od": "OD / CC",
             "secured_loans": "Loan (secured)", "unsecured_loans": "Loan (unsecured)", "capital": "Capital", "reserves": "Reserves",
             "fixed_assets": "Fixed assets", "duties": "Tax dene hain", "other_ca": "Advance / deposit", "provisions": "Provisions"}


class FinError(Exception):
    pass


class FinSettings(Base):
    __tablename__ = "fin_settings"
    tenant_id: Mapped[str] = mapped_column(String(64), primary_key=True)
    enabled: Mapped[bool] = mapped_column(Boolean, default=True)
    business: Mapped[str] = mapped_column(String(160), default="")
    entity: Mapped[str] = mapped_column(String(12), default="individual")     # individual | firm | llp | company
    tax_paid: Mapped[str] = mapped_column(Text, default="{}")                 # {"2026-27": 50000} advance tax paid so far
    budget: Mapped[str] = mapped_column(Text, default="{}")                   # monthly {bank category: amount}
    morning_time: Mapped[str] = mapped_column(String(5), default="10:15")
    last_morning: Mapped[str] = mapped_column(String(10), default="")


class FinSnapshot(Base):
    __tablename__ = "fin_snapshots"
    id: Mapped[str] = mapped_column(String(40), primary_key=True, default=lambda: new_id("fs"))
    tenant_id: Mapped[str] = mapped_column(String(64), index=True)
    label: Mapped[str] = mapped_column(String(40), default="")
    start: Mapped[str] = mapped_column(String(10), default="")
    end: Mapped[str] = mapped_column(String(10), default="")
    balances: Mapped[str] = mapped_column(Text, default="{}")                 # category -> amount
    ledgers: Mapped[str] = mapped_column(Text, default="[]")                  # [{name, group, amount, side, category}]
    source: Mapped[str] = mapped_column(String(20), default="manual")
    status: Mapped[str] = mapped_column(String(10), default="ready")          # ready | approved (ok to share)
    created_at: Mapped[datetime] = mapped_column(DateTime(timezone=True), default=now)


class FinTxn(Base):
    __tablename__ = "fin_txns"
    id: Mapped[str] = mapped_column(String(40), primary_key=True, default=lambda: new_id("ft"))
    tenant_id: Mapped[str] = mapped_column(String(64), index=True)
    account: Mapped[str] = mapped_column(String(60), default="")
    day: Mapped[str] = mapped_column(String(10), index=True)
    narration: Mapped[str] = mapped_column(String(300), default="")
    amount: Mapped[float] = mapped_column(Float, default=0.0)                 # + money in, - money out
    balance: Mapped[float | None] = mapped_column(Float, nullable=True)
    category: Mapped[str] = mapped_column(String(30), default="")
    fp: Mapped[str] = mapped_column(String(40), index=True)                   # duplicate guard


class FinRule(Base):
    """Her corrections: kind=bank (narration words -> category) | ledger (ledger name -> statement head)."""
    __tablename__ = "fin_rules"
    id: Mapped[str] = mapped_column(String(40), primary_key=True, default=lambda: new_id("fr"))
    tenant_id: Mapped[str] = mapped_column(String(64), index=True)
    kind: Mapped[str] = mapped_column(String(8), default="bank")
    pattern: Mapped[str] = mapped_column(String(120), default="")
    category: Mapped[str] = mapped_column(String(30), default="")


class FinItem(Base):
    """Bills to collect (receivable) or to pay (payable)."""
    __tablename__ = "fin_items"
    id: Mapped[str] = mapped_column(String(40), primary_key=True, default=lambda: new_id("fi"))
    tenant_id: Mapped[str] = mapped_column(String(64), index=True)
    kind: Mapped[str] = mapped_column(String(10), default="receivable")
    party: Mapped[str] = mapped_column(String(160), default="")
    phone: Mapped[str] = mapped_column(String(20), default="")
    bill: Mapped[str] = mapped_column(String(40), default="")
    bill_date: Mapped[str] = mapped_column(String(10), default="")
    amount: Mapped[float] = mapped_column(Float, default=0.0)
    paid: Mapped[float] = mapped_column(Float, default=0.0)
    credit_days: Mapped[int] = mapped_column(Integer, default=30)
    msme: Mapped[bool] = mapped_column(Boolean, default=False)                # payable to a micro / small enterprise
    agreed_days: Mapped[int] = mapped_column(Integer, default=0)              # written agreement (max 45)


class FinLoan(Base):
    __tablename__ = "fin_loans"
    id: Mapped[str] = mapped_column(String(40), primary_key=True, default=lambda: new_id("fl"))
    tenant_id: Mapped[str] = mapped_column(String(64), index=True)
    name: Mapped[str] = mapped_column(String(80), default="")
    outstanding: Mapped[float] = mapped_column(Float, default=0.0)
    rate: Mapped[float] = mapped_column(Float, default=0.0)
    emi: Mapped[float] = mapped_column(Float, default=0.0)
    months_left: Mapped[int] = mapped_column(Integer, default=0)


# ---------------------------------------------------------------- helpers
def _j(s: str, default):
    try:
        v = json.loads(s or "")
        return v if isinstance(v, type(default)) else default
    except (TypeError, ValueError):
        return default


def _rs(v) -> str:
    v = round(float(v or 0))
    s = str(abs(v))
    if len(s) > 3:
        s = re.sub(r"(\d)(?=(\d\d)+$)", r"\1,", s[:-3]) + "," + s[-3:]
    return ("-" if v < 0 else "") + "₹" + s


def _lakh(v) -> str:
    v = float(v or 0)
    if abs(v) >= 1e7:
        return f"₹{v / 1e7:.2f} cr"
    if abs(v) >= 1e5:
        return f"₹{v / 1e5:.2f} lakh"
    return _rs(v)


def _date(v) -> date | None:
    if v is None or v == "":
        return None
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    s = str(v).strip()
    for pat, order in ((r"(\d{4})-(\d{1,2})-(\d{1,2}).*", "ymd"), (r"(\d{1,2})[-/.](\d{1,2})[-/.](\d{2,4})", "dmy")):
        m = re.fullmatch(pat, s)
        if m:
            a, b, c = (int(x) for x in m.groups())
            try:
                if order == "ymd":
                    return date(a, b, c)
                return date(c + 2000 if c < 100 else c, b, a)
            except ValueError:
                return None
    m = re.fullmatch(r"(\d{1,2})[-/ ]([A-Za-z]{3})[A-Za-z]*[-/ ,]*(\d{2,4})", s)
    if m:
        mon = ("jan", "feb", "mar", "apr", "may", "jun", "jul", "aug", "sep", "oct", "nov", "dec")
        if m.group(2).lower() in mon:
            y = int(m.group(3))
            try:
                return date(y + 2000 if y < 100 else y, mon.index(m.group(2).lower()) + 1, int(m.group(1)))
            except ValueError:
                return None
    return None


def _tz(tenant: str):
    from ..life import tz_of
    return tz_of(tenant)


def today(tenant: str) -> date:
    return datetime.now(_tz(tenant)).date()


def _publish(tenant: str, kind: str, payload: dict, conv: str = "") -> None:
    from ..events import bus
    try:
        if not conv:
            from ..net.service import home_conv
            conv = home_conv(tenant)
        bus.publish(conv, kind, payload)
    except Exception:
        log.exception("analyst publish failed")


def _st(s, tenant: str) -> FinSettings:
    r = s.get(FinSettings, tenant)
    if not r:
        r = FinSettings(tenant_id=tenant)
        s.add(r)
        s.flush()
    return r


def settings(tenant: str) -> dict[str, Any]:
    with session() as s:
        r = _st(s, tenant)
        return {"enabled": r.enabled, "business": r.business, "entity": r.entity, "tax_paid": _j(r.tax_paid, {}),
                "budget": _j(r.budget, {}), "morning_time": r.morning_time}


def save_settings(tenant: str, **kw: Any) -> dict[str, Any]:
    with session() as s:
        r = _st(s, tenant)
        if kw.get("business") is not None:
            r.business = str(kw["business"])[:160]
        if kw.get("entity"):
            e = str(kw["entity"]).lower()
            e = {"proprietor": "individual", "proprietorship": "individual", "pvt ltd": "company", "partnership": "firm"}.get(e, e)
            if e not in ("individual", "firm", "llp", "company"):
                raise FinError("entity: individual | firm | llp | company")
            r.entity = e
        if kw.get("budget"):
            b = _j(r.budget, {})
            b |= {str(k): float(v) for k, v in dict(kw["budget"]).items()}
            r.budget = json.dumps(b)
        if kw.get("tax_paid") is not None:
            tp = _j(r.tax_paid, {})
            fy = str(kw.get("fy") or "")
            if not fy:
                fs = R.fy_start_of(today(tenant))
                fy = f"{fs}-{str(fs + 1)[2:]}"
            tp[fy] = float(kw["tax_paid"])
            r.tax_paid = json.dumps(tp)
        if kw.get("morning_time"):
            r.morning_time = str(kw["morning_time"]).zfill(5)[:5]
        if "enabled" in kw:
            r.enabled = bool(kw["enabled"])
    return settings(tenant)


# ---------------------------------------------------------------- tables from Excel / CSV
def _rows(name: str, data: bytes) -> list[list]:
    if name.lower().endswith((".xlsx", ".xlsm")):
        import openpyxl
        try:
            wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        except Exception as e:
            raise FinError(f"Excel nahi khula: {e}")
        return [list(r) for r in wb.worksheets[0].iter_rows(values_only=True)][:50000]
    text = data.decode("utf-8-sig", errors="replace")
    try:
        dialect = csv.Sniffer().sniff(text[:4000], delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    return list(csv.reader(io.StringIO(text), dialect))


def _norm(h) -> str:
    return re.sub(r"\s+", " ", str(h or "").strip().lower())


def _load_file(tenant: str, file_id: str) -> tuple[str, bytes]:
    from ..mailbox import service as ms
    try:
        name, _m, data = ms.load_files(tenant, [{"type": "file", "file_id": file_id}])[0]
    except Exception as e:
        raise FinError(f"File {file_id} nahi mili: {e}")
    return name, data


# ---------------------------------------------------------------- books (Tally trial balance / P&L / BS)
_DR = ("debit", "dr", "debit amount", "closing debit", "closing balance debit")
_CR = ("credit", "cr", "credit amount", "closing credit", "closing balance credit")
_BAL = ("closing balance", "balance", "amount", "closing")


def _ledger_rules(s, tenant: str) -> dict[str, str]:
    return {r.pattern: r.category for r in s.scalars(select(FinRule).where(FinRule.tenant_id == tenant, FinRule.kind == "ledger")).all()}


def parse_books(name: str, data: bytes, learned: dict[str, str] | None = None) -> list[dict]:
    rows = _rows(name, data)
    hi, cols = None, {}
    for i, r in enumerate(rows[:25]):
        h = [_norm(x) for x in r]
        if any(x in ("particulars", "ledger", "ledger name", "account", "name", "account name", "group") for x in h) and \
                any(x in _DR + _CR + _BAL for x in h):
            hi = i
            for j, x in enumerate(h):
                if x in ("particulars", "ledger", "ledger name", "account", "name", "account name") and "name" not in cols:
                    cols["name"] = j
                elif x in ("group", "under", "parent") and "group" not in cols:
                    cols["group"] = j
                elif x in _DR and "dr" not in cols:
                    cols["dr"] = j
                elif x in _CR and "cr" not in cols:
                    cols["cr"] = j
                elif x in _BAL and "bal" not in cols:
                    cols["bal"] = j
            # Tally puts "Debit / Credit" one row under "Closing Balance"
            if "dr" not in cols and i + 1 < len(rows):
                sub = [_norm(x) for x in rows[i + 1]]
                for j, x in enumerate(sub):
                    if x in _DR and "dr" not in cols:
                        cols["dr"] = j
                    elif x in _CR and "cr" not in cols:
                        cols["cr"] = j
                if "dr" in cols:
                    hi = i + 1
            break
    if hi is None or "name" not in cols:
        raise FinError("Trial balance / P&L ki table nahi mili ('Particulars' aur 'Debit / Credit' ya 'Closing Balance' column chahiye).")
    out = []
    learned = learned or {}
    context = ""
    # detailed export = group names appear as headers without amounts; then every amount row is a ledger
    detailed = False
    for r in rows[hi + 1:]:
        nm0 = str(r[cols["name"]] if cols["name"] < len(r) and r[cols["name"]] is not None else "").strip()
        vals = [r[cols[k]] for k in ("dr", "cr", "bal") if k in cols and cols[k] < len(r)]
        if _norm(nm0) in R.GROUPS and not any(R.clean_amount(v) for v in vals):
            detailed = True
            break
    for r in rows[hi + 1:]:
        nm = str(r[cols["name"]] if cols["name"] < len(r) and r[cols["name"]] is not None else "").strip()
        if not nm or _norm(nm) in ("grand total", "total", "difference in opening balances"):
            continue
        dr = R.clean_amount(r[cols["dr"]]) if "dr" in cols and cols["dr"] < len(r) else 0.0
        cr = R.clean_amount(r[cols["cr"]]) if "cr" in cols and cols["cr"] < len(r) else 0.0
        if "bal" in cols and cols["bal"] < len(r) and not dr and not cr:
            raw = str(r[cols["bal"]] or "")
            amt = R.clean_amount(raw)
            if raw.strip().lower().endswith("cr"):
                cr = abs(amt)
            else:
                dr = abs(amt)
        grp = str(r[cols["group"]]).strip() if "group" in cols and cols["group"] < len(r) and r[cols["group"]] else ""
        known_group = _norm(nm) in R.GROUPS
        if known_group and not dr and not cr:
            context = nm
            continue
        if not dr and not cr:
            continue
        is_group = known_group and not detailed
        if learned.get(_norm(nm)):
            cat = learned[_norm(nm)]
        elif known_group:
            cat = R.GROUPS[_norm(nm)]
        elif grp:
            cat = R.classify(nm, grp)
        else:
            words = R.classify(nm)
            ctx_cat = R.GROUPS.get(_norm(context), "") if context else ""
            # broad parents: the ledger's own words may narrow it, but only to a head inside that parent
            fits = {"other_ca": {"debtors", "cash", "bank", "closing_stock", "other_ca"},
                    "other_cl": {"creditors", "duties", "provisions", "bank_od", "other_cl"},
                    "secured_loans": {"bank_od", "secured_loans", "unsecured_loans"}}
            if ctx_cat in fits:
                cat = words if words in fits[ctx_cat] else ctx_cat
            elif ctx_cat:
                cat = words if (ctx_cat in ("indirect_exp", "direct_exp") and words in ("depreciation", "interest")) else ctx_cat
            else:
                cat = words
        out.append({"name": nm, "group": grp or (context if not is_group else ""), "is_group": is_group, "dr": dr, "cr": cr, "category": cat})
    if not out:
        raise FinError("File mein koi rakam nahi mili.")
    return out


def _pick(lines: list[dict]) -> list[dict]:
    """Ledger lines when they balance (detailed TB); else group lines, finest level (parents minus their sub-groups)."""
    led = [x for x in lines if not x["is_group"]]
    grp = [x for x in lines if x["is_group"]]
    if led:
        d, c = sum(x["dr"] for x in led), sum(x["cr"] for x in led)
        if d and abs(d - c) <= 0.01 * max(d, c) or not grp:
            return led
    sub = {"current assets": ("sundry debtors", "cash-in-hand", "cash in hand", "bank accounts", "stock-in-hand", "stock in hand",
                              "deposits (asset)", "loans & advances (asset)"),
           "current liabilities": ("sundry creditors", "duties & taxes", "duties and taxes", "provisions"),
           "loans (liability)": ("bank od a/c", "bank od accounts", "secured loans", "unsecured loans")}
    names = {_norm(x["name"]): x for x in grp}
    out = []
    for x in grp:
        n = _norm(x["name"])
        if n in sub:
            kids = [names[k] for k in sub[n] if k in names]
            if kids:
                rest = dict(x)
                rest["dr"] = max(0.0, x["dr"] - sum(k["dr"] for k in kids))
                rest["cr"] = max(0.0, x["cr"] - sum(k["cr"] for k in kids))
                if rest["dr"] or rest["cr"]:
                    out.append(rest)
                continue
        out.append(x)
    return out


def _balances(lines: list[dict]) -> dict[str, float]:
    b: dict[str, float] = {}
    for x in lines:
        if x["category"] in ("unknown", "suspense"):
            continue
        amt = abs(x["dr"] - x["cr"])
        b[x["category"]] = b.get(x["category"], 0.0) + amt
    return {k: round(v, 2) for k, v in b.items()}


def import_books(tenant: str, label: str, start: str, end: str, file_id: str = "", data: bytes | None = None,
                 name: str = "tb.xlsx") -> dict[str, Any]:
    if file_id:
        name, data = _load_file(tenant, file_id)
    if not data:
        raise FinError("Tally ka trial balance / P&L file chahiye.")
    sd, ed = _date(start), _date(end)
    if not sd or not ed or ed < sd:
        raise FinError("Kis samay ka hai? start aur end date chahiye (jaise 2026-04-01 se 2026-09-30).")
    with _LOCK, session() as s:
        lines = parse_books(name, data, _ledger_rules(s, tenant))
        use = _pick(lines)
        unknown = [x["name"] for x in use if x["category"] == "unknown"]
        snap = FinSnapshot(tenant_id=tenant, label=(label or f"{sd.isoformat()}..{ed.isoformat()}")[:40], start=sd.isoformat(),
                           end=ed.isoformat(), balances=json.dumps(_balances(use)), ledgers=json.dumps(use, ensure_ascii=False), source="tally")
        s.add(snap)
        s.flush()
        sid = snap.id
    out = analyse(tenant, sid, publish=True)
    if unknown:
        out["unknown"] = unknown[:30]
        out["say"] += f" {len(unknown)} ledger samajh nahi aaye ({', '.join(unknown[:5])}): bataiye kis head mein jaayein."
    return out


def set_figures(tenant: str, label: str, start: str, end: str, figures: dict) -> dict[str, Any]:
    """Figures read from a P&L / balance sheet photo or PDF (category -> amount)."""
    sd, ed = _date(start), _date(end)
    if not sd or not ed:
        raise FinError("start / end date chahiye.")
    b = {k: abs(R.clean_amount(v)) for k, v in (figures or {}).items() if k in CATS}
    if not b:
        raise FinError("Figures chahiye, jaise {revenue, purchases, indirect_exp, debtors, creditors, bank, capital...}")
    with session() as s:
        snap = FinSnapshot(tenant_id=tenant, label=(label or sd.isoformat())[:40], start=sd.isoformat(), end=ed.isoformat(),
                           balances=json.dumps(b), source="manual")
        s.add(snap)
        s.flush()
        sid = snap.id
    return analyse(tenant, sid, publish=True)


def map_ledger(tenant: str, ledger: str, category: str, snapshot: str = "") -> dict[str, Any]:
    if category not in CATS:
        raise FinError("Head: " + ", ".join(CATS))
    with session() as s:
        s.add(FinRule(tenant_id=tenant, kind="ledger", pattern=_norm(ledger)[:120], category=category))
        if snapshot:
            sn = s.get(FinSnapshot, snapshot)
            if sn and sn.tenant_id == tenant:
                led = _j(sn.ledgers, [])
                for x in led:
                    if _norm(x["name"]) == _norm(ledger):
                        x["category"] = category
                sn.ledgers = json.dumps(led, ensure_ascii=False)
                sn.balances = json.dumps(_balances(led))
    return {"say": f"'{ledger}' ab hamesha {CAT_LABEL.get(category, category)} mein jayega."}


def _snap(s, tenant: str, ref: str = "") -> FinSnapshot:
    if ref:
        r = s.get(FinSnapshot, ref)
        if r and r.tenant_id == tenant:
            return r
        r = s.scalars(select(FinSnapshot).where(FinSnapshot.tenant_id == tenant, FinSnapshot.label == ref)
                      .order_by(FinSnapshot.created_at.desc())).first()
        if r:
            return r
        raise FinError(f"'{ref}' ka hisaab nahi mila.")
    r = s.scalars(select(FinSnapshot).where(FinSnapshot.tenant_id == tenant).order_by(FinSnapshot.end.desc(),
                                                                                       FinSnapshot.created_at.desc())).first()
    if not r:
        raise FinError("Abhi koi hisaab nahi: Tally se trial balance (Excel) bhejiye ya P&L ke aankde bataiye.")
    return r


def _debt_service(s, tenant: str, days: int) -> float:
    loans = s.scalars(select(FinLoan).where(FinLoan.tenant_id == tenant)).all()
    yearly_principal = 0.0
    for ln in loans:
        interest_month = ln.outstanding * ln.rate / 1200
        yearly_principal += max(0.0, ln.emi - interest_month) * 12
    return yearly_principal * days / 365


def analyse(tenant: str, snapshot: str = "", conv: str = "", publish: bool = False) -> dict[str, Any]:
    with session() as s:
        sn = _snap(s, tenant, snapshot)
        b = _j(sn.balances, {})
        sd, ed = _date(sn.start), _date(sn.end)
        days = max(1, (ed - sd).days + 1)
        st = R.statements(b)
        ratios = R.ratios(st, days, _debt_service(s, tenant, days))
        judged = R.judge(ratios)
        prev = s.scalars(select(FinSnapshot).where(FinSnapshot.tenant_id == tenant, FinSnapshot.end < sn.start)
                         .order_by(FinSnapshot.end.desc())).first()
        trend = {}
        if prev:
            pd_ = max(1, (_date(prev.end) - _date(prev.start)).days + 1)
            pst = R.statements(_j(prev.balances, {}))
            scale = days / pd_
            for k in ("revenue", "gross_profit", "opex", "profit_before_tax"):
                trend[k] = R.pct_change(st["pl"][k], pst["pl"][k] * scale)
        sid, label = sn.id, sn.label
        payroll_cost = _payroll_cost(s, tenant, sd, ed)
    insights = _insights(st, ratios, trend, payroll_cost, days)
    pl = st["pl"]
    say = (f"{label}: bikri {_lakh(pl['revenue'])}, gross profit {_lakh(pl['gross_profit'])} "
           f"({ratios['gross_margin'] if ratios['gross_margin'] is not None else '-'}%), kharche {_lakh(pl['opex'])}, "
           f"tax se pehle munafa {_lakh(pl['profit_before_tax'])}. " + " ".join(i["text"] for i in insights[:3]))
    v = {"snapshot": {"id": sid, "label": label, "start": sd.isoformat(), "end": ed.isoformat(), "days": days},
         "pl": pl, "bs": st["bs"], "ratios": judged, "trend": trend, "insights": insights, "payroll_cost": payroll_cost, "say": say}
    if conv or publish:
        _publish(tenant, "fin_report", {"snapshot": sid, "label": label}, conv)
    return v


def _payroll_cost(s, tenant: str, sd: date, ed: date) -> float | None:
    try:
        from ..payroll.service import PayRun
    except Exception:
        return None
    tot = 0.0
    found = False
    for r in s.scalars(select(PayRun).where(PayRun.tenant_id == tenant, PayRun.kind == "monthly", PayRun.status.in_(("approved", "paid")))).all():
        try:
            d = date(int(r.period[:4]), int(r.period[5:7]), 1)
        except ValueError:
            continue
        if sd <= d <= ed:
            found = True
            tot += float(_j(r.totals, {}).get("cost") or 0)
    return round(tot) if found else None


def _insights(st: dict, r: dict, trend: dict, payroll_cost: float | None, days: int) -> list[dict]:
    out = []
    pl, bs = st["pl"], st["bs"]

    def add(level, text):
        out.append({"level": level, "text": text})
    if pl["profit_before_tax"] < 0:
        add("bad", f"Nuksaan: {_lakh(-pl['profit_before_tax'])}. Kharche bikri se zyada.")
    if r.get("current_ratio") is not None and r["current_ratio"] < 1:
        add("bad", f"Current ratio {r['current_ratio']}: chhote samay ke dene, haath ke paise + lene se zyada (cash ki tangi).")
    elif r.get("current_ratio") is not None and r["current_ratio"] < 1.33:
        add("warn", f"Current ratio {r['current_ratio']} (bank 1.33 dekhte hain).")
    if r.get("debtor_days") and r["debtor_days"] > 60:
        add("warn", f"Paisa {r['debtor_days']:.0f} din mein aata hai: vasooli tez karni padegi.")
    if r.get("dscr") is not None and r["dscr"] < 1.25:
        add("bad" if r["dscr"] < 1 else "warn", f"DSCR {r['dscr']}: kamai se loan ki kist mushkil se nikalti hai (bank 1.25+ chahte hain).")
    if r.get("tol_tnw") is not None and r["tol_tnw"] > 3:
        add("warn", f"TOL/TNW {r['tol_tnw']}: apne paise ke mukable udhaar zyada.")
    if trend.get("revenue") is not None:
        add("info" if trend["revenue"] >= 0 else "warn", f"Bikri pichhle hisaab se {trend['revenue']:+.1f}%.")
    if trend.get("opex") is not None and trend.get("revenue") is not None and trend["opex"] > trend["revenue"] + 10:
        add("warn", f"Kharche {trend['opex']:+.1f}% badhe, bikri se tez.")
    if r.get("gross_margin") is not None and r["gross_margin"] < 10:
        add("warn", f"Gross margin sirf {r['gross_margin']}%: daam ya kharid ka bhaav dekhiye.")
    if payroll_cost and pl["revenue"]:
        add("info", f"Salary ka kharcha (Hoola payroll) {_lakh(payroll_cost)} = bikri ka {payroll_cost * 100 / pl['revenue']:.1f}%.")
    if not out:
        add("good", "Sab ratio theek seema mein hain.")
    return out


def snapshots(tenant: str) -> list[dict]:
    with session() as s:
        return [{"id": r.id, "label": r.label, "start": r.start, "end": r.end, "source": r.source, "status": r.status}
                for r in s.scalars(select(FinSnapshot).where(FinSnapshot.tenant_id == tenant).order_by(FinSnapshot.end.desc())).all()]


# ---------------------------------------------------------------- bank statement -> cash
_BH = {"date": ("date", "txn date", "transaction date", "value date", "tran date", "value dt", "txn dt"),
       "narration": ("narration", "description", "particulars", "remarks", "transaction details", "details", "transaction remarks"),
       "debit": ("debit", "withdrawal", "withdrawal amt.", "withdrawal amount", "withdrawals", "debit amount", "dr", "withdrawal (dr)", "debit(inr)"),
       "credit": ("credit", "deposit", "deposit amt.", "deposit amount", "deposits", "credit amount", "cr", "deposit (cr)", "credit(inr)"),
       "amount": ("amount", "transaction amount", "amount(inr)"), "type": ("dr/cr", "cr/dr", "type"),
       "balance": ("balance", "closing balance", "balance (inr)", "available balance", "balance(inr)")}


def _bank_rules(s, tenant: str) -> list[tuple[str, str]]:
    return [(r.pattern, r.category) for r in s.scalars(select(FinRule).where(FinRule.tenant_id == tenant, FinRule.kind == "bank")).all()]


def _cat(rules: list[tuple[str, str]], narration: str, credit: bool) -> str:
    n = narration.lower()
    for pat, cat in rules:
        if pat and pat in n:
            return cat
    return R.bank_category(narration, credit)


def import_bank(tenant: str, file_id: str = "", data: bytes | None = None, name: str = "statement.csv", account: str = "") -> dict[str, Any]:
    if file_id:
        name, data = _load_file(tenant, file_id)
    if not data:
        raise FinError("Bank statement (Excel / CSV) chahiye. PDF ho to bank se Excel / CSV download kariye.")
    rows = _rows(name, data)
    hi, cols = None, {}
    for i, r in enumerate(rows[:40]):
        h = [_norm(x) for x in r]
        found = {}
        for j, x in enumerate(h):
            for k, names in _BH.items():
                if x in names and k not in found:
                    found[k] = j
        if "date" in found and "narration" in found and (("debit" in found and "credit" in found) or "amount" in found):
            hi, cols = i, found
            break
    if hi is None:
        raise FinError("Statement mein Date, Narration aur Debit / Credit (ya Amount) column nahi mile.")
    added, dups = 0, 0
    with _LOCK, session() as s:
        rules = _bank_rules(s, tenant)
        for r in rows[hi + 1:]:
            g = lambda k: r[cols[k]] if k in cols and cols[k] < len(r) else None  # noqa: E731
            d = _date(g("date"))
            if not d:
                continue
            narr = str(g("narration") or "").strip()
            if "amount" in cols and "debit" not in cols:
                amt = R.clean_amount(g("amount"))
                t = str(g("type") or "").strip().lower()
                if t.startswith("d") or (not t and amt < 0):
                    amt = -abs(amt)
                else:
                    amt = abs(amt)
            else:
                amt = abs(R.clean_amount(g("credit"))) - abs(R.clean_amount(g("debit")))
            if not amt:
                continue
            bal = R.clean_amount(g("balance")) if g("balance") not in (None, "") else None
            fp = hashlib.sha1(f"{account}|{d}|{amt:.2f}|{narr}|{bal}".encode()).hexdigest()[:40]
            if s.scalars(select(FinTxn.id).where(FinTxn.tenant_id == tenant, FinTxn.fp == fp)).first():
                dups += 1
                continue
            s.add(FinTxn(tenant_id=tenant, account=account[:60], day=d.isoformat(), narration=narr[:300], amount=round(amt, 2),
                         balance=bal, category=_cat(rules, narr, amt > 0), fp=fp))
            added += 1
    out = cash(tenant)
    out["imported"], out["duplicates"] = added, dups
    out["say"] = f"{added} entry li" + (f" ({dups} pehle se thi)" if dups else "") + ". " + out["say"]
    return out


def set_category(tenant: str, pattern: str, category: str, txn_id: str = "") -> dict[str, Any]:
    """Learn: narration words -> category (applies to old and new entries)."""
    pat = str(pattern or "").strip().lower()[:120]
    category = re.sub(r"[^a-z_]", "", str(category).lower().replace(" ", "_"))[:30]
    if not category:
        raise FinError("Category chahiye (jaise rent, salary, raw_material...).")
    with session() as s:
        if txn_id:
            t = s.get(FinTxn, txn_id)
            if t and t.tenant_id == tenant:
                t.category = category
                pat = pat or re.sub(r"\d+", "", t.narration.lower()).strip()[:40]
        if pat:
            s.add(FinRule(tenant_id=tenant, kind="bank", pattern=pat, category=category))
            n = 0
            for t in s.scalars(select(FinTxn).where(FinTxn.tenant_id == tenant)).all():
                if pat in t.narration.lower():
                    t.category = category
                    n += 1
        else:
            n = 1
    return {"updated": n, "say": f"'{pattern or 'yeh entry'}' ab {category} mein ({n} entry badli). Aage se yahi."}


def cash(tenant: str, months: int = 3) -> dict[str, Any]:
    on = today(tenant)
    with session() as s:
        txns = s.scalars(select(FinTxn).where(FinTxn.tenant_id == tenant).order_by(FinTxn.day)).all()
        st = _st(s, tenant)
        budget = _j(st.budget, {})
        items = s.scalars(select(FinItem).where(FinItem.tenant_id == tenant)).all()
        rows = [(t.day, t.amount, t.category, t.balance, t.narration) for t in txns]
        open_items = [(i.kind, i.party, i.amount - i.paid, i.bill_date, i.credit_days) for i in items if i.amount - i.paid > 0.5]
    if not rows:
        raise FinError("Bank statement abhi nahi mila: Excel / CSV bhejiye.")
    last_day = _date(rows[-1][0])
    by_month: dict[str, dict] = {}
    for d, amt, cat, _b, _n in rows:
        m = d[:7]
        x = by_month.setdefault(m, {"in": 0.0, "out": 0.0, "cats": {}})
        if amt > 0:
            x["in"] += amt
        else:
            x["out"] += -amt
        x["cats"][cat] = x["cats"].get(cat, 0.0) + amt
    months_sorted = sorted(by_month)
    full = [m for m in months_sorted if m != last_day.isoformat()[:7] or last_day.day >= 28]
    use = (full or months_sorted)[-months:]
    avg_in = sum(by_month[m]["in"] for m in use) / len(use)
    avg_out = sum(by_month[m]["out"] for m in use) / len(use)
    bal = next((b for _d, _a, _c, b, _n in reversed(rows) if b is not None), None)
    if bal is None:
        bal = sum(a for _d, a, _c, _b, _n in rows)
    burn = avg_out - avg_in
    runway = R.runway_months(bal, burn)
    # recurring outflows: a category paid in most of the months used
    recurring = []
    cats = {c for m in use for c in by_month[m]["cats"]}
    for c in sorted(cats):
        vals = [by_month[m]["cats"].get(c, 0.0) for m in use]
        outs = [-v for v in vals if v < 0]
        if len(outs) >= max(2, len(use) - 1) and c not in ("supplier_payment", "misc", "cash_withdrawal"):
            recurring.append({"category": c, "monthly": round(sum(outs) / len(outs))})
    # budget vs actual (latest full month)
    variance = []
    if budget and use:
        lm = use[-1]
        for c, bud in budget.items():
            act = -by_month[lm]["cats"].get(c, 0.0)
            variance.append({"category": c, "budget": round(bud), "actual": round(act), "diff": round(act - bud)})
    # 13 weeks: average weekly in / out, plus bills due to collect / pay
    weeks = []
    run = bal
    w_in, w_out = avg_in * 12 / 52, avg_out * 12 / 52
    for w in range(13):
        ws = on + timedelta(days=7 * w)
        we = ws + timedelta(days=6)
        col = sum(a for k, _p, a, bd, cd in open_items if k == "receivable" and _due(bd, cd) and ws <= _due(bd, cd) <= we)
        pay = sum(a for k, _p, a, bd, cd in open_items if k == "payable" and _due(bd, cd) and ws <= _due(bd, cd) <= we)
        run = run + w_in - w_out          # averages already hold usual collections; bills due are shown beside
        weeks.append({"week_from": ws.isoformat(), "inflow": round(w_in), "outflow": round(w_out), "collect_due": round(col),
                      "pay_due": round(pay), "closing": round(run)})
    low = min(weeks, key=lambda x: x["closing"])
    say = (f"Bank: aaj {_lakh(bal)}. Pichhle {len(use)} mahine ka average: aaya {_lakh(avg_in)}, gaya {_lakh(avg_out)} har mahine. ")
    if runway is not None:
        say += f"Is raftaar se paisa lagbhag {runway} mahine chalega. "
    else:
        say += "Har mahine paisa badh raha hai. "
    if low["closing"] < 0:
        say += f"Dhyaan: {low['week_from']} wale hafte tak bank minus mein ja sakta hai. "
    if recurring:
        say += "Har mahine ke pakke kharche: " + ", ".join(f"{x['category']} {_rs(x['monthly'])}" for x in recurring[:5]) + "."
    return {"balance": round(bal), "avg_in": round(avg_in), "avg_out": round(avg_out), "burn": round(burn), "runway_months": runway,
            "months": {m: {"in": round(by_month[m]["in"]), "out": round(by_month[m]["out"]),
                           "cats": {c: round(v) for c, v in sorted(by_month[m]["cats"].items(), key=lambda kv: kv[1])}} for m in months_sorted[-12:]},
            "recurring": recurring, "budget_vs_actual": variance, "forecast": weeks, "lowest_week": low, "say": say}


def _due(bill_date: str, days: int) -> date | None:
    d = _date(bill_date)
    return d + timedelta(days=days or 0) if d else None


def txns(tenant: str, category: str = "", limit: int = 200) -> list[dict]:
    with session() as s:
        q = select(FinTxn).where(FinTxn.tenant_id == tenant)
        if category:
            q = q.where(FinTxn.category == category)
        return [{"id": t.id, "day": t.day, "narration": t.narration, "amount": t.amount, "category": t.category}
                for t in s.scalars(q.order_by(FinTxn.day.desc()).limit(limit)).all()]


# ---------------------------------------------------------------- receivables / payables
def add_item(tenant: str, kind: str, party: str, amount, bill: str = "", bill_date: str = "", credit_days: int | None = None,
             phone: str = "", msme: bool = False, agreed_days: int | None = None, paid=None) -> dict[str, Any]:
    kind = "payable" if str(kind).lower() in ("payable", "pay", "dena", "creditor") else "receivable"
    d = _date(bill_date) or today(tenant)
    amt = abs(R.clean_amount(amount))
    if not amt or not str(party).strip():
        raise FinError("Party aur rakam chahiye.")
    with session() as s:
        it = FinItem(tenant_id=tenant, kind=kind, party=str(party).strip()[:160], phone=re.sub(r"[^\d+]", "", str(phone or ""))[:20],
                     bill=str(bill or "")[:40], bill_date=d.isoformat(), amount=amt, paid=abs(R.clean_amount(paid)) if paid else 0.0,
                     credit_days=int(credit_days) if credit_days is not None else (R.msme_limit(agreed_days) if kind == "payable" and msme else 30),
                     msme=bool(msme), agreed_days=int(agreed_days or 0))
        s.add(it)
        s.flush()
        iid = it.id
    say = f"{'Lena' if kind == 'receivable' else 'Dena'}: {party} {_rs(amt)} ({d.isoformat()})."
    if kind == "payable" and msme:
        lim = R.msme_limit(agreed_days)
        say += f" MSME supplier: {lim} din mein dena zaroori ({(d + timedelta(days=lim)).isoformat()} tak), warna kharcha tax mein nahi ghatega."
    return {"id": iid, "say": say}


def import_items(tenant: str, kind: str, file_id: str = "", data: bytes | None = None, name: str = "outstanding.csv") -> dict[str, Any]:
    if file_id:
        name, data = _load_file(tenant, file_id)
    rows = _rows(name, data or b"")
    head = {"party": ("party", "party name", "name", "customer", "supplier", "ledger", "particulars"),
            "bill": ("bill no", "bill", "invoice", "invoice no", "ref no", "ref. no.", "voucher no"),
            "bill_date": ("date", "bill date", "invoice date"), "amount": ("pending amount", "amount", "balance", "outstanding", "pending"),
            "credit_days": ("due days", "credit days", "credit period"), "phone": ("phone", "mobile")}
    hi, cols = None, {}
    for i, r in enumerate(rows[:15]):
        h = [_norm(x) for x in r]
        f = {}
        for j, x in enumerate(h):
            for k, ns in head.items():
                if x in ns and k not in f:
                    f[k] = j
        if "party" in f and "amount" in f:
            hi, cols = i, f
            break
    if hi is None:
        raise FinError("Outstanding list mein Party aur Amount column chahiye.")
    n = 0
    for r in rows[hi + 1:]:
        g = lambda k: r[cols[k]] if k in cols and cols[k] < len(r) else None  # noqa: E731
        if not g("party") or not R.clean_amount(g("amount")):
            continue
        try:
            add_item(tenant, kind, str(g("party")), g("amount"), str(g("bill") or ""), str(g("bill_date") or ""),
                     int(R.clean_amount(g("credit_days"))) if g("credit_days") else None, str(g("phone") or ""))
            n += 1
        except FinError:
            continue
    return {"saved": n, "say": f"{n} bill ki list li."}


def item_paid(tenant: str, item_id: str, amount=None) -> dict:
    with session() as s:
        it = s.get(FinItem, item_id)
        if not it or it.tenant_id != tenant:
            raise FinError("Bill nahi mila.")
        it.paid = min(it.amount, it.paid + (abs(R.clean_amount(amount)) if amount else it.amount - it.paid))
        return {"say": f"{it.party}: {_rs(it.paid)} aa gaya / de diya, baaki {_rs(it.amount - it.paid)}."}


def aging(tenant: str, kind: str = "receivable") -> dict[str, Any]:
    on = today(tenant)
    with session() as s:
        its = s.scalars(select(FinItem).where(FinItem.tenant_id == tenant, FinItem.kind == kind)).all()
        rows = []
        for i in its:
            left = round(i.amount - i.paid, 2)
            if left <= 0.5:
                continue
            bd = _date(i.bill_date) or on
            age = (on - bd).days
            due = bd + timedelta(days=i.credit_days)
            row = {"id": i.id, "party": i.party, "phone": i.phone, "bill": i.bill, "bill_date": i.bill_date, "pending": left, "age": age,
                   "bucket": R.aging_bucket(age), "due": due.isoformat(), "overdue_days": max(0, (on - due).days), "msme": i.msme}
            if kind == "payable" and i.msme:
                lim = R.msme_limit(i.agreed_days)
                last = bd + timedelta(days=lim)
                row["msme_last_day"] = last.isoformat()
                row["msme_days_left"] = (last - on).days
                if on > last:
                    row["msme_interest"] = R.msme_interest(left, (on - last).days)
                    row["tax_risk"] = "Sec 37(2)(g): jab tak nahi dete, yeh kharcha tax mein nahi ghatega."
            rows.append(row)
    rows.sort(key=lambda x: -x["age"])
    buckets = {b: round(sum(x["pending"] for x in rows if x["bucket"] == b)) for b in ("0-30", "31-60", "61-90", "90+")}
    total = round(sum(x["pending"] for x in rows))
    if kind == "receivable":
        say = f"Lene hain {_lakh(total)}: " + ", ".join(f"{b} din {_lakh(v)}" for b, v in buckets.items() if v) + "."
        big = [x for x in rows if x["overdue_days"] > 0][:3]
        if big:
            say += " Sabse purane: " + ", ".join(f"{x['party']} {_rs(x['pending'])} ({x['age']} din)" for x in big) + "."
    else:
        risk = [x for x in rows if x.get("msme") and x.get("msme_days_left", 99) <= 5]
        say = f"Dene hain {_lakh(total)}."
        if risk:
            say += " MSME supplier jinhe jaldi dena hai: " + ", ".join(
                f"{x['party']} {_rs(x['pending'])} ({'beet gaya' if x['msme_days_left'] < 0 else str(x['msme_days_left']) + ' din'})" for x in risk[:5]) + "."
    return {"kind": kind, "items": rows, "buckets": buckets, "total": total, "say": say}


def reminder_drafts(tenant: str, min_overdue: int = 1) -> dict[str, Any]:
    """Polite payment reminders to customers whose bills are overdue (drafts: she sends)."""
    a = aging(tenant, "receivable")
    with session() as s:
        biz = _st(s, tenant).business
    by: dict[str, list] = {}
    for x in a["items"]:
        if x["overdue_days"] >= min_overdue:
            by.setdefault(x["party"], []).append(x)
    out = []
    for party, xs in by.items():
        bills = ", ".join(f"{x['bill'] or x['bill_date']} ({_rs(x['pending'])})" for x in xs[:8])
        tot = sum(x["pending"] for x in xs)
        out.append({"party": party, "to": xs[0]["phone"], "amount": round(tot),
                    "text": f"Namaste {party} ji, {biz or 'hamare'} ke bill {bills} ka kul {_rs(tot)} baaki hai. Kripya jaldi bhej dijiye. Dhanyavaad."})
    return {"drafts": out, "say": f"{len(out)} graahak ke liye yaad dilane ka message draft (aapki haan ke baad hi jayega)."}


# ---------------------------------------------------------------- loans, loan check, advance tax
def add_loan(tenant: str, name: str, outstanding, rate, emi=None, months_left: int | None = None) -> dict[str, Any]:
    o, r = abs(R.clean_amount(outstanding)), float(rate or 0)
    e = abs(R.clean_amount(emi)) if emi else (R.emi(o, r, int(months_left or 60)))
    with session() as s:
        s.add(FinLoan(tenant_id=tenant, name=str(name)[:80], outstanding=o, rate=r, emi=e, months_left=int(months_left or 0)))
    return {"say": f"Loan {name}: baaki {_lakh(o)}, {r:g}% par, kist {_rs(e)}."}


def loan_check(tenant: str, amount, rate, years: float, projected_turnover=None, snapshot: str = "") -> dict[str, Any]:
    amt = abs(R.clean_amount(amount))
    months = int(round(float(years) * 12))
    new_emi = R.emi(amt, float(rate), months)
    with session() as s:
        try:
            sn = _snap(s, tenant, snapshot)
        except FinError:
            sn = None
        loans = s.scalars(select(FinLoan).where(FinLoan.tenant_id == tenant)).all()
        cur_emi = sum(l.emi for l in loans)
        if sn:
            st = R.statements(_j(sn.balances, {}))
            days = max(1, (_date(sn.end) - _date(sn.start)).days + 1)
            yearly = (st["pl"]["profit_before_tax"] + st["pl"]["depreciation"] + st["pl"]["interest"]) * 365 / days
            turnover = st["pl"]["revenue"] * 365 / days
        else:
            yearly, turnover = None, None
    out: dict[str, Any] = {"amount": amt, "rate": rate, "months": months, "emi": new_emi, "existing_emi": round(cur_emi)}
    if yearly is not None:
        ds_new = (cur_emi + new_emi) * 12
        out["dscr_after"] = round(yearly / ds_new, 2) if ds_new else None
        out["cash_for_debt_yearly"] = round(yearly)
    pt = abs(R.clean_amount(projected_turnover)) if projected_turnover else (turnover * 1.1 if turnover else None)
    if pt:
        out["nayak"] = R.nayak_limit(pt)
    verdict = "pata nahi"
    if out.get("dscr_after") is not None:
        verdict = "theek" if out["dscr_after"] >= 1.5 else "kheench ke" if out["dscr_after"] >= 1.25 else "mushkil"
    out["verdict"] = verdict
    say = f"{_lakh(amt)} ka loan {rate}% par {years:g} saal: kist {_rs(new_emi)} har mahine."
    if out.get("dscr_after") is not None:
        say += f" Sab kiston ke baad DSCR {out['dscr_after']} ({verdict}; bank 1.25+ dekhte hain)."
    if out.get("nayak"):
        say += f" Working capital (CC / OD) {_lakh(out['nayak']['bank_limit'])} tak ban sakta hai (bikri ka 20%)."
    say += " Yeh hisaab hai, salah nahi: final faisla bank aur aapka."
    out["say"] = say
    return out


def advance_tax(tenant: str, profit=None, snapshot: str = "") -> dict[str, Any]:
    on = today(tenant)
    fs = R.fy_start_of(on)
    fy = f"{fs}-{str(fs + 1)[2:]}"
    with session() as s:
        st = _st(s, tenant)
        entity = st.entity
        paid = float(_j(st.tax_paid, {}).get(fy, 0.0))
        est = profit
        basis = "aapka andaza"
        if est in (None, ""):
            sn = _snap(s, tenant, snapshot)
            stt = R.statements(_j(sn.balances, {}))
            days = max(1, (_date(sn.end) - _date(sn.start)).days + 1)
            est = stt["pl"]["profit_before_tax"] * 365 / days
            basis = f"{sn.label} ke munafe se poore saal ka andaza"
    t = R.tax_estimate(float(est), entity)
    plan = R.advance_tax_plan(t["tax"], fs, paid, on)
    say = f"Saal {fy} ka munafa ~{_lakh(t['profit'])} ({basis}), tax ~{_lakh(t['tax'])} ({entity})."
    if plan["needed"] and plan.get("next"):
        say += f" Agli kist {plan['next']['due']} tak kul {plan['next']['cumulative_pct']}% = {_rs(plan['next']['cumulative'])}; ab tak {_rs(paid)} bhara, to {_rs(plan['pay_by_next'])} aur."
    elif not plan["needed"]:
        say += " " + plan["note"]
    say += " Challan aap khud bharengi (ITNS 280); CA se pakka karwa lijiye."
    return {"fy": fy, "entity": entity, **t, "paid": paid, "plan": plan, "say": say}


# ---------------------------------------------------------------- MIS report + her Haan before sharing
def report(tenant: str, snapshot: str = "", conv: str = "") -> dict[str, Any]:
    import openpyxl
    from .. import chatfiles
    a = analyse(tenant, snapshot)
    try:
        c = cash(tenant)
    except FinError:
        c = None
    rec = aging(tenant, "receivable")
    pay = aging(tenant, "payable")
    with session() as s:
        biz = _st(s, tenant).business
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Summary"
    ws.append([biz or "MIS", a["snapshot"]["label"], f"{a['snapshot']['start']} to {a['snapshot']['end']}"])
    ws.append([])
    for k, v in a["pl"].items():
        ws.append([k.replace("_", " ").title(), v])
    ws.append([])
    for i in a["insights"]:
        ws.append([i["level"], i["text"]])
    w2 = wb.create_sheet("Balance sheet")
    for k, v in a["bs"].items():
        w2.append([k.replace("_", " ").title(), v])
    w3 = wb.create_sheet("Ratios")
    w3.append(["Ratio", "Value", "Benchmark", "Rule", "Verdict"])
    for r in a["ratios"]:
        w3.append([r["label"], r["value"], r["benchmark"], r["rule"], r["verdict"]])
    if c:
        w4 = wb.create_sheet("Cash")
        w4.append(["Month", "In", "Out"])
        for m, x in c["months"].items():
            w4.append([m, x["in"], x["out"]])
        w4.append([])
        w4.append(["Week from", "Inflow", "Outflow", "Closing"])
        for w in c["forecast"]:
            w4.append([w["week_from"], w["inflow"], w["outflow"], w["closing"]])
    for title, ag in (("Receivables", rec), ("Payables", pay)):
        w5 = wb.create_sheet(title)
        w5.append(["Party", "Bill", "Date", "Pending", "Days", "Bucket", "MSME"])
        for x in ag["items"]:
            w5.append([x["party"], x["bill"], x["bill_date"], x["pending"], x["age"], x["bucket"], "Y" if x.get("msme") else ""])
    bio = io.BytesIO()
    wb.save(bio)
    files = [chatfiles.store(tenant, f"MIS_{a['snapshot']['label'].replace(' ', '_')}.xlsx", bio.getvalue(),
                             "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")]
    files.append(chatfiles.store(tenant, f"MIS_{a['snapshot']['label'].replace(' ', '_')}.pdf", _pdf(biz, a, c, rec), "application/pdf"))
    for f in files:
        _publish(tenant, "pc_file", {**f, "source": "analyst", "title": f.get("name", "")}, conv)
    return {"files": files, "analysis": a, "say": "MIS report (Excel + 1 page PDF) taiyaar. Bank / CA ko bhejna ho to pehle aapki haan."}


def _pdf(biz: str, a: dict, c: dict | None, rec: dict) -> bytes:
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.units import mm
    from reportlab.pdfgen import canvas
    bio = io.BytesIO()
    cv = canvas.Canvas(bio, pagesize=A4)
    W, H = A4
    y = H - 20 * mm
    cv.setFont("Helvetica-Bold", 15)
    cv.drawString(20 * mm, y, f"{biz or 'Business'}: MIS {a['snapshot']['label']}")
    y -= 8 * mm
    cv.setFont("Helvetica", 9)
    cv.drawString(20 * mm, y, f"{a['snapshot']['start']} to {a['snapshot']['end']}  |  prepared by Hoola (analysis, not financial advice)")
    y -= 10 * mm
    cv.setFont("Helvetica-Bold", 11)
    cv.drawString(20 * mm, y, "Profit & loss")
    y -= 6 * mm
    cv.setFont("Helvetica", 10)
    for k, v in a["pl"].items():
        cv.drawString(25 * mm, y, k.replace("_", " ").title())
        cv.drawRightString(100 * mm, y, f"{v:,.0f}")
        y -= 5.5 * mm
    yy = H - 38 * mm
    cv.setFont("Helvetica-Bold", 11)
    cv.drawString(115 * mm, yy, "Ratios")
    yy -= 6 * mm
    cv.setFont("Helvetica", 9)
    for r in a["ratios"]:
        if r["value"] is None:
            continue
        cv.drawString(115 * mm, yy, r["label"][:34])
        cv.drawRightString(W - 20 * mm, yy, f"{r['value']}" + (" !" if r["verdict"] == "weak" else ""))
        yy -= 5 * mm
    y = min(y, yy) - 8 * mm
    cv.setFont("Helvetica-Bold", 11)
    cv.drawString(20 * mm, y, "What stands out")
    y -= 6 * mm
    cv.setFont("Helvetica", 9.5)
    for i in a["insights"][:8]:
        cv.drawString(25 * mm, y, ("- " + i["text"])[:110])
        y -= 5 * mm
    if c:
        y -= 4 * mm
        cv.setFont("Helvetica-Bold", 11)
        cv.drawString(20 * mm, y, "Cash")
        y -= 6 * mm
        cv.setFont("Helvetica", 9.5)
        cv.drawString(25 * mm, y, f"Bank balance {c['balance']:,.0f} | avg in {c['avg_in']:,.0f} / out {c['avg_out']:,.0f} a month"
                      + (f" | runway {c['runway_months']} months" if c["runway_months"] else ""))
        y -= 5 * mm
    y -= 4 * mm
    cv.setFont("Helvetica-Bold", 11)
    cv.drawString(20 * mm, y, "Receivables")
    y -= 6 * mm
    cv.setFont("Helvetica", 9.5)
    cv.drawString(25 * mm, y, f"Total {rec['total']:,.0f} | " + " | ".join(f"{k} days {v:,.0f}" for k, v in rec["buckets"].items()))
    cv.showPage()
    cv.save()
    return bio.getvalue()


def approve_share(tenant: str, snapshot: str, approve: bool, to: str = "") -> dict[str, Any]:
    """Her Haan to share this period's report outside (bank / CA / investor). Hoola then drafts the note."""
    with session() as s:
        sn = _snap(s, tenant, snapshot)
        sn.status = "approved" if approve else "ready"
        biz = _st(s, tenant).business
        label = sn.label
    if not approve:
        return {"say": "Theek, report bahar nahi jayegi."}
    note = (f"Namaste, {biz or 'hamari firm'} ki {label} ki MIS report (P&L, balance sheet, ratios, cash, receivables) saath mein hai. "
            f"Kuch poochna ho to bataiye. Dhanyavaad.")
    return {"approved": True, "draft": note, "to": to,
            "say": "Haan mili: report aur note taiyaar. Email / WhatsApp se aap bhejengi (ya Hoola email draft bana dega, bhejne se pehle aap dekh lengi)."}


# ---------------------------------------------------------------- morning
def alerts(tenant: str) -> dict[str, Any]:
    on = today(tenant)
    out = []
    try:
        pay = aging(tenant, "payable")
        for x in pay["items"]:
            if x.get("msme") and x.get("msme_days_left") is not None and x["msme_days_left"] <= 5:
                if x["msme_days_left"] < 0:
                    out.append(f"MSME {x['party']} ka {_rs(x['pending'])} {-x['msme_days_left']} din late: byaaj ~{_rs(x.get('msme_interest', 0))}, "
                               "aur jab tak nahi dete kharcha tax mein nahi ghatega (Sec 37(2)(g)).")
                else:
                    out.append(f"MSME {x['party']} ko {_rs(x['pending'])} {x['msme_last_day']} tak dena (Sec 37(2)(g)).")
        rec = aging(tenant, "receivable")
        old = [x for x in rec["items"] if x["overdue_days"] >= 30]
        if old:
            out.append(f"{len(old)} graahak 30+ din se late: kul {_rs(sum(x['pending'] for x in old))}.")
    except FinError:
        pass
    fs = R.fy_start_of(on)
    for d in (date(fs, 6, 15), date(fs, 9, 15), date(fs, 12, 15), date(fs + 1, 3, 15)):
        if 0 <= (d - on).days <= 5:
            out.append(f"Advance tax ki kist {d.strftime('%d %b')} ko.")
    return {"alerts": out, "text": " ".join(out)}


def due_mornings() -> list[str]:
    from ..life import tz_of
    res = []
    with session() as s:
        for r in s.scalars(select(FinSettings).where(FinSettings.enabled.is_(True))).all():
            local = datetime.now(tz_of(r.tenant_id))
            d = local.date().isoformat()
            try:
                hh, mm = [int(x) for x in (r.morning_time or "10:15").split(":")]
            except ValueError:
                hh, mm = 10, 15
            late = (local.hour * 60 + local.minute) - (hh * 60 + mm)
            if r.last_morning != d and 0 <= late <= 180:
                r.last_morning = d
                res.append(r.tenant_id)
    return res


def send_morning(tenant: str) -> dict[str, Any]:
    from ..events import notify
    from ..net.service import home_conv
    a = alerts(tenant)
    if not a["alerts"]:
        return a
    conv = home_conv(tenant)
    try:
        notify(conv, "Paisa: " + a["text"])
    except Exception:
        pass
    _publish(tenant, "fin_alert", {"text": a["text"]}, conv)
    return a
