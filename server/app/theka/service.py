"""v0.117 Theka Sahayak: Hoola as a contractor's coworker, end to end.

  1. Gets the job: a tender she found (portal, ID, last date, EMD), "PWD road ka kaam mila, BOQ yeh hai", the site
     engineer's measurements ("PCC 12.5 cum, chainage 1+200"), "RA bill 3 banao", "department ne 4.2 lakh pass kiye",
     a subcontractor's bill, "aaj 18 mazdoor, baarish se kaam ruka", "cement kitna bacha".
  2. Does the work, step by step: tender checklist + dates + EMD; project terms (value, dates, GST, retention, security
     deposit, TDS, GST TDS, labour cess, mobilisation advance, DLP, BG, insurance); BOQ; measurement book entries against
     each item; RA bill = cumulative measured x rate minus what was billed before, GST, every deduction, net; bill passed /
     paid tracking; retention held and when it comes back; subcontractor bills with Sec 393(1) TDS; material in / out;
     site diary with labour, work done and hindrances.
  3. Checked against the rules and earlier work: BOQ quantity and deviation limit, extra items, GST 18% works contract,
     GST TDS 2% above Rs 2.5 lakh, income-tax TDS 1% / 2%, BOCW cess 1%, mobilisation recovery not beyond what is left,
     time used vs work done, hindrance days for an extension, BG / insurance / registration / tender dates.
  4. Ready as a card, a file or a draft: RA bill PDF (abstract + measurements), deviation statement, EOT letter draft,
     subcontractor payment sheet, receivables list, project register Excel.
  5. Waits for her Haan: an RA bill is marked submitted and a letter goes to the department only on her Haan (owner).
     Hoola never submits a bid (her DSC), never pays anyone, never inflates a measurement.
  6. Delivers and learns: passed vs claimed (department cuts), TDS credits to check in 26AS / GST cash ledger, progress
     vs time per project, material use per bill.
"""
from __future__ import annotations

import io
import json
import logging
import re
import threading
from datetime import date, datetime, timedelta
from typing import Any

from sqlalchemy import Boolean, DateTime, Float, Integer, String, Text, select
from sqlalchemy.orm import Mapped, mapped_column

from ..db import Base, new_id, now, session
from . import rules as R

log = logging.getLogger("santosh.theka")
_LOCK = threading.Lock()


class ThekaError(Exception):
    pass


class ThekaSettings(Base):
    __tablename__ = "theka_settings"
    tenant_id: Mapped[str] = mapped_column(String(64), primary_key=True)
    firm: Mapped[str] = mapped_column(String(160), default="")
    gstin: Mapped[str] = mapped_column(String(15), default="")
    individual: Mapped[bool] = mapped_column(Boolean, default=False)       # proprietor / HUF -> TDS 1% on her bills
    udyam: Mapped[str] = mapped_column(String(30), default="")
    registration: Mapped[str] = mapped_column(String(120), default="")      # e.g. "PWD Class I-A, No ..."
    registration_valid_to: Mapped[str] = mapped_column(String(10), default="")
    address: Mapped[str] = mapped_column(Text, default="")
    morning_time: Mapped[str] = mapped_column(String(5), default="08:00")
    last_morning: Mapped[str] = mapped_column(String(10), default="")


class ThekaTender(Base):
    __tablename__ = "theka_tenders"
    id: Mapped[str] = mapped_column(String(40), primary_key=True, default=lambda: new_id("tt"))
    tenant_id: Mapped[str] = mapped_column(String(64), index=True)
    name: Mapped[str] = mapped_column(String(300), default="")
    dept: Mapped[str] = mapped_column(String(160), default="")
    portal: Mapped[str] = mapped_column(String(80), default="")
    ref: Mapped[str] = mapped_column(String(80), default="")
    value: Mapped[float] = mapped_column(Float, default=0.0)
    emd: Mapped[float] = mapped_column(Float, default=0.0)
    emd_mode: Mapped[str] = mapped_column(String(12), default="")          # dd | bg | online | declaration | exempt
    prebid: Mapped[str] = mapped_column(String(10), default="")
    last_date: Mapped[str] = mapped_column(String(16), default="")          # YYYY-MM-DD[ HH:MM]
    opening: Mapped[str] = mapped_column(String(10), default="")
    docs: Mapped[str] = mapped_column(Text, default="{}")
    status: Mapped[str] = mapped_column(String(12), default="watching")     # watching | preparing | submitted | won | lost | dropped
    our_quote: Mapped[float] = mapped_column(Float, default=0.0)
    l1: Mapped[float] = mapped_column(Float, default=0.0)
    emd_back: Mapped[bool] = mapped_column(Boolean, default=False)
    notes: Mapped[str] = mapped_column(Text, default="")


class ThekaProject(Base):
    __tablename__ = "theka_projects"
    id: Mapped[str] = mapped_column(String(40), primary_key=True, default=lambda: new_id("tj"))
    tenant_id: Mapped[str] = mapped_column(String(64), index=True)
    name: Mapped[str] = mapped_column(String(200), default="")
    client: Mapped[str] = mapped_column(String(160), default="")
    govt: Mapped[bool] = mapped_column(Boolean, default=True)
    agreement_no: Mapped[str] = mapped_column(String(80), default="")
    value: Mapped[float] = mapped_column(Float, default=0.0)
    start: Mapped[str] = mapped_column(String(10), default="")
    end: Mapped[str] = mapped_column(String(10), default="")                # stipulated completion
    completed_on: Mapped[str] = mapped_column(String(10), default="")
    dlp_months: Mapped[int] = mapped_column(Integer, default=12)
    terms: Mapped[str] = mapped_column(Text, default="{}")                 # gst_rate, retention_pct, sd_pct, it_tds_pct, gst_tds, cess_pct, mob_pct, deviation_pct
    mob_advance: Mapped[float] = mapped_column(Float, default=0.0)
    bg: Mapped[str] = mapped_column(Text, default="{}")                    # {no, amount, valid_to}
    insurance_valid_to: Mapped[str] = mapped_column(String(10), default="")
    status: Mapped[str] = mapped_column(String(10), default="running")
    created_at: Mapped[datetime] = mapped_column(DateTime(timezone=True), default=now)


class ThekaItem(Base):
    __tablename__ = "theka_boq"
    id: Mapped[str] = mapped_column(String(40), primary_key=True, default=lambda: new_id("tq"))
    tenant_id: Mapped[str] = mapped_column(String(64), index=True)
    project_id: Mapped[str] = mapped_column(String(40), index=True)
    code: Mapped[str] = mapped_column(String(20), default="")
    description: Mapped[str] = mapped_column(Text, default="")
    unit: Mapped[str] = mapped_column(String(12), default="")
    qty: Mapped[float] = mapped_column(Float, default=0.0)
    rate: Mapped[float] = mapped_column(Float, default=0.0)
    extra: Mapped[bool] = mapped_column(Boolean, default=False)


class ThekaMeasure(Base):
    __tablename__ = "theka_mb"
    id: Mapped[str] = mapped_column(String(40), primary_key=True, default=lambda: new_id("tm"))
    tenant_id: Mapped[str] = mapped_column(String(64), index=True)
    project_id: Mapped[str] = mapped_column(String(40), index=True)
    item_id: Mapped[str] = mapped_column(String(40), index=True)
    day: Mapped[str] = mapped_column(String(10), default="")
    qty: Mapped[float] = mapped_column(Float, default=0.0)
    detail: Mapped[str] = mapped_column(String(300), default="")           # L x B x D / chainage / location
    by: Mapped[str] = mapped_column(String(60), default="")
    bill_id: Mapped[str] = mapped_column(String(40), default="")            # set when billed


class ThekaBill(Base):
    __tablename__ = "theka_bills"
    id: Mapped[str] = mapped_column(String(40), primary_key=True, default=lambda: new_id("tb"))
    tenant_id: Mapped[str] = mapped_column(String(64), index=True)
    project_id: Mapped[str] = mapped_column(String(40), index=True)
    number: Mapped[int] = mapped_column(Integer, default=1)
    final: Mapped[bool] = mapped_column(Boolean, default=False)
    upto: Mapped[str] = mapped_column(String(10), default="")
    data: Mapped[str] = mapped_column(Text, default="{}")
    status: Mapped[str] = mapped_column(String(10), default="draft")       # draft | submitted | passed | paid
    submitted_on: Mapped[str] = mapped_column(String(10), default="")
    passed: Mapped[float] = mapped_column(Float, default=0.0)              # net passed
    paid: Mapped[float] = mapped_column(Float, default=0.0)
    paid_on: Mapped[str] = mapped_column(String(10), default="")


class ThekaSub(Base):
    __tablename__ = "theka_subs"
    id: Mapped[str] = mapped_column(String(40), primary_key=True, default=lambda: new_id("ts"))
    tenant_id: Mapped[str] = mapped_column(String(64), index=True)
    name: Mapped[str] = mapped_column(String(160), default="")
    phone: Mapped[str] = mapped_column(String(20), default="")
    individual: Mapped[bool] = mapped_column(Boolean, default=True)
    pan_ok: Mapped[bool] = mapped_column(Boolean, default=True)
    work: Mapped[str] = mapped_column(String(160), default="")
    retention_pct: Mapped[float] = mapped_column(Float, default=0.0)


class ThekaSubBill(Base):
    __tablename__ = "theka_sub_bills"
    id: Mapped[str] = mapped_column(String(40), primary_key=True, default=lambda: new_id("tx"))
    tenant_id: Mapped[str] = mapped_column(String(64), index=True)
    sub_id: Mapped[str] = mapped_column(String(40), index=True)
    project_id: Mapped[str] = mapped_column(String(40), default="")
    day: Mapped[str] = mapped_column(String(10), default="")
    amount: Mapped[float] = mapped_column(Float, default=0.0)
    tds: Mapped[float] = mapped_column(Float, default=0.0)
    tds_on: Mapped[float] = mapped_column(Float, default=0.0)
    retention: Mapped[float] = mapped_column(Float, default=0.0)
    net: Mapped[float] = mapped_column(Float, default=0.0)
    paid: Mapped[bool] = mapped_column(Boolean, default=False)
    note: Mapped[str] = mapped_column(String(300), default="")


class ThekaMaterial(Base):
    __tablename__ = "theka_material"
    id: Mapped[str] = mapped_column(String(40), primary_key=True, default=lambda: new_id("tl"))
    tenant_id: Mapped[str] = mapped_column(String(64), index=True)
    project_id: Mapped[str] = mapped_column(String(40), index=True)
    item: Mapped[str] = mapped_column(String(80), default="")
    unit: Mapped[str] = mapped_column(String(12), default="")
    qty: Mapped[float] = mapped_column(Float, default=0.0)                 # + in, - used
    rate: Mapped[float] = mapped_column(Float, default=0.0)
    day: Mapped[str] = mapped_column(String(10), default="")
    note: Mapped[str] = mapped_column(String(200), default="")


class ThekaDiary(Base):
    __tablename__ = "theka_diary"
    id: Mapped[str] = mapped_column(String(40), primary_key=True, default=lambda: new_id("td"))
    tenant_id: Mapped[str] = mapped_column(String(64), index=True)
    project_id: Mapped[str] = mapped_column(String(40), index=True)
    day: Mapped[str] = mapped_column(String(10), index=True)
    labour: Mapped[int] = mapped_column(Integer, default=0)
    work: Mapped[str] = mapped_column(Text, default="")
    hindrance: Mapped[str] = mapped_column(Text, default="")               # empty = none
    lost_hours: Mapped[float] = mapped_column(Float, default=0.0)
    by_client: Mapped[bool] = mapped_column(Boolean, default=False)        # hindrance on the department's side (counts for EOT)


# ---------------------------------------------------------------- helpers
def _j(s: str, default):
    try:
        v = json.loads(s or "")
        return v if isinstance(v, type(default)) else default
    except (TypeError, ValueError):
        return default


def _num(v) -> float:
    if v in (None, ""):
        return 0.0
    s = str(v).replace(",", "").replace("₹", "")
    m = re.search(r"-?\d+(\.\d+)?", s)
    n = float(m.group(0)) if m else 0.0
    low = s.lower()
    if "crore" in low or re.search(r"\bcr\b", low):
        n *= 1e7
    elif "lakh" in low or "lac" in low:
        n *= 1e5
    return n


def _rs(v) -> str:
    v = round(float(v or 0))
    s = str(abs(v))
    if len(s) > 3:
        s = re.sub(r"(\d)(?=(\d\d)+$)", r"\1,", s[:-3]) + "," + s[-3:]
    return ("-" if v < 0 else "") + "₹" + s


def today(tenant: str) -> date:
    from ..life import tz_of
    return datetime.now(tz_of(tenant)).date()


def _date(v, default: date | None = None) -> date | None:
    if isinstance(v, date):
        return v
    if v in (None, ""):
        return default
    from ..analyst.service import _date as d
    return d(str(v)[:10]) or default


def add_months(d: date, n: int) -> date:
    from calendar import monthrange
    m = d.month - 1 + n
    y = d.year + m // 12
    m = m % 12 + 1
    return date(y, m, min(d.day, monthrange(y, m)[1]))


def _fy_start(d: date) -> date:
    return date(d.year if d.month >= 4 else d.year - 1, 4, 1)


def _publish(tenant: str, kind: str, payload: dict, conv: str = "") -> None:
    from ..events import bus
    try:
        if not conv:
            from ..net.service import home_conv
            conv = home_conv(tenant)
        bus.publish(conv, kind, payload)
    except Exception:
        log.exception("theka publish failed")


def _st(s, tenant: str) -> ThekaSettings:
    r = s.get(ThekaSettings, tenant)
    if r is None:
        r = ThekaSettings(tenant_id=tenant)
        s.add(r)
        s.flush()
    return r


def settings(tenant: str) -> dict:
    with session() as s:
        r = _st(s, tenant)
        return {k: getattr(r, k) for k in ("firm", "gstin", "individual", "udyam", "registration", "registration_valid_to", "address", "morning_time")}


def save_settings(tenant: str, **kw: Any) -> dict:
    from ..gstp.rules import valid_gstin
    with session() as s:
        r = _st(s, tenant)
        if kw.get("gstin"):
            g = str(kw["gstin"]).strip().upper()
            if not valid_gstin(g):
                raise ThekaError(f"GSTIN galat lagta hai: {g}")
            r.gstin = g
        for k in ("firm", "udyam", "registration", "address"):
            if kw.get(k) is not None:
                setattr(r, k, str(kw[k]).strip()[:300])
        if "individual" in kw and kw["individual"] is not None:
            r.individual = bool(kw["individual"])
        if kw.get("registration_valid_to"):
            d = _date(kw["registration_valid_to"])
            if not d:
                raise ThekaError("Registration ki tareekh samajh nahi aayi.")
            r.registration_valid_to = d.isoformat()
        if kw.get("morning_time"):
            r.morning_time = str(kw["morning_time"]).zfill(5)[:5]
    return settings(tenant)


# ---------------------------------------------------------------- tenders
def _tender_view(t: ThekaTender, on: date) -> dict:
    ld = _date(t.last_date)
    docs = _j(t.docs, {})
    return {"id": t.id, "name": t.name, "dept": t.dept, "portal": t.portal, "ref": t.ref, "value": t.value, "emd": t.emd, "emd_mode": t.emd_mode,
            "prebid": t.prebid, "last_date": t.last_date, "opening": t.opening, "status": t.status, "our_quote": t.our_quote, "l1": t.l1,
            "emd_back": t.emd_back, "docs": docs, "missing": [d for d, ok in docs.items() if not ok],
            "alert": R.tender_alert(ld, on) if ld and t.status in ("watching", "preparing") else "", "days_left": (ld - on).days if ld else None,
            "notes": t.notes}


def save_tender(tenant: str, name: str = "", tender: str = "", **kw: Any) -> dict:
    on = today(tenant)
    with _LOCK, session() as s:
        t = None
        if tender:
            t = s.get(ThekaTender, tender)
            if not t or t.tenant_id != tenant:
                t = s.scalars(select(ThekaTender).where(ThekaTender.tenant_id == tenant, ThekaTender.ref == tender)).first()
            if not t:
                raise ThekaError("Tender nahi mila.")
        new = t is None
        if new:
            if not str(name).strip():
                raise ThekaError("Tender ka naam / kaam chahiye.")
            t = ThekaTender(tenant_id=tenant, docs=json.dumps({d: False for d in R.DOCS}, ensure_ascii=False))
            s.add(t)
        if name:
            t.name = str(name)[:300]
        for k in ("dept", "portal", "ref", "notes"):
            if kw.get(k):
                setattr(t, k, str(kw[k])[:300])
        for k in ("value", "emd", "our_quote", "l1"):
            if kw.get(k) not in (None, ""):
                setattr(t, k, _num(kw[k]))
        if kw.get("emd_mode"):
            t.emd_mode = str(kw["emd_mode"]).lower()[:12]
        for k in ("prebid", "opening"):
            if kw.get(k):
                d = _date(kw[k])
                if d:
                    setattr(t, k, d.isoformat())
        if kw.get("last_date"):
            raw = str(kw["last_date"])
            d = _date(raw)
            if not d:
                raise ThekaError("Aakhri tareekh samajh nahi aayi.")
            tm = re.search(r"(\d{1,2}):(\d{2})", raw)
            t.last_date = d.isoformat() + (f" {int(tm.group(1)):02d}:{tm.group(2)}" if tm else "")
        if kw.get("status") in ("watching", "preparing", "submitted", "won", "lost", "dropped"):
            t.status = kw["status"]
        if kw.get("have"):
            docs = _j(t.docs, {})
            for h in kw["have"]:
                for d in docs:
                    if str(h).lower() in d.lower():
                        docs[d] = True
            t.docs = json.dumps(docs, ensure_ascii=False)
        if kw.get("emd_back") is not None:
            t.emd_back = bool(kw["emd_back"])
        s.flush()
        v = _tender_view(t, on)
    say = f"{'Naya tender' if new else 'Update'}: {v['name'][:60]}" + (f" ({v['dept']})" if v["dept"] else "") + (f", aakhri tareekh {v['last_date']}" if v["last_date"] else "") + "."
    if v["emd"]:
        say += f" EMD {_rs(v['emd'])}" + (" - Udyam ho to NIT mein chhoot / bid security declaration dekhiye" if not v["emd_mode"] else f" ({v['emd_mode']})") + "."
    if v["alert"]:
        say += " " + v["alert"]
    if v["missing"] and v["status"] in ("watching", "preparing"):
        say += f" {len(v['missing'])} kaagaz baaki."
    if v["status"] == "lost" and v["emd"] and not v["emd_back"]:
        say += " EMD wapas lene ki arzi daaliye."
    if v["status"] == "won":
        say += " Mubarak! 'project' bana kar BOQ jodiye; performance guarantee ki tareekh NIT se dekhiye."
    return {"tender": v, "say": say}


def tenders(tenant: str, status: str = "") -> list[dict]:
    on = today(tenant)
    with session() as s:
        q = select(ThekaTender).where(ThekaTender.tenant_id == tenant)
        if status:
            q = q.where(ThekaTender.status == status)
        out = [_tender_view(t, on) for t in s.scalars(q).all()]
    return sorted(out, key=lambda t: (t["days_left"] if t["days_left"] is not None else 9999))


# ---------------------------------------------------------------- projects + BOQ
def _proj(s, tenant: str, ref: str) -> ThekaProject:
    rows = s.scalars(select(ThekaProject).where(ThekaProject.tenant_id == tenant)).all()
    if not rows:
        raise ThekaError("Pehle project / kaam jodiye.")
    if not ref:
        run = [r for r in rows if r.status == "running"]
        if len(run) == 1:
            return run[0]
        raise ThekaError("Kaun sa kaam? " + ", ".join(r.name for r in rows[:8]))
    for r in rows:
        if r.id == ref or r.name.lower() == ref.lower() or (r.agreement_no and r.agreement_no == ref):
            return r
    hits = [r for r in rows if ref.lower() in r.name.lower()]
    if len(hits) == 1:
        return hits[0]
    raise ThekaError(f"Kaam '{ref}' nahi mila." if not hits else "Kai kaam: " + ", ".join(h.name for h in hits))


_DEFAULT_TERMS = {"gst_rate": R.GST_WORKS, "retention_pct": 5.0, "sd_pct": 0.0, "it_tds_pct": None, "gst_tds": True, "cess_pct": R.LABOUR_CESS,
                  "mob_pct": 10.0, "deviation_pct": 25.0}


def save_project(tenant: str, name: str = "", project: str = "", **kw: Any) -> dict:
    with _LOCK, session() as s:
        st = _st(s, tenant)
        p = _proj(s, tenant, project) if project else None
        new = p is None
        if new:
            if not str(name).strip():
                raise ThekaError("Kaam ka naam chahiye.")
            terms = dict(_DEFAULT_TERMS)
            terms["it_tds_pct"] = R.IT_TDS["individual"] if st.individual else R.IT_TDS["other"]
            p = ThekaProject(tenant_id=tenant, terms=json.dumps(terms))
            s.add(p)
        if name:
            p.name = str(name)[:200]
        for k in ("client", "agreement_no"):
            if kw.get(k):
                setattr(p, k, str(kw[k])[:160])
        if kw.get("govt") is not None:
            p.govt = bool(kw["govt"])
        if kw.get("value") not in (None, ""):
            p.value = _num(kw["value"])
        if kw.get("mob_advance") not in (None, ""):
            p.mob_advance = _num(kw["mob_advance"])
        for k in ("start", "end", "completed_on", "insurance_valid_to"):
            if kw.get(k):
                d = _date(kw[k])
                if not d:
                    raise ThekaError(f"Tareekh samajh nahi aayi: {kw[k]}")
                setattr(p, k, d.isoformat())
        if kw.get("dlp_months") not in (None, ""):
            p.dlp_months = int(_num(kw["dlp_months"]))
        t = _j(p.terms, {})
        for k in ("gst_rate", "retention_pct", "sd_pct", "it_tds_pct", "cess_pct", "mob_pct", "deviation_pct"):
            if kw.get(k) not in (None, ""):
                t[k] = _num(kw[k])
        if kw.get("gst_tds") is not None:
            t["gst_tds"] = bool(kw["gst_tds"])
        if not p.govt:
            t["gst_tds"] = False
            if kw.get("cess_pct") in (None, ""):
                t["cess_pct"] = 0.0
        p.terms = json.dumps(t)
        if kw.get("bg_amount") not in (None, "") or kw.get("bg_valid_to"):
            bg = _j(p.bg, {})
            if kw.get("bg_no"):
                bg["no"] = str(kw["bg_no"])[:40]
            if kw.get("bg_amount") not in (None, ""):
                bg["amount"] = _num(kw["bg_amount"])
            if kw.get("bg_valid_to"):
                bg["valid_to"] = (_date(kw["bg_valid_to"]) or today(tenant)).isoformat()
            p.bg = json.dumps(bg)
        if kw.get("status") in ("running", "completed", "closed"):
            p.status = kw["status"]
        s.flush()
        pid = p.id
    v = project_view(tenant, pid)
    say = f"{'Naya kaam' if new else 'Update'}: {v['name']}" + (f" ({v['client']})" if v["client"] else "") + (f", {_rs(v['value'])}" if v["value"] else "") \
        + (f", {v['start']} se {v['end']}" if v["start"] and v["end"] else "") + "."
    t = v["terms"]
    say += (f" Kataayi: retention {t.get('retention_pct', 0):g}%, TDS {t.get('it_tds_pct', 0):g}%" + (", GST TDS 2%" if t.get("gst_tds") else "")
            + (f", labour cess {t.get('cess_pct', 0):g}%" if t.get("cess_pct") else "") + " (agreement se alag ho to bataiye).")
    return {"project": v, "say": say}


def project_view(tenant: str, ref: str) -> dict:
    on = today(tenant)
    with session() as s:
        p = _proj(s, tenant, ref)
        items = s.scalars(select(ThekaItem).where(ThekaItem.project_id == p.id)).all()
        mb = s.scalars(select(ThekaMeasure).where(ThekaMeasure.project_id == p.id)).all()
        bills = s.scalars(select(ThekaBill).where(ThekaBill.project_id == p.id).order_by(ThekaBill.number)).all()
        cum: dict[str, float] = {}
        for m in mb:
            cum[m.item_id] = cum.get(m.item_id, 0.0) + m.qty
        boq_value = sum(i.qty * i.rate for i in items if not i.extra) or p.value
        done_value = sum(cum.get(i.id, 0.0) * i.rate for i in items)
        progress = round(done_value * 100 / boq_value, 1) if boq_value else 0.0
        billed = sum(_j(b.data, {}).get("this_taxable", 0.0) for b in bills)
        retention = sum(_j(b.data, {}).get("deductions", {}).get("retention", 0.0) for b in bills if b.status != "draft")
        v = {"id": p.id, "name": p.name, "client": p.client, "govt": p.govt, "agreement_no": p.agreement_no, "value": p.value, "start": p.start,
             "end": p.end, "completed_on": p.completed_on, "dlp_months": p.dlp_months, "terms": _j(p.terms, {}), "mob_advance": p.mob_advance,
             "bg": _j(p.bg, {}), "insurance_valid_to": p.insurance_valid_to, "status": p.status, "items": len(items), "progress_pct": progress,
             "done_value": round(done_value, 2), "billed": round(billed, 2), "retention_held": round(retention, 2),
             "bills": [{"id": b.id, "number": b.number, "final": b.final, "status": b.status, "net": _j(b.data, {}).get("net", 0.0), "passed": b.passed,
                        "paid": b.paid} for b in bills]}
    if v["start"] and v["end"]:
        v["delay"] = R.delay(date.fromisoformat(v["start"]), date.fromisoformat(v["end"]), on, progress)
    rr = R.retention_release(_date(v["completed_on"]), v["dlp_months"], add_months)
    v["retention_release"] = rr.isoformat() if rr else ""
    return v


def projects(tenant: str, status: str = "") -> list[dict]:
    with session() as s:
        q = select(ThekaProject.id).where(ThekaProject.tenant_id == tenant)
        if status:
            q = q.where(ThekaProject.status == status)
        ids = s.scalars(q).all()
    return [project_view(tenant, i) for i in ids]


def _boq_rows(name: str, data: bytes) -> list[dict]:
    from ..analyst.service import _rows
    rows = _rows(name, data)
    hi, cols = None, {}
    names = {"code": ("item no", "item", "s.no", "sr no", "sr. no.", "code", "item code", "sl no"),
             "description": ("description", "particulars", "item description", "description of item"),
             "unit": ("unit", "uom"), "qty": ("qty", "quantity", "qty."), "rate": ("rate", "rate (rs)", "unit rate", "rate in rs")}
    for i, r in enumerate(rows[:15]):
        h = [re.sub(r"\s+", " ", str(x or "").strip().lower()) for x in r]
        f = {}
        for j, x in enumerate(h):
            for k, ns in names.items():
                if x in ns and k not in f:
                    f[k] = j
        if {"description", "qty", "rate"} <= set(f):
            hi, cols = i, f
            break
    if hi is None:
        raise ThekaError("BOQ mein description, qty aur rate ke column chahiye.")
    out = []
    for r in rows[hi + 1:]:
        g = lambda k: r[cols[k]] if k in cols and cols[k] < len(r) else None  # noqa: E731
        desc = str(g("description") or "").strip()
        q, rt = _num(g("qty")), _num(g("rate"))
        if not desc or q <= 0 or rt <= 0:
            continue
        out.append({"code": str(g("code") or "").strip()[:20], "description": desc[:1000], "unit": str(g("unit") or "").strip()[:12], "qty": q, "rate": rt})
    return out


def boq(tenant: str, project: str, items: list | None = None, data: bytes | None = None, name: str = "boq.csv", file_id: str = "",
        extra: bool = False) -> dict:
    if file_id:
        from ..analyst.service import _load_file
        name, data = _load_file(tenant, file_id)
    rows = list(items or []) if items else _boq_rows(name, data or b"")
    if not rows:
        raise ThekaError("BOQ ki koi line nahi mili.")
    with _LOCK, session() as s:
        p = _proj(s, tenant, project)
        have = {i.code: i for i in s.scalars(select(ThekaItem).where(ThekaItem.project_id == p.id)).all()}
        measured = {m.item_id for m in s.scalars(select(ThekaMeasure).where(ThekaMeasure.project_id == p.id)).all()}
        nums = [int(c) for c in have if c.isdigit()]
        nxt = max(nums, default=0)
        n = 0
        for x in rows:
            code = str(x.get("code") or "").strip()[:20]
            if not code:                                   # no item number in her file: the next free one, never an existing item
                nxt += 1
                while str(nxt) in have:
                    nxt += 1
                code = str(nxt)
            q, rt = _num(x.get("qty")), _num(x.get("rate"))
            if rt <= 0:
                raise ThekaError(f"Item {code}: rate chahiye.")
            old_it = have.get(code)
            if old_it is not None and old_it.id in measured and x.get("description") and \
                    str(x["description"]).strip().lower()[:60] != old_it.description.strip().lower()[:60]:
                raise ThekaError(f"Item {code} par naap pehle se hai ({old_it.description[:40]}); naya kaam alag item number se jodiye.")
            it = old_it or ThekaItem(tenant_id=tenant, project_id=p.id, code=code)
            have[code] = it
            s.add(it)
            it.description = str(x.get("description") or it.description)[:1000]
            it.unit = str(x.get("unit") or it.unit)[:12]
            it.qty, it.rate = q if not extra else 0.0, rt
            it.extra = bool(extra or x.get("extra"))
            n += 1
        pname = p.name
        total = sum(i.qty * i.rate for i in s.scalars(select(ThekaItem).where(ThekaItem.project_id == p.id)).all())
    return {"items": n, "boq_value": round(total, 2), "say": f"{pname}: {n} BOQ item" + (" (extra item)" if extra else "") + f", kul {_rs(total)}."}


def boq_items(tenant: str, project: str) -> list[dict]:
    with session() as s:
        p = _proj(s, tenant, project)
        cum: dict[str, float] = {}
        for m in s.scalars(select(ThekaMeasure).where(ThekaMeasure.project_id == p.id)).all():
            cum[m.item_id] = cum.get(m.item_id, 0.0) + m.qty
        dev = _j(p.terms, {}).get("deviation_pct", 25.0)
        out = []
        for i in s.scalars(select(ThekaItem).where(ThekaItem.project_id == p.id)).all():
            c = cum.get(i.id, 0.0)
            out.append({"id": i.id, "code": i.code, "description": i.description, "unit": i.unit, "qty": i.qty, "rate": i.rate, "extra": i.extra,
                        "done": round(c, 3), "done_pct": round(c * 100 / i.qty, 1) if i.qty else None, "flag": R.deviation(c, i.qty, dev)})
    return sorted(out, key=lambda x: (x["extra"], _code_key(x["code"])))


def _code_key(c: str):
    return [int(p) if p.isdigit() else p for p in re.split(r"(\d+)", c)]


def _item(s, p: ThekaProject, ref: str) -> ThekaItem:
    rows = s.scalars(select(ThekaItem).where(ThekaItem.project_id == p.id)).all()
    ref = str(ref or "").strip()
    for i in rows:
        if i.id == ref or i.code == ref:
            return i
    hits = [i for i in rows if ref.lower() in i.description.lower()]
    if len(hits) == 1:
        return hits[0]
    if hits:
        raise ThekaError("Kai item: " + "; ".join(f"{h.code} {h.description[:40]}" for h in hits[:5]) + ". Item number bataiye.")
    raise ThekaError(f"BOQ mein '{ref}' nahi. Extra item ho to pehle department se rate pakka karke 'extra' jodiye.")


def measure(tenant: str, project: str, item: str, qty, detail: str = "", day: str = "", by: str = "") -> dict:
    """One measurement-book entry (as recorded on site; Hoola never rounds it up)."""
    q = _num(qty)
    if q == 0:
        raise ThekaError("Naap (qty) bataiye.")
    d = _date(day, today(tenant))
    with _LOCK, session() as s:
        p = _proj(s, tenant, project)
        it = _item(s, p, item)
        s.add(ThekaMeasure(tenant_id=tenant, project_id=p.id, item_id=it.id, day=d.isoformat(), qty=q, detail=str(detail)[:300], by=str(by)[:60]))
        s.flush()
        cum = sum(m.qty for m in s.scalars(select(ThekaMeasure).where(ThekaMeasure.item_id == it.id)).all())
        flag = R.deviation(cum, it.qty, _j(p.terms, {}).get("deviation_pct", 25.0))
        code, desc, unit, boqq = it.code, it.description, it.unit, it.qty
    say = f"MB: item {code} ({desc[:40]}) {q:g} {unit}. Ab tak {cum:g}" + (f" / BOQ {boqq:g}" if boqq else "") + "."
    if flag:
        say += " " + flag["text"]
    return {"cumulative": cum, "flag": flag, "say": say}


# ---------------------------------------------------------------- RA bills (her Haan to submit)
def ra_bill(tenant: str, project: str, upto: str = "", final: bool = False, conv: str = "") -> dict:
    on = today(tenant)
    u = _date(upto, on)
    with _LOCK, session() as s:
        p = _proj(s, tenant, project)
        t = _j(p.terms, {})
        bills = s.scalars(select(ThekaBill).where(ThekaBill.project_id == p.id).order_by(ThekaBill.number)).all()
        if any(b.status == "draft" for b in bills):
            old = next(b for b in bills if b.status == "draft")
            s.delete(old)
            s.flush()
            bills = [b for b in bills if b is not old]
        if any(b.final for b in bills):
            raise ThekaError("Final bill ban chuka hai.")
        items = {i.id: i for i in s.scalars(select(ThekaItem).where(ThekaItem.project_id == p.id)).all()}
        mbs = s.scalars(select(ThekaMeasure).where(ThekaMeasure.project_id == p.id, ThekaMeasure.day <= u.isoformat())).all()
        prev_qty: dict[str, float] = {}
        for b in bills:
            for ln in _j(b.data, {}).get("lines", []):
                prev_qty[ln["item_id"]] = ln["upto_qty"]
        upto_qty: dict[str, float] = {}
        for m in mbs:
            upto_qty[m.item_id] = upto_qty.get(m.item_id, 0.0) + m.qty
        lines, flags = [], []
        for iid, uq in upto_qty.items():
            it = items.get(iid)
            if not it:
                continue
            pq = prev_qty.get(iid, 0.0)
            this = round(uq - pq, 4)
            if abs(this) < 1e-9:
                continue
            lines.append({"item_id": iid, "code": it.code, "description": it.description[:200], "unit": it.unit, "boq_qty": it.qty, "rate": it.rate,
                          "prev_qty": pq, "this_qty": this, "upto_qty": round(uq, 4), "amount": round(this * it.rate, 2), "extra": it.extra})
            fl = R.deviation(uq, it.qty, t.get("deviation_pct", 25.0))
            if fl:
                flags.append(f"{it.code}: {fl['text']}")
        if not lines:
            raise ThekaError("Pichhle bill ke baad koi naya naap nahi.")
        lines.sort(key=lambda x: _code_key(x["code"]))
        taxable = round(sum(x["amount"] for x in lines), 2)
        gst = round(taxable * t.get("gst_rate", R.GST_WORKS) / 100, 2)
        mob_recovered = sum(_j(b.data, {}).get("deductions", {}).get("mobilisation_recovery", 0.0) for b in bills)
        mob_left = max(0.0, p.mob_advance - mob_recovered)
        terms = {**t, "value": p.value}
        if final:
            terms["mob_pct"] = 100.0                       # whatever is left is recovered in the final bill
        ded = R.deductions(taxable, gst, terms, mob_left)
        number = (bills[-1].number + 1) if bills else 1
        data = {"lines": lines, "this_taxable": taxable, "gst": gst, "gst_rate": t.get("gst_rate", R.GST_WORKS), "gross": round(taxable + gst, 2),
                "deductions": ded["items"], "deductions_total": ded["total"], "net": ded["net"], "flags": flags,
                "upto_value": round(sum(x["upto_qty"] * x["rate"] for x in lines) + sum(
                    prev_qty.get(i, 0.0) * items[i].rate for i in prev_qty if i not in {x["item_id"] for x in lines} and i in items), 2)}
        b = ThekaBill(tenant_id=tenant, project_id=p.id, number=number, final=bool(final), upto=u.isoformat(), data=json.dumps(data, ensure_ascii=False))
        s.add(b)
        s.flush()
        bid, pname = b.id, p.name
    f = bill_file(tenant, bid, conv)
    label = "Final bill" if final else f"RA bill {number}"
    say = (f"{pname} {label} (upto {u.strftime('%d %b')}): kaam {_rs(taxable)} + GST {_rs(gst)} = {_rs(taxable + gst)}; kataayi {_rs(ded['total'])} "
           f"(" + ", ".join(f"{k.replace('_', ' ')} {_rs(v)}" for k, v in ded["items"].items()) + f"); haath mein ~{_rs(ded['net'])}.")
    if flags:
        say += " Dhyaan: " + " ".join(flags[:3])
    say += " Department ko dene se pehle aapki haan."
    _publish(tenant, "theka_bill", {"bill_id": bid, "project": pname, "label": label, "net": ded["net"]}, conv)
    return {"bill_id": bid, "number": number, "bill": data, "file": f["file"], "say": say}


def bill_pdf(tenant: str, bill_id: str) -> bytes:
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfgen import canvas
    st = settings(tenant)
    with session() as s:
        b = s.get(ThekaBill, bill_id)
        if not b or b.tenant_id != tenant:
            raise ThekaError("Bill nahi mila.")
        p = s.get(ThekaProject, b.project_id)
        d = _j(b.data, {})
        mbs = s.scalars(select(ThekaMeasure).where(ThekaMeasure.project_id == p.id, ThekaMeasure.day <= b.upto)).all()
        items = {i.id: i.code for i in s.scalars(select(ThekaItem).where(ThekaItem.project_id == p.id)).all()}
        num, final, upto, pname, client, agr = b.number, b.final, b.upto, p.name, p.client, p.agreement_no
    buf = io.BytesIO()
    pdf = canvas.Canvas(buf, pagesize=A4)
    w, h = A4
    y = h - 40

    def t(x, text, size=8.5, bold=False, right=False):
        pdf.setFont("Helvetica-Bold" if bold else "Helvetica", size)
        s_ = str(text).replace("₹", "Rs ").encode("latin-1", "replace").decode("latin-1")
        (pdf.drawRightString if right else pdf.drawString)(x, y, s_[:120])

    def nl(k=12):
        nonlocal y
        y -= k
        if y < 60:
            pdf.showPage()
            y = h - 40

    t(30, f"{'FINAL' if final else 'RUNNING ACCOUNT'} BILL No. {num}   (upto {upto})", 12, True)
    nl(18)
    t(30, f"Name of work: {pname}", 9, True)
    nl()
    t(30, f"Client: {client or '-'}   Agreement: {agr or '-'}")
    nl()
    t(30, f"Contractor: {st['firm'] or '-'}   GSTIN: {st['gstin'] or '-'}")
    nl(16)
    cols = [(30, "Item"), (70, "Description"), (300, "Unit"), (335, "Prev qty"), (385, "This qty"), (435, "Upto qty"), (485, "Rate"), (565, "Amount")]
    for x, n in cols:
        t(x, n, 8, True, x == 565)
    nl(4)
    pdf.line(30, y, w - 30, y)
    nl(10)
    for ln in d["lines"]:
        t(30, ln["code"] + (" (EI)" if ln.get("extra") else ""))
        t(70, ln["description"][:45])
        t(300, ln["unit"])
        t(335, f"{ln['prev_qty']:g}")
        t(385, f"{ln['this_qty']:g}")
        t(435, f"{ln['upto_qty']:g}")
        t(485, f"{ln['rate']:,.2f}")
        t(565, f"{ln['amount']:,.2f}", 8.5, False, True)
        nl()
    pdf.line(30, y + 6, w - 30, y + 6)
    nl(4)
    rows = [("Value of work this bill", d["this_taxable"]), (f"GST @ {d['gst_rate']:g}%", d["gst"]), ("Gross", d["gross"])]
    rows += [(f"Less: {k.replace('_', ' ')}", -v) for k, v in d["deductions"].items()]
    rows += [("NET PAYABLE", d["net"])]
    for k, v in rows:
        t(380, k, 9, k == "NET PAYABLE")
        t(565, f"{v:,.2f}", 9, k == "NET PAYABLE", True)
        nl(13)
    nl(6)
    t(30, "Measurement abstract (as recorded in the measurement book):", 9, True)
    nl()
    for m in sorted(mbs, key=lambda m: (m.day, items.get(m.item_id, ""))):
        t(30, f"{m.day}  item {items.get(m.item_id, '')}  qty {m.qty:g}  {m.detail[:70]}  {('by ' + m.by) if m.by else ''}", 7.5)
        nl(10)
    nl(24)
    t(30, "Contractor's signature", 8.5)
    t(w - 30, "Checked by (Engineer-in-charge)", 8.5, False, True)
    pdf.showPage()
    pdf.save()
    return buf.getvalue()


def bill_file(tenant: str, bill_id: str, conv: str = "") -> dict:
    from .. import chatfiles
    with session() as s:
        b = s.get(ThekaBill, bill_id)
        num, final = b.number, b.final
    f = chatfiles.store(tenant, f"{'Final' if final else 'RA'}-Bill-{num}.pdf", bill_pdf(tenant, bill_id), "application/pdf")
    _publish(tenant, "pc_file", {**f, "source": "theka", "title": f"{'Final' if final else 'RA'} bill {num}"}, conv)
    return {"file": f}


def bill_status(tenant: str, bill_id: str, status: str, amount=None, on: str = "") -> dict:
    """submitted (her Haan) | passed (net passed by the department) | paid (money reached her bank)."""
    if status not in ("submitted", "passed", "paid"):
        raise ThekaError("status: submitted | passed | paid")
    d = _date(on, today(tenant)).isoformat()
    with _LOCK, session() as s:
        b = s.get(ThekaBill, bill_id)
        if not b or b.tenant_id != tenant:
            raise ThekaError("Bill nahi mila.")
        data = _j(b.data, {})
        if status == "submitted":
            b.status, b.submitted_on = "submitted", d
        elif status == "passed":
            b.status, b.passed = "passed", _num(amount) if amount not in (None, "") else data["net"]
        else:
            b.status, b.paid, b.paid_on = "paid", _num(amount) if amount not in (None, "") else (b.passed or data["net"]), d
            if not b.passed:
                b.passed = b.paid
        num, net, passed, paid = b.number, data["net"], b.passed, b.paid
        ded = data.get("deductions", {})
    if status == "submitted":
        return {"say": f"RA bill {num} department ko diya ({d}). Pass hone par bataiye."}
    if status == "passed":
        cut = round(net - passed, 2)
        return {"say": f"RA bill {num}: {_rs(passed)} pass" + (f"; {_rs(cut)} kam - department se wajah (cut / recovery) likhit mein lijiye." if cut > 1 else ".")}
    say = f"RA bill {num}: {_rs(paid)} mila."
    if ded.get("income_tax_tds"):
        say += f" TDS {_rs(ded['income_tax_tds'])} 26AS mein aaya ki nahi, dekhiye."
    if ded.get("gst_tds"):
        say += f" GST TDS {_rs(ded['gst_tds'])}: GST portal par GSTR-7A accept kariye, cash ledger mein aayega."
    return {"say": say}


def receivables(tenant: str) -> dict:
    on = today(tenant)
    out, ret = [], []
    with session() as s:
        ps = {p.id: p for p in s.scalars(select(ThekaProject).where(ThekaProject.tenant_id == tenant)).all()}
        for b in s.scalars(select(ThekaBill).where(ThekaBill.tenant_id == tenant, ThekaBill.status.in_(("submitted", "passed")))).all():
            d = _date(b.submitted_on, on)
            out.append({"project": ps[b.project_id].name if b.project_id in ps else "", "bill": b.number, "status": b.status,
                        "amount": b.passed or _j(b.data, {}).get("net", 0.0), "days": (on - d).days})
    for p in projects(tenant):
        if p["retention_held"]:
            ret.append({"project": p["name"], "retention": p["retention_held"], "release": p["retention_release"]})
    total = round(sum(x["amount"] for x in out), 2)
    return {"bills": sorted(out, key=lambda x: -x["days"]), "retention": ret, "total": total,
            "say": f"Department se {_rs(total)} aana baaki ({len(out)} bill)" + (f"; retention mein {_rs(sum(r['retention'] for r in ret))} ruka." if ret else ".")}


# ---------------------------------------------------------------- subcontractors (she pays)
def save_sub(tenant: str, name: str, phone: str = "", individual: bool | None = None, pan_ok: bool | None = None, work: str = "",
             retention_pct=None) -> dict:
    if not str(name).strip():
        raise ThekaError("Naam chahiye.")
    with session() as s:
        x = next((r for r in s.scalars(select(ThekaSub).where(ThekaSub.tenant_id == tenant)).all() if r.name.lower() == name.strip().lower()), None)
        new = x is None
        if new:
            x = ThekaSub(tenant_id=tenant, name=name.strip()[:160])
            s.add(x)
        if phone:
            x.phone = re.sub(r"[^\d+]", "", phone)[:20]
        if individual is not None:
            x.individual = bool(individual)
        if pan_ok is not None:
            x.pan_ok = bool(pan_ok)
        if work:
            x.work = work[:160]
        if retention_pct not in (None, ""):
            x.retention_pct = _num(retention_pct)
    return {"say": f"{'Naya' if new else 'Update'} petty contractor: {name}."}


def sub_bill(tenant: str, sub: str, amount, project: str = "", day: str = "", note: str = "") -> dict:
    amt = _num(amount)
    if amt <= 0:
        raise ThekaError("Bill ki rakam bataiye.")
    d = _date(day, today(tenant))
    with _LOCK, session() as s:
        rows = s.scalars(select(ThekaSub).where(ThekaSub.tenant_id == tenant)).all()
        x = next((r for r in rows if r.name.lower() == sub.lower()), None) or next((r for r in rows if sub.lower() in r.name.lower()), None)
        if not x:
            raise ThekaError(f"'{sub}' nahi mila; pehle jodiye.")
        pid = _proj(s, tenant, project).id if project else ""
        fy0 = _fy_start(d).isoformat()
        prev = s.scalars(select(ThekaSubBill).where(ThekaSubBill.sub_id == x.id, ThekaSubBill.day >= fy0)).all()
        tds = R.it_tds(amt, sum(b.amount for b in prev), x.individual, x.pan_ok, sum(b.tds_on for b in prev))
        ret = round(amt * x.retention_pct / 100, 2)
        net = round(amt - tds["tds"] - ret, 2)
        s.add(ThekaSubBill(tenant_id=tenant, sub_id=x.id, project_id=pid, day=d.isoformat(), amount=amt, tds=tds["tds"], tds_on=tds.get("on", 0.0),
                           retention=ret, net=net, note=str(note)[:300]))
        name = x.name
    say = f"{name} ka bill {_rs(amt)}: " + (f"TDS {_rs(tds['tds'])} ({tds['rate']:g}%, Sec 393(1)), " if tds["tds"] else "TDS nahi (seema se kam), ") \
        + (f"retention {_rs(ret)}, " if ret else "") + f"dena {_rs(net)}. Payment aap karengi; TDS agle mahine ki 7 tak jama."
    return {"tds": tds, "net": net, "say": say}


# ---------------------------------------------------------------- material, diary, EOT
def material(tenant: str, project: str, item: str, qty, rate=None, used: bool = False, day: str = "", note: str = "") -> dict:
    q = abs(_num(qty))
    if q <= 0:
        raise ThekaError("Maatra bataiye.")
    with session() as s:
        p = _proj(s, tenant, project)
        s.add(ThekaMaterial(tenant_id=tenant, project_id=p.id, item=str(item).strip().lower()[:80], qty=-q if used else q, rate=_num(rate),
                            day=(_date(day, today(tenant))).isoformat(), note=str(note)[:200]))
        s.flush()
        bal = sum(m.qty for m in s.scalars(select(ThekaMaterial).where(ThekaMaterial.project_id == p.id, ThekaMaterial.item == str(item).strip().lower()[:80])).all())
        pname = p.name
    return {"balance": bal, "say": f"{pname}: {item} {'kharch' if used else 'aaya'} {q:g}; site par ab {bal:g}." + (" Khatam hone wala!" if bal <= 0 else "")}


def stock(tenant: str, project: str) -> list[dict]:
    with session() as s:
        p = _proj(s, tenant, project)
        agg: dict[str, dict] = {}
        for m in s.scalars(select(ThekaMaterial).where(ThekaMaterial.project_id == p.id)).all():
            a = agg.setdefault(m.item, {"item": m.item, "in": 0.0, "used": 0.0, "value_in": 0.0})
            if m.qty > 0:
                a["in"] += m.qty
                a["value_in"] += m.qty * m.rate
            else:
                a["used"] += -m.qty
    return [{**a, "balance": round(a["in"] - a["used"], 3)} for a in agg.values()]


def diary(tenant: str, project: str, labour: int = 0, work: str = "", hindrance: str = "", lost_hours=None, by_client: bool = False,
          day: str = "") -> dict:
    d = _date(day, today(tenant))
    with _LOCK, session() as s:
        p = _proj(s, tenant, project)
        row = s.scalars(select(ThekaDiary).where(ThekaDiary.project_id == p.id, ThekaDiary.day == d.isoformat())).first() or \
            ThekaDiary(tenant_id=tenant, project_id=p.id, day=d.isoformat())
        s.add(row)
        row.labour = int(_num(labour)) or row.labour
        if work:
            row.work = (row.work + " | " if row.work else "") + str(work)[:1000]
        if hindrance:
            row.hindrance = str(hindrance)[:1000]
            row.lost_hours = _num(lost_hours) if lost_hours not in (None, "") else 8.0
            row.by_client = bool(by_client)
        pname = p.name
    say = f"{pname} {d.strftime('%d %b')}: {row.labour} mazdoor" + (f", kaam: {work[:80]}" if work else "") + "."
    if hindrance:
        say += f" Rukawat likhi ({'department ki taraf se - EOT mein ginegi' if by_client else 'apni taraf'})."
    return {"say": say}


def eot_letter(tenant: str, project: str, conv: str = "") -> dict:
    """Extension-of-time request from the hindrance register (department-side days). A draft: goes out only on her Haan."""
    with session() as s:
        p = _proj(s, tenant, project)
        rows = s.scalars(select(ThekaDiary).where(ThekaDiary.project_id == p.id, ThekaDiary.by_client.is_(True), ThekaDiary.hindrance != "")
                         .order_by(ThekaDiary.day)).all()
        st = _st(s, tenant)
        hind = [(r.day, r.hindrance, r.lost_hours) for r in rows]
        pname, client, agr, end = p.name, p.client, p.agreement_no, p.end
    if not hind:
        raise ThekaError("Hindrance register mein department ki taraf ki koi rukawat nahi likhi.")
    days = round(sum(h for _, _, h in hind) / 8, 1)
    lines = "\n".join(f"{i}. {d}: {t} ({h:g} ghante)" for i, (d, t, h) in enumerate(hind, 1))
    text = (f"To,\nThe Engineer-in-charge,\n{client or '__________'}\n\nSub: Request for extension of time - {pname}" + (f" (Agreement {agr})" if agr else "")
            + f"\n\nSir / Madam,\nThe work could not proceed on the following dates for reasons not attributable to us, as recorded in the hindrance register:\n"
            f"{lines}\n\nThe total hindrance comes to about {days:g} days. We request an extension of time of {int(days + 0.999)} days beyond the stipulated date "
            f"{end or '__________'} without levy of compensation.\n\nThanking you,\n{st.firm or '__________'}")
    _publish(tenant, "theka_letter", {"project": pname, "title": "EOT request", "days": days}, conv)
    return {"text": text, "days": days, "say": f"EOT arzi ka draft: {len(hind)} rukawat, ~{days:g} din. Aapki haan ke baad aap department ko dijiye."}


# ---------------------------------------------------------------- register, day
def register_file(tenant: str, conv: str = "") -> dict:
    import openpyxl
    from .. import chatfiles
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Projects"
    ws.append(["Work", "Client", "Value", "Start", "End", "Progress %", "Billed", "Retention held", "Retention release", "BG valid to", "Insurance to"])
    for p in projects(tenant):
        ws.append([p["name"], p["client"], p["value"], p["start"], p["end"], p["progress_pct"], p["billed"], p["retention_held"], p["retention_release"],
                   p["bg"].get("valid_to", ""), p["insurance_valid_to"]])
    ws2 = wb.create_sheet("Bills")
    ws2.append(["Work", "Bill", "Status", "Net claimed", "Passed", "Paid"])
    for p in projects(tenant):
        for b in p["bills"]:
            ws2.append([p["name"], ("Final " if b["final"] else "RA ") + str(b["number"]), b["status"], b["net"], b["passed"], b["paid"]])
    ws3 = wb.create_sheet("Tenders")
    ws3.append(["Tender", "Dept", "Ref", "Value", "EMD", "Last date", "Status", "Our quote", "L1", "EMD back"])
    for t in tenders(tenant):
        ws3.append([t["name"], t["dept"], t["ref"], t["value"], t["emd"], t["last_date"], t["status"], t["our_quote"], t["l1"], t["emd_back"]])
    bio = io.BytesIO()
    wb.save(bio)
    f = chatfiles.store(tenant, "Theka_register.xlsx", bio.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    _publish(tenant, "pc_file", {**f, "source": "theka", "title": "Contracts register"}, conv)
    return {"file": f, "say": "Kaam, bill aur tender ka register Excel mein."}


def day(tenant: str) -> dict:
    on = today(tenant)
    parts = []
    for t in tenders(tenant):
        if t["status"] in ("watching", "preparing") and t["days_left"] is not None and 0 <= t["days_left"] <= 3:
            parts.append(f"Tender '{t['name'][:40]}': {t['alert']}" + (f" ({len(t['missing'])} kaagaz baaki)" if t["missing"] else ""))
        if t["status"] == "lost" and t["emd"] and not t["emd_back"]:
            parts.append(f"'{t['name'][:30]}' ka EMD {_rs(t['emd'])} wapas lena")
    st = settings(tenant)
    rv = _date(st["registration_valid_to"])
    if rv and (rv - on).days <= 60:
        parts.append(f"Contractor registration {rv.isoformat()} ko khatam")
    for p in projects(tenant, "running"):
        bg = _date(p["bg"].get("valid_to"))
        if bg and (bg - on).days <= 30:
            parts.append(f"{p['name'][:30]}: BG {bg.isoformat()} ko khatam - badhwaiye")
        ins = _date(p["insurance_valid_to"])
        if ins and (ins - on).days <= 30:
            parts.append(f"{p['name'][:30]}: CAR insurance {ins.isoformat()} ko khatam")
        dl = p.get("delay")
        if dl and (dl["behind_pct"] > 15 or dl["late_days"]):
            parts.append(f"{p['name'][:30]}: {dl['say']}")
    for p in projects(tenant, "completed"):
        rr = _date(p["retention_release"])
        if rr and rr <= on + timedelta(days=15) and p["retention_held"]:
            parts.append(f"{p['name'][:30]}: retention {_rs(p['retention_held'])} {rr.isoformat()} ko milna chahiye - arzi daaliye")
    rc = receivables(tenant)
    old = [b for b in rc["bills"] if b["days"] > 45]
    if old:
        parts.append(f"{len(old)} bill 45 din se zyada baaki ({_rs(sum(b['amount'] for b in old))})")
    return {"parts": parts, "say": ("; ".join(parts) + ".") if parts else "Theka: sab theek."}


def due_mornings() -> list[str]:
    from ..life import tz_of
    out = []
    with session() as s:
        for r in s.scalars(select(ThekaSettings)).all():
            has = s.scalars(select(ThekaProject.id).where(ThekaProject.tenant_id == r.tenant_id).limit(1)).first() or \
                s.scalars(select(ThekaTender.id).where(ThekaTender.tenant_id == r.tenant_id).limit(1)).first()
            if not has:
                continue
            local = datetime.now(tz_of(r.tenant_id))
            try:
                hh, mm = [int(x) for x in (r.morning_time or "08:00").split(":")]
            except ValueError:
                hh, mm = 8, 0
            late = (local.hour * 60 + local.minute) - (hh * 60 + mm)
            d = local.date().isoformat()
            if r.last_morning != d and 0 <= late <= 180:
                r.last_morning = d
                out.append(r.tenant_id)
    return out


def send_morning(tenant: str) -> dict:
    from ..events import notify
    from ..net.service import home_conv
    d = day(tenant)
    if not d["parts"]:
        return {"text": ""}
    text = "Theka: " + d["say"]
    conv = home_conv(tenant)
    try:
        notify(conv, text)
    except Exception:
        pass
    _publish(tenant, "theka_day", {"text": text}, conv)
    return {"text": text}
