"""Import leads from CSV or Excel and normalise Indian phone numbers for WhatsApp."""
from __future__ import annotations

import csv
import io
import re
from dataclasses import dataclass, field

from .db import Lead, LeadList, session

NAME_KEYS = ["name", "naam", "customer name", "customer", "full name", "contact name", "client"]
PHONE_KEYS = ["whatsapp", "mobile", "phone", "number", "contact", "mob", "cell"]
CONSENT_KEYS = ["opt_in", "optin", "opt in", "opted_in", "consent", "whatsapp_consent", "subscribed"]
YES = {"yes", "y", "true", "1", "haan", "ha", "han", "opted in", "agree", "agreed", "subscribed"}

# Numbers used in the bundled sample file; never sent to in live mode.
SAMPLE_PREFIX = "91900000"


def normalize_phone(raw) -> str | None:
    """Return E.164 digits without '+' (e.g. 919876543210), or None if not usable."""
    if raw is None:
        return None
    if isinstance(raw, float):
        if raw != raw:  # NaN
            return None
        raw = str(int(raw))
    elif isinstance(raw, int):
        raw = str(raw)
    s = str(raw).strip()
    if not s:
        return None
    if re.fullmatch(r"\d(\.\d+)?[eE]\+\d+", s):   # Excel scientific notation lost the digits
        return None
    had_plus = s.startswith("+") or s.startswith("00")
    d = re.sub(r"\D", "", s)
    if s.startswith("00"):
        d = d[2:]
    if len(d) == 10 and d[0] in "6789":
        return "91" + d
    if len(d) == 11 and d[0] == "0" and d[1] in "6789":
        return "91" + d[1:]
    if len(d) == 12 and d.startswith("91") and d[2] in "6789":
        return d
    if len(d) == 13 and d.startswith("091") and d[3] in "6789":
        return d[1:]
    if had_plus and 8 <= len(d) <= 15 and not d.startswith("91"):
        return d   # foreign number written with a country code
    return None


def _pick(columns: list[str], keys: list[str]) -> str | None:
    low = {c: c.strip().lower() for c in columns}
    for k in keys:
        for c, l in low.items():
            if l == k:
                return c
    for k in keys:
        for c, l in low.items():
            if k in l:
                return c
    return None


def _rows_from_file(filename: str, data: bytes) -> tuple[list[str], list[dict]]:
    name = filename.lower()
    if name.endswith((".xlsx", ".xlsm")):
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        ws = wb.worksheets[0]
        it = ws.iter_rows(values_only=True)
        header = None
        for row in it:
            if row and any(v not in (None, "") for v in row):
                header = [str(v).strip() if v is not None else f"col{i+1}" for i, v in enumerate(row)]
                break
        if header is None:
            return [], []
        rows = []
        for row in it:
            if row is None or all(v in (None, "") for v in row):
                continue
            rows.append({header[i]: row[i] for i in range(min(len(header), len(row)))})
        return header, rows
    if name.endswith(".xls"):
        raise ValueError("Old .xls files are not supported. Save the sheet as .xlsx or .csv and upload again.")
    text = data.decode("utf-8-sig", errors="replace")
    try:
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    reader = csv.DictReader(io.StringIO(text), dialect=dialect)
    header = [h.strip() for h in (reader.fieldnames or [])]
    rows = [{(k or "").strip(): v for k, v in r.items()} for r in reader]
    return header, rows


@dataclass
class ImportResult:
    list_id: str
    name: str
    total_rows: int
    valid: int
    invalid: int
    duplicates: int
    opted_in: int
    has_consent_column: bool
    columns: list[str]
    name_column: str | None
    phone_column: str | None
    invalid_examples: list[str] = field(default_factory=list)

    def as_dict(self) -> dict:
        return self.__dict__.copy()


def import_leads(filename: str, data: bytes, tenant_id: str, list_name: str | None = None) -> ImportResult:
    header, rows = _rows_from_file(filename, data)
    if not header:
        raise ValueError("The file looks empty. Make sure the first row has column names like Name and Mobile.")
    phone_col = _pick(header, PHONE_KEYS)
    if phone_col is None:
        raise ValueError(f"No phone column found. Columns seen: {', '.join(header)}. Name one of them 'Mobile' or 'Phone'.")
    name_col = _pick(header, NAME_KEYS)
    consent_col = _pick(header, CONSENT_KEYS)

    seen: set[str] = set()
    leads: list[Lead] = []
    invalid, dups, opted, bad = 0, 0, 0, []
    for r in rows:
        phone = normalize_phone(r.get(phone_col))
        if not phone:
            invalid += 1
            if len(bad) < 5 and r.get(phone_col) not in (None, ""):
                bad.append(str(r.get(phone_col)))
            continue
        if phone in seen:
            dups += 1
            continue
        seen.add(phone)
        consent = None
        if consent_col:
            consent = str(r.get(consent_col) or "").strip().lower() in YES
            opted += 1 if consent else 0
        nm = r.get(name_col) if name_col else None
        nm = str(nm).strip() if nm not in (None, "") else None
        extra = {k: (str(v) if v is not None else None) for k, v in r.items()
                 if k not in (phone_col, name_col, consent_col)}
        leads.append(Lead(name=nm, phone=phone, opted_in=consent, extra=extra))

    with session() as s:
        ll = LeadList(tenant_id=tenant_id, name=list_name or filename.rsplit(".", 1)[0], source_filename=filename,
                      columns=header, has_consent_column=consent_col is not None, total_rows=len(rows),
                      valid=len(leads), invalid=invalid, duplicates=dups, opted_in=opted)
        s.add(ll)
        s.flush()
        for l in leads:
            l.list_id = ll.id
        s.add_all(leads)
        return ImportResult(list_id=ll.id, name=ll.name, total_rows=len(rows), valid=len(leads), invalid=invalid,
                            duplicates=dups, opted_in=opted, has_consent_column=consent_col is not None,
                            columns=header, name_column=name_col, phone_column=phone_col, invalid_examples=bad)
